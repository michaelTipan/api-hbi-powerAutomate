"""Continua mini E2E desde revision ya generada (mismo process_date)."""
from __future__ import annotations

import base64
import io
import json
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
WORK = Path(r"D:\CMC\HBI_Capital\_work\sandbox_mini")
PROCESS_DATE = "2026-07-30"
BANK_CODE = "banco_bogota"
CLIENTS = (
    "INFORMACION CREDITOS-CLIENTES/"
    "03 COMWARE PRUEBAS- INFORMACION CREDITOS CLIENTES"
)
BANK_DIR = f"{CLIENTS}/01 CARGA TRANSACCIONES BANCO"
REV_DIR = f"{CLIENTS}/02 VALIDACION PAGOS/01 REVISION"
BANK_NAME = "BANCO_BOGOTA.xlsx"
REVIEW_NAME = (
    "validacion_pagos_banco_bogota_2026-07-30_"
    "f06d836c-16f6-4999-82cd-51582ab34576.xlsx"
)

CLEAN_ROWS = [
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 10_147_965.0,
        "concepto": "A&M CONSTRUCOL",
        "tipo": "PAGO",
        "trx": "MINI PAGO ATRASADO 215",
    },
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 11_178_814.0,
        "concepto": "CONSULTORA INGENIERIA & PROYECTOS MARTINEZ RONCANCIO SAS",
        "tipo": "PAGO",
        "trx": "MINI PAGO ADELANTADO 333",
    },
]


def api_key() -> str:
    env = Path(r"D:\CMC\HBI_Capital\api-hbi-powerAutomate.env").read_text(encoding="utf-8")
    m = re.search(r"^API_HTTP_KEY=(.+)$", env, re.M)
    if not m:
        raise RuntimeError("API_HTTP_KEY missing")
    return m.group(1).strip()


def log(name: str, data: object) -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    path = WORK / f"{name}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(name, json.dumps(data, ensure_ascii=False, default=str)[:550])


def poll(c: httpx.Client, url: str, timeout_s: int = 1800) -> dict:
    t0 = time.time()
    while time.time() - t0 < timeout_s:
        r = c.get(url, timeout=90)
        if r.status_code != 200:
            print("poll_http", r.status_code, r.text[:200])
            time.sleep(8)
            continue
        body = r.json()
        st = body.get("status")
        err = body.get("error")
        err_txt = err if isinstance(err, str) else json.dumps(err, ensure_ascii=False) if err else ""
        print(time.strftime("%H:%M:%S"), "status", st, err_txt[:120])
        if st in ("completed", "failed", "error"):
            return body
        time.sleep(8)
    raise TimeoutError(url)


def children(c: httpx.Client, drive: str, fid: str = "root") -> list[dict]:
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c: httpx.Client, drive: str, parts: list[str]) -> str:
    cur = "root"
    for p in parts:
        hit = next((i for i in children(c, drive, cur) if i.get("name") == p), None)
        if not hit:
            raise FileNotFoundError(p)
        cur = hit["id"]
    return cur


def download(c: httpx.Client, drive: str, item_id: str) -> bytes:
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        timeout=180,
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def upload_bytes(c: httpx.Client, drive: str, item_id: str, raw: bytes) -> None:
    r = c.put(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        json={"content_base64": base64.b64encode(raw).decode("ascii")},
        timeout=180,
    )
    r.raise_for_status()


def queue_and_poll(c: httpx.Client, path: str, payload: dict, tag: str) -> dict:
    r = c.post(f"{BASE}{path}", json=payload, timeout=60)
    log(f"{tag}_queue", {"status": r.status_code, "body": r.json() if r.content else None})
    if r.status_code not in (200, 202):
        return {"http_status": r.status_code, "body": r.json() if r.content else r.text}
    body = r.json()
    job_id = body.get("job_id")
    if not job_id:
        return body
    done = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{job_id}")
    log(f"{tag}_job", done)
    return done


def summarize_review(raw: bytes) -> dict:
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    out: dict = {"sheets": wb.sheetnames}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        # Detect header row: first non-empty row with several strings
        header_idx = 0
        for i, row in enumerate(rows[:8]):
            vals = [v for v in row if v is not None and str(v).strip()]
            if len(vals) >= 3 and sum(isinstance(v, str) for v in vals) >= 2:
                header_idx = i
                break
        headers = [str(h).strip() if h is not None else "" for h in rows[header_idx]]
        data = []
        for row in rows[header_idx + 1 :]:
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            item = {
                headers[i]: row[i]
                for i in range(min(len(headers), len(row)))
                if headers[i]
            }
            data.append(item)
        out[name] = {"header_row": header_idx + 1, "headers": headers, "rows": data, "count": len(data)}
    return out


def compact_rows(rows: list[dict], keys: list[str]) -> list[dict]:
    out = []
    for row in rows:
        item = {}
        for k in keys:
            for rk, rv in row.items():
                if rk.lower().replace(" ", "") == k.lower().replace(" ", "") or k.lower() in rk.lower():
                    item[rk] = rv
        if not item:
            item = {k: row.get(k) for k in list(row.keys())[:10]}
        out.append(item or row)
    return out


def main() -> None:
    results: dict = {"cases": {}}
    key = api_key()
    with httpx.Client(timeout=180, headers={"X-API-Key": key}) as c:
        drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]
        rev_id = walk(c, drive, REV_DIR.split("/"))
        rev_files = [i for i in children(c, drive, rev_id) if str(i.get("name", "")).endswith(".xlsx")]
        log("06_rev_files", [{"name": i.get("name"), "id": i.get("id")} for i in rev_files])
        target = next((i for i in rev_files if i.get("name") == REVIEW_NAME), None)
        if not target and rev_files:
            target = sorted(rev_files, key=lambda i: i.get("lastModifiedDateTime") or "", reverse=True)[0]
        if not target:
            raise RuntimeError("no review file")

        rev_raw = download(c, drive, target["id"])
        (WORK / target["name"]).write_bytes(rev_raw)
        summary = summarize_review(rev_raw)
        counts = {k: v.get("count") for k, v in summary.items() if isinstance(v, dict) and "count" in v}
        log("07_review_summary", {"file": target["name"], "sheets": summary.get("sheets"), "counts": counts})

        pagos = summary.get("Distribucion_Pagos") or {}
        abonos = summary.get("Distribucion_Abonos") or {}
        errores = summary.get("Errores") or {}
        # Keys flexibles
        for k, v in summary.items():
            if not isinstance(v, dict):
                continue
            lk = k.lower()
            if "error" in lk:
                errores = v
            elif "abono" in lk:
                abonos = v
            elif "pago" in lk and "distrib" in lk:
                pagos = v

        focus_pago = [
            "Cliente", "Concepto", "Monto", "Tipo", "Estado", "Validar", "Crédito", "Credito",
            "Observacion", "Observación", "Fecha", "ID",
        ]
        results["cases"]["review_content"] = {
            "file": target["name"],
            "counts": counts,
            "pagos": compact_rows(pagos.get("rows") or [], focus_pago),
            "abonos": compact_rows(abonos.get("rows") or [], focus_pago),
            "errores": compact_rows(errores.get("rows") or [], focus_pago + ["Error", "Detalle", "Codigo", "Código"]),
        }
        log("07b_review_focus", results["cases"]["review_content"])

        # Finalize con errores abiertos
        fin_err = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/finalize/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "08_finalize_with_errors",
        )
        results["cases"]["finalize_with_errors"] = {
            "status": fin_err.get("status"),
            "error": fin_err.get("error"),
            "result": fin_err.get("result"),
            "user_message": fin_err.get("user_message"),
            "severity": fin_err.get("severity"),
        }

        # Borrar excel y reintentar Generate (debe quedar bloqueado / recreate?)
        for it in list(rev_files):
            if str(it.get("name", "")).startswith("~$"):
                continue
            del_r = c.delete(
                f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/items/{it['id']}"
            )
            print("delete_rev", it.get("name"), del_r.status_code)

        gen_after_delete = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "09_generate_after_delete_excel",
        )
        results["cases"]["generate_after_delete_excel"] = {
            "status": gen_after_delete.get("status"),
            "error": gen_after_delete.get("error"),
            "result": gen_after_delete.get("result"),
            "already_generated": (gen_after_delete.get("result") or {}).get("already_generated")
            if isinstance(gen_after_delete.get("result"), dict)
            else None,
            "file_action": (gen_after_delete.get("result") or {}).get("file_action")
            if isinstance(gen_after_delete.get("result"), dict)
            else None,
        }

        # Cancel recovery
        cancel1 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/cancel-active-process/queue",
            {"bank_code": BANK_CODE},
            "10_cancel_recovery",
        )
        results["cases"]["cancel_recovery"] = {
            "status": cancel1.get("status"),
            "result": cancel1.get("result") or cancel1.get("error"),
        }

        # Banco limpio + regenerate
        bank_folder = walk(c, drive, BANK_DIR.split("/"))
        bank_item = next(i for i in children(c, drive, bank_folder) if i.get("name") == BANK_NAME)
        raw2 = download(c, drive, bank_item["id"])
        wb2 = load_workbook(io.BytesIO(raw2))
        ws2 = wb2.active
        for r in range(4, ws2.max_row + 1):
            for col in range(1, 6):
                ws2.cell(r, col).value = None
        for i, row in enumerate(CLEAN_ROWS):
            rr = 4 + i
            ws2.cell(rr, 1).value = row["fecha"]
            ws2.cell(rr, 2).value = float(row["monto"])
            ws2.cell(rr, 3).value = row["concepto"]
            ws2.cell(rr, 4).value = row["tipo"]
            ws2.cell(rr, 5).value = row["trx"]
        buf2 = io.BytesIO()
        wb2.save(buf2)
        upload_bytes(c, drive, bank_item["id"], buf2.getvalue())

        gen3 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "11_generate_after_cancel",
        )
        results["cases"]["generate_after_cancel"] = {
            "status": gen3.get("status"),
            "error": gen3.get("error"),
            "result": gen3.get("result"),
        }

        # Idempotencia limpia
        gen4 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "12_generate_idempotent_clean",
        )
        results["cases"]["generate_idempotent_clean"] = {
            "status": gen4.get("status"),
            "result": gen4.get("result"),
        }

        # Analizar revision limpia
        rev_files2 = [i for i in children(c, drive, rev_id) if str(i.get("name", "")).endswith(".xlsx")]
        if rev_files2:
            newest = sorted(rev_files2, key=lambda i: i.get("lastModifiedDateTime") or "", reverse=True)[0]
            clean_sum = summarize_review(download(c, drive, newest["id"]))
            clean_pagos = clean_sum.get("Distribucion_Pagos") or {}
            results["cases"]["clean_review"] = {
                "file": newest["name"],
                "counts": {
                    k: v.get("count")
                    for k, v in clean_sum.items()
                    if isinstance(v, dict) and "count" in v
                },
                "pagos": compact_rows(clean_pagos.get("rows") or [], focus_pago),
            }
            log("11b_clean_review", results["cases"]["clean_review"])

        # Cancel final
        cancel2 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/cancel-active-process/queue",
            {"bank_code": BANK_CODE},
            "13_cancel_final",
        )
        results["cases"]["cancel_final"] = {
            "status": cancel2.get("status"),
            "result": cancel2.get("result") or cancel2.get("error"),
        }

    # Merge con resultados previos si existen
    prev_path = WORK / "99_results.json"
    merged = {"phase1": {}, "phase2": results}
    if prev_path.exists():
        try:
            merged["phase1"] = json.loads(prev_path.read_text(encoding="utf-8"))
        except Exception:
            pass
    # Incluir generate/idempotent de phase1 desde archivos
    for name in ("04_generate_job", "05_generate_idempotent_job", "02_cancel_pre_job"):
        p = WORK / f"{name}.json"
        if p.exists():
            merged["phase1"][name] = json.loads(p.read_text(encoding="utf-8"))

    log("99_results", merged)
    print("DONE")


if __name__ == "__main__":
    main()
