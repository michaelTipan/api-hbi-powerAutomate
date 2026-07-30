"""Mini E2E sandbox: un ejemplo por escenario, lote minimo, sin Equinorte."""
from __future__ import annotations

import base64
import io
import json
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
WORK = Path(r"D:\CMC\HBI_Capital\_work\sandbox_mini")
PROCESS_DATE = "2026-07-30"
BANK_CODE = "banco_bogota"
CLIENTS = (
    "INFORMACION CREDITOS-CLIENTES/"
    "03 COMWARE PRUEBAS- INFORMACION CREDITOS CLIENTES"
)
BANK_DIR = f"{CLIENTS}/01 CARGA TRANSACCIONES BANCO"
REV_DIR = f"{CLIENTS}/02 VALIDACION PAGOS/01 REVISION"
BANK_NAME = "BANCO_BOGOTA.xlsx"

# Lote minimo: 1 ejemplo por tipo/estado/error (sin Equinorte).
ROWS = [
    # ATRASADO + PAGO (fecha banco > fecha limite 2026-03-15)
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 10_147_965.0,
        "concepto": "A&M CONSTRUCOL",
        "tipo": "PAGO",
        "trx": "MINI PAGO ATRASADO 215",
        "tag": "pago_atrasado",
    },
    # ADELANTADO + PAGO (fecha banco < fecha limite 2026-07-30)
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 11_178_814.0,
        "concepto": "CONSULTORA INGENIERIA & PROYECTOS MARTINEZ RONCANCIO SAS",
        "tipo": "PAGO",
        "trx": "MINI PAGO ADELANTADO 333",
        "tag": "pago_adelantado",
    },
    # NORMAL + PAGO (fecha = limite 2026-05-23)
    {
        "fecha": datetime(2026, 5, 23),
        "monto": 19_102_163.0,
        "concepto": "GEOEXCON",
        "tipo": "PAGO",
        "trx": "MINI PAGO NORMAL 231",
        "tag": "pago_normal",
    },
    # PAGO Y ABONO CAPITAL (monto > cuota extracto 231)
    {
        "fecha": datetime(2026, 5, 23),
        "monto": 24_000_000.0,
        "concepto": "GEOEXCON",
        "tipo": "PAGO Y ABONO CAPITAL",
        "trx": "MINI PAGO+CAPITAL 231",
        "tag": "pago_y_abono_capital",
    },
    # ABONO CAPITAL
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 500_000.0,
        "concepto": "TRANSMAQUIOBRAS",
        "tipo": "ABONO CAPITAL",
        "trx": "MINI ABONO CAPITAL 320",
        "tag": "abono_capital",
    },
    # ABONO MORA
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 200_000.0,
        "concepto": "SERVICIOS DE CONSULTORIA Y OBRA CIVIL SAS",
        "tipo": "ABONO MORA",
        "trx": "MINI ABONO MORA 232",
        "tag": "abono_mora",
    },
    # Extracto danado / fecha_limite no legible (#46REPUESTOS)
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 21_967_682.0,
        "concepto": "AGRECAR",
        "tipo": "PAGO",
        "trx": "MINI EXTRACTO DANADO 46REPUESTOS",
        "tag": "extracto_danado",
    },
    # Cliente inexistente -> Errores
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 1_000_000.0,
        "concepto": "CLIENTE_FANTASMA_MINI",
        "tipo": "PAGO",
        "trx": "MINI CLIENTE INEXISTENTE",
        "tag": "cliente_fantasma",
    },
]


def api_key() -> str:
    env = Path(r"D:\CMC\HBI_Capital\api-hbi-powerAutomate.env").read_text(encoding="utf-8")
    m = re.search(r"^API_HTTP_KEY=(.+)$", env, re.M)
    if not m:
        raise RuntimeError("API_HTTP_KEY missing")
    return m.group(1).strip()


def log(name: str, data: object) -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    path = WORK / f"{name}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(name, json.dumps(data, ensure_ascii=False, default=str)[:500])


def poll(c: httpx.Client, url: str, timeout_s: int = 1800) -> dict:
    t0 = time.time()
    while time.time() - t0 < timeout_s:
        r = c.get(url, timeout=90)
        if r.status_code != 200:
            print("poll_http", r.status_code, r.text[:200])
            time.sleep(8)
            continue
        body = r.json()
        st = body.get("status")
        err = body.get("error")
        err_txt = err if isinstance(err, str) else json.dumps(err, ensure_ascii=False) if err else ""
        print(time.strftime("%H:%M:%S"), "status", st, err_txt[:120])
        if st in ("completed", "failed", "error"):
            return body
        time.sleep(8)
    raise TimeoutError(url)


def children(c: httpx.Client, drive: str, fid: str = "root") -> list[dict]:
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c: httpx.Client, drive: str, parts: list[str]) -> str:
    cur = "root"
    for p in parts:
        hit = next((i for i in children(c, drive, cur) if i.get("name") == p), None)
        if not hit:
            raise FileNotFoundError(p)
        cur = hit["id"]
    return cur


def download(c: httpx.Client, drive: str, item_id: str) -> bytes:
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        timeout=180,
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def upload_bytes(c: httpx.Client, drive: str, item_id: str, raw: bytes) -> None:
    r = c.put(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        json={"content_base64": base64.b64encode(raw).decode("ascii")},
        timeout=180,
    )
    r.raise_for_status()


def queue_and_poll(c: httpx.Client, path: str, payload: dict, tag: str) -> dict:
    r = c.post(f"{BASE}{path}", json=payload, timeout=60)
    log(f"{tag}_queue", {"status": r.status_code, "body": r.json() if r.content else None})
    if r.status_code not in (200, 202):
        return {"http_status": r.status_code, "body": r.json() if r.content else r.text}
    body = r.json()
    job_id = body.get("job_id")
    if not job_id:
        return body
    done = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{job_id}")
    log(f"{tag}_job", done)
    return done


def summarize_review(raw: bytes) -> dict:
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    out: dict = {"sheets": wb.sheetnames}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        data = []
        for row in rows[1:]:
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            item = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
            data.append(item)
        out[name] = {"headers": headers, "rows": data, "count": len(data)}
    return out


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    results: dict = {"process_date": PROCESS_DATE, "rows": ROWS, "cases": {}}
    key = api_key()
    with httpx.Client(timeout=180, headers={"X-API-Key": key}) as c:
        health = c.get(f"{BASE}/health").json()
        log("00_health", health)
        if health.get("build") != "extract-damaged-failclosed-20260730":
            raise RuntimeError(f"build inesperado: {health}")

        drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]
        bank_folder = walk(c, drive, BANK_DIR.split("/"))
        bank_item = next(i for i in children(c, drive, bank_folder) if i.get("name") == BANK_NAME)
        bank_id = bank_item["id"]
        log("01_bank_item", {"id": bank_id, "size": bank_item.get("size")})

        # 0) Cancel previo (limpieza)
        cancel0 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/cancel-active-process/queue",
            {"bank_code": BANK_CODE},
            "02_cancel_pre",
        )
        results["cases"]["cancel_pre"] = {
            "status": cancel0.get("status"),
            "result": cancel0.get("result") or cancel0.get("error"),
        }

        # 1) Subir banco minimo
        raw = download(c, drive, bank_id)
        wb = load_workbook(io.BytesIO(raw))
        ws = wb.active
        for r in range(4, ws.max_row + 1):
            for col in range(1, 6):
                ws.cell(r, col).value = None
        for i, row in enumerate(ROWS):
            rr = 4 + i
            ws.cell(rr, 1).value = row["fecha"]
            ws.cell(rr, 2).value = float(row["monto"])
            ws.cell(rr, 3).value = row["concepto"]
            ws.cell(rr, 4).value = row["tipo"]
            ws.cell(rr, 5).value = row["trx"]
        buf = io.BytesIO()
        wb.save(buf)
        payload = buf.getvalue()
        (WORK / "BANCO_BOGOTA_mini.xlsx").write_bytes(payload)
        upload_bytes(c, drive, bank_id, payload)
        log("03_bank_uploaded", {"rows": len(ROWS)})

        # 2) Generate
        gen = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "04_generate",
        )
        results["cases"]["generate"] = {
            "status": gen.get("status"),
            "result_keys": list((gen.get("result") or {}).keys()) if isinstance(gen.get("result"), dict) else None,
            "error": gen.get("error"),
            "result": gen.get("result"),
        }

        # 3) Idempotencia: segundo Generate
        gen2 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "05_generate_idempotent",
        )
        results["cases"]["generate_idempotent"] = {
            "status": gen2.get("status"),
            "error": gen2.get("error"),
            "result": gen2.get("result"),
            "enrichment": gen2.get("enrichment") or gen2.get("user_guidance"),
        }

        # 4) Descargar revision y resumir hojas
        rev_id = walk(c, drive, REV_DIR.split("/"))
        rev_files = [i for i in children(c, drive, rev_id) if str(i.get("name", "")).endswith(".xlsx")]
        log("06_rev_files", [{"name": i.get("name"), "id": i.get("id")} for i in rev_files])
        review_summary = None
        if rev_files:
            newest = sorted(rev_files, key=lambda i: i.get("lastModifiedDateTime") or "", reverse=True)[0]
            rev_raw = download(c, drive, newest["id"])
            (WORK / newest["name"]).write_bytes(rev_raw)
            review_summary = summarize_review(rev_raw)
            log("07_review_summary", {
                "file": newest["name"],
                "sheets": review_summary.get("sheets"),
                "counts": {k: v.get("count") for k, v in review_summary.items() if isinstance(v, dict) and "count" in v},
            })
            # Extraer estados/tipos relevantes
            dist = review_summary.get("Distribucion") or review_summary.get("Distribución") or {}
            dist_rows = dist.get("rows") or []
            abonos = None
            for k, v in review_summary.items():
                if isinstance(v, dict) and "abono" in k.lower():
                    abonos = v
                    break
            errores = None
            for k, v in review_summary.items():
                if isinstance(v, dict) and "error" in k.lower():
                    errores = v
                    break
            results["cases"]["review_content"] = {
                "file": newest["name"],
                "distribucion_sample": dist_rows[:20],
                "abonos_sample": (abonos or {}).get("rows", [])[:20],
                "errores_sample": (errores or {}).get("rows", [])[:20],
                "errores_count": (errores or {}).get("count", 0),
            }

        # 5) Finalize con Errores abiertos (debe fallar)
        fin_err = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/finalize/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "08_finalize_with_errors",
        )
        results["cases"]["finalize_with_errors"] = {
            "status": fin_err.get("status"),
            "error": fin_err.get("error"),
            "result": fin_err.get("result"),
        }

        # 6) Borrar Excel revision y reintentar Generate (debe bloquear sin cancel)
        if rev_files:
            for it in rev_files:
                if str(it.get("name", "")).startswith("~$"):
                    continue
                del_r = c.delete(
                    f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/items/{it['id']}"
                )
                print("delete_rev", it.get("name"), del_r.status_code)
        gen_after_delete = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "09_generate_after_delete_excel",
        )
        results["cases"]["generate_after_delete_excel"] = {
            "status": gen_after_delete.get("status"),
            "error": gen_after_delete.get("error"),
            "result": gen_after_delete.get("result"),
        }

        # 7) Cancel + regenerar (recuperacion)
        cancel1 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/cancel-active-process/queue",
            {"bank_code": BANK_CODE},
            "10_cancel_recovery",
        )
        results["cases"]["cancel_recovery"] = {
            "status": cancel1.get("status"),
            "result": cancel1.get("result") or cancel1.get("error"),
        }

        # Banco limpio: solo 2 filas validas (sin errores) para cerrar ciclo Generate OK
        clean_rows = [ROWS[0], ROWS[1]]  # atrasado + adelantado
        raw2 = download(c, drive, bank_id)
        wb2 = load_workbook(io.BytesIO(raw2))
        ws2 = wb2.active
        for r in range(4, ws2.max_row + 1):
            for col in range(1, 6):
                ws2.cell(r, col).value = None
        for i, row in enumerate(clean_rows):
            rr = 4 + i
            ws2.cell(rr, 1).value = row["fecha"]
            ws2.cell(rr, 2).value = float(row["monto"])
            ws2.cell(rr, 3).value = row["concepto"]
            ws2.cell(rr, 4).value = row["tipo"]
            ws2.cell(rr, 5).value = row["trx"]
        buf2 = io.BytesIO()
        wb2.save(buf2)
        upload_bytes(c, drive, bank_id, buf2.getvalue())

        gen3 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "11_generate_after_cancel",
        )
        results["cases"]["generate_after_cancel"] = {
            "status": gen3.get("status"),
            "error": gen3.get("error"),
            "result": gen3.get("result"),
        }

        # Reintento Generate (idempotencia limpia)
        gen4 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "12_generate_idempotent_clean",
        )
        results["cases"]["generate_idempotent_clean"] = {
            "status": gen4.get("status"),
            "error": gen4.get("error"),
            "result": gen4.get("result"),
        }

        # Cancel final para dejar sandbox limpio
        cancel2 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/cancel-active-process/queue",
            {"bank_code": BANK_CODE},
            "13_cancel_final",
        )
        results["cases"]["cancel_final"] = {
            "status": cancel2.get("status"),
            "result": cancel2.get("result") or cancel2.get("error"),
        }

    log("99_results", results)
    print("DONE")


if __name__ == "__main__":
    main()
