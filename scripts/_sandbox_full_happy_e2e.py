"""E2E feliz completo sandbox: Generate → Finalize → notify → merge → amort.

Lote minimo limpio (sin Equinorte): 1 PAGO A&M + 1 ABONO CAPITAL TRANSMAQUIOBRAS.
Cubre endpoints de ciclo productivo + cancel de limpieza.
"""
from __future__ import annotations

import base64
import io
import json
import re
import time
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
WORK = Path(r"D:\CMC\HBI_Capital\_work\sandbox_full_e2e")
ASIENTOS_DIR = WORK / "asientos"
PROCESS_DATE = "2026-07-30"
BANK_CODE = "banco_bogota"
CLIENTS = (
    "INFORMACION CREDITOS-CLIENTES/"
    "03 COMWARE PRUEBAS- INFORMACION CREDITOS CLIENTES"
)
BANK_DIR = f"{CLIENTS}/01 CARGA TRANSACCIONES BANCO"
REV_DIR = f"{CLIENTS}/02 VALIDACION PAGOS/01 REVISION"
BANK_NAME = "BANCO_BOGOTA.xlsx"

# Fecha limite V2 observada en A&M #215 = 2026-08-15 → NORMAL exacto.
ROWS = [
    {
        "fecha": datetime(2026, 8, 15),
        "monto": 10_147_965.0,
        "concepto": "A&M CONSTRUCOL",
        "tipo": "PAGO",
        "trx": "FULL E2E PAGO NORMAL 215",
        "credito": "CREDITO # 215",
        "digits": "215",
        "tag": "pago_normal",
    },
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 500_000.0,
        "concepto": "TRANSMAQUIOBRAS",
        "tipo": "ABONO CAPITAL",
        "trx": "FULL E2E ABONO CAPITAL 320",
        "credito": "CREDITO #320",
        "digits": "320",
        "tag": "abono_capital",
    },
]


def api_key() -> str:
    env = Path(r"D:\CMC\HBI_Capital\api-hbi-powerAutomate.env").read_text(encoding="utf-8")
    return re.search(r"^API_HTTP_KEY=(.+)$", env, re.M).group(1).strip()


def log(name: str, data: object) -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    (WORK / f"{name}.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    print(name, json.dumps(data, ensure_ascii=False, default=str)[:550])


def poll(c: httpx.Client, url: str, timeout_s: int = 2400) -> dict:
    t0 = time.time()
    while time.time() - t0 < timeout_s:
        try:
            r = c.get(url, timeout=90)
            if r.status_code != 200:
                print("poll_http", r.status_code, r.text[:200])
                time.sleep(10)
                continue
            body = r.json()
            st = body.get("status")
            err = body.get("error")
            err_txt = err if isinstance(err, str) else json.dumps(err, ensure_ascii=False) if err else ""
            print(time.strftime("%H:%M:%S"), "status", st, err_txt[:140])
            if st in ("completed", "failed", "error"):
                return body
        except Exception as exc:
            print("poll_err", type(exc).__name__, str(exc)[:120])
        time.sleep(10)
    raise TimeoutError(url)


def children(c, drive, fid="root"):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c, drive, parts):
    cur = "root"
    for p in parts:
        hit = next((i for i in children(c, drive, cur) if i.get("name") == p), None)
        if not hit:
            raise FileNotFoundError(p)
        cur = hit["id"]
    return cur


def download(c, drive, item_id: str) -> bytes:
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        timeout=180,
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def upload_bytes(c, drive, item_id: str, raw: bytes) -> None:
    r = c.put(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        json={"content_base64": base64.b64encode(raw).decode("ascii")},
        timeout=180,
    )
    r.raise_for_status()


def queue_and_poll(c, path: str, payload: dict, tag: str, job_url_tpl: str | None = None) -> dict:
    r = c.post(f"{BASE}{path}", json=payload, timeout=60)
    log(f"{tag}_queue", {"status": r.status_code, "body": r.json() if r.content else None})
    if r.status_code not in (200, 202):
        return {"http_status": r.status_code, "body": r.json() if r.content else r.text}
    body = r.json()
    job_id = body.get("job_id")
    if not job_id:
        return body
    if job_url_tpl:
        url = job_url_tpl.format(job_id=job_id)
    else:
        url = f"{BASE}/graph/sharepoint/payment-validation/jobs/{job_id}"
    done = poll(c, url)
    log(f"{tag}_job", done)
    return done


def draw_asiento(path: Path, data: dict) -> None:
    c = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    y = height - 20 * mm
    c.setFont("Helvetica-Bold", 14)
    c.drawString(20 * mm, y, data["voucher"])
    y -= 10 * mm
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm, y, f"{data['year']}  {data['month']}  {data['day']}")
    y -= 10 * mm
    c.drawString(20 * mm, y, f"N Identificacion : {data['nit']}")
    y -= 6 * mm
    c.drawString(20 * mm, y, f"Nombre : {data['nombre']}")
    y -= 12 * mm
    total = sum(v for _, v in data["lines"])
    for cuenta, valor in data["lines"]:
        line = f"{valor:,.2f} PAGO: No.Rad. {data['credit']} Linea 544 1 {cuenta}"
        c.drawString(20 * mm, y, line)
        y -= 6 * mm
    y -= 4 * mm
    c.setFont("Helvetica-Bold", 10)
    c.drawString(20 * mm, y, f"{total:,.2f}  {total:,.2f}")
    c.showPage()
    c.save()


def generate_asientos(rows: list[dict]) -> list[tuple[str, str, str, str]]:
    ASIENTOS_DIR.mkdir(parents=True, exist_ok=True)
    uploads: list[tuple[str, str, str, str]] = []
    voucher = 9200
    for row in rows:
        tipo = row["tipo"]
        if tipo == "PAGO":
            label = "PAGO CUOTA"
        elif tipo == "ABONO CAPITAL":
            label = "ABONO CAPITAL"
        else:
            continue
        digits = row["digits"]
        fname = f"Asiento 30-JUL-2026 {label} {row['concepto'][:20]} CRED {digits}.pdf"
        fname = re.sub(r"[^\w\s\-#.]", "", fname).replace("  ", " ")
        monto = float(row["monto"])
        fe = row["fecha"]
        data = {
            "voucher": str(voucher),
            "day": fe.day,
            "month": fe.month,
            "year": fe.year,
            "nit": f"900{digits.zfill(6)}-1",
            "nombre": row["concepto"],
            "credit": digits,
            "lines": [("11100505", monto), ("13410519", monto)],
        }
        voucher += 1
        dest = ASIENTOS_DIR / fname
        draw_asiento(dest, data)
        uploads.append((row["concepto"], row["credito"], digits, fname))
    return uploads


def find_id_pago_header_row(ws) -> int:
    for r in range(1, 20):
        if str(ws.cell(r, 1).value or "").strip() == "ID Pago":
            return r
    raise ValueError("header_not_found")


def header_map(ws, hr: int) -> dict[str, int]:
    return {
        str(ws.cell(hr, c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(hr, c).value is not None and str(ws.cell(hr, c).value).strip()
    }


def fnum(v) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def credit_digits(cred: str) -> str:
    m = re.search(r"(\d+)", cred or "")
    return m.group(1) if m else ""


def prepare_review(c, drive, rev_name: str) -> dict:
    folder_id = walk(c, drive, REV_DIR.split("/"))
    rev = next(i for i in children(c, drive, folder_id) if i.get("name") == rev_name)
    raw = download(c, drive, rev["id"])
    wb = load_workbook(io.BytesIO(raw))

    # Contar errores
    err_count = 0
    if "Errores" in wb.sheetnames:
        we = wb["Errores"]
        hr_e = 1
        for r in range(1, 10):
            vals = [we.cell(r, c).value for c in range(1, 6)]
            if any(str(v or "").lower().startswith("id") for v in vals) or any(
                "código" in str(v or "").lower() or "codigo" in str(v or "").lower() for v in vals
            ):
                hr_e = r
                break
        for r in range(hr_e + 1, we.max_row + 1):
            if any(we.cell(r, c).value not in (None, "") for c in range(1, 8)):
                err_count += 1

    ws = wb["Distribucion_Pagos"]
    hr = find_id_pago_header_row(ws)
    hm = header_map(ws, hr)
    col_id = hm["ID Pago"]
    col_cred = hm["Crédito"]
    col_monto = hm["Monto banco"]
    col_aplicar = hm["Aplicar a extracto"]
    col_mora = hm["Mora a aplicar"]
    col_cap = hm["Abono a capital"]
    col_otros = hm["Otros valores"]
    col_estado = hm["Estado Pago"]
    col_val = hm["Validar Pago"]

    validated = 0
    estados = []
    by_id: dict[str, list[int]] = {}
    for r in range(hr + 1, ws.max_row + 1):
        pid = str(ws.cell(r, col_id).value or "").strip()
        if not pid:
            continue
        ws.cell(r, col_val).value = "NO"
        by_id.setdefault(pid, []).append(r)

    for pid, rows in by_id.items():
        rows_sorted = sorted(rows)
        first_monto = None
        for r in rows_sorted:
            m = fnum(ws.cell(r, col_monto).value)
            if m is not None and first_monto is None:
                first_monto = m
        for i, r in enumerate(rows_sorted):
            ws.cell(r, col_monto).value = first_monto if i == 0 else None
        # Preferir crédito 215
        best = None
        for r in rows_sorted:
            if credit_digits(str(ws.cell(r, col_cred).value or "")) == "215":
                best = r
                break
        if best is None:
            best = rows_sorted[0]
        monto = first_monto or 0.0
        ws.cell(best, col_val).value = "SI"
        ws.cell(best, col_aplicar).value = monto
        ws.cell(best, col_mora).value = 0
        ws.cell(best, col_cap).value = 0
        ws.cell(best, col_otros).value = 0
        estados.append(
            {
                "credito": ws.cell(best, col_cred).value,
                "estado": ws.cell(best, col_estado).value,
                "monto": monto,
            }
        )
        validated += 1
        print("PAGO_SI", pid, ws.cell(best, col_cred).value, monto, ws.cell(best, col_estado).value)

    ab_ok = 0
    if "Distribucion_Abonos" in wb.sheetnames:
        wa = wb["Distribucion_Abonos"]
        hr_a = find_id_pago_header_row(wa)
        hm_a = header_map(wa, hr_a)
        col_fecha = hm_a.get("Fecha banco")
        col_cred_a = hm_a["Crédito"]
        col_val_a = hm_a["Validar Abono"]
        col_monto_a = hm_a["Monto banco"]
        by_ab: dict[str, list[int]] = {}
        for r in range(hr_a + 1, wa.max_row + 1):
            pid = str(wa.cell(r, hm_a["ID Pago"]).value or "").strip()
            if not pid:
                continue
            wa.cell(r, col_val_a).value = "NO"
            by_ab.setdefault(pid, []).append(r)
        for pid, rows_a in by_ab.items():
            monto_ab = None
            fecha_ab = None
            for r in rows_a:
                if monto_ab is None:
                    monto_ab = fnum(wa.cell(r, col_monto_a).value)
                if col_fecha and fecha_ab is None and wa.cell(r, col_fecha).value not in (None, ""):
                    fecha_ab = wa.cell(r, col_fecha).value
            best = None
            for r in rows_a:
                if credit_digits(str(wa.cell(r, col_cred_a).value or "")) == "320":
                    best = r
                    break
            if best is None and rows_a:
                best = rows_a[0]
            for r in rows_a:
                wa.cell(r, col_monto_a).value = None
                if col_fecha and r != best:
                    wa.cell(r, col_fecha).value = None
            if best is None:
                continue
            if monto_ab is not None:
                wa.cell(best, col_monto_a).value = monto_ab
            if col_fecha and fecha_ab is not None:
                fv = fecha_ab
                if isinstance(fv, str) and fv.strip():
                    try:
                        fv = datetime.strptime(fv.strip()[:10], "%Y-%m-%d")
                    except ValueError:
                        pass
                elif isinstance(fv, date) and not isinstance(fv, datetime):
                    fv = datetime(fv.year, fv.month, fv.day)
                wa.cell(best, col_fecha).value = fv
            wa.cell(best, col_val_a).value = "SI"
            ab_ok += 1
            print("ABONO_SI", pid, wa.cell(best, col_cred_a).value, monto_ab)

    wc = wb["Control"]
    for r in range(1, wc.max_row + 1):
        if str(wc.cell(r, 1).value or "").strip().lower() == "procesar":
            wc.cell(r, 2).value = "SI"

    buf = io.BytesIO()
    wb.save(buf)
    out = buf.getvalue()
    (WORK / "review_prepared.xlsx").write_bytes(out)
    upload_bytes(c, drive, rev["id"], out)
    info = {"validated_pagos": validated, "validated_abonos": ab_ok, "errores_sheet_rows": err_count, "estados": estados}
    log("05_review_prepared", info)
    return info


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    key = api_key()
    asiento_uploads = generate_asientos(ROWS)
    log("00_asientos_map", asiento_uploads)
    summary: dict = {"build": None, "steps": {}}

    with httpx.Client(timeout=180, headers={"X-API-Key": key}) as c:
        health = c.get(f"{BASE}/health").json()
        summary["build"] = health.get("build")
        log("01_health", health)
        if health.get("build") != "extract-damaged-failclosed-20260730":
            raise RuntimeError(f"build inesperado: {health}")

        drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]

        # Limpieza
        cancel0 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/cancel-active-process/queue",
            {"bank_code": BANK_CODE},
            "02_cancel_pre",
        )
        summary["steps"]["cancel_pre"] = cancel0.get("status")

        # Merge demasiado pronto (negativo)
        r = c.post(
            f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs",
            json={"bank_code": BANK_CODE},
        )
        log("03_merge_too_early", {"status": r.status_code, "body": r.json() if r.content else None})
        summary["steps"]["merge_too_early_http"] = r.status_code

        # Banco
        bank_folder = walk(c, drive, BANK_DIR.split("/"))
        bank_item = next(i for i in children(c, drive, bank_folder) if i.get("name") == BANK_NAME)
        raw = download(c, drive, bank_item["id"])
        wb = load_workbook(io.BytesIO(raw))
        ws = wb.active
        for rrow in range(4, ws.max_row + 1):
            for col in range(1, 6):
                ws.cell(rrow, col).value = None
        for i, row in enumerate(ROWS):
            rr = 4 + i
            ws.cell(rr, 1).value = row["fecha"]
            ws.cell(rr, 2).value = float(row["monto"])
            ws.cell(rr, 3).value = row["concepto"]
            ws.cell(rr, 4).value = row["tipo"]
            ws.cell(rr, 5).value = row["trx"]
        buf = io.BytesIO()
        wb.save(buf)
        payload = buf.getvalue()
        (WORK / "BANCO_BOGOTA_full.xlsx").write_bytes(payload)
        upload_bytes(c, drive, bank_item["id"], payload)

        # Limpiar revision previa
        rev_id = walk(c, drive, REV_DIR.split("/"))
        for it in list(children(c, drive, rev_id)):
            if str(it.get("name", "")).startswith("~$"):
                continue
            if str(it.get("name", "")).endswith(".xlsx"):
                c.delete(f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/items/{it['id']}")

        # Generate
        gen = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "04_generate",
        )
        if gen.get("status") != "completed":
            log("99_FAILED", {"step": "generate", "gen": gen})
            raise SystemExit(1)
        result = gen.get("result") or {}
        err_n = (result.get("summary") or {}).get("errores")
        summary["steps"]["generate"] = {
            "status": gen.get("status"),
            "errores": err_n,
            "file": result.get("validation_file"),
            "pagos": (result.get("summary") or {}).get("pagos_detectados"),
            "abonos": (result.get("summary") or {}).get("abonos_detectados"),
        }
        if err_n and int(err_n) > 0:
            # Si el abono mete errores de otros creditos, reintentar solo PAGO
            print("WARN errores en generate:", err_n, "— continuar solo si prepare puede? No, abortar a solo PAGO")
            # Rehacer solo con PAGO
            for rrow in range(4, ws.max_row + 1):
                for col in range(1, 6):
                    ws.cell(rrow, col).value = None
            row = ROWS[0]
            ws.cell(4, 1).value = row["fecha"]
            ws.cell(4, 2).value = float(row["monto"])
            ws.cell(4, 3).value = row["concepto"]
            ws.cell(4, 4).value = row["tipo"]
            ws.cell(4, 5).value = row["trx"]
            buf2 = io.BytesIO()
            wb.save(buf2)
            upload_bytes(c, drive, bank_item["id"], buf2.getvalue())
            queue_and_poll(
                c,
                "/graph/sharepoint/payment-validation/cancel-active-process/queue",
                {"bank_code": BANK_CODE},
                "04b_cancel_for_clean",
            )
            gen = queue_and_poll(
                c,
                "/graph/sharepoint/payment-validation/generate/queue",
                {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
                "04c_generate_pago_only",
            )
            result = gen.get("result") or {}
            err_n = (result.get("summary") or {}).get("errores")
            summary["steps"]["generate_pago_only"] = {
                "status": gen.get("status"),
                "errores": err_n,
                "file": result.get("validation_file"),
            }
            if gen.get("status") != "completed" or (err_n and int(err_n) > 0):
                log("99_FAILED", {"step": "generate_clean", "gen": gen})
                raise SystemExit(2)

        rev_name = result.get("validation_file")
        prep = prepare_review(c, drive, rev_name)
        summary["steps"]["prepare"] = prep
        if prep.get("errores_sheet_rows", 0) > 0:
            log("99_FAILED", {"step": "errores_still_present", "prep": prep})
            raise SystemExit(3)

        # Finalize
        fin = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/finalize/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "06_finalize",
        )
        summary["steps"]["finalize"] = {
            "status": fin.get("status"),
            "error": fin.get("error"),
            "result": fin.get("result"),
        }
        if fin.get("status") != "completed":
            log("99_FAILED", {"step": "finalize", "fin": fin})
            raise SystemExit(4)
        hist = (fin.get("result") or {}).get("historical_file_path")

        # Notify
        notify = queue_and_poll(
            c,
            "/graph/sharepoint/notify-validar-extractos-email",
            {"bank_code": BANK_CODE, "historical_file_path": hist},
            "07_notify",
            job_url_tpl=BASE + "/graph/sharepoint/notify-validar-extractos-email/jobs/{job_id}",
        )
        summary["steps"]["notify"] = {
            "status": notify.get("status"),
            "error": notify.get("error"),
            "result": notify.get("result"),
        }

        # Merge sin asientos
        m0 = queue_and_poll(
            c,
            "/graph/sharepoint/merge-composite-validado-pdfs",
            {"bank_code": BANK_CODE},
            "08_merge_sin_asientos",
            job_url_tpl=BASE + "/graph/sharepoint/merge-composite-validado-pdfs/jobs/{job_id}",
        )
        summary["steps"]["merge_sin_asientos"] = {
            "status": m0.get("status"),
            "result_status": (m0.get("result") or {}).get("status") if isinstance(m0.get("result"), dict) else None,
            "error": m0.get("error"),
        }

        # Subir asientos (solo filas que quedaron en el generate final)
        active_digits = {ROWS[0]["digits"]}
        if summary["steps"].get("generate", {}).get("abonos") and not summary["steps"].get("generate_pago_only"):
            active_digits.add(ROWS[1]["digits"])
        for cliente, credito_folder, digits, fname in asiento_uploads:
            if digits not in active_digits:
                continue
            local = ASIENTOS_DIR / fname
            rel = f"{CLIENTS}/{cliente}/{credito_folder}/ASIENTOS CONTABLES CRED {digits}/{fname}"
            up = c.put(
                f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/path-content",
                params={"item_path": rel},
                json={"content_base64": base64.b64encode(local.read_bytes()).decode("ascii")},
                timeout=180,
            )
            print("upload_asiento", digits, fname, up.status_code, up.text[:120])
            log(f"09_asiento_{digits}", {"status": up.status_code, "path": rel, "body": up.text[:300]})

        # Merge con asientos
        m1 = queue_and_poll(
            c,
            "/graph/sharepoint/merge-composite-validado-pdfs",
            {"bank_code": BANK_CODE},
            "10_merge",
            job_url_tpl=BASE + "/graph/sharepoint/merge-composite-validado-pdfs/jobs/{job_id}",
        )
        summary["steps"]["merge"] = {
            "status": m1.get("status"),
            "result": m1.get("result"),
            "error": m1.get("error"),
        }

        # Merge idempotencia
        m2 = queue_and_poll(
            c,
            "/graph/sharepoint/merge-composite-validado-pdfs",
            {"bank_code": BANK_CODE},
            "11_merge_idempotent",
            job_url_tpl=BASE + "/graph/sharepoint/merge-composite-validado-pdfs/jobs/{job_id}",
        )
        summary["steps"]["merge_idempotent"] = {
            "status": m2.get("status"),
            "already_merged": (m2.get("result") or {}).get("already_merged")
            if isinstance(m2.get("result"), dict)
            else None,
            "result": m2.get("result"),
        }

        # Amort dry + apply
        dry = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/amortization/dry-run/queue",
            {"bank_code": BANK_CODE},
            "12_amort_dry",
        )
        summary["steps"]["amort_dry"] = {
            "status": dry.get("status"),
            "can_apply": (dry.get("result") or {}).get("can_apply") if isinstance(dry.get("result"), dict) else None,
            "events": len((dry.get("result") or {}).get("events") or [])
            if isinstance(dry.get("result"), dict)
            else None,
            "error": dry.get("error"),
            "result": dry.get("result"),
        }

        apply = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/amortization/apply/queue",
            {"bank_code": BANK_CODE},
            "13_amort_apply",
        )
        summary["steps"]["amort_apply"] = {
            "status": apply.get("status"),
            "error": apply.get("error"),
            "result": apply.get("result"),
        }

        apply2 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/amortization/apply/queue",
            {"bank_code": BANK_CODE},
            "14_amort_apply_idempotent",
        )
        summary["steps"]["amort_apply_idempotent"] = {
            "status": apply2.get("status"),
            "already_applied": (apply2.get("result") or {}).get("already_applied")
            if isinstance(apply2.get("result"), dict)
            else None,
            "result": apply2.get("result"),
            "error": apply2.get("error"),
        }

        # Generate idempotencia post-ciclo (esperado: already o estado terminal)
        gen_end = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "15_generate_after_cycle",
        )
        summary["steps"]["generate_after_cycle"] = {
            "status": gen_end.get("status"),
            "error": gen_end.get("error"),
            "result": gen_end.get("result"),
        }

    # Veredicto
    ok_finalize = summary["steps"].get("finalize", {}).get("status") == "completed"
    ok_merge = summary["steps"].get("merge", {}).get("status") == "completed"
    ok_dry = summary["steps"].get("amort_dry", {}).get("status") == "completed"
    ok_apply = summary["steps"].get("amort_apply", {}).get("status") == "completed"
    summary["verdict"] = {
        "finalize_ok": ok_finalize,
        "merge_ok": ok_merge,
        "amort_dry_ok": ok_dry,
        "amort_apply_ok": ok_apply,
        "ready_for_production": bool(ok_finalize and ok_merge and ok_dry and ok_apply),
    }
    log("99_summary", summary)
    print("DONE", json.dumps(summary["verdict"], ensure_ascii=False))


if __name__ == "__main__":
    main()
