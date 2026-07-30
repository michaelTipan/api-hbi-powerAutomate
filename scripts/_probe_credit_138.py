"""Busca CREDITO 138 en clientes prod y revisa Errores de la ultima revision."""
from __future__ import annotations

import base64
import io
import json
import re
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
WORK = Path(r"D:\CMC\HBI_Capital\_work\credit_138")
CLIENTS = "INFORMACION CREDITOS-CLIENTES"
REV = f"{CLIENTS}/02 VALIDACION PAGOS/01 REVISION"


def api_key() -> str:
    env = Path(r"D:\CMC\HBI_Capital\api-hbi-powerAutomate.env").read_text(encoding="utf-8")
    return re.search(r"^API_HTTP_KEY=(.+)$", env, re.M).group(1).strip()


def children(c, drive, fid="root"):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c, drive, parts):
    cur = "root"
    for p in parts:
        hit = next((i for i in children(c, drive, cur) if i.get("name") == p), None)
        if not hit:
            raise FileNotFoundError(p)
        cur = hit["id"]
    return cur


def download(c, drive, item_id):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        timeout=180,
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    key = api_key()
    out: dict = {"hits": [], "review": None}
    with httpx.Client(timeout=180, headers={"X-API-Key": key}) as c:
        print("health", c.get(f"{BASE}/health").json())
        drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]
        root = walk(c, drive, CLIENTS.split("/"))
        clients = [i for i in children(c, drive, root) if i.get("folder")]
        print("clients", len(clients))
        for cli in clients:
            name = str(cli.get("name") or "")
            if name.startswith("0") or "PRUEBAS" in name.upper() or "VALIDACION" in name.upper():
                continue
            try:
                units = children(c, drive, cli["id"])
            except Exception as exc:
                print("skip", name, exc)
                continue
            for u in units:
                un = str(u.get("name") or "")
                digs = re.findall(r"\d+", un)
                if "138" not in digs and not re.search(r"(?i)#\s*138\b", un):
                    continue
                # list extracts
                files = []
                extractos = []
                try:
                    items = children(c, drive, u["id"])
                except Exception as exc:
                    items = []
                    print("unit_list_fail", name, un, exc)
                for it in items:
                    iname = str(it.get("name") or "")
                    if it.get("folder") and iname.upper() == "EXTRACTOS":
                        try:
                            extractos = [
                                x.get("name")
                                for x in children(c, drive, it["id"])
                                if str(x.get("name", "")).lower().endswith(".pdf")
                            ]
                        except Exception:
                            extractos = ["<error listing>"]
                    elif str(iname).lower().endswith(".pdf") and "extracto" in iname.lower():
                        files.append(iname)
                hit = {
                    "client": name,
                    "credit": un,
                    "extracts_root": files,
                    "extracts_extractos": extractos,
                    "has_tabla": any("amort" in str(it.get("name", "")).lower() for it in items),
                }
                out["hits"].append(hit)
                print("HIT", json.dumps(hit, ensure_ascii=False))

        # ultima revision
        try:
            rev_id = walk(c, drive, REV.split("/"))
            revs = [
                i
                for i in children(c, drive, rev_id)
                if str(i.get("name", "")).endswith(".xlsx") and not str(i.get("name")).startswith("~$")
            ]
            revs = sorted(revs, key=lambda i: i.get("lastModifiedDateTime") or "", reverse=True)
            if revs:
                newest = revs[0]
                raw = download(c, drive, newest["id"])
                (WORK / newest["name"]).write_bytes(raw)
                wb = load_workbook(io.BytesIO(raw), data_only=True)
                focus = {"file": newest["name"], "sheets": wb.sheetnames, "errores_138": [], "dist_138": []}
                for sheet in ("Errores", "Distribucion_Pagos", "Distribucion_Abonos"):
                    if sheet not in wb.sheetnames:
                        continue
                    ws = wb[sheet]
                    rows = list(ws.iter_rows(values_only=True))
                    if not rows:
                        continue
                    # header
                    hr = 0
                    for i, row in enumerate(rows[:10]):
                        vals = [str(v or "") for v in row]
                        if any("crédito" in v.lower() or "credito" in v.lower() for v in vals) or any(
                            "código" in v.lower() or "codigo" in v.lower() for v in vals
                        ):
                            hr = i
                            break
                    headers = [str(h or "").strip() for h in rows[hr]]
                    for row in rows[hr + 1 :]:
                        joined = " ".join(str(v or "") for v in row)
                        if "138" not in joined:
                            continue
                        item = {
                            headers[i]: row[i]
                            for i in range(min(len(headers), len(row)))
                            if headers[i]
                        }
                        if sheet == "Errores":
                            focus["errores_138"].append(item)
                        else:
                            focus["dist_138"].append({"sheet": sheet, **item})
                out["review"] = focus
                print("REVIEW", newest["name"], "err", len(focus["errores_138"]), "dist", len(focus["dist_138"]))
        except Exception as exc:
            print("review_fail", exc)
            out["review_error"] = str(exc)

    (WORK / "report.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    print("DONE hits", len(out["hits"]))


if __name__ == "__main__":
    main()
