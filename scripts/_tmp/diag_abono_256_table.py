# -*- coding: utf-8 -*-
"""RO: inspect Tabla amortizacion ICOEQUIPOS for abono planning blockers."""
from __future__ import annotations

import asyncio
import base64
import io
import json
import os
import sys
from pathlib import Path
from urllib.parse import quote

from dotenv import load_dotenv
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
load_dotenv(Path(r"D:\CMC\HBI_Capital\api-hbi-powerAutomate.env"), override=False)
os.environ["UI_SHAREPOINT_SMOKE"] = "1"

from app.application.services.amortization_workbook import (  # noqa: E402
    AmortizationSheetNotFoundError,
    AmbiguousAmortizationHeadersError,
    detect_amortization_sheet,
    find_next_available_application_row,
    human_required_header_labels,
)
from app.application.services.accounting_pdf_parser import PaymentApplicationEvent  # noqa: E402
from datetime import date  # noqa: E402
from scripts.support.ui_appservice_graph_http import AppServiceGraphHttp  # noqa: E402

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
API = (os.getenv("API_HTTP_KEY") or "").strip()
TABLA = (
    "INFORMACION CREDITOS-CLIENTES/03 COMWARE PRUEBAS- INFORMACION CREDITOS CLIENTES/"
    "ICOEQUIPOS SAS/Tabla de amortizacion ICOEQUIPOS.xlsx"
)
OUT = ROOT / "scripts" / "_tmp" / "diag_abono_256_table.json"


async def children(http, drive, fid="root"):
    r = await http._client.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
    )
    r.raise_for_status()
    return list((r.json() or {}).get("value") or [])


async def walk(http, drive, parts):
    cur = "root"
    for i, p in enumerate(parts):
        kids = await children(http, drive, cur)
        hit = next((x for x in kids if x.get("name") == p), None)
        if not hit:
            raise FileNotFoundError("/".join(parts[: i + 1]))
        cur = str(hit["id"])
    return cur


async def download(http, drive, rel):
    parts = [p for p in rel.strip("/").split("/") if p]
    fid = await walk(http, drive, parts[:-1])
    kids = await children(http, drive, fid)
    hit = next((x for x in kids if x.get("name") == parts[-1]), None)
    if not hit:
        raise FileNotFoundError(rel)
    r = await http._client.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": str(hit["id"])},
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


async def main():
    http = AppServiceGraphHttp(base_url=BASE, api_key=API, timeout=180.0)
    report: dict = {"tabla": TABLA}
    try:
        diag = await http._client.get(f"{BASE}/graph/diagnostics")
        diag.raise_for_status()
        drive = str(((diag.json() or {}).get("operations_site") or {}).get("drive_id") or "")
        raw = await download(http, drive, TABLA)
        report["bytes"] = len(raw)
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
        report["sheets"] = wb.sheetnames
        try:
            match = detect_amortization_sheet(wb, tabla_amortizacion_path=TABLA)
            report["detect"] = {
                "ok": True,
                "sheet": match.worksheet.title,
                "header_row": match.header_row,
                "headers": match.headers,
            }
            # Try find application row for abono event
            event = PaymentApplicationEvent(
                id_pago="2c91424b",
                cliente="ICOEQUIPOS SAS",
                credito="256",
                asiento_pdf_path="x.pdf",
                comprobante="1",
                fecha_asiento=date(2026, 7, 12),
                valor_pagado_cliente=40_000_000.0,
                capital=39_141_787.0,
                intereses=858_213.0,
                mora=0.0,
                retenciones=0.0,
                saldos_menores=0.0,
                raw_text="",
            )
            try:
                app = find_next_available_application_row(
                    match.worksheet,
                    match.headers,
                    event,
                    reserved_rows=set(),
                    header_row=match.header_row,
                    payment_date=date(2026, 7, 12),
                    detected_codes=frozenset(),
                )
                report["application_row"] = {
                    "status": getattr(app, "status", None) or (app.get("status") if isinstance(app, dict) else str(type(app))),
                    "repr": repr(app)[:500],
                }
            except Exception as exc:
                report["application_row"] = {"error": f"{type(exc).__name__}: {exc}"}
        except AmortizationSheetNotFoundError as exc:
            report["detect"] = {
                "ok": False,
                "error": "AMORTIZATION_SHEET_NOT_FOUND",
                "missing": exc.required_headers_missing,
                "missing_labels": human_required_header_labels(exc.required_headers_missing),
                "found": exc.headers_found,
                "best_sheet": exc.best_sheet_name,
                "best_row": exc.best_header_row,
                "labels": exc.header_labels,
            }
        except AmbiguousAmortizationHeadersError as exc:
            report["detect"] = {
                "ok": False,
                "error": "AMBIGUOUS_HEADER_MAPPING",
                "key": exc.key,
                "labels": exc.labels,
            }

        # Dump header-ish rows from first sheet
        ws = wb[wb.sheetnames[0]]
        preview = []
        for r in range(1, min(8, (ws.max_row or 1) + 1)):
            vals = []
            for c in range(1, min(25, (ws.max_column or 1) + 1)):
                v = ws.cell(r, c).value
                if v is not None and str(v).strip():
                    vals.append(f"{c}:{v}")
            preview.append({"row": r, "cells": vals})
        report["header_preview"] = preview
    finally:
        await http.aclose()

    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    asyncio.run(main())
