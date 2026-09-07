# -*- coding: utf-8 -*-
"""RO: Maderpol 248 — histórico, jobs amortización, estructura SP."""
from __future__ import annotations

import asyncio
import base64
import json
import os
import re
import sys
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

import httpx
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
load_dotenv(Path(r"D:\CMC\HBI_Capital\api-hbi-powerAutomate.env"), override=False)
os.environ["UI_SHAREPOINT_SMOKE"] = "1"

from openpyxl import load_workbook

from app.application.config.payment_validation_settings import (  # noqa: E402
    get_payment_validation_paths,
    list_payment_banks,
    resolve_bank_control_file_path,
)
from app.application.services.payment_helpers import find_best_amortization_table
from app.application.use_cases.payment_validation_process_control import (  # noqa: E402
    parse_process_control_row2,
)
from scripts.support.ui_appservice_graph_http import AppServiceGraphHttp  # noqa: E402

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
OUT = ROOT / "scripts" / "_tmp" / "diag_maderpol_248_deep.json"
API_KEY = (os.getenv("API_HTTP_KEY") or "").strip()


async def _children(http: AppServiceGraphHttp, drive: str, fid: str = "root"):
    r = await http._client.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
    )
    r.raise_for_status()
    return list((r.json() or {}).get("value") or [])


async def _walk(http: AppServiceGraphHttp, drive: str, parts: list[str]) -> str:
    cur = "root"
    for i, p in enumerate(parts):
        kids = await _children(http, drive, cur)
        hit = next((x for x in kids if x.get("name") == p), None)
        if not hit:
            raise FileNotFoundError("/".join(parts[: i + 1]))
        cur = str(hit["id"])
    return cur


async def _download_by_path(http: AppServiceGraphHttp, drive: str, rel_path: str) -> bytes:
    parts = [p for p in rel_path.strip("/").split("/") if p]
    name = parts[-1]
    folder_id = await _walk(http, drive, parts[:-1]) if len(parts) > 1 else "root"
    kids = await _children(http, drive, folder_id)
    hit = next((x for x in kids if x.get("name") == name), None)
    if not hit:
        raise FileNotFoundError(rel_path)
    r = await http._client.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": str(hit["id"])},
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def _rows_248_from_historical(raw: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    out: list[dict] = []
    for sheet in wb.sheetnames:
        if "aplicacion" not in sheet.lower():
            continue
        ws = wb[sheet]
        headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {h: i for i, h in enumerate(headers) if h}
        for row in ws.iter_rows(min_row=2, values_only=True):
            cred = str(row[idx.get("Crédito", idx.get("Credito", -1))] or "") if idx else ""
            if "248" not in cred and "MADERPOL" not in str(row).upper():
                continue
            out.append(
                {
                    "sheet": sheet,
                    "id_pago": row[idx.get("ID Pago", -1)] if "ID Pago" in idx else None,
                    "credito": cred,
                    "cliente": row[idx.get("Cliente", -1)] if "Cliente" in idx else None,
                    "tabla": row[idx.get("Link tabla amortización", idx.get("Link tabla amortizacion", -1))]
                    if idx
                    else None,
                    "ruta_asientos": row[idx.get("RutaAsientosContables", -1)] if idx else None,
                }
            )
    wb.close()
    return out


async def main() -> None:
    report: dict = {}
    http = AppServiceGraphHttp(base_url=BASE, api_key=API_KEY, timeout=180.0)
    try:
        diag = await http._client.get(f"{BASE}/graph/diagnostics")
        diag.raise_for_status()
        drive = str(((diag.json() or {}).get("operations_site") or {}).get("drive_id") or "")

        # Active control historical
        for bank in list_payment_banks():
            cpath = resolve_bank_control_file_path(bank.bank_code)
            try:
                raw = await _download_by_path(http, drive, cpath)
                snap = parse_process_control_row2(raw, control_file_path=cpath)
                if not snap.is_active or not snap.historical_file_path:
                    continue
                hist_raw = await _download_by_path(http, drive, snap.historical_file_path)
                rows = _rows_248_from_historical(hist_raw)
                report["active_bank"] = bank.bank_code
                report["process_key"] = snap.process_key
                report["historical_path"] = snap.historical_file_path
                report["historical_rows_248"] = rows
                report["last_amort_json_len"] = len(snap.last_amortization_attempt_json or "")
            except Exception as exc:
                report[f"bank_{bank.bank_code}_error"] = str(exc)

        # MADERPOL tree
        clients_base = os.getenv("GRAPH_CLIENTS_BASE_PATH", "").strip("/")
        base_parts = [p for p in clients_base.split("/") if p]
        folder_id = await _walk(http, drive, base_parts)
        kids = await _children(http, drive, folder_id)
        client = next(
            (k for k in kids if "MADERPOL" in str(k.get("name") or "").upper()),
            None,
        )
        if client:
            tree: dict = {"client": client.get("name"), "children": []}
            for child in await _children(http, drive, str(client["id"])):
                cname = str(child.get("name") or "")
                node = {"name": cname, "is_folder": bool(child.get("folder"))}
                if child.get("folder"):
                    sub = await _children(http, drive, str(child["id"]))
                    node["files"] = sorted(str(s.get("name") or "") for s in sub)[:40]
                    xlsx = [f for f in node["files"] if f.lower().endswith((".xlsx", ".xlsm"))]
                    if xlsx:
                        try:
                            find_best_amortization_table(
                                xlsx, client.get("name") or "MADERPOL", "248"
                            )
                            node["find_best_result"] = "ok"
                        except ValueError as exc:
                            node["find_best_error"] = str(exc)
                tree["children"].append(node)
            report["maderpol_tree"] = tree

        # All amortization jobs (recent)
        paths = get_payment_validation_paths()
        logs_parts = [p for p in paths.logs.strip("/").split("/") if p]
        logs_id = await _walk(http, drive, logs_parts)
        amort_jobs = []
        for lf in await _children(http, drive, logs_id):
            fname = str(lf.get("name") or "")
            if not fname.lower().endswith(".json"):
                continue
            try:
                r = await http._client.get(
                    f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
                    params={"item_id": str(lf["id"])},
                )
                r.raise_for_status()
                job = json.loads(
                    base64.b64decode(r.json()["content_base64"]).decode("utf-8", errors="replace")
                )
                jtype = str(job.get("type") or "")
                if "amortization" not in jtype:
                    continue
                result = job.get("result") if isinstance(job.get("result"), dict) else {}
                amort_jobs.append(
                    {
                        "file": fname,
                        "job_id": job.get("job_id"),
                        "type": jtype,
                        "status": job.get("status"),
                        "modified": lf.get("lastModifiedDateTime"),
                        "outcome": result.get("outcome"),
                        "can_apply": result.get("can_apply"),
                        "issues": len(result.get("operational_issues") or []),
                        "items_with_error": sum(
                            1
                            for it in (result.get("items") or [])
                            if isinstance(it, dict) and it.get("error_code")
                        ),
                        "mentions_248": "248" in json.dumps(result, ensure_ascii=False),
                    }
                )
            except Exception:
                continue
        report["amort_jobs"] = sorted(
            amort_jobs, key=lambda x: x.get("modified") or "", reverse=True
        )[:15]
    finally:
        await http.aclose()

    OUT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    asyncio.run(main())
