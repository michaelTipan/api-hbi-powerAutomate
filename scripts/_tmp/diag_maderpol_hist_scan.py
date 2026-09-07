# -*- coding: utf-8 -*-
"""RO: buscar Maderpol/248 en históricos recientes sandbox + errores generate."""
from __future__ import annotations

import asyncio
import base64
import json
import os
import re
import sys
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
load_dotenv(Path(r"D:\CMC\HBI_Capital\api-hbi-powerAutomate.env"), override=False)
os.environ["UI_SHAREPOINT_SMOKE"] = "1"

from openpyxl import load_workbook

from app.application.config.payment_validation_settings import get_payment_validation_paths
from scripts.support.ui_appservice_graph_http import AppServiceGraphHttp

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
API_KEY = (os.getenv("API_HTTP_KEY") or "").strip()


async def _children(http, drive, fid="root"):
    r = await http._client.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
    )
    r.raise_for_status()
    return list((r.json() or {}).get("value") or [])


async def _walk(http, drive, parts):
    cur = "root"
    for i, p in enumerate(parts):
        kids = await _children(http, drive, cur)
        hit = next((x for x in kids if x.get("name") == p), None)
        if not hit:
            raise FileNotFoundError("/".join(parts[: i + 1]))
        cur = str(hit["id"])
    return cur


async def _download(http, drive, rel_path):
    parts = [p for p in rel_path.strip("/").split("/") if p]
    name = parts[-1]
    folder_id = await _walk(http, drive, parts[:-1]) if len(parts) > 1 else "root"
    kids = await _children(http, drive, folder_id)
    hit = next((x for x in kids if x.get("name") == name), None)
    if not hit:
        raise FileNotFoundError(rel_path)
    r = await http._client.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": str(hit["id"])},
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def scan_workbook(raw: bytes, label: str) -> dict:
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    out = {"label": label, "maderpol_rows": [], "errores_248": []}
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(c or "").strip() for c in rows[0]]
        hlow = [h.lower() for h in headers]
        if "errores" in sheet.lower() or sheet.lower() == "errores":
            code_i = next((i for i, h in enumerate(hlow) if "codigo" in h or h == "code"), None)
            desc_i = next((i for i, h in enumerate(hlow) if "descripcion" in h or "descrip" in h), None)
            cred_i = next((i for i, h in enumerate(hlow) if "credito" in h or "crédito" in h), None)
            for row in rows[1:]:
                blob = " ".join(str(c or "") for c in row).upper()
                if "248" not in blob and "MADERPOL" not in blob:
                    continue
                out["errores_248"].append(
                    {
                        "sheet": sheet,
                        "code": row[code_i] if code_i is not None else None,
                        "desc": row[desc_i] if desc_i is not None else None,
                        "credito": row[cred_i] if cred_i is not None else None,
                        "row_text": blob[:300],
                    }
                )
        if "aplicacion" in sheet.lower():
            idx = {headers[i]: i for i in range(len(headers)) if headers[i]}
            for row in rows[1:]:
                blob = " ".join(str(c or "") for c in row).upper()
                if "248" not in blob and "MADERPOL" not in blob:
                    continue
                out["maderpol_rows"].append(
                    {
                        "sheet": sheet,
                        "id_pago": row[idx.get("ID Pago", -1)] if "ID Pago" in idx else None,
                        "cliente": row[idx.get("Cliente", -1)] if "Cliente" in idx else None,
                        "credito": row[idx.get("Crédito", idx.get("Credito", -1))]
                        if ("Crédito" in idx or "Credito" in idx)
                        else None,
                        "tabla": row[
                            idx.get(
                                "Link tabla amortización",
                                idx.get("Link tabla amortizacion", -1),
                            )
                        ]
                        if ("Link tabla amortización" in idx or "Link tabla amortizacion" in idx)
                        else None,
                    }
                )
    wb.close()
    return out


async def main():
    report = {"historicos": [], "reviews": []}
    http = AppServiceGraphHttp(base_url=BASE, api_key=API_KEY, timeout=180.0)
    try:
        diag = await http._client.get(f"{BASE}/graph/diagnostics")
        drive = str(((diag.json() or {}).get("operations_site") or {}).get("drive_id") or "")
        paths = get_payment_validation_paths()
        hist_root = paths.historical
        parts = [p for p in hist_root.strip("/").split("/") if p]
        # walk year/month folders — last 2 months
        y_id = await _walk(http, drive, parts)
        for y in sorted(await _children(http, drive, y_id), reverse=True)[:1]:
            yname = str(y.get("name") or "")
            m_id = str(y["id"])
            for m in sorted(await _children(http, drive, m_id), reverse=True)[:2]:
                mname = str(m.get("name") or "")
                d_id = str(m["id"])
                for d in sorted(await _children(http, drive, d_id), reverse=True)[:5]:
                    for f in await _children(http, drive, str(d["id"])):
                        fname = str(f.get("name") or "")
                        if not fname.endswith(".xlsx") or "cartera_validada" not in fname.lower():
                            continue
                        rel = f"{hist_root}/{yname}/{mname}/{d.get('name')}/{fname}".replace("//", "/")
                        try:
                            raw = await _download(http, drive, rel)
                            scanned = scan_workbook(raw, rel)
                            if scanned["maderpol_rows"] or scanned["errores_248"]:
                                report["historicos"].append(scanned)
                        except Exception as exc:
                            report.setdefault("errors", []).append(f"{rel}: {exc}")
        rev_root = paths.review
        rev_parts = [p for p in rev_root.strip("/").split("/") if p]
        rev_id = await _walk(http, drive, rev_parts)
        for f in sorted(await _children(http, drive, rev_id), key=lambda x: x.get("lastModifiedDateTime") or "", reverse=True)[:15]:
            fname = str(f.get("name") or "")
            if not fname.endswith(".xlsx"):
                continue
            rel = f"{rev_root}/{fname}".replace("//", "/")
            try:
                raw = await _download(http, drive, rel)
                scanned = scan_workbook(raw, rel)
                if scanned["maderpol_rows"] or scanned["errores_248"]:
                    report["reviews"].append(scanned)
            except Exception:
                pass
    finally:
        await http.aclose()
    out = ROOT / "scripts" / "_tmp" / "diag_maderpol_hist_scan.json"
    out.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    asyncio.run(main())
