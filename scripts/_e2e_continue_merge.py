"""Continúa E2E: subir asientos a rutas del histórico, notify, merge, amort."""
from __future__ import annotations

import base64
import json
import time
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook
import io

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
WORK = Path(r"D:\CMC\HBI_Capital\_work")
ASIENTOS = Path(r"D:\CMC\HBI_Capital\asientos_prueba\generados")
HIST = (
    "INFORMACION CREDITOS-CLIENTES/"
    "02 COMWARE AUTOMATIZACION - INFORMACION CREDITOS CLIENTES/"
    "01 VALIDACION PAGOS/03 HISTORICO/cartera_validada_banco_bogota_2026-07-27.xlsx"
)


def log(name: str, data: object) -> None:
    (WORK / f"{name}.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    print(name, json.dumps(data, ensure_ascii=False, default=str)[:350])


def poll(c: httpx.Client, url: str, timeout_s: int = 900) -> dict:
    t0 = time.time()
    while time.time() - t0 < timeout_s:
        try:
            r = c.get(url, timeout=90)
            body = r.json()
            print(time.strftime("%H:%M:%S"), body.get("status"))
            if body.get("status") in ("completed", "failed", "error"):
                return body
        except Exception as exc:
            print("poll_err", type(exc).__name__)
        time.sleep(8)
    raise TimeoutError(url)


def children(c, drive, fid="root"):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c, drive, parts):
    cur = "root"
    for p in parts:
        hit = next((i for i in children(c, drive, cur) if i.get("name") == p), None)
        if not hit:
            raise FileNotFoundError(p)
        cur = hit["id"]
    return cur


def download(c, drive, item_id):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        timeout=180,
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def credit_digits(label: str) -> str:
    import re

    m = re.search(r"(\d{1,6})", label or "")
    return m.group(1) if m else ""


def pick_asiento_file(digits: str, tipo: str) -> Path | None:
    tipo_u = (tipo or "").upper()
    cands = sorted(ASIENTOS.glob(f"*CRED {digits}.pdf")) + sorted(
        ASIENTOS.glob(f"*CRED {digits} *.pdf")
    )
    # prefer matching tipo keywords
    scored = []
    for p in ASIENTOS.glob("Asiento*.pdf"):
        if f" {digits}." in p.name or p.name.endswith(f" {digits}.pdf") or f" CRED {digits}." in p.name:
            score = 0
            n = p.name.upper()
            if "ABONO CAPITAL" in tipo_u and "ABONO CAPITAL" in n and "PAGO Y" not in n:
                score += 5
            if "ABONO MORA" in tipo_u and "ABONO MORA" in n:
                score += 5
            if "PAGO Y ABONO" in tipo_u and "PAGO Y ABONO" in n:
                score += 5
            if tipo_u.startswith("PAGO") and "PAGO CUOTA" in n:
                score += 3
            scored.append((score, p))
    if not scored:
        return None
    scored.sort(key=lambda x: (-x[0], x[1].name))
    return scored[0][1]


def main() -> None:
    with httpx.Client(timeout=180) as c:
        drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]
        parts = HIST.split("/")
        folder = walk(c, drive, parts[:-1])
        item = next(i for i in children(c, drive, folder) if i.get("name") == parts[-1])
        raw = download(c, drive, item["id"])
        (WORK / "historical.xlsx").write_bytes(raw)
        wb = load_workbook(io.BytesIO(raw), data_only=True)
        print("hist sheets", wb.sheetnames)

        # Buscar hoja con RutaAsientosContables
        rutas: list[tuple[str, str, str]] = []
        for name in wb.sheetnames:
            ws = wb[name]
            headers = {}
            hr = None
            for r in range(1, 15):
                for col in range(1, min(40, ws.max_column + 1)):
                    v = ws.cell(r, col).value
                    if v and "rutaasientos" in str(v).lower().replace(" ", ""):
                        hr = r
                    if v and str(v).strip():
                        headers[str(v).strip()] = col
                if hr:
                    break
            if not hr:
                # try exact
                for r in range(1, 15):
                    vals = [str(ws.cell(r, col).value or "") for col in range(1, 30)]
                    if any("Asientos" in v for v in vals):
                        hr = r
                        headers = {
                            str(ws.cell(r, col).value).strip(): col
                            for col in range(1, ws.max_column + 1)
                            if ws.cell(r, col).value
                        }
                        break
            if not hr:
                continue
            col_ruta = None
            col_cred = None
            col_tipo = None
            for k, col in headers.items():
                kl = k.lower()
                if "asiento" in kl and "ruta" in kl:
                    col_ruta = col
                if k in ("Crédito", "Credito"):
                    col_cred = col
                if "tipoaplicacionoriginal" in kl.replace(" ", "") or k == "TipoAplicacionOriginal":
                    col_tipo = col
            print("sheet", name, "hr", hr, "ruta", col_ruta, "cred", col_cred)
            if not col_ruta:
                continue
            for r in range(hr + 1, ws.max_row + 1):
                ruta = ws.cell(r, col_ruta).value
                if not ruta:
                    continue
                cred = ws.cell(r, col_cred).value if col_cred else ""
                tipo = ws.cell(r, col_tipo).value if col_tipo else ""
                rutas.append((str(ruta).strip(), str(cred or ""), str(tipo or "")))

        print("rutas asientos", len(rutas))
        for ruta, cred, tipo in rutas:
            digits = credit_digits(cred) or credit_digits(ruta)
            pdf = pick_asiento_file(digits, tipo)
            if not pdf:
                print("no pdf for", digits, tipo)
                continue
            dest = f"{ruta.rstrip('/')}/{pdf.name}"
            up = c.put(
                f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/path-content",
                params={"item_path": dest},
                json={"content_base64": base64.b64encode(pdf.read_bytes()).decode("ascii")},
                timeout=180,
            )
            print("upload", digits, up.status_code, dest[-80:], up.text[:100])

        # Notify retry
        r = c.post(
            f"{BASE}/graph/sharepoint/notify-validar-extractos-email",
            json={"bank_code": "banco_bogota", "historical_file_path": HIST},
        )
        log("notify2_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        r.raise_for_status()
        n = poll(c, f"{BASE}/graph/sharepoint/notify-validar-extractos-email/jobs/{r.json()['job_id']}")
        log("notify2_result", n)

        # Merge + retry
        r = c.post(f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs", json={})
        log("mergeA_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        r.raise_for_status()
        m1 = poll(c, f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs/jobs/{r.json()['job_id']}")
        log("mergeA_result", m1)

        r = c.post(f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs", json={})
        log("mergeB_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        if r.status_code == 202:
            m2 = poll(c, f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs/jobs/{r.json()['job_id']}")
            log("mergeB_result", m2)

        r = c.post(
            f"{BASE}/graph/sharepoint/payment-validation/amortization/dry-run/queue",
            json={},
        )
        if r.status_code == 202:
            ad = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{r.json()['job_id']}")
            log("amort_dry2", ad)
        r = c.post(
            f"{BASE}/graph/sharepoint/payment-validation/amortization/apply/queue",
            json={},
        )
        if r.status_code == 202:
            aa = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{r.json()['job_id']}")
            log("amort_apply2", aa)


if __name__ == "__main__":
    main()
