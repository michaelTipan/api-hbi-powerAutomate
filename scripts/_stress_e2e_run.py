"""Batería de estrés E2E (créditos NUEVOS) — Generate → Amort + idempotencia/reintentos/errores.

No reutiliza la batería GEOEXCON/EQUINORTE/AGRECAR#37 del E2E previo.
Montos/fechas salen de extract_winners.json (regla max fecha_limite V2).
"""
from __future__ import annotations

import base64
import io
import json
import re
import time
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
WORK = Path(r"D:\CMC\HBI_Capital\_work\stress_e2e")
ASIENTOS_DIR = WORK / "asientos"
CLIENTS = (
    "INFORMACION CREDITOS-CLIENTES/"
    "02 COMWARE AUTOMATIZACION - INFORMACION CREDITOS CLIENTES"
)
BANK_ITEM_ID = "01UDV3W2ZAP2QTC7ZDR5CJGYVIWOYLPOKK"
REV_DIR = f"{CLIENTS}/01 VALIDACION PAGOS/02 REVISION"
PROCESS_DATE = "2026-07-28"
BANK_CODE = "banco_bogota"

# Clientes que NO se validan (casos negativos / error de usuario)
SKIP_VALIDATE = {
    "CLIENTE_FANTASMA_STRESS",
    "ACIMOR",
    "MINCIVIL",
    "INVERSIONES Y PROYECTOS MIOS",
}


def log(name: str, data: object) -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    path = WORK / f"{name}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(name, json.dumps(data, ensure_ascii=False, default=str)[:450])


def poll(c: httpx.Client, url: str, timeout_s: int = 2400) -> dict:
    t0 = time.time()
    while time.time() - t0 < timeout_s:
        try:
            r = c.get(url, timeout=90)
            if r.status_code != 200:
                print("poll_http", r.status_code, r.text[:200])
                time.sleep(10)
                continue
            body = r.json()
            st = body.get("status")
            print(time.strftime("%H:%M:%S"), "status", st)
            if st in ("completed", "failed", "error"):
                return body
        except Exception as exc:
            print("poll_err", type(exc).__name__, str(exc)[:120])
        time.sleep(10)
    raise TimeoutError(url)


def children(c, drive, fid="root"):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c, drive, parts):
    cur = "root"
    for p in parts:
        hit = next((i for i in children(c, drive, cur) if i.get("name") == p), None)
        if not hit:
            raise FileNotFoundError(p)
        cur = hit["id"]
    return cur


def download(c, drive, item_id: str) -> bytes:
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        timeout=180,
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def credit_digits(credito_folder: str) -> str:
    m = re.search(r"#\s*(\d+)", credito_folder)
    if m:
        return m.group(1)
    m = re.search(r"(\d{2,})", credito_folder)
    return m.group(1) if m else ""


def _norm_cli(s: str) -> str:
    return (
        str(s or "")
        .upper()
        .replace("Ñ", "N")
        .replace("ñ", "n")
        .encode("ascii", "ignore")
        .decode("ascii")
    )


def load_winners() -> dict[str, dict]:
    """Indexa ganadores por carpeta de crédito (única en sandbox)."""
    raw = json.loads((WORK / "extract_winners.json").read_text(encoding="utf-8"))
    out: dict[str, dict] = {}
    for row in raw:
        w = row.get("winner")
        if not w:
            continue
        cli = row.get("cliente_folder_real") or row["cliente"]
        out[row["credito"]] = {
            **w,
            "cliente": cli,
            "credito": row["credito"],
            "digits": credit_digits(row["credito"]),
        }
    return out


def build_bank_rows(winners: dict[str, dict]) -> list[dict]:
    """25 filas: ~18% atrasados en PAGOS principales; resto adelantados; todos los tipos."""

    def w(cli: str, cred: str) -> dict:
        val = winners.get(cred)
        if not val:
            raise KeyError(f"{cli}|{cred}")
        # Preferir nombre de carpeta real del inventario
        return val

    am = w("A&M CONSTRUCOL", "CREDITO # 215")
    d32 = w("DIEGO RAMIRO PAZMIÑO DIAZ", "1 CREDITO # 32")
    d92 = w("DIEGO RAMIRO PAZMIÑO DIAZ", "2 CREDITO # 92")
    d94 = w("DIEGO RAMIRO PAZMIÑO DIAZ", "3 CREDITO # 94")
    g224 = w("G&J INGENIERIA", "CREDITO # 224 VIGENTE")
    g284 = w("G&J INGENIERIA", "CREDITO # 284 VIGENTE")
    g296 = w("G&J INGENIERIA", "CREDITO # 296 VIGENTE")
    i150 = w("INDUCAB", "CREDITO # 150")
    i88 = w("INDUCAB", "CREDITO # 88")
    a71 = w("AGRECAR", "4 CREDITO # 71")

    rows: list[dict] = [
        # --- PAGOS adelantados (fecha banco < fecha límite) ---
        {
            "fecha": datetime(2026, 5, 10),
            "monto": am["total_a_pagar"],
            "concepto": am["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 215 ADELANTADO",
            "tag": "pago_adelantado",
            "credito": am["credito"],
            "digits": am["digits"],
        },
        {
            "fecha": datetime(2026, 4, 1),
            "monto": d32["total_a_pagar"],
            "concepto": d32["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 32 ADELANTADO",
            "tag": "pago_adelantado",
            "credito": d32["credito"],
            "digits": d32["digits"],
        },
        {
            "fecha": datetime(2026, 3, 20),
            "monto": d92["total_a_pagar"],
            "concepto": d92["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 92 ADELANTADO",
            "tag": "pago_adelantado",
            "credito": d92["credito"],
            "digits": d92["digits"],
        },
        {
            "fecha": datetime(2026, 3, 10),
            "monto": d94["total_a_pagar"],
            "concepto": d94["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 94 ADELANTADO",
            "tag": "pago_adelantado",
            "credito": d94["credito"],
            "digits": d94["digits"],
        },
        {
            "fecha": datetime(2026, 4, 15),
            "monto": g224["total_a_pagar"],
            "concepto": g224["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 224 ADELANTADO",
            "tag": "pago_adelantado",
            "credito": g224["credito"],
            "digits": g224["digits"],
        },
        {
            "fecha": datetime(2026, 4, 10),
            "monto": g284["total_a_pagar"],
            "concepto": g284["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 284 ADELANTADO",
            "tag": "pago_adelantado",
            "credito": g284["credito"],
            "digits": g284["digits"],
        },
        {
            "fecha": datetime(2026, 4, 10),
            "monto": g296["total_a_pagar"],
            "concepto": g296["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 296 ADELANTADO",
            "tag": "pago_adelantado",
            "credito": g296["credito"],
            "digits": g296["digits"],
        },
        {
            "fecha": datetime(2025, 9, 15),
            "monto": a71["total_a_pagar"],
            "concepto": a71["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 71 ADELANTADO",
            "tag": "pago_adelantado",
            "credito": a71["credito"],
            "digits": a71["digits"],
        },
        # --- PAGOS atrasados (~18% de los 11 PAGO principales = 2) ---
        {
            "fecha": datetime(2026, 7, 27),
            "monto": i150["total_a_pagar"],
            "concepto": i150["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 150 ATRASADO",
            "tag": "pago_atrasado",
            "credito": i150["credito"],
            "digits": i150["digits"],
        },
        {
            "fecha": datetime(2026, 7, 27),
            "monto": i88["total_a_pagar"],
            "concepto": i88["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO exacto 88 ATRASADO",
            "tag": "pago_atrasado",
            "credito": i88["credito"],
            "digits": i88["digits"],
        },
        # --- PAGO Y ABONO CAPITAL ---
        {
            "fecha": datetime(2026, 5, 8),
            "monto": float(am["total_a_pagar"]) + 2_000_000.0,
            "concepto": am["cliente"],
            "tipo": "PAGO Y ABONO CAPITAL",
            "trx": "STRESS cuota+capital 215",
            "tag": "pago_y_abono",
            "credito": am["credito"],
            "digits": am["digits"],
            "extra_capital": 2_000_000.0,
        },
        {
            "fecha": datetime(2026, 4, 2),
            "monto": float(d32["total_a_pagar"]) + 1_500_000.0,
            "concepto": d32["cliente"],
            "tipo": "PAGO Y ABONO CAPITAL",
            "trx": "STRESS cuota+capital 32",
            "tag": "pago_y_abono",
            "credito": d32["credito"],
            "digits": d32["digits"],
            "extra_capital": 1_500_000.0,
        },
        {
            "fecha": datetime(2026, 4, 12),
            "monto": float(g224["total_a_pagar"]) + 800_000.0,
            "concepto": g224["cliente"],
            "tipo": "PAGO Y ABONO CAPITAL",
            "trx": "STRESS cuota+capital 224",
            "tag": "pago_y_abono",
            "credito": g224["credito"],
            "digits": g224["digits"],
            "extra_capital": 800_000.0,
        },
        # --- ABONO CAPITAL ---
        {
            "fecha": datetime(2026, 7, 20),
            "monto": 500_000.0,
            "concepto": d92["cliente"],
            "tipo": "ABONO CAPITAL",
            "trx": "STRESS abono capital 92",
            "tag": "abono_capital",
            "credito": d92["credito"],
            "digits": d92["digits"],
        },
        {
            "fecha": datetime(2026, 7, 21),
            "monto": 350_000.0,
            "concepto": g284["cliente"],
            "tipo": "ABONO CAPITAL",
            "trx": "STRESS abono capital 284",
            "tag": "abono_capital",
            "credito": g284["credito"],
            "digits": g284["digits"],
        },
        {
            "fecha": datetime(2026, 7, 22),
            "monto": 250_000.0,
            "concepto": g296["cliente"],
            "tipo": "ABONO CAPITAL",
            "trx": "STRESS abono capital 296",
            "tag": "abono_capital",
            "credito": g296["credito"],
            "digits": g296["digits"],
        },
        {
            "fecha": datetime(2026, 7, 18),
            "monto": 400_000.0,
            "concepto": d94["cliente"],
            "tipo": "ABONO CAPITAL",
            "trx": "STRESS abono capital 94",
            "tag": "abono_capital",
            "credito": d94["credito"],
            "digits": d94["digits"],
        },
        # --- ABONO MORA ---
        {
            "fecha": datetime(2026, 7, 19),
            "monto": 180_000.0,
            "concepto": d32["cliente"],
            "tipo": "ABONO MORA",
            "trx": "STRESS abono mora 32",
            "tag": "abono_mora",
            "credito": d32["credito"],
            "digits": d32["digits"],
        },
        {
            "fecha": datetime(2026, 7, 23),
            "monto": 120_000.0,
            "concepto": g224["cliente"],
            "tipo": "ABONO MORA",
            "trx": "STRESS abono mora 224",
            "tag": "abono_mora",
            "credito": g224["credito"],
            "digits": g224["digits"],
        },
        {
            "fecha": datetime(2026, 7, 24),
            "monto": 90_000.0,
            "concepto": i150["cliente"],
            "tipo": "ABONO MORA",
            "trx": "STRESS abono mora 150",
            "tag": "abono_mora",
            "credito": i150["credito"],
            "digits": i150["digits"],
        },
        # --- Errores de usuario / negativos ---
        {
            "fecha": datetime(2026, 7, 15),
            "monto": 9_999_999.0,
            "concepto": "CLIENTE_FANTASMA_STRESS",
            "tipo": "PAGO",
            "trx": "STRESS cliente inexistente",
            "tag": "error_cliente",
            "credito": None,
            "digits": None,
        },
        {
            "fecha": datetime(2026, 4, 20),
            "monto": 20_469_925.01,
            "concepto": "ACIMOR",
            "tipo": "PAGO",
            "trx": "STRESS sin CREDITO tipico",
            "tag": "error_estructura",
            "credito": None,
            "digits": None,
        },
        {
            "fecha": datetime(2026, 5, 16),
            "monto": 5_000_000.0,
            "concepto": "INVERSIONES Y PROYECTOS MIOS",
            "tipo": "PAGO",
            "trx": "STRESS sin extracto util 318",
            "tag": "error_sin_extracto",
            "credito": None,
            "digits": None,
        },
        {
            "fecha": datetime(2026, 4, 25),
            "monto": 30_469_925.02,
            "concepto": "MINCIVIL",
            "tipo": "PAGO",
            "trx": "STRESS negativo mincivil",
            "tag": "error_estructura",
            "credito": None,
            "digits": None,
        },
        # --- Extra PAGO adelantado (25ª fila) ---
        {
            "fecha": datetime(2026, 4, 5),
            "monto": d92["total_a_pagar"],
            "concepto": d92["cliente"],
            "tipo": "PAGO",
            "trx": "STRESS PAGO segundo 92 ADELANTADO",
            "tag": "pago_adelantado_extra",
            "credito": d92["credito"],
            "digits": d92["digits"],
        },
    ]
    assert len(rows) >= 25, len(rows)
    return rows


def upload_bank(c: httpx.Client, drive: str, rows: list[dict]) -> None:
    raw = download(c, drive, BANK_ITEM_ID)
    wb = load_workbook(io.BytesIO(raw))
    ws = wb.active
    for r in range(4, ws.max_row + 1):
        for col in range(1, 6):
            ws.cell(r, col).value = None
    for i, row in enumerate(rows):
        rr = 4 + i
        ws.cell(rr, 1).value = row["fecha"]
        ws.cell(rr, 2).value = float(row["monto"])
        ws.cell(rr, 3).value = row["concepto"]
        ws.cell(rr, 4).value = row["tipo"]
        ws.cell(rr, 5).value = row["trx"]
    buf = io.BytesIO()
    wb.save(buf)
    payload = buf.getvalue()
    (WORK / "BANCO_BOGOTA_stress.xlsx").write_bytes(payload)
    up = c.put(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": BANK_ITEM_ID},
        json={"content_base64": base64.b64encode(payload).decode("ascii")},
        timeout=180,
    )
    up.raise_for_status()
    print("bank_uploaded", len(rows))


def clear_revision(c: httpx.Client, drive: str) -> None:
    rid = walk(c, drive, REV_DIR.split("/"))
    for it in list(children(c, drive, rid)):
        if str(it.get("name", "")).startswith("~$"):
            continue
        url = f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/items/{it['id']}"
        r = c.delete(url)
        print("delete_rev", it.get("name"), r.status_code)
        r.raise_for_status()


def draw_asiento(path: Path, data: dict) -> None:
    c = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    y = height - 20 * mm
    c.setFont("Helvetica-Bold", 14)
    c.drawString(20 * mm, y, data["voucher"])
    y -= 10 * mm
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm, y, f"{data['year']}  {data['month']}  {data['day']}")
    y -= 10 * mm
    c.drawString(20 * mm, y, f"N° Identificación : {data['nit']}")
    y -= 6 * mm
    c.drawString(20 * mm, y, f"Nombre : {data['nombre']}")
    y -= 12 * mm
    total = sum(v for _, v in data["lines"])
    for cuenta, valor in data["lines"]:
        line = f"{valor:,.2f} PAGO: No.Rad. {data['credit']} Linea 544 1 {cuenta}"
        c.drawString(20 * mm, y, line)
        y -= 6 * mm
    y -= 4 * mm
    c.setFont("Helvetica-Bold", 10)
    c.drawString(20 * mm, y, f"{total:,.2f}  {total:,.2f}")
    c.showPage()
    c.save()


def generate_asientos(rows: list[dict]) -> list[tuple[str, str, str, str]]:
    """Retorna lista (cliente, credito_folder, digits, filename) para subir."""
    ASIENTOS_DIR.mkdir(parents=True, exist_ok=True)
    uploads: list[tuple[str, str, str, str]] = []
    voucher = 9100
    for row in rows:
        if not row.get("digits") or row["concepto"] in SKIP_VALIDATE:
            continue
        tipo = row["tipo"]
        if tipo == "PAGO":
            label = "PAGO CUOTA"
        elif tipo == "PAGO Y ABONO CAPITAL":
            label = "PAGO Y ABONO CAPITAL"
        elif tipo == "ABONO CAPITAL":
            label = "ABONO CAPITAL"
        elif tipo == "ABONO MORA":
            label = "ABONO MORA"
        else:
            continue
        digits = row["digits"]
        fname = f"Asiento 28-JUL-2026 {label} {row['concepto'][:20]} CRED {digits}.pdf"
        fname = re.sub(r"[^\w\s\-#.]", "", fname).replace("  ", " ")
        monto = float(row["monto"])
        lines = [("11100505", monto), ("13410519", monto)]
        fe = row["fecha"]
        data = {
            "filename": fname,
            "voucher": str(voucher),
            "day": fe.day,
            "month": fe.month,
            "year": fe.year,
            "nit": f"900{digits.zfill(6)}-1",
            "nombre": row["concepto"],
            "credit": digits,
            "lines": lines,
        }
        voucher += 1
        dest = ASIENTOS_DIR / fname
        draw_asiento(dest, data)
        uploads.append((row["concepto"], row["credito"], digits, fname))
    print("asientos_generados", len(uploads))
    return uploads


def find_id_pago_header_row(ws) -> int:
    for r in range(1, 20):
        if str(ws.cell(r, 1).value or "").strip() == "ID Pago":
            return r
    raise ValueError("header_not_found")


def header_map(ws, hr: int) -> dict[str, int]:
    return {
        str(ws.cell(hr, c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(hr, c).value is not None and str(ws.cell(hr, c).value).strip()
    }


def fnum(v) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def prepare_review(
    c: httpx.Client,
    drive: str,
    rev_name: str,
    abono_targets: list[tuple[str, str]] | None = None,
    pago_targets: list[dict] | None = None,
) -> None:
    folder_id = walk(c, drive, REV_DIR.split("/"))
    rev = next(i for i in children(c, drive, folder_id) if i.get("name") == rev_name)
    raw = download(c, drive, rev["id"])
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Distribucion_Pagos"]
    hr = find_id_pago_header_row(ws)
    hm = header_map(ws, hr)
    col_id = hm["ID Pago"]
    col_cli = hm["Cliente"]
    col_cred = hm["Crédito"]
    col_monto = hm["Monto banco"]
    col_ext = hm["Valor extracto"]
    col_aplicar = hm["Aplicar a extracto"]
    col_mora = hm["Mora a aplicar"]
    col_cap = hm["Abono a capital"]
    col_otros = hm["Otros valores"]
    col_estado = hm["Estado Pago"]
    col_val = hm["Validar Pago"]
    col_obs = hm["Observación"]
    col_tipo = hm.get("TipoAplicacionOriginal")

    for r in range(hr + 1, ws.max_row + 1):
        if ws.cell(r, col_id).value:
            ws.cell(r, col_val).value = "NO"

    by_id: dict[str, list[int]] = {}
    validated = 0
    for r in range(hr + 1, ws.max_row + 1):
        pid = str(ws.cell(r, col_id).value or "").strip()
        if not pid:
            continue
        cli = str(ws.cell(r, col_cli).value or "").strip()
        if any(s.lower() in cli.lower() for s in SKIP_VALIDATE):
            if str(ws.cell(r, col_estado).value or "").upper() == "NORMAL":
                ws.cell(r, col_obs).value = "STRESS: no validar (caso negativo)"
            continue
        by_id.setdefault(pid, []).append(r)

    # targets de pago: preferir dígito planificado
    unused_pagos = [dict(t) for t in (pago_targets or [])]

    for pid, rows in by_id.items():
        rows_sorted = sorted(rows)
        # Sanear: Monto banco solo en la primera fila del grupo (regla finalize)
        first_monto = None
        for r in rows_sorted:
            m = fnum(ws.cell(r, col_monto).value)
            if m is not None and first_monto is None:
                first_monto = m
        for i, r in enumerate(rows_sorted):
            if i == 0 and first_monto is not None:
                ws.cell(r, col_monto).value = first_monto
            else:
                ws.cell(r, col_monto).value = None

        monto = first_monto
        cli = ""
        tipo_row = ""
        for r in rows_sorted:
            if not cli:
                cli = str(ws.cell(r, col_cli).value or "").strip()
            if col_tipo and not tipo_row:
                tipo_row = str(ws.cell(r, col_tipo).value or "").upper()
        if monto is None:
            continue

        # Emparejar target por cliente+tipo+monto aproximado
        want_digits = None
        chosen_t = None
        for i, t in enumerate(unused_pagos):
            if t["tipo"] not in ("PAGO", "PAGO Y ABONO CAPITAL"):
                continue
            if abs(float(t["monto"]) - monto) > 1.0:
                continue
            if not (
                _norm_cli(t["concepto"]).split()[0] in _norm_cli(cli)
                or _norm_cli(cli).split()[0] in _norm_cli(t["concepto"])
            ):
                continue
            want_digits = t["digits"]
            chosen_t = i
            break

        best_r = None
        if want_digits:
            for r in rows_sorted:
                cred = str(ws.cell(r, col_cred).value or "")
                digs = credit_digits(cred) or re.sub(r"\D", "", cred)
                if digs == want_digits:
                    best_r = r
                    break
        if best_r is None:
            best_score = None
            for r in rows_sorted:
                e = fnum(ws.cell(r, col_ext).value)
                if e is None:
                    continue
                score = abs(monto - e)
                if best_score is None or score < best_score:
                    best_score, best_r = score, r
        if best_r is None:
            best_r = rows_sorted[0]
        if chosen_t is not None:
            unused_pagos.pop(chosen_t)

        ws.cell(best_r, col_val).value = "SI"
        validated += 1
        ext = fnum(ws.cell(best_r, col_ext).value)
        tipo = str(ws.cell(best_r, col_tipo).value or "").upper() if col_tipo else tipo_row
        if tipo == "PAGO Y ABONO CAPITAL":
            if ext is not None and ext > 0 and monto > ext:
                cuota = float(ext)
                capital = round(monto - cuota, 2)
            else:
                capital = max(round(monto * 0.15, 2), 1_000.0)
                if capital >= monto:
                    capital = max(round(monto * 0.1, 2), 1.0)
                cuota = round(monto - capital, 2)
            ws.cell(best_r, col_aplicar).value = cuota
            ws.cell(best_r, col_mora).value = 0
            ws.cell(best_r, col_cap).value = capital
            ws.cell(best_r, col_otros).value = 0
            print("PAGO_Y_ABONO", pid, "cred", ws.cell(best_r, col_cred).value, "ext", ext, "cuota", cuota, "cap", capital)
        else:
            ws.cell(best_r, col_aplicar).value = monto
            ws.cell(best_r, col_mora).value = 0
            ws.cell(best_r, col_cap).value = 0
            ws.cell(best_r, col_otros).value = 0
            print("PAGO_SI", pid, "cred", ws.cell(best_r, col_cred).value, "monto", monto)
        for r in rows_sorted:
            if r == best_r:
                continue
            if str(ws.cell(r, col_estado).value or "").upper() == "NORMAL":
                ws.cell(r, col_obs).value = "STRESS: otro crédito mismo pago; no validar"

    # Abonos: un crédito SI por ID Pago según targets (cliente, digits)
    targets = list(abono_targets or [])
    if "Distribucion_Abonos" in wb.sheetnames:
        wa = wb["Distribucion_Abonos"]
        hr_a = find_id_pago_header_row(wa)
        hm_a = header_map(wa, hr_a)
        col_fecha = hm_a.get("Fecha banco")
        col_cred_a = hm_a["Crédito"]
        col_val_a = hm_a["Validar Abono"]
        col_cli_a = hm_a["Cliente"]
        col_monto_a = hm_a["Monto banco"]
        by_ab: dict[str, list[int]] = {}
        for r in range(hr_a + 1, wa.max_row + 1):
            pid = str(wa.cell(r, hm_a["ID Pago"]).value or "").strip()
            if not pid:
                continue
            wa.cell(r, col_val_a).value = "NO"
            cli = str(wa.cell(r, col_cli_a).value or "").strip()
            if any(s.lower() in cli.lower() for s in SKIP_VALIDATE):
                continue
            by_ab.setdefault(pid, []).append(r)

        unused_targets = list(targets)
        ab_ok = 0
        for pid, rows_a in by_ab.items():
            monto_ab = None
            fecha_ab = None
            cli_ab = ""
            for r in rows_a:
                if monto_ab is None:
                    monto_ab = fnum(wa.cell(r, col_monto_a).value)
                if col_fecha and fecha_ab is None and wa.cell(r, col_fecha).value not in (None, ""):
                    fecha_ab = wa.cell(r, col_fecha).value
                if not cli_ab:
                    cli_ab = str(wa.cell(r, col_cli_a).value or "").strip()
            best = None
            chosen_idx = None
            for i, (t_cli, t_dig) in enumerate(unused_targets):
                if not (
                    _norm_cli(t_cli).split()[0] in _norm_cli(cli_ab)
                    or _norm_cli(cli_ab).split()[0] in _norm_cli(t_cli)
                ):
                    continue
                for r in rows_a:
                    cred = str(wa.cell(r, col_cred_a).value or "")
                    digs = credit_digits(cred) or re.sub(r"\D", "", cred)
                    if digs == t_dig:
                        best = r
                        chosen_idx = i
                        break
                if best is not None:
                    break
            if best is None and rows_a:
                best = rows_a[0]
            if chosen_idx is not None:
                unused_targets.pop(chosen_idx)
            if best is None:
                continue
            # Monto/fecha solo en la fila SI (finalize lee esa fila)
            for r in rows_a:
                wa.cell(r, col_monto_a).value = None
                if col_fecha:
                    # conservar fecha solo en SI más abajo
                    pass
            if monto_ab is not None:
                wa.cell(best, col_monto_a).value = monto_ab
            if col_fecha:
                for r in rows_a:
                    if r != best:
                        wa.cell(r, col_fecha).value = None
                fv = fecha_ab
                if isinstance(fv, str) and fv.strip():
                    try:
                        fv = datetime.strptime(fv.strip()[:10], "%Y-%m-%d")
                    except ValueError:
                        pass
                elif isinstance(fv, date) and not isinstance(fv, datetime) and fv is not None:
                    fv = datetime(fv.year, fv.month, fv.day)
                if fv is not None:
                    wa.cell(best, col_fecha).value = fv
            wa.cell(best, col_val_a).value = "SI"
            ab_ok += 1
            print("abono SI", pid, wa.cell(best, col_cred_a).value, wa.cell(best, col_monto_a).value)
        print("abonos_SI", ab_ok)

    wc = wb["Control"]
    for r in range(1, wc.max_row + 1):
        if str(wc.cell(r, 1).value or "").strip().lower() == "procesar":
            wc.cell(r, 2).value = "SI"

    buf = io.BytesIO()
    wb.save(buf)
    out = buf.getvalue()
    (WORK / "review_prepared.xlsx").write_bytes(out)
    up = c.put(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": rev["id"]},
        json={"content_base64": base64.b64encode(out).decode("ascii")},
        timeout=180,
    )
    up.raise_for_status()
    print("review_prepared validated_pagos", validated)


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    winners = load_winners()
    rows = build_bank_rows(winners)
    plan = {
        "process_date": PROCESS_DATE,
        "rows": len(rows),
        "by_tag": {},
        "credits": sorted({r["digits"] for r in rows if r.get("digits")}),
        "clients": sorted({r["concepto"] for r in rows}),
    }
    for r in rows:
        plan["by_tag"][r["tag"]] = plan["by_tag"].get(r["tag"], 0) + 1
    log("00_plan", plan)

    asiento_uploads = generate_asientos(rows)
    log("00_asientos_map", asiento_uploads)

    with httpx.Client(timeout=180) as c:
        # Health
        h = c.get(f"{BASE}/health")
        log("01_health", {"status": h.status_code, "body": h.json()})

        drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]

        # Paso adelantado ANTES de generate (debe fallar control)
        r = c.post(
            f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs",
            json={"bank_code": BANK_CODE},
        )
        log("02_merge_too_early", {"status": r.status_code, "body": r.json() if r.content else {}})

        upload_bank(c, drive, rows)
        clear_revision(c, drive)

        # Generate
        r = c.post(
            f"{BASE}/graph/sharepoint/payment-validation/generate/queue",
            json={"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
        )
        log("03_generate_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        r.raise_for_status()
        gen = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{r.json()['job_id']}")
        log("03_generate_result", gen)
        if gen.get("status") != "completed":
            print("GENERATE FAILED")
            return

        # Idempotencia generate
        gen2: dict = {}
        r = c.post(
            f"{BASE}/graph/sharepoint/payment-validation/generate/queue",
            json={"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
        )
        log("04_generate2_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        if r.status_code == 202:
            gen2 = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{r.json()['job_id']}")
            log("04_generate2_result", gen2)

        result = gen.get("result") or {}
        rev_name = result.get("validation_file") or f"validacion_pagos_banco_bogota_{PROCESS_DATE}.xlsx"
        print("validation_file", rev_name, "pagos", (result.get("summary") or {}).get("pagos_banco"))

        # Notify demasiado pronto (antes de finalize)
        r = c.post(
            f"{BASE}/graph/sharepoint/notify-validar-extractos-email",
            json={"bank_code": BANK_CODE},
        )
        log("05_notify_too_early", {"status": r.status_code, "body": r.json() if r.content else {}})
        if r.status_code == 202:
            ne = poll(
                c,
                f"{BASE}/graph/sharepoint/notify-validar-extractos-email/jobs/{r.json()['job_id']}",
                timeout_s=300,
            )
            log("05_notify_too_early_result", ne)

        prepare_review(
            c,
            drive,
            rev_name,
            abono_targets=[
                (r["concepto"], r["digits"])
                for r in rows
                if r.get("digits") and r["tipo"] in ("ABONO CAPITAL", "ABONO MORA")
            ],
        )

        # Finalize
        r = c.post(
            f"{BASE}/graph/sharepoint/payment-validation/finalize/queue",
            json={"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
        )
        log("06_finalize_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        r.raise_for_status()
        fin = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{r.json()['job_id']}")
        log("06_finalize_result", fin)
        if fin.get("status") != "completed":
            print("FINALIZE FAILED")
            return
        hist = (fin.get("result") or {}).get("historical_file_path")

        # Notify
        r = c.post(
            f"{BASE}/graph/sharepoint/notify-validar-extractos-email",
            json={"bank_code": BANK_CODE, "historical_file_path": hist},
        )
        log("07_notify_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        r.raise_for_status()
        notify = poll(
            c,
            f"{BASE}/graph/sharepoint/notify-validar-extractos-email/jobs/{r.json()['job_id']}",
            timeout_s=900,
        )
        log("07_notify_result", notify)

        # Amort demasiado pronto
        r = c.post(
            f"{BASE}/graph/sharepoint/payment-validation/amortization/dry-run/queue",
            json={"bank_code": BANK_CODE},
        )
        log("08_amort_too_early", {"status": r.status_code, "body": r.json() if r.content else {}})
        if r.status_code == 202:
            ae = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{r.json()['job_id']}")
            log("08_amort_too_early_result", ae)

        # Merge SIN asientos (espera MERGE_PARCIAL)
        r = c.post(
            f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs",
            json={"bank_code": BANK_CODE},
        )
        log("10_merge_sin_asientos_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        if r.status_code == 202:
            m0 = poll(
                c,
                f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs/jobs/{r.json()['job_id']}",
                timeout_s=1200,
            )
            log("10_merge_sin_asientos_result", m0)

        # Subir asientos
        for cliente, credito_folder, digits, fname in asiento_uploads:
            local = ASIENTOS_DIR / fname
            if not local.exists() or not credito_folder:
                print("skip_asiento", fname)
                continue
            rel = f"{CLIENTS}/{cliente}/{credito_folder}/ASIENTOS CONTABLES CRED {digits}/{fname}"
            up = c.put(
                f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/path-content",
                params={"item_path": rel},
                json={"content_base64": base64.b64encode(local.read_bytes()).decode("ascii")},
                timeout=180,
            )
            print("upload_asiento", digits, fname, up.status_code, up.text[:100])

        # Merge con asientos
        r = c.post(
            f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs",
            json={"bank_code": BANK_CODE},
        )
        log("11_merge_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        r.raise_for_status()
        m1 = poll(
            c,
            f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs/jobs/{r.json()['job_id']}",
            timeout_s=1800,
        )
        log("11_merge_result", m1)

        # Merge idempotencia
        m2: dict = {}
        r = c.post(
            f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs",
            json={"bank_code": BANK_CODE},
        )
        log("12_merge2_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        if r.status_code == 202:
            m2 = poll(
                c,
                f"{BASE}/graph/sharepoint/merge-composite-validado-pdfs/jobs/{r.json()['job_id']}",
                timeout_s=900,
            )
            log("12_merge2_result", m2)

        # Dry-run + apply
        r = c.post(
            f"{BASE}/graph/sharepoint/payment-validation/amortization/dry-run/queue",
            json={"bank_code": BANK_CODE},
        )
        log("13_amort_dry_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        r.raise_for_status()
        dry = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{r.json()['job_id']}", timeout_s=1800)
        log("13_amort_dry_result", dry)

        r = c.post(
            f"{BASE}/graph/sharepoint/payment-validation/amortization/apply/queue",
            json={"bank_code": BANK_CODE},
        )
        log("14_amort_apply_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        r.raise_for_status()
        apply = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{r.json()['job_id']}", timeout_s=1800)
        log("14_amort_apply_result", apply)

        # Apply idempotencia
        apply2: dict = {}
        r = c.post(
            f"{BASE}/graph/sharepoint/payment-validation/amortization/apply/queue",
            json={"bank_code": BANK_CODE},
        )
        log("15_amort_apply2_queue", {"status": r.status_code, "body": r.json() if r.content else {}})
        if r.status_code == 202:
            apply2 = poll(
                c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{r.json()['job_id']}", timeout_s=900
            )
            log("15_amort_apply2_result", apply2)

        summary = {
            "generate_status": gen.get("status"),
            "generate_pagos": (result.get("summary") or {}).get("pagos_banco"),
            "generate_errores": (result.get("summary") or {}).get("errores"),
            "generate_already_2": (gen2.get("result") or {}).get("already_generated"),
            "finalize_status": fin.get("status"),
            "finalize_validated": (fin.get("result") or {}).get("validated_rows"),
            "notify_status": notify.get("status"),
            "merge_status": (m1.get("result") or {}).get("status") or m1.get("status"),
            "merge_already": (m2.get("result") or {}).get("already_merged"),
            "dry_can_apply": (dry.get("result") or {}).get("can_apply"),
            "dry_events": len((dry.get("result") or {}).get("events") or []),
            "apply_status": (apply.get("result") or {}).get("status") or apply.get("status"),
            "apply_tables": (apply.get("result") or {}).get("tables_uploaded_count"),
            "apply_already": (apply2.get("result") or {}).get("already_applied"),
            "plan_rows": len(rows),
            "plan_credits": plan["credits"],
            "plan_clients": plan["clients"],
        }
        log("99_summary", summary)
        print("DONE", json.dumps(summary, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
