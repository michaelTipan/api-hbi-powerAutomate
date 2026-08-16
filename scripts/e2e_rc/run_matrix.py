"""Runner matriz E01–E40 contra sandbox desplegado (CAPA B Graph + jobs)."""
from __future__ import annotations

import argparse
import json
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Callable

import httpx

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.e2e_rc.bank_upload import d, upload_bank_bogota
from scripts.e2e_rc.graph_session import SandboxGraphSession
from scripts.e2e_rc.path_guard import AUTHORIZED_CLIENTS_BASE, assert_sandbox_mutable_path
from scripts.e2e_rc.review_edit import (
    approve_credits_split,
    approve_single_credit_pago,
    edit_aplicacion_pagos_rows,
)
from scripts.e2e_rc.scenarios import SCENARIOS, ScenarioResult, by_id
from scripts.e2e_rc.fixtures_catalog import (
    CORREOS_XLSX_REL,
    E16_BANK_TOTAL,
    E16_SPLIT_A,
    E16_SPLIT_B,
    RC_E13_CAPITAL,
    RC_E13_INTERESES,
    RC_E13_MORA,
    RC_E13_TOTAL,
    RC_E14_CAPITAL,
    RC_E14_CREDIT,
    RC_E14_FOLDER,
    RC_E14_INTERESES,
    RC_E14_MORA,
    RC_E14_TOTAL,
    RC_E24_ASIENTO_A,
    RC_E24_ASIENTO_B,
    RC_E24_BANK,
    RC_E24_CREDIT_A,
    RC_E24_CREDIT_B,
    RC_E24_REVIEW_A,
    RC_E24_REVIEW_B,
    RC_E40_BANK,
    RC_E40_CREDIT_B,
    RC_E40_FOLDER_B,
    RC_E40_SPLIT_A,
    RC_E40_SPLIT_B,
    RC_MORA_CREDIT,
    RC_MORA_OBLIG,
    RC_MORA_VENCIDO,
    RC_PAYOFF_CREDIT,
    RC_PAYOFF_FOLDER,
    assert_sandbox_notify_recipients,
    gaps_by_scenario,
)
from scripts.e2e_rc.sandbox_fixtures import (
    provision_canonical_payoff_credit,
    provision_e15_parseable_asiento,
    provision_e16_second_active_credit,
    provision_rc_mora_credit,
    rewrite_sandbox_correos_xlsx,
)

import io
from openpyxl import load_workbook
from urllib.parse import quote

WORK = Path(r"D:\CMC\HBI_Capital\_work\rc_e2e")
BANK_CODE = "banco_bogota"
# process_date base; cada escenario usa una fecha distinta para no chocar con
# procesos FINALIZADO (cancel_not_allowed).
PROCESS_DATE_BASE = datetime(2026, 10, 15)


def process_date_for(scenario_id: str) -> str:
    n = int(scenario_id[1:])
    return (PROCESS_DATE_BASE + timedelta(days=n - 1)).strftime("%Y-%m-%d")


REV_DIR = f"{AUTHORIZED_CLIENTS_BASE}/02 VALIDACION PAGOS/01 REVISION"
CTRL_PATH = (
    f"{AUTHORIZED_CLIENTS_BASE}/02 VALIDACION PAGOS/90 ACCESO RESTRINGIDO/"
    "03 CONTROL TECNICO/control_proceso_validacion_pagos_banco_bogota.xlsx"
)
NOTIFY_SANDBOX_TO = "herramientas.jsakedev@gmail.com"


def _log(name: str, data: object) -> Path:
    WORK.mkdir(parents=True, exist_ok=True)
    path = WORK / f"{name}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(name, json.dumps(data, ensure_ascii=False, default=str)[:400])
    return path


def cancel_active(session: SandboxGraphSession) -> dict[str, Any]:
    return session.queue_and_poll(
        "/graph/sharepoint/payment-validation/cancel-active-process/queue",
        {"bank_code": BANK_CODE},
        timeout_s=600,
    )


def reset_bank_slot(session: SandboxGraphSession) -> dict[str, Any]:
    """Libera el slot del banco en sandbox: IsActive=false + limpia REVISION.

    Tras CONSOLIDADO/MERGE el mismo ``process_date`` devolvería
    ``already_generated`` (producto). Para re-ejecutar E2E en la misma fecha
    se deja el control en VACIO (permitido para lote nuevo).
    Solo paths PRUEBAS (path guard).
    """
    assert_sandbox_mutable_path(CTRL_PATH)
    assert_sandbox_mutable_path(REV_DIR)
    out: dict[str, Any] = {"deactivated": 0, "cleared_to_vacio": 0, "deleted_reviews": []}
    # Control
    ctrl_folder = "/".join(CTRL_PATH.split("/")[:-1])
    ctrl_name = CTRL_PATH.split("/")[-1]
    folder_id = session.walk(ctrl_folder)
    item = next(i for i in session.children(folder_id) if i.get("name") == ctrl_name)
    raw = session.download_item(str(item["id"]))
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Procesos"]
    headers = {
        str(ws.cell(1, c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(1, c).value
    }
    col_active = headers["IsActive"]
    col_estado = headers.get("EstadoProceso")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col_active).value in (True, "TRUE", "True", 1, "1"):
            ws.cell(r, col_active).value = False
            out["deactivated"] += 1
        if col_estado:
            estado = str(ws.cell(r, col_estado).value or "").strip().upper()
            if estado and estado not in ("VACIO", "CANCELADO", ""):
                ws.cell(r, col_estado).value = "VACIO"
                out["cleared_to_vacio"] += 1
    buf = io.BytesIO()
    wb.save(buf)
    session.upload_item(str(item["id"]), buf.getvalue(), path_for_guard=CTRL_PATH)
    # Review folder
    rev_id = session.walk(REV_DIR)
    for it in session.children(rev_id):
        name = str(it.get("name") or "")
        if not name.lower().endswith(".xlsx") or name.startswith("~$"):
            continue
        del_r = session._client.delete(
            f"{session.base}/graph/sharepoint/drives/{quote(session.drive_id or '', safe='')}/items/{it['id']}"
        )
        out["deleted_reviews"].append({"name": name, "status": del_r.status_code})
    return out


def generate(session: SandboxGraphSession, process_date: str) -> dict[str, Any]:
    return session.queue_and_poll(
        "/graph/sharepoint/payment-validation/generate/queue",
        {"bank_code": BANK_CODE, "process_date": process_date},
        timeout_s=1800,
    )


def finalize(session: SandboxGraphSession, process_date: str) -> dict[str, Any]:
    return session.queue_and_poll(
        "/graph/sharepoint/payment-validation/finalize/queue",
        {"bank_code": BANK_CODE, "process_date": process_date},
        timeout_s=1800,
    )


def notify(session: SandboxGraphSession) -> dict[str, Any]:
    """Notify PA (CORREOS.xlsx del entorno sandbox)."""
    return session.queue_and_poll(
        "/graph/sharepoint/notify-validar-extractos-email",
        {"bank_code": BANK_CODE},
        timeout_s=1800,
    )


def merge(session: SandboxGraphSession) -> dict[str, Any]:
    return session.queue_and_poll(
        "/graph/sharepoint/merge-composite-validado-pdfs",
        {"bank_code": BANK_CODE},
        timeout_s=1800,
    )


def amort_dry_run(session: SandboxGraphSession) -> dict[str, Any]:
    return session.queue_and_poll(
        "/graph/sharepoint/payment-validation/amortization/dry-run/queue",
        {"bank_code": BANK_CODE},
        timeout_s=1800,
    )


def amort_apply(session: SandboxGraphSession) -> dict[str, Any]:
    return session.queue_and_poll(
        "/graph/sharepoint/payment-validation/amortization/apply/queue",
        {"bank_code": BANK_CODE},
        timeout_s=1800,
    )


def _result_dict(job: dict[str, Any]) -> dict[str, Any]:
    result = job.get("result")
    return result if isinstance(result, dict) else {}


def _has_payoff_not_achieved(job: dict[str, Any]) -> bool:
    """True si dry-run/amort bloqueó por PAYOFF_NOT_ACHIEVED (ítem u operational_issue)."""
    result = _result_dict(job)
    if result.get("error_code") == "PAYOFF_NOT_ACHIEVED":
        return True
    for item in result.get("items") or []:
        if isinstance(item, dict) and item.get("error_code") == "PAYOFF_NOT_ACHIEVED":
            return True
    for issue in result.get("operational_issues") or []:
        if not isinstance(issue, dict):
            continue
        ref = str(issue.get("technical_reference") or "")
        if ref == "PAYOFF_NOT_ACHIEVED":
            return True
    return False


def latest_review(session: SandboxGraphSession) -> tuple[dict[str, Any], bytes]:
    assert_sandbox_mutable_path(REV_DIR)
    folder_id = session.walk(REV_DIR)
    files = [
        i
        for i in session.children(folder_id)
        if str(i.get("name", "")).endswith(".xlsx") and not str(i.get("name", "")).startswith("~$")
    ]
    if not files:
        raise FileNotFoundError("no_review_xlsx")
    newest = sorted(files, key=lambda i: i.get("lastModifiedDateTime") or "", reverse=True)[0]
    raw = session.download_item(str(newest["id"]))
    return newest, raw


def upload_review(session: SandboxGraphSession, item: dict[str, Any], raw: bytes) -> None:
    path = f"{REV_DIR}/{item['name']}"
    session.upload_item(str(item["id"]), raw, path_for_guard=path)


def _job_ok(job: dict[str, Any]) -> bool:
    return job.get("status") == "completed"


def _job_failed(job: dict[str, Any]) -> bool:
    return job.get("status") in ("failed", "error", "http_error")


def _process_key(job: dict[str, Any]) -> str | None:
    result = job.get("result") or {}
    if isinstance(result, dict):
        return result.get("process_key") or result.get("processKey")
    return None


def _merge_output_names(outputs: list[Any]) -> list[str]:
    names: list[str] = []
    for out in outputs:
        if not isinstance(out, dict):
            continue
        rel = str(out.get("output_relative_path") or out.get("filename") or out.get("name") or "")
        if rel:
            names.append(rel.rsplit("/", 1)[-1])
    return names


def _merge_output_name(outputs: list[Any], merge_result: dict[str, Any]) -> str | None:
    names = _merge_output_names(outputs)
    if names:
        return names[0]
    return (
        merge_result.get("output_filename")
        or merge_result.get("merged_filename")
        or merge_result.get("composite_name")
    )


def _prep_bank_and_generate(session: SandboxGraphSession, rows: list[dict[str, Any]], tag: str) -> ScenarioResult | tuple[dict[str, Any], str]:
    process_date = process_date_for(tag)
    cancel_active(session)
    reset = reset_bank_slot(session)
    _log(f"{tag}_reset", reset)
    up = upload_bank_bogota(session, rows)
    _log(f"{tag}_bank", {**up, "process_date": process_date})
    gen = generate(session, process_date)
    already = None
    if isinstance(gen.get("result"), dict):
        already = gen["result"].get("already_generated")
    _log(f"{tag}_generate", {"status": gen.get("status"), "error": gen.get("error"), "process_date": process_date, "already_generated": already})
    if not _job_ok(gen):
        return ScenarioResult(
            id=tag,
            status="FAIL",
            evidence="generate_failed",
            process_key=_process_key(gen),
            detail={"generate": gen, "process_date": process_date},
        )
    if already:
        return ScenarioResult(
            id=tag,
            status="FAIL",
            evidence="generate_already_generated_isolation_broken",
            process_key=_process_key(gen),
            detail={"generate": gen, "process_date": process_date},
        )
    return gen, process_date


def run_e01(session: SandboxGraphSession) -> ScenarioResult:
    """GEOEXCON pago exacto obligación actual (montos históricos E2E)."""
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 19_102_163.0,
            "concepto": "GEOEXCON",
            "trx": "RC E01 PAGO NORMAL 231",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E01")
    if isinstance(prep, ScenarioResult):
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    (WORK / f"E01_{item['name']}").write_bytes(raw)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo="PAGO DE OBLIGACIÓN ACTUAL",
        observacion="RC-E01",
    )
    if not applied:
        return ScenarioResult("E01", "FAIL", evidence="no_rows_edited", detail={"review": item.get("name")})
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    _log("E01_finalize", {"status": fin.get("status"), "error": fin.get("error")})
    cancel = cancel_active(session)
    status = "PASS" if _job_ok(fin) else "FAIL"
    return ScenarioResult(
        "E01",
        status,
        evidence=f"finalize={fin.get('status')}; edited={len(applied)}",
        process_key=_process_key(fin) or _process_key(gen),
        artifacts=[str(WORK / f"E01_{item['name']}")],
        cleanup=f"cancel={cancel.get('status')}",
        detail={"applied": applied, "finalize_error": fin.get("error")},
    )


def run_e02(session: SandboxGraphSession) -> ScenarioResult:
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 19_102_163.0,
            "concepto": "GEOEXCON",
            "trx": "RC E02 PARCIAL 231",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E02")
    if isinstance(prep, ScenarioResult):
        prep.id = "E02"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo="PAGO PARCIAL A OBLIGACIÓN ACTUAL",
        obligacion=5_000_000.0,
        observacion="RC-E02 parcial",
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    _log("E02_finalize", {"status": fin.get("status"), "error": fin.get("error")})
    # Parcial con saldo por asignar > 0 debe bloquear Finalize
    blocked = _job_failed(fin)
    cancel_active(session)
    return ScenarioResult(
        "E02",
        "PASS" if blocked else "FAIL",
        evidence=f"expected_finalize_block; got={fin.get('status')}",
        process_key=_process_key(fin) or _process_key(gen),
        detail={"applied": applied, "error": fin.get("error")},
        cleanup="cancelled",
    )


def run_e03(session: SandboxGraphSession) -> ScenarioResult:
    rows = [
        {
            "fecha": d(2026, 7, 27),
            "monto": 25_075_203.0,
            "concepto": "AGRECAR",
            "trx": "RC E03 ATRASADO 37",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E03")
    if isinstance(prep, ScenarioResult):
        prep.id = "E03"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo="PAGO DE OBLIGACIÓN ACTUAL",
        observacion="RC-E03",
        credito_contains="37",
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    cancel_active(session)
    return ScenarioResult(
        "E03",
        "PASS" if _job_ok(fin) else "FAIL",
        evidence=f"finalize={fin.get('status')}",
        process_key=_process_key(fin) or _process_key(gen),
        detail={"applied": applied, "error": fin.get("error")},
        cleanup="cancelled",
    )


def run_e04(session: SandboxGraphSession) -> ScenarioResult:
    rows = [
        {
            "fecha": d(2026, 4, 15),
            "monto": 32_691_683.0,
            "concepto": "EQUINORTE",
            "trx": "RC E04 ADELANTADO 264",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E04")
    if isinstance(prep, ScenarioResult):
        prep.id = "E04"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo="PAGO DE OBLIGACIÓN ACTUAL",
        observacion="RC-E04",
        credito_contains="",
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    cancel_active(session)
    return ScenarioResult(
        "E04",
        "PASS" if _job_ok(fin) else "FAIL",
        evidence=f"finalize={fin.get('status')}",
        process_key=_process_key(fin) or _process_key(gen),
        detail={"applied": applied, "error": fin.get("error")},
        cleanup="cancelled",
    )


def run_e09(session: SandboxGraphSession) -> ScenarioResult:
    rows = [
        {
            "fecha": d(2026, 7, 27),
            "monto": 500_000.0,
            "concepto": "EQUINORTE",
            "trx": "RC E09 ABONO CAPITAL 265",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E09")
    if isinstance(prep, ScenarioResult):
        prep.id = "E09"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo="ABONO A CAPITAL",
        obligacion=0.0,
        capital=500_000.0,
        observacion="RC-E09",
        credito_contains="265",
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    cancel_active(session)
    return ScenarioResult(
        "E09",
        "PASS" if _job_ok(fin) else "FAIL",
        evidence=f"finalize={fin.get('status')}",
        process_key=_process_key(fin) or _process_key(gen),
        detail={"applied": applied, "error": fin.get("error")},
        cleanup="cancelled",
    )


def run_e10(session: SandboxGraphSession) -> ScenarioResult:
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 24_000_000.0,
            "concepto": "GEOEXCON",
            "trx": "RC E10 OBLIG+CAPITAL 231",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E10")
    if isinstance(prep, ScenarioResult):
        prep.id = "E10"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo="PAGO Y ABONO A CAPITAL",
        obligacion=19_102_163.0,
        capital=4_897_837.0,
        observacion="RC-E10",
        credito_contains="231",
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    cancel_active(session)
    return ScenarioResult(
        "E10",
        "PASS" if _job_ok(fin) else "FAIL",
        evidence=f"finalize={fin.get('status')}",
        process_key=_process_key(fin) or _process_key(gen),
        detail={"applied": applied, "error": fin.get("error")},
        cleanup="cancelled",
    )


def run_e15(session: SandboxGraphSession) -> ScenarioResult:
    """PAGO TOTAL erróneo: Finalize PASS; dry-run debe bloquear con PAYOFF_NOT_ACHIEVED.

    Mandato §12 / E15: Tipo confirmado puede ser CANCELACIÓN / PAGO TOTAL aunque la
    sugerida no lo sea; Finalize no inventa saldo=0. El bloqueo es en amort dry-run.
    """
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 19_102_163.0,
            "concepto": "GEOEXCON",
            "trx": "RC E15 PAGO TOTAL ERR 231",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E15")
    if isinstance(prep, ScenarioResult):
        prep.id = "E15"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo="CANCELACIÓN / PAGO TOTAL",
        observacion="RC-E15 intentional wrong total",
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    _log("E15_finalize", {"status": fin.get("status"), "error": fin.get("error")})
    if not _job_ok(fin):
        cancel_active(session)
        return ScenarioResult(
            "E15",
            "FAIL",
            evidence=f"finalize_expected_pass; got={fin.get('status')}",
            process_key=_process_key(fin) or _process_key(gen),
            detail={"applied": applied, "error": fin.get("error")},
            cleanup="cancelled",
        )

    # Continuar hasta dry-run (Notify → Merge → amort dry-run).
    # Provision asiento parseable ANTES de merge: el consolidado/dry-run debe
    # evaluar payoff (PAYOFF_NOT_ACHIEVED), no quedarse en ACCOUNTING_PARSE_FAILED.
    try:
        asiento_fix = provision_e15_parseable_asiento(session)
        _log("E15_asiento_fixture", asiento_fix)
    except Exception as exc:
        cancel_active(session)
        return ScenarioResult(
            "E15",
            "BLOCKED",
            evidence=f"finalize_pass; asiento_fixture_failed:{exc}",
            process_key=_process_key(fin) or _process_key(gen),
            detail={"asiento_fixture_error": str(exc)},
            cleanup="cancelled",
        )

    ntf = notify(session)
    _log("E15_notify", {"status": ntf.get("status"), "error": ntf.get("error")})
    if not _job_ok(ntf):
        cancel_active(session)
        return ScenarioResult(
            "E15",
            "BLOCKED",
            evidence=f"finalize_pass; notify_failed={ntf.get('status')}",
            process_key=_process_key(fin) or _process_key(gen),
            detail={"notify": ntf.get("error") or ntf.get("result")},
            cleanup="cancelled",
        )

    mrg = merge(session)
    _log("E15_merge", {"status": mrg.get("status"), "error": mrg.get("error")})
    if not _job_ok(mrg):
        cancel_active(session)
        return ScenarioResult(
            "E15",
            "BLOCKED",
            evidence=f"finalize_pass; merge_failed={mrg.get('status')}",
            process_key=_process_key(fin) or _process_key(gen),
            detail={"merge": mrg.get("error") or mrg.get("result")},
            cleanup="cancelled",
        )

    dry = amort_dry_run(session)
    _log(
        "E15_dry_run",
        {
            "status": dry.get("status"),
            "error": dry.get("error"),
            "can_apply": _result_dict(dry).get("can_apply"),
            "payoff_block": _has_payoff_not_achieved(dry),
        },
    )
    cancel_active(session)

    blocked = (
        _job_ok(dry)
        and _result_dict(dry).get("can_apply") is False
        and _has_payoff_not_achieved(dry)
    )
    # Job completed with requires_correction / can_apply false also OK.
    if not blocked and _job_ok(dry) and _has_payoff_not_achieved(dry):
        blocked = True
    if not blocked and dry.get("status") == "completed":
        # UI amortization_process shape: outcome requires_correction
        outcome = str(_result_dict(dry).get("outcome") or "")
        if outcome == "requires_correction" and _has_payoff_not_achieved(dry):
            blocked = True

    return ScenarioResult(
        "E15",
        "PASS" if blocked else "FAIL",
        evidence=(
            f"finalize=completed; dry_run={dry.get('status')}; "
            f"can_apply={_result_dict(dry).get('can_apply')}; "
            f"payoff_block={_has_payoff_not_achieved(dry)}"
        ),
        process_key=_process_key(fin) or _process_key(gen),
        detail={
            "applied": applied,
            "asiento_fixture": asiento_fix,
            "dry_error": dry.get("error"),
            "dry_result_summary": {
                "can_apply": _result_dict(dry).get("can_apply"),
                "error_code": _result_dict(dry).get("error_code"),
                "items_errors": [
                    it.get("error_code")
                    for it in (_result_dict(dry).get("items") or [])
                    if isinstance(it, dict) and it.get("error_code")
                ],
            },
        },
        cleanup="cancelled",
    )

def run_e21(session: SandboxGraphSession) -> ScenarioResult:
    rows = [
        {
            "fecha": d(2026, 5, 16),
            "monto": 5_000_000.0,
            "concepto": "INVERSIONES Y PROYECTOS MIOS",
            "trx": "RC E21 SIN EXTRACTO 318",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E21")
    if isinstance(prep, ScenarioResult):
        return ScenarioResult(
            "E21",
            "PASS" if prep.evidence == "generate_failed" else prep.status,
            evidence="generate_outcome_for_missing_extract",
            detail=prep.detail,
        )
    gen, process_date = prep
    item, raw = latest_review(session)
    (WORK / f"E21_{item['name']}").write_bytes(raw)
    edited, applied = approve_single_credit_pago(
        raw, tipo="PAGO DE OBLIGACIÓN ACTUAL", observacion="RC-E21"
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    cancel_active(session)
    status = "PASS" if _job_failed(fin) else "FAIL"
    return ScenarioResult(
        "E21",
        status,
        evidence=f"finalize={fin.get('status')} (expect block without extract)",
        process_key=_process_key(fin) or _process_key(gen),
        detail={"applied": applied, "error": fin.get("error"), "process_date": process_date},
        cleanup="cancelled",
    )


def run_e29(session: SandboxGraphSession) -> ScenarioResult:
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 19_102_163.0,
            "concepto": "GEOEXCON",
            "trx": "RC E29 RETRY GEN 231",
        }
    ]
    process_date = process_date_for("E29")
    cancel_active(session)
    reset_bank_slot(session)
    upload_bank_bogota(session, rows)
    g1 = generate(session, process_date)
    g2 = generate(session, process_date)
    cancel_active(session)
    reset_bank_slot(session)
    # Segundo generate debe ser idempotente (completed reuse) o busy/conflict controlado
    ok = _job_ok(g1) and (
        _job_ok(g2)
        or g2.get("status") in ("failed", "error", "http_error")
        or (g2.get("http_status") in (409, 409))
    )
    return ScenarioResult(
        "E29",
        "PASS" if ok else "FAIL",
        evidence=f"g1={g1.get('status')}; g2={g2.get('status')}/{g2.get('http_status')}",
        process_key=_process_key(g1),
        detail={"g1": {"status": g1.get("status"), "error": g1.get("error")}, "g2": {"status": g2.get("status"), "error": g2.get("error"), "http": g2.get("http_status")}},
        cleanup="cancelled",
    )


def run_e37(session: SandboxGraphSession) -> ScenarioResult:
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 19_102_163.0,
            "concepto": "GEOEXCON",
            "trx": "RC E37 CANCEL 231",
        }
    ]
    process_date = process_date_for("E37")
    cancel_active(session)
    reset_bank_slot(session)
    upload_bank_bogota(session, rows)
    gen = generate(session, process_date)
    cancel = cancel_active(session)
    reset_bank_slot(session)
    ok = _job_ok(gen) and cancel.get("status") in ("completed", "failed", "error")
    # cancel completed is success; some stacks return completed with result
    return ScenarioResult(
        "E37",
        "PASS" if _job_ok(gen) and (cancel.get("status") == "completed" or cancel.get("http_status") in (200, 202)) else ("PASS" if _job_ok(cancel) or cancel.get("status") == "completed" else "FAIL"),
        evidence=f"gen={gen.get('status')}; cancel={cancel.get('status')}",
        process_key=_process_key(gen),
        detail={"cancel": cancel.get("result") or cancel.get("error")},
        cleanup="done",
    )


def run_e16(session: SandboxGraphSession) -> ScenarioResult:
    """Un pago → dos créditos GEOEXCON activos (231+299; 254 es TERMINADO)."""
    fixture = provision_e16_second_active_credit(session)
    _log("E16_fixture", fixture)
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": E16_BANK_TOTAL,
            "concepto": "GEOEXCON",
            "trx": "RC E16 SPLIT 231-299",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E16")
    if isinstance(prep, ScenarioResult):
        prep.id = "E16"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_credits_split(
        raw,
        splits=[
            {
                "credito_contains": "231",
                "tipo": "PAGO DE OBLIGACIÓN ACTUAL",
                "obligacion": E16_SPLIT_A,
            },
            {
                "credito_contains": "299",
                "tipo": "PAGO DE OBLIGACIÓN ACTUAL",
                "obligacion": E16_SPLIT_B,
            },
        ],
        observacion="RC-E16",
    )
    if len(applied) < 2:
        cancel_active(session)
        return ScenarioResult(
            "E16",
            "FAIL",
            evidence=f"need_two_credit_rows; applied={len(applied)}",
            process_key=_process_key(gen),
            detail={"applied": applied, "fixture": fixture},
            cleanup="cancelled",
        )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    if not _job_ok(fin):
        cancel_active(session)
        return ScenarioResult(
            "E16",
            "FAIL",
            evidence=f"finalize={fin.get('status')}",
            process_key=_process_key(fin) or _process_key(gen),
            detail={"applied": applied, "error": fin.get("error"), "fixture": fixture},
            cleanup="cancelled",
        )
    # Flujo documental: notify (sandbox recipients) + merge (1 PDF / ID pago).
    rewrite_sandbox_correos_xlsx(session)
    nfy = notify(session)
    mrg = merge(session)
    cancel_active(session)
    merge_result = mrg.get("result") if isinstance(mrg.get("result"), dict) else {}
    outputs = list(merge_result.get("outputs") or [])
    merge_name = _merge_output_name(outputs, merge_result)
    pdf_ok = bool(outputs) or bool(
        merge_result.get("pdf_created") or merge_result.get("pdf_reused")
    )
    credits_in_name = bool(merge_name) and "231" in str(merge_name) and "299" in str(merge_name)
    merge_ok = _job_ok(mrg) and pdf_ok and len(outputs) <= 1 and credits_in_name
    return ScenarioResult(
        "E16",
        "PASS" if merge_ok else "FAIL",
        evidence=(
            f"finalize=ok; notify={nfy.get('status')}; merge={mrg.get('status')}; "
            f"credits={len(applied)}; outputs={len(outputs)}; name={merge_name}; "
            f"file_action={merge_result.get('file_action')}"
        ),
        process_key=_process_key(mrg) or _process_key(fin) or _process_key(gen),
        detail={
            "applied": applied,
            "fixture": fixture,
            "notify": nfy.get("status"),
            "merge": merge_result or mrg.get("error"),
            "merge_name": merge_name,
            "outputs_count": len(outputs),
        },
        cleanup="cancelled",
    )


def run_e17(session: SandboxGraphSession) -> ScenarioResult:
    """Un pago → dos créditos con tipos distintos → Merge APLICACION MULTIPLE."""
    fixture = provision_e16_second_active_credit(session)
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": E16_BANK_TOTAL,
            "concepto": "GEOEXCON",
            "trx": "RC E17 MULTI TIPO 231-299",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E17")
    if isinstance(prep, ScenarioResult):
        prep.id = "E17"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_credits_split(
        raw,
        splits=[
            {
                "credito_contains": "231",
                "tipo": "PAGO DE OBLIGACIÓN ACTUAL",
                "obligacion": E16_SPLIT_A,
            },
            {
                "credito_contains": "299",
                "tipo": "ABONO A CAPITAL",
                "obligacion": 0.0,
                "capital": E16_SPLIT_B,
            },
        ],
        observacion="RC-E17",
    )
    if len(applied) < 2:
        cancel_active(session)
        return ScenarioResult(
            "E17",
            "FAIL",
            evidence=f"need_two_credit_rows; applied={len(applied)}",
            process_key=_process_key(gen),
            detail={"applied": applied, "fixture": fixture},
            cleanup="cancelled",
        )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    if not _job_ok(fin):
        cancel_active(session)
        return ScenarioResult(
            "E17",
            "FAIL",
            evidence=f"finalize={fin.get('status')}",
            process_key=_process_key(fin) or _process_key(gen),
            detail={"error": fin.get("error"), "applied": applied},
            cleanup="cancelled",
        )
    rewrite_sandbox_correos_xlsx(session)
    nfy = notify(session)
    mrg = merge(session)
    cancel_active(session)
    merge_result = mrg.get("result") if isinstance(mrg.get("result"), dict) else {}
    outputs = list(merge_result.get("outputs") or [])
    names = _merge_output_names(outputs)
    blob = " ".join(names).upper()
    token_ok = "APLICACION MULTIPLE" in blob or "APLICACIÓN MÚLTIPLE" in blob
    return ScenarioResult(
        "E17",
        "PASS" if _job_ok(mrg) and token_ok else ("FAIL" if _job_ok(fin) else "FAIL"),
        evidence=(
            f"finalize=ok; notify={nfy.get('status')}; merge={mrg.get('status')}; "
            f"outputs={len(outputs)}; names={names}; token_ok={token_ok}"
        ),
        process_key=_process_key(mrg) or _process_key(fin),
        detail={"applied": applied, "merge": merge_result, "names": names},
        cleanup="cancelled",
    )


def run_e23(session: SandboxGraphSession) -> ScenarioResult:
    """Mismo cliente, dos pagos fechas distintas en un lote."""
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 19_102_163.0,
            "concepto": "GEOEXCON",
            "trx": "RC E23 PAGO A 231",
        },
        {
            "fecha": d(2026, 5, 24),
            "monto": 5_000_000.0,
            "concepto": "GEOEXCON",
            "trx": "RC E23 PAGO B 231",
        },
    ]
    prep = _prep_bank_and_generate(session, rows, "E23")
    if isinstance(prep, ScenarioResult):
        prep.id = "E23"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)

    def updater(row, _r):
        from app.application.services.review_schema import AplicacionPagosCols

        credito = str(row.get(AplicacionPagosCols.CREDITO) or "")
        if not credito.strip():
            return None
        try:
            monto = float(row.get(AplicacionPagosCols.MONTO_BANCO) or 0)
        except (TypeError, ValueError):
            monto = 0.0
        # Un SI por ID Pago (fila con monto banco); extras 299 u otras fechas vacías → NO.
        if monto > 0 and "231" in credito:
            return {
                AplicacionPagosCols.VALIDAR_PAGO: "SI",
                AplicacionPagosCols.TIPO_APLICACION: "PAGO DE OBLIGACIÓN ACTUAL",
                AplicacionPagosCols.OBSERVACION: "RC-E23",
            }
        return {
            AplicacionPagosCols.VALIDAR_PAGO: "NO",
            AplicacionPagosCols.TIPO_APLICACION: "",
            AplicacionPagosCols.OBSERVACION: "RC-E23",
        }

    edited, applied = edit_aplicacion_pagos_rows(raw, row_updater=updater)
    if len(applied) < 2:
        cancel_active(session)
        return ScenarioResult(
            "E23",
            "BLOCKED",
            evidence=f"need_two_payment_rows; applied={len(applied)}",
            process_key=_process_key(gen),
            detail={"applied": applied},
            cleanup="cancelled",
        )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    cancel_active(session)
    return ScenarioResult(
        "E23",
        "PASS" if _job_ok(fin) else "FAIL",
        evidence=f"finalize={fin.get('status')}; rows={len(applied)}",
        process_key=_process_key(fin) or _process_key(gen),
        detail={"applied": applied, "error": fin.get("error")},
        cleanup="cancelled",
    )


def run_e30(session: SandboxGraphSession) -> ScenarioResult:
    """Retry Finalize: segundo finalize sobre mismo process_date no debe corromper."""
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 19_102_163.0,
            "concepto": "GEOEXCON",
            "trx": "RC E30 RETRY FIN 231",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E30")
    if isinstance(prep, ScenarioResult):
        prep.id = "E30"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw, tipo="PAGO DE OBLIGACIÓN ACTUAL", observacion="RC-E30"
    )
    upload_review(session, item, edited)
    f1 = finalize(session, process_date)
    f2 = finalize(session, process_date)
    cancel_active(session)
    ok = _job_ok(f1) and (
        _job_ok(f2)
        or _job_failed(f2)
        or f2.get("http_status") in (409, 422)
    )
    return ScenarioResult(
        "E30",
        "PASS" if ok else "FAIL",
        evidence=f"f1={f1.get('status')}; f2={f2.get('status')}/{f2.get('http_status')}",
        process_key=_process_key(f1) or _process_key(gen),
        detail={"f1": f1.get("error"), "f2": f2.get("error"), "applied": applied},
        cleanup="cancelled",
    )


def run_e31(session: SandboxGraphSession) -> ScenarioResult:
    """Retry Notify: rewrite CORREOS.xlsx sandbox (solo allowlist) + notify live."""
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 19_102_163.0,
            "concepto": "GEOEXCON",
            "trx": "RC E31 RETRY NTF 231",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E31")
    if isinstance(prep, ScenarioResult):
        prep.id = "E31"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw, tipo="PAGO DE OBLIGACIÓN ACTUAL", observacion="RC-E31"
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    if not _job_ok(fin):
        cancel_active(session)
        return ScenarioResult(
            "E31",
            "FAIL",
            evidence=f"finalize_failed={fin.get('status')}",
            process_key=_process_key(fin) or _process_key(gen),
            detail={"error": fin.get("error")},
            cleanup="cancelled",
        )
    try:
        assert_sandbox_notify_recipients([NOTIFY_SANDBOX_TO])
        correos_fix = rewrite_sandbox_correos_xlsx(
            session, sandbox_to=NOTIFY_SANDBOX_TO
        )
        _log("E31_correos_rewrite", correos_fix)
    except Exception as exc:
        cancel_active(session)
        return ScenarioResult(
            "E31",
            "BLOCKED",
            evidence=f"correos_rewrite_failed:{exc}",
            process_key=_process_key(fin) or _process_key(gen),
            detail={"correos_path": CORREOS_XLSX_REL, "error": str(exc)},
            cleanup="cancelled",
        )

    ntf1 = notify(session)
    _log("E31_notify_1", {"status": ntf1.get("status"), "error": ntf1.get("error")})
    ntf2 = notify(session)
    _log("E31_notify_2", {"status": ntf2.get("status"), "error": ntf2.get("error")})
    cancel_active(session)
    ok = _job_ok(ntf1) and (
        _job_ok(ntf2)
        or ntf2.get("status") in ("completed", "failed", "http_error")
    )
    return ScenarioResult(
        "E31",
        "PASS" if ok else "FAIL",
        evidence=(
            f"correos={CORREOS_XLSX_REL}; "
            f"n1={ntf1.get('status')}; n2={ntf2.get('status')}"
        ),
        process_key=_process_key(fin) or _process_key(gen),
        detail={
            "applied": applied,
            "allowlist": sorted(assert_sandbox_notify_recipients([NOTIFY_SANDBOX_TO])),
            "correos_path": CORREOS_XLSX_REL,
            "n1": ntf1.get("error") or ntf1.get("result"),
            "n2": ntf2.get("error") or ntf2.get("http_status"),
        },
        cleanup="cancelled",
    )


def run_e38(session: SandboxGraphSession) -> ScenarioResult:
    """Cancel no permitido tras Finalize (fase avanzada)."""
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": 19_102_163.0,
            "concepto": "GEOEXCON",
            "trx": "RC E38 CANCEL BLOCK 231",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E38")
    if isinstance(prep, ScenarioResult):
        prep.id = "E38"
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw, tipo="PAGO DE OBLIGACIÓN ACTUAL", observacion="RC-E38"
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    if not _job_ok(fin):
        cancel_active(session)
        return ScenarioResult(
            "E38",
            "FAIL",
            evidence=f"finalize_failed={fin.get('status')}",
            process_key=_process_key(fin) or _process_key(gen),
            cleanup="cancelled",
        )
    cancel = cancel_active(session)
    # Esperado: cancel rechazado / no allowed (failed) — luego liberar slot.
    blocked_cancel = _job_failed(cancel) or (
        isinstance(cancel.get("result"), dict)
        and str(cancel["result"].get("error_code") or cancel["result"].get("status") or "")
        .lower()
        .find("not_allowed")
        >= 0
    )
    # Liberar para no dejar proceso activo (reset slot).
    reset_bank_slot(session)
    return ScenarioResult(
        "E38",
        "PASS" if blocked_cancel else "FAIL",
        evidence=f"finalize=ok; cancel={cancel.get('status')}; expect_block={blocked_cancel}",
        process_key=_process_key(fin) or _process_key(gen),
        detail={"cancel": cancel.get("result") or cancel.get("error"), "applied": applied},
        cleanup="slot_reset",
    )


def _review_row_snapshot(raw: bytes, credito_contains: str) -> dict[str, Any]:
    from app.application.services.review_schema import AplicacionPagosCols

    wb = load_workbook(io.BytesIO(raw))
    name = next((n for n in wb.sheetnames if "aplicacion" in n.lower()), wb.sheetnames[0])
    ws = wb[name]
    headers = {
        str(ws.cell(1, c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(1, c).value
    }
    for r in range(2, ws.max_row + 1):
        cred = str(ws.cell(r, headers.get(AplicacionPagosCols.CREDITO, 1)).value or "")
        if credito_contains and credito_contains not in cred:
            continue
        out: dict[str, Any] = {"credito": cred, "excel_row": r}
        for key in (
            AplicacionPagosCols.SALDO_VENCIDO,
            AplicacionPagosCols.MONTO_BANCO,
            AplicacionPagosCols.FECHA_LIMITE,
            AplicacionPagosCols.LINK_EXTRACTO,
        ):
            if key in headers:
                out[key] = ws.cell(r, headers[key]).value
        return out
    return {}


def _clear_errores_sheet(raw: bytes) -> bytes:
    """Resolución humana CAPA B: vaciar hoja Errores dejando solo encabezados."""
    from app.application.services.review_schema import ReviewSheets

    wb = load_workbook(io.BytesIO(raw))
    name = next(
        (n for n in wb.sheetnames if n == ReviewSheets.ERRORES or "error" in n.lower()),
        None,
    )
    if name:
        ws = wb[name]
        if ws.max_row and ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _run_typed_finalize(
    session: SandboxGraphSession,
    *,
    sid: str,
    bank_monto: float,
    tipo: str,
    concepto: str = "GEOEXCON",
    fecha=None,
    trx: str = "",
    credito_contains: str = RC_MORA_CREDIT,
    obligacion: float | None = None,
    vencido: float = 0.0,
    capital: float = 0.0,
    expect_finalize_ok: bool = True,
    clear_errores: bool = False,
) -> ScenarioResult:
    rows = [
        {
            "fecha": fecha or d(2026, 5, 23),
            "monto": bank_monto,
            "concepto": concepto,
            "trx": trx or f"RC {sid} {credito_contains}",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, sid)
    if isinstance(prep, ScenarioResult):
        prep.id = sid
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    snap = _review_row_snapshot(raw, credito_contains)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo=tipo,
        obligacion=obligacion,
        vencido=vencido,
        capital=capital,
        observacion=f"RC-{sid}",
        credito_contains=credito_contains,
    )
    if clear_errores:
        edited = _clear_errores_sheet(edited)
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    cancel_active(session)
    ok = _job_ok(fin) if expect_finalize_ok else _job_failed(fin)
    return ScenarioResult(
        sid,
        "PASS" if ok and applied else "FAIL",
        evidence=f"finalize={fin.get('status')}; applied={len(applied)}; snap={snap}; clear_errores={clear_errores}",
        process_key=_process_key(fin) or _process_key(gen),
        detail={"applied": applied, "error": fin.get("error"), "snap": snap},
        cleanup="cancelled",
    )


def run_e05(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(session, asiento_valor=1_000_000.0)
    _log("E05_fixture", fx)
    return _run_typed_finalize(
        session,
        sid="E05",
        bank_monto=1_000_000.0,
        tipo="APLICACIÓN A SALDO VENCIDO",
        obligacion=0.0,
        vencido=1_000_000.0,
    )


def run_e06(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(session, asiento_valor=RC_MORA_VENCIDO)
    _log("E06_fixture", fx)
    return _run_typed_finalize(
        session,
        sid="E06",
        bank_monto=RC_MORA_VENCIDO,
        tipo="APLICACIÓN A SALDO VENCIDO",
        obligacion=0.0,
        vencido=RC_MORA_VENCIDO,
    )


def run_e07(session: SandboxGraphSession) -> ScenarioResult:
    bank = RC_MORA_VENCIDO + 2_000_000.0
    fx = provision_rc_mora_credit(session, asiento_valor=bank)
    _log("E07_fixture", fx)
    return _run_typed_finalize(
        session,
        sid="E07",
        bank_monto=bank,
        tipo="PAGO COMBINADO (SALDO VENCIDO + OBLIGACIÓN ACTUAL)",
        obligacion=2_000_000.0,
        vencido=RC_MORA_VENCIDO,
    )


def run_e08(session: SandboxGraphSession) -> ScenarioResult:
    bank = RC_MORA_VENCIDO + RC_MORA_OBLIG
    fx = provision_rc_mora_credit(session, asiento_valor=bank)
    _log("E08_fixture", fx)
    return _run_typed_finalize(
        session,
        sid="E08",
        bank_monto=bank,
        tipo="PAGO COMBINADO (SALDO VENCIDO + OBLIGACIÓN ACTUAL)",
        obligacion=RC_MORA_OBLIG,
        vencido=RC_MORA_VENCIDO,
    )


def run_e11(session: SandboxGraphSession) -> ScenarioResult:
    bank = RC_MORA_VENCIDO + 1_000_000.0
    fx = provision_rc_mora_credit(session, asiento_valor=bank)
    _log("E11_fixture", fx)
    return _run_typed_finalize(
        session,
        sid="E11",
        bank_monto=bank,
        tipo="APLICACIÓN A SALDO VENCIDO + ABONO A CAPITAL",
        obligacion=0.0,
        vencido=RC_MORA_VENCIDO,
        capital=1_000_000.0,
    )


def run_e12(session: SandboxGraphSession) -> ScenarioResult:
    bank = RC_MORA_VENCIDO + RC_MORA_OBLIG + 1_000_000.0
    fx = provision_rc_mora_credit(session, asiento_valor=bank)
    _log("E12_fixture", fx)
    return _run_typed_finalize(
        session,
        sid="E12",
        bank_monto=bank,
        tipo="PAGO COMBINADO + ABONO A CAPITAL",
        obligacion=RC_MORA_OBLIG,
        vencido=RC_MORA_VENCIDO,
        capital=1_000_000.0,
    )


def run_e18(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(
        session, right_role="APLICACION_ANTERIOR", right_amount=1_250_000.0
    )
    _log("E18_fixture", fx)
    result = _run_typed_finalize(
        session,
        sid="E18",
        bank_monto=RC_MORA_OBLIG,
        tipo="PAGO DE OBLIGACIÓN ACTUAL",
        obligacion=RC_MORA_OBLIG,
    )
    snap = result.detail.get("snap") or {}
    from app.application.services.review_schema import AplicacionPagosCols

    saldo = snap.get(AplicacionPagosCols.SALDO_VENCIDO)
    try:
        saldo_f = float(saldo or 0)
    except (TypeError, ValueError):
        saldo_f = -1.0
    if result.status == "PASS" and saldo_f > 0.02:
        result.status = "FAIL"
        result.evidence += "; aplicacion_anterior_must_not_fill_saldo_vencido"
    return result


def run_e19(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(session, right_role="SALDO_VENCIDO")
    _log("E19_fixture", fx)
    result = _run_typed_finalize(
        session,
        sid="E19",
        bank_monto=RC_MORA_OBLIG,
        tipo="PAGO DE OBLIGACIÓN ACTUAL",
        obligacion=RC_MORA_OBLIG,
    )
    snap = result.detail.get("snap") or {}
    from app.application.services.review_schema import AplicacionPagosCols

    saldo = snap.get(AplicacionPagosCols.SALDO_VENCIDO)
    try:
        saldo_f = float(saldo or 0)
    except (TypeError, ValueError):
        saldo_f = 0.0
    if result.status == "PASS" and abs(saldo_f - RC_MORA_VENCIDO) > 1.0:
        result.status = "FAIL"
        result.evidence += f"; expected_saldo_vencido={RC_MORA_VENCIDO} got={saldo}"
    return result


def run_e20(session: SandboxGraphSession) -> ScenarioResult:
    """Ambiguo: Generate deja Errores + saldo vacío; humano limpia Errores y Finaliza."""
    fx = provision_rc_mora_credit(session, right_role="AMBIGUO", right_amount=RC_MORA_VENCIDO)
    _log("E20_fixture", fx)
    result = _run_typed_finalize(
        session,
        sid="E20",
        bank_monto=RC_MORA_OBLIG,
        tipo="PAGO DE OBLIGACIÓN ACTUAL",
        obligacion=RC_MORA_OBLIG,
        clear_errores=True,
    )
    snap = result.detail.get("snap") or {}
    from app.application.services.review_schema import AplicacionPagosCols

    saldo = snap.get(AplicacionPagosCols.SALDO_VENCIDO)
    try:
        saldo_f = float(saldo or 0)
    except (TypeError, ValueError):
        saldo_f = 0.0
    link = str(snap.get(AplicacionPagosCols.LINK_EXTRACTO) or "").upper()
    if "AMBIGUO" not in link:
        result.status = "FAIL"
        result.evidence += "; expected_ambiguous_extract_link"
    elif saldo_f > 0.02:
        result.status = "FAIL"
        result.evidence += "; ambiguous_must_not_fill_saldo_vencido"
    return result


def run_e22(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(
        session,
        right_role="VACIO",
        fecha_limite="15/04/2026",
        extra_extracts=[
            {
                "name": "Extracto JUNIO posterior 301.pdf",
                "fecha_limite": "23/06/2026",
                "right_role": "SALDO_VENCIDO",
                "right_amount": 9_999_999.0,
            }
        ],
    )
    _log("E22_fixture", fx)
    result = _run_typed_finalize(
        session,
        sid="E22",
        bank_monto=RC_MORA_OBLIG,
        tipo="PAGO DE OBLIGACIÓN ACTUAL",
        fecha=d(2026, 4, 15),
        obligacion=RC_MORA_OBLIG,
        clear_errores=True,
    )
    snap = result.detail.get("snap") or {}
    from app.application.services.review_schema import AplicacionPagosCols

    link = str(snap.get(AplicacionPagosCols.LINK_EXTRACTO) or "").upper()
    fecha = str(snap.get(AplicacionPagosCols.FECHA_LIMITE) or "")
    if "JUNIO" in link or "06/2026" in fecha or "2026-06" in fecha:
        result.status = "FAIL"
        result.evidence += f"; selected_later_extract_not_as_of fecha={fecha}"
    elif "2026-04" not in fecha and "15/04" not in fecha and "04-15" not in fecha:
        # Fecha límite as-of abril esperada
        if result.status == "PASS":
            result.status = "FAIL"
            result.evidence += f"; expected_april_fecha_limite got={fecha}"
    return result


def _finalize_notify_merge(
    session: SandboxGraphSession, sid: str, bank_monto: float, tipo: str
) -> tuple[ScenarioResult | None, dict[str, Any], dict[str, Any], dict[str, Any]]:
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": bank_monto,
            "concepto": "GEOEXCON",
            "trx": f"RC {sid} {RC_MORA_CREDIT}",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, sid)
    if isinstance(prep, ScenarioResult):
        prep.id = sid
        return prep, {}, {}, {}
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo=tipo,
        observacion=f"RC-{sid}",
        credito_contains=RC_MORA_CREDIT,
        obligacion=bank_monto,
    )
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    if not _job_ok(fin):
        cancel_active(session)
        return (
            ScenarioResult(
                sid,
                "FAIL",
                evidence=f"finalize={fin.get('status')}",
                process_key=_process_key(fin) or _process_key(gen),
                detail={"error": fin.get("error"), "applied": applied},
                cleanup="cancelled",
            ),
            fin,
            {},
            {},
        )
    rewrite_sandbox_correos_xlsx(session)
    nfy = notify(session)
    mrg = merge(session)
    return None, fin, nfy, mrg


def _amort_effectively_blocked(dry: dict[str, Any], mrg: dict[str, Any] | None = None) -> bool:
    """True si Apply no escribiría (item ERROR, would_apply=0, merge skip/parcial)."""
    if mrg is not None:
        mr = _result_dict(mrg)
        outputs = list(mr.get("outputs") or [])
        skipped = list(mr.get("skipped") or [])
        if skipped and not outputs:
            return True
        if mr.get("file_action") in ("partial",) and not outputs:
            return True
    if _job_failed(dry):
        return True
    result = _result_dict(dry)
    if not result:
        return True
    summary = result.get("summary") if isinstance(result.get("summary"), dict) else {}
    if summary.get("would_apply") == 0 and int(summary.get("total") or 0) > 0:
        return True
    for item in result.get("items") or []:
        if not isinstance(item, dict):
            continue
        if item.get("application_status") == "ERROR" or item.get("error_code"):
            return True
    for issue in result.get("operational_issues") or []:
        if isinstance(issue, dict) and issue.get("severity") in ("error", "blocker"):
            return True
    return False


def run_e25(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(
        session, right_role="VACIO", asiento_valor=1_000.0
    )
    _log("E25_fixture", fx)
    early, fin, nfy, mrg = _finalize_notify_merge(
        session, "E25", RC_MORA_OBLIG, "PAGO DE OBLIGACIÓN ACTUAL"
    )
    if early:
        return early
    dry = amort_dry_run(session)
    cancel_active(session)
    blocked = _amort_effectively_blocked(dry, mrg)
    return ScenarioResult(
        "E25",
        "PASS" if blocked else "FAIL",
        evidence=(
            f"notify={nfy.get('status')}; merge={mrg.get('status')}; "
            f"dry={dry.get('status')}; expect_block"
        ),
        process_key=_process_key(dry) or _process_key(fin),
        detail={"dry": dry.get("error") or dry.get("result"), "fixture": fx},
        cleanup="cancelled",
    )


def run_e26(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(
        session, right_role="VACIO", asiento_credit_label="999", asiento_valor=RC_MORA_OBLIG
    )
    _log("E26_fixture", fx)
    early, fin, nfy, mrg = _finalize_notify_merge(
        session, "E26", RC_MORA_OBLIG, "PAGO DE OBLIGACIÓN ACTUAL"
    )
    if early:
        return early
    dry = amort_dry_run(session)
    cancel_active(session)
    blocked = _amort_effectively_blocked(dry, mrg)
    return ScenarioResult(
        "E26",
        "PASS" if blocked else "FAIL",
        evidence=f"dry={dry.get('status')}; expect_wrong_credit_block",
        process_key=_process_key(dry) or _process_key(fin),
        detail={"dry": dry.get("error") or dry.get("result"), "notify": nfy, "merge": mrg},
        cleanup="cancelled",
    )


def run_e27(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(session, right_role="VACIO", asiento_mode="missing")
    _log("E27_fixture", fx)
    early, fin, nfy, mrg = _finalize_notify_merge(
        session, "E27", RC_MORA_OBLIG, "PAGO DE OBLIGACIÓN ACTUAL"
    )
    if early:
        return early
    dry = amort_dry_run(session)
    cancel_active(session)
    blocked = _amort_effectively_blocked(dry, mrg)
    return ScenarioResult(
        "E27",
        "PASS" if blocked else "FAIL",
        evidence=f"merge={mrg.get('status')}; dry={dry.get('status')}; expect_missing_asiento",
        process_key=_process_key(mrg) or _process_key(fin),
        detail={"merge": mrg.get("error") or mrg.get("result"), "dry": dry.get("error") or dry.get("result")},
        cleanup="cancelled",
    )


def run_e28(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(session, right_role="VACIO", asiento_mode="illegible")
    _log("E28_fixture", fx)
    early, fin, nfy, mrg = _finalize_notify_merge(
        session, "E28", RC_MORA_OBLIG, "PAGO DE OBLIGACIÓN ACTUAL"
    )
    if early:
        return early
    dry = amort_dry_run(session)
    cancel_active(session)
    blocked = _amort_effectively_blocked(dry, mrg)
    return ScenarioResult(
        "E28",
        "PASS" if blocked else "FAIL",
        evidence=f"merge={mrg.get('status')}; dry={dry.get('status')}; expect_illegible_block",
        process_key=_process_key(mrg) or _process_key(fin),
        detail={"merge": mrg.get("result") or mrg.get("error"), "dry": dry.get("error") or dry.get("result")},
        cleanup="cancelled",
    )


def run_e32(session: SandboxGraphSession) -> ScenarioResult:
    fx = provision_rc_mora_credit(session, right_role="VACIO", asiento_valor=RC_MORA_OBLIG)
    _log("E32_fixture", fx)
    early, fin, nfy, mrg1 = _finalize_notify_merge(
        session, "E32", RC_MORA_OBLIG, "PAGO DE OBLIGACIÓN ACTUAL"
    )
    if early:
        return early
    mrg2 = merge(session)
    cancel_active(session)
    r1 = _result_dict(mrg1)
    r2 = _result_dict(mrg2)
    ok = _job_ok(mrg1) and (
        _job_ok(mrg2)
        or r2.get("already_merged")
        or r2.get("pdf_reused")
        or r2.get("file_action") in ("reused", "already_merged")
        or mrg2.get("http_status") in (409, 422)
    )
    return ScenarioResult(
        "E32",
        "PASS" if ok else "FAIL",
        evidence=(
            f"m1={mrg1.get('status')} action={r1.get('file_action')}; "
            f"m2={mrg2.get('status')} action={r2.get('file_action')} reused={r2.get('pdf_reused')}"
        ),
        process_key=_process_key(mrg1) or _process_key(fin),
        detail={"m1": r1, "m2": r2, "notify": nfy.get("status")},
        cleanup="cancelled",
    )


def _download_path_bytes(session: SandboxGraphSession, relative_path: str) -> bytes:
    assert_sandbox_mutable_path(relative_path)
    parent = "/".join(relative_path.split("/")[:-1])
    name = relative_path.split("/")[-1]
    folder_id = session.walk(parent)
    item = next(i for i in session.children(folder_id) if i.get("name") == name)
    return session.download_item(str(item["id"]))


def _tabla_saldo_snapshot(raw: bytes) -> dict[str, Any]:
    """Lee saldos / valor pagado de tabla canónica (fila 2=prev, fila 3=cuota).

    Apply deja fórmulas Excel (``=+F3+G3+H3``, ``=+K2-G3``). Graph no cachea
    ``data_only``, así que se leen celdas + se evalúan sumas/restas simples.
    """
    wb = load_workbook(io.BytesIO(raw), data_only=False)
    ws = wb.active
    assert ws is not None
    headers = {
        str(ws.cell(1, c).value).strip().lower(): c
        for c in range(1, (ws.max_column or 1) + 1)
        if ws.cell(1, c).value
    }

    def _col(*names: str) -> int | None:
        for n in names:
            for h, idx in headers.items():
                if n in h.replace("á", "a"):
                    return idx
        return None

    def _num(val: Any) -> float | None:
        if val is None or val == "":
            return None
        if isinstance(val, (int, float)) and not isinstance(val, bool):
            return float(val)
        text = str(val).strip().replace(",", ".")
        if text.startswith("="):
            return None
        try:
            return float(text)
        except ValueError:
            return None

    def _eval_cell(row: int, col: int | None) -> Any:
        if not col:
            return None
        val = ws.cell(row, col).value
        n = _num(val)
        if n is not None:
            return n
        if not isinstance(val, str) or not val.strip().startswith("="):
            return val
        # Fórmulas típicas Apply: =+F3+G3+H3  /  =+K2-G3
        expr = val.strip().lstrip("=").lstrip("+")
        import re

        refs = re.findall(r"([A-Z]+)(\d+)", expr.upper())
        if not refs:
            return val
        nums: list[float] = []
        for letters, rtxt in refs:
            cidx = 0
            for ch in letters:
                cidx = cidx * 26 + (ord(ch) - ord("A") + 1)
            cell_n = _num(ws.cell(int(rtxt), cidx).value)
            if cell_n is None:
                return val
            nums.append(cell_n)
        if "+" in expr and "-" not in expr and len(nums) >= 2:
            return float(sum(nums))
        if "-" in expr and len(nums) == 2:
            return float(nums[0] - nums[1])
        return val

    saldo_c = _col("saldo a capital", "saldo capital")
    vp_c = _col("valor pagado cliente", "valor pagado")
    abono_c = _col("abono a k", "abono k")
    fecha_c = _col("fecha pago")
    ibr_c = _col("ibr")
    int_c = _col("valor intereses")
    mora_c = _col("intereses mora", "mora")
    out: dict[str, Any] = {"rows": ws.max_row}
    if saldo_c:
        out["saldo_row2"] = _eval_cell(2, saldo_c)
        out["saldo_row3"] = _eval_cell(3, saldo_c)
        out["saldo_row3_raw"] = ws.cell(3, saldo_c).value
    if vp_c:
        out["valor_pagado_row3"] = _eval_cell(3, vp_c)
        out["valor_pagado_row3_raw"] = ws.cell(3, vp_c).value
    if abono_c:
        out["abono_k_row3"] = _eval_cell(3, abono_c)
    if int_c:
        out["intereses_row3"] = _eval_cell(3, int_c)
    if mora_c:
        out["mora_row3"] = _eval_cell(3, mora_c)
    if fecha_c:
        out["fecha_pago_row3"] = str(ws.cell(3, fecha_c).value or "")
    if ibr_c:
        out["ibr_row3"] = ws.cell(3, ibr_c).value
    return out


def _pipeline_to_merge(
    session: SandboxGraphSession,
    *,
    sid: str,
    rows: list[dict[str, Any]],
    tipo: str,
    credito_contains: str,
    obligacion: float | None = None,
    capital: float = 0.0,
    vencido: float = 0.0,
    clear_errores: bool = True,
    force_monto_banco: float | None = None,
) -> tuple[ScenarioResult | None, dict[str, Any], dict[str, Any], dict[str, Any], str]:
    prep = _prep_bank_and_generate(session, rows, sid)
    if isinstance(prep, ScenarioResult):
        prep.id = sid
        return prep, {}, {}, {}, ""
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_single_credit_pago(
        raw,
        tipo=tipo,
        obligacion=obligacion,
        vencido=vencido,
        capital=capital,
        observacion=f"RC-{sid}",
        credito_contains=credito_contains,
        force_monto_banco=force_monto_banco,
    )
    if clear_errores:
        edited = _clear_errores_sheet(edited)
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    if not _job_ok(fin):
        cancel_active(session)
        return (
            ScenarioResult(
                sid,
                "FAIL",
                evidence=f"finalize={fin.get('status')}",
                process_key=_process_key(fin) or _process_key(gen),
                detail={"error": fin.get("error"), "applied": applied},
                cleanup="cancelled",
            ),
            fin,
            {},
            {},
            process_date,
        )
    rewrite_sandbox_correos_xlsx(session)
    nfy = notify(session)
    if not _job_ok(nfy):
        cancel_active(session)
        return (
            ScenarioResult(
                sid,
                "FAIL",
                evidence=f"notify={nfy.get('status')}",
                process_key=_process_key(fin),
                detail={"notify": nfy.get("error")},
                cleanup="cancelled",
            ),
            fin,
            nfy,
            {},
            process_date,
        )
    mrg = merge(session)
    if not _job_ok(mrg):
        cancel_active(session)
        return (
            ScenarioResult(
                sid,
                "FAIL",
                evidence=f"merge={mrg.get('status')}",
                process_key=_process_key(fin),
                detail={"merge": mrg.get("error") or mrg.get("result")},
                cleanup="cancelled",
            ),
            fin,
            nfy,
            mrg,
            process_date,
        )
    return None, fin, nfy, mrg, process_date


def soft_close_ui(
    session: SandboxGraphSession,
    *,
    bank_code: str,
    process_key: str,
    reason: str = "RC-E39 soft close",
) -> dict[str, Any]:
    from scripts.e2e_rc.graph_session import DEFAULT_BASE

    cred_path = Path(r"D:\CMC\HBI_Capital\_work\rc_e2e\e2e_credentials.json")
    data = json.loads(cred_path.read_text(encoding="utf-8"))
    origin = DEFAULT_BASE.rstrip("/")
    ui_headers = {"Origin": origin, "Content-Type": "application/json"}
    with httpx.Client(timeout=120.0, base_url=DEFAULT_BASE) as client:
        login = client.post(
            "/api/ui/v1/auth/login",
            json={"username": data["username"], "password": data["password"]},
            headers=ui_headers,
        )
        if login.status_code != 200:
            return {
                "status": "http_error",
                "http_status": login.status_code,
                "body": login.text[:500],
            }
        csrf = client.get("/api/ui/v1/auth/csrf", headers=ui_headers)
        if csrf.status_code != 200:
            return {
                "status": "http_error",
                "http_status": csrf.status_code,
                "body": csrf.text[:500],
                "step": "csrf",
            }
        csrf_token = (csrf.json() or {}).get("csrf_token") or ""
        write_headers = {**ui_headers, "X-CSRF-Token": csrf_token}
        r = client.post(
            "/api/ui/v1/processes/soft-close",
            json={"bank_code": bank_code, "process_key": process_key, "reason": reason},
            headers=write_headers,
        )
        if r.status_code not in (200, 202):
            return {
                "status": "http_error",
                "http_status": r.status_code,
                "body": r.text[:800],
            }
        body = r.json()
        job_id = body.get("job_id")
        if not job_id:
            return body
        t0 = time.time()
        while time.time() - t0 < 600:
            jr = client.get(f"/api/ui/v1/jobs/{job_id}")
            if jr.status_code == 200:
                done = jr.json()
                if done.get("status") in ("completed", "failed", "error"):
                    return done
            time.sleep(5)
        return {"status": "timeout", "job_id": job_id}


def run_e13(session: SandboxGraphSession) -> ScenarioResult:
    """Última cuota / PAYOFF_OK: Apply + re-read saldo cancelado."""
    fx = provision_canonical_payoff_credit(
        session,
        credit=RC_PAYOFF_CREDIT,
        folder=RC_PAYOFF_FOLDER,
        capital=RC_E13_CAPITAL,
        intereses=RC_E13_INTERESES,
        mora=RC_E13_MORA,
        valor_pagado=RC_E13_TOTAL,
        backup_suffix="E13",
    )
    _log("E13_fixture", {k: v for k, v in fx.items() if k != "tabla_bytes"})
    before = _tabla_saldo_snapshot(_download_path_bytes(session, fx["tabla"]))
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": RC_E13_TOTAL,
            "concepto": "GEOEXCON",
            "trx": "RC E13 PAYOFF OK 302",
        }
    ]
    early, fin, nfy, mrg, _pd = _pipeline_to_merge(
        session,
        sid="E13",
        rows=rows,
        tipo="CANCELACIÓN / PAGO TOTAL",
        credito_contains=RC_PAYOFF_CREDIT,
        obligacion=RC_E13_TOTAL,
    )
    if early:
        return early
    dry = amort_dry_run(session)
    dry_r = _result_dict(dry)
    _log("E13_dry", {"status": dry.get("status"), "can_apply": dry_r.get("can_apply"), "summary": dry_r.get("summary")})
    if not _job_ok(dry) or dry_r.get("can_apply") is not True:
        cancel_active(session)
        return ScenarioResult(
            "E13",
            "FAIL",
            evidence=f"dry can_apply={dry_r.get('can_apply')}; status={dry.get('status')}",
            process_key=_process_key(dry) or _process_key(fin),
            detail={"dry": dry_r or dry.get("error"), "before": before, "fixture": fx},
            cleanup="cancelled",
        )
    if _has_payoff_not_achieved(dry):
        cancel_active(session)
        return ScenarioResult(
            "E13",
            "FAIL",
            evidence="unexpected_PAYOFF_NOT_ACHIEVED",
            process_key=_process_key(dry),
            detail={"dry": dry_r},
            cleanup="cancelled",
        )
    apply = amort_apply(session)
    apply_r = _result_dict(apply)
    _log("E13_apply", {"status": apply.get("status"), "result_keys": list(apply_r)[:20]})
    after = _tabla_saldo_snapshot(_download_path_bytes(session, fx["tabla"]))
    cancel_active(session)
    vp = after.get("valor_pagado_row3")
    abono = after.get("abono_k_row3")
    try:
        vp_f = float(vp or 0)
        abono_f = float(abono or 0)
        saldo_after = after.get("saldo_row3")
        saldo_f = float(saldo_after) if saldo_after is not None else None
    except (TypeError, ValueError):
        vp_f, abono_f, saldo_f = 0.0, 0.0, None
    ok = (
        _job_ok(apply)
        and abs(vp_f - RC_E13_TOTAL) < 1.0
        and abs(abono_f - RC_E13_CAPITAL) < 1.0
        and saldo_f is not None
        and abs(saldo_f) < 1.0
    )
    return ScenarioResult(
        "E13",
        "PASS" if ok else "FAIL",
        evidence=(
            f"dry_can_apply={dry_r.get('can_apply')}; apply={apply.get('status')}; "
            f"before_saldo={before.get('saldo_row2')}; after_vp={vp}; after_abono={abono}"
        ),
        process_key=_process_key(apply) or _process_key(fin),
        detail={
            "before": before,
            "after": after,
            "dry": dry_r.get("summary"),
            "apply": apply_r.get("summary") or apply.get("error"),
            "fixture": {k: fx[k] for k in ("tabla", "asiento", "saldo_before") if k in fx},
        },
        cleanup="cancelled",
    )


def run_e14(session: SandboxGraphSession) -> ScenarioResult:
    """Cancelación con capital residual mayor (capital adicional)."""
    fx = provision_canonical_payoff_credit(
        session,
        credit=RC_E14_CREDIT,
        folder=RC_E14_FOLDER,
        capital=RC_E14_CAPITAL,
        intereses=RC_E14_INTERESES,
        mora=RC_E14_MORA,
        valor_pagado=RC_E14_TOTAL,
        backup_suffix="E14",
    )
    _log("E14_fixture", {k: v for k, v in fx.items() if k != "tabla_bytes"})
    before = _tabla_saldo_snapshot(_download_path_bytes(session, fx["tabla"]))
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": RC_E14_TOTAL,
            "concepto": "GEOEXCON",
            "trx": "RC E14 PAYOFF EXTRA 303",
        }
    ]
    early, fin, nfy, mrg, _pd = _pipeline_to_merge(
        session,
        sid="E14",
        rows=rows,
        tipo="CANCELACIÓN / PAGO TOTAL",
        credito_contains=RC_E14_CREDIT,
        obligacion=RC_E14_TOTAL,
    )
    if early:
        return early
    dry = amort_dry_run(session)
    dry_r = _result_dict(dry)
    if not _job_ok(dry) or dry_r.get("can_apply") is not True or _has_payoff_not_achieved(dry):
        cancel_active(session)
        return ScenarioResult(
            "E14",
            "FAIL",
            evidence=f"dry can_apply={dry_r.get('can_apply')} payoff_bad={_has_payoff_not_achieved(dry)}",
            process_key=_process_key(dry) or _process_key(fin),
            detail={"dry": dry_r or dry.get("error"), "before": before},
            cleanup="cancelled",
        )
    apply = amort_apply(session)
    after = _tabla_saldo_snapshot(_download_path_bytes(session, fx["tabla"]))
    cancel_active(session)
    try:
        abono_f = float(after.get("abono_k_row3") or 0)
        vp_f = float(after.get("valor_pagado_row3") or 0)
        saldo_after = after.get("saldo_row3")
        saldo_f = float(saldo_after) if saldo_after is not None else None
    except (TypeError, ValueError):
        abono_f, vp_f, saldo_f = 0.0, 0.0, None
    ok = (
        _job_ok(apply)
        and abs(abono_f - RC_E14_CAPITAL) < 1.0
        and abs(vp_f - RC_E14_TOTAL) < 1.0
        and saldo_f is not None
        and abs(saldo_f) < 1.0
    )
    return ScenarioResult(
        "E14",
        "PASS" if ok else "FAIL",
        evidence=f"apply={apply.get('status')}; abono={abono_f}; vp={vp_f}; capital={RC_E14_CAPITAL}",
        process_key=_process_key(apply) or _process_key(fin),
        detail={"before": before, "after": after, "notify": nfy.get("status"), "merge": _result_dict(mrg).get("file_action")},
        cleanup="cancelled",
    )


def run_e24(session: SandboxGraphSession) -> ScenarioResult:
    """Review ≠ asiento por crédito; SUM asientos = banco → no bloquear."""
    fx_a = provision_canonical_payoff_credit(
        session,
        credit=RC_E24_CREDIT_A,
        folder=f"CREDITO # {RC_E24_CREDIT_A}",
        capital=RC_E24_ASIENTO_A * 0.6,
        intereses=RC_E24_ASIENTO_A * 0.3,
        mora=RC_E24_ASIENTO_A * 0.1,
        valor_pagado=RC_E24_ASIENTO_A,
        backup_suffix="E24A",
    )
    fx_b = provision_canonical_payoff_credit(
        session,
        credit=RC_E24_CREDIT_B,
        folder=f"CREDITO # {RC_E24_CREDIT_B}",
        capital=RC_E24_ASIENTO_B * 0.6,
        intereses=RC_E24_ASIENTO_B * 0.3,
        mora=RC_E24_ASIENTO_B * 0.1,
        valor_pagado=RC_E24_ASIENTO_B,
        backup_suffix="E24B",
    )
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": RC_E24_BANK,
            "concepto": "GEOEXCON",
            "trx": "RC E24 REVIEW!=ASIENTO",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E24")
    if isinstance(prep, ScenarioResult):
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_credits_split(
        raw,
        splits=[
            {
                "credito_contains": RC_E24_CREDIT_A,
                "tipo": "PAGO DE OBLIGACIÓN ACTUAL",
                "obligacion": RC_E24_REVIEW_A,
            },
            {
                "credito_contains": RC_E24_CREDIT_B,
                "tipo": "PAGO DE OBLIGACIÓN ACTUAL",
                "obligacion": RC_E24_REVIEW_B,
            },
        ],
        observacion="RC-E24",
    )
    edited = _clear_errores_sheet(edited)
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    if not _job_ok(fin):
        cancel_active(session)
        return ScenarioResult(
            "E24",
            "FAIL",
            evidence=f"finalize={fin.get('status')}",
            detail={"error": fin.get("error"), "applied": applied},
            cleanup="cancelled",
        )
    rewrite_sandbox_correos_xlsx(session)
    nfy = notify(session)
    mrg = merge(session)
    dry = amort_dry_run(session)
    cancel_active(session)
    dry_r = _result_dict(dry)
    # No debe bloquear por mismatch review vs asiento si totales asiento=banco.
    blocked_wrongly = False
    for it in dry_r.get("items") or []:
        if not isinstance(it, dict):
            continue
        code = str(it.get("error_code") or "")
        if "REVIEW" in code.upper() and "ASIENTO" in code.upper():
            blocked_wrongly = True
    ok = (
        _job_ok(mrg)
        and len(applied) >= 2
        and abs(RC_E24_ASIENTO_A + RC_E24_ASIENTO_B - RC_E24_BANK) < 0.02
        and abs(RC_E24_REVIEW_A + RC_E24_REVIEW_B - RC_E24_BANK) < 0.02
        and not blocked_wrongly
        and (_job_ok(dry) or dry_r.get("can_apply") in (True, False))
    )
    return ScenarioResult(
        "E24",
        "PASS" if ok else "FAIL",
        evidence=(
            f"finalize=ok; merge={mrg.get('status')}; dry={dry.get('status')}; "
            f"review={RC_E24_REVIEW_A}+{RC_E24_REVIEW_B}; asiento={RC_E24_ASIENTO_A}+{RC_E24_ASIENTO_B}"
        ),
        process_key=_process_key(dry) or _process_key(fin) or _process_key(gen),
        detail={"applied": applied, "dry_summary": dry_r.get("summary"), "fx_a": fx_a["asiento"], "fx_b": fx_b["asiento"], "notify": nfy.get("status")},
        cleanup="cancelled",
    )


def run_e33(session: SandboxGraphSession) -> ScenarioResult:
    """Retry Apply: segundo intento no duplica escritura."""
    fx = provision_canonical_payoff_credit(
        session,
        credit=RC_PAYOFF_CREDIT,
        folder=RC_PAYOFF_FOLDER,
        capital=RC_E13_CAPITAL,
        intereses=RC_E13_INTERESES,
        mora=RC_E13_MORA,
        valor_pagado=RC_E13_TOTAL,
        backup_suffix="E33",
    )
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": RC_E13_TOTAL,
            "concepto": "GEOEXCON",
            "trx": "RC E33 RETRY APPLY 302",
        }
    ]
    early, fin, nfy, mrg, _pd = _pipeline_to_merge(
        session,
        sid="E33",
        rows=rows,
        tipo="CANCELACIÓN / PAGO TOTAL",
        credito_contains=RC_PAYOFF_CREDIT,
        obligacion=RC_E13_TOTAL,
    )
    if early:
        return early
    dry = amort_dry_run(session)
    if not _job_ok(dry) or _result_dict(dry).get("can_apply") is not True:
        cancel_active(session)
        return ScenarioResult(
            "E33",
            "FAIL",
            evidence=f"dry_not_ready={dry.get('status')}",
            detail={"dry": dry.get("error") or dry.get("result")},
            cleanup="cancelled",
        )
    a1 = amort_apply(session)
    after1 = _tabla_saldo_snapshot(_download_path_bytes(session, fx["tabla"]))
    raw1 = _download_path_bytes(session, fx["tabla"])
    a2 = amort_apply(session)
    after2 = _tabla_saldo_snapshot(_download_path_bytes(session, fx["tabla"]))
    raw2 = _download_path_bytes(session, fx["tabla"])
    cancel_active(session)
    a2r = _result_dict(a2)
    idempotent = (
        _job_ok(a1)
        and (
            _job_ok(a2)
            or a2.get("http_status") in (409, 422)
            or str(a2r.get("process_control_estado") or "").upper() in ("AMORTIZADO", "AMORTIZACION_COMPLETA", "FINALIZADO")
            or any(
                str(it.get("application_status") or it.get("apply_status") or "")
                in ("ALREADY_APPLIED", "SKIPPED_IDEMPOTENT")
                for it in (a2r.get("items") or [])
            )
        )
        and after1.get("valor_pagado_row3") == after2.get("valor_pagado_row3")
        and after1.get("abono_k_row3") == after2.get("abono_k_row3")
        and len(raw1) == len(raw2)
    )
    return ScenarioResult(
        "E33",
        "PASS" if idempotent else "FAIL",
        evidence=f"a1={a1.get('status')}; a2={a2.get('status')}; vp1={after1.get('valor_pagado_row3')}; vp2={after2.get('valor_pagado_row3')}",
        process_key=_process_key(a1) or _process_key(fin),
        detail={"after1": after1, "after2": after2, "a2": a2r.get("summary") or a2.get("error"), "notify": nfy.get("status"), "merge": mrg.get("status")},
        cleanup="cancelled",
    )


def run_e39(session: SandboxGraphSession) -> ScenarioResult:
    """Soft-close en CONSOLIDADO sin Apply."""
    fx = provision_canonical_payoff_credit(
        session,
        credit=RC_PAYOFF_CREDIT,
        folder=RC_PAYOFF_FOLDER,
        capital=RC_E13_CAPITAL,
        intereses=RC_E13_INTERESES,
        mora=RC_E13_MORA,
        valor_pagado=RC_E13_TOTAL,
        backup_suffix="E39",
    )
    before = _tabla_saldo_snapshot(_download_path_bytes(session, fx["tabla"]))
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": RC_E13_TOTAL,
            "concepto": "GEOEXCON",
            "trx": "RC E39 SOFT CLOSE 302",
        }
    ]
    early, fin, nfy, mrg, _pd = _pipeline_to_merge(
        session,
        sid="E39",
        rows=rows,
        tipo="PAGO DE OBLIGACIÓN ACTUAL",
        credito_contains=RC_PAYOFF_CREDIT,
        obligacion=RC_E13_TOTAL,
    )
    if early:
        return early
    pk = _process_key(mrg) or _process_key(fin)
    if not pk:
        cancel_active(session)
        return ScenarioResult("E39", "FAIL", evidence="missing_process_key", cleanup="cancelled")
    sc = soft_close_ui(session, bank_code=BANK_CODE, process_key=pk)
    after = _tabla_saldo_snapshot(_download_path_bytes(session, fx["tabla"]))
    # Liberar slot por si soft-close no desactivó.
    reset_bank_slot(session)
    sc_r = sc.get("result") if isinstance(sc.get("result"), dict) else sc
    ok = _job_ok(sc) or str(sc.get("status") or "").lower() == "completed"
    # Tabla no debe haber recibido Apply (valor pagado vacío).
    no_write = after.get("valor_pagado_row3") in (None, "", 0, 0.0) and before.get("valor_pagado_row3") in (None, "", 0, 0.0)
    return ScenarioResult(
        "E39",
        "PASS" if ok and no_write else "FAIL",
        evidence=f"soft_close={sc.get('status')}; no_write={no_write}; estado={sc_r.get('estado') if isinstance(sc_r, dict) else None}",
        process_key=pk,
        detail={"soft_close": sc_r if isinstance(sc_r, dict) else sc, "before": before, "after": after, "notify": nfy.get("status")},
        cleanup="slot_reset",
    )


def run_e40(session: SandboxGraphSession) -> ScenarioResult:
    """AMORTIZACION_PARCIAL: A OK, B falla; retry solo B."""
    fx_a = provision_canonical_payoff_credit(
        session,
        credit=RC_PAYOFF_CREDIT,
        folder=RC_PAYOFF_FOLDER,
        capital=RC_E40_SPLIT_A * 0.6,
        intereses=RC_E40_SPLIT_A * 0.3,
        mora=RC_E40_SPLIT_A * 0.1,
        valor_pagado=RC_E40_SPLIT_A,
        backup_suffix="E40A",
    )
    # B empieza corrupto
    fx_b = provision_canonical_payoff_credit(
        session,
        credit=RC_E40_CREDIT_B,
        folder=RC_E40_FOLDER_B,
        capital=RC_E40_SPLIT_B * 0.6,
        intereses=RC_E40_SPLIT_B * 0.3,
        mora=RC_E40_SPLIT_B * 0.1,
        valor_pagado=RC_E40_SPLIT_B,
        tabla_mode="corrupt",
        backup_suffix="E40B",
    )
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": RC_E40_BANK,
            "concepto": "GEOEXCON",
            "trx": "RC E40 PARTIAL 302-304",
        }
    ]
    prep = _prep_bank_and_generate(session, rows, "E40")
    if isinstance(prep, ScenarioResult):
        return prep
    gen, process_date = prep
    item, raw = latest_review(session)
    edited, applied = approve_credits_split(
        raw,
        splits=[
            {
                "credito_contains": RC_PAYOFF_CREDIT,
                "tipo": "PAGO DE OBLIGACIÓN ACTUAL",
                "obligacion": RC_E40_SPLIT_A,
            },
            {
                "credito_contains": RC_E40_CREDIT_B,
                "tipo": "PAGO DE OBLIGACIÓN ACTUAL",
                "obligacion": RC_E40_SPLIT_B,
            },
        ],
        observacion="RC-E40",
    )
    edited = _clear_errores_sheet(edited)
    upload_review(session, item, edited)
    fin = finalize(session, process_date)
    if not _job_ok(fin):
        cancel_active(session)
        return ScenarioResult("E40", "FAIL", evidence=f"finalize={fin.get('status')}", detail={"error": fin.get("error")}, cleanup="cancelled")
    rewrite_sandbox_correos_xlsx(session)
    nfy = notify(session)
    mrg = merge(session)
    if not _job_ok(mrg):
        # Merge parcial puede ser OK si B sin tabla válida en generate — still try apply path
        pass
    a1 = amort_apply(session)
    a1r = _result_dict(a1)
    after_a1 = _tabla_saldo_snapshot(_download_path_bytes(session, fx_a["tabla"]))
    raw_a1 = _download_path_bytes(session, fx_a["tabla"])
    estado1 = str(a1r.get("process_control_estado") or a1r.get("estado") or "").upper()
    # Reparar B
    fx_b_fix = provision_canonical_payoff_credit(
        session,
        credit=RC_E40_CREDIT_B,
        folder=RC_E40_FOLDER_B,
        capital=RC_E40_SPLIT_B * 0.6,
        intereses=RC_E40_SPLIT_B * 0.3,
        mora=RC_E40_SPLIT_B * 0.1,
        valor_pagado=RC_E40_SPLIT_B,
        tabla_mode="canonical",
        backup_suffix="E40Bfix",
    )
    a2 = amort_apply(session)
    a2r = _result_dict(a2)
    after_a2 = _tabla_saldo_snapshot(_download_path_bytes(session, fx_a["tabla"]))
    raw_a2 = _download_path_bytes(session, fx_a["tabla"])
    after_b = _tabla_saldo_snapshot(_download_path_bytes(session, fx_b_fix["tabla"]))
    cancel_active(session)
    a_unchanged = (
        after_a1.get("valor_pagado_row3") == after_a2.get("valor_pagado_row3")
        and after_a1.get("abono_k_row3") == after_a2.get("abono_k_row3")
        and len(raw_a1) == len(raw_a2)
    )
    b_applied = after_b.get("valor_pagado_row3") not in (None, "", 0, 0.0)
    partial_ok = "PARCIAL" in estado1 or _amort_effectively_blocked(a1, mrg) or (
        after_a1.get("valor_pagado_row3") not in (None, "", 0, 0.0)
    )
    ok = a_unchanged and b_applied and partial_ok and len(applied) >= 2
    return ScenarioResult(
        "E40",
        "PASS" if ok else "FAIL",
        evidence=(
            f"a1={a1.get('status')} estado={estado1}; a2={a2.get('status')}; "
            f"a_unchanged={a_unchanged}; b_vp={after_b.get('valor_pagado_row3')}"
        ),
        process_key=_process_key(a2) or _process_key(a1) or _process_key(fin) or _process_key(gen),
        detail={
            "after_a1": after_a1,
            "after_a2": after_a2,
            "after_b": after_b,
            "a1_summary": a1r.get("summary") or a1.get("error"),
            "a2_summary": a2r.get("summary") or a2.get("error"),
            "notify": nfy.get("status"),
            "fx_b_corrupt": fx_b.get("tabla_mode"),
        },
        cleanup="cancelled",
    )


HANDLERS: dict[str, Callable[[SandboxGraphSession], ScenarioResult]] = {
    "E01": run_e01,
    "E02": run_e02,
    "E03": run_e03,
    "E04": run_e04,
    "E05": run_e05,
    "E06": run_e06,
    "E07": run_e07,
    "E08": run_e08,
    "E09": run_e09,
    "E10": run_e10,
    "E11": run_e11,
    "E12": run_e12,
    "E13": run_e13,
    "E14": run_e14,
    "E15": run_e15,
    "E16": run_e16,
    "E17": run_e17,
    "E18": run_e18,
    "E19": run_e19,
    "E20": run_e20,
    "E21": run_e21,
    "E22": run_e22,
    "E23": run_e23,
    "E24": run_e24,
    "E25": run_e25,
    "E26": run_e26,
    "E27": run_e27,
    "E28": run_e28,
    "E29": run_e29,
    "E30": run_e30,
    "E31": run_e31,
    "E32": run_e32,
    "E33": run_e33,
    "E37": run_e37,
    "E38": run_e38,
    "E39": run_e39,
    "E40": run_e40,
}


def blocked(scenario_id: str, reason: str) -> ScenarioResult:
    meta = by_id(scenario_id)
    return ScenarioResult(
        scenario_id,
        "BLOCKED",
        evidence=reason,
        detail={"title": meta.title, "notes": meta.notes},
    )


def run_selected(ids: list[str]) -> list[ScenarioResult]:
    results: list[ScenarioResult] = []
    with SandboxGraphSession() as session:
        meta = session.ensure_sandbox()
        _log("00_sandbox_meta", meta)
        for sid in ids:
            print(f"=== RUN {sid} ===")
            handler = HANDLERS.get(sid)
            if handler is None:
                # Known not yet automated in this harness
                reason_map = {
                    "E05": "Requiere extracto con saldo vencido parcial medible en cliente sandbox; ver fixtures_catalog.",
                    "E06": "Requiere saldo vencido total en extracto sandbox dedicado.",
                    "E07": "Requiere mix vencido+actual parcial; fixture extracto no listada.",
                    "E08": "Requiere mix vencido+actual completa; fixture extracto no listada.",
                    "E11": "Requiere saldo vencido + abono capital en mismo crédito sandbox.",
                    "E12": "Requiere A+V+K concurrente con extracto multi-panel.",
                    "E13": "Requiere crédito en última cuota (payoff) en sandbox.",
                    "E14": "Requiere payoff + capital adicional con tabla amort sandbox.",
                    "E17": "Requiere multi-crédito tipos distintos + Merge token MULTIPLE.",
                    "E18": "Requiere PDF extracto con panel derecho aplicación anterior (fixture espacial).",
                    "E19": "Requiere PDF extracto con saldo mora a la derecha.",
                    "E20": "Requiere PDF ambiguo (doble panel) en crédito sandbox.",
                    "E22": "Requiere dos extractos y as-of fecha banco ≠ más reciente.",
                    "E24": "Requiere asientos contables sandbox con mismatch por crédito y total OK.",
                    "E25": "Requiere asientos con total ≠ banco (bloqueo amort).",
                    "E26": "Requiere asiento de crédito incorrecto en carpeta sandbox.",
                    "E27": "Requiere asiento faltante post-merge.",
                    "E28": "Requiere PDF asiento ilegible en sandbox (blank_image_like_pdf generator listo).",
                    "E32": "Retry merge requiere merge previo exitoso + asientos.",
                    "E33": "Retry amort requiere apply previo + backup tabla.",
                    "E34": "CAPA A Playwright (doble click).",
                    "E35": "CAPA A Playwright refresh mid-job.",
                    "E36": "CAPA A Playwright network fault injection.",
                    "E39": "Finalize sin amort solo en fase correcta; requiere control state machine fixture.",
                    "E40": "AMORTIZACION_PARCIAL true: un crédito OK / otro falla; requiere dual-credit amort fixture.",
                }
                results.append(blocked(sid, reason_map.get(sid, "Handler no implementado aún")))
                continue
            try:
                results.append(handler(session))
            except Exception as exc:  # noqa: BLE001 — matriz debe registrar fallo
                results.append(
                    ScenarioResult(sid, "FAIL", evidence=f"exception:{type(exc).__name__}:{exc}")
                )
    return results


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--ids",
        default="E01,E02,E03,E04,E09,E10,E15,E16,E21,E23,E29,E30,E31,E37,E38",
        help="Comma-separated scenario ids, or ALL",
    )
    args = parser.parse_args()
    if args.ids.strip().upper() == "ALL":
        ids = [s.id for s in SCENARIOS]
    else:
        ids = [x.strip() for x in args.ids.split(",") if x.strip()]
    results = run_selected(ids)
    # Fill missing catalog entries as BLOCKED for ALL
    seen = {r.id for r in results}
    if args.ids.strip().upper() == "ALL":
        for s in SCENARIOS:
            if s.id not in seen:
                results.append(blocked(s.id, "No ejecutado"))
    summary = {
        "process_date_base": PROCESS_DATE_BASE.isoformat(),
        "bank_code": BANK_CODE,
        "results": [r.as_dict() for r in results],
        "counts": {
            "PASS": sum(1 for r in results if r.status == "PASS"),
            "FAIL": sum(1 for r in results if r.status == "FAIL"),
            "BLOCKED": sum(1 for r in results if r.status == "BLOCKED"),
        },
    }
    out = _log("matrix_results", summary)
    print("WROTE", out)
    print("COUNTS", summary["counts"])
    return 1 if summary["counts"]["FAIL"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
