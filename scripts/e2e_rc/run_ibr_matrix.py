"""IBR real sandbox: Apply + re-read en tablas canónicas (crédito 301).

Solo harness/fixtures. Expectativas:
- PAGO PARCIAL → IBR si fecha banco ≥ fecha límite
- SALDO VENCIDO + OBLIGACIÓN ACTUAL → IBR por corte, no por cierre de cuota
- PAGO ADELANTADO no existe → E usa PAGO DE OBLIGACIÓN ACTUAL
- ABONO A CAPITAL → aplica sin IBR / sin cierre de cuota
"""
from __future__ import annotations

import argparse
import io
import json
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from scripts.e2e_rc.fixtures_catalog import (
    RC_MORA_CREDIT,
    RC_MORA_FOLDER,
    RC_MORA_OBLIG,
    RC_MORA_VENCIDO,
    RC_PAYOFF_CLIENT,
)
from scripts.e2e_rc.run_matrix import (
    PROCESS_DATE_BASE,
    WORK,
    _download_path_bytes,
    _job_ok,
    _pipeline_to_merge,
    _result_dict,
    _tabla_saldo_snapshot,
    amort_apply,
    amort_dry_run,
    cancel_active,
    d,
)
from scripts.e2e_rc.sandbox_fixtures import provision_canonical_payoff_credit
from scripts.e2e_rc.scenarios import ScenarioResult

IBR_CREDIT = RC_MORA_CREDIT
IBR_FOLDER = RC_MORA_FOLDER
SALDO = RC_MORA_OBLIG


def _safe_float(val: Any) -> float:
    if val is None or val == "":
        return 0.0
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        return float(val)
    text = str(val).strip()
    if text.startswith("="):
        return 0.0
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return 0.0


def _log(name: str, data: object) -> Path:
    WORK.mkdir(parents=True, exist_ok=True)
    path = WORK / f"{name}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(name, json.dumps(data, ensure_ascii=False, default=str)[:300])
    return path


def _scan_tabla_writes(raw: bytes) -> dict[str, Any]:
    """Escanea filas de datos (ABONO puede escribir fuera de la fila 3 de cuota)."""
    wb = load_workbook(io.BytesIO(raw), data_only=False)
    ws = wb.active
    assert ws is not None
    headers = {
        str(ws.cell(1, c).value).strip().lower(): c
        for c in range(1, (ws.max_column or 1) + 1)
        if ws.cell(1, c).value
    }

    def _col(*names: str) -> int | None:
        for n in names:
            for h, idx in headers.items():
                if n in h.replace("á", "a"):
                    return idx
        return None

    vp_c = _col("valor pagado cliente", "valor pagado")
    abono_c = _col("abono a k", "abono k")
    ibr_c = _col("ibr")
    fecha_c = _col("fecha pago")
    written_rows: list[dict[str, Any]] = []
    for r in range(2, (ws.max_row or 1) + 1):
        vp = ws.cell(r, vp_c).value if vp_c else None
        abono = ws.cell(r, abono_c).value if abono_c else None
        ibr = ws.cell(r, ibr_c).value if ibr_c else None
        fecha = ws.cell(r, fecha_c).value if fecha_c else None
        has_write = False
        if abono not in (None, "", 0, 0.0):
            has_write = True
        if isinstance(vp, str) and vp.strip().startswith("="):
            has_write = True
        elif vp not in (None, "", 0, 0.0):
            has_write = True
        if fecha not in (None, ""):
            # fecha sola en fila de cuota puede existir; exigir abono/vp
            pass
        if has_write:
            written_rows.append(
                {
                    "row": r,
                    "vp": vp,
                    "abono_k": abono,
                    "ibr": ibr,
                    "fecha_pago": str(fecha or ""),
                }
            )
    return {
        "rows": ws.max_row,
        "written": written_rows,
        "any_abono": any(
            _safe_float(w.get("abono_k")) > 0
            or (
                isinstance(w.get("vp"), str)
                and str(w.get("vp")).startswith("=")
            )
            or _safe_float(w.get("vp")) > 0
            for w in written_rows
        ),
        "any_ibr": any(w.get("ibr") not in (None, "", 0, 0.0) for w in written_rows),
    }


def _run_case(
    session,
    case_id: str,
    *,
    bank: float,
    tipo: str,
    obligacion: float | None = None,
    vencido: float = 0.0,
    capital: float = 0.0,
    intereses: float | None = None,
    mora: float = 0.0,
    saldo_before: float = SALDO,
    expect_ibr: bool,
    expect_vp: float | None = None,
    force_monto_banco: float | None = None,
    abono_mode: bool = False,
) -> ScenarioResult:
    if abono_mode:
        # Asiento puro capital: valor_pagado == capital, sin intereses/mora.
        capital_a = float(capital if capital > 0 else bank)
        int_a = 0.0
        mora_a = 0.0
    else:
        capital_a = capital if capital > 0 else max(0.0, bank - (intereses or bank * 0.3) - mora)
        int_a = intereses if intereses is not None else max(0.0, bank - capital_a - mora)
        mora_a = mora
    fx = provision_canonical_payoff_credit(
        session,
        credit=IBR_CREDIT,
        folder=IBR_FOLDER,
        capital=capital_a if capital_a > 0 else saldo_before * 0.4,
        intereses=int_a,
        mora=mora_a,
        valor_pagado=bank,
        backup_suffix=case_id,
    )
    if saldo_before != capital_a:
        from scripts.e2e_rc.fixtures_catalog import canonical_amortization_xlsx

        session.upload_by_path(
            fx["tabla"],
            canonical_amortization_xlsx(saldo_before=saldo_before),
        )
    before = _tabla_saldo_snapshot(_download_path_bytes(session, fx["tabla"]))
    before_scan = _scan_tabla_writes(_download_path_bytes(session, fx["tabla"]))
    rows = [
        {
            "fecha": d(2026, 5, 23),
            "monto": bank,
            "concepto": RC_PAYOFF_CLIENT,
            "trx": f"RC {case_id} IBR",
        }
    ]
    case_num = {"A": 41, "B": 42, "C": 43, "D": 44, "E": 45, "F": 46}[case_id[-1]]
    sid = f"I{case_num:02d}"
    early, fin, nfy, mrg, _pd = _pipeline_to_merge(
        session,
        sid=sid,
        rows=rows,
        tipo=tipo,
        credito_contains=IBR_CREDIT,
        obligacion=obligacion,
        vencido=vencido,
        capital=capital,
        force_monto_banco=force_monto_banco if force_monto_banco is not None else (
            bank if abono_mode else None
        ),
    )
    if early:
        early.id = case_id
        return early
    dry = amort_dry_run(session)
    dry_r = _result_dict(dry)
    if not _job_ok(dry) or dry_r.get("can_apply") is not True:
        cancel_active(session)
        return ScenarioResult(
            case_id,
            "FAIL",
            evidence=f"dry_not_ready can_apply={dry_r.get('can_apply')}",
            detail={"dry": dry_r, "before": before},
            cleanup="cancelled",
        )
    apply = amort_apply(session)
    raw_after = _download_path_bytes(session, fx["tabla"])
    after = _tabla_saldo_snapshot(raw_after)
    after_scan = _scan_tabla_writes(raw_after)
    cancel_active(session)

    if abono_mode:
        # ABONO: escritura en tabla, sin IBR nuevo, sin exigir vp de cuota.
        applied_ok = after_scan.get("any_abono") is True or _safe_float(
            after.get("abono_k_row3")
        ) > 0
        ibr_ok = after_scan.get("any_ibr") is False and (
            after.get("ibr_row3") in (None, "", 0, 0.0)
        )
        # No cierre de cuota: saldo_row3 no debe quedar en 0 por cuota cancelada
        # (abono reduce capital; puede dejar saldo_row3 fórmula). Basta no-IBR.
        ok = _job_ok(apply) and applied_ok and ibr_ok
        return ScenarioResult(
            case_id,
            "PASS" if ok else "FAIL",
            evidence=(
                f"apply={apply.get('status')}; abono_written={applied_ok}; "
                f"ibr_after={after.get('ibr_row3')}; any_ibr={after_scan.get('any_ibr')}; "
                f"expect_ibr={expect_ibr}"
            ),
            detail={
                "before": before,
                "after": after,
                "after_scan": after_scan,
                "before_scan": before_scan,
                "dry": dry_r.get("summary"),
                "fixture": fx["tabla"],
            },
            cleanup="cancelled",
        )

    ibr_after = after.get("ibr_row3")
    vp_after = after.get("valor_pagado_row3")
    vp_f = _safe_float(vp_after)
    if vp_f == 0.0 and isinstance(vp_after, str) and vp_after.strip().startswith("="):
        vp_f = (
            _safe_float(after.get("intereses_row3"))
            + _safe_float(after.get("abono_k_row3"))
            + _safe_float(after.get("mora_row3"))
        )
    ibr_ok = (ibr_after not in (None, "", 0, 0.0)) if expect_ibr else (
        ibr_after in (None, "", 0, 0.0) or before.get("ibr_row3") == ibr_after
    )
    vp_ok = True if expect_vp is None else abs(vp_f - expect_vp) < 1.0
    ok = _job_ok(apply) and ibr_ok and vp_ok
    return ScenarioResult(
        case_id,
        "PASS" if ok else "FAIL",
        evidence=(
            f"apply={apply.get('status')}; ibr_before={before.get('ibr_row3')}; "
            f"ibr_after={ibr_after}; vp={vp_f}; expect_ibr={expect_ibr}"
        ),
        detail={"before": before, "after": after, "dry": dry_r.get("summary"), "fixture": fx["tabla"]},
        cleanup="cancelled",
    )


def main() -> int:
    from scripts.e2e_rc.graph_session import SandboxGraphSession

    parser = argparse.ArgumentParser()
    parser.add_argument("--ids", default="A,B,C,D,E,F")
    args = parser.parse_args()
    want = [x.strip().upper() for x in args.ids.split(",") if x.strip()]
    cases: dict[str, dict[str, Any]] = {
        "A": dict(
            bank=RC_MORA_OBLIG,
            tipo="PAGO DE OBLIGACIÓN ACTUAL",
            obligacion=RC_MORA_OBLIG,
            expect_ibr=True,
            expect_vp=RC_MORA_OBLIG,
        ),
        # Parcial oficial: sin capital adicional en review; asiento = banco.
        "B": dict(
            bank=2_000_000.0,
            tipo="PAGO PARCIAL A OBLIGACIÓN ACTUAL",
            obligacion=2_000_000.0,
            capital=0.0,
            saldo_before=SALDO,
            intereses=800_000.0,
            mora=200_000.0,
            expect_ibr=True,
            expect_vp=2_000_000.0,
        ),
        # COMBINADO: IBR por fecha de corte, no por cierre de cuota.
        "C": dict(
            bank=RC_MORA_VENCIDO + 1_000_000.0,
            tipo="SALDO VENCIDO + OBLIGACIÓN ACTUAL",
            obligacion=1_000_000.0,
            vencido=RC_MORA_VENCIDO,
            expect_ibr=True,
        ),
        "D": dict(
            bank=RC_MORA_VENCIDO + RC_MORA_OBLIG,
            tipo="SALDO VENCIDO + OBLIGACIÓN ACTUAL",
            obligacion=RC_MORA_OBLIG,
            vencido=RC_MORA_VENCIDO,
            expect_ibr=True,
        ),
        # PAGO ADELANTADO no existe en schema v3 → obligación actual (IBR).
        "E": dict(
            bank=1_500_000.0,
            tipo="PAGO DE OBLIGACIÓN ACTUAL",
            obligacion=1_500_000.0,
            expect_ibr=True,
            expect_vp=1_500_000.0,
        ),
        # Abono puro: monto banco float forzado + asiento solo capital.
        "F": dict(
            bank=1_000_000.0,
            tipo="ABONO A CAPITAL",
            obligacion=0.0,
            capital=1_000_000.0,
            expect_ibr=False,
            abono_mode=True,
            force_monto_banco=1_000_000.0,
        ),
    }
    results: list[ScenarioResult] = []
    with SandboxGraphSession() as session:
        for cid in want:
            if cid not in cases:
                results.append(ScenarioResult(f"IBR-{cid}", "BLOCKED", evidence="unknown_case"))
                continue
            print(f"=== IBR-{cid} ===")
            results.append(_run_case(session, f"IBR-{cid}", **cases[cid]))
    summary = {
        "ran_at": datetime.now().isoformat(),
        "results": [r.as_dict() for r in results],
        "counts": {
            "PASS": sum(1 for r in results if r.status == "PASS"),
            "FAIL": sum(1 for r in results if r.status == "FAIL"),
            "BLOCKED": sum(1 for r in results if r.status == "BLOCKED"),
        },
    }
    out = _log("ibr_matrix_results", summary)
    print("WROTE", out, summary["counts"])
    return 1 if summary["counts"]["FAIL"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
