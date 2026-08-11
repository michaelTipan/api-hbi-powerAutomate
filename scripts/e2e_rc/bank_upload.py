"""Sube plantilla HBI 4 cols (Fecha/Crédito/Concepto/Transacción) al banco sandbox."""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook, load_workbook

from scripts.e2e_rc.path_guard import AUTHORIZED_CLIENTS_BASE, assert_sandbox_mutable_path

BANK_REL = f"{AUTHORIZED_CLIENTS_BASE}/01 CARGA TRANSACCIONES BANCO"
HEADERS = ["Fecha", "Crédito", "Concepto", "Transacción"]


def build_bank_xlsx(rows: list[dict[str, Any]]) -> bytes:
    """rows: fecha (datetime|str), monto (float), concepto, trx."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Movimientos"
    for c, h in enumerate(HEADERS, start=1):
        ws.cell(1, c).value = h
    for i, row in enumerate(rows):
        r = 2 + i
        fecha = row["fecha"]
        ws.cell(r, 1).value = fecha
        ws.cell(r, 2).value = float(row["monto"])
        ws.cell(r, 3).value = str(row["concepto"])
        ws.cell(r, 4).value = str(row["trx"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload_bank_bogota(session, rows: list[dict[str, Any]], *, filename: str = "BANCO_BOGOTA.xlsx") -> dict[str, Any]:
    folder_path = assert_sandbox_mutable_path(BANK_REL)
    folder_id = session.walk(folder_path)
    item = next((i for i in session.children(folder_id) if i.get("name") == filename), None)
    if not item:
        raise FileNotFoundError(filename)
    raw = build_bank_xlsx(rows)
    full = f"{folder_path}/{filename}"
    session.upload_item(str(item["id"]), raw, path_for_guard=full)
    return {"path": full, "item_id": item["id"], "rows": len(rows), "bytes": len(raw)}


def d(year: int, month: int, day: int) -> datetime:
    return datetime(year, month, day)
