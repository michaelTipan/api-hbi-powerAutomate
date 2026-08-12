"""Catálogo de fixtures sandbox PRUEBAS para matriz E01–E40 (CAPA B).

No escribe SharePoint por sí solo: documenta requisitos y genera artefactos
locales (PDF mínimos, plans de upload) para que el harness/E2E provisione
bajo ``AUTHORIZED_CLIENTS_BASE`` cuando corra contra sandbox.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from pypdf import PdfWriter

from scripts.e2e_rc.path_guard import AUTHORIZED_CLIENTS_BASE

# Destinatarios permitidos en Notify E2E (nunca @hbi.com.co reales).
SANDBOX_NOTIFY_ALLOWLIST = frozenset(
    {
        "herramientas.jsakedev@gmail.com",
    }
)
HBI_DOMAIN_BLOCK = "@hbi.com.co"

# Path canónico sandbox (overlay sandbox-ui-enabled). Solo mutar bajo PRUEBAS.
CORREOS_XLSX_REL = (
    f"{AUTHORIZED_CLIENTS_BASE}/02 VALIDACION PAGOS/"
    "02 CONTROL OPERATIVO/CORREOS.xlsx"
)

# Subcarpeta bajo ASIENTOS: merge solo lista PDF en el nivel actual.
ASIENTO_QUARANTINE_SUBFOLDER = "_RC_CUARENTENA"

# Cliente/crédito E2E canónico (harness GEOEXCON).
E15_CLIENT = "GEOEXCON"
E15_CREDIT = "231"
E15_ASIENTO_FILENAME = "Asiento RC-E15 PAGO TOTAL GEOEXCON CRED 231.pdf"

# E16: segundo crédito ACTIVO (254 en sandbox es TERMINADO → Generate lo omite).
E16_CLIENT = "GEOEXCON"
E16_CREDIT_A = "231"
E16_CREDIT_B = "299"
E16_CREDIT_B_FOLDER = "CREDITO # 299"
E16_BANK_TOTAL = 25_000_000.0
E16_SPLIT_A = 19_102_163.0
E16_SPLIT_B = 5_897_837.0  # A+B == BANK_TOTAL
E16_ASIENTO_B_FILENAME = "Asiento RC-E16 GEOEXCON CRED 299.pdf"
E16_TABLA_B_FILENAME = "Tabla de amortizacion GEOEXCON CRED 299.xlsx"

# Crédito E2E dedicado a mora / as-of / bloqueos (no reutilizar 231/299).
RC_MORA_CLIENT = "GEOEXCON"
RC_MORA_CREDIT = "301"
RC_MORA_FOLDER = "CREDITO # 301"
RC_MORA_OBLIG = 5_000_000.0
RC_MORA_VENCIDO = 2_000_000.0


@dataclass(frozen=True)
class FixtureNeed:
    scenario_id: str
    kind: str  # extract | asiento | tabla | correos | multi_credit | dual_amort | capa_a
    detail: str
    client_hint: str = ""
    credit_hint: str = ""
    local_generator: str = ""  # nombre helper en este módulo, si aplica


# Escenarios que hoy quedan BLOCKED sin fixture sandbox.
CRITICAL_FIXTURE_GAPS: tuple[FixtureNeed, ...] = (
    FixtureNeed("E05", "extract", "Saldo vencido parcial medible", "AGRECAR|GEOEXCON", ""),
    FixtureNeed("E06", "extract", "Saldo vencido total en extracto", "", ""),
    FixtureNeed("E07", "extract", "Vencido total + obligación parcial", "", ""),
    FixtureNeed("E08", "extract", "Vencido total + obligación completa", "", ""),
    FixtureNeed("E11", "extract", "Vencido + abono capital mismo crédito", "", ""),
    FixtureNeed("E12", "extract", "A+V+K concurrente multi-panel", "", ""),
    FixtureNeed("E13", "tabla", "Crédito en última cuota (payoff real)", "", ""),
    FixtureNeed("E14", "tabla", "Payoff + capital adicional", "", ""),
    FixtureNeed(
        "E15",
        "asiento",
        "Asiento parseable GEOEXCON/231 para dry-run PAYOFF_NOT_ACHIEVED",
        "GEOEXCON",
        "231",
        "parseable_asiento_pdf",
    ),
    FixtureNeed(
        "E16",
        "multi_credit",
        "Un pago → dos créditos activos reconciliables (231+299; 254 es TERMINADO)",
        "GEOEXCON",
        "231+299",
        "provision_e16_second_active_credit",
    ),
    FixtureNeed("E17", "multi_credit", "Tipos distintos → Merge MULTIPLE", "GEOEXCON", ""),
    FixtureNeed("E18", "extract", "Panel derecho aplicación anterior", "", "", "minimal_text_pdf"),
    FixtureNeed("E19", "extract", "Saldo mora a la derecha", "", "", "minimal_text_pdf"),
    FixtureNeed("E20", "extract", "PDF ambiguo doble panel", "", "", "minimal_text_pdf"),
    FixtureNeed("E22", "extract", "Dos extractos; as-of ≠ más reciente", "", ""),
    FixtureNeed("E23", "multi_credit", "Dos pagos mismo cliente fechas distintas", "GEOEXCON", ""),
    FixtureNeed("E24", "asiento", "Mismatch por crédito; total banco OK", "", ""),
    FixtureNeed("E25", "asiento", "Total asientos ≠ banco (bloqueo amort)", "GEOEXCON", "231"),
    FixtureNeed("E26", "asiento", "Asiento de crédito incorrecto", "", ""),
    FixtureNeed("E27", "asiento", "Asiento faltante post-merge", "", ""),
    FixtureNeed("E28", "asiento", "PDF asiento ilegible (sin texto)", "", "", "blank_image_like_pdf"),
    FixtureNeed("E30", "control", "Retry finalize sobre FINALIZADO estable", "", ""),
    FixtureNeed("E31", "correos", "CORREOS.xlsx solo recipients sandbox", "", "", "assert_sandbox_notify_recipients"),
    FixtureNeed("E32", "asiento", "Retry merge + asientos previos", "", ""),
    FixtureNeed("E33", "tabla", "Retry amort + backup tabla", "", ""),
    FixtureNeed("E34", "capa_a", "Playwright doble click", "", ""),
    FixtureNeed("E35", "capa_a", "Playwright refresh mid-job", "", ""),
    FixtureNeed("E36", "capa_a", "Playwright network fault", "", ""),
    FixtureNeed("E38", "control", "Cancel no permitido post-fase avanzada", "", ""),
    FixtureNeed("E39", "control", "Finalize sin amort en fase correcta", "", ""),
    FixtureNeed("E40", "dual_amort", "AMORTIZACION_PARCIAL true dual-crédito", "", ""),
)


def clients_base() -> str:
    return AUTHORIZED_CLIENTS_BASE


def asientos_folder_rel(client: str, credit_folder: str, credit_num: str) -> str:
    """Ruta relativa PRUEBAS a carpeta ASIENTOS (path-guard la valida en upload)."""
    return (
        f"{AUTHORIZED_CLIENTS_BASE}/{client}/{credit_folder}/"
        f"ASIENTOS CONTABLES CRED {credit_num}"
    )


def e15_asientos_folder_rel() -> str:
    return asientos_folder_rel(E15_CLIENT, f"CREDITO # {E15_CREDIT}", E15_CREDIT)


def minimal_text_pdf(*, title: str = "RC fixture") -> bytes:
    """PDF con texto seleccionable (útil como asiento/extracto mínimo)."""
    from reportlab.pdfgen import canvas  # type: ignore

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 720, title)
    c.drawString(72, 700, "544111100505 1000000.00")
    c.drawString(72, 680, "544113410519 1000000.00")
    c.save()
    return buf.getvalue()


def parseable_asiento_pdf(
    *,
    credit: str = E15_CREDIT,
    valor_pagado: float = 19_102_163.0,
    capital: float = 5_000_000.0,
    intereses: float = 10_000_000.0,
    mora: float = 4_102_163.0,
    title: str = "RC-E15 asiento parseable PAGO TOTAL",
    fecha: str = "22/05/2026",
    comprobante: str = "9915",
) -> bytes:
    """Asiento E2E parseable (texto seleccionable, cuentas ERP) sin PII.

    Capital deliberadamente < saldo típico de tabla → dry-run debe llegar a
    ``PAYOFF_NOT_ACHIEVED`` (no quedarse en ``ACCOUNTING_PARSE_FAILED``).
    """
    from reportlab.pdfgen import canvas  # type: ignore

    # Normalizar suma: si el caller pasa valor_pagado, alinear componentes.
    components = capital + intereses + mora
    if abs(components - valor_pagado) > 0.02:
        # Mantener capital bajo (payoff fail) y repartir el resto en intereses.
        intereses = max(0.0, valor_pagado - capital - mora)

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    y = 750
    # Un solo formato (code-before-amount): duplicar Linea sumaría montos 2×.
    lines = [
        title,
        f"Comprobante {comprobante} Fecha {fecha}",
        f"No.Rad. {credit}",
        f"544111100505 {valor_pagado:,.2f}",
        f"544113410519 {capital:,.2f}",
        f"544113430501 {intereses:,.2f}",
        f"544141502030 {mora:,.2f}",
    ]
    for line in lines:
        c.drawString(72, y, line[:110])
        y -= 18
    c.save()
    return buf.getvalue()


def minimal_amortization_xlsx(
    *,
    client: str = E16_CLIENT,
    credit: str = E16_CREDIT_B,
    saldo: float = E16_SPLIT_B,
) -> bytes:
    """Tabla mínima para señal operativa de crédito E2E (sin PII)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Amortizacion"
    headers = [
        "Cuota",
        "Fecha",
        "Capital",
        "Intereses",
        "Cuota total",
        "Saldo",
        "Cliente",
        "Credito",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(1, col).value = h
    ws.cell(2, 1).value = 1
    ws.cell(2, 2).value = "2026-05-23"
    ws.cell(2, 3).value = round(saldo * 0.4, 2)
    ws.cell(2, 4).value = round(saldo * 0.6, 2)
    ws.cell(2, 5).value = round(saldo, 2)
    ws.cell(2, 6).value = round(saldo, 2)
    ws.cell(2, 7).value = client
    ws.cell(2, 8).value = credit
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def minimal_extract_pdf(
    *,
    credit: str,
    fecha_limite: str = "23/05/2026",
    valor_obligacion: float,
    client: str = E16_CLIENT,
) -> bytes:
    """Extracto sintético mínimo (texto seleccionable) sin PII real."""
    from reportlab.pdfgen import canvas  # type: ignore

    def _co(n: float) -> str:
        # 1234567.89 -> 1.234.567,89
        s = f"{n:,.2f}"
        return s.replace(",", "X").replace(".", ",").replace("X", ".")

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    y = 750
    for line in (
        f"EXTRACTO DE OBLIGACION No. {credit}",
        f"Cliente {client}",
        f"Fecha limite de pago: {fecha_limite}",
        f"Valor de la obligacion actual: {_co(valor_obligacion)}",
        f"Total a pagar: {_co(valor_obligacion)}",
    ):
        c.drawString(72, y, line[:110])
        y -= 18
    c.save()
    return buf.getvalue()


def _co_latin(n: float) -> str:
    s = f"{n:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def spatial_extract_pdf(
    *,
    credit: str,
    fecha_limite: str = "23/05/2026",
    valor_obligacion: float,
    client: str = E16_CLIENT,
    right_role: str = "VACIO",
    right_amount: float = 0.0,
    left_intereses_mora: float | None = None,
) -> bytes:
    """Extracto con columnas LEFT/RIGHT (umbral 55% del ancho) para el parser espacial."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas  # type: ignore

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    y = 750
    for line in (
        f"EXTRACTO DE OBLIGACION No. {credit}",
        f"Cliente {client}",
        f"Fecha limite de pago: {fecha_limite}",
        f"Valor de la obligacion actual: {_co_latin(valor_obligacion)}",
        f"Total a pagar: {_co_latin(valor_obligacion)}",
    ):
        c.drawString(50, y, line[:90])
        y -= 18
    if left_intereses_mora is not None:
        c.drawString(50, y, f"Intereses de mora: {_co_latin(left_intereses_mora)}")
    role = (right_role or "VACIO").upper()
    if role == "SALDO_VENCIDO":
        c.drawString(400, 720, f"SALDO MORA $ {_co_latin(right_amount)}")
    elif role == "APLICACION_ANTERIOR":
        c.drawString(400, 720, f"APLICACION ANTERIOR $ {_co_latin(right_amount)}")
    elif role == "AMBIGUO":
        c.drawString(400, 720, f"SALDO MORA $ {_co_latin(right_amount)}")
        c.drawString(400, 700, f"APLICACION ANTERIOR $ {_co_latin(right_amount)}")
    c.save()
    return buf.getvalue()


def blank_image_like_pdf() -> bytes:
    """PDF sin texto extraíble (asiento ilegible E28)."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def build_sandbox_correos_xlsx(
    *,
    emisor: str = "herramientas.jsakedev@gmail.com",
    receptores: list[str] | None = None,
) -> bytes:
    """CORREOS.xlsx mínimo con EMISOR/RECEPTORES solo allowlist (sin @hbi.com.co)."""
    recs = assert_sandbox_notify_recipients(
        list(receptores) if receptores is not None else ["herramientas.jsakedev@gmail.com"]
    )
    sender = str(emisor or "").strip().lower()
    if not sender:
        raise ValueError("correos_emisor_empty")
    if HBI_DOMAIN_BLOCK in sender:
        raise ValueError(f"correos_emisor_blocked_hbi:{sender}")

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "CORREOS"
    ws.cell(1, 1).value = "EMISOR"
    ws.cell(1, 2).value = "RECEPTORES"
    ws.cell(2, 1).value = sender
    ws.cell(2, 2).value = "; ".join(recs)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rewrite_correos_recipients_bytes(
    raw: bytes,
    *,
    sandbox_to: str = "herramientas.jsakedev@gmail.com",
) -> bytes:
    """Reescribe RECEPTORES de un CORREOS.xlsx existente → solo allowlist.

    Conserva EMISOR si no es @hbi.com.co; si lo es, lo sustituye por sandbox_to.
    No toca otras hojas/columnas innecesarias.
    """
    from openpyxl import load_workbook

    allow = assert_sandbox_notify_recipients([sandbox_to])
    replacement = allow[0]
    wb = load_workbook(io.BytesIO(raw))
    changed = False
    for ws in wb.worksheets:
        header_row = None
        col_em = col_rec = None
        max_col = ws.max_column or 1
        max_scan = min(ws.max_row or 1, 80)
        for row_idx in range(1, max_scan + 1):
            cells = [
                str(ws.cell(row_idx, c).value or "").strip().upper()
                for c in range(1, max_col + 1)
            ]
            if "EMISOR" in cells and ("RECEPTORES" in cells or "RECEPTOR" in cells):
                header_row = row_idx
                col_em = cells.index("EMISOR") + 1
                if "RECEPTORES" in cells:
                    col_rec = cells.index("RECEPTORES") + 1
                else:
                    col_rec = cells.index("RECEPTOR") + 1
                break
        if header_row is None or col_em is None or col_rec is None:
            continue
        last = ws.max_row or header_row
        for r in range(header_row + 1, last + 1):
            em_val = ws.cell(r, col_em).value
            if em_val is not None and HBI_DOMAIN_BLOCK in str(em_val).lower():
                ws.cell(r, col_em).value = replacement
                changed = True
            rec_val = ws.cell(r, col_rec).value
            if rec_val is None or not str(rec_val).strip():
                continue
            text = str(rec_val)
            if HBI_DOMAIN_BLOCK in text.lower() or replacement not in text.lower():
                ws.cell(r, col_rec).value = replacement
                changed = True
        # Garantizar al menos una fila de datos.
        if last <= header_row:
            ws.cell(header_row + 1, col_em).value = replacement
            ws.cell(header_row + 1, col_rec).value = replacement
            changed = True
    if not changed:
        # Forzar fila 2 si estructura mínima.
        ws0 = wb.worksheets[0]
        ws0.cell(1, 1).value = "EMISOR"
        ws0.cell(1, 2).value = "RECEPTORES"
        ws0.cell(2, 1).value = replacement
        ws0.cell(2, 2).value = replacement
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def assert_sandbox_notify_recipients(emails: list[str]) -> list[str]:
    """Fail-closed: solo allowlist sandbox; bloquea cualquier @hbi.com.co."""
    cleaned: list[str] = []
    for raw in emails:
        addr = str(raw or "").strip().lower()
        if not addr:
            continue
        if HBI_DOMAIN_BLOCK in addr:
            raise ValueError(f"notify_recipient_blocked_hbi:{addr}")
        if addr not in SANDBOX_NOTIFY_ALLOWLIST:
            raise ValueError(f"notify_recipient_not_sandbox:{addr}")
        cleaned.append(addr)
    if not cleaned:
        raise ValueError("notify_recipients_empty")
    return cleaned


def gaps_by_scenario() -> dict[str, list[FixtureNeed]]:
    out: dict[str, list[FixtureNeed]] = {}
    for need in CRITICAL_FIXTURE_GAPS:
        out.setdefault(need.scenario_id, []).append(need)
    return out


def catalog_as_dicts() -> list[dict[str, Any]]:
    return [
        {
            "scenario_id": n.scenario_id,
            "kind": n.kind,
            "detail": n.detail,
            "client_hint": n.client_hint,
            "credit_hint": n.credit_hint,
            "local_generator": n.local_generator,
            "clients_base": AUTHORIZED_CLIENTS_BASE,
            "correos_xlsx_rel": CORREOS_XLSX_REL,
        }
        for n in CRITICAL_FIXTURE_GAPS
    ]
