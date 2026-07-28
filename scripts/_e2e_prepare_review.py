"""Rellena montos contables + Validar selectivo + Procesar=SI y sube revisión."""
from __future__ import annotations

import base64
import io
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
WORK = Path(r"D:\CMC\HBI_Capital\_work")
REV_NAME = "validacion_pagos_banco_bogota_2026-07-27.xlsx"
REV_DIR = (
    "INFORMACION CREDITOS-CLIENTES/"
    "02 COMWARE AUTOMATIZACION - INFORMACION CREDITOS CLIENTES/"
    "01 VALIDACION PAGOS/02 REVISION"
)
SKIP_CLIENTES = {
    "ACIMOR",
    "MINCIVIL",
    "INVERSIONES Y PROYECTOS MIOS",
    "CLIENTE_INEXISTENTE_XYZ",
}


def children(c, drive, fid="root"):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c, drive, parts):
    cur = "root"
    for p in parts:
        hit = next(i for i in children(c, drive, cur) if i.get("name") == p)
        cur = hit["id"]
    return cur


def download(c, drive, item_id):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        timeout=180,
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def find_id_pago_header_row(ws) -> int:
    for r in range(1, 15):
        if str(ws.cell(r, 1).value or "").strip() == "ID Pago":
            return r
    raise ValueError("header_not_found")


def header_map(ws, hr: int) -> dict[str, int]:
    return {
        str(ws.cell(hr, c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(hr, c).value is not None and str(ws.cell(hr, c).value).strip()
    }


def fnum(v) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def main() -> None:
    with httpx.Client(timeout=180) as c:
        drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]
        folder_id = walk(c, drive, REV_DIR.split("/"))
        rev = next(i for i in children(c, drive, folder_id) if i.get("name") == REV_NAME)
        raw = download(c, drive, rev["id"])
        wb = load_workbook(io.BytesIO(raw))
        ws = wb["Distribucion_Pagos"]
        hr = find_id_pago_header_row(ws)
        hm = header_map(ws, hr)
        col_id = hm["ID Pago"]
        col_cli = hm["Cliente"]
        col_cred = hm["Crédito"]
        col_monto = hm["Monto banco"]
        col_ext = hm["Valor extracto"]
        col_aplicar = hm["Aplicar a extracto"]
        col_mora = hm["Mora a aplicar"]
        col_cap = hm["Abono a capital"]
        col_otros = hm["Otros valores"]
        col_estado = hm["Estado Pago"]
        col_val = hm["Validar Pago"]
        col_obs = hm["Observación"]
        col_tipo = hm.get("TipoAplicacionOriginal")

        for r in range(hr + 1, ws.max_row + 1):
            if ws.cell(r, col_id).value:
                ws.cell(r, col_val).value = "NO"

        by_id: dict[str, list[int]] = {}
        for r in range(hr + 1, ws.max_row + 1):
            pid = str(ws.cell(r, col_id).value or "").strip()
            if not pid:
                continue
            cli = str(ws.cell(r, col_cli).value or "").strip()
            if any(s.lower() in cli.lower() for s in SKIP_CLIENTES):
                # dejar NO; si NORMAL exigir obs
                if str(ws.cell(r, col_estado).value or "").upper() == "NORMAL":
                    ws.cell(r, col_obs).value = "E2E: no validar (cliente de prueba negativo)"
                continue
            by_id.setdefault(pid, []).append(r)

        for pid, rows in by_id.items():
            monto = None
            for r in rows:
                monto = fnum(ws.cell(r, col_monto).value)
                if monto is not None:
                    break
            if monto is None:
                continue
            best_r, best_score = None, None
            for r in rows:
                e = fnum(ws.cell(r, col_ext).value)
                if e is None:
                    continue
                score = abs(monto - e)
                if best_score is None or score < best_score:
                    best_score, best_r = score, r
            if best_r is None:
                continue

            # Marcar SI y rellenar contable
            ws.cell(best_r, col_val).value = "SI"
            ext = fnum(ws.cell(best_r, col_ext).value) or 0.0
            tipo = str(ws.cell(best_r, col_tipo).value or "").upper() if col_tipo else ""
            if tipo == "PAGO Y ABONO CAPITAL":
                cuota = min(ext, monto)
                capital = round(monto - cuota, 2)
                ws.cell(best_r, col_aplicar).value = cuota
                ws.cell(best_r, col_mora).value = 0
                ws.cell(best_r, col_cap).value = capital
                ws.cell(best_r, col_otros).value = 0
            else:
                # PAGO: aplicar el monto banco completo (saldo cero)
                ws.cell(best_r, col_aplicar).value = monto
                ws.cell(best_r, col_mora).value = 0
                ws.cell(best_r, col_cap).value = 0
                ws.cell(best_r, col_otros).value = 0
            print(
                "SI",
                pid,
                ws.cell(best_r, col_cred).value,
                "aplicar",
                ws.cell(best_r, col_aplicar).value,
                "cap",
                ws.cell(best_r, col_cap).value,
            )

            for r in rows:
                if r == best_r:
                    continue
                est = str(ws.cell(r, col_estado).value or "").upper()
                if est == "NORMAL":
                    ws.cell(r, col_obs).value = "E2E: otro crédito del mismo pago; no validar"

        # Abonos 265
        if "Distribucion_Abonos" in wb.sheetnames:
            wa = wb["Distribucion_Abonos"]
            hr_a = find_id_pago_header_row(wa)
            hm_a = header_map(wa, hr_a)
            col_fecha = hm_a.get("Fecha banco")
            # Normalizar todas las fechas abono a datetime (compat. coerce antiguo)
            if col_fecha:
                for r in range(hr_a + 1, wa.max_row + 1):
                    fv = wa.cell(r, col_fecha).value
                    if isinstance(fv, str) and fv.strip():
                        try:
                            wa.cell(r, col_fecha).value = datetime.strptime(fv.strip()[:10], "%Y-%m-%d")
                        except ValueError:
                            pass

            by_ab: dict[str, list[int]] = {}
            for r in range(hr_a + 1, wa.max_row + 1):
                pid = str(wa.cell(r, hm_a["ID Pago"]).value or "").strip()
                if pid:
                    by_ab.setdefault(pid, []).append(r)
                    wa.cell(r, hm_a["Validar Abono"]).value = "NO"
            for pid, rows in by_ab.items():
                monto = None
                fecha = None
                for r in rows:
                    if monto is None:
                        monto = fnum(wa.cell(r, hm_a["Monto banco"]).value)
                    if fecha is None and col_fecha:
                        fecha = wa.cell(r, col_fecha).value
                for r in rows:
                    cred = str(wa.cell(r, hm_a["Crédito"]).value or "")
                    if "265" not in cred:
                        continue
                    if monto is not None:
                        wa.cell(r, hm_a["Monto banco"]).value = monto
                    if col_fecha and fecha is not None:
                        if isinstance(fecha, str):
                            try:
                                fecha_dt = datetime.strptime(str(fecha)[:10], "%Y-%m-%d")
                            except ValueError:
                                fecha_dt = fecha
                        elif isinstance(fecha, date) and not isinstance(fecha, datetime):
                            fecha_dt = datetime(fecha.year, fecha.month, fecha.day)
                        else:
                            fecha_dt = fecha
                        wa.cell(r, col_fecha).value = fecha_dt
                    wa.cell(r, hm_a["Validar Abono"]).value = "SI"
                    print(
                        "abono SI",
                        pid,
                        cred,
                        monto,
                        wa.cell(r, col_fecha).value if col_fecha else None,
                    )

        wc = wb["Control"]
        for r in range(1, wc.max_row + 1):
            if str(wc.cell(r, 1).value or "").strip().lower() == "procesar":
                wc.cell(r, 2).value = "SI"

        # Estado debe seguir EN_REVISION para finalize
        for r in range(1, wc.max_row + 1):
            if str(wc.cell(r, 1).value or "").strip().lower() == "estado":
                print("Control Estado=", wc.cell(r, 2).value)

        buf = io.BytesIO()
        wb.save(buf)
        out = buf.getvalue()
        (WORK / "review_prepared.xlsx").write_bytes(out)
        up = c.put(
            f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
            params={"item_id": rev["id"]},
            json={"content_base64": base64.b64encode(out).decode("ascii")},
            timeout=180,
        )
        print("upload", up.status_code)
        up.raise_for_status()
        print("OK")


if __name__ == "__main__":
    main()
