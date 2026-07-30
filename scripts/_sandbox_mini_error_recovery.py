"""Caso critico: Finalize con Procesar=SI y hoja Errores abierta + cancel recovery."""
from __future__ import annotations

import base64
import io
import json
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
WORK = Path(r"D:\CMC\HBI_Capital\_work\sandbox_mini")
PROCESS_DATE = "2026-07-30"
BANK_CODE = "banco_bogota"
CLIENTS = (
    "INFORMACION CREDITOS-CLIENTES/"
    "03 COMWARE PRUEBAS- INFORMACION CREDITOS CLIENTES"
)
BANK_DIR = f"{CLIENTS}/01 CARGA TRANSACCIONES BANCO"
REV_DIR = f"{CLIENTS}/02 VALIDACION PAGOS/01 REVISION"
BANK_NAME = "BANCO_BOGOTA.xlsx"

ROWS = [
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 10_147_965.0,
        "concepto": "A&M CONSTRUCOL",
        "tipo": "PAGO",
        "trx": "ERRREC PAGO OK 215",
    },
    {
        "fecha": datetime(2026, 7, 20),
        "monto": 1_234_567.0,
        "concepto": "CLIENTE_FANTASMA_MINI",
        "tipo": "PAGO",
        "trx": "ERRREC FANTASMA",
    },
]


def api_key() -> str:
    env = Path(r"D:\CMC\HBI_Capital\api-hbi-powerAutomate.env").read_text(encoding="utf-8")
    return re.search(r"^API_HTTP_KEY=(.+)$", env, re.M).group(1).strip()


def log(name: str, data: object) -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    (WORK / f"{name}.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    print(name, json.dumps(data, ensure_ascii=False, default=str)[:500])


def poll(c: httpx.Client, url: str, timeout_s: int = 1800) -> dict:
    t0 = time.time()
    while time.time() - t0 < timeout_s:
        body = c.get(url, timeout=90).json()
        st = body.get("status")
        err = body.get("error")
        err_txt = err if isinstance(err, str) else json.dumps(err, ensure_ascii=False) if err else ""
        print(time.strftime("%H:%M:%S"), "status", st, err_txt[:120])
        if st in ("completed", "failed", "error"):
            return body
        time.sleep(8)
    raise TimeoutError(url)


def children(c, drive, fid="root"):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c, drive, parts):
    cur = "root"
    for p in parts:
        hit = next(i for i in children(c, drive, cur) if i.get("name") == p)
        cur = hit["id"]
    return cur


def download(c, drive, item_id):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        timeout=180,
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def upload_bytes(c, drive, item_id, raw):
    r = c.put(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        json={"content_base64": base64.b64encode(raw).decode("ascii")},
        timeout=180,
    )
    r.raise_for_status()


def queue_and_poll(c, path, payload, tag):
    r = c.post(f"{BASE}{path}", json=payload, timeout=60)
    log(f"{tag}_queue", {"status": r.status_code, "body": r.json() if r.content else None})
    body = r.json()
    done = poll(c, f"{BASE}/graph/sharepoint/payment-validation/jobs/{body['job_id']}")
    log(f"{tag}_job", done)
    return done


def set_procesar_si(raw: bytes) -> bytes:
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Control"]
    # Buscar celda con etiqueta Procesar y setear valor SI en columna adyacente
    found = False
    for row in ws.iter_rows(min_row=1, max_row=40, max_col=6):
        for cell in row:
            val = str(cell.value or "").strip().lower()
            if val in {"procesar", "procesar?"}:
                # valor suele estar a la derecha
                right = ws.cell(cell.row, cell.column + 1)
                right.value = "SI"
                found = True
                print("set_procesar", cell.coordinate, "->", right.coordinate, "SI")
                break
        if found:
            break
    if not found:
        # fallback tipico fila/col conocida en plantillas previas
        for r in range(1, 30):
            for col in range(1, 5):
                if str(ws.cell(r, col).value or "").strip().lower().startswith("procesar"):
                    ws.cell(r, col + 1).value = "SI"
                    found = True
                    print("set_procesar_fallback", r, col + 1)
                    break
            if found:
                break
    if not found:
        raise RuntimeError("No se encontro celda Procesar en Control")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> None:
    key = api_key()
    out: dict = {}
    with httpx.Client(timeout=180, headers={"X-API-Key": key}) as c:
        drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]

        # cancel limpio
        queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/cancel-active-process/queue",
            {"bank_code": BANK_CODE},
            "20_cancel_pre",
        )

        bank_folder = walk(c, drive, BANK_DIR.split("/"))
        bank_item = next(i for i in children(c, drive, bank_folder) if i.get("name") == BANK_NAME)
        raw = download(c, drive, bank_item["id"])
        wb = load_workbook(io.BytesIO(raw))
        ws = wb.active
        for r in range(4, ws.max_row + 1):
            for col in range(1, 6):
                ws.cell(r, col).value = None
        for i, row in enumerate(ROWS):
            rr = 4 + i
            ws.cell(rr, 1).value = row["fecha"]
            ws.cell(rr, 2).value = float(row["monto"])
            ws.cell(rr, 3).value = row["concepto"]
            ws.cell(rr, 4).value = row["tipo"]
            ws.cell(rr, 5).value = row["trx"]
        buf = io.BytesIO()
        wb.save(buf)
        upload_bytes(c, drive, bank_item["id"], buf.getvalue())

        gen = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "21_generate_with_errors",
        )
        out["generate"] = {
            "status": gen.get("status"),
            "errores": ((gen.get("result") or {}).get("summary") or {}).get("errores"),
            "validation_file": (gen.get("result") or {}).get("validation_file"),
        }

        rev_id = walk(c, drive, REV_DIR.split("/"))
        rev_files = [i for i in children(c, drive, rev_id) if str(i.get("name", "")).endswith(".xlsx")]
        target = sorted(rev_files, key=lambda i: i.get("lastModifiedDateTime") or "", reverse=True)[0]
        rev_raw = download(c, drive, target["id"])
        patched = set_procesar_si(rev_raw)
        upload_bytes(c, drive, target["id"], patched)
        (WORK / ("patched_" + target["name"])).write_bytes(patched)
        log("22_procesar_si", {"file": target["name"]})

        fin = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/finalize/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "23_finalize_open_errors",
        )
        out["finalize_open_errors"] = {
            "status": fin.get("status"),
            "error": fin.get("error"),
            "user_message": fin.get("user_message") or (fin.get("error") or {}).get("user_message")
            if isinstance(fin.get("error"), dict)
            else fin.get("user_message"),
            "error_code": (fin.get("error") or {}).get("error_code")
            if isinstance(fin.get("error"), dict)
            else None,
        }

        # Intento regenerar sin cancel (debe bloquear o recreate same day)
        gen_block = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "24_generate_while_active",
        )
        out["generate_while_active"] = {
            "status": gen_block.get("status"),
            "already_generated": (gen_block.get("result") or {}).get("already_generated")
            if isinstance(gen_block.get("result"), dict)
            else None,
            "error": gen_block.get("error"),
            "file_action": (gen_block.get("result") or {}).get("file_action")
            if isinstance(gen_block.get("result"), dict)
            else None,
        }

        cancel = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/cancel-active-process/queue",
            {"bank_code": BANK_CODE},
            "25_cancel_after_stuck",
        )
        out["cancel_after_stuck"] = {
            "status": cancel.get("status"),
            "result": cancel.get("result"),
        }

        gen2 = queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/generate/queue",
            {"bank_code": BANK_CODE, "process_date": PROCESS_DATE},
            "26_generate_after_cancel",
        )
        out["generate_after_cancel"] = {
            "status": gen2.get("status"),
            "already_generated": (gen2.get("result") or {}).get("already_generated")
            if isinstance(gen2.get("result"), dict)
            else None,
            "errores": ((gen2.get("result") or {}).get("summary") or {}).get("errores"),
            "file_action": (gen2.get("result") or {}).get("file_action")
            if isinstance(gen2.get("result"), dict)
            else None,
        }

        # dejar limpio
        queue_and_poll(
            c,
            "/graph/sharepoint/payment-validation/cancel-active-process/queue",
            {"bank_code": BANK_CODE},
            "27_cancel_final",
        )

    log("28_error_recovery_results", out)
    print("DONE")


if __name__ == "__main__":
    main()
