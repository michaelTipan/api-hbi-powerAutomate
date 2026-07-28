"""Verifica fecha en revisión SP y coerce local."""
from __future__ import annotations

import base64
import io
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
REV_DIR = (
    "INFORMACION CREDITOS-CLIENTES/"
    "02 COMWARE AUTOMATIZACION - INFORMACION CREDITOS CLIENTES/"
    "01 VALIDACION PAGOS/02 REVISION"
)
NAME = "validacion_pagos_banco_bogota_2026-07-27.xlsx"


def children(c, drive, fid="root"):
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c, drive, parts):
    cur = "root"
    for p in parts:
        hit = next(i for i in children(c, drive, cur) if i.get("name") == p)
        cur = hit["id"]
    return cur


def coerce(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    text = str(v).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


with httpx.Client(timeout=180) as c:
    drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]
    fid = walk(c, drive, REV_DIR.split("/"))
    rev = next(i for i in children(c, drive, fid) if i.get("name") == NAME)
    raw = base64.b64decode(
        c.get(
            f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
            params={"item_id": rev["id"]},
        ).json()["content_base64"]
    )
    Path(r"D:\CMC\HBI_Capital\_work\review_from_sp.xlsx").write_bytes(raw)
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["Distribucion_Abonos"]
    for r in range(4, 10):
        fecha = ws.cell(r, 5).value
        print(r, "fecha", repr(fecha), type(fecha), "coerce", coerce(fecha), "val", ws.cell(r, 6).value)

# check deployed source contains ISO parse
import re
# can't easily; print local file snippet
src = Path(
    r"D:\CMC\HBI_Capital\api-hbi-powerAutomate\app\application\use_cases\payment_validation_finalize.py"
).read_text(encoding="utf-8")
print("local_has_isoformat_parse", "%Y-%m-%d" in src[src.find("_coerce_fecha_banco_to_date") : src.find("_coerce_fecha_banco_to_date") + 500])
