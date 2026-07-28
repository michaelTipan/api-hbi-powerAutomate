"""Inspecciona BANCO_BOGOTA.xlsx y control en sandbox vía API."""
from __future__ import annotations

import base64
import io
from pathlib import Path
from urllib.parse import quote

import httpx
from openpyxl import load_workbook

BASE = "https://app-hbiauto-prod-001-afawg2g7frgte8c5.eastus-01.azurewebsites.net"
CLIENTS = (
    "INFORMACION CREDITOS-CLIENTES/"
    "02 COMWARE AUTOMATIZACION - INFORMACION CREDITOS CLIENTES"
)
BANK_DIR = f"{CLIENTS}/00 CARGA TRANSACCIONES BANCO"
CTRL_DIR = f"{CLIENTS}/01 VALIDACION PAGOS/01 CONTROL"
REV_DIR = f"{CLIENTS}/01 VALIDACION PAGOS"
WORK = Path(r"D:\CMC\HBI_Capital\_work")


def children(c: httpx.Client, drive: str, fid: str = "root") -> list[dict]:
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/children",
        params={"folder_item_id": fid},
        timeout=120,
    )
    r.raise_for_status()
    return r.json().get("value") or []


def walk(c: httpx.Client, drive: str, parts: list[str]) -> str:
    cur = "root"
    for p in parts:
        items = children(c, drive, cur)
        hit = next((i for i in items if i.get("name") == p), None)
        if not hit:
            raise FileNotFoundError(f"missing {p} in {'/'.join(parts)}")
        cur = hit["id"]
    return cur


def download(c: httpx.Client, drive: str, item_id: str) -> bytes:
    r = c.get(
        f"{BASE}/graph/sharepoint/drives/{quote(drive, safe='')}/item-content",
        params={"item_id": item_id},
        timeout=180,
    )
    r.raise_for_status()
    return base64.b64decode(r.json()["content_base64"])


def main() -> None:
    WORK.mkdir(exist_ok=True)
    with httpx.Client(timeout=180) as c:
        drive = c.get(f"{BASE}/graph/sharepoint/resolve-env").json()["resolved"]["drive_id"]
        bank_folder = walk(c, drive, BANK_DIR.split("/"))
        bank_items = children(c, drive, bank_folder)
        print("BANK FOLDER FILES:")
        for it in bank_items:
            print(" -", it.get("name"), "id=", it.get("id"))

        bog = next(i for i in bank_items if i.get("name") == "BANCO_BOGOTA.xlsx")
        raw = download(c, drive, bog["id"])
        (WORK / "BANCO_BOGOTA.xlsx").write_bytes(raw)
        wb = load_workbook(io.BytesIO(raw), data_only=False)
        print("sheets", wb.sheetnames)
        ws = wb.active
        print("active", ws.title, "max_row", ws.max_row, "max_col", ws.max_column)
        for r in range(1, min(20, ws.max_row + 1)):
            vals = [ws.cell(r, col).value for col in range(1, min(12, ws.max_column + 1))]
            print(r, vals)

        ctrl_folder = walk(c, drive, CTRL_DIR.split("/"))
        ctrl_items = children(c, drive, ctrl_folder)
        print("CONTROL FILES:")
        for it in ctrl_items:
            print(" -", it.get("name"))
        ctrl = next(
            i
            for i in ctrl_items
            if i.get("name") == "control_proceso_validacion_pagos_banco_bogota.xlsx"
        )
        craw = download(c, drive, ctrl["id"])
        (WORK / "control_bogota.xlsx").write_bytes(craw)
        cwb = load_workbook(io.BytesIO(craw), data_only=True)
        print("control sheets", cwb.sheetnames)
        for name in cwb.sheetnames:
            s = cwb[name]
            print("---", name)
            for r in range(1, min(30, s.max_row + 1)):
                vals = [s.cell(r, col).value for col in range(1, min(8, s.max_column + 1))]
                if any(v is not None and str(v).strip() for v in vals):
                    print(r, vals)

        val_root = walk(c, drive, REV_DIR.split("/"))
        print("VALIDACION children:")
        for it in children(c, drive, val_root):
            print(" -", it.get("name"), "folder" if "folder" in it else "file")


if __name__ == "__main__":
    main()
