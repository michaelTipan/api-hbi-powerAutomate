"""Apply Fase 5: ABONO escribe fila de aplicación sin tocar IBR."""

import asyncio
import hashlib
import io
from datetime import date

import openpyxl
import pytest

from app.application.services.abono_dry_run import ABONO_APPLICATION_ROW_STRATEGY
from app.application.services.amortization_workbook import (
    AUTOMATION_LOG_SHEET,
    detect_headers,
    find_header_row,
    load_automation_log_idempotency_keys,
)
from app.application.use_cases.amortization_fill_apply import run_amortization_fill_apply
from tests.test_abono_dry_run import _abono_dry_run_files, _abono_manifest_output
from tests.test_amortization_fill_apply import MockGraphApply
from tests.test_amortization_fill_dry_run import (
    _accounting_text,
    _amort_table_date_at_row,
    _asiento_pdf_placeholder,
    _base_files,
    _hist_bytes,
    _ibr_bytes,
)


@pytest.fixture(autouse=True)
def env_sharepoint(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")


@pytest.fixture(autouse=True)
def stub_process_control_writes(monkeypatch):
    async def _ok_update(*_a, **_k):
        return None

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.update_process_control_row2",
        _ok_update,
    )


def _patch_pdf_extract(monkeypatch):
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    monkeypatch.setattr(
        "app.application.services.abono_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )


def _ibr_cell_value(blob: bytes, row: int) -> object:
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    ws = wb.active
    headers = detect_headers(ws, header_row=find_header_row(ws))
    col = headers.get("ibr_i")
    if not col:
        return None
    return ws.cell(row, col).value


def test_apply_abono_writes_row_without_ibr(monkeypatch):
    fecha = date(2026, 5, 22)
    out = _abono_manifest_output()
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    tabla = "TABLAS/amort_258.xlsx"
    ibr_before = _ibr_cell_value(files[tabla], 2)
    asiento = out["credit_items"][0]["asiento_pdf_paths"][0]
    g = MockGraphApply(files)
    _patch_pdf_extract(monkeypatch)

    result = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )

    assert result["status"] == "ok"
    assert result["apply_wrote_changes"] is True
    assert result["abono_application_rows_written"] == 1
    assert result["abono_ibr_updates_skipped"] >= 1
    assert "IBR no fue modificado" in result.get("user_message", "")
    assert len(g.put_calls) == 1
    ibr_after = _ibr_cell_value(g.uploaded[tabla], 2)
    assert ibr_after == ibr_before
    item = [i for i in result["items"] if i.get("tipo_aplicacion") == "ABONO"][0]
    assert item["apply_status"] == "APPLIED"
    assert item["apply_ibr_written"] is False
    assert item["application_row_strategy"] == ABONO_APPLICATION_ROW_STRATEGY
    assert item["payment_application"]["valor_pagado_cliente"] == 50_000_000.0
    wb = openpyxl.load_workbook(io.BytesIO(g.uploaded[tabla]), data_only=False)
    ws = wb.active
    headers = detect_headers(ws, header_row=find_header_row(ws))
    app_row = int(item["application_row"])
    vp_cell = ws.cell(app_row, headers["valor_pagado_cliente"]).value
    assert vp_cell is not None
    assert load_automation_log_idempotency_keys(
        openpyxl.load_workbook(io.BytesIO(g.uploaded[tabla]), data_only=False)
    )
    assert asiento not in g.files
    assert result["accounting_pdfs_moved_count"] == 1


def test_apply_abono_retry_idempotent(monkeypatch):
    fecha = date(2026, 5, 22)
    out = _abono_manifest_output()
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    g = MockGraphApply(files)
    _patch_pdf_extract(monkeypatch)

    async def _noop_move(*_a, **_k):
        from app.application.services.accounting_pdf_processed_move import (
            empty_accounting_pdf_move_summary,
        )

        return empty_accounting_pdf_move_summary()

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.process_used_accounting_pdfs_after_apply",
        _noop_move,
    )

    out1 = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out1["summary"]["applied"] == 1
    log_rows_after_first = _automation_log_rows(g, "TABLAS/amort_258.xlsx")

    out2 = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out2["status"] == "ok"
    assert out2["summary"]["applied"] == 0
    assert out2["summary"]["skipped_idempotent"] == 1
    assert _automation_log_rows(g, "TABLAS/amort_258.xlsx") == log_rows_after_first


def _automation_log_rows(g: MockGraphApply, tabla: str) -> int:
    blob = g.uploaded.get(tabla) or g.files.get(tabla)
    if not blob:
        return 0
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if AUTOMATION_LOG_SHEET not in wb.sheetnames:
        return 0
    return max(0, wb[AUTOMATION_LOG_SHEET].max_row - 1)


def test_apply_mixed_pago_ibr_abono_no_ibr(monkeypatch):
    fecha = date(2026, 5, 22)
    pago = {
        "id_pago": "P1",
        "cliente": "EQUINORTE",
        "credito": "CREDITO # 258",
        "tipo_aplicacion": "PAGO",
        "requiere_extracto": True,
        "asiento_pdf_path": "clientes/EQUINORTE/CREDITO # 258/ASIENTOS CONTABLES CRED 258/asiento.pdf",
        "extracto_pdf_path": "clientes/EQUINORTE/extracto.pdf",
        "output_relative_path": "OUT/pago.pdf",
    }
    abono = _abono_manifest_output(id_pago="AB9")
    files = _abono_dry_run_files(abono_output=abono, pago_manifest_output=pago, fecha=fecha)
    files["TABLAS/amort_258.xlsx"] = _amort_table_date_at_row(fecha, 8)
    g = MockGraphApply(files)
    _patch_pdf_extract(monkeypatch)

    result = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )

    assert result["status"] == "ok"
    assert result["apply_wrote_changes"] is True
    pago_items = [i for i in result["items"] if i.get("tipo_aplicacion") != "ABONO"]
    abono_items = [i for i in result["items"] if i.get("tipo_aplicacion") == "ABONO"]
    assert pago_items[0]["apply_ibr_written"] is True
    assert abono_items[0]["apply_ibr_written"] is False
    assert pago_items[0]["apply_status"] == "APPLIED"
    assert abono_items[0]["apply_status"] == "APPLIED"


def test_apply_pago_only_still_works(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    _patch_pdf_extract(monkeypatch)
    result = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert result["status"] == "ok"
    assert result["summary"]["applied"] == 1
