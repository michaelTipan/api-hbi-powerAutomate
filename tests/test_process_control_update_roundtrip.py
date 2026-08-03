"""Integridad Excel: parche fila 2 del control + roundtrip ZIP + concurrencia eTag."""

from __future__ import annotations

import asyncio
import io
import zipfile
from urllib.parse import unquote

import httpx
import openpyxl
import pytest

from app.application.use_cases.payment_validation_process_control import (
    parse_process_control_row2,
    update_process_control_row2,
)
from app.application.use_cases.setup_merge_control_workbook import (
    BANK_CODE_BOGOTA,
    PROCESS_CONTROL_BANK_FILE_BOGOTA,
    SHEET_NAME,
    _build_process_control_workbook_bytes,
)


def _http_error(status: int) -> httpx.HTTPStatusError:
    req = httpx.Request("PUT", "https://graph.microsoft.com/test")
    resp = httpx.Response(status, request=req, text=f"HTTP {status}")
    return httpx.HTTPStatusError("err", request=req, response=resp)


def _assert_valid_xlsx_zip(raw: bytes) -> None:
    assert raw[:2] == b"PK", "Un .xlsx válido es un ZIP (firma PK)"
    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        names = set(zf.namelist())
        assert "[Content_Types].xml" in names
        assert "xl/workbook.xml" in names


def test_process_control_update_roundtrip_preserves_xlsx_structure():
    raw = _build_process_control_workbook_bytes(BANK_CODE_BOGOTA, "Banco de Bogotá")
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        ws = wb[SHEET_NAME]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(headers)}
        ws.cell(2, col["EstadoProceso"], value="REVISION_CREADA")
        buf = io.BytesIO()
        wb.save(buf)
        patched = buf.getvalue()
    finally:
        wb.close()

    _assert_valid_xlsx_zip(patched)
    snap = parse_process_control_row2(patched, control_file_path=PROCESS_CONTROL_BANK_FILE_BOGOTA)
    assert snap.estado_proceso == "REVISION_CREADA"


class _ControlGraphMock:
    def __init__(self, initial: bytes) -> None:
        self._bytes = initial
        self._etag = '"v1"'
        self.put_calls: list[tuple[str | None, bytes]] = []
        self.get_meta_calls = 0

    def _path(self, endpoint: str) -> str:
        if "/root:/" not in endpoint:
            return ""
        part = endpoint.split("/root:/", 1)[1]
        if part.endswith(":/content"):
            part = part[: -len(":/content")]
        else:
            part = part.rstrip(":")
        return unquote(part).strip("/")

    async def get(self, endpoint: str, params=None):
        self.get_meta_calls += 1
        return {"eTag": self._etag, "name": "control.xlsx"}

    async def get_bytes(self, endpoint: str, params=None):
        return self._bytes

    async def put_bytes(
        self,
        endpoint: str,
        content: bytes,
        content_type: str = "",
        if_match: str | None = None,
    ):
        self.put_calls.append((if_match, content))
        if len(self.put_calls) == 1:
            raise _http_error(412)
        self._bytes = content
        self._etag = f'"v{len(self.put_calls) + 1}"'
        return {"webUrl": "https://example/control.xlsx", "eTag": self._etag}


def test_update_process_control_row2_retries_on_etag_conflict():
    raw = _build_process_control_workbook_bytes(BANK_CODE_BOGOTA, "Banco de Bogotá")
    graph = _ControlGraphMock(raw)

    asyncio.run(
        update_process_control_row2(
            graph,  # type: ignore[arg-type]
            "site",
            "drive",
            bank_code=BANK_CODE_BOGOTA,
            updates={"EstadoProceso": "FINALIZADO", "IsActive": True},
        )
    )

    assert graph.get_meta_calls >= 2
    assert len(graph.put_calls) == 2
    assert graph.put_calls[0][0] == '"v1"'
    _assert_valid_xlsx_zip(graph._bytes)
    snap = parse_process_control_row2(graph._bytes, control_file_path=PROCESS_CONTROL_BANK_FILE_BOGOTA)
    assert snap.estado_proceso == "FINALIZADO"
    assert snap.is_active is True
