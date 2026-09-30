"""Golden antes/después: casos simples (1 consignación, 1 crédito, 1 asiento).

Los snapshots se capturaron con el motor previo a la conciliación por grupos.
Filas, valores escritos, fechas y claves de idempotencia deben ser idénticos.
Regenerar solo con ``UPDATE_GOLDEN=1`` y justificación explícita.
"""

from __future__ import annotations

import asyncio
import io
import json
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from app.application.services.amortization_workbook import AUTOMATION_LOG_SHEET
from app.application.use_cases.amortization_fill_apply import run_amortization_fill_apply
from app.application.use_cases.amortization_fill_dry_run import run_amortization_fill_dry_run
from tests.test_abono_dry_run import _abono_dry_run_files, _abono_manifest_output
from tests.test_amortization_fill_apply import MockGraphApply
from tests.test_amortization_fill_dry_run import (
    _accounting_text,
    _amort_table_date_at_row,
    _asiento_pdf_placeholder,
    _base_files,
    _hist_bytes,
    _ibr_bytes,
)

GOLDEN_PATH = Path(__file__).parent / "fixtures" / "golden_conciliacion_simple.json"

_VOLATILE_LOG_COLUMNS = frozenset({"Timestamp", "AsientoPdfHash", "AsientoPdfETag"})
_ITEM_FIELDS = (
    "id_pago",
    "credito",
    "tipo_aplicacion",
    "application_status",
    "error_code",
    "application_row",
    "due_date_row",
    "payment_date_iso",
    "payment_application",
    "idempotency_key",
)


@pytest.fixture(autouse=True)
def env_sharepoint(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")


@pytest.fixture(autouse=True)
def stub_process_control_writes(monkeypatch):
    async def _ok_update(*_a, **_k):
        return None

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.update_process_control_row2",
        _ok_update,
    )


def _patch_text(monkeypatch, text: str) -> None:
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: text,
    )
    monkeypatch.setattr(
        "app.application.services.abono_dry_run.extract_text_from_pdf",
        lambda _b: text,
    )


def _cell_repr(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return round(value, 6)
    return value


def _sheet_dump(blob: bytes) -> dict[str, Any]:
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=False)
    out: dict[str, Any] = {}
    for ws in wb.worksheets:
        if ws.title == AUTOMATION_LOG_SHEET:
            headers = [str(c.value or "") for c in ws[1]]
            rows = []
            for r in range(2, (ws.max_row or 1) + 1):
                row = {}
                for idx, name in enumerate(headers, start=1):
                    if not name or name in _VOLATILE_LOG_COLUMNS:
                        continue
                    val = ws.cell(r, idx).value
                    if val not in (None, ""):
                        row[name] = _cell_repr(val)
                rows.append(row)
            out[ws.title] = rows
            continue
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value not in (None, ""):
                    cells[c.coordinate] = _cell_repr(c.value)
        out[ws.title] = cells
    return out


def _item_snapshot(item: dict[str, Any]) -> dict[str, Any]:
    snap = {k: item.get(k) for k in _ITEM_FIELDS}
    pa = snap.get("payment_application")
    if isinstance(pa, dict):
        snap["payment_application"] = {k: pa.get(k) for k in sorted(pa)}
    return snap


def _identity(manifest: dict[str, Any]) -> dict[str, Any]:
    return manifest


def _as_manifest_v2(manifest: dict[str, Any]) -> dict[str, Any]:
    """Mismo lote con el formato de consolidados (un consolidado por pago en casos 1↔1↔1)."""
    outputs = []
    for out in manifest.get("outputs") or []:
        key = f"{out.get('credito') or out.get('id_pago')}|{out.get('fecha_banco') or ''}"
        nombre = str(out.get("output_relative_path") or "x.pdf").rsplit("/", 1)[-1]
        payment = {**out, "consolidado_key": key, "group_key": key.split("|")[0], "nombre_consolidado": nombre}
        outputs.append(
            {
                **{k: v for k, v in out.items() if k not in ("credit_items", "asiento_pdf_paths")},
                "id_pagos": [out.get("id_pago")],
                "consolidado_key": key,
                "nombre_consolidado": nombre,
                "payments": [payment],
            }
        )
    return {**manifest, "manifest_version": 2, "outputs": outputs}


_MANIFEST = {"transform": _identity}


def _run_case(g: MockGraphApply, fecha: date, tables: list[str]) -> dict[str, Any]:
    for path in [p for p in g.files if "merge_manifest" in p and p.endswith(".json")]:
        g.files[path] = json.dumps(_MANIFEST["transform"](json.loads(g.files[path]))).encode("utf-8")
    dry = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
            update_process_control=False,
        )
    )
    applied = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    return {
        "dry_run": {
            "can_apply": dry.get("can_apply"),
            "items": [_item_snapshot(it) for it in dry.get("items") or []],
        },
        "apply": {
            "status": applied.get("status"),
            "applied": (applied.get("summary") or {}).get("applied"),
            "items": [
                {
                    **_item_snapshot(it),
                    "apply_status": it.get("apply_status"),
                }
                for it in applied.get("items") or []
            ],
        },
        "tables": {t: _sheet_dump(g.uploaded[t]) for t in tables if t in g.uploaded},
    }


def _case_pago_legacy(monkeypatch) -> dict[str, Any]:
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    _patch_text(monkeypatch, _accounting_text())
    return _run_case(g, fecha, ["TABLAS/amort.xlsx"])


def _case_pago_credit_items(monkeypatch, tipo: str) -> dict[str, Any]:
    fecha = date(2026, 4, 22)
    asiento = "clientes/EQUINORTE/CREDITO # 258/ASIENTOS CONTABLES CRED 258/4635.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "P1",
                "cliente": "EQUINORTE",
                "credito": "258",
                "tipo_aplicacion_original": tipo,
                "monto_banco": 50_000_000.0,
                "fecha_banco": fecha.isoformat(),
                "creditos_seleccionados": ["258"],
                "credit_items": [
                    {
                        "credito": "258",
                        "tipo_aplicacion_original": tipo,
                        "ruta_tabla_amortizacion": "TABLAS/amort.xlsx",
                        "asiento_pdf_paths": [asiento],
                        "extracto_pdf_paths": ["clientes/EQUINORTE/extracto.pdf"],
                    }
                ],
                "output_relative_path": "OUT/p1.pdf",
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": _hist_bytes(
            "P1", "258", "TABLAS/amort.xlsx", fecha, tipo_aplicacion=tipo
        ),
        "TABLAS/amort.xlsx": _amort_table_date_at_row(fecha, 8),
        asiento: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphApply(files)
    _patch_text(monkeypatch, _accounting_text())
    return _run_case(g, fecha, ["TABLAS/amort.xlsx"])


def _case_abono(monkeypatch) -> dict[str, Any]:
    fecha = date(2026, 5, 22)
    out = _abono_manifest_output()
    g = MockGraphApply(_abono_dry_run_files(abono_output=out, fecha=fecha))
    _patch_text(monkeypatch, _accounting_text())
    return _run_case(g, fecha, ["TABLAS/amort_258.xlsx"])


_CASES = {
    "pago_legacy": lambda mp: _case_pago_legacy(mp),
    "pago_cuota_credit_items": lambda mp: _case_pago_credit_items(mp, "PAGO CUOTA"),
    "pago_cuota_mas_abono_capital": lambda mp: _case_pago_credit_items(
        mp, "PAGO CUOTA Y ABONO A CAPITAL"
    ),
    "abono_simple": lambda mp: _case_abono(mp),
}


def _load_golden() -> dict[str, Any]:
    if not GOLDEN_PATH.exists():
        return {}
    return json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))


@pytest.mark.parametrize("case", sorted(_CASES))
def test_golden_simple_case_unchanged(case, monkeypatch):
    actual = json.loads(json.dumps(_CASES[case](monkeypatch), default=str))
    golden = _load_golden()
    if os.environ.get("UPDATE_GOLDEN") == "1":
        golden[case] = actual
        GOLDEN_PATH.parent.mkdir(parents=True, exist_ok=True)
        GOLDEN_PATH.write_text(
            json.dumps(golden, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        return
    assert case in golden, f"golden faltante para {case}; ejecute con UPDATE_GOLDEN=1"
    assert actual == golden[case]


def test_golden_manifest_con_literal_anterior_mas_da_el_mismo_resultado(monkeypatch):
    actual = json.loads(
        json.dumps(_case_pago_credit_items(monkeypatch, "PAGO CUOTA + ABONO A CAPITAL"), default=str)
    )
    assert actual == _load_golden()["pago_cuota_mas_abono_capital"]


@pytest.mark.parametrize("case", sorted(_CASES))
def test_golden_manifest_v2_da_el_mismo_resultado(case, monkeypatch):
    monkeypatch.setitem(_MANIFEST, "transform", _as_manifest_v2)
    seen: list[int] = []
    orig = MockGraphApply.get_bytes

    async def _get_bytes(self, endpoint, params=None):
        blob = await orig(self, endpoint, params)
        if "merge_manifest" in endpoint:
            seen.append(json.loads(blob).get("manifest_version"))
        return blob

    monkeypatch.setattr(MockGraphApply, "get_bytes", _get_bytes)
    actual = json.loads(json.dumps(_CASES[case](monkeypatch), default=str))
    assert seen and set(seen) == {2}
    assert actual == _load_golden()[case]
