"""Soft-close: CERRADO_SIN_AMORTIZAR sin borrar artefactos; mueve asientos usados."""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
from urllib.parse import unquote

import httpx
import openpyxl
import pytest

from app.application.job_status_enrichment import enrich_job_for_http_response
from app.application.ui.generate_capabilities import bank_blocks_new_generate
from app.application.ui.process_control_capabilities import (
    compute_cancel_lote_availability,
    compute_soft_close_availability,
)
from app.application.use_cases.payment_validation_soft_close import (
    SOFT_CLOSE_ESTADO,
    soft_close_payment_validation,
)
from app.application.use_cases.setup_merge_control_workbook import (
    PROCESS_CONTROL_BANK_FILE_BOGOTA,
    SHEET_NAME,
    _build_process_control_workbook_bytes,
)


def _http_error(status: int) -> httpx.HTTPStatusError:
    req = httpx.Request("GET", "https://graph.microsoft.com/test")
    resp = httpx.Response(status, request=req, text=f"HTTP {status} body")
    return httpx.HTTPStatusError("err", request=req, response=resp)


class MockGraphClientSoftClose:
    def __init__(self) -> None:
        self.downloaded_files: dict[str, bytes] = {}
        self.put_calls: list[tuple[str, bytes]] = []
        self.delete_calls: list[str] = []
        self.patch_calls: list[tuple[str, dict]] = []
        self.post_json_calls: list[tuple[str, dict]] = []

    def _item_path(self, endpoint: str) -> str | None:
        if "/root:/" not in endpoint or ":/content" in endpoint or "children" in endpoint:
            return None
        return unquote(endpoint.split("/root:/", 1)[1].rstrip(":")).strip("/")

    async def get(self, endpoint: str, params=None):
        if endpoint == "/sites":
            return {"value": [{"id": "dummy_site"}]}
        if endpoint == "/sites/dummy_site/drives":
            return {"value": [{"id": "dummy_drive", "name": "DRIVE"}]}
        if endpoint.endswith(":/children"):
            return {"value": []}
        path = self._item_path(endpoint)
        if path is not None:
            if path not in self.downloaded_files:
                raise _http_error(404)
            blob = self.downloaded_files.get(path, b"")
            digest = hashlib.sha256(blob).hexdigest()[:16]
            return {
                "id": "item1",
                "name": path.rsplit("/", 1)[-1],
                "webUrl": "https://example/item",
                "eTag": f'"{digest}"',
                "size": len(blob),
            }
        return {"value": []}

    async def get_bytes(self, endpoint: str, params=None):
        if not endpoint.endswith(":/content"):
            raise _http_error(404)
        path = unquote(endpoint.split("/root:/", 1)[1].rsplit(":/content", 1)[0]).strip("/")
        if path in self.downloaded_files:
            return self.downloaded_files[path]
        raise _http_error(404)

    async def put_bytes(self, endpoint: str, content: bytes, content_type: str = ""):
        path = unquote(endpoint.split("/root:/", 1)[1].rsplit(":/content", 1)[0]).strip("/")
        self.put_calls.append((endpoint, content))
        self.downloaded_files[path] = content
        return {"webUrl": "https://example/upload"}

    async def post_json(self, endpoint: str, body: dict):
        self.post_json_calls.append((endpoint, body))
        if ":/children" in endpoint:
            parent = unquote(endpoint.split("/root:/", 1)[1].rsplit(":/children", 1)[0])
            name = str(body.get("name") or "").strip()
            if name:
                self.downloaded_files.setdefault(f"{parent}/{name}", b"")
        return {}, 201

    async def patch_json(self, endpoint: str, body: dict):
        self.patch_calls.append((endpoint, body))
        source = self._item_path(endpoint)
        if not source or source not in self.downloaded_files:
            raise _http_error(404)
        parent_path = str(body.get("parentReference", {}).get("path") or "")
        parent = parent_path.replace("/drive/root:/", "").strip("/")
        name = str(body.get("name") or "").strip()
        dest = f"{parent}/{name}"
        self.downloaded_files[dest] = self.downloaded_files.pop(source)
        return {"id": dest}

    async def delete(self, endpoint: str, params=None):
        self.delete_calls.append(endpoint)
        return None


def _set_env() -> None:
    os.environ["GRAPH_SHAREPOINT_SITE_SEARCH"] = "SITIO"
    os.environ["GRAPH_SHAREPOINT_DRIVE_NAME"] = "DRIVE"
    os.environ["GRAPH_PAYMENT_VALIDATION_REVIEW_PATH"] = "revision"
    os.environ["GRAPH_CLIENTS_BASE_PATH"] = "clientes"
    os.environ["GRAPH_PROCESS_ARCHIVE_PATH"] = "archivo"


def _control_late_phase(
    *,
    estado: str = "CONSOLIDADO",
    process_key: str = "payment-validation|banco_bogota|2026-06-01|abc",
    merge_manifest_path: str = "merge/m.json",
) -> bytes:
    raw = _build_process_control_workbook_bytes("banco_bogota", "Banco de Bogotá")
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        ws = wb[SHEET_NAME]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(headers)}
        ws.cell(2, col["ProcessKey"], value=process_key)
        ws.cell(2, col["EstadoProceso"], value=estado)
        ws.cell(2, col["IsActive"], value=True)
        ws.cell(2, col["HistoricalFilePath"], value="historico/h.xlsx")
        ws.cell(2, col["EmailPdfPath"], value="correo/c.pdf")
        ws.cell(2, col["MergeManifestPath"], value=merge_manifest_path)
        ws.cell(2, col["ValidationFilePath"], value="revision/val.xlsx")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()


def _read_control_row(raw: bytes) -> dict[str, object]:
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    try:
        ws = wb[SHEET_NAME]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(headers) if h}
        return {name: ws.cell(2, idx).value for name, idx in col.items()}
    finally:
        wb.close()


def test_soft_close_sets_terminal_free_state_without_deleting_artifacts():
    _set_env()
    client = MockGraphClientSoftClose()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_late_phase()

    res = asyncio.run(
        soft_close_payment_validation(
            client,
            bank_code="banco_bogota",
            process_key="payment-validation|banco_bogota|2026-06-01|abc",
            reason="Amortización manual",
        )
    )
    assert res["already_closed"] is False
    assert res["process_control_estado"] == SOFT_CLOSE_ESTADO
    assert res["artifacts_deleted"] is False
    assert not client.delete_calls
    assert res["accounting_pdfs_moved_count"] == 0

    row = _read_control_row(client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA])
    assert row["EstadoProceso"] == SOFT_CLOSE_ESTADO
    assert str(row["IsActive"]).strip().lower() in ("false", "falso", "0")
    assert row["HistoricalFilePath"] == "historico/h.xlsx"
    assert row["EmailPdfPath"] == "correo/c.pdf"
    assert row["MergeManifestPath"] == "merge/m.json"
    assert row["LastErrorUserMessage"] == "Amortización manual"
    assert row["LastCompletedStep"] == "SOFT_CLOSE"


def test_soft_close_moves_used_asientos_to_procesados():
    _set_env()
    asiento = "clientes/E/CREDITO # 258/ASIENTOS CONTABLES CRED 258/asiento.pdf"
    unused = "clientes/E/CREDITO # 258/ASIENTOS CONTABLES CRED 258/otro.pdf"
    manifest_path = "merge/m.json"
    manifest = {
        "report_date_iso": "2026-06-01",
        "outputs": [
            {
                "status": "COMPLETE",
                "id_pago": "pago1",
                "fecha_banco": "2026-06-01",
                "credito": "CREDITO # 258",
                "asiento_pdf_paths": [asiento],
            }
        ],
    }
    client = MockGraphClientSoftClose()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_late_phase(
        merge_manifest_path=manifest_path
    )
    client.downloaded_files[manifest_path] = json.dumps(manifest).encode("utf-8")
    client.downloaded_files[asiento] = b"%PDF-asiento"
    client.downloaded_files[unused] = b"%PDF-unused"

    res = asyncio.run(
        soft_close_payment_validation(
            client,
            bank_code="banco_bogota",
            process_key="payment-validation|banco_bogota|2026-06-01|abc",
            reason="",
        )
    )
    assert res["already_closed"] is False
    assert res["process_control_estado"] == SOFT_CLOSE_ESTADO
    assert res["accounting_pdfs_moved_count"] == 1
    assert asiento not in client.downloaded_files
    assert unused in client.downloaded_files
    dests = [
        p
        for p in client.downloaded_files
        if "PROCESADOS" in p and p.endswith(".pdf")
    ]
    assert len(dests) == 1
    assert "ASIENTOS CONTABLES CRED 258" in dests[0]
    assert not client.delete_calls


def test_soft_close_completes_when_asiento_move_warns():
    """Si el PDF ya no está, soft-close igual cierra el banco (best-effort)."""
    _set_env()
    asiento = "clientes/E/CREDITO # 258/ASIENTOS CONTABLES CRED 258/missing.pdf"
    manifest_path = "merge/m.json"
    manifest = {
        "report_date_iso": "2026-06-01",
        "outputs": [
            {
                "status": "COMPLETE",
                "id_pago": "pago1",
                "credito": "258",
                "asiento_pdf_path": asiento,
            }
        ],
    }
    client = MockGraphClientSoftClose()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_late_phase(
        merge_manifest_path=manifest_path
    )
    client.downloaded_files[manifest_path] = json.dumps(manifest).encode("utf-8")
    # asiento ausente a propósito

    res = asyncio.run(
        soft_close_payment_validation(
            client,
            bank_code="banco_bogota",
            process_key="payment-validation|banco_bogota|2026-06-01|abc",
            reason="",
        )
    )
    assert res["process_control_estado"] == SOFT_CLOSE_ESTADO
    assert res["accounting_pdfs_moved_count"] == 0
    assert res["accounting_pdfs_move_warnings_count"] == 1
    row = _read_control_row(client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA])
    assert row["EstadoProceso"] == SOFT_CLOSE_ESTADO


def test_soft_close_refuses_pre_finalize():
    _set_env()
    client = MockGraphClientSoftClose()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_late_phase(
        estado="REVISION_CREADA"
    )
    with pytest.raises(ValueError, match="soft_close_not_allowed\\|REVISION_CREADA"):
        asyncio.run(
            soft_close_payment_validation(
                client,
                bank_code="banco_bogota",
                process_key="payment-validation|banco_bogota|2026-06-01|abc",
                reason="no aplica",
            )
        )


def test_soft_close_refuses_merge_phase_states():
    _set_env()
    client = MockGraphClientSoftClose()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_late_phase(
        estado="PENDIENTE_ASIENTOS"
    )
    with pytest.raises(ValueError, match="soft_close_not_allowed\\|PENDIENTE_ASIENTOS"):
        asyncio.run(
            soft_close_payment_validation(
                client,
                bank_code="banco_bogota",
                process_key="payment-validation|banco_bogota|2026-06-01|abc",
                reason="",
            )
        )


def test_soft_close_allows_empty_reason():
    _set_env()
    client = MockGraphClientSoftClose()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_late_phase()
    res = asyncio.run(
        soft_close_payment_validation(
            client,
            bank_code="banco_bogota",
            process_key="payment-validation|banco_bogota|2026-06-01|abc",
            reason="  ",
        )
    )
    assert res["already_closed"] is False
    assert res["reason"] == ""
    row = _read_control_row(client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA])
    assert row["EstadoProceso"] == SOFT_CLOSE_ESTADO
    assert row["LastErrorUserMessage"] == "Cerrado sin amortizar."


def test_soft_close_unlocks_generate_capabilities():
    assert (
        bank_blocks_new_generate(
            process_key="payment-validation|banco_bogota|2026-06-01|abc",
            control_estado=SOFT_CLOSE_ESTADO,
            operational_status=SOFT_CLOSE_ESTADO,
            is_active=False,
        )
        is False
    )


def test_capabilities_cancel_vs_soft_close_phases():
    cancel = compute_cancel_lote_availability(
        write_allowed=True,
        mutation_active=False,
        control_estado="REVISION_CREADA",
        is_active=True,
    )
    soft = compute_soft_close_availability(
        write_allowed=True,
        mutation_active=False,
        control_estado="REVISION_CREADA",
    )
    assert cancel.allowed is True
    assert soft.allowed is False

    cancel2 = compute_cancel_lote_availability(
        write_allowed=True,
        mutation_active=False,
        control_estado="CONSOLIDADO",
        is_active=True,
    )
    soft2 = compute_soft_close_availability(
        write_allowed=True,
        mutation_active=False,
        control_estado="CONSOLIDADO",
    )
    assert cancel2.allowed is True
    assert soft2.allowed is True

    # Merge / asientos: soft-close no aplica (solo amortización).
    for merge_estado in ("PENDIENTE_ASIENTOS", "MERGE_PARCIAL", "ERROR_MERGE"):
        soft_merge = compute_soft_close_availability(
            write_allowed=True,
            mutation_active=False,
            control_estado=merge_estado,
        )
        assert soft_merge.allowed is False, merge_estado

    soft_partial = compute_soft_close_availability(
        write_allowed=True,
        mutation_active=False,
        control_estado="AMORTIZACION_PARCIAL",
    )
    assert soft_partial.allowed is True


def test_enrichment_soft_close_completed_and_failed():
    completed = enrich_job_for_http_response(
        {
            "job_id": "s1",
            "type": "soft_close_process",
            "status": "completed",
            "result": {
                "already_closed": False,
                "bank_code": "banco_bogota",
                "bank_name": "Banco de Bogotá",
            },
        }
    )
    assert completed["severity"] == "success"
    assert "sin amortizar" in completed["user_message"].lower()

    failed = enrich_job_for_http_response(
        {
            "job_id": "s2",
            "type": "soft_close_process",
            "status": "failed",
            "error": {
                "type": "ValueError",
                "message": "soft_close_not_allowed|REVISION_CREADA|pk",
            },
        }
    )
    assert failed["severity"] == "error"
    assert failed["error"]["error_code"] == "soft_close_not_allowed"
