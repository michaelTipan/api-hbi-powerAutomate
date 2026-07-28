"""Finalize (Phase 2): auto-detect banco + control por banco + escribe FINALIZADO."""

import asyncio
import io
import os
from datetime import date
from urllib.parse import unquote

import httpx
import openpyxl
import pytest

from app.application.services.review_schema import (
    REVIEW_SCHEMA_VERSION,
    CasosPagoCols,
    ControlCols,
    DistribucionCols,
    ReviewSheets,
)
from app.application.use_cases.payment_validation_finalize import finalize_payment_validation
from app.application.use_cases.setup_merge_control_workbook import (
    PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA,
    PROCESS_CONTROL_BANK_FILE_BOGOTA,
    _build_process_control_workbook_bytes,
)


def _http_error(status: int) -> httpx.HTTPStatusError:
    req = httpx.Request("GET", "https://graph.microsoft.com/test")
    resp = httpx.Response(status, request=req, text=f"HTTP {status} body")
    return httpx.HTTPStatusError("err", request=req, response=resp)


def _set_env() -> None:
    os.environ["GRAPH_SHAREPOINT_SITE_SEARCH"] = "SITIO"
    os.environ["GRAPH_SHAREPOINT_DRIVE_NAME"] = "DRIVE"
    os.environ["GRAPH_PAYMENT_VALIDATION_REVIEW_PATH"] = "revision"
    os.environ["GRAPH_PAYMENT_VALIDATION_HISTORY_PATH"] = "history"
    os.environ["GRAPH_VALIDATION_FILE_PREFIX"] = "validacion_pagos"
    os.environ["GRAPH_CLIENTS_BASE_PATH"] = "clientes"


def _review_xlsx_ready() -> bytes:
    wb = openpyxl.Workbook()
    ws_ctrl = wb.active
    ws_ctrl.title = ReviewSheets.CONTROL
    ws_ctrl.append([ControlCols.ROW_PROCESAR, ControlCols.VAL_PROCESAR_SI])
    ws_ctrl.append([ControlCols.ROW_REVIEW_SCHEMA_VERSION, REVIEW_SCHEMA_VERSION])
    ws_ctrl.append([ControlCols.ROW_ESTADO, "EN_REVISION"])
    ws_casos = wb.create_sheet(ReviewSheets.CASOS_PAGO)
    ws_casos.append(CasosPagoCols.HEADERS)
    ws_casos.append(["ID1", None, "CLI", "concepto", 100, ""])
    ws_dist = wb.create_sheet(ReviewSheets.DISTRIBUCION)
    ws_dist.append(DistribucionCols.HEADERS)
    ws_dist.append(
        [
            "ID1",
            "CLI",
            "CRED",
            100,
            None,
            None,
            0,
            100,
            70,
            10,
            20,
            0,
            100,
            0,
            "NORMAL",
            "SI",
            "",
            "clientes/CLI/CRED/extractos/e1.pdf",
            "",
            "",
            "",
            "clientes/CLI/CRED",
            "",
            "",
            "",
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _control_ready(bank_code: str, bank_name: str, validation_path: str) -> bytes:
    raw = _build_process_control_workbook_bytes(bank_code, bank_name)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        ws = wb["Procesos"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(headers)}
        ws.cell(2, col["EstadoProceso"], value="REVISION_CREADA")
        ws.cell(2, col["IsActive"], value=True)
        ws.cell(2, col["ValidationFilePath"], value=validation_path)
        ws.cell(2, col["ProcessKey"], value=f"payment-validation|{bank_code}|2026-06-01")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()


class MockGraphFinalize:
    def __init__(self) -> None:
        self.downloaded_files: dict[str, bytes] = {}
        self.uploaded_files: dict[str, bytes] = {}
        self.requested_endpoints: list[str] = []
        self.children: list[dict] = []

    async def get(self, endpoint: str, params=None):
        self.requested_endpoints.append(endpoint)
        if endpoint == "/sites":
            return {"value": [{"id": "dummy_site"}]}
        if endpoint == "/sites/dummy_site/drives":
            return {"value": [{"id": "dummy_drive", "name": "DRIVE"}]}
        if endpoint.endswith(":/children"):
            return {"value": self.children}
        return {}

    async def get_bytes(self, endpoint: str, params=None):
        self.requested_endpoints.append(endpoint)
        if not endpoint.endswith(":/content"):
            raise _http_error(404)
        path = unquote(endpoint.split("/root:/", 1)[1].rsplit(":/content", 1)[0]).strip("/")
        if path not in self.downloaded_files:
            raise _http_error(404)
        return self.downloaded_files[path]

    async def put_bytes(self, endpoint: str, content: bytes, content_type: str = ""):
        path = unquote(endpoint.split("/root:/", 1)[1].rsplit(":/content", 1)[0]).strip("/")
        self.uploaded_files[path] = content
        return {"webUrl": "https://example/upload.xlsx"}

    async def post_json(self, endpoint: str, body: dict):
        return {}, 201


def test_finalize_without_bank_code_no_ready_process():
    _set_env()
    client = MockGraphFinalize()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _build_process_control_workbook_bytes(
        "banco_bogota", "Banco de Bogotá"
    )
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA] = _build_process_control_workbook_bytes(
        "banco_bancolombia", "Bancolombia"
    )
    with pytest.raises(ValueError, match="NO_READY_PROCESS"):
        asyncio.run(finalize_payment_validation(client))


def test_finalize_without_bank_code_detects_bogota():
    _set_env()
    client = MockGraphFinalize()
    review_path = "revision/validacion_pagos_banco_bogota_2026-06-01.xlsx"
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_ready(
        "banco_bogota", "Banco de Bogotá", review_path
    )
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA] = _build_process_control_workbook_bytes(
        "banco_bancolombia", "Bancolombia"
    )
    client.downloaded_files[review_path] = _review_xlsx_ready()
    res = asyncio.run(finalize_payment_validation(client, process_date=date(2026, 6, 1)))
    assert res["bank_code"] == "banco_bogota"
    assert res["process_control_updated"] is True
    assert res["process_control_estado"] == "FINALIZADO"
    assert "cartera_validada_banco_bogota_2026-06-01_" in res["historical_file_path"]
    assert res["historical_file_path"].endswith(".xlsx")
    assert "soporte_asientos_contables_banco_bogota_2026-06-01_" in res["secretary_file_path"]
    assert res["secretary_file_path"].endswith(".xlsx")
    assert res["process_key"].startswith("payment-validation|banco_bogota|2026-06-01")
    # Lote único: ProcessKey incluye UUID (4 segmentos).
    assert len(res["process_key"].split("|")) >= 4

