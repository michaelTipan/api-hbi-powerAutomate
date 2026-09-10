"""Tests utilidades tabla de amortización."""

import io
from datetime import date

import openpyxl
import pytest

from app.application.services.accounting_pdf_parser import (
    ACCOUNT_CAPITAL,
    ACCOUNT_VALOR_PAGADO_CLIENTE,
    PaymentApplicationEvent,
)
from app.application.services.amortization_workbook import (
    ADOPTADO_EXISTENTE,
    APLICADO,
    APPLICATION_GROWTH_BLOCKED,
    AUTOMATION_LOG_SHEET,
    REVISION_MANUAL,
    AmbiguousAmortizationHeadersError,
    AmortizationHeadersSplitAcrossRowsError,
    AmortizationSheetNotFoundError,
    append_automation_log,
    build_target_row_search_debug,
    compare_existing_application,
    detect_amortization_sheet,
    detect_headers,
    ensure_automation_log,
    build_amortization_idempotency_key,
    find_application_row_detailed,
    find_bottom_most_occupied_payment_row,
    find_row_by_due_date,
    find_row_by_due_date_detailed,
    human_required_header_labels,
    is_complete_applied_payment_row,
    is_incomplete_payment_orphan_row,
    is_payment_application_empty,
    load_automation_log_idempotency_keys,
    _row_due_date,
    write_ibr,
    PaymentApplicationWriteOptions,
    write_payment_application,
    infer_cierra_cuota_from_schedule,
    _is_formula_value,
    _default_saldo_capital_formula,
)


def _sample_event(**kwargs) -> PaymentApplicationEvent:
    base = dict(
        id_pago="P1",
        cliente="EQUINORTE",
        credito="258",
        asiento_pdf_path="a.pdf",
        comprobante="1",
        fecha_asiento=date(2026, 5, 10),
        valor_pagado_cliente=50_000_000.0,
        capital=49_118_143.0,
        intereses=0.0,
        mora=881_857.0,
        retenciones=0.0,
        saldos_menores=0.0,
        raw_text="",
    )
    base.update(kwargs)
    return PaymentApplicationEvent(**base)


def _amort_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabla"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "IBR +i",
            "Spread",
            "Cuota",
            "Cuota + i",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
            "Saldos Menores",
        ]
    )
    ws.append([15, 5, 2026, None, None, 1000, 1100, None, None, None, None, None, None])
    ws.append([22, 5, 2026, None, None, 1000, 1100, None, None, None, None, None, None])
    ws.append([22, 6, 2026, None, None, 1000, 1100, None, None, None, None, None, None])
    ws.append([22, 7, 2026, None, None, 1000, 1100, None, None, None, None, None, None])
    ws.append([22, 8, 2026, None, None, 1000, 1100, None, None, None, None, None, None])
    ws.append([22, 9, 2026, None, None, 1000, 1100, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _equinorte_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EQUINORTE"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "IBR +i",
            "Spread",
            "Cuota + i",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Intereses de mora",
            "Valor pagado cliente",
            "Saldos Menores",
        ]
    )
    ws.append([22, 5, 2026, None, None, 1000, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _headers_on_row_workbook_bytes(header_row: int) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    for _ in range(header_row - 1):
        ws.append([None] * 12)
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "IBR +i",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
            "Saldos Menores",
        ]
    )
    ws.append([22, 5, 2026, None, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _duplicate_dia_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EQUINORTE"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "dia",
            "Causac Inter Mes",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
            "Saldos Menores",
        ]
    )
    ws.append([22, 4, 2026, 99, None, None, None, None, None, None, None])
    ws.append([15, 4, 2026, 88, None, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _float_date_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabla"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([22.0, 4.0, 2026.0, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_duplicate_dia_uses_first_contiguous_group():
    wb = openpyxl.load_workbook(io.BytesIO(_duplicate_dia_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    assert match.headers["dia"] == 1
    assert match.headers["mes"] == 2
    assert match.headers["anio"] == 3
    row = find_row_by_due_date(
        match.worksheet,
        match.headers,
        date(2026, 4, 22),
        header_row=match.header_row,
    )
    assert row == 2
    assert match.worksheet.cell(row, match.headers["dia"]).value == 22


def test_find_row_april_22_2026():
    wb = openpyxl.load_workbook(io.BytesIO(_duplicate_dia_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    assert find_row_by_due_date(
        match.worksheet,
        match.headers,
        date(2026, 4, 22),
        header_row=match.header_row,
    ) == 2


def test_find_row_float_date_components():
    wb = openpyxl.load_workbook(io.BytesIO(_float_date_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    row = find_row_by_due_date(
        match.worksheet,
        match.headers,
        date(2026, 4, 22),
        header_row=match.header_row,
    )
    assert row == 2


def _monthly_schedule_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EQUINORTE"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([22, 12, 2025, None, None, None, None])
    for r in range(3, 8):
        ws.cell(r, 1, f'=IF(A{r}<>" ",1,1)')
        ws.cell(r, 2, f'=IF(A{r}<>" ",1,1)')
        ws.cell(r, 3, f'=IF(A{r}<>" ",1,1)')
        ws.cell(r, 7, 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_formula_cells_do_not_produce_false_dates():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.cell(2, 1, '=IF(A5<>" ",+L4," ")')
    ws.cell(2, 2, "=IF(AND(A5<>\" \")*(M4=12),1,1+M4)")
    ws.cell(2, 3, "=IF(AND(A5<>\" \")*(M5=1),1,1+N4,N4)")
    h = detect_headers(ws, header_row=1)
    assert _row_due_date(ws, 2, h) is None
    assert find_row_by_due_date(ws, h, date(2005, 5, 5), header_row=1) is None


def test_monthly_schedule_finds_april_22_2026():
    wb = openpyxl.load_workbook(
        io.BytesIO(_monthly_schedule_workbook_bytes()), data_only=True
    )
    match = detect_amortization_sheet(wb)
    result = find_row_by_due_date_detailed(
        match.worksheet,
        match.headers,
        date(2026, 4, 22),
        header_row=match.header_row,
    )
    assert result.monthly_schedule_fallback_used is True
    assert result.row == 6
    assert result.monthly_schedule_base_date == "2025-12-22"


def test_monthly_schedule_out_of_range_returns_none():
    wb = openpyxl.load_workbook(
        io.BytesIO(_monthly_schedule_workbook_bytes()), data_only=True
    )
    match = detect_amortization_sheet(wb)
    result = find_row_by_due_date_detailed(
        match.worksheet,
        match.headers,
        date(2030, 1, 22),
        header_row=match.header_row,
    )
    assert result.row is None


def test_target_row_debug_includes_date_columns():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    debug = build_target_row_search_debug(
        match.worksheet,
        match.headers,
        date(2099, 1, 1),
        header_row=match.header_row,
        sheet_name=match.worksheet.title,
        tabla_amortizacion_path="TABLAS/x.xlsx",
    )
    assert debug["target_date"] == "2099-01-01"
    assert debug["date_header_columns"] == {"dia": 1, "mes": 2, "anio": 3}
    assert debug["header_row"] == match.header_row
    assert debug["sheet_name"] == "Tabla"
    assert debug["tabla_amortizacion_path"] == "TABLAS/x.xlsx"
    assert any(c["date"] == "2026-05-22" for c in debug["candidate_date_rows"])


def test_detect_headers_and_sheet():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws = match.worksheet
    h = match.headers
    assert h["dia"] == 1
    assert h["mes"] == 2
    assert h["anio"] == 3
    assert h["ibr_i"] == 4
    assert "cuota" in h


def test_detect_sheet_named_equinorte_by_headers():
    wb = openpyxl.load_workbook(io.BytesIO(_equinorte_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    assert match.worksheet.title == "EQUINORTE"
    assert "dia" in match.headers
    assert match.headers["intereses_mora"] is not None
    assert "ibr_i" in match.headers


def test_detect_headers_on_row_4():
    wb = openpyxl.load_workbook(io.BytesIO(_headers_on_row_workbook_bytes(4)))
    match = detect_amortization_sheet(wb)
    assert match.header_row == 4
    assert match.headers["valor_pagado_cliente"] > 0
    row = find_row_by_due_date(
        match.worksheet,
        match.headers,
        date(2026, 5, 22),
        header_row=match.header_row,
    )
    assert row == 5


def test_detect_intereses_de_mora_and_ibr_plus_i_aliases():
    wb = openpyxl.load_workbook(io.BytesIO(_equinorte_workbook_bytes()))
    h = detect_headers(wb["EQUINORTE"])
    assert "intereses_mora" in h
    assert "ibr_i" in h


def test_detect_amortization_sheet_raises_with_debug():
    wb = openpyxl.Workbook()
    wb.active.title = "SoloResumen"
    wb.active.append(["Total", "Monto"])
    wb.active.append([1, 2])
    with pytest.raises(AmortizationSheetNotFoundError) as excinfo:
        detect_amortization_sheet(wb, tabla_amortizacion_path="TABLAS/x.xlsx")
    exc = excinfo.value
    assert "SoloResumen" in exc.workbook_sheets
    assert exc.tabla_amortizacion_path == "TABLAS/x.xlsx"
    assert exc.required_headers_missing


def test_detect_amortization_sheet_reports_renamed_required_column():
    """Renombrar Fecha pago no debe ocultar el resto: el error nombra lo que falta."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VICTOR SOTO"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "BASTDE",  # era Fecha pago
            "Valor intereses",
            "Abono a K",
            "QFWER",
            "Intereses de mora",
            "Valor pagado cliente",
        ]
    )
    ws.append([22, 5, 2026, None, 100, 200, None, 0, 300])
    with pytest.raises(AmortizationSheetNotFoundError) as excinfo:
        detect_amortization_sheet(wb, tabla_amortizacion_path="TABLAS/victor.xlsx")
    exc = excinfo.value
    assert exc.required_headers_missing == ["fecha_pago"]
    assert "fecha_pago" not in exc.headers_found
    assert "valor_intereses" in exc.headers_found
    assert "abono_k" in exc.headers_found
    assert "valor_pagado_cliente" in exc.headers_found
    assert exc.best_sheet_name == "VICTOR SOTO"
    assert human_required_header_labels(exc.required_headers_missing) == ["Fecha pago"]


def test_find_row_by_due_date():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    row = find_row_by_due_date(
        match.worksheet, match.headers, date(2026, 5, 22), header_row=match.header_row
    )
    assert row == 3


def test_find_row_by_due_date_excludes_used_rows():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    first = find_row_by_due_date(
        match.worksheet,
        match.headers,
        date(2026, 5, 22),
        header_row=match.header_row,
    )
    second = find_row_by_due_date(
        match.worksheet,
        match.headers,
        date(2026, 5, 22),
        exclude_rows=frozenset({first}),
        header_row=match.header_row,
    )
    assert first == 3
    assert second is None


def test_build_amortization_idempotency_key_includes_asiento_and_comprobante():
    key = build_amortization_idempotency_key(
        "7785e37e",
        "CREDITO # 258",
        "clientes/E/asiento.pdf",
        "99-1",
        pdf_hash="abc",
    )
    assert "7785e37e" in key
    assert "asiento.pdf" in key
    assert "99-1" in key
    assert "abc" in key


def test_ensure_automation_log_created():
    wb = openpyxl.Workbook()
    ws = ensure_automation_log(wb)
    assert ws.title == AUTOMATION_LOG_SHEET
    assert ws.cell(1, 1).value == "Timestamp"


def test_write_empty_row_aplicado():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    row = find_row_by_due_date(ws, h, date(2026, 5, 22), header_row=match.header_row)
    assert row is not None
    ev = _sample_event()
    assert compare_existing_application(ws, row, h, ev) == APLICADO
    opts = PaymentApplicationWriteOptions(
        payment_date=date(2026, 5, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    write_payment_application(ws, row, h, ev, write_options=opts, header_row=match.header_row)
    assert not is_payment_application_empty(ws, row, h)


def test_existing_matching_adoptado():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    row = find_row_by_due_date(ws, h, date(2026, 5, 22), header_row=match.header_row)
    ev = _sample_event()
    opts = PaymentApplicationWriteOptions(
        payment_date=date(2026, 5, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    write_payment_application(ws, row, h, ev, write_options=opts, header_row=match.header_row)
    assert compare_existing_application(
        ws, row, h, ev, payment_date=date(2026, 5, 10), detected_codes=opts.detected_codes
    ) == ADOPTADO_EXISTENTE


def test_existing_mismatch_revision_manual():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    row = find_row_by_due_date(ws, h, date(2026, 5, 22), header_row=match.header_row)
    ev = _sample_event()
    opts = PaymentApplicationWriteOptions(
        payment_date=date(2026, 5, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    write_payment_application(ws, row, h, ev, write_options=opts, header_row=match.header_row)
    other = _sample_event(capital=1.0)
    assert compare_existing_application(ws, row, h, other) == REVISION_MANUAL


def test_write_ibr_on_corte_row():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    row = find_row_by_due_date(ws, h, date(2026, 5, 22), header_row=match.header_row)
    write_ibr(ws, row, h, 0.10580)
    assert ws.cell(row, h["ibr_i"]).value == pytest.approx(0.10580)


def test_append_automation_log():
    wb = openpyxl.Workbook()
    ensure_automation_log(wb)
    append_automation_log(
        wb,
        {
            "id_pago": "P1",
            "cliente": "X",
            "credito": "258",
            "fila": 3,
            "accion": APLICADO,
            "detalle": "ok",
            "idempotency_key": "k1",
            "asiento_pdf_path": "a.pdf",
            "application_row": 3,
            "ibr_row": 3,
        },
    )
    ws = wb[AUTOMATION_LOG_SHEET]
    assert ws.cell(2, 2).value == "P1"
    assert ws.cell(2, 6).value == APLICADO
    assert ws.cell(2, 8).value == "k1"
    assert load_automation_log_idempotency_keys(wb) == {"k1"}


def test_find_application_row_skips_occupied_and_uses_next_free():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.header_row
    headers = match.headers
    due = find_row_by_due_date(ws, headers, date(2026, 5, 22), header_row=h)
    assert due == 3
    # Pago completo en la due (Fecha pago + monto de asiento), como deja el apply real.
    ws.cell(3, headers["fecha_pago"]).value = date(2026, 5, 10)
    ws.cell(3, headers["abono_k"]).value = 99.0
    ws.cell(4, headers["valor_pagado_cliente"]).value = None
    ev = _sample_event()
    result = find_application_row_detailed(
        ws, headers, ev, due_date_row=due, header_row=h
    )
    assert result.row == 4
    assert result.compare_status == APLICADO


def test_find_application_row_adopts_matching_row_below_due():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    due = find_row_by_due_date(ws, match.headers, date(2026, 5, 22), header_row=match.header_row)
    ev = _sample_event()
    opts = PaymentApplicationWriteOptions(
        payment_date=date(2026, 5, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    write_payment_application(
        ws, 4, match.headers, ev, write_options=opts, header_row=match.header_row
    )
    result = find_application_row_detailed(
        ws, match.headers, ev, due_date_row=due, header_row=match.header_row
    )
    assert result.row == 4
    assert result.compare_status == ADOPTADO_EXISTENTE


def test_facturacion_label_row_is_not_empty():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    label_row = 4
    ws.cell(label_row, h["fecha_pago"]).value = "FACTURACION"
    assert not is_payment_application_empty(ws, label_row, h)
    ev = _sample_event()
    assert compare_existing_application(
        ws,
        label_row,
        h,
        ev,
        payment_date=date(2026, 8, 6),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    ) == REVISION_MANUAL


def test_find_application_row_blocks_facturacion_in_growth_zone():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    due = find_row_by_due_date(ws, h, date(2026, 5, 22), header_row=match.header_row)
    assert due == 3
    ws.cell(3, h["fecha_pago"]).value = date(2026, 5, 10)
    ws.cell(3, h["abono_k"]).value = 99.0
    ws.cell(4, h["fecha_pago"]).value = "FACTURACION"
    ev = _sample_event()
    result = find_application_row_detailed(
        ws,
        h,
        ev,
        due_date_row=due,
        header_row=match.header_row,
        payment_date=date(2026, 5, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    assert result.row is None
    assert result.growth_blocked is True
    assert result.growth_block_row == 4
    assert result.growth_block_end_row == 4
    assert result.growth_block_value == "Fila 4"
    assert result.growth_block_code == "APPLICATION_GROWTH_BLOCKED"
    assert result.search_start_row == 4


def test_find_application_row_does_not_adopt_facturacion_label():
    """Rótulo FACTURACION no debe adoptarse como pago."""
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    due = find_row_by_due_date(ws, h, date(2026, 5, 22), header_row=match.header_row)
    ws.cell(4, h["fecha_pago"]).value = "FACTURACION"
    ev = _sample_event()
    opts = PaymentApplicationWriteOptions(
        payment_date=date(2026, 5, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    result = find_application_row_detailed(
        ws,
        h,
        ev,
        due_date_row=due,
        header_row=match.header_row,
        payment_date=opts.payment_date,
        detected_codes=opts.detected_codes,
    )
    # Sin pagos ocupados: start=due; fila due vacía → escribe ahí.
    # FACTURACION en 4 queda por debajo y no se toca si hay vacía antes.
    assert result.row == due
    assert result.compare_status == APLICADO
    assert result.growth_blocked is False


def test_infer_cierra_cuota_false_when_asiento_below_cuota():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "CUOTA + I",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([22, 5, 2026, 10_000_000, None, None, None, None])
    headers = {
        "dia": 1,
        "mes": 2,
        "anio": 3,
        "cuota_i": 4,
        "fecha_pago": 5,
        "valor_intereses": 6,
        "abono_k": 7,
        "valor_pagado_cliente": 8,
    }
    ev = _sample_event(intereses=1_000_000.0, capital=1_000_000.0, mora=0.0)
    assert (
        infer_cierra_cuota_from_schedule(
            ws=ws,
            headers=headers,
            due_date_row=2,
            application_row=2,
            event=ev,
            subtipo_aplicacion="CUOTA",
            policy_cierra_cuota=True,
        )
        is False
    )


def test_infer_cierra_cuota_true_when_asiento_covers_cuota():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "CUOTA + I",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([22, 5, 2026, 2_000_000, None, None, None, None])
    headers = {
        "dia": 1,
        "mes": 2,
        "anio": 3,
        "cuota_i": 4,
        "fecha_pago": 5,
        "valor_intereses": 6,
        "abono_k": 7,
        "valor_pagado_cliente": 8,
    }
    ev = _sample_event(intereses=800_000.0, capital=1_200_000.0, mora=0.0)
    assert (
        infer_cierra_cuota_from_schedule(
            ws=ws,
            headers=headers,
            due_date_row=2,
            application_row=2,
            event=ev,
            subtipo_aplicacion="CUOTA",
            policy_cierra_cuota=False,
        )
        is True
    )


def test_infer_cierra_cuota_ignores_mora_and_non_cuota_subtipo():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["dia", "mes", "año", "CUOTA", "Valor intereses", "Abono a K"])
    ws.append([22, 5, 2026, 5_000_000, None, None])
    headers = {
        "dia": 1,
        "mes": 2,
        "anio": 3,
        "cuota": 4,
        "valor_intereses": 5,
        "abono_k": 6,
    }
    ev = _sample_event(intereses=0.0, capital=0.0, mora=5_000_000.0)
    assert (
        infer_cierra_cuota_from_schedule(
            ws=ws,
            headers=headers,
            due_date_row=2,
            application_row=2,
            event=ev,
            subtipo_aplicacion="SALDO_VENCIDO",
            policy_cierra_cuota=False,
        )
        is False
    )
    assert (
        infer_cierra_cuota_from_schedule(
            ws=ws,
            headers=headers,
            due_date_row=2,
            application_row=2,
            event=ev,
            subtipo_aplicacion="CUOTA",
            policy_cierra_cuota=True,
        )
        is False
    )


def test_saldo_formula_anchors_to_last_occupied_across_gap():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
            "Saldo a capital",
        ]
    )
    ws.append([1, 5, 2026, date(2026, 5, 1), 0, 1000, 1000, 9000])
    ws.cell(2, 8).number_format = "#,##0.00"
    ws.cell(2, 6).number_format = "#,##0.00"
    ws.cell(2, 4).number_format = "DD/MM/YYYY"
    ws.append([15, 6, 2026, None, None, None, None, None])
    ws.append([20, 7, 2026, None, None, None, None, None])
    headers = detect_headers(ws, header_row=1)
    assert find_bottom_most_occupied_payment_row(ws, headers, header_row=1) == 2
    formula = _default_saldo_capital_formula(headers, 5, prev_occupied_row=2)
    assert formula == "=+H2-F5"

    ev = _sample_event(capital=500.0, intereses=0.0, mora=0.0, valor_pagado_cliente=500.0)
    plan = write_payment_application(
        ws,
        5,
        headers,
        ev,
        write_options=PaymentApplicationWriteOptions(
            payment_date=date(2026, 8, 1),
            detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
        ),
        header_row=1,
    )
    assert plan["saldo_a_capital"] == "formula"
    assert ws.cell(5, headers["saldo_a_capital"]).value == "=+H2-F5"
    assert ws.cell(5, headers["abono_k"]).number_format == "#,##0.00"
    assert ws.cell(5, headers["fecha_pago"]).number_format == "DD/MM/YYYY"


def test_application_row_first_empty_after_bottom_most_skips_historical_hole():
    """Estilo INGEOROZCOL: hueco 33-34; due en 35; último pago en 32 → escribe 33."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
            "Saldo a capital",
        ]
    )
    for _ in range(2, 33):
        ws.append([1, 1, 2025, None, None, None, None, None])
    ws.cell(20, 4).value = date(2026, 1, 10)
    ws.cell(20, 6).value = 100.0
    ws.cell(20, 7).value = 100.0
    ws.cell(32, 4).value = date(2026, 6, 10)
    ws.cell(32, 6).value = 200.0
    ws.cell(32, 7).value = 200.0
    ws.append([20, 7, 2026, None, None, None, None, None])  # 33
    ws.append([20, 7, 2026, None, None, None, None, None])  # 34
    ws.append([20, 7, 2026, None, None, None, None, None])  # 35 due
    headers = detect_headers(ws, header_row=1)
    assert find_bottom_most_occupied_payment_row(ws, headers, header_row=1) == 32
    ev = _sample_event()
    result = find_application_row_detailed(
        ws, headers, ev, due_date_row=35, header_row=1
    )
    assert result.search_start_row == 33
    assert result.row == 33
    assert result.compare_status == APLICADO
    assert result.row != 21


def test_application_row_does_not_adopt_far_below_contractual_anchor():
    """Basura bajo FACTURACION (fila 66) no debe ADOPTARSE ni mover el último pago."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
            "Saldo a capital",
        ]
    )
    for _ in range(2, 33):
        ws.append([1, 1, 2025, None, None, None, None, None])
    ws.cell(32, 4).value = date(2026, 6, 23)
    ws.cell(32, 5).value = 50.0
    ws.cell(32, 6).value = 100.0
    ws.cell(32, 7).value = 150.0
    for r in range(33, 52):
        ws.append([20, 7, 2026, None, None, None, None, None])
    # Cierra el bloque contractual (como tablas reales).
    ws.append([None, None, None, "FACTURACION", None, None, None, None])  # 52
    for _ in range(53, 66):
        ws.append([None, None, None, None, None, None, None, None])
    # Fila 66: resto lejano con montos/fecha del evento → antes ADOPTADO silencioso.
    ws.append([1, 1, 2020, date(2026, 8, 6), 500.0, 1000.0, 1500.0, None])
    headers = detect_headers(ws, header_row=1)
    assert find_bottom_most_occupied_payment_row(ws, headers, header_row=1) == 32
    ev = _sample_event(
        intereses=500.0,
        capital=1000.0,
        mora=0.0,
        valor_pagado_cliente=1500.0,
    )
    result = find_application_row_detailed(
        ws,
        headers,
        ev,
        due_date_row=35,
        header_row=1,
        payment_date=date(2026, 8, 6),
    )
    assert result.search_start_row == 33
    assert result.row == 33
    assert result.compare_status == APLICADO
    assert result.row != 66


def test_header_aliases_saldo_de_capital_and_abono_a_capital():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Abono a capital",
            "Valor pagado cliente",
        ]
    )
    ws.append([1, 1, 2026, None, None, None, None, None])
    h = detect_headers(ws, header_row=1)
    assert h["abono_k"] == 6
    assert h["saldo_a_capital"] == 7


@pytest.mark.parametrize(
    "label",
    ["RETENCIONES", "Retenciones", "retenciones", "ReTeNcIoNeS", " retenciones "],
)
def test_retenciones_header_is_case_insensitive(label: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
            label,
        ]
    )
    ws.append([1, 1, 2026, None, None, None, None, None])
    h = detect_headers(ws, header_row=1)
    assert h["retenciones"] == 8


def test_header_capital_alone_still_maps_abono_k():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Capital",
            "Valor pagado cliente",
            "Saldo de capital",
        ]
    )
    ws.append([1, 1, 2026, None, None, None, None, None])
    h = detect_headers(ws, header_row=1)
    assert h["abono_k"] == 6
    assert h["saldo_a_capital"] == 8


def test_ambiguous_duplicate_abono_k_headers_fail():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([1, 1, 2026, None, None, None, None, None])
    with pytest.raises(AmbiguousAmortizationHeadersError) as exc:
        detect_headers(ws, header_row=1)
    assert exc.value.key == "abono_k"


def test_contiguous_table_still_writes_next_row_noop_shape():
    """Casos sanos contiguos: último pago N → siguiente N+1 (sin huecos)."""
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    due = find_row_by_due_date(ws, h, date(2026, 5, 22), header_row=match.header_row)
    ws.cell(due, h["fecha_pago"]).value = date(2026, 5, 10)
    ws.cell(due, h["abono_k"]).value = 50.0
    # Fila libre contigua debajo del último pago.
    ws.append([22, 6, 2026, None, None, None, None, None, None, None, None, None, None])
    ev = _sample_event()
    result = find_application_row_detailed(
        ws, h, ev, due_date_row=due, header_row=match.header_row
    )
    assert result.search_start_row == due + 1
    assert result.row == due + 1
    assert result.compare_status == APLICADO
    assert result.growth_blocked is False


def test_growth_blocked_by_other_section_label():
    wb = openpyxl.load_workbook(io.BytesIO(_amort_workbook_bytes()))
    match = detect_amortization_sheet(wb)
    ws, h = match.worksheet, match.headers
    due = 3
    ws.cell(3, h["fecha_pago"]).value = date(2026, 5, 10)
    ws.cell(3, h["abono_k"]).value = 99.0
    ws.cell(4, h["fecha_pago"]).value = "OTROS CONCEPTOS"
    ev = _sample_event()
    result = find_application_row_detailed(
        ws, h, ev, due_date_row=due, header_row=match.header_row
    )
    assert result.growth_blocked is True
    assert result.growth_block_code == APPLICATION_GROWTH_BLOCKED
    assert result.growth_block_row == 4
    assert result.growth_block_end_row == 4
    assert result.growth_block_value == "Fila 4"

def test_left_side_bound_ignores_far_garbage_without_schedule():
    """Basura en zona de pagos sin cronograma no mueve el último pago (fila 79)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
            "Saldo a capital",
        ]
    )
    for _ in range(2, 35):
        ws.append([1, 1, 2025, None, None, None, None, None])
    ws.cell(34, 4).value = date(2026, 6, 10)
    ws.cell(34, 6).value = 200.0
    ws.cell(34, 7).value = 200.0
    for r in range(35, 63):
        ws.append([20, 7, 2026, None, None, None, None, None])
    # Basura lejana sin cronograma (como ene-26 pegado en zona de pagos).
    for _ in range(63, 79):
        ws.append([None, None, None, date(2026, 1, 1), 10.0, 20.0, 30.0, None])
    headers = detect_headers(ws, header_row=1)
    assert find_bottom_most_occupied_payment_row(ws, headers, header_row=1) == 34
    ev = _sample_event()
    result = find_application_row_detailed(
        ws, headers, ev, due_date_row=35, header_row=1
    )
    assert result.search_start_row == 35
    assert result.row == 35
    assert result.compare_status == APLICADO
    assert result.growth_blocked is False


def test_left_side_bound_requires_new_row_when_schedule_exhausted():
    """Sin cronograma en la primera fila vacía → ampliar bloque; no saltar a basura."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
            "Saldo a capital",
        ]
    )
    ws.append([1, 1, 2026, date(2026, 1, 10), 10.0, 20.0, 30.0, None])  # 2
    ws.append([None, None, None, None, None, None, None, None])  # 3 sin cronograma
    # Basura abajo
    ws.append([None, None, None, "TOTAL", 1, 2, 3, None])
    headers = detect_headers(ws, header_row=1)
    assert find_bottom_most_occupied_payment_row(ws, headers, header_row=1) == 2
    result = find_application_row_detailed(
        ws, headers, _sample_event(), due_date_row=2, header_row=1
    )
    assert result.search_start_row == 3
    assert result.row is None
    assert result.requires_new_row is True
    assert result.suggested_row == 3


def test_detect_amortization_sheet_rejects_multiple_valid_sheets():
    wb = openpyxl.Workbook()
    for title in ("HojaA", "HojaB"):
        ws = wb.create_sheet(title) if title != "HojaA" else wb.active
        ws.title = title
        ws.append(
            [
                "dia",
                "mes",
                "año",
                "Fecha pago",
                "Valor intereses",
                "Abono a K",
                "Valor pagado cliente",
            ]
        )
        ws.append([1, 1, 2026, None, None, None, None])
    # openpyxl creates default sheet; ensure only our two
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    with pytest.raises(AmbiguousAmortizationHeadersError) as exc:
        detect_amortization_sheet(wb, tabla_amortizacion_path="TABLAS/x.xlsx")
    assert exc.value.key == "multiple_sheets"
    assert exc.value.error_code == "AMBIGUOUS_AMORTIZATION_HEADERS"


def test_detect_amortization_sheet_headers_split_across_rows():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SPLIT"
    ws.append(["dia", "mes", "año", "Valor intereses", "Abono a K"])
    ws.append(["Fecha pago", "Valor pagado cliente", None, None, None])
    with pytest.raises(AmortizationHeadersSplitAcrossRowsError):
        detect_amortization_sheet(wb, tabla_amortizacion_path="TABLAS/split.xlsx")


def test_orphan_vp_without_fecha_pago_does_not_push_write_past_due():
    """Caso Correa: basura en VP sin Fecha pago no mueve el cursor; escribe en la due."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Intereses de mora",
            "Valor pagado cliente",
            "Saldo a capital",
        ]
    )
    # Filas 2-4: pagos completos (análogo a 16-18).
    for i, (d, m, y, pay) in enumerate(
        (
            (25, 6, 2026, date(2026, 6, 4)),
            (25, 7, 2026, date(2026, 7, 7)),
            (25, 8, 2026, date(2026, 8, 11)),
        ),
        start=2,
    ):
        ws.append([d, m, y, pay, 1000.0 + i, 2000.0 + i, 50.0, 3000.0 + i, None])
    # Fila 5 = due (análogo a 19): cronograma, aplicación vacía.
    ws.append([25, 9, 2026, None, None, None, None, None, None])
    # Filas 6-8: cronograma vacío de aplicación.
    for m in (10, 11, 12):
        ws.append([25, m, 2026, None, None, None, None, None, None])
    # Fila 9 (análogo a 39): día/mes/año + solo Valor pagado huérfano (sin Fecha pago).
    ws.append([25, 4, 2028, None, None, None, None, 3_414_431.67, None])
    # Fila 10: sin cronograma (no debe elegirse).
    ws.append([None, None, None, None, None, None, None, None, None])

    headers = detect_headers(ws, header_row=1)
    assert is_complete_applied_payment_row(ws, 4, headers) is True
    assert is_incomplete_payment_orphan_row(ws, 9, headers) is True
    assert is_complete_applied_payment_row(ws, 9, headers) is False
    assert find_bottom_most_occupied_payment_row(ws, headers, header_row=1) == 4

    result = find_application_row_detailed(
        ws,
        headers,
        _sample_event(valor_pagado_cliente=3_985_120.67, capital=8_922_320.0, intereses=2_664_600.0),
        due_date_row=5,
        header_row=1,
        payment_date=date(2026, 9, 7),
    )
    assert result.search_start_row == 5
    assert result.row == 5
    assert result.compare_status == APLICADO
    assert result.requires_new_row is False


def test_fecha_pago_plus_only_valor_pagado_is_not_complete_occupancy():
    """VP numérico suelto (sin desglose de asiento) no marca la fila como ocupada."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([25, 8, 2026, date(2026, 8, 11), 10.0, 20.0, 30.0])  # 2 completo
    ws.append([25, 9, 2026, date(2026, 9, 7), None, None, 999.0])  # 3 solo fecha+VP
    headers = detect_headers(ws, header_row=1)
    assert is_complete_applied_payment_row(ws, 2, headers) is True
    assert is_complete_applied_payment_row(ws, 3, headers) is False
    assert is_incomplete_payment_orphan_row(ws, 3, headers) is True
    assert find_bottom_most_occupied_payment_row(ws, headers, header_row=1) == 2
    result = find_application_row_detailed(
        ws, headers, _sample_event(), due_date_row=3, header_row=1
    )
    assert result.row is None
    assert result.growth_blocked is True
    assert result.growth_block_row == 3
    assert result.growth_block_value == "Fila 3"
    assert result.growth_block_code == APPLICATION_GROWTH_BLOCKED
    assert result.compare_status != REVISION_MANUAL


def test_orphan_on_due_date_row_requires_revision_not_write_below():
    """Huérfano en la propia due → bloqueo de basura; no saltar a fila sin cronograma."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([25, 8, 2026, date(2026, 8, 11), 10.0, 20.0, 30.0])  # 2 pagado
    ws.append([25, 9, 2026, None, None, None, 999.0])  # 3 due huérfana (solo VP)
    ws.append([None, None, None, None, None, None, None])  # 4 sin cronograma

    headers = detect_headers(ws, header_row=1)
    assert is_incomplete_payment_orphan_row(ws, 3, headers) is True
    result = find_application_row_detailed(
        ws, headers, _sample_event(), due_date_row=3, header_row=1
    )
    assert result.row is None
    assert result.growth_blocked is True
    assert result.growth_block_row == 3
    assert result.growth_block_end_row == 3
    assert result.growth_block_value == "Fila 3"
    assert result.growth_block_code == APPLICATION_GROWTH_BLOCKED
    assert result.requires_new_row is False
    assert result.compare_status != REVISION_MANUAL


def test_formula_only_vp_without_fecha_is_not_complete_payment():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([25, 8, 2026, date(2026, 8, 1), 1.0, 2.0, 3.0])
    ws.append([25, 9, 2026, None, None, None, "=+AI10+AI11"])
    headers = detect_headers(ws, header_row=1)
    assert is_payment_application_empty(ws, 3, headers) is True
    assert is_incomplete_payment_orphan_row(ws, 3, headers) is False
    assert find_bottom_most_occupied_payment_row(ws, headers, header_row=1) == 2
    result = find_application_row_detailed(
        ws, headers, _sample_event(), due_date_row=3, header_row=1
    )
    assert result.row == 3
    assert result.compare_status == APLICADO


def test_garbage_between_last_valid_and_first_empty_blocks_and_does_not_adopt_below():
    """Basura tipo 295 (filas 5-8) bloquea; no se adopta ni escribe debajo."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
            "Saldo a capital",
        ]
    )
    ws.append([10, 8, 2026, date(2026, 8, 10), 10.0, 20.0, None, 30.0, None])  # 2
    ws.append([10, 9, 2026, date(2026, 9, 10), 11.0, 21.0, None, 32.0, None])  # 3
    ws.append([10, 10, 2026, date(2026, 10, 10), 12.0, 22.0, None, 34.0, None])  # 4 último válido
    ws.append([None, None, None, "FACTURACION", None, None, None, None, None])  # 5
    ws.append([None, None, None, "FECHA DESEM", None, 1.0, None, None, None])  # 6
    ws.append([10, 11, 2026, None, None, None, None, 999.0, None])  # 7 huérfano
    ws.append([None, None, None, None, None, 2.0, None, None, None])  # 8
    ws.append([10, 12, 2026, None, None, None, None, None, None])  # 9 vacía con cronograma
    ws.append([10, 1, 2027, None, None, None, None, None, None])  # 10

    headers = detect_headers(ws, header_row=1)
    ev = _sample_event(
        intereses=100.0,
        capital=200.0,
        mora=0.0,
        valor_pagado_cliente=300.0,
    )
    opts = PaymentApplicationWriteOptions(
        payment_date=date(2026, 12, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    write_payment_application(ws, 10, headers, ev, write_options=opts, header_row=1)

    result = find_application_row_detailed(
        ws,
        headers,
        ev,
        due_date_row=5,
        header_row=1,
        payment_date=opts.payment_date,
        detected_codes=opts.detected_codes,
    )
    assert result.row is None
    assert result.growth_blocked is True
    assert result.growth_block_row == 5
    assert result.growth_block_end_row == 8
    assert result.growth_block_value == "Filas 5 a 8"
    assert result.compare_status != ADOPTADO_EXISTENTE
    assert result.compare_status != APLICADO


def test_matching_last_valid_payment_adopts_despite_garbage_below():
    """Reintento idempotente: el pago ya aplicado se adopta aunque haya basura debajo."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
            "Saldo a capital",
        ]
    )
    ws.append([10, 8, 2026, date(2026, 8, 10), 10.0, 20.0, None, 30.0, None])  # 2
    ws.append([10, 9, 2026, date(2026, 9, 10), 11.0, 21.0, None, 32.0, None])  # 3
    ws.append([10, 10, 2026, date(2026, 10, 10), 12.0, 22.0, None, 34.0, None])  # 4
    ws.append([None, None, None, "FACTURACION", None, None, None, None, None])  # 5
    ws.append([None, None, None, "FECHA DESEM", None, 1.0, None, None, None])  # 6
    ws.append([10, 11, 2026, None, None, None, None, 999.0, None])  # 7
    ws.append([None, None, None, None, None, 2.0, None, None, None])  # 8
    ws.append([10, 12, 2026, None, None, None, None, None, None])  # 9

    headers = detect_headers(ws, header_row=1)
    ev = _sample_event(
        intereses=12.0,
        capital=22.0,
        mora=0.0,
        valor_pagado_cliente=34.0,
    )
    result = find_application_row_detailed(
        ws,
        headers,
        ev,
        due_date_row=5,
        header_row=1,
        payment_date=date(2026, 10, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    assert result.growth_blocked is False
    assert result.row == 4
    assert result.compare_status == ADOPTADO_EXISTENTE


def test_incomplete_junk_without_section_label_still_blocks_write():
    """Basura tipo 295 sin rótulo FACTURACION también cierra la ventana."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([10, 8, 2026, date(2026, 8, 10), 10.0, 20.0, 30.0])  # 2 válido
    ws.append([None, None, None, "jun-26", None, None, None])  # 3 basura
    ws.append([None, None, None, "FECHA DESEM", None, 1.0, None])  # 4
    ws.append([10, 9, 2026, None, None, None, 999.0])  # 5 huérfano
    ws.append([10, 10, 2026, None, None, None, None])  # 6 vacía con cronograma

    headers = detect_headers(ws, header_row=1)
    assert find_bottom_most_occupied_payment_row(ws, headers, header_row=1) == 2
    result = find_application_row_detailed(
        ws, headers, _sample_event(), due_date_row=6, header_row=1
    )
    assert result.row is None
    assert result.growth_blocked is True
    assert result.growth_block_row == 3
    assert result.growth_block_end_row == 5
    assert result.growth_block_value == "Filas 3 a 5"
    assert result.compare_status != ADOPTADO_EXISTENTE


def test_first_empty_with_schedule_still_writes_when_distant_junk_exists():
    """Primera vacía con cronograma es destino; basura lejana tras el hueco se ignora."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([10, 8, 2026, date(2026, 8, 10), 10.0, 20.0, 30.0])  # 2 válido
    ws.append([10, 9, 2026, None, None, None, None])  # 3 vacía con cronograma
    ws.append([10, 10, 2026, None, None, None, None])  # 4 vacía
    ws.append([None, None, None, None, None, None, None])  # 5
    ws.append([None, None, None, "FACTURACION", None, 1.0, 999.0])  # 6 basura lejana

    headers = detect_headers(ws, header_row=1)
    result = find_application_row_detailed(
        ws, headers, _sample_event(), due_date_row=3, header_row=1
    )
    assert result.growth_blocked is False
    assert result.row == 3
    assert result.compare_status == APLICADO

