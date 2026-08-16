"""Finalize v4: decisiones humanas sin distribución manual A/V/K."""
from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest

from app.application.services.finalize_aplicacion_pagos import (
    collect_aplicacion_pagos_issues,
    parse_editable_money,
    InvalidMonetaryValue,
)
from app.application.services.review_schema import (
    REVIEW_SCHEMA_INCONSISTENT,
    REVIEW_SCHEMA_REQUIRES_REGENERATION,
    AplicacionPagosCols,
    AplicacionPagosColsV3,
    ReviewSheets,
    TipoAplicacionConfirmado,
    ValidarPago,
    require_review_schema_v4,
)
from tests.fixtures.review_workbook_v3_legacy import build_review_workbook_v3_bytes
from datetime import date


def _row(**overrides):
    base = {
        AplicacionPagosCols.ID_PAGO: "p1",
        AplicacionPagosCols.CLIENTE: "C",
        AplicacionPagosCols.CREDITO: "1",
        AplicacionPagosCols.MONTO_BANCO: 100,
        AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.SI,
        AplicacionPagosCols.TIPO_APLICACION: TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL,
        "_excel_row": 4,
    }
    base.update(overrides)
    return base


def test_por_definir_blocks():
    issues = collect_aplicacion_pagos_issues(
        [_row(**{AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.POR_DEFINIR})]
    )
    assert any(i["error_code"] == "validar_pago_por_definir" for i in issues)


def test_si_with_valid_tipo_ok():
    assert collect_aplicacion_pagos_issues([_row()]) == []


def test_si_without_tipo_blocks():
    issues = collect_aplicacion_pagos_issues(
        [_row(**{AplicacionPagosCols.TIPO_APLICACION: ""})]
    )
    assert any(i["error_code"] == "tipo_aplicacion_required" for i in issues)


def test_no_with_tipo_blocks():
    issues = collect_aplicacion_pagos_issues(
        [
            _row(
                **{
                    AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.NO,
                    AplicacionPagosCols.TIPO_APLICACION: TipoAplicacionConfirmado.ABONO_A_CAPITAL,
                }
            )
        ]
    )
    assert any(i["error_code"] == "no_row_must_have_empty_tipo" for i in issues)


def test_no_with_empty_tipo_ok_when_another_si():
    no = _row(
        **{
            AplicacionPagosCols.CREDITO: "2",
            AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.NO,
            AplicacionPagosCols.TIPO_APLICACION: "",
            AplicacionPagosCols.MONTO_BANCO: 100,
            "_excel_row": 4,
        }
    )
    si = _row(
        **{
            AplicacionPagosCols.CREDITO: "1",
            AplicacionPagosCols.MONTO_BANCO: None,
            "_excel_row": 5,
        }
    )
    assert collect_aplicacion_pagos_issues([no, si]) == []


def test_f01_canonical_bank_amount_from_no_row():
    """F-01: monto canónico por ID Pago (primera cifra válida), no solo fila SI."""
    no = _row(
        **{
            AplicacionPagosCols.CREDITO: "A",
            AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.NO,
            AplicacionPagosCols.TIPO_APLICACION: "",
            AplicacionPagosCols.MONTO_BANCO: 250_000,
            "_excel_row": 4,
        }
    )
    si = _row(
        **{
            AplicacionPagosCols.CREDITO: "B",
            AplicacionPagosCols.MONTO_BANCO: None,
            "_excel_row": 5,
        }
    )
    assert collect_aplicacion_pagos_issues([no, si]) == []
    assert parse_editable_money(no[AplicacionPagosCols.MONTO_BANCO]) == 250_000.0


def test_multiple_credits_same_id_pago():
    rows = [
        _row(**{AplicacionPagosCols.CREDITO: "1", "_excel_row": 4}),
        _row(
            **{
                AplicacionPagosCols.CREDITO: "2",
                AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.NO,
                AplicacionPagosCols.TIPO_APLICACION: "",
                AplicacionPagosCols.MONTO_BANCO: None,
                "_excel_row": 5,
            }
        ),
        _row(
            **{
                AplicacionPagosCols.CREDITO: "3",
                AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.NO,
                AplicacionPagosCols.TIPO_APLICACION: "",
                AplicacionPagosCols.MONTO_BANCO: None,
                "_excel_row": 6,
            }
        ),
    ]
    assert collect_aplicacion_pagos_issues(rows) == []


def test_all_no_blocks_payment_without_selected_credit():
    rows = [
        _row(
            **{
                AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.NO,
                AplicacionPagosCols.TIPO_APLICACION: "",
                AplicacionPagosCols.CREDITO: "1",
            }
        ),
        _row(
            **{
                AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.NO,
                AplicacionPagosCols.TIPO_APLICACION: "",
                AplicacionPagosCols.CREDITO: "2",
                "_excel_row": 5,
            }
        ),
    ]
    issues = collect_aplicacion_pagos_issues(rows)
    assert any(i["error_code"] == "payment_without_selected_credit" for i in issues)


def test_v3_workbook_requires_regeneration():
    raw = build_review_workbook_v3_bytes(
        process_id="old",
        process_date=date(2026, 5, 20),
        bank_code="banco_bogota",
        aplicacion_rows=[],
        error_records=[],
    )
    wb = openpyxl.load_workbook(BytesIO(raw))
    with pytest.raises(ValueError, match=REVIEW_SCHEMA_REQUIRES_REGENERATION):
        require_review_schema_v4(wb)
    headers = [c.value for c in wb[ReviewSheets.APLICACION_PAGOS][3][:21]]
    assert headers == list(AplicacionPagosColsV3.HEADERS)


def test_unknown_schema_fail_closed():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(["ID Pago", "Cliente"])
    with pytest.raises(ValueError, match="unsupported_review_schema_version"):
        require_review_schema_v4(wb)


def test_finalize_locates_headers_not_positions():
    """No depende de posiciones históricas: header en fila 3, datos desde 4."""
    from app.application.use_cases.payment_validation_finalize import _find_table_header_row

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(["banner"])
    ws.append(["help"])
    ws.append(list(AplicacionPagosCols.HEADERS))
    ws.append(["p1", "C", "1", 100])
    assert _find_table_header_row(ws, AplicacionPagosCols.ID_PAGO) == 3


def test_invalid_bank_amount_text_still_fail_closed():
    with pytest.raises(InvalidMonetaryValue):
        parse_editable_money("5O00000")


def _wb_meta_and_headers(*, meta_version: int | None, headers: list[str]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(["banner"])
    ws.append(["help"])
    ws.append(headers)
    if meta_version is not None:
        ws_meta = wb.create_sheet(ReviewSheets.META)
        ws_meta.append(["Campo", "Valor"])
        ws_meta.append(["ReviewSchemaVersion", meta_version])
    return wb


def test_require_v4_ok_when_meta_and_headers_match():
    wb = _wb_meta_and_headers(meta_version=4, headers=list(AplicacionPagosCols.HEADERS))
    assert require_review_schema_v4(wb) == 4


def test_meta_3_requires_regeneration():
    wb = _wb_meta_and_headers(meta_version=3, headers=list(AplicacionPagosColsV3.HEADERS))
    with pytest.raises(ValueError, match=REVIEW_SCHEMA_REQUIRES_REGENERATION):
        require_review_schema_v4(wb)


def test_meta_4_with_v3_headers_inconsistent():
    wb = _wb_meta_and_headers(meta_version=4, headers=list(AplicacionPagosColsV3.HEADERS))
    with pytest.raises(ValueError, match=REVIEW_SCHEMA_INCONSISTENT):
        require_review_schema_v4(wb)


def test_meta_4_missing_required_header_inconsistent():
    headers = [h for h in AplicacionPagosCols.HEADERS if h != AplicacionPagosCols.TIPO_APLICACION]
    wb = _wb_meta_and_headers(meta_version=4, headers=headers)
    with pytest.raises(ValueError, match=REVIEW_SCHEMA_INCONSISTENT):
        require_review_schema_v4(wb)


def test_unknown_meta_version_unsupported():
    wb = _wb_meta_and_headers(meta_version=99, headers=list(AplicacionPagosCols.HEADERS))
    with pytest.raises(ValueError, match="unsupported_review_schema_version"):
        require_review_schema_v4(wb)
