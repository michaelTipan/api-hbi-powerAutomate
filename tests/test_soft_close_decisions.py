"""Cierre sin amortizar con registro de decisiones + compatibilidad del Control antiguo."""

from __future__ import annotations

import asyncio
import io
import json
from types import SimpleNamespace

import openpyxl
import pytest

from app.application.services import amortization_decisions as decisions
from app.application.services.amortization_decisions import (
    MOTIVO_OTRO,
    MOTIVOS_EXCLUSION,
    add_exclusion,
    add_unused_asiento,
    empty_registry,
    registry_filename,
)
from app.application.ui.schemas import (
    UiAmortizationDecisions,
    UiExcludedCredit,
    UiLink,
)
from app.application.use_cases.payment_validation_process_control import (
    read_process_control_snapshot,
    update_process_control_row2,
)
from app.application.use_cases.payment_validation_soft_close import (
    SOFT_CLOSE_ESTADO,
    SOFT_CLOSE_PARTIAL_ESTADO,
    soft_close_payment_validation,
)
from app.application.use_cases.setup_merge_control_workbook import (
    PROCESS_CONTROL_BANK_FILE_BOGOTA,
    SHEET_NAME,
    _build_process_control_workbook_bytes,
)
from tests.test_soft_close_process_control import (
    MockGraphClientSoftClose,
    _control_late_phase,
    _read_control_row,
    _set_env,
)

PK = "payment-validation|banco_bogota|2026-06-01|abc12345"
REGISTRY_PATH = f"merge/{registry_filename('banco_bogota', '2026-06-01', PK)}"
MANIFEST_PATH = "merge/m.json"
FOLDER_258 = "clientes/EMPRESA A/CREDITO # 258/ASIENTOS CONTABLES CRED 258"
FOLDER_300 = "clientes/EMPRESA B/CREDITO # 300/ASIENTOS CONTABLES CRED 300"
ASIENTO_258 = f"{FOLDER_258}/asiento_258.pdf"
ASIENTO_300 = f"{FOLDER_300}/asiento_300.pdf"
ASIENTO_OTRO = f"{FOLDER_258}/asiento_de_otro_credito.pdf"
MOTIVO = next(code for code in MOTIVOS_EXCLUSION if code != MOTIVO_OTRO)


class _GraphWithIfMatch(MockGraphClientSoftClose):
    async def put_bytes(self, endpoint: str, content: bytes, content_type: str = "", if_match=None):
        return await super().put_bytes(endpoint, content, content_type)


@pytest.fixture(autouse=True)
def _clean_registry_cache():
    decisions.clear_registry_cache()
    yield
    decisions.clear_registry_cache()


def _manifest() -> dict:
    return {
        "report_date_iso": "2026-06-01",
        "outputs": [
            {
                "status": "COMPLETE",
                "id_pago": "P1",
                "fecha_banco": "2026-06-01",
                "credito": "CREDITO # 258",
                "asiento_pdf_paths": [ASIENTO_258, ASIENTO_OTRO],
            },
            {
                "status": "COMPLETE",
                "id_pago": "P2",
                "fecha_banco": "2026-06-01",
                "credito": "CREDITO # 300",
                "asiento_pdf_paths": [ASIENTO_300],
            },
        ],
    }


def _registry() -> dict:
    reg = empty_registry(process_key=PK, bank_code="banco_bogota", report_date_iso="2026-06-01")
    reg, _ = add_exclusion(
        reg,
        credits=["300"],
        clientes={"300": "EMPRESA B"},
        id_pagos=["P2"],
        tablas={"300": "clientes/EMPRESA B/CREDITO # 300/tabla_300.xlsx"},
        motivo_codigo=MOTIVO,
        nota="",
        quien="operador",
    )
    reg, _ = add_unused_asiento(
        reg,
        path=ASIENTO_OTRO,
        credito_carpeta="258",
        nota="Asiento del crédito 999 guardado en la carpeta del 258",
        quien="operador",
    )
    return reg


def _control_with_registry(*, estado: str, registry_path: str = REGISTRY_PATH) -> bytes:
    raw = _control_late_phase(
        estado=estado,
        process_key=PK,
        merge_manifest_path=MANIFEST_PATH,
        apply_idempotency_key=PK if estado == "AMORTIZACION_PARCIAL" else "",
    )
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        ws = wb[SHEET_NAME]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        ws.cell(2, headers.index("DecisionsRegistryPath") + 1, value=registry_path)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()


def _client(*, estado: str = "AMORTIZACION_PARCIAL", registry: dict | None = None) -> _GraphWithIfMatch:
    client = _GraphWithIfMatch()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_with_registry(estado=estado)
    client.downloaded_files[MANIFEST_PATH] = json.dumps(_manifest()).encode("utf-8")
    client.downloaded_files[ASIENTO_258] = b"%PDF-258"
    client.downloaded_files[ASIENTO_300] = b"%PDF-300"
    client.downloaded_files[ASIENTO_OTRO] = b"%PDF-otro"
    if registry is not None:
        client.downloaded_files[REGISTRY_PATH] = json.dumps(registry).encode("utf-8")
    return client


def _procesados(client, folder: str) -> dict[str, bytes]:
    prefix = f"{folder}/PROCESADOS/"
    return {
        path[len(prefix):]: blob
        for path, blob in client.downloaded_files.items()
        if path.startswith(prefix) and "/" not in path[len(prefix):] and path.endswith(".pdf")
    }


def _close(client) -> dict:
    return asyncio.run(
        soft_close_payment_validation(client, bank_code="banco_bogota", process_key=PK, reason="")
    )


def test_soft_close_with_excluded_credit_moves_asientos_and_records_closure():
    _set_env()
    client = _client(registry=_registry())

    res = _close(client)

    assert res["process_control_estado"] == SOFT_CLOSE_PARTIAL_ESTADO
    assert not client.delete_calls
    for src in (ASIENTO_258, ASIENTO_300, ASIENTO_OTRO):
        assert src not in client.downloaded_files
    assert _procesados(client, FOLDER_258) == {"2026-06-01_P1__asiento_258.pdf": b"%PDF-258"}
    assert _procesados(client, FOLDER_300) == {"2026-06-01_P2__asiento_300.pdf": b"%PDF-300"}
    assert (
        client.downloaded_files[f"{FOLDER_258}/PROCESADOS/NO USADOS/asiento_de_otro_credito.pdf"]
        == b"%PDF-otro"
    )

    assert [c["credito"] for c in res["creditos_sin_amortizar"]] == ["300"]
    assert res["creditos_sin_amortizar"][0]["cliente"] == "EMPRESA B"
    assert res["creditos_sin_amortizar"][0]["motivo"] == MOTIVOS_EXCLUSION[MOTIVO]
    moves = res["asientos_no_usados_moves"]
    assert len(moves) == 1 and moves[0]["status"] == "moved"
    assert "crédito 999" in moves[0]["nota"]

    saved = json.loads(client.downloaded_files[REGISTRY_PATH].decode("utf-8"))
    assert saved["closure"]["estado"] == SOFT_CLOSE_PARTIAL_ESTADO
    assert [c["credito"] for c in saved["closure"]["creditos_sin_amortizar"]] == ["300"]
    assert saved["closure"]["movimientos_asientos"][0]["status"] == "moved"
    assert saved["exclusions"][0]["status"] == "active"

    row = _read_control_row(client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA])
    assert row["EstadoProceso"] == SOFT_CLOSE_PARTIAL_ESTADO
    assert row["DecisionsRegistryPath"] == REGISTRY_PATH


def _payment_v2(id_pago: str, credito: str, paths: list[str]) -> dict:
    return {
        "status": "COMPLETE",
        "id_pago": id_pago,
        "fecha_banco": "2026-06-01",
        "credit_items": [{"credito": credito, "asiento_pdf_paths": paths}],
    }


def _consolidado_v2(*payments: dict) -> dict:
    return {
        "status": "COMPLETE",
        "id_pagos": [p["id_pago"] for p in payments],
        "fecha_banco": "2026-06-01",
        "consolidado_key": "x|2026-06-01",
        "nombre_consolidado": "01 JUNIO PAGO CUOTA EMPRESA CRED 1.pdf",
        "payments": list(payments),
    }


def _assert_close_moved_everything(client) -> None:
    assert not client.delete_calls
    for src in (ASIENTO_258, ASIENTO_300, ASIENTO_OTRO):
        assert src not in client.downloaded_files
    assert _procesados(client, FOLDER_258) == {"2026-06-01_P1__asiento_258.pdf": b"%PDF-258"}
    assert _procesados(client, FOLDER_300) == {"2026-06-01_P2__asiento_300.pdf": b"%PDF-300"}
    assert (
        client.downloaded_files[f"{FOLDER_258}/PROCESADOS/NO USADOS/asiento_de_otro_credito.pdf"]
        == b"%PDF-otro"
    )


def test_soft_close_manifest_v2_exclusion_en_fase_4_mueve_igual_que_v1():
    _set_env()
    client = _client(registry=_registry())
    manifest = {
        "manifest_version": 2,
        "report_date_iso": "2026-06-01",
        "outputs": [
            _consolidado_v2(_payment_v2("P1", "258", [ASIENTO_258])),
            _consolidado_v2(_payment_v2("P2", "300", [ASIENTO_300])),
        ],
        "excluded_groups": [],
    }
    client.downloaded_files[MANIFEST_PATH] = json.dumps(manifest).encode("utf-8")

    res = _close(client)

    assert res["process_control_estado"] == SOFT_CLOSE_PARTIAL_ESTADO
    _assert_close_moved_everything(client)
    assert [c["credito"] for c in res["creditos_sin_amortizar"]] == ["300"]


def test_soft_close_exclusion_en_fase_3_mueve_sus_asientos_a_procesados():
    _set_env()
    reg = empty_registry(process_key=PK, bank_code="banco_bogota", report_date_iso="2026-06-01")
    reg, _ = add_exclusion(
        reg, credits=["300"], clientes={"300": "EMPRESA B"}, id_pagos=["P2"], tablas={},
        motivo_codigo=MOTIVO, nota="", quien="operador", fase="consolidacion",
    )
    reg, _ = add_unused_asiento(
        reg, path=ASIENTO_OTRO, credito_carpeta="258", nota="de otro crédito", quien="operador",
    )
    client = _client(estado="CONSOLIDADO", registry=reg)
    manifest = {
        "manifest_version": 2,
        "report_date_iso": "2026-06-01",
        "outputs": [_consolidado_v2(_payment_v2("P1", "258", [ASIENTO_258]))],
        "excluded_groups": [
            {
                "id_pago": "P2",
                "creditos": ["300"],
                "fecha_banco": "2026-06-01",
                "asiento_items": [{"path": ASIENTO_300, "credito": "300"}],
            }
        ],
    }
    client.downloaded_files[MANIFEST_PATH] = json.dumps(manifest).encode("utf-8")

    res = _close(client)

    assert res["process_control_estado"] == SOFT_CLOSE_ESTADO
    _assert_close_moved_everything(client)
    assert res["accounting_pdfs_moved_count"] == 2
    assert [c["credito"] for c in res["creditos_sin_amortizar"]] == ["300"]


def test_soft_close_never_processed_with_exclusions_is_sin_amortizar():
    _set_env()
    client = _client(estado="CONSOLIDADO", registry=_registry())

    res = _close(client)

    assert res["process_control_estado"] == SOFT_CLOSE_ESTADO
    assert [c["credito"] for c in res["creditos_sin_amortizar"]] == ["300"]
    saved = json.loads(client.downloaded_files[REGISTRY_PATH].decode("utf-8"))
    assert saved["closure"]["estado"] == SOFT_CLOSE_ESTADO


def test_soft_close_unreadable_registry_does_not_block_close():
    _set_env()
    client = _client(registry=None)
    client.downloaded_files[REGISTRY_PATH] = b"{no es json"

    res = _close(client)

    assert res["process_control_estado"] == SOFT_CLOSE_PARTIAL_ESTADO
    assert res["creditos_sin_amortizar"] == []
    assert res["asientos_no_usados_moves"] == []
    assert not client.delete_calls
    # Sin registro, el asiento vuelve al flujo normal de usados (nunca se borra).
    assert "2026-06-01_P1__asiento_de_otro_credito.pdf" in _procesados(client, FOLDER_258)
    assert not any("NO USADOS" in path for path in client.downloaded_files)
    assert client.downloaded_files[REGISTRY_PATH] == b"{no es json"


def test_soft_close_missing_registry_file_does_not_block_close():
    _set_env()
    client = _client(registry=None)

    res = _close(client)

    assert res["process_control_estado"] == SOFT_CLOSE_PARTIAL_ESTADO
    assert res["creditos_sin_amortizar"] == []
    assert REGISTRY_PATH not in client.downloaded_files


def test_soft_close_ignores_registry_of_previous_process():
    _set_env()
    stale_path = f"merge/{registry_filename('banco_bogota', '2026-05-01', 'payment-validation|banco_bogota|2026-05-01|zzz99999')}"
    client = _client(registry=None)
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_with_registry(
        estado="AMORTIZACION_PARCIAL", registry_path=stale_path
    )
    stale = _registry()
    stale["process_key"] = "payment-validation|banco_bogota|2026-05-01|zzz99999"
    client.downloaded_files[stale_path] = json.dumps(stale).encode("utf-8")

    res = _close(client)

    assert res["creditos_sin_amortizar"] == []
    assert res["asientos_no_usados_moves"] == []
    assert json.loads(client.downloaded_files[stale_path].decode("utf-8"))["closure"] is None


# --- Control antiguo sin columna DecisionsRegistryPath ----------------------------


def _old_control_without_column() -> bytes:
    raw = _build_process_control_workbook_bytes("banco_bogota", "Banco de Bogotá")
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        ws = wb[SHEET_NAME]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        ws.delete_cols(headers.index("DecisionsRegistryPath") + 1)
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        ws.cell(2, headers.index("ProcessKey") + 1, value=PK)
        ws.cell(2, headers.index("EstadoProceso") + 1, value="CONSOLIDADO")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()


def test_old_control_workbook_without_column_reads_and_appends():
    _set_env()
    client = MockGraphClientSoftClose()
    old = _old_control_without_column()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = old
    old_row = _read_control_row(old)
    assert "DecisionsRegistryPath" not in old_row

    snap = asyncio.run(
        read_process_control_snapshot(client, "dummy_site", "dummy_drive", bank_code="banco_bogota")
    )
    assert snap.decisions_registry_path == ""
    assert snap.estado_proceso == "CONSOLIDADO"

    asyncio.run(
        update_process_control_row2(
            client,
            "dummy_site",
            "dummy_drive",
            bank_code="banco_bogota",
            updates={"DecisionsRegistryPath": REGISTRY_PATH},
        )
    )
    new_row = _read_control_row(client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA])
    assert new_row["DecisionsRegistryPath"] == REGISTRY_PATH
    for key, value in old_row.items():
        assert new_row[key] == value
    snap2 = asyncio.run(
        read_process_control_snapshot(client, "dummy_site", "dummy_drive", bank_code="banco_bogota")
    )
    assert snap2.decisions_registry_path == REGISTRY_PATH


def test_old_control_workbook_soft_close_without_registry():
    _set_env()
    client = MockGraphClientSoftClose()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _old_control_without_column()

    res = _close(client)

    assert res["process_control_estado"] == SOFT_CLOSE_ESTADO
    assert res["creditos_sin_amortizar"] == []


def test_local_control_payload_without_key(monkeypatch, tmp_path):
    from app.application.process_control_store import persist_control_payload
    from app.application.use_cases.payment_validation_process_control import (
        _payload_from_cells,
        canonical_empty_control_cells,
    )

    monkeypatch.setenv("PROCESS_CONTROL_STORE", "local")
    monkeypatch.setenv("ACTIVE_ENVIRONMENT", "sandbox")
    monkeypatch.setenv("PAYMENT_VALIDATION_CONTROL_STORE_DIR", str(tmp_path))
    cells = canonical_empty_control_cells("banco_bogota")
    cells.pop("DecisionsRegistryPath", None)
    cells["ProcessKey"] = PK
    persist_control_payload(
        _payload_from_cells(bank_code="banco_bogota", control_file_path="control.xlsx", cells=cells)
    )
    graph = MockGraphClientSoftClose()

    async def _run():
        snap = await read_process_control_snapshot(graph, "s", "d", bank_code="banco_bogota")
        assert snap.decisions_registry_path == ""
        await update_process_control_row2(
            graph, "s", "d", bank_code="banco_bogota", updates={"DecisionsRegistryPath": REGISTRY_PATH}
        )
        snap2 = await read_process_control_snapshot(graph, "s", "d", bank_code="banco_bogota")
        assert snap2.decisions_registry_path == REGISTRY_PATH
        assert snap2.process_key == PK

    asyncio.run(_run())
    assert not graph.put_calls


# --- Detalle: enriquecimiento con decisiones -------------------------------------


class _Reader:
    def __init__(self) -> None:
        self.calls: list[str] = []

    async def get_web_url(self, path: str):
        self.calls.append(path)
        return f"https://sp/{path}"


def _query_service(reader, *, graph=object(), assess=True):
    from app.application.ui.process_query import UiProcessQueryService

    return UiProcessQueryService(reader, graph=graph, assess_amortization=assess)


def _control(registry_path: str = REGISTRY_PATH, estado: str = "AMORTIZACION_PARCIAL"):
    snap = SimpleNamespace(
        decisions_registry_path=registry_path, process_key=PK, estado_proceso=estado
    )
    return SimpleNamespace(snapshot=snap)


def _view() -> UiAmortizationDecisions:
    link = UiLink(rel="table", label="Abrir tabla", path="clientes/EMPRESA B/tabla_300.xlsx")
    return UiAmortizationDecisions(
        excluded=[
            UiExcludedCredit(credit="300", motivo="m", motivo_codigo=MOTIVO, table_link=link)
        ]
    )


@pytest.fixture
def _patch_detail(monkeypatch):
    import app.application.sharepoint_resolution as sp
    import app.application.ui.amortization_decisions_service as svc

    seen: dict = {}

    async def _ctx(_graph):
        return {"site_id": "s", "drive_id": "d"}

    async def _load(graph, **kwargs):
        seen.update(kwargs)
        return seen.get("view", _view())

    monkeypatch.setattr(sp, "resolve_sharepoint_from_env", _ctx)
    monkeypatch.setattr(svc, "load_decisions_for_detail", _load)
    return seen


def test_detail_enrichment_fills_view_and_table_urls(_patch_detail):
    reader = _Reader()
    detail = SimpleNamespace(last_amortization_attempt=None, amortization_decisions=None)

    asyncio.run(_query_service(reader)._enrich_amortization_decisions(detail, _control()))

    view = detail.amortization_decisions
    assert view is not None
    assert view.excluded[0].table_link.web_url == "https://sp/clientes/EMPRESA B/tabla_300.xlsx"
    assert _patch_detail["closed"] is False
    assert _patch_detail["site_id"] == "s"


def test_detail_enrichment_marks_closed_states(_patch_detail):
    detail = SimpleNamespace(last_amortization_attempt=None, amortization_decisions=None)
    asyncio.run(
        _query_service(_Reader())._enrich_amortization_decisions(
            detail, _control(estado=SOFT_CLOSE_PARTIAL_ESTADO)
        )
    )
    assert _patch_detail["closed"] is True


@pytest.mark.parametrize(
    "control,assess,graph",
    [
        (_control(registry_path=""), True, object()),
        (_control(registry_path="merge/decisiones_banco_bogota_2026-05-01_zzz99999.json"), True, object()),
        (_control(), False, object()),
        (_control(), True, None),
    ],
)
def test_detail_enrichment_skips(_patch_detail, control, assess, graph):
    detail = SimpleNamespace(last_amortization_attempt=None, amortization_decisions=None)
    asyncio.run(
        _query_service(_Reader(), graph=graph, assess=assess)._enrich_amortization_decisions(
            detail, control
        )
    )
    assert detail.amortization_decisions is None
    assert "closed" not in _patch_detail


def test_detail_enrichment_load_error_is_silent(monkeypatch, _patch_detail):
    import app.application.ui.amortization_decisions_service as svc

    async def _boom(graph, **kwargs):
        raise RuntimeError("graph down")

    monkeypatch.setattr(svc, "load_decisions_for_detail", _boom)
    detail = SimpleNamespace(last_amortization_attempt=None, amortization_decisions=None)
    asyncio.run(_query_service(_Reader())._enrich_amortization_decisions(detail, _control()))
    assert detail.amortization_decisions is None
