"""Notify: parse CORREOS even when data_only cache is empty."""
from __future__ import annotations

import asyncio
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import Workbook, load_workbook

from app.application.use_cases.send_validar_extractos_notification import (
    _load_sender_and_recipients_from_correos_xlsx,
    _parse_correos_workbook,
    _recipients_excluding_sender,
)
from scripts.e2e_rc.fixtures_catalog import build_sandbox_correos_xlsx


def _empty_header_wb() -> Workbook:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "CORREOS"
    ws["A1"] = "EMISOR"
    ws["B1"] = "RECEPTORES"
    return wb


def test_parse_correos_from_canonical_builder() -> None:
    raw = build_sandbox_correos_xlsx()
    wb = load_workbook(BytesIO(raw), data_only=True)
    sender, recs = _parse_correos_workbook(wb)
    assert sender.endswith("@hbicapital.com.co")
    assert recs == ["herramientas.jsakedev@gmail.com"]


def test_parse_correos_literal_values_when_cached_empty_would_fail() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "CORREOS"
    ws["A1"] = "EMISOR"
    ws["B1"] = "RECEPTORES"
    ws["A2"] = "herramientas.jsakedev@gmail.com"
    ws["B2"] = "herramientas.jsakedev@gmail.com"
    buf = BytesIO()
    wb.save(buf)
    parsed = _parse_correos_workbook(load_workbook(BytesIO(buf.getvalue()), data_only=False))
    assert parsed[0] == "herramientas.jsakedev@gmail.com"
    assert parsed[1] == ["herramientas.jsakedev@gmail.com"]


def test_load_correos_falls_back_when_data_only_cache_empty() -> None:
    empty = _empty_header_wb()
    ok = Workbook()
    ws = ok.active
    assert ws is not None
    ws.title = "CORREOS"
    ws["A1"] = "EMISOR"
    ws["B1"] = "RECEPTORES"
    ws["A2"] = "herramientas.jsakedev@gmail.com"
    ws["B2"] = "herramientas.jsakedev@gmail.com"

    async def _run() -> tuple[str, list[str]]:
        with (
            patch(
                "app.application.use_cases.send_validar_extractos_notification._graph_download_by_path",
                new=AsyncMock(return_value=b"xlsx"),
            ),
            patch(
                "app.application.config.payment_validation_settings.resolve_correos_xlsx_path",
                return_value="CORREOS.xlsx",
            ),
            patch(
                "app.application.use_cases.send_validar_extractos_notification.load_workbook",
                side_effect=[empty, ok],
            ),
        ):
            return await _load_sender_and_recipients_from_correos_xlsx(
                SimpleNamespace(), "site", "drive"
            )

    sender, recs = asyncio.run(_run())
    assert sender == "herramientas.jsakedev@gmail.com"
    assert recs == ["herramientas.jsakedev@gmail.com"]


def test_recipients_keep_sender_when_it_is_the_only_mailbox() -> None:
    only = "herramientas.jsakedev@gmail.com"
    assert _recipients_excluding_sender(only, [only]) == [only]
    assert _recipients_excluding_sender(
        only, [only, "otro@example.com"]
    ) == ["otro@example.com"]
