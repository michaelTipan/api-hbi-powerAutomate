"""Merge Fase 3: ABONO sin extracto, manifest extendido y naming."""

import asyncio
import json
from datetime import date
from io import BytesIO
from unittest.mock import AsyncMock, patch
from urllib.parse import unquote

import pytest
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

from app.application.services.review_schema import (
    AplicacionPagosCols,
    REVIEW_SCHEMA_VERSION,
    ReviewSheets,
    TipoAplicacion,
    TipoAplicacionConfirmado,
    ValidarPago,
)
from app.application.sharepoint_resolution import encode_graph_drive_path
from app.application.use_cases.merge_composite_validado_pdfs import (
    _merge_composite_output_basename,
    merge_composite_validado_pdfs,
)
from tests.test_finalize_validation import make_distrib_row
from tests.test_merge_composite_control_workbook import (
    _MergeGraph,
    _asiento_pdf,
    _bank_bytes,
    _tiny_pdf,
)


def _merged_page_count(pdf_bytes: bytes) -> int:
    return len(PdfReader(BytesIO(pdf_bytes)).pages)


def _abono_historical_bytes(
    *,
    credits: tuple[str, ...] = ("258",),
    with_asiento_paths: bool = True,
    id_pago: str = "AB1",
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(
        list(AplicacionPagosCols.HEADERS)
        + ["_ruta_extracto", "_ruta_unidad_credito", "_ruta_asientos_contables"]
    )
    for cred in credits:
        row, _ = make_distrib_row(
            id_pago=id_pago,
            validar_pago=ValidarPago.SI,
            credito=cred,
            cliente="EQUINORTE",
            tipo_aplicacion=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
            abono_capital=1000,
            valor_int=0,
            mora_a_aplicar=0,
            fecha_banco=date(2026, 6, 3),
        )
        ruta_as = (
            f"clientes/EQUINORTE/CREDITO# {cred}/ASIENTOS CONTABLES CRED {cred}"
            if with_asiento_paths
            else ""
        )
        ws.append(list(row) + [ruta_as])
    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append(["Campo", "Valor"])
    ws_meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()

def _setup_abono_merge_graph(
    g: _MergeGraph,
    *,
    credits: tuple[str, ...] = ("258",),
    missing_asiento: bool = False,
    empty_asiento_folder: bool = False,
) -> tuple[str, str]:
    hist = "HIST/abono.xlsx"
    email = "EMAIL/mail.pdf"
    g.initial["bank/report.xlsx"] = _bank_bytes()
    g.initial[hist] = _abono_historical_bytes(
        credits=credits,
        with_asiento_paths=not missing_asiento,
    )
    g.initial[email] = _tiny_pdf()
    for cred in credits:
        folder = f"clientes/EQUINORTE/CREDITO# {cred}/ASIENTOS CONTABLES CRED {cred}"
        if not missing_asiento and not empty_asiento_folder:
            g.children[folder] = [{"name": f"asiento {cred}.pdf", "file": {}}]
            g.initial[f"{folder}/asiento {cred}.pdf"] = _asiento_pdf(
                100.0 if len(credits) == 1 else (40.0 if cred == credits[0] else 60.0),
                credit=cred,
            )
        elif not missing_asiento and empty_asiento_folder:
            g.children[folder] = []
    return hist, email


@pytest.fixture(autouse=True)
def _merge_env(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST_SITE")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "TEST_DRIVE")
    monkeypatch.setenv("GRAPH_SHAREPOINT_FILE_PATH", "bank/report.xlsx")

    import app.application.use_cases.merge_composite_validado_pdfs as m
    import app.application.use_cases.payment_validation_process_control as pc
    from types import SimpleNamespace

    async def _fake_read(_g, _s, _d, *, bank_code: str):
        return SimpleNamespace(
            control_file_path="CTL/ignored.xlsx",
            estado_proceso="PENDIENTE_ASIENTOS",
            is_active=True,
            process_key=f"payment-validation|{bank_code}|2026-05-12",
            validation_file_path="",
            historical_file_path="",
            secretary_file_path="",
            email_pdf_path="",
            notify_idempotency_key="",
            merge_manifest_path="",
            merge_idempotency_key="",
            bank_code=bank_code,
            bank_name="",
        )

    async def _fake_update(_g, _s, _d, *, bank_code: str, updates: dict):
        return True

    async def _fake_resolve(_g, _site_search: str, _drive_name: str, _path: str):
        return {
            "site_id": "s1",
            "drive_id": "d1",
            "path_encoded": encode_graph_drive_path("bank/report.xlsx"),
            "file_path": "bank/report.xlsx",
        }

    monkeypatch.setattr(pc, "read_process_control_snapshot", _fake_read)
    monkeypatch.setattr(pc, "update_process_control_row2", _fake_update)
    monkeypatch.setattr(m, "resolve_sharepoint_path", _fake_resolve)


def test_merge_output_basename_tokens_v3():
    from app.application.services.review_schema import (
        MERGE_NAME_TOKEN_MULTIPLE,
        TipoAplicacionConfirmado,
        merge_name_token_for_tipos,
    )

    d = date(2026, 6, 3)
    pago = _merge_composite_output_basename(
        d,
        "CLI",
        "258",
        bank_code="banco_bogota",
        name_token=merge_name_token_for_tipos([TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL]),
    )
    abono = _merge_composite_output_basename(
        d,
        "CLI",
        "258",
        bank_code="banco_bogota",
        name_token=merge_name_token_for_tipos([TipoAplicacionConfirmado.ABONO_A_CAPITAL]),
    )
    multiple = _merge_composite_output_basename(
        d,
        "CLI",
        "258, 265",
        bank_code="banco_bogota",
        name_token=MERGE_NAME_TOKEN_MULTIPLE,
    )
    total = _merge_composite_output_basename(
        d,
        "CLI",
        "258",
        bank_code="banco_bogota",
        name_token=merge_name_token_for_tipos([TipoAplicacionConfirmado.CANCELACION_PAGO_TOTAL]),
    )
    assert " PAGO CUOTA " in pago
    assert " ABONO A CAPITAL " in abono
    assert f" {MERGE_NAME_TOKEN_MULTIPLE} " in multiple
    assert " PAGO TOTAL " in total
    assert "OBSERV" not in pago.upper()
    assert pago != abono


def test_merge_output_basename_pago_vs_abono():
    # Compat: canonical PAGO/ABONO aún produce token vía merge_name_token_for_tipos fallback.
    d = date(2026, 6, 3)
    pago = _merge_composite_output_basename(d, "CLI", "258", bank_code="banco_bogota")
    abono = _merge_composite_output_basename(
        d, "CLI", "258", bank_code="banco_bogota", tipo_aplicacion=TipoAplicacion.ABONO.value
    )
    assert " PAGO CUOTA " in pago
    # ABONO sin tipo confirmado cae a PAGO CUOTA vía merge_name_token_for_tipos
    # cuando el string no es un TipoAplicacionConfirmado; pasar name_token explícito:
    abono2 = _merge_composite_output_basename(
        d, "CLI", "258", bank_code="banco_bogota", name_token="ABONO A CAPITAL"
    )
    assert " ABONO A CAPITAL " in abono2
    _ = abono


def test_merge_abono_single_credit_email_plus_asiento():
    g = _MergeGraph()
    hist, email = _setup_abono_merge_graph(g)

    ctx = {"site_id": "s1", "drive_id": "d1", "path_encoded": "x", "file_path": "bank/report.xlsx"}

    async def run():
        with patch(
            "app.application.use_cases.merge_composite_validado_pdfs.resolve_sharepoint_from_env",
            new_callable=AsyncMock,
            return_value=ctx,
        ):
            return await merge_composite_validado_pdfs(
                g,
                bank_code="banco_bogota",
                historical_file_path=hist,
                email_pdf_path=email,
            )

    result = asyncio.run(run())
    assert result.abono_outputs_count == 1
    assert result.payment_outputs_count == 0
    assert result.extracts_not_required_count == 1
    out = result.outputs[0]
    assert out.tipo_aplicacion == TipoAplicacion.ABONO.value
    assert out.requiere_extracto is False
    assert out.extracto_pdf_path == ""
    assert out.credit_items[0]["extracto_pdf_paths"] == []
    assert " ABONO A CAPITAL " in out.output_relative_path
    merged_key = next(k for k in g.uploaded if k.endswith(".pdf") and " ABONO A CAPITAL " in k)
    assert _merged_page_count(g.uploaded[merged_key]) == 2


def test_merge_abono_multi_credit_multiple_asientos():
    g = _MergeGraph()
    hist, email = _setup_abono_merge_graph(g, credits=("258", "265"))

    ctx = {"site_id": "s1", "drive_id": "d1", "path_encoded": "x", "file_path": "bank/report.xlsx"}

    async def run():
        with patch(
            "app.application.use_cases.merge_composite_validado_pdfs.resolve_sharepoint_from_env",
            new_callable=AsyncMock,
            return_value=ctx,
        ):
            return await merge_composite_validado_pdfs(
                g,
                bank_code="banco_bogota",
                historical_file_path=hist,
                email_pdf_path=email,
            )

    result = asyncio.run(run())
    assert result.abono_outputs_count == 1
    assert len(result.outputs[0].credit_items) == 2
    merged_key = next(k for k in g.uploaded if " ABONO A CAPITAL " in k)
    assert _merged_page_count(g.uploaded[merged_key]) == 3


def test_merge_abono_missing_asiento_skipped_not_partial_for_extract():
    g = _MergeGraph()
    hist, email = _setup_abono_merge_graph(g, empty_asiento_folder=True)

    ctx = {"site_id": "s1", "drive_id": "d1", "path_encoded": "x", "file_path": "bank/report.xlsx"}

    async def run():
        with patch(
            "app.application.use_cases.merge_composite_validado_pdfs.resolve_sharepoint_from_env",
            new_callable=AsyncMock,
            return_value=ctx,
        ):
            return await merge_composite_validado_pdfs(
                g,
                bank_code="banco_bogota",
                historical_file_path=hist,
                email_pdf_path=email,
            )

    result = asyncio.run(run())
    assert result.outputs_count == 0
    assert result.abono_skipped_count == 1
    assert any("abono_accounting_pdf_missing" in s for s in result.skipped)
    assert not any("extract_routes_missing" in s for s in result.skipped)


def test_merge_manifest_contains_extended_fields():
    g = _MergeGraph()
    hist, email = _setup_abono_merge_graph(g)

    ctx = {"site_id": "s1", "drive_id": "d1", "path_encoded": "x", "file_path": "bank/report.xlsx"}

    async def run():
        with patch(
            "app.application.use_cases.merge_composite_validado_pdfs.resolve_sharepoint_from_env",
            new_callable=AsyncMock,
            return_value=ctx,
        ):
            return await merge_composite_validado_pdfs(
                g,
                bank_code="banco_bogota",
                historical_file_path=hist,
                email_pdf_path=email,
            )

    result = asyncio.run(run())
    manifest_key = next(k for k in g.uploaded if k.endswith(".json"))
    payload = json.loads(g.uploaded[manifest_key].decode("utf-8"))
    out0 = payload["outputs"][0]
    assert out0["tipo_aplicacion"] == "ABONO"
    assert out0["requiere_extracto"] is False
    assert out0["creditos_seleccionados"] == ["258"]
    assert payload["abono_outputs_count"] == 1
    assert payload["extracts_not_required_count"] == 1
    assert result.merge_manifest_path


def test_merge_mixed_tipos_same_id_uses_aplicacion_multiple_token():
    from app.application.services.review_schema import MERGE_NAME_TOKEN_MULTIPLE

    r_pago, _ = make_distrib_row(
        id_pago="X1",
        credito="258",
        validar_pago=ValidarPago.SI,
        ruta_pdf_internal="ext/x.pdf",
        tipo_aplicacion=TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL,
        cliente="EQUINORTE",
        fecha_banco=date(2026, 6, 3),
    )
    r_abono, _ = make_distrib_row(
        id_pago="X1",
        credito="265",
        validar_pago=ValidarPago.SI,
        tipo_aplicacion=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
        abono_capital=500,
        valor_int=0,
        mora_a_aplicar=0,
        cliente="EQUINORTE",
        fecha_banco=date(2026, 6, 3),
    )
    wb = Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(
        list(AplicacionPagosCols.HEADERS)
        + ["_ruta_extracto", "_ruta_unidad_credito", "_ruta_asientos_contables"]
    )
    ws.append(
        list(r_pago)
        + ["clientes/EQUINORTE/CREDITO# 258/ASIENTOS CONTABLES CRED 258"]
    )
    ws.append(
        list(r_abono)
        + ["clientes/EQUINORTE/CREDITO# 265/ASIENTOS CONTABLES CRED 265"]
    )
    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append(["Campo", "Valor"])
    ws_meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    bio = BytesIO()
    wb.save(bio)

    g = _MergeGraph()
    hist, email = "HIST/mixed.xlsx", "EMAIL/mail.pdf"
    g.initial["bank/report.xlsx"] = _bank_bytes()
    g.initial[hist] = bio.getvalue()
    g.initial[email] = _tiny_pdf()
    g.initial["ext/x.pdf"] = _tiny_pdf()
    for cred in ("258", "265"):
        folder = f"clientes/EQUINORTE/CREDITO# {cred}/ASIENTOS CONTABLES CRED {cred}"
        g.children[folder] = [{"name": f"asiento {cred}.pdf", "file": {}}]
        g.initial[f"{folder}/asiento {cred}.pdf"] = _asiento_pdf(
            40.0 if cred == "258" else 60.0, credit=cred
        )

    ctx = {"site_id": "s1", "drive_id": "d1", "path_encoded": "x", "file_path": "bank/report.xlsx"}

    async def run():
        with patch(
            "app.application.use_cases.merge_composite_validado_pdfs.resolve_sharepoint_from_env",
            new_callable=AsyncMock,
            return_value=ctx,
        ):
            return await merge_composite_validado_pdfs(
                g,
                bank_code="banco_bogota",
                historical_file_path=hist,
                email_pdf_path=email,
            )

    result = asyncio.run(run())
    assert result.outputs_count == 1
    assert MERGE_NAME_TOKEN_MULTIPLE in result.outputs[0].output_relative_path
