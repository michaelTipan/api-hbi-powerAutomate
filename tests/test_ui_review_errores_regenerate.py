"""Hoja Errores → issues específicos + available_actions.regenerate."""
from __future__ import annotations

from openpyxl import Workbook

from app.application.services.review_schema import ErroresCols, ReviewSheets
from app.application.ui.process_projection import (
    PaymentProcessProjectionService,
    ProjectionSources,
)
from app.application.ui.review_errores_read import (
    ReviewErrorRow,
    build_operational_issues_from_review_errores,
    parse_review_errores_workbook,
)
from tests.ui_fixtures import make_snap


def _errores_workbook(*, with_row: bool = True) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = ReviewSheets.ERRORES
    for col, header in enumerate(ErroresCols.HEADERS, start=1):
        ws.cell(1, col, header)
    if with_row:
        ws.cell(2, 1, "PAY-9")
        ws.cell(2, 2, "CLIENTE DEMO")
        ws.cell(2, 3, "215")
        ws.cell(2, 4, "Extracto ilegible")
        ws.cell(2, 5, "La fecha del extracto no se pudo leer.")
        ws.cell(2, 6, "Corrija o retire el PDF en SharePoint y regenere.")
        ws.cell(2, 7, "Sí")
        ext = ws.cell(2, 8, "extracto_malo.pdf")
        ext.hyperlink = "https://sharepoint.example/extracto_malo.pdf"
        fold = ws.cell(2, 9, "CREDITO 215")
        fold.hyperlink = "https://sharepoint.example/credito-215"
        ws.cell(2, 10, "fecha_limite_extracto_not_readable")
    return wb


def test_parse_review_errores_workbook_incluye_contexto_de_archivo() -> None:
    rows = parse_review_errores_workbook(_errores_workbook())
    assert len(rows) == 1
    row = rows[0]
    assert row.id_pago == "PAY-9"
    assert row.credito == "215"
    assert row.cliente == "CLIENTE DEMO"
    assert row.extract_label == "extracto_malo.pdf"
    assert row.folder_label == "CREDITO 215"
    assert row.extract_url and "extracto_malo" in row.extract_url
    assert row.folder_url and "credito-215" in row.folder_url
    assert row.codigo_tecnico == "fecha_limite_extracto_not_readable"


def test_parse_review_errores_workbook_vacia() -> None:
    assert parse_review_errores_workbook(_errores_workbook(with_row=False)) == []


def test_build_operational_issues_expone_links_y_ubicacion() -> None:
    row = ReviewErrorRow(
        row_number=2,
        id_pago="PAY-9",
        cliente="CLIENTE DEMO",
        credito="215",
        tipo_caso="Extracto ilegible",
        descripcion="La fecha del extracto no se pudo leer.",
        que_debe_hacer="Corrija el PDF.",
        requiere_soporte="Sí",
        codigo_tecnico="fecha_limite_extracto_not_readable",
        extract_label="extracto_malo.pdf",
        folder_label="CREDITO 215",
        extract_url="https://sharepoint.example/extracto_malo.pdf",
        folder_url="https://sharepoint.example/credito-215",
    )
    issues = build_operational_issues_from_review_errores(
        [row], file_name="rev.xlsx"
    )
    assert len(issues) == 1
    issue = issues[0]
    assert issue.retry is not None
    assert issue.retry.action == "regenerate"
    assert issue.location is not None
    assert issue.location.sheet == "Errores"
    assert issue.location.credit == "215"
    assert issue.location.payment_id == "PAY-9"
    assert "extracto_malo.pdf" in (issue.user_message or "")
    rels = {lnk.rel for lnk in issue.links}
    # review_excel queda en fase 1 / banner; no se repite por issue de Errores.
    assert "review_excel" not in rels
    assert "error_extract" in rels
    assert "error_folder" in rels


def test_projection_regenerate_allowed_with_open_errores(
    monkeypatch,
) -> None:
    monkeypatch.setenv("UI_ENABLED", "true")
    monkeypatch.setenv("UI_WRITE_ENABLED", "true")
    monkeypatch.setenv("UI_FINALIZE_ENABLED", "true")
    monkeypatch.setenv("UI_AUTH_MODE", "mock")
    monkeypatch.setenv("ACTIVE_ENVIRONMENT", "sandbox")
    monkeypatch.delenv("WEBSITE_INSTANCE_ID", raising=False)
    monkeypatch.delenv("WEBSITE_SITE_NAME", raising=False)

    snap = make_snap(
        estado_proceso="REVISION_CREADA",
        validation_file_path="revision/validacion_pagos_demo.xlsx",
    )
    row = ReviewErrorRow(
        row_number=2,
        id_pago="PAY-9",
        cliente="CLIENTE DEMO",
        credito="215",
        tipo_caso="Extracto ilegible",
        descripcion="PDF ilegible.",
        que_debe_hacer="Corrija y regenere.",
        requiere_soporte="Sí",
        codigo_tecnico="fecha_limite_extracto_not_readable",
        extract_label="extracto_malo.pdf",
        folder_label="CREDITO 215",
        extract_url="https://sharepoint.example/a.pdf",
        folder_url="https://sharepoint.example/folder",
    )
    detail = PaymentProcessProjectionService().project(
        ProjectionSources(snapshot=snap, review_errores_rows=(row,))
    )
    assert detail.operational_status == "CORRECCION_REQUERIDA"
    assert "Errores" in (detail.operational_message or "")
    assert detail.available_actions["regenerate"].allowed is True
    assert detail.available_actions["finalize"].allowed is False
    assert "Errores" in (detail.available_actions["finalize"].reason or "")
    assert any(i.retry and i.retry.action == "regenerate" for i in detail.operational_issues)
    assert any(
        "extracto_malo.pdf" in (i.user_message or "") for i in detail.operational_issues
    )


def test_projection_regenerate_blocked_when_writes_off(monkeypatch) -> None:
    monkeypatch.setenv("UI_ENABLED", "true")
    monkeypatch.setenv("UI_WRITE_ENABLED", "false")
    monkeypatch.setenv("ACTIVE_ENVIRONMENT", "sandbox")

    snap = make_snap(estado_proceso="REVISION_CREADA")
    row = ReviewErrorRow(
        row_number=2,
        id_pago="1",
        cliente="X",
        credito="1",
        tipo_caso="Caso",
        descripcion="Desc",
        que_debe_hacer="Hacer",
        requiere_soporte="",
        codigo_tecnico="code",
        extract_label="",
        folder_label="",
        extract_url=None,
        folder_url=None,
    )
    detail = PaymentProcessProjectionService().project(
        ProjectionSources(snapshot=snap, review_errores_rows=(row,))
    )
    assert detail.available_actions["regenerate"].allowed is False
    assert detail.available_actions["regenerate"].reason


def test_projection_regenerate_allowed_without_errores_pre_finalize(monkeypatch) -> None:
    """Releer banco / rehacer revisión sin Errores ni archivo faltante."""
    monkeypatch.setenv("UI_ENABLED", "true")
    monkeypatch.setenv("UI_WRITE_ENABLED", "true")
    monkeypatch.setenv("UI_FINALIZE_ENABLED", "true")
    monkeypatch.setenv("UI_AUTH_MODE", "mock")
    monkeypatch.setenv("ACTIVE_ENVIRONMENT", "sandbox")
    monkeypatch.delenv("WEBSITE_INSTANCE_ID", raising=False)
    monkeypatch.delenv("WEBSITE_SITE_NAME", raising=False)

    path = "revision/validacion_pagos_demo.xlsx"
    snap = make_snap(
        estado_proceso="REVISION_CREADA",
        validation_file_path=path,
    )
    detail = PaymentProcessProjectionService().project(
        ProjectionSources(snapshot=snap, artifact_exists={path: True})
    )
    assert detail.operational_status != "CORRECCION_REQUERIDA"
    assert detail.available_actions["regenerate"].allowed is True
    assert detail.available_actions["finalize"].allowed is True


def test_build_operational_issues_fills_guide_when_excel_descripcion_empty() -> None:
    row = ReviewErrorRow(
        row_number=4,
        id_pago="PAY-A",
        cliente="EQUINORTE",
        credito="CREDITO # 258",
        tipo_caso="",
        descripcion="",
        que_debe_hacer="",
        requiere_soporte="",
        codigo_tecnico="extract_as_of_not_found",
        extract_label="",
        folder_label="Ver carpeta crédito 258",
        extract_url=None,
        folder_url="https://sharepoint.example/credito-258",
    )
    issues = build_operational_issues_from_review_errores(
        [row], file_name="rev.xlsx"
    )
    assert len(issues) == 1
    issue = issues[0]
    assert "caso pendiente en la hoja Errores" not in (issue.user_message or "").lower()
    assert "extracto" in (issue.user_message or "").lower()
    assert "EQUINORTE" in (issue.user_message or "")
    assert "generar" in (issue.next_action or "").lower()
    assert any(lnk.rel == "error_folder" for lnk in issue.links)
    assert issue.title.lower().startswith("extracto")


def test_projection_regenerate_blocked_after_finalize(monkeypatch) -> None:
    monkeypatch.setenv("UI_ENABLED", "true")
    monkeypatch.setenv("UI_WRITE_ENABLED", "true")
    monkeypatch.setenv("ACTIVE_ENVIRONMENT", "sandbox")
    monkeypatch.delenv("WEBSITE_INSTANCE_ID", raising=False)
    monkeypatch.delenv("WEBSITE_SITE_NAME", raising=False)

    snap = make_snap(
        estado_proceso="FINALIZADO",
        validation_file_path="revision/x.xlsx",
        historical_file_path="historico/h.xlsx",
    )
    detail = PaymentProcessProjectionService().project(ProjectionSources(snapshot=snap))
    assert detail.available_actions["regenerate"].allowed is False
    assert "antes de finalizar" in (detail.available_actions["regenerate"].reason or "").lower()


def test_projection_regenerate_when_review_file_missing(monkeypatch) -> None:
    monkeypatch.setenv("UI_ENABLED", "true")
    monkeypatch.setenv("UI_WRITE_ENABLED", "true")
    monkeypatch.setenv("UI_AUTH_MODE", "mock")
    monkeypatch.setenv("ACTIVE_ENVIRONMENT", "sandbox")
    monkeypatch.delenv("WEBSITE_INSTANCE_ID", raising=False)
    monkeypatch.delenv("WEBSITE_SITE_NAME", raising=False)

    path = "revision/validacion_pagos_demo.xlsx"
    snap = make_snap(
        estado_proceso="REVISION_CREADA",
        validation_file_path=path,
    )
    detail = PaymentProcessProjectionService().project(
        ProjectionSources(snapshot=snap, artifact_exists={path: False})
    )
    assert detail.operational_status == "CORRECCION_REQUERIDA"
    assert detail.available_actions["regenerate"].allowed is True
    assert any(i.issue_id == "review-file-missing" for i in detail.operational_issues)
    assert "SharePoint" in (detail.operational_message or "")
    gen = next(s for s in detail.steps if s.name == "generate")
    assert gen.status == "failed_business"
    assert gen.retry_action == "retry_generate"
