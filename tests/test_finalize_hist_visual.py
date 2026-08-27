"""Visual del histórico (cartera validada): formato Monto/Fecha y bordes."""
from __future__ import annotations

from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Side

from app.application.services.review_schema import AplicacionPagosCols
from app.application.use_cases.payment_validation_finalize import (
    _explode_vertical_merges_in_named_columns,
    _normalize_hist_aplicacion_row_borders,
)


def test_explode_merges_copies_number_format_to_child_rows() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = AplicacionPagosCols.MONTO_BANCO
    ws["B1"] = AplicacionPagosCols.FECHA_BANCO
    ws["A2"] = 13687114.0
    ws["A2"].number_format = "#,##0.00"
    ws["A2"].alignment = Alignment(vertical="center", horizontal="center")
    ws["B2"] = date(2026, 8, 25)
    ws["B2"].number_format = "yyyy-mm-dd"
    ws.merge_cells("A2:A4")
    ws.merge_cells("B2:B4")

    _explode_vertical_merges_in_named_columns(
        ws,
        1,
        (AplicacionPagosCols.MONTO_BANCO, AplicacionPagosCols.FECHA_BANCO),
    )

    assert ws.merged_cells.ranges == set()
    for r in (2, 3, 4):
        assert ws.cell(r, 1).value == 13687114.0
        assert ws.cell(r, 1).number_format == "#,##0.00"
        assert ws.cell(r, 1).alignment.horizontal == "left"
        assert ws.cell(r, 2).value == date(2026, 8, 25)
        assert ws.cell(r, 2).number_format == "yyyy-mm-dd"
        assert ws.cell(r, 2).alignment.horizontal == "left"


def test_normalize_hist_borders_removes_thick_navy_edges() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ID Pago"
    thick = Side(style="thick", color="002060")
    thin = Side(style="thin", color="C8C8C8")
    ws["A2"] = "pay-1"
    ws["A2"].border = Border(left=thin, right=thin, top=thick, bottom=thin)
    ws["A3"] = "pay-1"
    ws["A3"].border = Border(left=thin, right=thin, top=thin, bottom=thick)

    _normalize_hist_aplicacion_row_borders(ws, 1)

    for r in (2, 3):
        b = ws.cell(r, 1).border
        assert b.top.style == "thin"
        assert b.bottom.style == "thin"
        assert b.left.style == "thin"
        assert b.right.style == "thin"
