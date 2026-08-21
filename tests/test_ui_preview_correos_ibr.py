"""Tests de vista previa CORREOS / IBR para confirmación UI."""

from __future__ import annotations

import asyncio
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.application.services.ibr_workbook import find_ibr_for_date
from app.application.ui.ibr_preview import (
    _list_ibr_ranges,
    build_ibr_operator_message,
    format_operator_date,
    load_ibr_preview,
    resolve_ibr_rate_rows,
    resolve_ibr_rates_for_dates,
)
from app.application.ui.notify_recipients_preview import (
    _parse_correos_with_sheet,
    load_notify_recipients_preview,
)
from app.application.use_cases.setup_ibr_workbook import IBR_COLUMNS, IBR_SHEET_NAME


def _correos_bytes(*, emisor: str, receptores: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "CORREOS"
    ws["A1"] = "EMISOR"
    ws["B1"] = "RECEPTORES"
    ws["A2"] = emisor
    for i, r in enumerate(receptores, start=2):
        ws.cell(row=i, column=2, value=r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ibr_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = IBR_SHEET_NAME
    for i, h in enumerate(IBR_COLUMNS, start=1):
        ws.cell(1, i, h)
    ws.cell(2, 1, date(2026, 1, 1))
    ws.cell(2, 2, date(2026, 7, 31))
    ws.cell(2, 3, 0.09)
    ws.cell(3, 1, date(2026, 8, 1))
    ws.cell(3, 2, date(2026, 12, 31))
    ws.cell(3, 3, 0.1058)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_correos_excludes_emisor_from_effective_list_logic():
    raw = _correos_bytes(
        emisor="from@hbi.test",
        receptores=["from@hbi.test", "a@hbi.test", "b@hbi.test"],
    )
    wb = load_workbook(BytesIO(raw), data_only=True)
    sender, recipients, sheet = _parse_correos_with_sheet(wb)
    assert sender == "from@hbi.test"
    assert sheet == "CORREOS"
    assert "a@hbi.test" in recipients
    effective = [e for e in recipients if e.lower() != sender.lower()]
    assert effective == ["a@hbi.test", "b@hbi.test"]


def test_list_ibr_ranges_and_find_for_process_date():
    raw = _ibr_bytes()
    ranges = _list_ibr_ranges(raw)
    assert len(ranges) == 2
    assert find_ibr_for_date(raw, date(2026, 8, 11)) == pytest.approx(0.1058)
    assert find_ibr_for_date(raw, date(2026, 7, 15)) == pytest.approx(0.09)


def test_resolve_ibr_rate_rows_skips_adelantado_and_keeps_on_time():
    raw = _ibr_bytes()
    rows = resolve_ibr_rate_rows(
        raw,
        [
            {
                "cut": date(2026, 7, 15),
                "credito": "99",
                "updates_ibr": False,
                "skip_reason": "adelantado",
            },
            {
                "cut": date(2026, 8, 4),
                "credito": "215",
                "updates_ibr": True,
                "skip_reason": None,
            },
        ],
    )
    assert rows[0]["status"] == "skipped_adelantado"
    assert rows[0]["updates_ibr"] is False
    assert rows[0]["credito"] == "99"
    assert rows[0]["credito_label"] == "Crédito 99"
    assert rows[1]["status"] == "found"
    assert rows[1]["updates_ibr"] is True
    assert rows[1]["credito_label"] == "Crédito 215"
    msg = build_ibr_operator_message(rows)
    assert "pago antes del corte" in msg
    assert "10,58 %" in msg
    assert "Crédito 99" in msg
    assert "Crédito 215" in msg


def test_resolve_ibr_rate_rows_keeps_same_cut_for_distinct_credits():
    raw = _ibr_bytes()
    rows = resolve_ibr_rate_rows(
        raw,
        [
            {"cut": date(2026, 8, 4), "credito": "128", "updates_ibr": True},
            {"cut": date(2026, 8, 4), "credito": "215", "updates_ibr": True},
        ],
    )
    assert len(rows) == 2
    assert [r["credito"] for r in rows] == ["128", "215"]
    assert all(r["status"] == "found" for r in rows)


def test_resolve_ibr_rates_for_multiple_cut_dates():
    raw = _ibr_bytes()
    rates = resolve_ibr_rates_for_dates(raw, [date(2026, 8, 4), date(2026, 7, 15)])
    assert [row["date"] for row in rates] == ["2026-07-15", "2026-08-04"]
    assert rates[0]["rate_label"] == "9 %"
    assert rates[1]["rate_label"] == "10,58 %"
    assert rates[0]["updates_ibr"] is True
    msg = build_ibr_operator_message(rates)
    assert "distintos cortes" in msg
    assert "15 jul 2026" in msg
    assert "4 ago 2026" in msg
    assert "ProcessKey" not in msg
    assert "rango" not in msg.lower()


def test_format_operator_date_is_short_spanish():
    assert format_operator_date(date(2026, 8, 4)) == "4 ago 2026"


def test_load_notify_recipients_preview(monkeypatch):
    raw = _correos_bytes(emisor="ops@hbi.test", receptores=["ops@hbi.test", "c@hbi.test"])

    class G:
        pass

    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "ops")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "Documents")
    monkeypatch.setenv("GRAPH_VALIDAR_NOTIFY_CORREOS_XLSX_PATH", "CTL/CORREOS.xlsx")

    async def _resolve(*_a, **_k):
        return {
            "site_id": "s",
            "drive_id": "d",
            "path_encoded": "x",
            "file_path": "CTL/CORREOS.xlsx",
        }

    async def _dl(*_a, **_k):
        return raw

    async def _meta(*_a, **_k):
        return {"lastModifiedDateTime": "2026-08-11T15:00:00Z"}

    monkeypatch.setattr(
        "app.application.ui.notify_recipients_preview.require_operations_site_config",
        lambda: None,
    )
    monkeypatch.setattr(
        "app.application.ui.notify_recipients_preview.resolve_sharepoint_path",
        _resolve,
    )
    monkeypatch.setattr(
        "app.application.ui.notify_recipients_preview._graph_download_by_path",
        _dl,
    )
    monkeypatch.setattr(
        "app.application.ui.notify_recipients_preview._graph_get_item_metadata_by_path",
        _meta,
    )

    out = asyncio.run(load_notify_recipients_preview(G()))  # type: ignore[arg-type]
    assert out["emisor"] == "ops@hbi.test"
    assert out["receptores"] == ["c@hbi.test"]
    assert out["file_last_modified"] == "11 ago 2026, 10:00 a. m."
    assert any("Excel Online" in w for w in out["warnings"])
    assert "acaba de editar" in out["warnings"][0]


def test_load_ibr_preview(monkeypatch):
    raw = _ibr_bytes()

    class G:
        pass

    monkeypatch.setattr(
        "app.application.ui.ibr_preview.require_operations_site_config",
        lambda: None,
    )

    async def _resolve(*_a, **_k):
        return {
            "site_id": "s",
            "drive_id": "d",
            "path_encoded": "x",
            "file_path": "CTL/IBR.xlsx",
        }

    async def _dl(*_a, **_k):
        return raw

    async def _meta(*_a, **_k):
        return {"lastModifiedDateTime": "2026-08-11T16:00:00Z"}

    async def _no_hist(*_a, **_k):
        raise RuntimeError("no historical in this test")

    monkeypatch.setattr("app.application.ui.ibr_preview.resolve_sharepoint_path", _resolve)
    monkeypatch.setattr("app.application.ui.ibr_preview._graph_download_by_path", _dl)
    monkeypatch.setattr(
        "app.application.ui.ibr_preview._graph_get_item_metadata_by_path",
        _meta,
    )
    monkeypatch.setattr(
        "app.application.use_cases.payment_validation_process_control.read_process_control_snapshot",
        _no_hist,
    )
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "ops")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "Documents")

    out = asyncio.run(
        load_ibr_preview(
            G(),  # type: ignore[arg-type]
            process_key=(
                "payment-validation|banco_bogota|2026-08-11|"
                "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"
            ),
        )
    )
    assert out["rate_status"] == "found"
    assert out["rate_pct"] == pytest.approx(10.58)
    assert out["process_date"] == "2026-08-11"
    assert out["rates"][0]["date"] == "2026-08-11"
    assert out["file_last_modified"] == "11 ago 2026, 11:00 a. m."
    assert "corte" in out["user_message"].lower()
