"""CAPA B v4: un crédito SI y extras NO; sin columnas de distribución."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.application.services.review_schema import AplicacionPagosCols, ReviewSheets
from scripts.e2e_rc.review_edit import approve_single_credit_pago


def _two_credit_review() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = ReviewSheets.APLICACION_PAGOS
    headers = list(AplicacionPagosCols.HEADERS)
    for col, h in enumerate(headers, start=1):
        ws.cell(1, col).value = h
    idx = {h: i + 1 for i, h in enumerate(headers)}
    ws.cell(2, idx[AplicacionPagosCols.ID_PAGO]).value = "P1"
    ws.cell(2, idx[AplicacionPagosCols.CLIENTE]).value = "GEOEXCON"
    ws.cell(2, idx[AplicacionPagosCols.CREDITO]).value = "CREDITO # 231"
    ws.cell(2, idx[AplicacionPagosCols.MONTO_BANCO]).value = 19_102_163
    ws.cell(2, idx[AplicacionPagosCols.VALIDAR_PAGO]).value = "POR DEFINIR"
    ws.cell(3, idx[AplicacionPagosCols.ID_PAGO]).value = "P1"
    ws.cell(3, idx[AplicacionPagosCols.CLIENTE]).value = "GEOEXCON"
    ws.cell(3, idx[AplicacionPagosCols.CREDITO]).value = "CREDITO # 299"
    ws.cell(3, idx[AplicacionPagosCols.VALIDAR_PAGO]).value = "POR DEFINIR"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_approve_single_marks_extra_credit_no() -> None:
    raw = _two_credit_review()
    edited, applied = approve_single_credit_pago(
        raw, tipo="PAGO DE OBLIGACIÓN ACTUAL", credito_contains="231"
    )
    assert len(applied) == 2
    wb = load_workbook(BytesIO(edited))
    ws = wb[ReviewSheets.APLICACION_PAGOS]
    headers = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    assert ws.cell(2, idx[AplicacionPagosCols.VALIDAR_PAGO]).value == "SI"
    assert ws.cell(3, idx[AplicacionPagosCols.VALIDAR_PAGO]).value == "NO"
    assert ws.cell(3, idx[AplicacionPagosCols.TIPO_APLICACION]).value in ("", None)
    for banned in (
        "Aplicar a obligación actual",
        "Total asignado al crédito",
        "Saldo por asignar",
    ):
        assert banned not in headers
