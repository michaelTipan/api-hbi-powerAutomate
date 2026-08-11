"""Tests lectura compartida de Aplicacion_Pagos v3 en histórico."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.application.services.historical_application_rows import (
    detect_application_type_group_conflict,
    group_abono_rows_for_email,
    group_rows_by_id_pago,
    read_validated_abono_rows,
    read_validated_application_rows,
    read_validated_payment_rows,
)
from app.application.services.review_schema import (
    AplicacionPagosCols,
    REVIEW_SCHEMA_VERSION,
    ReviewSheets,
    TipoAplicacion,
    TipoAplicacionConfirmado,
    ValidarPago,
)
from tests.test_finalize_validation import make_distrib_row


def _v3_historical_bytes(*, rows: list | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(list(AplicacionPagosCols.HEADERS) + ["_ruta_extracto", "_ruta_unidad_credito", "_ruta_asientos_contables"])
    for row in rows or []:
        ws.append(row)
    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append(["Campo", "Valor"])
    ws_meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_read_application_rows_pago_and_abono_unified():
    r_pago, _ = make_distrib_row(
        id_pago="P1",
        validar_pago=ValidarPago.SI,
        ruta_pdf_internal="ext/p1.pdf",
        tipo_aplicacion=TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL,
    )
    r_abono, _ = make_distrib_row(
        id_pago="AB1",
        credito="258",
        validar_pago=ValidarPago.SI,
        ruta_pdf_internal="",
        tipo_aplicacion=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
        abono_capital=1000,
        valor_int=0,
        mora_a_aplicar=0,
    )
    # append asientos path as 3rd technical col
    r_pago = list(r_pago) + ["clientes/CLI/CREDITO# 1/ASIENTOS"]
    r_abono = list(r_abono) + ["clientes/CLI/CREDITO# 258/ASIENTOS"]
    # make_distrib_row already appends ruta + unidad → need careful length
    # Recreate cleanly:
    r_pago, _ = make_distrib_row(
        id_pago="P1",
        validar_pago=ValidarPago.SI,
        ruta_pdf_internal="ext/p1.pdf",
        tipo_aplicacion=TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL,
    )
    r_abono, _ = make_distrib_row(
        id_pago="AB1",
        credito="258",
        validar_pago=ValidarPago.SI,
        ruta_pdf_internal="",
        tipo_aplicacion=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
        abono_capital=1000,
        valor_int=0,
        mora_a_aplicar=0,
    )
    # make_distrib_row returns HEADERS + ruta + unidad; add asientos
    r_pago = list(r_pago) + ["clientes/CLI/P1/ASIENTOS"]
    r_abono = list(r_abono) + ["clientes/CLI/CREDITO# 258/ASIENTOS"]

    wb = load_workbook(BytesIO(_v3_historical_bytes(rows=[r_pago, r_abono])), data_only=True)
    all_rows = read_validated_application_rows(wb)
    assert len(all_rows) == 2
    pagos = read_validated_payment_rows(wb)
    abonos = read_validated_abono_rows(wb)
    assert len(pagos) == 1
    assert pagos[0]["tipo_aplicacion"] == TipoAplicacion.PAGO.value
    assert pagos[0]["include_extract_in_composite"] is True
    assert len(abonos) == 1
    assert abonos[0]["tipo_aplicacion"] == TipoAplicacion.ABONO.value
    assert abonos[0]["include_extract_in_composite"] is False


def test_group_abono_single_row_per_id_pago():
    rows = []
    for cred in ("258", "265"):
        r, _ = make_distrib_row(
            id_pago="AB1",
            credito=cred,
            validar_pago=ValidarPago.SI,
            tipo_aplicacion=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
            abono_capital=500,
            valor_int=0,
            mora_a_aplicar=0,
            monto_banco=2_000_000,
            fecha_banco=date(2026, 6, 3),
            cliente="EQUINORTE",
        )
        rows.append(list(r) + [f"clientes/EQUINORTE/CREDITO# {cred}/ASIENTOS"])
    wb = load_workbook(BytesIO(_v3_historical_bytes(rows=rows)), data_only=True)
    abono_rows = read_validated_abono_rows(wb)
    assert len(abono_rows) == 2
    groups = group_abono_rows_for_email(abono_rows)
    assert len(groups) == 1
    assert groups[0].id_pago == "AB1"
    assert groups[0].creditos_seleccionados == ("258", "265")


def test_unified_groups_allow_mixed_tipos_same_id():
    """v3: mismo ID con tipos distintos → un grupo (APLICACION MULTIPLE en Merge)."""
    r1, _ = make_distrib_row(
        id_pago="X1",
        validar_pago=ValidarPago.SI,
        tipo_aplicacion=TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL,
    )
    r2, _ = make_distrib_row(
        id_pago="X1",
        credito="265",
        validar_pago=ValidarPago.SI,
        tipo_aplicacion=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
        abono_capital=100,
        valor_int=0,
        mora_a_aplicar=0,
    )
    wb = load_workbook(
        BytesIO(_v3_historical_bytes(rows=[list(r1) + ["a"], list(r2) + ["b"]])),
        data_only=True,
    )
    groups = group_rows_by_id_pago(read_validated_application_rows(wb))
    assert set(groups.keys()) == {"X1"}
    assert len(groups["X1"]) == 2
    # compat no-op
    detect_application_type_group_conflict({"X1": groups["X1"]}, {})
