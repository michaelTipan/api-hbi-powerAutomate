"""Tests apply real de amortización (mock SharePoint con put_bytes)."""

import asyncio
import hashlib
import io
import json
from datetime import date
from urllib.parse import unquote

import httpx
import openpyxl
import pytest

from app.application.services.accounting_pdf_parser import (
    ACCOUNT_CAPITAL,
    ACCOUNT_SALDOS_MENORES,
)
from app.application.services.amortization_workbook import (
    AUTOMATION_LOG_SHEET,
    load_automation_log_idempotency_keys,
    _is_formula_value,
)
from app.application.use_cases.amortization_fill_apply import (
    AmortizationPreflightError,
    run_amortization_fill_apply,
    validate_amortization_preflight,
)
from app.application.use_cases.amortization_fill_dry_run import run_amortization_fill_dry_run
from tests.test_amortization_fill_dry_run import (
    MockGraphDryRun,
    _accounting_text,
    _amort_table_date_at_row,
    _amort_table_displaced_application,
    _amort_table_two_dates_at_rows,
    _amort_table_with_op_formulas,
    _asiento_pdf_placeholder,
    _base_files,
    _hist_bytes,
    _hist_bytes_multi_credit,
    _ibr_bytes,
)


class MockGraphApply(MockGraphDryRun):
    """Mock Graph con almacén mutable para simular round-trip de tablas."""

    def __init__(self, files: dict[str, bytes], *, fail_upload_423: bool = False) -> None:
        super().__init__(files)
        self.uploaded: dict[str, bytes] = {}
        self.fail_upload_423 = fail_upload_423
        self.patch_calls: list[tuple[str, dict]] = []
        self.post_json_calls: list[tuple[str, dict]] = []

    def _item_path(self, endpoint: str) -> str | None:
        if "/root:/" not in endpoint or ":/content" in endpoint or "children" in endpoint:
            return None
        return unquote(endpoint.split("/root:/", 1)[1].rstrip(":"))

    async def get(self, endpoint: str, params=None):
        path = self._item_path(endpoint)
        if path is not None:
            if path not in self.files:
                request = httpx.Request("GET", "https://graph.test/item")
                response = httpx.Response(404, request=request)
                raise httpx.HTTPStatusError("404", request=request, response=response)
            blob = self.files.get(path, b"")
            digest = hashlib.sha256(blob).hexdigest()[:16]
            return {
                "eTag": f'"{digest}"',
                "size": len(blob),
                "lastModifiedDateTime": "2026-04-23T12:00:00Z",
                "name": path.rsplit("/", 1)[-1],
                "webUrl": f"https://sharepoint.test/{path}",
            }
        return await super().get(endpoint, params)

    async def post_json(self, endpoint: str, body: dict):
        self.post_json_calls.append((endpoint, body))
        if ":/children" in endpoint:
            parent = unquote(endpoint.split("/root:/", 1)[1].rsplit(":/children", 1)[0])
            name = str(body.get("name") or "").strip()
            if name:
                folder_key = f"{parent}/{name}"
                self.files.setdefault(folder_key, b"")
        return {}, 201

    async def patch_json(self, endpoint: str, body: dict):
        self.patch_calls.append((endpoint, body))
        source = self._item_path(endpoint)
        if not source or source not in self.files:
            request = httpx.Request("PATCH", "https://graph.test/move")
            response = httpx.Response(404, request=request)
            raise httpx.HTTPStatusError("404", request=request, response=response)
        parent_path = str(body.get("parentReference", {}).get("path") or "")
        parent = parent_path.replace("/drive/root:/", "").strip("/")
        name = str(body.get("name") or "").strip()
        dest = f"{parent}/{name}"
        self.files[dest] = self.files.pop(source)
        return {"id": dest}

    async def get_bytes(self, endpoint: str, params=None):
        key = self._key(endpoint)
        if key and key in self.uploaded:
            return self.uploaded[key]
        return await super().get_bytes(endpoint, params)

    async def put_bytes(self, endpoint: str, content: bytes, content_type: str = ""):
        if self.fail_upload_423:
            request = httpx.Request("PUT", "https://graph.test/upload")
            response = httpx.Response(423, request=request, text="Locked")
            raise httpx.HTTPStatusError("423 Locked", request=request, response=response)
        key = self._key(endpoint)
        if key:
            self.uploaded[key] = content
            self.files[key] = content
        self.put_calls.append((endpoint, content))
        return {"id": "uploaded"}


def _saldos_menores_accounting_text() -> str:
    return f"""
    Comprobante 1 Fecha 22/04/2026
    {ACCOUNT_SALDOS_MENORES} 100.00
    {ACCOUNT_CAPITAL} 100.00
    """


@pytest.fixture(autouse=True)
def env_sharepoint(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")


def test_preflight_rejects_errors_when_not_partial_ready():
    dry = {"summary": {"errors": 1, "revision_manual": 0}, "can_apply": False}
    with pytest.raises(AmortizationPreflightError) as excinfo:
        validate_amortization_preflight(dry)
    assert excinfo.value.error_code == "preflight_errors"


def test_preflight_allows_partial_apply_when_can_apply_true():
    dry = {
        "summary": {"errors": 1, "revision_manual": 0, "would_apply": 1},
        "can_apply": True,
        "items": [],
    }
    validate_amortization_preflight(dry)


def test_apply_single_table_writes_and_logs(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["status"] == "ok"
    assert out["summary"]["applied"] == 1
    assert len(g.put_calls) == 1
    item = out["items"][0]
    assert item["apply_status"] == "APPLIED"
    assert item["due_date_row"] == 8
    assert item["application_row"] == 8
    assert item.get("payment_date_iso") == fecha.isoformat()
    wb = openpyxl.load_workbook(io.BytesIO(g.uploaded["TABLAS/amort.xlsx"]), data_only=False)
    ws = wb["EQUINORTE"]
    from app.application.services.amortization_workbook import detect_headers, find_header_row, _is_formula_value

    headers = detect_headers(ws, header_row=find_header_row(ws))
    fecha_cell = ws.cell(8, headers["fecha_pago"]).value
    if hasattr(fecha_cell, "date"):
        fecha_cell = fecha_cell.date()
    assert fecha_cell == fecha
    assert _is_formula_value(ws.cell(8, headers["valor_pagado_cliente"]).value)
    assert item["write_plan"].get("valor_pagado_cliente") == "formula"
    log_ws = wb[AUTOMATION_LOG_SHEET]
    assert log_ws.max_row >= 2
    assert log_ws.protection.sheet is True
    assert load_automation_log_idempotency_keys(wb)
    assert out.get("automation_log_protected") is True
    assert out["accounting_pdfs_moved_count"] == 1
    move = out["accounting_pdfs_moves"][0]
    assert move["status"] == "moved"
    assert "PROCESADOS" in move["destination_path"]
    assert "ASIENTOS CONTABLES CRED 258" in move["processed_folder_path"]


def test_apply_multiple_asientos_different_application_rows(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 265", "TABLAS/amort_265.xlsx", fecha)
    asiento_a = "clientes/E/CREDITO # 265/ASIENTOS CONTABLES CRED 265/a1.pdf"
    asiento_b = "clientes/E/CREDITO # 265/ASIENTOS CONTABLES CRED 265/a2.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_paths": [asiento_a, asiento_b],
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_265.xlsx": _amort_table_two_dates_at_rows(fecha, 8, 9),
        asiento_a: _asiento_pdf_placeholder(),
        asiento_b: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphApply(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["status"] == "ok"
    assert out["summary"]["applied"] == 2
    rows = sorted(it["application_row"] for it in out["items"])
    assert rows == [8, 9]
    assert all(it["due_date_row"] == 8 for it in out["items"])
    assert all(it["ibr_row"] == 8 for it in out["items"])
    assert out["accounting_pdfs_moved_count"] == 2
    dests = [m["destination_path"] for m in out["accounting_pdfs_moves"]]
    assert all("PROCESADOS" in d for d in dests)
    assert any("_evento-1" in d for d in dests)
    assert any("_evento-2" in d for d in dests)


def test_apply_displaced_application_row(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 265", "TABLAS/amort_265.xlsx", fecha)
    amort = _amort_table_displaced_application(
        fecha, due_row=8, occupied_application_rows=[8, 9], free_application_row=10
    )
    asiento = "clientes/E/asiento.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_path": asiento,
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_265.xlsx": amort,
        asiento: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphApply(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["items"][0]["due_date_row"] == 8
    assert out["items"][0]["application_row"] == 10


def test_apply_retry_already_applied_skips_preflight_and_move(monkeypatch):
    """Segundo apply con control AMORTIZACION_APLICADA: sin dry-run ni movimiento de PDFs."""
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("851eb0f6", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    asiento = "clientes/EQUINORTE/CREDITO # 258/ASIENTOS CONTABLES CRED 258/asiento.pdf"
    process_key = f"payment-validation|banco_bogota|{fecha.isoformat()}"
    manifest_key = f"LOGS/merge_manifest_banco_bogota_{fecha.isoformat()}.json"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "851eb0f6",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 258",
                "asiento_pdf_path": asiento,
            }
        ],
    }
    files = _base_files(
        hist=hist,
        amort=_amort_table_date_at_row(fecha, 8),
        asiento_pdf=_asiento_pdf_placeholder(),
        ibr=_ibr_bytes(),
        fecha=fecha,
    )
    files[manifest_key] = json.dumps(manifest).encode("utf-8")
    g = MockGraphApply(files)
    dry_run_calls: list[int] = []
    move_calls: list[int] = []
    control_phase = {"done": False}

    async def _track_dry_run(graph, **kwargs):
        dry_run_calls.append(1)
        return await run_amortization_fill_dry_run(graph, **kwargs)

    async def _track_move(*_a, **_k):
        move_calls.append(1)
        from app.application.services.accounting_pdf_processed_move import (
            process_used_accounting_pdfs_after_apply,
        )

        return await process_used_accounting_pdfs_after_apply(*_a, **_k)

    from app.application.use_cases.payment_validation_process_control import (
        ProcessControlSnapshot,
    )

    consolidado = ProcessControlSnapshot(
        control_file_path="CTL/bogota.xlsx",
        estado_proceso="CONSOLIDADO",
        is_active=True,
        process_key=process_key,
        process_id="",
        validation_file_path="",
        historical_file_path="HIST/cartera.xlsx",
        secretary_file_path="",
        email_pdf_path="",
        notify_idempotency_key="",
        merge_manifest_path=manifest_key,
        merge_idempotency_key="",
        apply_idempotency_key="",
        bank_code="banco_bogota",
        bank_name="Banco de Bogotá",
    )
    aplicada = ProcessControlSnapshot(
        control_file_path=consolidado.control_file_path,
        estado_proceso="AMORTIZACION_APLICADA",
        is_active=True,
        process_key=process_key,
        process_id="",
        validation_file_path="",
        historical_file_path=consolidado.historical_file_path,
        secretary_file_path="",
        email_pdf_path="",
        notify_idempotency_key="",
        merge_manifest_path=consolidado.merge_manifest_path,
        merge_idempotency_key="",
        apply_idempotency_key=process_key,
        bank_code="banco_bogota",
        bank_name="Banco de Bogotá",
    )

    async def _read_snap(_g, _s, _d, *, bank_code: str):
        if control_phase["done"]:
            return aplicada
        return consolidado

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.read_process_control_snapshot",
        _read_snap,
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.run_amortization_fill_dry_run",
        _track_dry_run,
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.process_used_accounting_pdfs_after_apply",
        _track_move,
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )

    out1 = asyncio.run(
        run_amortization_fill_apply(
            g,
            merge_manifest_path=manifest_key,
            historical_file_path="HIST/cartera.xlsx",
            bank_code="banco_bogota",
        )
    )
    control_phase["done"] = True
    assert out1["already_applied"] is False
    assert out1["summary"]["applied"] == 1
    assert out1["accounting_pdfs_moved_count"] == 1
    assert asiento not in g.files

    out2 = asyncio.run(
        run_amortization_fill_apply(
            g,
            merge_manifest_path=manifest_key,
            historical_file_path="HIST/cartera.xlsx",
            bank_code="banco_bogota",
        )
    )
    assert out2["status"] == "ok"
    assert out2["already_applied"] is True
    assert out2["file_action"] == "reused"
    assert out2["apply_wrote_changes"] is False
    assert out2["process_control_estado"] == "AMORTIZACION_APLICADA"
    assert out2["accounting_pdfs_moved_count"] == 0
    assert out2["accounting_pdfs_processed_count"] == 0
    assert "user_message" in out2
    assert len(dry_run_calls) == 1
    assert len(move_calls) == 1


def test_apply_idempotent_on_second_run(monkeypatch):
    async def _noop_move(*_a, **_k):
        from app.application.services.accounting_pdf_processed_move import (
            empty_accounting_pdf_move_summary,
        )

        return empty_accounting_pdf_move_summary()

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.process_used_accounting_pdfs_after_apply",
        _noop_move,
    )
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out1 = asyncio.run(
        run_amortization_fill_apply(
            g, report_date_iso=fecha.isoformat(), historical_file_path="HIST/cartera.xlsx"
        )
    )
    assert out1["summary"]["applied"] == 1
    out2 = asyncio.run(
        run_amortization_fill_apply(
            g, report_date_iso=fecha.isoformat(), historical_file_path="HIST/cartera.xlsx"
        )
    )
    assert out2["summary"]["skipped_idempotent"] == 1
    assert out2["summary"]["applied"] == 0


def test_apply_allows_bank_inferred_saldos_menores(monkeypatch):
    fecha = date(2026, 4, 23)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _saldos_menores_accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso="2026-04-23",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["status"] == "ok"
    assert out["summary"]["applied"] == 1
    item = out["items"][0]
    assert "BANK_VALUE_INFERRED_OR_MISSING" in (item.get("warnings") or [])


def test_apply_blocks_disallowed_warning_preflight(monkeypatch):
    fecha = date(2026, 4, 23)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )

    def _bad_text(_b):
        return """
        Comprobante 1 Fecha 22/04/2026
        544113410519 100.00
        544113430501 50.00
        544141502030 10.00
        """

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        _bad_text,
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso="2026-04-23",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["status"] == "preflight_failed"
    assert out["preflight_error_code"] == "preflight_warnings_not_allowed"
    assert g.put_calls == []


def test_apply_pdf_changed_same_path_not_idempotent(monkeypatch):
    async def _noop_move(*_a, **_k):
        from app.application.services.accounting_pdf_processed_move import (
            empty_accounting_pdf_move_summary,
        )

        return empty_accounting_pdf_move_summary()

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.process_used_accounting_pdfs_after_apply",
        _noop_move,
    )
    fecha = date(2026, 4, 22)
    asiento = "clientes/EQUINORTE/CREDITO # 258/ASIENTOS CONTABLES CRED 258/asiento.pdf"
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    files = _base_files(
        hist=hist,
        amort=_amort_table_date_at_row(fecha, 8),
        asiento_pdf=b"pdf-v1",
        ibr=_ibr_bytes(),
        fecha=fecha,
    )
    files[asiento] = b"pdf-v1"
    g = MockGraphApply(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out1 = asyncio.run(
        run_amortization_fill_apply(
            g, report_date_iso=fecha.isoformat(), historical_file_path="HIST/cartera.xlsx"
        )
    )
    assert out1["summary"]["applied"] == 1
    g.files[asiento] = b"pdf-v2-changed"
    out2 = asyncio.run(
        run_amortization_fill_apply(
            g, report_date_iso=fecha.isoformat(), historical_file_path="HIST/cartera.xlsx"
        )
    )
    assert out2["items"][0]["apply_status"] == "ERROR"
    assert out2["items"][0]["apply_error_code"] == "PDF_CHANGED_SAME_PATH"


def test_apply_post_upload_verification_ok_includes_tables_summary(monkeypatch):
    fecha = date(2026, 4, 23)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g, report_date_iso="2026-04-23", historical_file_path="HIST/cartera.xlsx"
        )
    )
    assert out["tables_summary"]
    ts = out["tables_summary"][0]
    assert ts["verification_status"] == "ok"
    assert ts["upload_status"] == "uploaded"
    assert ts["eventos_aplicados"] == 1
    assert out["tables_updated_links"] == [
        {
            "label": "TABLAS · amort.xlsx",
            "path": "TABLAS/amort.xlsx",
            "file_url": "https://sharepoint.test/TABLAS/amort.xlsx?web=1",
        }
    ]
    assert "TABLAS · amort.xlsx" in out["tables_updated_links_html"]
    assert "web=1" in out["tables_updated_links_html"]
    assert 'target="_blank"' in out["tables_updated_links_html"]
    assert out["report_date_iso"] == "2026-04-23"
    assert "finalizó correctamente" in out["user_message"]


def test_apply_post_upload_verification_formula_mismatch(monkeypatch):
    fecha = date(2026, 4, 23)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )

    def _fake_write(ws, row, headers, event, **kwargs):
        from app.application.services.amortization_workbook import write_payment_application as real_write

        plan = real_write(ws, row, headers, event, **kwargs)
        vp_col = headers.get("valor_pagado_cliente")
        if vp_col is not None and plan.get("valor_pagado_cliente") == "formula":
            ws.cell(row, vp_col, value=50_000_000.0)
        return plan

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.write_payment_application",
        _fake_write,
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g, report_date_iso="2026-04-23", historical_file_path="HIST/cartera.xlsx"
        )
    )
    assert out["items"][0]["apply_error_code"] == "POST_UPLOAD_VERIFICATION_FAILED_FORMULA_MISMATCH"


def test_apply_post_upload_verification_failed(monkeypatch):
    fecha = date(2026, 4, 23)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.verify_uploaded_table",
        lambda *_a, **_k: ["IBR no coincide"],
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g, report_date_iso="2026-04-23", historical_file_path="HIST/cartera.xlsx"
        )
    )
    assert out["items"][0]["apply_error_code"] == "POST_UPLOAD_VERIFICATION_FAILED"
    assert out["tables_summary"][0]["verification_status"] == "POST_UPLOAD_VERIFICATION_FAILED"


def test_apply_excel_locked_does_not_count_as_applied(monkeypatch):
    fecha = date(2026, 4, 23)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        ),
        fail_upload_423=True,
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g, report_date_iso="2026-04-23", historical_file_path="HIST/cartera.xlsx"
        )
    )
    assert out["summary"]["applied"] == 0
    assert out["summary"]["errors"] == 1
    assert out["items"][0]["apply_error_code"] == "EXCEL_LOCKED"
    assert out["tables_uploaded"] == []


def _op_snapshot(xlsx_bytes: bytes, *, sheet: str = "EQUINORTE", last_row: int = 12):
    """Contenido literal de las columnas O:P, para comparar antes/después de Apply."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    ws = wb[sheet]
    return [(ws.cell(r, 15).value, ws.cell(r, 16).value) for r in range(1, last_row + 1)]


def test_apply_leaves_op_columns_untouched(monkeypatch):
    """O (dia) y P (Causac Inter Mes) las administra contabilidad: Apply no las toca."""
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    amort = _amort_table_with_op_formulas(fecha, due_row=8, formula_through_row=7, max_row=9)
    before = _op_snapshot(amort)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=amort,
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["summary"]["applied"] == 1
    assert out["formula_fill_rows_count"] == 0
    assert out["formula_fill_last_row"] == 0
    assert out["formula_fill_columns"] == ""
    assert _op_snapshot(g.uploaded["TABLAS/amort.xlsx"]) == before


def test_apply_leaves_op_columns_untouched_with_two_events(monkeypatch):
    """Dos asientos en filas 8 y 9: tampoco se crea O9:P9 ni O10:P10."""
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 265", "TABLAS/amort_265.xlsx", fecha)
    asiento_a = "clientes/E/CREDITO # 265/ASIENTOS CONTABLES CRED 265/a1.pdf"
    asiento_b = "clientes/E/CREDITO # 265/ASIENTOS CONTABLES CRED 265/a2.pdf"
    amort = _amort_table_two_dates_at_rows(fecha, 8, 9)
    wb = openpyxl.load_workbook(io.BytesIO(amort), data_only=False)
    ws = wb["EQUINORTE"]
    for r in range(4, 8):
        ws.cell(r, 3, 30.0)
        ws.cell(r, 15, f"=+C{r}/30")
        ws.cell(r, 16, f"=+O{r}*10")
    buf = io.BytesIO()
    wb.save(buf)
    before = _op_snapshot(buf.getvalue())
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_paths": [asiento_a, asiento_b],
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_265.xlsx": buf.getvalue(),
        asiento_a: _asiento_pdf_placeholder(),
        asiento_b: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphApply(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["summary"]["applied"] == 2
    assert out["formula_fill_rows_count"] == 0
    assert _op_snapshot(g.uploaded["TABLAS/amort_265.xlsx"]) == before


def test_apply_skips_op_formula_fill_without_applied_events(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    amort = _amort_table_with_op_formulas(fecha, due_row=8, formula_through_row=7, max_row=9)
    wb = openpyxl.load_workbook(io.BytesIO(amort), data_only=False)
    ws = wb["EQUINORTE"]
    ws.cell(8, 5, fecha)
    ws.cell(8, 6, 0.0)
    ws.cell(8, 7, 49_118_143.0)
    ws.cell(8, 8, 881_857.0)
    ws.cell(8, 9, 50_000_000.0)
    buf = io.BytesIO()
    wb.save(buf)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=buf.getvalue(),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["summary"]["applied"] == 0
    assert out["summary"]["adopted"] == 1
    wb_out = openpyxl.load_workbook(io.BytesIO(g.uploaded["TABLAS/amort.xlsx"]), data_only=False)
    ws_out = wb_out["EQUINORTE"]
    assert not _is_formula_value(ws_out.cell(8, 15).value)
    assert out.get("formula_fill_rows_count", 0) == 0
