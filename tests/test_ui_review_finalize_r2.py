"""R2: Finalize atómico desde revisión (etag → preflight → Procesar=SI → enqueue)."""
from __future__ import annotations

import io
from unittest.mock import AsyncMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.adapters.primary.http.deps import init_graph_client
from app.adapters.primary.http.ui.router_v1 import (
    configure_ui_router_for_tests,
    reset_ui_router_test_hooks,
)
from app.application.services import finalize_queue_service as finalize_queue_module
from app.application.services.finalize_queue_service import (
    reset_finalize_queue_service_for_tests,
)
from app.application.services.review_schema import ControlCols, ReviewSheets
from app.application.ui.feature_flags import reset_ui_fail_closed_log_for_tests
from app.application.ui.login_rate_limit import reset_login_rate_limiter_for_tests
from app.application.ui.password_hash import hash_password
from app.application.ui.review_finalize import set_control_procesar_si
from app.application.ui.session_repository import (
    InMemorySessionRepository,
    set_session_repository_for_tests,
)
from tests.test_ui_review_edit_r1 import (
    ORIGIN,
    PROCESS_KEY,
    REVIEW_PATH,
    ReviewFakeReader,
    _MockGraph,
    _client,
    _enable_local,
    _ready_reader,
    _wb_bytes,
    _write_headers,
)
from tests.test_ui_review_read_r0 import _build_minimal_review_wb
from tests.ui_test_app import create_ui_test_app


@pytest.fixture(autouse=True)
def _reset_queues():
    reset_finalize_queue_service_for_tests()
    reset_ui_fail_closed_log_for_tests()
    reset_login_rate_limiter_for_tests()
    yield
    reset_finalize_queue_service_for_tests()


def test_set_control_procesar_si() -> None:
    wb = _build_minimal_review_wb()
    assert set_control_procesar_si(wb) is True
    ws = wb[ReviewSheets.CONTROL]
    found = False
    for r in range(1, (ws.max_row or 1) + 1):
        if str(ws.cell(r, 1).value or "").strip() == ControlCols.ROW_PROCESAR:
            assert str(ws.cell(r, 2).value) == ControlCols.VAL_PROCESAR_SI
            found = True
        if str(ws.cell(r, 1).value or "").strip() == ControlCols.ROW_ESTADO:
            assert str(ws.cell(r, 2).value).upper() == "EN_REVISION"
    assert found


def _enable_r2(monkeypatch: pytest.MonkeyPatch, *, review_edit: bool = True, finalize: bool = True) -> None:
    _enable_local(monkeypatch, review_edit=review_edit, write=True)
    monkeypatch.setenv("UI_FINALIZE_ENABLED", "true" if finalize else "false")


def test_atomic_finalize_happy_path(monkeypatch: pytest.MonkeyPatch) -> None:
    reader = _ready_reader()
    graph = _MockGraph(reader)
    _enable_r2(monkeypatch)
    init_graph_client(graph)  # type: ignore[arg-type]
    reset_ui_router_test_hooks()
    configure_ui_router_for_tests(sharepoint_reader=reader)
    monkeypatch.setattr(
        "app.application.ui.review_finalize.resolve_sharepoint_from_env",
        AsyncMock(return_value={"site_id": "site-1", "drive_id": "drive-1"}),
    )
    monkeypatch.setattr(
        "app.application.ui.review_finalize.collect_preflight_issues_from_workbook",
        lambda _wb: [],
    )

    async def _fake_fin(graph, **k):
        return {"status": "ok", "process_key": PROCESS_KEY, "already_finalized": False}

    monkeypatch.setattr(finalize_queue_module, "finalize_payment_validation", _fake_fin)

    app = create_ui_test_app()
    client = TestClient(app, base_url=ORIGIN)
    client.post(
        "/api/ui/v1/auth/login",
        json={"username": "operator", "password": "CorrectHorseBattery!"},
        headers={"Origin": ORIGIN},
    )
    csrf = client.get("/api/ui/v1/auth/csrf", headers={"Origin": ORIGIN}).json()[
        "csrf_token"
    ]
    etag = reader.current_etag()
    res = client.post(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review/finalize",
        json={"changes": []},
        headers={**_write_headers(csrf), "If-Match": etag},
    )
    assert res.status_code == 202, res.text
    body = res.json()
    assert body["accepted"] is True
    assert body["action"] == "finalize"
    assert body["job_id"]
    assert body["poll_url"].endswith(body["job_id"])
    assert len(graph.put_calls) == 1

    wb = load_workbook(io.BytesIO(reader.files[REVIEW_PATH]), data_only=False)
    ws = wb[ReviewSheets.CONTROL]
    procesar = None
    for r in range(1, (ws.max_row or 1) + 1):
        if str(ws.cell(r, 1).value or "").strip() == ControlCols.ROW_PROCESAR:
            procesar = str(ws.cell(r, 2).value or "").strip()
    assert procesar == "SI"


def test_atomic_finalize_preflight_blocks_no_put_no_enqueue(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    reader = _ready_reader()
    graph = _MockGraph(reader)
    _enable_r2(monkeypatch)
    init_graph_client(graph)  # type: ignore[arg-type]
    reset_ui_router_test_hooks()
    configure_ui_router_for_tests(sharepoint_reader=reader)
    monkeypatch.setattr(
        "app.application.ui.review_finalize.resolve_sharepoint_from_env",
        AsyncMock(return_value={"site_id": "site-1", "drive_id": "drive-1"}),
    )
    # Forzar descuadre real en el workbook
    wb = _build_minimal_review_wb()
    ws = wb[ReviewSheets.DISTRIBUCION_PAGOS]
    ws.cell(4, 5).value = 1.0  # aplicar extracto
    ws.cell(4, 7).value = 0.0
    reader.set_file(REVIEW_PATH, _wb_bytes(wb), bump=False)

    enqueue_calls: list[str] = []

    async def _no_enqueue(*a, **k):
        enqueue_calls.append("enqueue")
        raise AssertionError("enqueue no debe llamarse")

    monkeypatch.setattr(
        "app.application.ui.review_finalize.get_finalize_queue_service",
        lambda: type("S", (), {"enqueue": _no_enqueue})(),
    )

    client, csrf = _client(monkeypatch, reader, graph, review_edit=True)
    # _client sets UI_REVIEW but may not set FINALIZE — force again
    monkeypatch.setenv("UI_FINALIZE_ENABLED", "true")
    etag = reader.current_etag()
    before = reader.files[REVIEW_PATH]
    res = client.post(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review/finalize",
        json={"changes": []},
        headers={**_write_headers(csrf), "If-Match": etag},
    )
    assert res.status_code == 422, res.text
    detail = res.json()["detail"]
    assert detail["error_code"] == "review_preflight_blocked"
    assert detail["issue_count"] >= 1
    assert graph.put_calls == []
    assert reader.files[REVIEW_PATH] == before
    assert enqueue_calls == []


def test_atomic_finalize_etag_conflict(monkeypatch: pytest.MonkeyPatch) -> None:
    reader = _ready_reader()
    graph = _MockGraph(reader)
    _enable_r2(monkeypatch)
    client, csrf = _client(monkeypatch, reader, graph)
    monkeypatch.setenv("UI_FINALIZE_ENABLED", "true")
    monkeypatch.setattr(
        "app.application.ui.review_finalize.resolve_sharepoint_from_env",
        AsyncMock(return_value={"site_id": "site-1", "drive_id": "drive-1"}),
    )
    res = client.post(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review/finalize",
        json={"changes": []},
        headers={**_write_headers(csrf), "If-Match": '"stale"'},
    )
    assert res.status_code == 409
    assert res.json()["detail"]["error_code"] == "review_etag_conflict"
    assert graph.put_calls == []


def test_atomic_finalize_requires_review_edit_flag(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    reader = _ready_reader()
    graph = _MockGraph(reader)
    _enable_r2(monkeypatch, review_edit=False, finalize=True)
    client, csrf = _client(monkeypatch, reader, graph, review_edit=False)
    monkeypatch.setenv("UI_FINALIZE_ENABLED", "true")
    res = client.post(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review/finalize",
        json={"changes": []},
        headers={**_write_headers(csrf), "If-Match": reader.current_etag()},
    )
    assert res.status_code == 403
    assert res.json()["detail"]["error_code"] == "ui_review_edit_disabled"


def test_legacy_finalize_still_works(monkeypatch: pytest.MonkeyPatch) -> None:
    """Regresión: POST /processes/finalize clásico no requiere review edit."""
    from tests.ui_fixtures import make_snap
    from app.adapters.primary.http.ui.router_v1 import configure_ui_router_for_tests

    async def _fake_fin(graph, **k):
        return {"status": "ok", "process_key": PROCESS_KEY, "already_finalized": False}

    monkeypatch.setattr(finalize_queue_module, "finalize_payment_validation", _fake_fin)
    encoded = hash_password("CorrectHorseBattery!")
    monkeypatch.setenv("UI_ENABLED", "true")
    monkeypatch.setenv("UI_WRITE_ENABLED", "true")
    monkeypatch.setenv("UI_FINALIZE_ENABLED", "true")
    monkeypatch.setenv("UI_REVIEW_EDIT_ENABLED", "false")
    monkeypatch.setenv("UI_AUTH_MODE", "local_session")
    monkeypatch.setenv("ACTIVE_ENVIRONMENT", "sandbox")
    monkeypatch.setenv("UI_LOCAL_USERNAME", "operator")
    monkeypatch.setenv("UI_LOCAL_PASSWORD_HASH", encoded)
    monkeypatch.setenv("UI_LOCAL_ROLE", "operator")
    monkeypatch.setenv("UI_SESSION_TTL_MINUTES", "480")
    monkeypatch.setenv("UI_SESSION_IDLE_MINUTES", "60")
    monkeypatch.setenv("UI_LOGIN_MAX_ATTEMPTS", "5")
    monkeypatch.setenv("UI_LOGIN_WINDOW_SECONDS", "900")
    monkeypatch.setenv("UI_COOKIE_SECURE", "true")
    monkeypatch.setenv("UI_COOKIE_HTTPONLY", "true")
    monkeypatch.setenv("UI_COOKIE_SAMESITE", "strict")
    monkeypatch.setenv("UI_ALLOWED_ORIGINS", ORIGIN)
    reset_ui_fail_closed_log_for_tests()
    reset_login_rate_limiter_for_tests()
    set_session_repository_for_tests(InMemorySessionRepository())
    init_graph_client(_MockGraph(_ready_reader()))  # type: ignore[arg-type]
    reset_ui_router_test_hooks()

    snap = make_snap(
        process_key=PROCESS_KEY,
        bank_code="banco_bogota",
        bank_name="Bogotá",
        estado_proceso="REVISION_CREADA",
        validation_file_path=REVIEW_PATH,
        is_active=True,
    )
    configure_ui_router_for_tests(control_loader=lambda _bc: snap)
    client = TestClient(create_ui_test_app(), base_url=ORIGIN)
    client.post(
        "/api/ui/v1/auth/login",
        json={"username": "operator", "password": "CorrectHorseBattery!"},
        headers={"Origin": ORIGIN},
    )
    csrf = client.get("/api/ui/v1/auth/csrf", headers={"Origin": ORIGIN}).json()[
        "csrf_token"
    ]
    res = client.post(
        "/api/ui/v1/processes/finalize",
        json={"bank_code": "banco_bogota", "process_key": PROCESS_KEY},
        headers=_write_headers(csrf),
    )
    assert res.status_code == 202, res.text
    assert res.json()["job_id"]
