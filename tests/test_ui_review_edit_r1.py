"""R1: PATCH review (paridad Decimal/whitelist/ETag) + preflight dry-run."""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.adapters.primary.http.deps import init_graph_client
from app.adapters.primary.http.ui.router_v1 import (
    configure_ui_router_for_tests,
    reset_ui_router_test_hooks,
)
from app.application.services.review_schema import (
    DistribucionAbonosCols,
    DistribucionCols,
    ReviewSheets,
)
from app.application.ui.path_guard import UiAllowedRoots
from app.application.ui.password_hash import hash_password
from app.application.ui.review_preflight import collect_preflight_issues_from_workbook
from app.application.ui.review_read import parse_review_workbook
from app.application.ui.review_write import (
    ReviewPatchValidationError,
    apply_patches_to_workbook,
    etags_match,
)
from app.application.ui.schemas import UiReviewRowPatch
from app.application.ui.session_repository import (
    InMemorySessionRepository,
    set_session_repository_for_tests,
)
from app.application.ui.feature_flags import reset_ui_fail_closed_log_for_tests
from app.application.ui.login_rate_limit import reset_login_rate_limiter_for_tests
from app.application.ui.ports import UiDriveItemMeta, UiFileContent
from tests.fakes.ui_sharepoint_fake import FakeUiSharePointRead, make_fake_control
from tests.test_ui_review_read_r0 import _build_minimal_review_wb
from tests.ui_fixtures import make_snap
from tests.ui_test_app import create_ui_test_app

ORIGIN = "https://testserver"
ROOT = "INFORMACION CREDITOS-CLIENTES/03 COMWARE PRUEBAS- INFORMACION CREDITOS CLIENTES"
REVIEW_PATH = f"{ROOT}/02 VALIDACION PAGOS/01 REVISION/validacion_pagos_demo.xlsx"
CONTROL_PATH = (
    f"{ROOT}/02 VALIDACION PAGOS/90 ACCESO RESTRINGIDO/03 CONTROL TECNICO/control.xlsx"
)
PROCESS_KEY = "payment-validation|banco_bogota|2026-07-30|bb40fcea-r1"


def _wb_bytes(wb: Workbook | None = None) -> bytes:
    book = wb or _build_minimal_review_wb()
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


class ReviewFakeReader(FakeUiSharePointRead):
    """Fake con etag mutable para concurrencia."""

    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self._etag_n = 1

    def current_etag(self) -> str:
        return f'"etag-{self._etag_n}"'

    def bump_etag(self) -> str:
        self._etag_n += 1
        return self.current_etag()

    def set_file(self, path: str, content: bytes, *, bump: bool = True) -> str:
        safe = self._guard(path)
        self.files[safe] = content
        etag = self.bump_etag() if bump else self.current_etag()
        self.metas[safe] = UiDriveItemMeta(
            path=safe,
            name=safe.rsplit("/", 1)[-1],
            web_url=f"https://sharepoint.example/{safe}?web=1",
            etag=etag,
            exists=True,
        )
        return etag

    async def get_item_meta(self, relative_path: str) -> UiDriveItemMeta:
        safe = self._guard(relative_path)
        self.meta_calls.append(safe)
        if safe in self.metas:
            return self.metas[safe]
        return await super().get_item_meta(relative_path)

    async def download_bytes(self, relative_path: str) -> UiFileContent:
        safe = self._guard(relative_path)
        self.download_calls.append(safe)
        if safe not in self.files:
            raise FileNotFoundError(safe)
        etag = self.metas[safe].etag if safe in self.metas else self.current_etag()
        return UiFileContent(path=safe, content=self.files[safe], etag=etag)


class _MockGraph:
    def __init__(self, reader: ReviewFakeReader) -> None:
        self.reader = reader
        self.put_calls: list[tuple[str, bytes]] = []

    async def get(self, *a, **k):
        return {}

    async def get_bytes(self, *a, **k):
        return b""

    async def put_bytes(self, endpoint, content, content_type=None):
        self.put_calls.append((endpoint, content))
        # Simula escritura SharePoint → el reader ve el nuevo contenido + etag
        self.reader.set_file(REVIEW_PATH, content, bump=True)
        return {"id": "item-1", "eTag": self.reader.current_etag()}

    async def delete(self, *a, **k):
        return None

    async def post_json(self, *a, **k):
        return {}, 202


def _enable_local(
    monkeypatch: pytest.MonkeyPatch,
    *,
    review_edit: bool = True,
    write: bool = True,
) -> None:
    encoded = hash_password("CorrectHorseBattery!")
    monkeypatch.setenv("UI_ENABLED", "true")
    monkeypatch.setenv("UI_WRITE_ENABLED", "true" if write else "false")
    monkeypatch.setenv("UI_REVIEW_EDIT_ENABLED", "true" if review_edit else "false")
    monkeypatch.setenv("UI_AUTH_MODE", "local_session")
    monkeypatch.setenv("ACTIVE_ENVIRONMENT", "sandbox")
    monkeypatch.setenv("UI_LOCAL_USERNAME", "operator")
    monkeypatch.setenv("UI_LOCAL_PASSWORD_HASH", encoded)
    monkeypatch.setenv("UI_LOCAL_ROLE", "operator")
    monkeypatch.setenv("UI_SESSION_TTL_MINUTES", "480")
    monkeypatch.setenv("UI_SESSION_IDLE_MINUTES", "60")
    monkeypatch.setenv("UI_LOGIN_MAX_ATTEMPTS", "5")
    monkeypatch.setenv("UI_LOGIN_WINDOW_SECONDS", "900")
    monkeypatch.setenv("UI_COOKIE_SECURE", "true")
    monkeypatch.setenv("UI_COOKIE_HTTPONLY", "true")
    monkeypatch.setenv("UI_COOKIE_SAMESITE", "strict")
    monkeypatch.setenv("UI_ALLOWED_ORIGINS", ORIGIN)
    reset_ui_fail_closed_log_for_tests()
    reset_login_rate_limiter_for_tests()
    set_session_repository_for_tests(InMemorySessionRepository())


def _write_headers(csrf: str) -> dict[str, str]:
    return {
        "Origin": ORIGIN,
        "X-CSRF-Token": csrf,
        "Content-Type": "application/json",
    }


def _client(
    monkeypatch: pytest.MonkeyPatch,
    reader: ReviewFakeReader,
    graph: _MockGraph,
    *,
    review_edit: bool = True,
) -> tuple[TestClient, str]:
    _enable_local(monkeypatch, review_edit=review_edit)
    init_graph_client(graph)  # type: ignore[arg-type]
    reset_ui_router_test_hooks()
    configure_ui_router_for_tests(sharepoint_reader=reader)
    monkeypatch.setattr(
        "app.application.ui.review_write.resolve_sharepoint_from_env",
        _fake_resolve_sp,
    )
    app = create_ui_test_app()
    client = TestClient(app, base_url=ORIGIN)
    client.post(
        "/api/ui/v1/auth/login",
        json={"username": "operator", "password": "CorrectHorseBattery!"},
        headers={"Origin": ORIGIN},
    )
    csrf = client.get("/api/ui/v1/auth/csrf", headers={"Origin": ORIGIN}).json()[
        "csrf_token"
    ]
    return client, csrf


async def _fake_resolve_sp(_graph):
    return {"site_id": "site-1", "drive_id": "drive-1"}


def _ready_reader(content: bytes | None = None) -> ReviewFakeReader:
    snap = make_snap(
        control_file_path=CONTROL_PATH,
        validation_file_path=REVIEW_PATH,
        process_key=PROCESS_KEY,
        bank_code="banco_bogota",
        bank_name="Banco de Bogotá",
        estado_proceso="REVISION_CREADA",
    )
    reader = ReviewFakeReader(
        roots=UiAllowedRoots(environment="sandbox", roots=(ROOT,)),
        controls={
            "banco_bogota": make_fake_control(snap),
            "banco_bancolombia": make_fake_control(
                make_snap(
                    control_file_path=CONTROL_PATH.replace("bogota", "bancolombia"),
                    process_key="other",
                    bank_code="banco_bancolombia",
                    validation_file_path=REVIEW_PATH,
                )
            ),
        },
    )
    reader.set_file(REVIEW_PATH, content or _wb_bytes(), bump=False)
    return reader


# ─── Unit: paridad Decimal / whitelist ───────────────────────────────────────


def test_apply_patch_parity_ui_excel_get() -> None:
    wb = _build_minimal_review_wb()
    row_key = "Distribucion_Pagos|pago-1|265"
    updated = apply_patches_to_workbook(
        wb,
        [
            UiReviewRowPatch(
                row_key=row_key,
                fields={
                    "aplicar_a_extracto": "123.45",
                    "mora_a_aplicar": 10,
                    "abono_a_capital": "0.50",
                    "otros_valores": 0,
                    "observacion": "borrador incompleto ok",
                    "validar_pago": "NO",
                },
            )
        ],
    )
    assert updated == [row_key]
    pagos, _abonos, _err, _schema = parse_review_workbook(wb)
    assert len(pagos) == 1
    p = pagos[0]
    assert p.aplicar_a_extracto == 123.45
    assert p.mora_a_aplicar == 10.0
    assert p.abono_a_capital == 0.5
    assert p.otros_valores == 0.0
    assert p.observacion == "borrador incompleto ok"
    assert p.validar_pago == "NO"
    # Excel cell types: float money
    ws = wb[ReviewSheets.DISTRIBUCION_PAGOS]
    assert isinstance(ws.cell(4, 5).value, float)


def test_whitelist_rejects_unknown_field() -> None:
    wb = _build_minimal_review_wb()
    with pytest.raises(ReviewPatchValidationError) as ei:
        apply_patches_to_workbook(
            wb,
            [
                UiReviewRowPatch(
                    row_key="Distribucion_Pagos|pago-1|265",
                    fields={"monto_banco": 999},
                )
            ],
        )
    assert ei.value.code == "field_not_editable"


def test_etags_match_normalizes_quotes() -> None:
    assert etags_match('"abc"', "abc")
    assert etags_match('W/"abc"', '"abc"')
    assert not etags_match("a", "b")


def test_preflight_reports_amount_mismatch_without_mutating() -> None:
    wb = _build_minimal_review_wb()
    # Forzar descuadre: Validar SI + NORMAL + total != monto
    ws = wb[ReviewSheets.DISTRIBUCION_PAGOS]
    # headers row 3; data row 4 — aplicar extracto 1 (descuadre vs 1000.5)
    ws.cell(4, 5).value = 1.0
    ws.cell(4, 7).value = 0.0
    before = _wb_bytes(wb)
    issues = collect_preflight_issues_from_workbook(wb)
    codes = {i["error_code"] for i in issues}
    assert "amount_mismatch" in codes
    # Sin mutación de Control (preflight no escribe Procesar=SI)
    assert wb[ReviewSheets.CONTROL].cell(3, 2).value == "NO"


# ─── HTTP: flag / ETag / paridad / preflight ─────────────────────────────────


def test_bootstrap_review_edit_default_off(monkeypatch: pytest.MonkeyPatch) -> None:
    _enable_local(monkeypatch, review_edit=False)
    init_graph_client(_MockGraph(_ready_reader()))  # type: ignore[arg-type]
    reset_ui_router_test_hooks()
    client = TestClient(create_ui_test_app(), base_url=ORIGIN)
    res = client.get("/api/ui/v1/bootstrap")
    assert res.status_code == 200
    assert res.json()["review_edit_allowed"] is False


def test_patch_review_flag_off_403(monkeypatch: pytest.MonkeyPatch) -> None:
    reader = _ready_reader()
    graph = _MockGraph(reader)
    client, csrf = _client(monkeypatch, reader, graph, review_edit=False)
    etag = reader.current_etag()
    res = client.patch(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review",
        json={
            "changes": [
                {
                    "row_key": "Distribucion_Pagos|pago-1|265",
                    "fields": {"observacion": "x"},
                }
            ]
        },
        headers={**_write_headers(csrf), "If-Match": etag},
    )
    assert res.status_code == 403
    assert res.json()["detail"]["error_code"] == "ui_review_edit_disabled"
    assert graph.put_calls == []


def test_patch_review_parity_and_get(monkeypatch: pytest.MonkeyPatch) -> None:
    reader = _ready_reader()
    graph = _MockGraph(reader)
    client, csrf = _client(monkeypatch, reader, graph)
    etag = reader.current_etag()
    patch = client.patch(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review",
        json={
            "changes": [
                {
                    "row_key": "Distribucion_Pagos|pago-1|265",
                    "fields": {
                        "aplicar_a_extracto": 50.25,
                        "observacion": "desde UI",
                        "validar_pago": "NO",
                    },
                }
            ]
        },
        headers={**_write_headers(csrf), "If-Match": etag},
    )
    assert patch.status_code == 200, patch.text
    body = patch.json()
    assert "pago-1" in body["updated_row_keys"][0]
    assert body["review"]["pagos"][0]["aplicar_a_extracto"] == 50.25
    assert body["review"]["pagos"][0]["observacion"] == "desde UI"
    assert body["etag"]
    assert len(graph.put_calls) == 1

    # GET refleja lo mismo
    got = client.get(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review",
        headers={"Origin": ORIGIN},
    )
    assert got.status_code == 200
    pago = got.json()["pagos"][0]
    assert pago["aplicar_a_extracto"] == 50.25
    assert pago["observacion"] == "desde UI"
    assert pago["validar_pago"] == "NO"

    # Bytes Excel tienen el valor
    wb2 = load_workbook(io.BytesIO(reader.files[REVIEW_PATH]), data_only=False)
    pagos, *_ = parse_review_workbook(wb2)
    assert pagos[0].aplicar_a_extracto == 50.25


def test_patch_review_etag_conflict_409(monkeypatch: pytest.MonkeyPatch) -> None:
    reader = _ready_reader()
    graph = _MockGraph(reader)
    client, csrf = _client(monkeypatch, reader, graph)
    res = client.patch(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review",
        json={
            "changes": [
                {
                    "row_key": "Distribucion_Pagos|pago-1|265",
                    "fields": {"observacion": "stale"},
                }
            ]
        },
        headers={**_write_headers(csrf), "If-Match": '"etag-stale"'},
    )
    assert res.status_code == 409
    assert res.json()["detail"]["error_code"] == "review_etag_conflict"
    assert graph.put_calls == []


def test_patch_review_missing_if_match(monkeypatch: pytest.MonkeyPatch) -> None:
    reader = _ready_reader()
    graph = _MockGraph(reader)
    client, csrf = _client(monkeypatch, reader, graph)
    res = client.patch(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review",
        json={"changes": []},
        headers=_write_headers(csrf),
    )
    assert res.status_code == 428
    assert res.json()["detail"]["error_code"] == "missing_if_match"


def test_preflight_dry_run_no_put(monkeypatch: pytest.MonkeyPatch) -> None:
    reader = _ready_reader()
    graph = _MockGraph(reader)
    client, csrf = _client(monkeypatch, reader, graph)
    before = reader.files[REVIEW_PATH]
    res = client.post(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review/preflight",
        json={},
        headers=_write_headers(csrf),
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert "ok" in body
    assert "issues" in body
    assert graph.put_calls == []
    assert reader.files[REVIEW_PATH] == before


def test_r0_get_still_works_with_flag_off(monkeypatch: pytest.MonkeyPatch) -> None:
    """Regresión R0: GET review sin flag de edición."""
    reader = _ready_reader()
    graph = _MockGraph(reader)
    _enable_local(monkeypatch, review_edit=False)
    init_graph_client(graph)  # type: ignore[arg-type]
    reset_ui_router_test_hooks()
    configure_ui_router_for_tests(sharepoint_reader=reader)
    client = TestClient(create_ui_test_app(), base_url=ORIGIN)
    client.post(
        "/api/ui/v1/auth/login",
        json={"username": "operator", "password": "CorrectHorseBattery!"},
        headers={"Origin": ORIGIN},
    )
    res = client.get(
        f"/api/ui/v1/processes/{PROCESS_KEY}/review",
        headers={"Origin": ORIGIN},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["read_only"] is True
    assert data["pagos"][0]["row_key"] == "Distribucion_Pagos|pago-1|265"
    assert "aplicar_a_extracto" in data["pagos"][0]["editable_fields"]
