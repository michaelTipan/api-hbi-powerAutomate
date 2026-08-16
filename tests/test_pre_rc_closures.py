"""Cierres pre-RC: retenciones, fórmula VP, IBR anticipado, filas, hard blockers."""

from __future__ import annotations

import asyncio
import io
from datetime import date, timedelta

import openpyxl
import pytest

from app.application.services.accounting_pdf_parser import (
    ACCOUNT_CAPITAL,
    ACCOUNT_VALOR_PAGADO_CLIENTE,
    PaymentApplicationEvent,
)
from app.application.services.amortization_workbook import (
    APLICADO,
    _default_valor_pagado_formula,
    detect_amortization_sheet,
    find_application_row_detailed,
    write_payment_application,
    PaymentApplicationWriteOptions,
)
from app.application.services.review_schema import (
    TipoAplicacionConfirmado,
    resolve_actualiza_ibr,
    resolve_policy_from_tipo_confirmado,
)
from app.application.ui.amortization_operational_issues import (
    attach_operational_issues_to_amortization_result,
)
from app.application.use_cases.amortization_application_plan import (
    AmortizationPreparedPlan,
    verify_amortization_plan_freshness,
)
from app.application.use_cases.amortization_fill_apply import (
    _item_actualiza_ibr,
    execute_amortization_from_prepared,
)
from app.application.use_cases.amortization_fill_dry_run import (
    HARD_STRUCTURAL_APPLY_BLOCKERS,
    RETENCIONES_COLUMN_MISSING,
    _policy_requires_ibr,
)


def _event(**kwargs) -> PaymentApplicationEvent:
    base = dict(
        id_pago="P1",
        cliente="CLI",
        credito="327",
        asiento_pdf_path="a.pdf",
        comprobante="c1",
        fecha_asiento=date(2026, 5, 10),
        valor_pagado_cliente=1000.0,
        capital=700.0,
        intereses=200.0,
        mora=100.0,
        retenciones=0.0,
        saldos_menores=0.0,
        raw_text="",
    )
    base.update(kwargs)
    return PaymentApplicationEvent(**base)


def test_fallback_vp_formula_subtracts_retenciones_when_header_exists():
    headers = {
        "valor_pagado_cliente": 12,
        "valor_intereses": 9,
        "abono_k": 10,
        "intereses_mora": 11,
        "retenciones": 13,
    }
    assert _default_valor_pagado_formula(headers, 51) == "=+I51+J51+K51-M51"


def test_fallback_vp_formula_without_retenciones_header():
    headers = {
        "valor_pagado_cliente": 12,
        "valor_intereses": 9,
        "abono_k": 10,
        "intereses_mora": 11,
    }
    assert _default_valor_pagado_formula(headers, 5) == "=+I5+J5+K5"


def test_template_formula_has_priority_over_fallback():
    from app.application.services.amortization_workbook import _find_formula_template

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["h1", "h2", "h3", "h4", "Valor pagado cliente"])
    ws.append([None, None, None, None, "=+B2+C2+D2-X2"])
    ws.append([None, None, None, None, None])
    templ = _find_formula_template(ws, 5, 3, 1)
    assert templ is not None
    assert "X3" in templ or "X2" in templ or templ.startswith("=")
    fallback = _default_valor_pagado_formula(
        {
            "valor_pagado_cliente": 5,
            "valor_intereses": 2,
            "abono_k": 3,
            "intereses_mora": 4,
            "retenciones": 6,
        },
        3,
    )
    assert fallback != templ


def test_hard_structural_includes_retenciones():
    assert RETENCIONES_COLUMN_MISSING in HARD_STRUCTURAL_APPLY_BLOCKERS


def test_five_items_hard_blocker_blocks_all_writes():
    items = []
    for i in range(1, 6):
        item = {
            "id_pago": f"P{i}",
            "cliente": "CLI",
            "credito": str(i),
            "tipo_aplicacion": "PAGO",
            "application_status": "WOULD_APPLY" if i != 4 else "ERROR",
            "error_code": None if i != 4 else RETENCIONES_COLUMN_MISSING,
            "tabla_amortizacion_path": f"TABLAS/{i}.xlsx",
            "asiento_pdf_path": f"AS/{i}.pdf",
            "valor_retenciones": 4_378_159 if i == 4 else 0,
            "payment_application": {"retenciones": 4_378_159 if i == 4 else 0},
        }
        items.append(item)

    payment_applicable = any(
        it.get("application_status") in ("WOULD_APPLY", "WOULD_ADOPT_EXISTING")
        and not it.get("error_code")
        for it in items
    )
    has_hard = any(
        str(it.get("error_code") or "") in HARD_STRUCTURAL_APPLY_BLOCKERS for it in items
    )
    can_apply = payment_applicable and not has_hard
    assert payment_applicable is True
    assert can_apply is False

    plan = AmortizationPreparedPlan(
        can_apply=True,
        dry_run={"items": items, "can_apply": True, "summary": {}},
        fingerprints={},
        site_id="s1",
        drive_id="d1",
        resolved_bank_code="banco_bogota",
    )
    puts: list[str] = []

    class G:
        async def get(self, *a, **k):
            return {}

        async def get_bytes(self, *a, **k):
            raise FileNotFoundError()

        async def put_bytes(self, endpoint, content, content_type="", if_match=None):
            puts.append(endpoint)
            return {}

        async def post_json(self, *a, **k):
            return {}, 202

        async def patch_json(self, *a, **k):
            return {}

        async def delete(self, *a, **k):
            return None

    result = asyncio.run(execute_amortization_from_prepared(G(), plan))
    assert result.get("can_apply") is False
    assert result.get("apply_wrote_changes") is False
    assert puts == []
    assert (result.get("tables_uploaded") or []) == []


def test_retenciones_operational_issue_copy_cta_tabla_not_asientos():
    result = {
        "status": "blocked",
        "mode": "apply",
        "outcome": "requires_correction",
        "can_apply": False,
        "items": [
            {
                "id_pago": "P9",
                "cliente": "EL CONDOR",
                "credito": "248",
                "application_status": "ERROR",
                "error_code": RETENCIONES_COLUMN_MISSING,
                "tabla_amortizacion_path": "clientes/CONDOR/CREDITO# 248/tabla.xlsx",
                "asiento_pdf_path": "clientes/CONDOR/ASIENTOS/asiento.pdf",
                "ruta_asientos_contables": "clientes/CONDOR/ASIENTOS",
                "valor_retenciones": 4_378_159,
                "payment_application": {"retenciones": 4_378_159},
            }
        ],
    }
    out = attach_operational_issues_to_amortization_result(result)
    issues = out.get("operational_issues") or []
    assert issues
    issue = issues[0]
    msg = issue["user_message"]
    assert "RETENCIONES_COLUMN_MISSING" not in msg
    assert "columna RETENCIONES" in msg
    assert "$4.378.159" in msg
    assert "No se realizó ninguna modificación" in msg
    assert "Abra la tabla de amortización" in (issue.get("next_action") or "")
    labels = [lk.get("label") for lk in (issue.get("links") or [])]
    assert "Abrir tabla de amortización" in labels
    assert not any("ASIENTOS" in (lb or "").upper() and "tabla" not in (lb or "").lower() for lb in labels)
    assert issue["location"]["credit"] == "248"
    assert issue["location"]["payment_id"] == "P9"
    assert issue["recoverable"] is True
    assert issue["technical_reference"] == RETENCIONES_COLUMN_MISSING


IBR_TIPOS = list(TipoAplicacionConfirmado.OPTIONS_ORDERED)


@pytest.mark.parametrize("tipo", IBR_TIPOS)
@pytest.mark.parametrize("delta_days", [-20, -1, 0, 1])
def test_ibr_anticipado_canonical_dry_run_matches_apply(tipo, delta_days):
    policy = resolve_policy_from_tipo_confirmado(tipo)
    fecha_limite = date(2026, 6, 15)
    payment_date = fecha_limite + timedelta(days=delta_days)
    expected = resolve_actualiza_ibr(
        policy, payment_date=payment_date, fecha_limite=fecha_limite
    )
    dry = _policy_requires_ibr(
        policy, payment_date=payment_date, fecha_limite=fecha_limite
    )
    item = {
        "tipo_aplicacion_original": tipo,
        "fecha_limite_pago": fecha_limite.isoformat(),
        "payment_date_iso": payment_date.isoformat(),
    }
    apply_dec = _item_actualiza_ibr(item)
    assert dry is expected
    assert apply_dec is expected
    if payment_date < fecha_limite:
        assert expected is False


def test_two_distinct_pdfs_same_amounts_get_distinct_rows():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
        ]
    )
    for day in range(1, 25):
        ws.append([day, 5, 2026, None, None, None, None, None])
    match = detect_amortization_sheet(wb)
    ev1 = _event(asiento_pdf_path="p1.pdf", comprobante="c1")
    ev2 = _event(asiento_pdf_path="p2.pdf", comprobante="c2")
    r1 = find_application_row_detailed(
        match.worksheet,
        match.headers,
        ev1,
        due_date_row=3,
        header_row=match.header_row,
        exclude_rows=frozenset(),
    )
    assert r1.row is not None
    opts = PaymentApplicationWriteOptions(
        payment_date=date(2026, 5, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    write_payment_application(
        match.worksheet,
        r1.row,
        match.headers,
        ev1,
        write_options=opts,
        header_row=match.header_row,
    )
    r2 = find_application_row_detailed(
        match.worksheet,
        match.headers,
        ev2,
        due_date_row=3,
        header_row=match.header_row,
        exclude_rows=frozenset({r1.row}),
    )
    assert r2.row is not None
    assert r2.row != r1.row


def test_retry_same_event_adopts_without_exclude():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
        ]
    )
    ws.append([22, 5, 2026, None, None, None, None, None])
    match = detect_amortization_sheet(wb)
    ev = _event()
    first = find_application_row_detailed(
        match.worksheet, match.headers, ev, due_date_row=2, header_row=match.header_row
    )
    opts = PaymentApplicationWriteOptions(
        payment_date=date(2026, 5, 10),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    write_payment_application(
        match.worksheet, first.row, match.headers, ev, write_options=opts, header_row=match.header_row
    )
    retry = find_application_row_detailed(
        match.worksheet, match.headers, ev, due_date_row=2, header_row=match.header_row
    )
    assert retry.row == first.row
    from app.application.services.amortization_workbook import ADOPTADO_EXISTENTE

    assert retry.compare_status == ADOPTADO_EXISTENTE


def test_twenty_payments_unique_rows():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
        ]
    )
    for day in range(1, 25):
        ws.append([day, 5, 2026, None, None, None, None, None])
    match = detect_amortization_sheet(wb)
    reserved: set[int] = set()
    rows: list[int] = []
    for i in range(20):
        ev = _event(asiento_pdf_path=f"p{i}.pdf", comprobante=f"c{i}")
        found = find_application_row_detailed(
            match.worksheet,
            match.headers,
            ev,
            due_date_row=2,
            header_row=match.header_row,
            exclude_rows=frozenset(reserved),
        )
        assert found.row is not None
        assert found.compare_status == APLICADO
        reserved.add(found.row)
        rows.append(found.row)
    assert len(set(rows)) == 20


def test_table_freshness_change_blocks_apply():
    tabla = "TABLAS/x.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["dia", "mes", "año", "Fecha pago", "Abono a K"])
    ws.append([1, 5, 2026, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    original = buf.getvalue()
    plan = AmortizationPreparedPlan(
        can_apply=True,
        dry_run={
            "items": [
                {
                    "tabla_amortizacion_path": tabla,
                    "tabla_sha256": "deadbeef",
                    "payment_application": {"retenciones": 0},
                }
            ]
        },
        fingerprints={
            "process_key": "pk",
            "table_fingerprints": [{"path": tabla, "sha256": "deadbeef"}],
        },
        site_id="s1",
        drive_id="d1",
        resolved_bank_code="banco_bogota",
        apply_idempotency_key="pk",
    )

    class G:
        async def get(self, *a, **k):
            return {"eTag": "etag-1"}

        async def get_bytes(self, *a, **k):
            return original

    stale = asyncio.run(verify_amortization_plan_freshness(G(), plan))
    assert stale is not None
    assert stale.get("error_code") == "AMORTIZATION_TABLE_CHANGED_REQUIRES_REVALIDATION"
    assert stale.get("can_apply") is False


def test_freshness_etag_before_after_mismatch_blocks():
    import hashlib

    tabla = "TABLAS/x.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["dia", "mes", "año", "Fecha pago", "Abono a K"])
    ws.append([1, 5, 2026, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    original = buf.getvalue()
    sha = hashlib.sha256(original).hexdigest()
    plan = AmortizationPreparedPlan(
        can_apply=True,
        dry_run={
            "items": [
                {
                    "tabla_amortizacion_path": tabla,
                    "tabla_sha256": sha,
                    "payment_application": {"retenciones": 0},
                }
            ]
        },
        fingerprints={
            "process_key": "pk",
            "table_fingerprints": [{"path": tabla, "sha256": sha}],
        },
        site_id="s1",
        drive_id="d1",
        resolved_bank_code="banco_bogota",
        apply_idempotency_key="pk",
    )

    class G:
        def __init__(self) -> None:
            self.gets = 0

        async def get(self, *a, **k):
            self.gets += 1
            return {"eTag": f'"etag-{self.gets}"'}

        async def get_bytes(self, *a, **k):
            return original

    stale = asyncio.run(verify_amortization_plan_freshness(G(), plan))
    assert stale is not None
    assert stale.get("error_code") == "AMORTIZATION_TABLE_CHANGED_REQUIRES_REVALIDATION"


def test_apply_one_table_http_412_maps_to_revalidation():
    import httpx

    from app.application.use_cases.amortization_fill_apply import (
        UPLOAD_STATUS_PRECONDITION_FAILED,
        _apply_one_table,
    )
    from tests.test_amortization_workbook import _amort_workbook_bytes

    raw = _amort_workbook_bytes()

    class G:
        async def get_bytes(self, *a, **k):
            return raw

        async def put_bytes(self, *a, **k):
            req = httpx.Request("PUT", "https://graph.test/content")
            resp = httpx.Response(412, request=req)
            raise httpx.HTTPStatusError("Precondition Failed", request=req, response=resp)

    item = {
        "application_status": "WOULD_APPLY",
        "error_code": None,
        "tabla_amortizacion_path": "TABLAS/x.xlsx",
        "application_row": 3,
        "ibr_row": 3,
        "due_date_row": 3,
        "id_pago": "P1",
        "credito": "258",
        "cliente": "EQUINORTE",
        "asiento_pdf_path": "a.pdf",
        "idempotency_key": "k-412",
        "payment_date_iso": "2026-05-10",
        "payment_application": {
            "valor_pagado_cliente": 50_000_000.0,
            "capital": 49_118_143.0,
            "intereses": 0.0,
            "mora": 881_857.0,
            "retenciones": 0.0,
            "saldos_menores": 0.0,
        },
        "detected_codes": ["544111100505", "544113410519"],
        "actualiza_ibr": False,
    }
    out = asyncio.run(
        _apply_one_table(
            G(),
            "s1",
            "d1",
            "TABLAS/x.xlsx",
            [item],
            tabla_bytes=raw,
            tabla_etag='"v1"',
        )
    )
    assert out["uploaded"] is False
    assert out["upload_status"] == UPLOAD_STATUS_PRECONDITION_FAILED
    assert out["error_code"] == "AMORTIZATION_TABLE_CHANGED_REQUIRES_REVALIDATION"


def test_apply_second_table_http_412_partial_no_retry_no_later_puts():
    import httpx
    from urllib.parse import unquote

    from app.application.use_cases.amortization_fill_apply import (
        APPLY_STATUS_SKIPPED_IDEMPOTENT,
        execute_amortization_from_prepared,
    )
    from tests.test_amortization_workbook import _amort_workbook_bytes

    raw_a = _amort_workbook_bytes()
    raw_b = _amort_workbook_bytes()
    raw_c = _amort_workbook_bytes()
    blobs = {
        "TABLAS/A.xlsx": raw_a,
        "TABLAS/B.xlsx": raw_b,
        "TABLAS/C.xlsx": raw_c,
    }
    puts: list[str] = []

    def _planned(tabla: str, id_pago: str) -> dict:
        return {
            "application_status": "WOULD_APPLY",
            "error_code": None,
            "tabla_amortizacion_path": tabla,
            "application_row": 3,
            "ibr_row": 3,
            "due_date_row": 3,
            "id_pago": id_pago,
            "credito": "258",
            "cliente": "EQUINORTE",
            "asiento_pdf_path": f"{id_pago}.pdf",
            "idempotency_key": f"k-{id_pago}",
            "payment_date_iso": "2026-05-10",
            "payment_application": {
                "valor_pagado_cliente": 50_000_000.0,
                "capital": 49_118_143.0,
                "intereses": 0.0,
                "mora": 881_857.0,
                "retenciones": 0.0,
                "saldos_menores": 0.0,
            },
            "detected_codes": ["544111100505", "544113410519"],
            "actualiza_ibr": False,
        }

    items = [
        _planned("TABLAS/A.xlsx", "PA"),
        _planned("TABLAS/B.xlsx", "PB"),
        _planned("TABLAS/C.xlsx", "PC"),
    ]
    plan = AmortizationPreparedPlan(
        can_apply=True,
        dry_run={"items": items, "can_apply": True, "summary": {}},
        fingerprints={},
        site_id="s1",
        drive_id="d1",
        resolved_bank_code="banco_bogota",
        apply_idempotency_key="pk-412-multi",
        table_bytes=dict(blobs),
        table_etags={
            "TABLAS/A.xlsx": '"ea"',
            "TABLAS/B.xlsx": '"eb"',
            "TABLAS/C.xlsx": '"ec"',
        },
    )

    class G:
        async def get(self, *a, **k):
            return {}

        async def get_bytes(self, endpoint, *a, **k):
            decoded = unquote(str(endpoint))
            for path, data in blobs.items():
                if path in decoded:
                    return data
            raise FileNotFoundError(endpoint)

        async def put_bytes(self, endpoint, content, content_type="", if_match=None):
            decoded = unquote(str(endpoint))
            puts.append(decoded)
            if "TABLAS/B.xlsx" in decoded:
                req = httpx.Request("PUT", "https://graph.test/content")
                resp = httpx.Response(412, request=req)
                raise httpx.HTTPStatusError(
                    "Precondition Failed", request=req, response=resp
                )
            for path in blobs:
                if path in decoded:
                    blobs[path] = content
                    return {}
            return {}

        async def post_json(self, *a, **k):
            return {}, 202

        async def patch_json(self, *a, **k):
            return {}

        async def delete(self, *a, **k):
            return None

    first = asyncio.run(execute_amortization_from_prepared(G(), plan))
    b_puts = [p for p in puts if "TABLAS/B.xlsx" in p]
    c_puts = [p for p in puts if "TABLAS/C.xlsx" in p]
    a_puts = [p for p in puts if "TABLAS/A.xlsx" in p]
    assert len(a_puts) == 1
    assert len(b_puts) == 1
    assert c_puts == []
    assert first.get("error_code") == "AMORTIZATION_TABLE_CHANGED_REQUIRES_REVALIDATION"
    assert first.get("apply_wrote_changes") is True
    assert "TABLAS/A.xlsx" in (first.get("tables_uploaded") or [])
    assert first.get("status") == "partial"
    um = str(first.get("user_message") or "").casefold()
    assert "no se modificó ninguna tabla" not in um
    assert "no se realizó ninguna modificación" not in um or "tabla que cambió" in um
    assert first.get("can_apply") is False
    assert any(
        e.get("error_code") == "AMORTIZATION_TABLE_CHANGED_REQUIRES_REVALIDATION"
        and "TABLAS/B.xlsx" in str(e.get("tabla_amortizacion_path") or "")
        for e in (first.get("apply_errors") or [])
    )

    puts.clear()
    retry_plan = AmortizationPreparedPlan(
        can_apply=True,
        dry_run={"items": items, "can_apply": True, "summary": {}},
        fingerprints={},
        site_id="s1",
        drive_id="d1",
        resolved_bank_code="banco_bogota",
        apply_idempotency_key="pk-412-multi",
        table_bytes=dict(blobs),
        table_etags={
            "TABLAS/A.xlsx": '"ea2"',
            "TABLAS/B.xlsx": '"eb"',
            "TABLAS/C.xlsx": '"ec"',
        },
    )
    second = asyncio.run(execute_amortization_from_prepared(G(), retry_plan))
    assert [p for p in puts if "TABLAS/A.xlsx" in p] == []
    assert len([p for p in puts if "TABLAS/B.xlsx" in p]) == 1
    assert [p for p in puts if "TABLAS/C.xlsx" in p] == []
    a_items = [
        it
        for it in (second.get("items") or [])
        if str(it.get("tabla_amortizacion_path") or "") == "TABLAS/A.xlsx"
    ]
    assert a_items
    assert all(
        it.get("apply_status") == APPLY_STATUS_SKIPPED_IDEMPOTENT for it in a_items
    )


def test_248_policy_flags_untouched():
    p = resolve_policy_from_tipo_confirmado(
        TipoAplicacionConfirmado.CANCELACION_PAGO_TOTAL
    )
    assert p.requiere_extracto is True
    assert p.include_extract_in_composite is False
    assert p.payoff_expected is True


def test_freshness_missing_etag_blocks_without_writes():
    import hashlib

    tabla = "TABLAS/x.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["dia", "mes", "año", "Fecha pago", "Abono a K"])
    ws.append([1, 5, 2026, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    original = buf.getvalue()
    sha = hashlib.sha256(original).hexdigest()
    plan = AmortizationPreparedPlan(
        can_apply=True,
        dry_run={
            "items": [
                {
                    "tabla_amortizacion_path": tabla,
                    "tabla_sha256": sha,
                    "payment_application": {"retenciones": 0},
                }
            ]
        },
        fingerprints={
            "process_key": "pk",
            "table_fingerprints": [{"path": tabla, "sha256": sha}],
        },
        site_id="s1",
        drive_id="d1",
        resolved_bank_code="banco_bogota",
        apply_idempotency_key="pk",
    )

    class G:
        async def get(self, *a, **k):
            return {}

        async def get_bytes(self, *a, **k):
            return original

    stale = asyncio.run(verify_amortization_plan_freshness(G(), plan))
    assert stale is not None
    assert stale.get("error_code") == "AMORTIZATION_TABLE_CHANGED_REQUIRES_REVALIDATION"
    assert stale.get("can_apply") is False


def test_apply_one_table_without_etag_does_not_put():
    from app.application.use_cases.amortization_fill_apply import (
        UPLOAD_STATUS_PRECONDITION_FAILED,
        _apply_one_table,
    )
    from tests.test_amortization_workbook import _amort_workbook_bytes

    raw = _amort_workbook_bytes()
    puts: list[str] = []

    class G:
        async def get_bytes(self, *a, **k):
            return raw

        async def put_bytes(self, *a, **k):
            puts.append("put")
            return {}

    item = {
        "application_status": "WOULD_APPLY",
        "error_code": None,
        "tabla_amortizacion_path": "TABLAS/x.xlsx",
        "application_row": 3,
        "ibr_row": 3,
        "due_date_row": 3,
        "id_pago": "P1",
        "credito": "258",
        "cliente": "EQUINORTE",
        "asiento_pdf_path": "a.pdf",
        "idempotency_key": "k-no-etag",
        "payment_date_iso": "2026-05-10",
        "payment_application": {
            "valor_pagado_cliente": 50_000_000.0,
            "capital": 49_118_143.0,
            "intereses": 0.0,
            "mora": 881_857.0,
            "retenciones": 0.0,
            "saldos_menores": 0.0,
        },
        "detected_codes": ["544111100505", "544113410519"],
        "actualiza_ibr": False,
    }
    out = asyncio.run(
        _apply_one_table(
            G(),
            "s1",
            "d1",
            "TABLAS/x.xlsx",
            [item],
            tabla_bytes=raw,
            tabla_etag=None,
        )
    )
    assert puts == []
    assert out["uploaded"] is False
    assert out["upload_status"] == UPLOAD_STATUS_PRECONDITION_FAILED


def test_asiento_metadata_failure_is_parse_failed_stage():
    from app.application.services.asiento_lote_assignment import CandidateParseFailure
    from app.application.use_cases.merge_composite_validado_pdfs import (
        _parse_asiento_candidate_cached,
    )

    class G:
        async def get_bytes(self, *a, **k):
            return b"%PDF-1.4 dummy"

        async def get(self, *a, **k):
            raise RuntimeError("graph metadata unavailable")

    parsed = asyncio.run(
        _parse_asiento_candidate_cached(G(), "s1", "d1", "AS/a-327.pdf", "327", {})
    )
    assert isinstance(parsed, CandidateParseFailure)
    assert parsed.stage == "metadata"
    assert parsed.credit == "327"
