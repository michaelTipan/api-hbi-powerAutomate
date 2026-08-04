"""Apply Fase 4B: fail-closed cuando hay grupos ABONO no habilitados."""

import asyncio
import hashlib
import io
import json
from datetime import date

import openpyxl
import pytest

from app.application.services.abono_dry_run import (
    ABONO_ASIENTO_DUPLICADO,
    ABONO_ASIENTOS_NO_CUADRAN,
    build_abono_group_from_manifest_output,
    reconcile_abono_group,
)
from app.application.services.amortization_workbook import AUTOMATION_LOG_SHEET
from app.application.use_cases.amortization_fill_apply import run_amortization_fill_apply
from tests.test_abono_dry_run import (
    _abono_dry_run_files,
    _abono_manifest_output,
    _asiento_pdf,
)
from tests.test_amortization_fill_apply import MockGraphApply
from tests.test_amortization_fill_dry_run import (
    _accounting_text,
    _amort_table_bytes,
    _amort_table_date_at_row,
    _asiento_pdf_placeholder,
    _base_files,
    _hist_bytes,
    _ibr_bytes,
)


@pytest.fixture(autouse=True)
def env_sharepoint(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")


def _patch_pdf_extract(monkeypatch):
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    monkeypatch.setattr(
        "app.application.services.abono_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )


def _assert_no_side_effects(g: MockGraphApply, *, asiento_paths: tuple[str, ...] = ()) -> None:
    assert g.put_calls == []
    assert g.uploaded == {}
    assert not g.patch_calls
    for path in asiento_paths:
        assert path in g.files


def _automation_log_rows(g: MockGraphApply, tabla: str) -> int:
    blob = g.files.get(tabla)
    if not blob:
        return 0
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    if AUTOMATION_LOG_SHEET not in wb.sheetnames:
        return 0
    return max(0, wb[AUTOMATION_LOG_SHEET].max_row - 1)


def test_apply_abono_not_reconciled_blocks_before_write(monkeypatch):
    fecha = date(2026, 5, 22)
    out = _abono_manifest_output(monto_banco=10_000_000.0)
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    g = MockGraphApply(files)
    _patch_pdf_extract(monkeypatch)

    result = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )

    assert result["status"] == "blocked"
    assert result["apply_wrote_changes"] is False
    assert result["can_apply"] is False
    gr = result["blocking_abono_groups"][0]
    assert gr["reconciliation_status"] == "FAILED"
    assert any(
        e.get("error_code") == ABONO_ASIENTOS_NO_CUADRAN for e in gr.get("blocking_errors") or []
    )
    _assert_no_side_effects(g)


def test_apply_abono_missing_asiento_blocks(monkeypatch):
    fecha = date(2026, 5, 22)
    out = _abono_manifest_output()
    out["credit_items"][0]["asiento_pdf_paths"] = []
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    g = MockGraphApply(files)
    _patch_pdf_extract(monkeypatch)

    result = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )

    assert result["status"] == "blocked"
    assert result["apply_wrote_changes"] is False
    assert result["blocking_abono_groups"]
    assert result["blocking_abono_groups"][0]["blocking_errors"]
    _assert_no_side_effects(g)


def test_apply_mixed_pago_ready_abono_not_reconciled_blocks_all(monkeypatch):
    fecha = date(2026, 5, 22)
    pago = {
        "id_pago": "P1",
        "cliente": "EQUINORTE",
        "credito": "CREDITO # 258",
        "tipo_aplicacion": "PAGO",
        "requiere_extracto": True,
        "asiento_pdf_path": "clientes/EQUINORTE/CREDITO # 258/ASIENTOS CONTABLES CRED 258/asiento.pdf",
        "extracto_pdf_path": "clientes/EQUINORTE/extracto.pdf",
        "output_relative_path": "OUT/pago.pdf",
    }
    abono = _abono_manifest_output(id_pago="AB9", monto_banco=10_000_000.0)
    files = _abono_dry_run_files(abono_output=abono, pago_manifest_output=pago, fecha=fecha)
    files["TABLAS/amort_258.xlsx"] = _amort_table_date_at_row(fecha, 8)
    pago_tabla_hash = hashlib.sha256(files["TABLAS/amort_258.xlsx"]).hexdigest()
    g = MockGraphApply(files)
    _patch_pdf_extract(monkeypatch)

    result = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )

    assert result["status"] == "blocked"
    assert result["apply_wrote_changes"] is False
    assert result["abono_groups_blocked"] >= 1
    um = str(result.get("user_message") or "").lower()
    assert "problema" in um
    assert "ninguna tabla" in um
    issues = result.get("operational_issues") or []
    assert issues, "expected operational_issues when abono blocks apply"
    assert any("cuadra" in str(i.get("user_message") or "").lower() for i in issues)
    _assert_no_side_effects(g)
    assert hashlib.sha256(g.files["TABLAS/amort_258.xlsx"]).hexdigest() == pago_tabla_hash


def test_apply_pago_only_no_regression(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    _patch_pdf_extract(monkeypatch)

    result = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )

    assert result["status"] == "ok"
    assert result["apply_wrote_changes"] is True
    assert result["summary"]["applied"] == 1
    assert len(g.put_calls) == 1


def test_apply_already_applied_before_dry_run(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("851eb0f6", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    asiento = "clientes/EQUINORTE/CREDITO # 258/ASIENTOS CONTABLES CRED 258/asiento.pdf"
    process_key = f"payment-validation|banco_bogota|{fecha.isoformat()}"
    manifest_key = f"LOGS/merge_manifest_banco_bogota_{fecha.isoformat()}.json"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "851eb0f6",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 258",
                "asiento_pdf_path": asiento,
            }
        ],
    }
    files = _base_files(
        hist=hist,
        amort=_amort_table_date_at_row(fecha, 8),
        asiento_pdf=_asiento_pdf_placeholder(),
        ibr=_ibr_bytes(),
        fecha=fecha,
    )
    files[manifest_key] = json.dumps(manifest).encode("utf-8")
    g = MockGraphApply(files)
    dry_run_calls: list[int] = []

    async def _track_dry_run(graph, **kwargs):
        dry_run_calls.append(1)
        from app.application.use_cases.amortization_fill_dry_run import run_amortization_fill_dry_run

        return await run_amortization_fill_dry_run(graph, **kwargs)

    from app.application.use_cases.payment_validation_process_control import ProcessControlSnapshot

    aplicada = ProcessControlSnapshot(
        control_file_path="CTL/bogota.xlsx",
        estado_proceso="AMORTIZACION_APLICADA",
        is_active=True,
        process_key=process_key,
        process_id="",
        validation_file_path="",
        historical_file_path="HIST/cartera.xlsx",
        secretary_file_path="",
        email_pdf_path="",
        notify_idempotency_key="",
        merge_manifest_path=manifest_key,
        merge_idempotency_key="",
        apply_idempotency_key=process_key,
        bank_code="banco_bogota",
        bank_name="Banco de Bogotá",
    )

    async def _read_snap(_g, _s, _d, *, bank_code: str):
        return aplicada

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.read_process_control_snapshot",
        _read_snap,
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.run_amortization_fill_dry_run",
        _track_dry_run,
    )

    result = asyncio.run(
        run_amortization_fill_apply(
            g,
            merge_manifest_path=manifest_key,
            historical_file_path="HIST/cartera.xlsx",
            bank_code="banco_bogota",
        )
    )

    assert result["status"] == "ok"
    assert result["already_applied"] is True
    assert len(dry_run_calls) == 0
    _assert_no_side_effects(g)


def test_reconcile_abono_same_pdf_path_not_double_counted(monkeypatch):
    """Un asiento_pdf_path repetido en el grupo no suma valor_pagado_cliente dos veces."""
    out = _abono_manifest_output(monto_banco=50_000_000.0)
    path = out["credit_items"][0]["asiento_pdf_paths"][0]
    out["credit_items"][0]["asiento_pdf_paths"] = [path, path]
    group = build_abono_group_from_manifest_output(out, bank_code="banco_bogota", process_key="pk")
    monkeypatch.setattr(
        "app.application.services.abono_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )

    async def _run():
        async def _dl(_path: str) -> bytes:
            return _asiento_pdf()

        return await reconcile_abono_group(group, _dl)

    result = asyncio.run(_run())
    assert result.reconciliation_status == "FAILED"
    assert any(e["error_code"] == ABONO_ASIENTO_DUPLICADO for e in result.blocking_errors)
    assert result.total_asientos != 100_000_000
