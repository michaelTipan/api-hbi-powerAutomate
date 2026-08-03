"""Tests de contrato de errores de negocio y preflight de revisión."""

from __future__ import annotations

import io

import openpyxl
import pytest

from app.application.job_status_enrichment import resolve_business_error_for_job
from app.application.services.business_error_contract import (
    job_error_storage_dict,
    parse_business_error_message,
)
from app.application.services.excel_write_guard import ExcelWriteGuardError
from app.application.services.review_schema import (
    ControlCols,
    DistribucionCols,
    ErroresCols,
    ReviewSheets,
    ValidarPago,
)
from app.application.services.review_workbook_preflight import (
    collect_review_workbook_preflight_issues,
    count_open_errores_rows,
)


def test_parse_business_error_message_json_details() -> None:
    code, details = parse_business_error_message(
        'missing_mora_a_aplicar|{"excel_row": 12, "field": "Mora a aplicar"}'
    )
    assert code == "missing_mora_a_aplicar"
    assert details["excel_row"] == 12
    assert details["field"] == "Mora a aplicar"


def test_job_error_storage_dict_finalize_amount_mismatch() -> None:
    err = job_error_storage_dict("finalize", ValueError("amount_mismatch"))
    assert err["error_code"] == "amount_mismatch"
    assert "no coinciden" in err["user_message"].lower()
    assert "Distribucion_Pagos" in err["next_action"]


def test_job_error_storage_dict_excel_write_guard() -> None:
    exc = ExcelWriteGuardError("empty_cliente", excel_row=7)
    err = job_error_storage_dict("finalize", exc)
    assert err["error_code"] == "empty_cliente"
    assert err["excel_row"] == 7
    assert err["user_message"]


def test_resolve_business_error_review_has_open_errors() -> None:
    payload = resolve_business_error_for_job(
        "finalize",
        message="review_has_open_errors|2",
        exc_type="ValueError",
    )
    assert payload["error_code"] == "review_has_open_errors"
    assert "errores" in payload["user_message"].casefold()


def _minimal_review_workbook(
    *,
    procesar: str = "NO",
    errores_rows: int = 1,
) -> bytes:
    wb = openpyxl.Workbook()
    ws_ctrl = wb.active
    ws_ctrl.title = ReviewSheets.CONTROL
    ws_ctrl.append([ControlCols.ROW_PROCESAR, procesar])
    ws_ctrl.append([ControlCols.ROW_ESTADO, "EN_REVISION"])

    ws_dist = wb.create_sheet(ReviewSheets.DISTRIBUCION_PAGOS)
    ws_dist.append([DistribucionCols.ID_PAGO, DistribucionCols.VALIDAR_PAGO, DistribucionCols.ESTADO_PAGO])

    ws_err = wb.create_sheet(ReviewSheets.ERRORES)
    ws_err.append([ErroresCols.ID_PAGO, ErroresCols.DESCRIPCION, ErroresCols.CODIGO_TECNICO])
    for i in range(errores_rows):
        ws_err.append([f"P{i+1}", "Falta extracto", "extract_not_found"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_preflight_flags_procesar_and_errores() -> None:
    raw = _minimal_review_workbook(procesar="NO", errores_rows=2)
    issues = collect_review_workbook_preflight_issues(raw)
    codes = {item["error_code"] for item in issues}
    assert "process_not_approved" in codes
    assert "review_has_open_errors" in codes
    assert count_open_errores_rows(openpyxl.load_workbook(io.BytesIO(raw), data_only=True)) == 2


def test_preflight_missing_distribucion_fields() -> None:
    wb = openpyxl.Workbook()
    ws_ctrl = wb.active
    ws_ctrl.title = ReviewSheets.CONTROL
    ws_ctrl.append([ControlCols.ROW_PROCESAR, ControlCols.VAL_PROCESAR_SI])
    ws_ctrl.append([ControlCols.ROW_ESTADO, "EN_REVISION"])
    ws_dist = wb.create_sheet(ReviewSheets.DISTRIBUCION_PAGOS)
    ws_dist.append(
        [
            DistribucionCols.ID_PAGO,
            DistribucionCols.VALIDAR_PAGO,
            DistribucionCols.ESTADO_PAGO,
            DistribucionCols.APLICAR_A_EXTRACTO,
            DistribucionCols.MORA_A_APLICAR,
            DistribucionCols.ABONO_A_CAPITAL,
            DistribucionCols.OTROS_VALORES,
            DistribucionCols.RUTA_UNIDAD_CREDITO,
        ]
    )
    ws_dist.append(["P1", ValidarPago.SI, "NORMAL", None, None, None, None, "ruta/credito"])
    buf = io.BytesIO()
    wb.save(buf)
    issues = collect_review_workbook_preflight_issues(buf.getvalue())
    codes = {item["error_code"] for item in issues}
    assert "missing_valor_intereses" in codes or "missing_mora_a_aplicar" in codes
