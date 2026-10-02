"""Amortización colgada en APLICANDO_AMORTIZACION tras un reinicio: salida segura."""

from __future__ import annotations

import asyncio
import io
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest

from app.application.job_manager import JobManager
from app.application.services.amortization_stale import (
    STALE_APPLYING_AFTER,
    applying_is_stale,
)
from app.application.ui.process_control_capabilities import compute_soft_close_availability
from app.application.ui.process_projection import (
    derive_operational_guidance,
    derive_operational_status,
    derive_steps_from_control,
)
from app.application.use_cases.payment_validation_process_control import (
    snapshot_from_control_cells,
)
from app.application.use_cases.payment_validation_soft_close import (
    SOFT_CLOSE_PARTIAL_ESTADO,
    soft_close_had_amortization_writes,
    soft_close_payment_validation,
)
from app.application.use_cases.setup_merge_control_workbook import (
    PROCESS_CONTROL_BANK_FILE_BOGOTA,
    SHEET_NAME,
)
from tests.test_soft_close_process_control import (
    MockGraphClientSoftClose,
    _control_late_phase,
    _read_control_row,
    _set_env,
)
from tests.ui_fixtures import make_snap

PK = "payment-validation|banco_bogota|2026-06-01|abc"
NOW = datetime(2026, 6, 1, 15, 0, tzinfo=timezone.utc)


class FakeJobs:
    def __init__(self, *, active: bool = False, latest: dict | None = None) -> None:
        self.active = active
        self.latest = latest

    def is_amortization_active(self) -> bool:
        return self.active

    def find_latest_job_by_process_and_types(self, process_key, types, *, statuses=None):
        return self.latest


def _applying(**overrides):
    base = dict(
        estado_proceso="APLICANDO_AMORTIZACION",
        process_key=PK,
        last_updated_at_proceso=(NOW - timedelta(minutes=2)).isoformat(),
    )
    base.update(overrides)
    return make_snap(**base)


# ─── Detección ──────────────────────────────────────────────────────────────


def test_otros_estados_nunca_son_colgados():
    for estado in ("CONSOLIDADO", "AMORTIZACION_PARCIAL", "ERROR_APPLY"):
        snap = make_snap(estado_proceso=estado, last_updated_at_proceso="")
        assert applying_is_stale(snap, job_manager=FakeJobs(), now=NOW) is False


def test_apply_vivo_en_este_servidor_no_es_colgado():
    snap = _applying(last_updated_at_proceso="")
    assert applying_is_stale(snap, job_manager=FakeJobs(active=True), now=NOW) is False
    running = FakeJobs(latest={"status": "running"})
    assert applying_is_stale(snap, job_manager=running, now=NOW) is False


def test_job_interrumpido_por_reinicio_es_colgado_de_inmediato():
    interrupted = FakeJobs(
        latest={"status": "failed", "error": {"type": "JobInterruptedByProcessRestart"}}
    )
    assert applying_is_stale(_applying(), job_manager=interrupted, now=NOW) is True


def test_sin_job_espera_el_plazo_antes_de_ofrecer_salida():
    recent = _applying()
    assert applying_is_stale(recent, job_manager=FakeJobs(), now=NOW) is False
    old = _applying(
        last_updated_at_proceso=(NOW - STALE_APPLYING_AFTER - timedelta(seconds=1)).isoformat()
    )
    assert applying_is_stale(old, job_manager=FakeJobs(), now=NOW) is True


def test_sin_marca_de_tiempo_y_sin_job_vivo_es_colgado():
    snap = _applying(last_updated_at_proceso="")
    assert applying_is_stale(snap, job_manager=FakeJobs(), now=NOW) is True


def test_marca_de_tiempo_con_offset_colombia_y_sin_zona():
    colombia = timezone(timedelta(hours=-5))
    recent_local = (NOW - timedelta(minutes=1)).astimezone(colombia)
    snap = _applying(last_updated_at_proceso=recent_local.isoformat())
    assert applying_is_stale(snap, job_manager=FakeJobs(), now=NOW) is False
    naive = recent_local.replace(tzinfo=None).isoformat()
    snap = _applying(last_updated_at_proceso=naive)
    assert applying_is_stale(snap, job_manager=FakeJobs(), now=NOW) is False


def test_control_lee_last_updated_at_proceso():
    snap = snapshot_from_control_cells(
        {"EstadoProceso": "APLICANDO_AMORTIZACION", "LastUpdatedAtProceso": "2026-06-01T10:00:00-05:00"},
        control_file_path="c.xlsx",
    )
    assert snap.last_updated_at_proceso == "2026-06-01T10:00:00-05:00"
    snap = snapshot_from_control_cells(
        {"LastUpdatedAtProceso": datetime(2026, 6, 1, 10, 0)}, control_file_path="c.xlsx"
    )
    assert snap.last_updated_at_proceso == "2026-06-01T10:00:00"


# ─── Proyección: estado, mensaje y salidas ──────────────────────────────────


def test_colgado_se_muestra_como_interrumpido_con_mensaje_claro():
    snap = _applying(merge_manifest_path="m.json", merge_idempotency_key=PK)
    steps = derive_steps_from_control(snap, applying_stale=True)
    apply = next(s for s in steps if s.name == "apply")
    assert apply.status == "failed_retryable"
    assert apply.can_retry is True
    status = derive_operational_status(snap, steps)
    assert status == "ERROR_RECUPERABLE"
    title, message, ref = derive_operational_guidance(
        status, control_estado="APLICANDO_AMORTIZACION"
    )
    assert title == "Amortización interrumpida"
    assert "Vuelva a procesar la amortización" in message
    assert "Cerrar con amortización incompleta" in message
    assert ref is None


def test_aplicando_vivo_sigue_como_aplicando():
    snap = _applying(merge_manifest_path="m.json", merge_idempotency_key=PK)
    steps = derive_steps_from_control(snap, applying_stale=False)
    assert derive_operational_status(snap, steps) == "APLICANDO"


def test_cerrar_disponible_solo_si_esta_colgado():
    stale = compute_soft_close_availability(
        write_allowed=True,
        mutation_active=False,
        control_estado="APLICANDO_AMORTIZACION",
        applying_stale=True,
    )
    assert stale.allowed is True
    live = compute_soft_close_availability(
        write_allowed=True,
        mutation_active=False,
        control_estado="APLICANDO_AMORTIZACION",
    )
    assert live.allowed is False
    busy = compute_soft_close_availability(
        write_allowed=True,
        mutation_active=True,
        control_estado="APLICANDO_AMORTIZACION",
        applying_stale=True,
    )
    assert busy.allowed is False


def test_cierre_tras_interrupcion_cuenta_como_amortizacion_incompleta():
    assert soft_close_had_amortization_writes(
        estado="APLICANDO_AMORTIZACION", process_key=PK, apply_idempotency_key=""
    ) is True


# ─── Cerrar con amortización incompleta desde APLICANDO colgado ─────────────


@pytest.fixture
def fresh_jobs(tmp_path, monkeypatch):
    monkeypatch.setenv("PAYMENT_VALIDATION_JOBS_DIR", str(tmp_path))
    JobManager._instance = None
    yield
    JobManager._instance = None


def _control_applying(updated_at: str) -> bytes:
    raw = _control_late_phase(estado="APLICANDO_AMORTIZACION", process_key=PK)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        ws = wb[SHEET_NAME]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(headers)}
        ws.cell(2, col["LastUpdatedAtProceso"], value=updated_at)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()


def test_soft_close_desde_aplicando_colgado_cierra_incompleto(fresh_jobs):
    _set_env()
    client = MockGraphClientSoftClose()
    old = (datetime.now(timezone.utc) - timedelta(hours=2)).isoformat()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_applying(old)
    res = asyncio.run(
        soft_close_payment_validation(client, bank_code="banco_bogota", process_key=PK, reason="")
    )
    assert res["process_control_estado"] == SOFT_CLOSE_PARTIAL_ESTADO
    row = _read_control_row(client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA])
    assert row["EstadoProceso"] == SOFT_CLOSE_PARTIAL_ESTADO


def test_soft_close_rechaza_aplicando_reciente(fresh_jobs):
    _set_env()
    client = MockGraphClientSoftClose()
    recent = datetime.now(timezone.utc).isoformat()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _control_applying(recent)
    with pytest.raises(ValueError, match="soft_close_not_allowed\\|APLICANDO_AMORTIZACION"):
        asyncio.run(
            soft_close_payment_validation(
                client, bank_code="banco_bogota", process_key=PK, reason=""
            )
        )
