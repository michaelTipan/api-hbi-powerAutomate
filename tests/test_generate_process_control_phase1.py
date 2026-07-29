"""Generate (Phase 1): bank_code + control por banco + idempotencia inicial."""

import asyncio
import io
import os
from datetime import date
from urllib.parse import unquote

import httpx
import openpyxl
import pytest

from app.application.use_cases.payment_validation_generate import generate_payment_validation
from app.application.use_cases.setup_merge_control_workbook import (
    PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA,
    PROCESS_CONTROL_BANK_FILE_BOGOTA,
    _build_process_control_workbook_bytes,
)


def _http_error(status: int) -> httpx.HTTPStatusError:
    req = httpx.Request("GET", "https://graph.microsoft.com/test")
    resp = httpx.Response(status, request=req, text=f"HTTP {status} body")
    return httpx.HTTPStatusError("err", request=req, response=resp)


class MockGraphClientProcessControl:
    def __init__(self) -> None:
        self.children = []
        self.downloaded_files: dict[str, bytes] = {}
        self.folder_children: dict[str, list[dict]] = {}
        self.requested_endpoints: list[str] = []
        self.put_calls: list[tuple[str, bytes]] = []
        # Rutas relativas que responden 404 en GET de item (existencia de archivo).
        self.force_404_paths: set[str] = set()

    async def get(self, endpoint: str, params=None):
        self.requested_endpoints.append(endpoint)
        if endpoint == "/sites":
            return {"value": [{"id": "dummy_site"}]}
        if endpoint == "/sites/dummy_site/drives":
            return {"value": [{"id": "dummy_drive", "name": "DRIVE"}]}
        if endpoint.endswith(":/children"):
            path = endpoint.split("/root:/", 1)[1].rsplit(":/children", 1)[0]
            path = unquote(path)
            if path == "revision":
                return {"value": self.children}
            return {"value": self.folder_children.get(path, [])}
        if "/root:/" in endpoint and ":/children" not in endpoint and ":/content" not in endpoint:
            path = unquote(endpoint.split("/root:/", 1)[1].split(":", 1)[0]).strip("/")
            if path in self.force_404_paths:
                raise _http_error(404)
            return {"id": "item1", "name": path.rsplit("/", 1)[-1], "webUrl": "https://example/item.xlsx"}
        return {"value": []}

    async def get_bytes(self, endpoint: str, params=None):
        self.requested_endpoints.append(endpoint)
        if not endpoint.endswith(":/content"):
            raise _http_error(404)
        path = unquote(endpoint.split("/root:/", 1)[1].rsplit(":/content", 1)[0]).strip("/")
        if path in self.downloaded_files:
            return self.downloaded_files[path]
        raise _http_error(404)

    async def put_bytes(self, endpoint: str, content: bytes, content_type: str = ""):
        self.put_calls.append((endpoint, content))
        return {"webUrl": "https://example/upload.xlsx"}


def _set_env():
    os.environ["GRAPH_SHAREPOINT_SITE_SEARCH"] = "SITIO"
    os.environ["GRAPH_SHAREPOINT_DRIVE_NAME"] = "DRIVE"
    os.environ["GRAPH_PAYMENT_VALIDATION_REVIEW_PATH"] = "revision"
    os.environ["GRAPH_CLIENTS_BASE_PATH"] = "clientes"
    os.environ["GRAPH_VALIDATION_FILE_PREFIX"] = "val"
    os.environ["GRAPH_BANK_PAYMENTS_FILE_PATH"] = "banco.xlsx"
    os.environ["GRAPH_BANK_PAYMENTS_FILE_PATH_BANCOLOMBIA"] = "banco.xlsx"


def _minimal_bank_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Fecha", "Crédito", "Concepto", "Tipo Aplicación", "Transacción"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_generate_invalid_bank_code_fails():
    _set_env()
    client = MockGraphClientProcessControl()
    client.downloaded_files["banco.xlsx"] = _minimal_bank_xlsx()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _build_process_control_workbook_bytes(
        "banco_bogota", "Banco de Bogotá"
    )
    with pytest.raises(ValueError, match="invalid_bank_code"):
        asyncio.run(generate_payment_validation(client, date(2026, 6, 1), bank_code="otro_banco"))


def test_generate_without_bank_code_raises_required():
    _set_env()
    client = MockGraphClientProcessControl()
    with pytest.raises(ValueError, match="bank_code_required"):
        asyncio.run(generate_payment_validation(client, date(2026, 6, 1), bank_code=""))


def test_generate_with_bank_code_writes_revision_creada():
    _set_env()
    client = MockGraphClientProcessControl()
    client.children = []
    client.folder_children["clientes"] = []
    client.downloaded_files["banco.xlsx"] = _minimal_bank_xlsx()
    client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _build_process_control_workbook_bytes(
        "banco_bogota", "Banco de Bogotá"
    )

    res = asyncio.run(generate_payment_validation(client, date(2026, 6, 1), bank_code="banco_bogota"))
    assert res["bank_code"] == "banco_bogota"
    assert res["process_control_updated"] is True
    assert res["process_control_estado"] == "REVISION_CREADA"
    assert res["validation_file"].startswith("val_banco_bogota_2026-06-01")


def test_generate_idempotency_reuses_when_same_process_key_and_validation_path_present():
    _set_env()
    client = MockGraphClientProcessControl()
    client.children = []
    client.folder_children["clientes"] = []
    client.downloaded_files["banco.xlsx"] = _minimal_bank_xlsx()
    # Build control with ProcessKey + ValidationFilePath already set
    raw = _build_process_control_workbook_bytes("banco_bogota", "Banco de Bogotá")
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        ws = wb["Procesos"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(headers)}
        ws.cell(2, col["ProcessKey"], value="payment-validation|banco_bogota|2026-06-01")
        ws.cell(2, col["EstadoProceso"], value="REVISION_CREADA")
        ws.cell(2, col["IsActive"], value=True)
        ws.cell(2, col["ValidationFilePath"], value="revision/existing.xlsx")
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = buf.getvalue()
    finally:
        wb.close()

    res = asyncio.run(generate_payment_validation(client, date(2026, 6, 1), bank_code="banco_bogota"))
    assert res["already_generated"] is True
    assert res["file_action"] == "reused"
    assert res["validation_file_path"] == "revision/existing.xlsx"


def test_generate_recreates_when_registered_review_file_missing():
    """Si el Excel registrado ya no está y el lote sigue en REVISION_CREADA, se recrea."""
    _set_env()
    client = MockGraphClientProcessControl()
    client.children = []
    client.folder_children["clientes"] = []
    client.downloaded_files["banco.xlsx"] = _minimal_bank_xlsx()
    client.force_404_paths.add("revision/gone.xlsx")
    raw = _build_process_control_workbook_bytes("banco_bogota", "Banco de Bogotá")
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        ws = wb["Procesos"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(headers)}
        ws.cell(2, col["ProcessKey"], value="payment-validation|banco_bogota|2026-06-01|abc")
        ws.cell(2, col["EstadoProceso"], value="REVISION_CREADA")
        ws.cell(2, col["IsActive"], value=True)
        ws.cell(2, col["ValidationFilePath"], value="revision/gone.xlsx")
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = buf.getvalue()
    finally:
        wb.close()

    res = asyncio.run(generate_payment_validation(client, date(2026, 6, 1), bank_code="banco_bogota"))
    assert res["already_generated"] is False
    assert res["file_action"] == "recreated"
    assert res["process_control_estado"] == "REVISION_CREADA"
    assert res["validation_file"].startswith("val_banco_bogota_2026-06-01")
    assert res["validation_file_path"] != "revision/gone.xlsx"
    assert client.put_calls, "debió subir un nuevo Excel de revisión"


def test_generate_blocks_when_active_process_exists():
    _set_env()
    client = MockGraphClientProcessControl()
    client.children = []
    client.folder_children["clientes"] = []
    client.downloaded_files["banco.xlsx"] = _minimal_bank_xlsx()
    raw = _build_process_control_workbook_bytes("banco_bogota", "Banco de Bogotá")
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    try:
        ws = wb["Procesos"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(headers)}
        ws.cell(2, col["ProcessKey"], value="payment-validation|banco_bogota|2026-05-31")
        ws.cell(2, col["EstadoProceso"], value="REVISION_CREADA")
        ws.cell(2, col["IsActive"], value=True)
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = buf.getvalue()
    finally:
        wb.close()

    with pytest.raises(ValueError, match="active_process_exists\\|payment-validation\\|banco_bogota\\|2026-05-31"):
        asyncio.run(generate_payment_validation(client, date(2026, 6, 1), bank_code="banco_bogota"))

