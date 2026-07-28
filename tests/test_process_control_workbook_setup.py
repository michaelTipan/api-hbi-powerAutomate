"""Setup de controles de proceso por banco (control_proceso_validacion_pagos_*.xlsx)."""

import asyncio
import io

import httpx
import openpyxl
import pytest

from app.application.use_cases import setup_merge_control_workbook as smcw
from app.application.use_cases.setup_merge_control_workbook import (
    BANK_CODE_BANCOLOMBIA,
    BANK_CODE_BOGOTA,
    MERGE_CONTROL_COLUMNS,
    PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA,
    PROCESS_CONTROL_BANK_FILE_BOGOTA,
    PROCESS_CONTROL_COLUMNS,
    PROCESS_CONTROL_EXTENSION_COLUMNS,
    PROCESS_CONTROL_TABLE_DISPLAY_NAME,
    SHEET_NAME,
    _build_workbook_with_base_columns_only,
    _build_process_control_workbook_bytes,
    build_payment_validation_process_key,
    setup_merge_control_workbook,
)


def _http_error(status: int) -> httpx.HTTPStatusError:
    req = httpx.Request("GET", "https://graph.microsoft.com/test")
    resp = httpx.Response(status, request=req, text=f"HTTP {status} body")
    return httpx.HTTPStatusError("err", request=req, response=resp)


class MultiFileMockGraph:
    def __init__(self) -> None:
        self.files: dict[str, bytes] = {}
        self.put_count = 0
        self.post_json_count = 0
        self.web_url = "https://comwareec.sharepoint.com/sites/x/control.xlsx"
        self.put_response: dict = {"webUrl": self.web_url}

    def _path_from_endpoint(self, endpoint: str) -> str | None:
        low = endpoint.lower()
        if ":/content" not in low and not low.endswith(":"):
            return None
        for path in (
            PROCESS_CONTROL_BANK_FILE_BOGOTA,
            PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA,
        ):
            key = path.replace(" ", "%20").lower()
            if path.lower() in low or key.split("/")[-1] in low:
                return path
        for stored in self.files:
            if stored.lower() in low or stored.split("/")[-1].lower() in low:
                return stored
        return None

    async def get(self, endpoint: str, params=None):
        return {"value": []}

    async def get_bytes(self, endpoint: str, params=None):
        path = self._path_from_endpoint(endpoint)
        if path and path in self.files:
            return self.files[path]
        if path:
            raise _http_error(404)
        raise _http_error(404)

    async def put_bytes(self, endpoint: str, content: bytes, content_type: str = ""):
        path = self._path_from_endpoint(endpoint)
        if not path:
            for known in (
                PROCESS_CONTROL_BANK_FILE_BOGOTA,
                PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA,
            ):
                if known.split("/")[-1].lower() in endpoint.lower():
                    path = known
                    break
        if not path:
            raise AssertionError(f"Unexpected write endpoint: {endpoint!r}")
        self.put_count += 1
        self.files[path] = content
        return dict(self.put_response)

    async def post_json(self, endpoint: str, body: dict):
        self.post_json_count += 1
        return {}, 201


@pytest.fixture
def mock_resolve(monkeypatch):
    async def _fake_resolve(client, site_search, drive_name, path):
        return {
            "site_id": "site-1",
            "drive_id": "drive-1",
            "path_encoded": "enc",
            "file_path": path,
        }

    monkeypatch.setattr(smcw, "resolve_sharepoint_path", _fake_resolve)


@pytest.fixture(autouse=True)
def env_site(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST_SITE")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")


def _bank_result(out: dict, bank_code: str) -> dict:
    for b in out["banks"]:
        if b["bank_code"] == bank_code:
            return b
    raise AssertionError(f"bank {bank_code} not in {out['banks']}")


def test_build_payment_validation_process_key_examples():
    assert (
        build_payment_validation_process_key("banco_bogota", "2026-06-01")
        == "payment-validation|banco_bogota|2026-06-01"
    )
    assert (
        build_payment_validation_process_key("banco_bancolombia", "2026-06-01")
        == "payment-validation|banco_bancolombia|2026-06-01"
    )
    assert (
        build_payment_validation_process_key(
            "banco_bogota", "2026-06-01", "4df53868-eeb1-428f-9c92-98e0efcad7ec"
        )
        == "payment-validation|banco_bogota|2026-06-01|4df53868-eeb1-428f-9c92-98e0efcad7ec"
    )


def test_build_process_artifact_filename_includes_process_id():
    from app.application.use_cases.setup_merge_control_workbook import (
        build_process_artifact_filename,
    )

    name = build_process_artifact_filename(
        kind="cartera_validada",
        bank_code="banco_bogota",
        process_date="2026-07-28",
        process_id="4df53868-eeb1-428f-9c92-98e0efcad7ec",
    )
    assert name == (
        "cartera_validada_banco_bogota_2026-07-28_4df53868-eeb1-428f-9c92-98e0efcad7ec.xlsx"
    )


def test_build_payment_validation_process_key_requires_inputs():
    with pytest.raises(ValueError, match="bank_code_and_process_date_required"):
        build_payment_validation_process_key("", "2026-06-01")


def test_setup_creates_bogota_control(mock_resolve):
    g = MultiFileMockGraph()
    out = asyncio.run(setup_merge_control_workbook(g))
    bogota = _bank_result(out, BANK_CODE_BOGOTA)
    assert bogota["created"] is True
    assert bogota["status"] == "success"
    assert PROCESS_CONTROL_BANK_FILE_BOGOTA in g.files


def test_setup_creates_bancolombia_control(mock_resolve):
    g = MultiFileMockGraph()
    out = asyncio.run(setup_merge_control_workbook(g))
    bancol = _bank_result(out, BANK_CODE_BANCOLOMBIA)
    assert bancol["created"] is True
    assert bancol["status"] == "success"
    assert PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA in g.files


def test_setup_does_not_overwrite_existing_bank_controls(mock_resolve):
    g = MultiFileMockGraph()
    asyncio.run(setup_merge_control_workbook(g))
    first_puts = g.put_count
    out2 = asyncio.run(setup_merge_control_workbook(g))
    assert _bank_result(out2, BANK_CODE_BOGOTA)["created"] is False
    assert _bank_result(out2, BANK_CODE_BANCOLOMBIA)["created"] is False
    assert g.put_count == first_puts


def test_process_control_workbook_has_all_columns(mock_resolve):
    g = MultiFileMockGraph()
    asyncio.run(setup_merge_control_workbook(g))
    wb = openpyxl.load_workbook(io.BytesIO(g.files[PROCESS_CONTROL_BANK_FILE_BOGOTA]), data_only=True)
    try:
        ws = wb[SHEET_NAME]
        headers = [ws.cell(1, c).value for c in range(1, len(PROCESS_CONTROL_COLUMNS) + 1)]
        assert headers == list(PROCESS_CONTROL_COLUMNS)
        for legacy_name in MERGE_CONTROL_COLUMNS:
            assert legacy_name in headers
        for ext_name in PROCESS_CONTROL_EXTENSION_COLUMNS:
            assert ext_name in headers
        assert PROCESS_CONTROL_TABLE_DISPLAY_NAME in ws.tables
    finally:
        wb.close()


def test_process_control_initial_bank_values_bogota(mock_resolve):
    g = MultiFileMockGraph()
    asyncio.run(setup_merge_control_workbook(g))
    wb = openpyxl.load_workbook(io.BytesIO(g.files[PROCESS_CONTROL_BANK_FILE_BOGOTA]), data_only=True)
    try:
        ws = wb[SHEET_NAME]
        col_map = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        assert ws.cell(2, col_map["BankCode"]).value == BANK_CODE_BOGOTA
        assert ws.cell(2, col_map["BankName"]).value == "Banco de Bogotá"
        assert ws.cell(2, col_map["EstadoProceso"]).value == "VACIO"
        assert str(ws.cell(2, col_map["IsActive"]).value).lower() == "false"
    finally:
        wb.close()


def test_process_control_initial_bank_values_bancolombia(mock_resolve):
    g = MultiFileMockGraph()
    asyncio.run(setup_merge_control_workbook(g))
    wb = openpyxl.load_workbook(io.BytesIO(g.files[PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA]), data_only=True)
    try:
        ws = wb[SHEET_NAME]
        col_map = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        assert ws.cell(2, col_map["BankCode"]).value == BANK_CODE_BANCOLOMBIA
        assert ws.cell(2, col_map["BankName"]).value == "Bancolombia"
        assert ws.cell(2, col_map["EstadoProceso"]).value == "VACIO"
    finally:
        wb.close()


def test_setup_adds_missing_extension_columns_to_legacy_twelve_col_workbook(mock_resolve):
    g = MultiFileMockGraph()
    g.files[PROCESS_CONTROL_BANK_FILE_BOGOTA] = _build_workbook_with_base_columns_only()
    g.put_count = 0
    out = asyncio.run(setup_merge_control_workbook(g))
    bogota = _bank_result(out, BANK_CODE_BOGOTA)
    assert bogota["created"] is False
    assert bogota["repaired"] is True
    assert set(PROCESS_CONTROL_EXTENSION_COLUMNS).issubset(set(bogota["columns_added"]))
    wb = openpyxl.load_workbook(io.BytesIO(g.files[PROCESS_CONTROL_BANK_FILE_BOGOTA]), data_only=True)
    try:
        ws = wb[SHEET_NAME]
        headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
        assert headers[: len(MERGE_CONTROL_COLUMNS)] == list(MERGE_CONTROL_COLUMNS)
        for name in PROCESS_CONTROL_EXTENSION_COLUMNS:
            assert name in headers
    finally:
        wb.close()


def test_setup_endpoint_response_includes_banks_summary(mock_resolve):
    g = MultiFileMockGraph()
    out = asyncio.run(setup_merge_control_workbook(g))
    assert out["status"] == "success"
    assert len(out["banks"]) == 2
    assert out["process_control_columns"] == list(PROCESS_CONTROL_COLUMNS)


def test_build_process_control_bytes_standalone():
    raw = _build_process_control_workbook_bytes(BANK_CODE_BOGOTA, "Banco de Bogotá")
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    try:
        ws = wb[SHEET_NAME]
        assert list(ws.tables) == [PROCESS_CONTROL_TABLE_DISPLAY_NAME]
    finally:
        wb.close()


def test_setup_only_creates_official_bank_controls(mock_resolve):
    g = MultiFileMockGraph()
    out = asyncio.run(setup_merge_control_workbook(g))
    assert out["status"] == "success"
    assert set(g.files) == {
        PROCESS_CONTROL_BANK_FILE_BOGOTA,
        PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA,
    }
