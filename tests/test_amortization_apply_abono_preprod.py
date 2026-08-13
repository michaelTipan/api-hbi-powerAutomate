"""Endurecimiento preproducción ABONO: retry parcial híbrido y sección llena."""

import asyncio
import io
import json
from datetime import date

import openpyxl
import pytest

from app.application.services.abono_dry_run import (
    ABONO_APPLICATION_STATUS_ALREADY_APPLIED,
    APPLICATION_PAYMENT_SECTION_FULL,
)
from app.application.services.amortization_workbook import (
    AUTOMATION_LOG_SHEET,
    APLICADO,
    append_automation_log,
    build_amortization_idempotency_key,
    ensure_automation_log,
)
from app.application.use_cases import amortization_fill_apply as apply_mod
from app.application.use_cases.amortization_fill_apply import (
    VERIFICATION_FAILED,
    run_amortization_fill_apply,
)
from app.application.use_cases.amortization_fill_dry_run import run_amortization_fill_dry_run
from tests.test_abono_dry_run import _abono_dry_run_files, _abono_manifest_output, _asiento_pdf
from tests.test_amortization_fill_apply import MockGraphApply
from tests.test_amortization_fill_dry_run import (
    MockGraphDryRun,
    _accounting_text,
    _amort_table_date_at_row,
    _hist_bytes_multi_credit,
    _ibr_bytes,
)


@pytest.fixture(autouse=True)
def env_sharepoint(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")


@pytest.fixture(autouse=True)
def stub_process_control_writes(monkeypatch):
    async def _ok_update(*_a, **_k):
        return None

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.update_process_control_row2",
        _ok_update,
    )


def _patch_pdf_extract(monkeypatch):
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    monkeypatch.setattr(
        "app.application.services.abono_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )


def _two_credit_abono_files(fecha: date) -> tuple[dict, dict[str, bytes]]:
    out = _abono_manifest_output(
        id_pago="AB-PARTIAL",
        monto_banco=100_000_000.0,
        credits=("258", "265"),
    )
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    hist = _hist_bytes_multi_credit(
        [
            ("AB-PARTIAL", "EQUINORTE", "CREDITO # 258", "TABLAS/amort_258.xlsx", fecha),
            ("AB-PARTIAL", "EQUINORTE", "CREDITO # 265", "TABLAS/amort_265.xlsx", fecha),
        ]
    )
    files["HIST/cartera.xlsx"] = hist
    files["TABLAS/amort_258.xlsx"] = _amort_table_date_at_row(fecha, 8)
    files["TABLAS/amort_265.xlsx"] = _amort_table_date_at_row(fecha, 8)
    return out, files


def _log_rows_for_table(blob: bytes) -> int:
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=False)
    if AUTOMATION_LOG_SHEET not in wb.sheetnames:
        return 0
    return max(0, (wb[AUTOMATION_LOG_SHEET].max_row or 1) - 1)


def test_abono_partial_retry_recognizes_already_applied_from_log(monkeypatch):
    """Retry tras AMORTIZACION_PARCIAL: crédito aplicado vía log, pendiente vía PDF."""
    fecha = date(2026, 5, 22)
    out, files = _two_credit_abono_files(fecha)
    asiento_258 = out["credit_items"][0]["asiento_pdf_paths"][0]
    asiento_265 = out["credit_items"][1]["asiento_pdf_paths"][0]
    g = MockGraphApply(files)
    _patch_pdf_extract(monkeypatch)

    fail_265_once = {"done": False}
    original_apply_one = apply_mod._apply_one_table

    async def flaky_apply_one_table(graph, site_id, drive_id, tabla_path, *args, **kwargs):
        if "amort_265" in tabla_path and not fail_265_once["done"]:
            fail_265_once["done"] = True
            return {
                "items": [],
                "uploaded": False,
                "tabla_path": tabla_path,
                "upload_status": "failed",
                "verification_status": VERIFICATION_FAILED,
            }
        return await original_apply_one(
            graph, site_id, drive_id, tabla_path, *args, **kwargs
        )

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply._apply_one_table",
        flaky_apply_one_table,
    )

    first = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert first["status"] == "partial"
    assert first["process_control_estado"] == "AMORTIZACION_PARCIAL"
    assert g.uploaded.get("TABLAS/amort_258.xlsx")
    assert "TABLAS/amort_265.xlsx" not in g.uploaded
    log_258_first = _log_rows_for_table(g.uploaded["TABLAS/amort_258.xlsx"])
    assert log_258_first == 1
    assert asiento_258 not in g.files

    second = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert second["status"] == "ok"
    assert second["process_control_estado"] == "AMORTIZACION_APLICADA"
    assert g.uploaded.get("TABLAS/amort_265.xlsx")
    log_258_second = _log_rows_for_table(g.uploaded["TABLAS/amort_258.xlsx"])
    assert log_258_second == log_258_first

    abono_items = [i for i in second["items"] if i.get("tipo_aplicacion") == "ABONO"]
    already = [i for i in abono_items if i.get("application_status") == ABONO_APPLICATION_STATUS_ALREADY_APPLIED]
    applied = [i for i in abono_items if i.get("apply_status") == "APPLIED"]
    assert len(already) >= 1 or any(
        i.get("apply_status") == "SKIPPED_IDEMPOTENT" for i in abono_items if "258" in str(i.get("credito"))
    )
    assert any("265" in str(i.get("credito")) for i in applied)
    assert asiento_265 not in g.files or any(
        "PROCESADOS" in p for p in g.files if "265" in p
    )


def test_abono_hybrid_reconcile_blocks_when_amounts_dont_match(monkeypatch):
    fecha = date(2026, 5, 22)
    out = _abono_manifest_output(monto_banco=100_000_000.0, credits=("258",))
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    tabla = "TABLAS/amort_258.xlsx"
    wb = openpyxl.load_workbook(io.BytesIO(files[tabla]), data_only=False)
    ensure_automation_log(wb)
    asiento = out["credit_items"][0]["asiento_pdf_paths"][0]
    append_automation_log(
        wb,
        {
            "id_pago": out["id_pago"],
            "credito": "258",
            "application_row": 8,
            "accion": APLICADO,
            "idempotency_key": build_amortization_idempotency_key(
                out["id_pago"], "258", asiento, ""
            ),
            "asiento_pdf_path": asiento,
            "valor_pagado_cliente": 30_000_000.0,
            "tipo_aplicacion": "ABONO",
        },
    )
    buf = io.BytesIO()
    wb.save(buf)
    files[tabla] = buf.getvalue()
    del files[asiento]

    g = MockGraphDryRun(files)
    _patch_pdf_extract(monkeypatch)
    dry = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert dry["can_apply"] is False
    gr = dry["abono_group_results"][0]
    assert gr["reconciliation_status"] == "FAILED"


def test_abono_applied_without_resolvable_amount_blocks(monkeypatch):
    fecha = date(2026, 5, 22)
    out = _abono_manifest_output(monto_banco=50_000_000.0, credits=("258",))
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    tabla = "TABLAS/amort_258.xlsx"
    wb = openpyxl.load_workbook(io.BytesIO(files[tabla]), data_only=False)
    ensure_automation_log(wb)
    asiento = out["credit_items"][0]["asiento_pdf_paths"][0]
    append_automation_log(
        wb,
        {
            "id_pago": out["id_pago"],
            "credito": "258",
            "application_row": 99,
            "accion": APLICADO,
            "idempotency_key": build_amortization_idempotency_key(
                out["id_pago"], "258", asiento, ""
            ),
            "asiento_pdf_path": asiento,
            "tipo_aplicacion": "ABONO",
        },
    )
    buf = io.BytesIO()
    wb.save(buf)
    files[tabla] = buf.getvalue()
    del files[asiento]

    g = MockGraphDryRun(files)
    _patch_pdf_extract(monkeypatch)
    dry = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert dry["can_apply"] is False
    assert any(
        i.get("error_code") == "APPLIED_ABONO_AMOUNT_UNRESOLVABLE"
        for i in dry["items"]
        if i.get("tipo_aplicacion") == "ABONO"
    )


def test_abono_dry_run_application_section_full(monkeypatch):
    from app.application.services.amortization_workbook import FindApplicationRowResult

    fecha = date(2026, 5, 22)
    out = _abono_manifest_output(monto_banco=50_000_000.0, credits=("258",))
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    files["TABLAS/amort_258.xlsx"] = _amort_table_date_at_row(fecha, 8)

    def _full_section(*_a, **_k):
        return FindApplicationRowResult(
            row=None,
            requires_new_row=True,
            suggested_row=9,
            search_start_row=8,
        )

    monkeypatch.setattr(
        "app.application.services.abono_dry_run.find_next_available_application_row",
        _full_section,
    )

    g = MockGraphDryRun(files)
    _patch_pdf_extract(monkeypatch)
    dry = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    abono_item = next(i for i in dry["items"] if i.get("tipo_aplicacion") == "ABONO")
    assert abono_item["error_code"] == APPLICATION_PAYMENT_SECTION_FULL
    assert abono_item["application_status"] == "ERROR"
    assert abono_item.get("application_section_full", {}).get("suggested_row") == 9
    assert dry["can_apply"] is False


def test_abono_apply_blocked_when_section_full(monkeypatch):
    from app.application.services.amortization_workbook import FindApplicationRowResult

    fecha = date(2026, 5, 22)
    out = _abono_manifest_output(monto_banco=50_000_000.0, credits=("258",))
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    files["TABLAS/amort_258.xlsx"] = _amort_table_date_at_row(fecha, 8)

    monkeypatch.setattr(
        "app.application.services.abono_dry_run.find_next_available_application_row",
        lambda *_a, **_k: FindApplicationRowResult(
            row=None, requires_new_row=True, suggested_row=9, search_start_row=8
        ),
    )

    g = MockGraphApply(files)
    _patch_pdf_extract(monkeypatch)
    result = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert result["status"] in ("failed", "blocked", "partial")
    assert result.get("apply_wrote_changes") is not True
    assert not g.put_calls
    assert not g.patch_calls


def test_abono_available_row_still_works_after_section_full_change(monkeypatch):
    fecha = date(2026, 5, 22)
    out = _abono_manifest_output(monto_banco=50_000_000.0, credits=("258",))
    files = _abono_dry_run_files(abono_output=out, fecha=fecha)
    files["TABLAS/amort_258.xlsx"] = _amort_table_date_at_row(fecha, 8)

    g = MockGraphDryRun(files)
    _patch_pdf_extract(monkeypatch)
    dry = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    abono_item = next(i for i in dry["items"] if i.get("tipo_aplicacion") == "ABONO")
    assert abono_item["application_status"] == "WOULD_APPLY"
    assert dry["can_apply"] is True


def test_pago_requires_application_row_unchanged(monkeypatch):
    """PAGO conserva REQUIRES_APPLICATION_ROW cuando la sección está llena."""
    fecha = date(2026, 4, 22)
    from tests.test_amortization_fill_dry_run import _asiento_pdf_placeholder, _hist_bytes

    hist = _hist_bytes("7785e37e", "CREDITO # 265", "TABLAS/amort_265.xlsx", fecha)
    asiento_a = "clientes/E/a1.pdf"
    asiento_b = "clientes/E/a2.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_paths": [asiento_a, asiento_b],
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_265.xlsx": _amort_table_date_at_row(fecha, 8),
        asiento_a: _asiento_pdf_placeholder(),
        asiento_b: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    second = out["items"][1]
    assert second["error_code"] == "REQUIRES_APPLICATION_ROW"
