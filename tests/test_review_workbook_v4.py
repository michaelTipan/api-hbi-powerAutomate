"""Workbook de revisión v4: 15 columnas, visual operador, sin distribución manual."""
from __future__ import annotations

from datetime import date
from io import BytesIO

import openpyxl

from app.application.services.review_schema import (
    REVIEW_SCHEMA_VERSION,
    AplicacionPagosCols,
    AplicacionPagosColsV3,
    ErroresCols,
    MANUAL_DISTRIBUTION_HEADERS_V3,
    MetaCols,
    ReviewSheets,
    ValidarPago,
    require_review_schema_v4,
)
from app.application.services.review_workbook_v4 import (
    REVIEW_FIRST_DATA_ROW,
    REVIEW_HEADER_ROW,
    build_aplicacion_pagos_row,
    build_review_workbook_v4_bytes,
)

_EXPECTED_V4 = [
    "ID Pago",
    "Cliente",
    "Crédito",
    "Monto banco",
    "Fecha banco",
    "Fecha límite",
    "Días respecto vencimiento",
    "Valor obligación actual",
    "Saldo vencido",
    "Validar Pago",
    "Tipo de aplicación",
    "Link extracto",
    "Link tabla amortización",
    "Link carpeta crédito",
    "Observación",
]


def _sample_row() -> dict:
    payment = {
        "id_pago": "PAY-1",
        "cliente": "CLIENTE SYN",
        "monto_banco": 1_500_000,
        "fecha_banco": date(2026, 5, 20),
    }
    candidate = {
        "credito": "258",
        "fecha_limite": date(2026, 5, 23),
        "valor_obligacion_actual": 1_500_000,
        "saldo_vencido_visible": None,
        "link_extracto": "https://example/extracto",
        "link_tabla": "https://example/tabla",
        "link_carpeta_credito": "https://example/carpeta",
        "ruta_extracto_pdf": "clientes/CLIENTE/CREDITO 258/EXTRACTOS/extracto.pdf",
        "ruta_unidad_credito": "clientes/CLIENTE/CREDITO 258",
        "ruta_tabla_amortizacion": "clientes/CLIENTE/CREDITO 258/tabla.xlsx",
        "credito_normalizado": "258",
        "right_panel_role": "VACIO",
        "parser_status": "OK",
        "extract_evidence": {},
    }
    return build_aplicacion_pagos_row(payment, candidate)


def _load_v4():
    raw = build_review_workbook_v4_bytes(
        process_id="proc-1",
        process_date=date(2026, 5, 20),
        bank_code="banco_bogota",
        aplicacion_rows=[_sample_row()],
        error_records=[],
    )
    wb = openpyxl.load_workbook(BytesIO(raw))
    return raw, wb


def test_review_schema_version_is_4():
    assert REVIEW_SCHEMA_VERSION == 4


def test_v4_has_exactly_15_visible_columns_in_order():
    assert list(AplicacionPagosCols.HEADERS) == _EXPECTED_V4
    assert len(AplicacionPagosCols.HEADERS) == 15
    assert "Aplicación sugerida" not in AplicacionPagosCols.HEADERS


def test_removed_manual_distribution_headers_not_in_v4():
    names = set(AplicacionPagosCols.HEADERS)
    assert names.isdisjoint(MANUAL_DISTRIBUTION_HEADERS_V3)
    assert MANUAL_DISTRIBUTION_HEADERS_V3.issubset(set(AplicacionPagosColsV3.HEADERS))
    assert len(AplicacionPagosColsV3.HEADERS) == 21


def test_workbook_v4_opens_with_openpyxl_and_layout():
    _raw, wb = _load_v4()
    ws = wb[ReviewSheets.APLICACION_PAGOS]
    assert ws.cell(1, 1).value == "APLICACIÓN DE PAGOS"
    assert ws.cell(2, 1).value
    headers = [ws.cell(REVIEW_HEADER_ROW, c).value for c in range(1, 16)]
    assert headers == _EXPECTED_V4
    assert ws.cell(REVIEW_FIRST_DATA_ROW, 1).value == "PAY-1"
    assert ws.freeze_panes == "D4"
    assert ws.auto_filter.ref
    assert ws.sheet_view.showGridLines is False
    assert str(ws.sheet_properties.tabColor.rgb).endswith("00B050")
    assert wb[ReviewSheets.META].sheet_state == "hidden"
    assert wb[ReviewSheets.LISTAS].sheet_state == "hidden"


def test_workbook_v4_does_not_contain_removed_headers_or_avk_formulas():
    _raw, wb = _load_v4()
    ws = wb[ReviewSheets.APLICACION_PAGOS]
    visible = [ws.cell(REVIEW_HEADER_ROW, c).value for c in range(1, 40)]
    visible = [v for v in visible if v]
    for banned in MANUAL_DISTRIBUTION_HEADERS_V3:
        assert banned not in visible
    data_row = REVIEW_FIRST_DATA_ROW
    for c in range(1, 16):
        val = ws.cell(data_row, c).value
        if isinstance(val, str):
            assert not val.startswith("=")
            assert "Aplicación sugerida" not in str(val)
            assert "POR DISTRIBUIR" not in val


def test_workbook_v4_editables_and_observacion_optional():
    _raw, wb = _load_v4()
    ws = wb[ReviewSheets.APLICACION_PAGOS]
    r = REVIEW_FIRST_DATA_ROW
    col_vp = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.VALIDAR_PAGO) + 1
    col_tipo = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.TIPO_APLICACION) + 1
    col_obs = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.OBSERVACION) + 1
    col_id = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.ID_PAGO) + 1
    assert ws.cell(r, col_vp).protection.locked is False
    assert ws.cell(r, col_tipo).protection.locked is False
    assert ws.cell(r, col_obs).protection.locked is False
    assert ws.cell(r, col_id).protection.locked is True
    assert ws.cell(r, col_vp).value == ValidarPago.NO


def test_workbook_v4_meta_version_and_finalize_accepts():
    _raw, wb = _load_v4()
    meta = {row[0]: row[1] for row in wb[ReviewSheets.META].iter_rows(values_only=True) if row}
    assert meta.get(MetaCols.ROW_REVIEW_SCHEMA_VERSION) == 4
    assert require_review_schema_v4(wb) == 4


def test_workbook_v4_friendly_links_client_separator_and_no_control_copy():
    payment_a = {
        "id_pago": "PAY-A",
        "cliente": "EQUINORTE",
        "monto_banco": 100,
        "fecha_banco": date(2026, 5, 20),
    }
    payment_b = {
        "id_pago": "PAY-B",
        "cliente": "GEOEXCON",
        "monto_banco": 200,
        "fecha_banco": date(2026, 5, 20),
    }

    def _cand(credit: str) -> dict:
        return {
            "credito": credit,
            "fecha_limite": date(2026, 5, 23),
            "valor_obligacion_actual": 100,
            "saldo_vencido_visible": None,
            "link_extracto": "https://example/extracto",
            "link_tabla": "https://example/tabla",
            "link_carpeta_credito": "https://example/carpeta",
            "ruta_extracto_pdf": f"cli/{credit}/e.pdf",
            "ruta_unidad_credito": f"cli/{credit}",
            "ruta_tabla_amortizacion": f"cli/{credit}/t.xlsx",
            "credito_normalizado": credit,
            "right_panel_role": "VACIO",
            "parser_status": "OK",
            "extract_evidence": {},
        }

    raw = build_review_workbook_v4_bytes(
        process_id="proc-vis",
        process_date=date(2026, 5, 20),
        bank_code="banco_bogota",
        aplicacion_rows=[
            build_aplicacion_pagos_row(payment_a, _cand("258")),
            build_aplicacion_pagos_row(payment_b, _cand("301")),
        ],
        error_records=[
            {
                "id_pago": "PAY-A",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 258",
                "code": "extract_not_found",
                "link_extracto_url": "",
                "link_carpeta_credito_url": "https://example/carpeta258",
            }
        ],
    )
    wb = openpyxl.load_workbook(BytesIO(raw))
    ws = wb[ReviewSheets.APLICACION_PAGOS]
    help_txt = str(ws.cell(2, 1).value or "")
    assert "Control" not in help_txt
    assert "Procesar a SI" not in help_txt
    assert "Aplicación sugerida" not in [
        ws.cell(REVIEW_HEADER_ROW, c).value for c in range(1, 20)
    ]
    col_tipo = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.TIPO_APLICACION) + 1
    assert ws.column_dimensions[openpyxl.utils.get_column_letter(col_tipo)].width >= 48
    col_ext = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
    cell = ws.cell(REVIEW_FIRST_DATA_ROW, col_ext)
    assert "Ver extracto" in str(cell.value)
    assert "258" in str(cell.value)
    assert cell.hyperlink is not None
    r_geo = REVIEW_FIRST_DATA_ROW + 1
    assert ws.cell(r_geo, 1).value == "PAY-B"
    assert ws.cell(r_geo, col_ext).border.top.style == "medium"

    ws_err = wb[ReviewSheets.ERRORES]
    assert ws_err.sheet_state != "hidden"
    col_desc = ErroresCols.HEADERS.index(ErroresCols.DESCRIPCION) + 1
    assert "extracto" in str(ws_err.cell(REVIEW_FIRST_DATA_ROW, col_desc).value).lower()
    col_hacer = ErroresCols.HEADERS.index(ErroresCols.QUE_DEBE_HACER) + 1
    assert str(ws_err.cell(REVIEW_FIRST_DATA_ROW, col_hacer).value or "").strip()
    col_tipo = ErroresCols.HEADERS.index(ErroresCols.TIPO_CASO) + 1
    assert str(ws_err.cell(REVIEW_FIRST_DATA_ROW, col_tipo).value or "").strip()
    col_fold = ErroresCols.HEADERS.index(ErroresCols.LINK_CARPETA_CREDITO) + 1
    fold = ws_err.cell(REVIEW_FIRST_DATA_ROW, col_fold)
    assert "Ver carpeta" in str(fold.value)
    assert fold.hyperlink is not None


def test_workbook_v4_canonical_monto_on_first_id_pago_row_only():
    payment = {
        "id_pago": "PAY-CANON",
        "cliente": "CLI",
        "monto_banco": 1_500_000,
        "fecha_banco": date(2026, 5, 20),
    }

    def _cand(credito: str, oblig: float) -> dict:
        return {
            "credito": credito,
            "fecha_limite": date(2026, 5, 23),
            "valor_obligacion_actual": oblig,
            "saldo_vencido_visible": None,
            "link_extracto": "https://example/extracto",
            "link_tabla": "https://example/tabla",
            "link_carpeta_credito": "https://example/carpeta",
            "ruta_extracto_pdf": f"clientes/CLI/CREDITO {credito}/EXTRACTOS/e.pdf",
            "ruta_unidad_credito": f"clientes/CLI/CREDITO {credito}",
            "ruta_tabla_amortizacion": f"clientes/CLI/CREDITO {credito}/tabla.xlsx",
            "credito_normalizado": credito,
            "right_panel_role": "VACIO",
            "parser_status": "OK",
            "extract_evidence": {},
        }

    from app.application.services.review_schema import TipoAplicacionConfirmado

    row_no = build_aplicacion_pagos_row(payment, _cand("1", 900_000))
    row_si = build_aplicacion_pagos_row(payment, _cand("2", 1_500_000))
    row_no[AplicacionPagosCols.VALIDAR_PAGO] = ValidarPago.NO
    row_si[AplicacionPagosCols.VALIDAR_PAGO] = ValidarPago.SI
    row_si[AplicacionPagosCols.TIPO_APLICACION] = TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL
    row_si[AplicacionPagosCols.MONTO_BANCO] = None

    raw = build_review_workbook_v4_bytes(
        process_id="proc-canon",
        process_date=date(2026, 5, 20),
        bank_code="banco_bogota",
        aplicacion_rows=[row_no, row_si],
        error_records=[],
    )
    wb = openpyxl.load_workbook(BytesIO(raw))
    ws = wb[ReviewSheets.APLICACION_PAGOS]
    col_monto = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.MONTO_BANCO) + 1
    r1 = REVIEW_FIRST_DATA_ROW
    r2 = r1 + 1
    assert ws.cell(r1, col_monto).value == 1_500_000
    assert ws.cell(r2, col_monto).value in (None, "")
    assert "POR DISTRIBUIR" not in (ws.cell(1, 1).value or "")
    assert "POR DISTRIBUIR" not in (ws.cell(2, 1).value or "")
