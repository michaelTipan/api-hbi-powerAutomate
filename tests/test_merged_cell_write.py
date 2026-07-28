"""Unit smoke: escritura segura en MergedCell (descombinar y escribir en destino)."""
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell

from app.application.services.amortization_workbook import _set_cell_value


def test_set_cell_value_unmerges_and_writes_target_cell():
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:B2")
    assert isinstance(ws.cell(2, 2), MergedCell)
    _set_cell_value(ws, 2, 2, 123.45)
    assert ws.cell(2, 2).value == 123.45
    assert not isinstance(ws.cell(2, 2), MergedCell)


def test_set_cell_value_vertical_merge_keeps_row():
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("Q10:Q12")
    _set_cell_value(ws, 11, 17, "2026-07-19")
    assert ws.cell(11, 17).value == "2026-07-19"
