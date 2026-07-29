"""Protección de _AUTOMATION_LOG y no intervención de la API en las columnas O:P."""

from __future__ import annotations

import openpyxl

from app.application.services import amortization_workbook
from app.application.services.amortization_workbook import (
    AUTOMATION_LOG_SHEET,
    append_automation_log,
    protect_automation_log_sheet,
    ensure_automation_log,
)


def test_workbook_has_no_op_formula_fill_capability():
    """
    Las columnas O (dia) y P (Causac Inter Mes) las llena contabilidad: el
    multiplicador de días de causación es criterio de negocio y no se deriva del
    soporte. El servicio no debe reexponer una capacidad de arrastre.
    """
    for removed in (
        "ensure_application_related_formulas",
        "application_formula_fill_observability",
        "ApplicationFormulaFillResult",
        "APPLICATION_RELATED_FORMULA_COLUMNS",
    ):
        assert not hasattr(amortization_workbook, removed), removed


def test_automation_log_created_and_protected():
    wb = openpyxl.Workbook()
    ensure_automation_log(wb)
    protected, warning = protect_automation_log_sheet(wb)
    assert protected is True
    assert warning is None
    ws = wb[AUTOMATION_LOG_SHEET]
    assert ws.protection.sheet is True
    assert ws.cell(1, 1).protection.locked is True


def test_append_automation_log_on_protected_sheet(monkeypatch):
    monkeypatch.delenv("AMORTIZATION_LOG_SHEET_PROTECTION_PASSWORD", raising=False)
    wb = openpyxl.Workbook()
    ensure_automation_log(wb)
    protect_automation_log_sheet(wb)
    append_automation_log(
        wb,
        {
            "id_pago": "p1",
            "cliente": "c",
            "credito": "cr",
            "application_row": 8,
            "accion": "APLICADO",
            "idempotency_key": "k1",
        },
    )
    ws = wb[AUTOMATION_LOG_SHEET]
    assert ws.max_row >= 2
    protect_automation_log_sheet(wb)
    assert ws.protection.sheet is True


def test_automation_log_protection_password_from_env(monkeypatch):
    monkeypatch.setenv("AMORTIZATION_LOG_SHEET_PROTECTION_PASSWORD", "test-secret")
    wb = openpyxl.Workbook()
    ensure_automation_log(wb)
    protect_automation_log_sheet(wb)
    ws = wb[AUTOMATION_LOG_SHEET]
    assert ws.protection.sheet is True
    assert bool(ws.protection.password)
