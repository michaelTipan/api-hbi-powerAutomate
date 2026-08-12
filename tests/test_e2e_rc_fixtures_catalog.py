"""Tests locales del catálogo de fixtures RC y guards Notify sandbox."""

import io

import pytest
from openpyxl import Workbook, load_workbook

from scripts.e2e_rc.fixtures_catalog import (
    CORREOS_XLSX_REL,
    CRITICAL_FIXTURE_GAPS,
    HBI_DOMAIN_BLOCK,
    SANDBOX_NOTIFY_ALLOWLIST,
    assert_sandbox_notify_recipients,
    blank_image_like_pdf,
    build_sandbox_correos_xlsx,
    catalog_as_dicts,
    gaps_by_scenario,
    minimal_text_pdf,
    parseable_asiento_pdf,
    rewrite_correos_recipients_bytes,
)
from scripts.e2e_rc.review_edit import approve_credits_split
from scripts.e2e_rc.path_guard import AUTHORIZED_CLIENTS_BASE


def test_fixture_catalog_covers_known_blocked_ids():
    ids = {n.scenario_id for n in CRITICAL_FIXTURE_GAPS}
    for sid in ("E05", "E13", "E15", "E25", "E28", "E40"):
        assert sid in ids
    assert all(AUTHORIZED_CLIENTS_BASE in d["clients_base"] for d in catalog_as_dicts())
    e15 = next(n for n in CRITICAL_FIXTURE_GAPS if n.scenario_id == "E15")
    assert e15.local_generator == "parseable_asiento_pdf"
    assert e15.client_hint == "GEOEXCON"
    assert CORREOS_XLSX_REL.startswith(AUTHORIZED_CLIENTS_BASE)
    assert "PRUEBAS" in CORREOS_XLSX_REL
    assert CORREOS_XLSX_REL.endswith("CORREOS.xlsx")


def test_gaps_by_scenario_groups():
    grouped = gaps_by_scenario()
    assert "E28" in grouped
    assert grouped["E28"][0].local_generator == "blank_image_like_pdf"
    assert "E15" in grouped


def test_minimal_and_blank_pdf_generators():
    text_pdf = minimal_text_pdf(title="RC unit")
    blank = blank_image_like_pdf()
    assert text_pdf[:4] == b"%PDF"
    assert blank[:4] == b"%PDF"
    assert len(text_pdf) > len(blank)


def test_canonical_amortization_xlsx_detectable_by_parser():
    from app.application.services.amortization_workbook import detect_amortization_sheet
    from scripts.e2e_rc.fixtures_catalog import canonical_amortization_xlsx

    raw = canonical_amortization_xlsx(saldo_before=3_000_000.0)
    wb = load_workbook(io.BytesIO(raw))
    match = detect_amortization_sheet(wb)
    assert match.worksheet is not None
    headers = match.headers
    for key in (
        "dia",
        "mes",
        "anio",
        "fecha_pago",
        "valor_intereses",
        "abono_k",
        "valor_pagado_cliente",
        "saldo_a_capital",
    ):
        assert key in headers
    assert match.worksheet.cell(2, headers["saldo_a_capital"]).value == pytest.approx(
        3_000_000.0
    )


def test_spatial_extract_pdf_right_panel_roles():
    from app.application.services.extract_snapshot_parser import (
        RightPanelRole,
        parse_extract_snapshot,
    )
    from scripts.e2e_rc.fixtures_catalog import spatial_extract_pdf

    mora = parse_extract_snapshot(
        spatial_extract_pdf(
            credit="301",
            valor_obligacion=5_000_000.0,
            right_role="SALDO_VENCIDO",
            right_amount=2_000_000.0,
            left_intereses_mora=99_000.0,
        )
    )
    assert mora.right_panel_role == RightPanelRole.SALDO_VENCIDO
    assert mora.saldo_vencido_visible == pytest.approx(2_000_000.0)
    assert mora.valor_obligacion_actual == pytest.approx(5_000_000.0)

    hist = parse_extract_snapshot(
        spatial_extract_pdf(
            credit="301",
            valor_obligacion=5_000_000.0,
            right_role="APLICACION_ANTERIOR",
            right_amount=1_250_000.0,
        )
    )
    assert hist.right_panel_role == RightPanelRole.APLICACION_ANTERIOR
    assert hist.saldo_vencido_visible is None

    amb = parse_extract_snapshot(
        spatial_extract_pdf(
            credit="301",
            valor_obligacion=5_000_000.0,
            right_role="AMBIGUO",
            right_amount=2_000_000.0,
        )
    )
    assert amb.right_panel_role == RightPanelRole.AMBIGUO
    assert amb.saldo_vencido_visible is None


def test_parseable_asiento_pdf_roundtrip_parser():
    from app.application.services.accounting_pdf_parser import (
        ACCOUNT_CAPITAL,
        ACCOUNT_VALOR_PAGADO_CLIENTE,
        extract_text_from_pdf,
        parse_accounting_text,
    )

    raw = parseable_asiento_pdf(capital=5_000_000.0, valor_pagado=19_102_163.0)
    text = extract_text_from_pdf(raw)
    ev = parse_accounting_text(
        text,
        {
            "id_pago": "rc",
            "cliente": "GEOEXCON",
            "credito": "CREDITO # 231",
            "asiento_pdf_path": "x.pdf",
        },
    )
    assert ev.valor_pagado_cliente == pytest.approx(19_102_163.0, abs=0.02)
    assert ev.capital == pytest.approx(5_000_000.0, abs=0.02)
    assert ACCOUNT_VALOR_PAGADO_CLIENTE in (ev.detected_codes or []) or ev.valor_pagado_cliente > 0
    assert ACCOUNT_CAPITAL in (ev.detected_codes or []) or ev.capital > 0


def test_notify_allowlist_blocks_hbi_domain():
    assert "herramientas.jsakedev@gmail.com" in SANDBOX_NOTIFY_ALLOWLIST
    out = assert_sandbox_notify_recipients(["herramientas.jsakedev@gmail.com"])
    assert out == ["herramientas.jsakedev@gmail.com"]
    with pytest.raises(ValueError, match="blocked_hbi|not_sandbox"):
        assert_sandbox_notify_recipients([f"ops{HBI_DOMAIN_BLOCK}"])


def test_build_and_rewrite_correos_xlsx_strips_hbi():
    raw = build_sandbox_correos_xlsx()
    wb = load_workbook(io.BytesIO(raw))
    assert wb.active["A1"].value == "EMISOR"
    assert "@hbicapital.com.co" in str(wb.active["A2"].value).lower()
    assert wb.active["B2"].value == "herramientas.jsakedev@gmail.com"

    dirty = Workbook()
    ws = dirty.active
    ws["A1"] = "EMISOR"
    ws["B1"] = "RECEPTORES"
    ws["A2"] = f"ops{HBI_DOMAIN_BLOCK}"
    ws["B2"] = f"finanzas{HBI_DOMAIN_BLOCK}; otro@example.com"
    buf = io.BytesIO()
    dirty.save(buf)
    cleaned = rewrite_correos_recipients_bytes(buf.getvalue())
    wb2 = load_workbook(io.BytesIO(cleaned))
    assert HBI_DOMAIN_BLOCK not in str(wb2.active["A2"].value).lower()
    assert HBI_DOMAIN_BLOCK not in str(wb2.active["B2"].value).lower()
    assert "herramientas.jsakedev@gmail.com" in str(wb2.active["B2"].value).lower()


def test_approve_credits_split_edits_two_rows():
    from app.application.services.review_schema import (
        AplicacionPagosCols,
        REVIEW_SCHEMA_VERSION,
        ReviewSheets,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(list(AplicacionPagosCols.HEADERS))
    base = {h: "" for h in AplicacionPagosCols.HEADERS}
    for cred, monto in (("CREDITO # 231", 10.0), ("CREDITO # 254", 5.0)):
        row = dict(base)
        row.update(
            {
                AplicacionPagosCols.ID_PAGO: "p1",
                AplicacionPagosCols.CLIENTE: "GEOEXCON",
                AplicacionPagosCols.CREDITO: cred,
                AplicacionPagosCols.MONTO_BANCO: monto,
            }
        )
        ws.append([row[h] for h in AplicacionPagosCols.HEADERS])
    meta = wb.create_sheet(ReviewSheets.META)
    meta.append(["Campo", "Valor"])
    meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    buf = io.BytesIO()
    wb.save(buf)
    edited, applied = approve_credits_split(
        buf.getvalue(),
        splits=[
            {"credito_contains": "231", "tipo": "PAGO DE OBLIGACIÓN ACTUAL", "obligacion": 10.0},
            {"credito_contains": "254", "tipo": "ABONO A CAPITAL", "obligacion": 0.0, "capital": 5.0},
        ],
    )
    assert len(applied) == 2
    assert edited[:2] == b"PK"  # xlsx zip
