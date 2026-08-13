"""Cierres F-01 integrado / F-02 fault-injection / F-03 orden / F-04 manifest."""
from __future__ import annotations

import asyncio
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import pytest
from openpyxl import Workbook, load_workbook

from app.application.services.review_schema import (
    AplicacionPagosCols,
    REVIEW_SCHEMA_VERSION,
    ReviewSheets,
    TipoAplicacion,
    TipoAplicacionConfirmado,
    ValidarPago,
)
from app.application.use_cases.amortization_fill_dry_run import (
    BANK_ASIENTOS_NO_CUADRAN,
    BANK_MONTO_BANCO_MISSING,
    _enrich_payment_outputs_monto_from_hist,
    _reconcile_bank_vs_asientos_for_payment_outputs,
)
from app.application.use_cases.merge_composite_validado_pdfs import _group_meta_from_rows
from app.application.use_cases.payment_validation_finalize import (
    _coerce_abono_bank_amount,
    _materialize_group_monto_banco_on_si_rows,
)
from app.application.use_cases.send_validar_extractos_notification import (
    NOTIFY_MAIL_SENT_STEP,
    NOTIFY_SENDING_STEP,
    _control_snap_already_notified,
    _control_snap_notify_mail_uncertain,
    _notify_sending_marker,
    send_validar_extractos_notification_email,
)
from tests.test_finalize_validation import make_distrib_row


# ─── F-01: NO (con monto) → SI (sin monto) ───────────────────────────────────


def _build_no_si_workbook(
    *,
    id_pago: str = "ID-F01",
    monto_on_no: float = 10_000_000.0,
    tipo_si: str = TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL,
) -> bytes:
    no_row, _ = make_distrib_row(
        id_pago=id_pago,
        validar_pago=ValidarPago.NO,
        monto_banco=monto_on_no,
        fecha_banco=date(2026, 4, 22),
        tipo_aplicacion=tipo_si,
        credito="200",
    )
    si_row, _ = make_distrib_row(
        id_pago=id_pago,
        validar_pago=ValidarPago.SI,
        monto_banco=None,
        fecha_banco=date(2026, 4, 22),
        tipo_aplicacion=tipo_si,
        credito="100",
    )
    wb = Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(list(AplicacionPagosCols.HEADERS))
    ws.append(list(no_row))
    ws.append(list(si_row))
    meta = wb.create_sheet(ReviewSheets.META)
    meta.append(["Campo", "Valor"])
    meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_f01_no_then_si_materialize_merge_dryrun_pipeline() -> None:
    """Primera candidata NO con monto → SI vacía → canónico en SI → Merge → Dry-run."""
    raw = _build_no_si_workbook()
    wb = load_workbook(BytesIO(raw), data_only=True)
    ws = wb[ReviewSheets.APLICACION_PAGOS]
    headers = [str(c.value or "").strip() for c in ws[1]]
    distributions: list[dict] = []
    monto_casos: dict[str, float] = {}
    for r_idx in range(2, ws.max_row + 1):
        row_dict = {
            headers[i]: ws.cell(r_idx, i + 1).value for i in range(len(headers))
        }
        row_dict["_excel_row"] = r_idx
        id_pago = str(row_dict.get(AplicacionPagosCols.ID_PAGO) or "").strip()
        coerced = _coerce_abono_bank_amount(row_dict.get(AplicacionPagosCols.MONTO_BANCO))
        if id_pago and coerced is not None and coerced > 0:
            if id_pago not in monto_casos or monto_casos[id_pago] <= 0:
                monto_casos[id_pago] = float(coerced)
        distributions.append(row_dict)

    assert monto_casos["ID-F01"] == 10_000_000.0

    _materialize_group_monto_banco_on_si_rows(ws, 1, distributions, monto_casos)
    col_monto = headers.index(AplicacionPagosCols.MONTO_BANCO) + 1
    # Fila 2 = NO (ya tenía monto); fila 3 = SI (debe materializar).
    assert float(ws.cell(3, col_monto).value) == 10_000_000.0

    si_rows = [
        {
            "cliente": "CLI",
            "monto_banco": ws.cell(3, col_monto).value,
            "fecha_banco": date(2026, 4, 22),
            "credito_digits": "100",
        }
    ]
    _cli, monto_meta, _f, _creds = _group_meta_from_rows(
        si_rows, tipo_aplicacion=TipoAplicacion.PAGO.value
    )
    assert monto_meta == 10_000_000.0

    payment_outputs = [{"id_pago": "ID-F01", "monto_banco": None}]
    hist_index = {("ID-F01", "100"): {"monto_banco": float(ws.cell(3, col_monto).value)}}
    _enrich_payment_outputs_monto_from_hist(payment_outputs, hist_index)
    assert payment_outputs[0]["monto_banco"] == 10_000_000.0

    items = [
        {
            "id_pago": "ID-F01",
            "tipo_aplicacion": TipoAplicacion.PAGO.value,
            "asiento_pdf_path": "A/asiento.pdf",
            "payment_application": {"valor_pagado_cliente": 10_000_000},
            "error_code": None,
        }
    ]
    out = _reconcile_bank_vs_asientos_for_payment_outputs(payment_outputs, items)
    assert out[0].get("error_code") not in {
        BANK_MONTO_BANCO_MISSING,
        BANK_ASIENTOS_NO_CUADRAN,
    }


def test_f01_abono_reuses_same_monto_canonical_path() -> None:
    raw = _build_no_si_workbook(
        id_pago="AB-F01",
        monto_on_no=2_500_000.0,
        tipo_si=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
    )
    wb = load_workbook(BytesIO(raw), data_only=True)
    ws = wb[ReviewSheets.APLICACION_PAGOS]
    headers = [str(c.value or "").strip() for c in ws[1]]
    distributions = []
    monto_casos: dict[str, float] = {}
    for r_idx in range(2, ws.max_row + 1):
        row_dict = {
            headers[i]: ws.cell(r_idx, i + 1).value for i in range(len(headers))
        }
        row_dict["_excel_row"] = r_idx
        id_pago = str(row_dict.get(AplicacionPagosCols.ID_PAGO) or "").strip()
        coerced = _coerce_abono_bank_amount(row_dict.get(AplicacionPagosCols.MONTO_BANCO))
        if id_pago and coerced is not None and coerced > 0:
            monto_casos.setdefault(id_pago, float(coerced))
        distributions.append(row_dict)

    _materialize_group_monto_banco_on_si_rows(ws, 1, distributions, monto_casos)
    col_monto = headers.index(AplicacionPagosCols.MONTO_BANCO) + 1
    assert float(ws.cell(3, col_monto).value) == 2_500_000.0
    _cli, monto_meta, _f, _c = _group_meta_from_rows(
        [
            {
                "cliente": "CLI",
                "monto_banco": ws.cell(3, col_monto).value,
                "fecha_banco": None,
                "credito_digits": "100",
            }
        ],
        tipo_aplicacion=TipoAplicacion.ABONO.value,
    )
    assert monto_meta == 2_500_000.0


# ─── F-02: checkpoint NOTIFY_SENDING ─────────────────────────────────────────


def test_f02_uncertain_checkpoint_blocks_already_notified_false() -> None:
    snap = SimpleNamespace(
        estado_proceso="FINALIZADO",
        notify_idempotency_key=_notify_sending_marker("pk-1"),
        email_pdf_path="",
        last_completed_step=NOTIFY_SENDING_STEP,
    )
    assert _control_snap_notify_mail_uncertain(snap) is True
    assert _control_snap_already_notified(snap) is False


def test_f02_mail_sent_checkpoint_is_already_notified() -> None:
    snap = SimpleNamespace(
        estado_proceso="PENDIENTE_ASIENTOS",
        notify_idempotency_key="pk-1",
        email_pdf_path="",
        last_completed_step=NOTIFY_MAIL_SENT_STEP,
    )
    assert _control_snap_notify_mail_uncertain(snap) is False
    assert _control_snap_already_notified(snap) is True


def test_f02_sendmail_exactly_once_when_mail_sent_persist_fails(monkeypatch) -> None:
    """sendMail 202 → falla persist MAIL_SENT → retry NO llama sendMail otra vez."""
    from openpyxl import Workbook as _Wb

    PK = "payment-validation|banco_bogota|2026-06-03"
    monkeypatch.setenv("GRAPH_VALIDAR_NOTIFY_EXPORT_EMAIL_PDF", "false")
    monkeypatch.setenv("GRAPH_VALIDAR_NOTIFY_ATTACH_PDFS", "false")

    control = {
        "process_key": PK,
        "estado_proceso": "FINALIZADO",
        "is_active": True,
        "notify_idempotency_key": "",
        "last_completed_step": "",
        "email_pdf_path": "",
        "historical_file_path": "HIST/cartera.xlsx",
        "control_file_path": "CTL/x.xlsx",
        "bank_code": "banco_bogota",
        "bank_name": "Banco Bogotá",
    }
    send_calls: list[str] = []

    def _snap():
        return SimpleNamespace(
            process_key=control["process_key"],
            estado_proceso=control["estado_proceso"],
            is_active=control["is_active"],
            notify_idempotency_key=control["notify_idempotency_key"],
            last_completed_step=control["last_completed_step"],
            email_pdf_path=control["email_pdf_path"],
            historical_file_path=control["historical_file_path"],
            control_file_path=control["control_file_path"],
            bank_code=control["bank_code"],
            bank_name=control["bank_name"],
            validation_file_path="",
            secretary_file_path="",
            merge_manifest_path="",
            merge_idempotency_key="",
            apply_idempotency_key="",
            process_id="",
        )

    async def _read(*_a, **_k):
        return _snap()

    async def _update(_g, _s, _d, *, bank_code: str, updates: dict):
        step = str(updates.get("LastCompletedStep") or "")
        if step == NOTIFY_MAIL_SENT_STEP:
            raise RuntimeError("injected: persist NOTIFY_MAIL_SENT failed")
        if "NotifyIdempotencyKey" in updates:
            control["notify_idempotency_key"] = str(updates["NotifyIdempotencyKey"])
        if "LastCompletedStep" in updates:
            control["last_completed_step"] = str(updates["LastCompletedStep"])
        if "EstadoProceso" in updates:
            control["estado_proceso"] = str(updates["EstadoProceso"])
        return True

    def _bank_bytes() -> bytes:
        wb = _Wb()
        ws = wb.active
        ws.append(["Fecha", "Concepto", "Crédito", "Monto"])
        ws.append([date(2026, 6, 3), "pago", "258", "1000"])
        bio = BytesIO()
        wb.save(bio)
        return bio.getvalue()

    def _hist_bytes() -> bytes:
        r, _ = make_distrib_row(
            id_pago="P1",
            validar_pago=ValidarPago.SI,
            ruta_pdf_internal="clientes/CLI/extractos/e1.pdf",
            tipo_aplicacion=TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL,
            monto_banco=1_000_000,
            fecha_banco=date(2026, 6, 3),
        )
        wb = _Wb()
        ws = wb.active
        ws.title = ReviewSheets.APLICACION_PAGOS
        ws.append(list(AplicacionPagosCols.HEADERS) + ["_ruta_extracto"])
        ws.append(list(r) + ["clientes/CLI/e.pdf"])
        meta = wb.create_sheet(ReviewSheets.META)
        meta.append(["Campo", "Valor"])
        meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
        bio = BytesIO()
        wb.save(bio)
        return bio.getvalue()

    class _Graph:
        async def post_json(self, endpoint: str, body: dict):
            if "sendMail" in endpoint:
                send_calls.append(endpoint)
                return {}, 202
            return {}, 201

        async def get_bytes(self, *_a, **_k):
            return _bank_bytes()

        async def put_bytes(self, *_a, **_k):
            return {}

    async def _dl(_g, _s, _d, path: str):
        if "HIST" in path or "cartera" in path:
            return _hist_bytes()
        return _bank_bytes()

    g = _Graph()
    patches = [
        patch(
            "app.application.use_cases.send_validar_extractos_notification.resolve_sharepoint_from_env",
            new_callable=AsyncMock,
            return_value={"site_id": "s1", "drive_id": "d1", "path_encoded": "x"},
        ),
        patch(
            "app.application.use_cases.send_validar_extractos_notification._load_sender_and_recipients_from_correos_xlsx",
            new_callable=AsyncMock,
            return_value=("sender@hbi.test", ["to@hbi.test"]),
        ),
        patch(
            "app.application.use_cases.send_validar_extractos_notification._graph_download_by_path",
            new_callable=AsyncMock,
            side_effect=_dl,
        ),
        patch(
            "app.application.use_cases.send_validar_extractos_notification.resolve_sharepoint_path",
            new_callable=AsyncMock,
            return_value={
                "site_id": "s1",
                "drive_id": "d1",
                "path_encoded": "bank/report.xlsx",
                "file_path": "banco.xlsx",
            },
        ),
        patch(
            "app.application.use_cases.payment_validation_process_control.read_process_control_snapshot",
            new_callable=AsyncMock,
            side_effect=_read,
        ),
        patch(
            "app.application.use_cases.payment_validation_process_control.update_process_control_row2",
            new_callable=AsyncMock,
            side_effect=_update,
        ),
    ]

    with patches[0], patches[1], patches[2], patches[3], patches[4], patches[5]:
        first = asyncio.run(
            send_validar_extractos_notification_email(
                g,
                bank_code="banco_bogota",
                historical_file_path="HIST/cartera.xlsx",
            )
        )
        assert len(send_calls) == 1
        assert control["last_completed_step"] == NOTIFY_SENDING_STEP
        assert str(control["notify_idempotency_key"]).startswith("NOTIFY_SENDING|")

        second = asyncio.run(
            send_validar_extractos_notification_email(
                g,
                bank_code="banco_bogota",
                historical_file_path="HIST/cartera.xlsx",
            )
        )
        assert len(send_calls) == 1, "retry no debe volver a llamar sendMail"
        assert second.merge_control_error_code == "notify_mail_uncertain"


# ─── F-03: orden Control → cleanup ───────────────────────────────────────────


def test_f03_control_failure_skips_review_delete(monkeypatch) -> None:
    """Si Control final falla, no se elimina el Excel de revisión."""
    from datetime import date as _date

    from tests.test_amortization_fill_apply import MockGraphApply
    from tests.test_amortization_fill_dry_run import (
        _accounting_text,
        _amort_table_date_at_row,
        _asiento_pdf_placeholder,
        _base_files,
        _hist_bytes,
        _ibr_bytes,
    )
    from app.application.use_cases.amortization_fill_apply import run_amortization_fill_apply
    from app.application.use_cases.amortization_application_plan import (
        prepare_amortization_application,
    )

    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")

    deleted = {"n": 0}

    async def _fail_control(*_a, **_k):
        raise RuntimeError("injected control final failure")

    async def _track_delete(*_a, **_k):
        deleted["n"] += 1
        return {"deleted": True, "path": "REV/x.xlsx"}

    fecha = _date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.update_process_control_row2",
        _fail_control,
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply._delete_review_validation_file",
        _track_delete,
    )

    async def _prep(*a, **k):
        plan = await prepare_amortization_application(*a, **k)
        object.__setattr__(plan, "review_validation_path", "REV/review.xlsx")
        return plan

    monkeypatch.setattr(
        "app.application.use_cases.amortization_application_plan.prepare_amortization_application",
        _prep,
    )

    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out.get("outcome") == "applied_control_pending"
    assert deleted["n"] == 0
    cleanup = out.get("review_validation_file_cleanup") or {}
    assert cleanup.get("reason") == "skipped_control_not_finalized"


def test_f03_apply_deletes_review_only_after_control(monkeypatch) -> None:
    """Orden observable: update Control final → luego delete review."""
    from datetime import date as _date

    from tests.test_amortization_fill_apply import MockGraphApply
    from tests.test_amortization_fill_dry_run import (
        _accounting_text,
        _amort_table_date_at_row,
        _asiento_pdf_placeholder,
        _base_files,
        _hist_bytes,
        _ibr_bytes,
    )
    from app.application.use_cases.amortization_fill_apply import run_amortization_fill_apply
    from app.application.use_cases.amortization_application_plan import (
        prepare_amortization_application,
    )

    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")

    events: list[str] = []

    async def _ok_update(*_a, **_k):
        events.append("control_update")
        return None

    async def _track_delete(*_a, **_k):
        events.append("review_delete")
        return {"deleted": True, "path": "REV/x.xlsx"}

    fecha = _date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.update_process_control_row2",
        _ok_update,
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply._delete_review_validation_file",
        _track_delete,
    )

    async def _prep(*a, **k):
        plan = await prepare_amortization_application(*a, **k)
        object.__setattr__(plan, "review_validation_path", "REV/review.xlsx")
        return plan

    monkeypatch.setattr(
        "app.application.use_cases.amortization_application_plan.prepare_amortization_application",
        _prep,
    )

    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out.get("outcome") == "applied"
    assert "control_update" in events
    assert "review_delete" in events
    assert events.index("control_update") < events.index("review_delete")


def test_f03_cleanup_failure_after_control_keeps_applied(monkeypatch) -> None:
    """Cleanup falla tras Control OK → outcome applied (no applied_control_pending)."""
    from datetime import date as _date

    from tests.test_amortization_fill_apply import MockGraphApply
    from tests.test_amortization_fill_dry_run import (
        _accounting_text,
        _amort_table_date_at_row,
        _asiento_pdf_placeholder,
        _base_files,
        _hist_bytes,
        _ibr_bytes,
    )
    from app.application.use_cases.amortization_fill_apply import run_amortization_fill_apply
    from app.application.use_cases.amortization_application_plan import (
        prepare_amortization_application,
    )

    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")

    async def _ok_update(*_a, **_k):
        return None

    async def _fail_delete(*_a, **_k):
        raise RuntimeError("injected cleanup failure")

    fecha = _date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphApply(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.update_process_control_row2",
        _ok_update,
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply._delete_review_validation_file",
        _fail_delete,
    )

    async def _prep(*a, **k):
        plan = await prepare_amortization_application(*a, **k)
        object.__setattr__(plan, "review_validation_path", "REV/review.xlsx")
        return plan

    monkeypatch.setattr(
        "app.application.use_cases.amortization_application_plan.prepare_amortization_application",
        _prep,
    )

    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out.get("process_control_finalized") is True
    assert out.get("outcome") == "applied"
    cleanup = out.get("review_validation_file_cleanup") or {}
    assert cleanup.get("reason") == "cleanup_failed_after_control"


# ─── F-04: manifest upload fail ──────────────────────────────────────────────


def test_f04_manifest_upload_failure_yields_error_merge(monkeypatch) -> None:
    from app.application.sharepoint_resolution import encode_graph_drive_path
    from app.application.use_cases.merge_composite_validado_pdfs import (
        merge_composite_validado_pdfs,
    )
    from tests.test_merge_composite_control_workbook import (
        _MergeGraph,
        _bank_bytes,
        _hist_workbook_bytes,
        _tiny_pdf,
    )
    import app.application.use_cases.payment_validation_process_control as pc

    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST_SITE")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "TEST_DRIVE")
    monkeypatch.setenv("GRAPH_SHAREPOINT_FILE_PATH", "bank/report.xlsx")
    monkeypatch.setenv("GRAPH_MERGE_COMPOSITE_OUTPUT_FOLDER_PATH", "OUT/PDFS")

    hist = "HIST/hist.xlsx"
    email = "EMAIL/mail.pdf"
    extract = "clientes/ACME/CREDITO# 264/Extracto.pdf"
    asiento_dir = "clientes/ACME/CREDITO# 264/ASIENTOS CONTABLES CRED 264"
    asiento_rel = f"{asiento_dir}/asiento_264.pdf"

    g = _MergeGraph()
    g.initial["bank/report.xlsx"] = _bank_bytes()
    g.initial[hist] = _hist_workbook_bytes(
        [["VALIDAR", "", "G1", "ACME", "264", asiento_dir]]
    )
    g.initial[email] = _tiny_pdf()
    g.initial[extract] = _tiny_pdf()
    g.initial[asiento_rel] = _tiny_pdf()
    g.children[asiento_dir] = [{"name": "asiento_264.pdf", "file": {}}]

    ctx = {
        "site_id": "s1",
        "drive_id": "d1",
        "path_encoded": encode_graph_drive_path("bank/report.xlsx"),
        "file_path": "bank/report.xlsx",
    }
    captured: list[dict] = []

    async def _fake_read(_g, _s, _d, *, bank_code: str):
        return type(
            "Snap",
            (),
            {
                "control_file_path": "CTL/ignored.xlsx",
                "estado_proceso": "PENDIENTE_ASIENTOS",
                "is_active": True,
                "process_key": f"payment-validation|{bank_code}|2026-06-01",
                "validation_file_path": "",
                "historical_file_path": "",
                "secretary_file_path": "",
                "email_pdf_path": "",
                "notify_idempotency_key": "",
                "merge_manifest_path": "",
                "merge_idempotency_key": "",
                "bank_code": bank_code,
                "bank_name": "",
                "apply_idempotency_key": "",
                "process_id": "",
                "last_completed_step": "",
            },
        )()

    async def capture_update(_g, _s, _d, *, bank_code: str, updates: dict):
        captured.append(dict(updates))
        return True

    async def fake_collect(_g, _s, _d, _cell):
        return [extract]

    async def boom_manifest(*_a, **_k):
        raise RuntimeError("injected manifest upload failure")

    monkeypatch.setattr(pc, "read_process_control_snapshot", _fake_read)
    monkeypatch.setattr(pc, "update_process_control_row2", capture_update)

    with (
        patch(
            "app.application.use_cases.merge_composite_validado_pdfs.resolve_sharepoint_from_env",
            new_callable=AsyncMock,
            return_value=ctx,
        ),
        patch(
            "app.application.use_cases.merge_composite_validado_pdfs.resolve_sharepoint_path",
            new_callable=AsyncMock,
            return_value=ctx,
        ),
        patch(
            "app.application.use_cases.merge_composite_validado_pdfs._collect_pdf_paths_from_ruta_cell",
            new_callable=AsyncMock,
            side_effect=fake_collect,
        ),
        patch(
            "app.application.use_cases.merge_composite_validado_pdfs._upload_merge_manifest",
            new_callable=AsyncMock,
            side_effect=boom_manifest,
        ),
    ):
        result = asyncio.run(
            merge_composite_validado_pdfs(
                g,
                force_rebuild=True,
                bank_code="banco_bogota",
                historical_file_path=hist,
                email_pdf_path=email,
            )
        )

    assert result.merge_control_status == "ERROR_MERGE" or any(
        u.get("EstadoProceso") == "ERROR_MERGE" for u in captured
    )
    assert result.eligible_for_dry_run is False
    assert not (result.merge_idempotency_key or "").strip()
    err_updates = [u for u in captured if u.get("EstadoProceso") == "ERROR_MERGE"]
    assert err_updates
    assert err_updates[-1].get("MergeIdempotencyKey") in ("", None)
    assert err_updates[-1].get("LastStepErrorCode") == "missing_merge_manifest"
