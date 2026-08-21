from app.application.services.review_schema import AplicacionPagosCols, ValidarPago
"""Notify validar extractos (Phase 3): auto-resolve por banco vía control cuando no hay body."""

import asyncio
import time
from datetime import date
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch

import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.adapters.primary.http.deps import init_graph_client
from app.adapters.primary.http.routers.sharepoint import router
from app.application.job_status_enrichment import enrich_job_for_http_response
from app.application.use_cases.send_validar_extractos_notification import (
    _find_distribucion_header_row,
    _process_date_from_process_key,
    send_validar_extractos_notification_email,
)
from app.application.use_cases.setup_merge_control_workbook import (
    PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA,
    PROCESS_CONTROL_BANK_FILE_BOGOTA,
    _build_process_control_workbook_bytes,
)

def _minimal_bank_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Fecha", "Concepto", "Crédito", "Extra"])
    ws.append([date(2026, 5, 12), "abono", "C1", "x"])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _minimal_historico_xlsx(
    estado_header: str = "Estado línea",
    ruta_header: str = "Ruta",
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Aplicacion_Pagos"
    ws.append([estado_header, ruta_header])
    ws.append(["VALIDAR", ""])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


@pytest.fixture  # type: ignore[name-defined]
def client():
    import os
    os.environ["GRAPH_SHAREPOINT_SITE_SEARCH"] = "SITIO"
    os.environ["GRAPH_SHAREPOINT_DRIVE_NAME"] = "DRIVE"
    # Necesario para resolve_sharepoint_from_env, aunque los tests parcheen descargas.
    os.environ["GRAPH_SHAREPOINT_FILE_PATH"] = "banco.xlsx"
    os.environ["GRAPH_BANK_PAYMENTS_FILE_PATH"] = "banco.xlsx"
    os.environ["GRAPH_BANK_PAYMENTS_FILE_PATH_BANCOLOMBIA"] = "banco.xlsx"

    class _GraphStub:
        async def get(self, endpoint, params=None):
            if endpoint == "/sites":
                return {"value": [{"id": "s1"}]}
            if endpoint == "/sites/s1/drives":
                return {"value": [{"id": "d1", "name": "DRIVE"}]}
            return {}

        async def get_bytes(self, endpoint, params=None):
            # Resolve sharepoint paths in notify + controls by bank.
            if endpoint.endswith("banco.xlsx:/content"):
                return _minimal_bank_xlsx()
            if "control_proceso_validacion_pagos_banco_bogota.xlsx" in endpoint:
                return _build_process_control_workbook_bytes("banco_bogota", "Banco de Bogotá")
            if "control_proceso_validacion_pagos_banco_bancolombia.xlsx" in endpoint:
                return _build_process_control_workbook_bytes("banco_bancolombia", "Bancolombia")
            return b""

        async def post_json(self, *_a, **_k):
            return {}, 202

        async def put_bytes(self, *_a, **_k):
            return {}

    app = FastAPI()
    init_graph_client(_GraphStub())
    app.include_router(router)
    from app.application.job_manager import JobManager
    from app.application.services.merge_queue_service import (
        reset_merge_queue_service_for_tests,
    )
    from app.application.services.notify_queue_service import (
        reset_notify_queue_service_for_tests,
    )

    reset_notify_queue_service_for_tests()
    reset_merge_queue_service_for_tests()
    jm = JobManager()
    jm._validation_jobs.clear()
    jm._generate_active = False
    jm._finalize_active = False
    jm._notify_active = False
    jm._merge_active = False
    yield TestClient(app, raise_server_exceptions=False)
    reset_notify_queue_service_for_tests()
    reset_merge_queue_service_for_tests()
    jm._validation_jobs.clear()
    jm._generate_active = False
    jm._finalize_active = False
    jm._notify_active = False
    jm._merge_active = False


def _poll_job(client: TestClient, jid: str, timeout: float = 5.0) -> dict:
    deadline = time.perf_counter() + timeout
    last: dict = {}
    while time.perf_counter() < deadline:
        res = client.get(f"/graph/sharepoint/notify-validar-extractos-email/jobs/{jid}")
        last = res.json()
        if last.get("status") in ("completed", "failed"):
            return last
        time.sleep(0.02)
    return last


def test_notify_requires_historical_file_path(client):
    r = client.post("/graph/sharepoint/notify-validar-extractos-email", json={})
    assert r.status_code == 202
    body = _poll_job(client, r.json()["job_id"])
    assert body.get("status") == "failed"
    err = body.get("error")
    assert isinstance(err, dict)
    enriched = enrich_job_for_http_response(body)
    e = enriched["error"]
    # Sin body ahora intenta auto-detectar banco por control; si no hay ninguno listo -> NO_READY_PROCESS.
    assert e["error_code"] in ("NO_READY_PROCESS", "missing_historical_file_path")


def test_notify_rejects_blank_historical_file_path(client):
    r = client.post(
        "/graph/sharepoint/notify-validar-extractos-email",
        json={"historical_file_path": "   "},
    )
    assert r.status_code == 202
    body = _poll_job(client, r.json()["job_id"])
    assert body.get("status") == "failed"
    enriched = enrich_job_for_http_response(body)
    assert enriched["error"]["error_code"] in ("NO_READY_PROCESS", "missing_historical_file_path")


def test_notify_failed_explicit_history_missing_returns_standard_error():
    """404 al descargar histórico → historical_file_not_found enriquecido."""

    async def fail_hist(_g, _s, _d, _path):
        req = httpx.Request("GET", "https://graph.microsoft.com/v1.0/x")
        resp = httpx.Response(404, request=req, text="not found")
        raise httpx.HTTPStatusError("404", request=req, response=resp)

    class _GraphOk:
        async def get_bytes(self, *_a, **_k):
            return _minimal_bank_xlsx()

        async def post_json(self, *_a, **_k):
            return {}, 202

        async def put_bytes(self, *_a, **_k):
            return {}

    async def run():
        with (
            patch(
                "app.application.use_cases.send_validar_extractos_notification.resolve_sharepoint_from_env",
                new_callable=AsyncMock,
                return_value={
                    "site_id": "s1",
                    "drive_id": "d1",
                    "path_encoded": "bank/report.xlsx",
                },
            ),
            patch(
                "app.application.use_cases.send_validar_extractos_notification._load_sender_and_recipients_from_correos_xlsx",
                new_callable=AsyncMock,
                return_value=("sender@example.com", ["to@example.com"]),
            ),
            patch(
                "app.application.use_cases.send_validar_extractos_notification._graph_download_by_path",
                new_callable=AsyncMock,
                side_effect=fail_hist,
            ),
                patch(
                    "app.application.use_cases.send_validar_extractos_notification.resolve_sharepoint_path",
                    new_callable=AsyncMock,
                    return_value={
                        "site_id": "s1",
                        "drive_id": "d1",
                        "path_encoded": "bank/report.xlsx",
                        "file_path": "banco.xlsx",
                    },
                ),
        ):
            g = _GraphOk()
            with pytest.raises(ValueError) as ei:
                await send_validar_extractos_notification_email(
                    g,
                    historical_file_path="path/missing.xlsx",
                )
            assert str(ei.value).startswith("historical_file_not_found")

        raw = {
            "job_id": "j",
            "type": "notify_validar_extractos",
            "status": "failed",
            "error": str(ei.value),
        }
        out = enrich_job_for_http_response(raw)
        e = out["error"]
        assert e["error_code"] == "historical_file_not_found"
        assert e["user_message"]
        # El paso de finalize se nombra "Flujo 2" en los mensajes al operador.
        next_action = e["next_action"].lower()
        assert "flujo 2" in next_action or "historico" in next_action

    asyncio.run(run())


def test_notify_standard_job_response_fields_remain_backward_compatible(client):
    r = client.post("/graph/sharepoint/notify-validar-extractos-email", json={})
    assert r.status_code == 202
    jid = r.json()["job_id"]
    body = _poll_job(client, jid)
    assert body.get("job_id") == jid
    assert body.get("type") == "notify_validar_extractos"
    enriched = enrich_job_for_http_response(body)
    assert enriched.get("severity") == "error"
    err = enriched.get("error") or {}
    assert err.get("user_message")
    assert err.get("next_action")
    assert err.get("error_code") in ("NO_READY_PROCESS", "missing_historical_file_path")


def test_notify_completed_result_includes_historical_file_source_explicit():
    raw = {
        "job_id": "j",
        "type": "notify_validar_extractos",
        "status": "completed",
        "result": {
            "status": "ok",
            "historical_file_path": "HIST/x.xlsx",
            "historico_excel_path": "HIST/x.xlsx",
            "historical_file_source": "explicit",
            "attachments_count": 0,
        },
    }
    out = enrich_job_for_http_response(raw)
    assert out["result"]["historical_file_source"] == "explicit"
    assert out["result"]["historical_file_path"] == "HIST/x.xlsx"


def test_notify_enrichment_maps_missing_historical_file_path_string():
    raw = {
        "job_id": "j",
        "type": "notify_validar_extractos",
        "status": "failed",
        "error": "missing_historical_file_path",
    }
    out = enrich_job_for_http_response(raw)
    assert out["error"]["error_code"] == "missing_historical_file_path"
    assert "histórico" in out["error"]["user_message"].lower()
    assert out["error"]["next_action"]


def test_find_distribucion_header_row_accepts_estado_nuevo():
    wb = Workbook()
    ws = wb.active
    ws.title = "Aplicacion_Pagos"
    ws.append(["ID Pago", "Estado", "Ruta"])
    row, hmap = _find_distribucion_header_row(ws)
    assert row == 1
    assert hmap["ESTADO"] == 2
    assert hmap["RUTA"] == 3


def test_find_distribucion_header_row_accepts_estado_linea_legacy():
    wb = Workbook()
    ws = wb.active
    ws.title = "Aplicacion_Pagos"
    ws.append(["Estado línea", "Rutas"])
    row, hmap = _find_distribucion_header_row(ws)
    assert row == 1
    assert hmap["ESTADO LINEA"] == 1
    assert hmap["RUTAS"] == 2


def test_find_distribucion_header_row_missing_status_column():
    wb = Workbook()
    ws = wb.active
    ws.append(["Ruta", "Cliente"])
    with pytest.raises(ValueError, match="missing_distribucion_status_column"):
        _find_distribucion_header_row(ws)


def test_find_distribucion_header_row_missing_route_column():
    wb = Workbook()
    ws = wb.active
    ws.append(["Estado", "Cliente"])
    with pytest.raises(ValueError, match="missing_distribucion_route_column"):
        _find_distribucion_header_row(ws)


def test_find_distribucion_header_row_missing_headers_entirely():
    wb = Workbook()
    ws = wb.active
    ws.append(["Cliente", "Crédito"])
    with pytest.raises(ValueError, match="missing_distribucion_headers"):
        _find_distribucion_header_row(ws)


def test_notify_enrichment_maps_missing_distribucion_status_column():
    raw = {
        "job_id": "j",
        "type": "notify_validar_extractos",
        "status": "failed",
        "error": "missing_distribucion_status_column",
    }
    out = enrich_job_for_http_response(raw)
    assert out["error"]["error_code"] == "missing_distribucion_status_column"
    assert out["error"]["error_code"] != "unknown_error"
    assert "Estado" in out["error"]["user_message"]


def test_notify_enrichment_maps_missing_distribucion_headers():
    raw = {
        "job_id": "j",
        "type": "notify_validar_extractos",
        "status": "failed",
        "error": "missing_distribucion_headers",
    }
    out = enrich_job_for_http_response(raw)
    assert out["error"]["error_code"] == "missing_distribucion_headers"
    assert out["error"]["error_code"] != "unknown_error"




def test_notify_enrichment_maps_historical_file_not_found_prefix():
    raw = {
        "job_id": "j",
        "type": "notify_validar_extractos",
        "status": "failed",
        "error": "historical_file_not_found|HTTP 404 url=x detail='y'",
    }
    out = enrich_job_for_http_response(raw)
    assert out["error"]["error_code"] == "historical_file_not_found"


def test_notify_use_case_requires_historical_file_path():
    async def run():
        # Sin historical_file_path ahora intenta auto-detectar por control; sin procesos listos -> NO_READY_PROCESS.
        g = MagicMock()
        with (
            patch(
                "app.application.use_cases.send_validar_extractos_notification.resolve_sharepoint_from_env",
                new_callable=AsyncMock,
                return_value={"site_id": "s1", "drive_id": "d1", "path_encoded": "bank/report.xlsx"},
            ),
            patch(
                "app.application.use_cases.payment_validation_process_control.read_process_control_snapshot",
                new_callable=AsyncMock,
                return_value=MagicMock(
                    estado_proceso="VACIO",
                    is_active=False,
                    historical_file_path="",
                    email_pdf_path="",
                    notify_idempotency_key="",
                    process_key="",
                    validation_file_path="",
                    secretary_file_path="",
                    bank_code="banco_bogota",
                    bank_name="Banco de Bogotá",
                    control_file_path="CTL.xlsx",
                ),
            ),
        ):
            with pytest.raises(ValueError, match="NO_READY_PROCESS"):
                await send_validar_extractos_notification_email(g, historical_file_path=None)

    asyncio.run(run())


def test_notify_use_case_rejects_whitespace_only_path():
    async def run():
        g = MagicMock()
        with (
            patch(
                "app.application.use_cases.send_validar_extractos_notification.resolve_sharepoint_from_env",
                new_callable=AsyncMock,
                return_value={"site_id": "s1", "drive_id": "d1", "path_encoded": "bank/report.xlsx"},
            ),
            patch(
                "app.application.use_cases.payment_validation_process_control.read_process_control_snapshot",
                new_callable=AsyncMock,
                return_value=MagicMock(
                    estado_proceso="VACIO",
                    is_active=False,
                    historical_file_path="",
                    email_pdf_path="",
                    notify_idempotency_key="",
                    process_key="",
                    validation_file_path="",
                    secretary_file_path="",
                    bank_code="banco_bogota",
                    bank_name="Banco de Bogotá",
                    control_file_path="CTL.xlsx",
                ),
            ),
        ):
            with pytest.raises(ValueError, match="NO_READY_PROCESS"):
                await send_validar_extractos_notification_email(g, historical_file_path="  \t  ")

    asyncio.run(run())








def test_process_date_from_process_key_extracts_iso():
    assert _process_date_from_process_key(
        "payment-validation|banco_bogota|2026-07-27"
    ) == date(2026, 7, 27)
    assert _process_date_from_process_key(
        "payment-validation|banco_bogota|2026-07-27|4df53868-eeb1-428f-9c92-98e0efcad7ec"
    ) == date(2026, 7, 27)
    assert _process_date_from_process_key("payment-validation|banco_bogota") is None
    assert _process_date_from_process_key("") is None


def test_format_fechas_validacion_lists_all_dates():
    from app.application.use_cases.send_validar_extractos_notification import (
        _dates_from_bank_email_table,
        _format_fechas_validacion_es,
        _intro_fechas_clause,
        _split_saludo,
    )

    assert _format_fechas_validacion_es([date(2026, 4, 1)]) == "01/04/2026"
    assert (
        _format_fechas_validacion_es([date(2026, 5, 10), date(2026, 4, 1)])
        == "01/04/2026 y 10/05/2026"
    )
    assert (
        _format_fechas_validacion_es(
            [date(2026, 7, 27), date(2026, 4, 1), date(2025, 9, 15)]
        )
        == "15/09/2025, 01/04/2026 y 27/07/2026"
    )
    assert _intro_fechas_clause("01/04/2026", plural=False) == "El día 01/04/2026"
    assert (
        _intro_fechas_clause("01/04/2026 y 10/05/2026", plural=True)
        == "Los días 01/04/2026 y 10/05/2026"
    )
    saludo, resto = _split_saludo(
        "Buen día. Los días 01/04/2026 ingresaron a la cuenta BANCO BOGOTA."
    )
    assert saludo.lower().startswith("buen día")
    assert "Los días" in resto
    saludo_legacy, _ = _split_saludo("Buenos días. El día 01/04/2026 ingresaron.")
    assert saludo_legacy.lower().startswith("buenos días")

    bank_dates = _dates_from_bank_email_table(
        ["Fecha", "Crédito", "Concepto", "Transacción"],
        [
            ["04/08/2026", "327", "SERGIO", "Cr Ach"],
            ["31/07/2026", "88888", "PARTE Identificar", "Cr Ach"],
        ],
    )
    assert date(2026, 7, 31) in bank_dates
    assert date(2026, 8, 4) in bank_dates
    merged = _format_fechas_validacion_es(
        [date(2026, 8, 4), date(2026, 8, 6)] + bank_dates
    )
    assert "31/07/2026" in merged
    assert "04/08/2026" in merged


def test_resolve_body_intro_template_rejects_legacy_and_mojibake():
    from app.application.use_cases.send_validar_extractos_notification import (
        _resolve_body_intro_template,
    )

    default = (
        "Buen día. {fechas_clause} ingresaron a la cuenta {banco} los siguientes valores, "
        "que corresponden a:"
    )
    ok = (
        "Buen día. {fechas_clause} ingresaron a la cuenta {banco} los siguientes valores, "
        "que corresponden a:"
    )
    assert _resolve_body_intro_template(ok, default=default) == ok
    assert _resolve_body_intro_template("", default=default) == default
    assert (
        _resolve_body_intro_template(
            "Buenos días. El día {fecha} ingresaron a la cuenta {banco}.",
            default=default,
        )
        == default
    )
    assert (
        _resolve_body_intro_template(
            "Buenos dÃ­as. El dÃ­a {fecha} ingresaron a la cuenta {banco}.",
            default=default,
        )
        == default
    )


def test_parse_bank_report_skips_template_example_row():
    from openpyxl import Workbook

    from app.application.use_cases.send_validar_extractos_notification import (
        _is_bank_template_example_row,
        _parse_bank_report_table_and_min_date,
    )

    assert _is_bank_template_example_row(
        ["ejemplo: 23-abr", "ejemplo: $1", "ejemplo: CLIENTE", "PAGO", "ejemplo: txt"]
    )
    assert not _is_bank_template_example_row(
        ["10/05/2026", "100", "CLIENTE", "PAGO", "STRESS"]
    )

    wb = Workbook()
    ws = wb.active
    ws.append(["Fecha", "Crédito", "Concepto", "Tipo Aplicación", "Transacción"])
    ws.append(
        [
            "ejemplo: 23-abr",
            "ejemplo: $49.538.473,00",
            "ejemplo: EQUINORTE",
            "PAGO",
            "ejemplo: Ach Bancolombia",
        ]
    )
    ws.append(["10/05/2026", "10252756", "A&M CONSTRUCOL", "PAGO", "STRESS PAGO"])
    buf = BytesIO()
    wb.save(buf)
    report_d, headers, rows = _parse_bank_report_table_and_min_date(buf.getvalue())
    assert report_d == date(2026, 5, 10)
    assert headers[0].casefold() == "fecha"
    assert len(rows) == 1
    assert rows[0][2] == "A&M CONSTRUCOL"
    assert all(not str(c).casefold().startswith("ejemplo") for row in rows for c in row)

