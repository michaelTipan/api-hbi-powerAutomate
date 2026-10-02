"""Tablas con la estructura real (fórmulas de cronograma, guardadas sin caché).

Las tablas que sube Apply las guarda openpyxl: pierden los valores calculados de las
fórmulas. Estos tests reproducen la estructura de las tablas reales (fila de
encabezado 4, número de cuota en A con referencia a otra hoja, día/mes/año en L:N con
fórmulas encadenadas) para cubrir:

- crédito ya aplicado → ALREADY_APPLIED antes de buscar la fila de la fecha límite;
- evaluación de fórmulas simples cuando falta la caché;
- día 31 en meses cortos (último día del mes, día habitual de cuota);
- fila de la fecha límite no encontrada con / sin pagos previos;
- la columna IBR nunca se modifica.
"""

from __future__ import annotations

import asyncio
import io
import json
from calendar import monthrange
from datetime import date, datetime

import openpyxl
import pytest
from openpyxl.utils.datetime import from_excel

from app.application.services.accounting_pdf_parser import ACCOUNT_VALOR_PAGADO_CLIENTE
from app.application.services.amortization_formula_eval import (
    FormulaNotEvaluable,
    SimpleFormulaEvaluator,
)
from app.application.services.amortization_workbook import (
    AUTOMATION_LOG_SHEET,
    _build_monthly_schedule,
    _row_due_date,
    detect_headers,
    enable_formula_fallback,
)
from app.application.use_cases.amortization_fill_apply import run_amortization_fill_apply
from app.application.use_cases.amortization_fill_dry_run import run_amortization_fill_dry_run
from tests.test_amortization_fill_apply import MockGraphApply
from tests.test_amortization_fill_dry_run import (
    MockGraphDryRun,
    _asiento_pdf_placeholder,
    _hist_bytes,
)

SHEET = " ASFALTOS MEDELLIN SAS"
HEADER_ROW = 4
FIRST_ROW = 5
LAST_ROW = 40  # 36 cuotas
TABLA = "TABLAS/amort.xlsx"
CREDITO = "CREDITO # 258"
ID_PAGO = "7785e37e"
ASIENTO = "clientes/EQUINORTE/CREDITO # 258/ASIENTOS CONTABLES CRED 258/asiento.pdf"
FECHA = date(2026, 9, 30)
DUE_ROW = 24  # cuota 20: 30/09/2026
IBR_PAST = 0.11232
IBR_FUTURE = 0.08851

_HEADERS = {
    2: "Ki",
    3: "i",
    4: "k",
    5: "Cuota",
    6: "KF",
    7: "IBR +i",
    8: "Spread",
    9: "Tasa + i",
    10: "Cuota + i",
    11: "Plazo + i",
    12: "dia",
    13: "mes",
    14: "año",
    15: "dia",
    16: "Causac Inter Mes",
    17: "Fecha pago",
    18: "Valor intereses ",
    19: "Abono a K",
    20: "Abono a capital",
    21: "Intereses de mora",
    22: "Valor pagado cliente",
    23: "Saldos menores",
}


@pytest.fixture(autouse=True)
def env_sharepoint(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")


@pytest.fixture(autouse=True)
def stub_process_control_writes(monkeypatch):
    async def _ok_update(*_a, **_k):
        return None

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_apply.update_process_control_row2",
        _ok_update,
    )


def _cuota_date(index: int, *, second_day: int) -> date:
    """Fecha de la cuota ``index`` (0 = 28/02/2025) con el día habitual ``second_day``."""
    if index == 0:
        return date(2025, 2, 28)
    month0 = 1 + index  # febrero (índice 1) + index
    year = 2025 + month0 // 12
    month = month0 % 12 + 1
    return date(year, month, min(second_day, monthrange(year, month)[1]))


def _table_bytes(
    *,
    second_day: int = 30,
    paid_through_row: int = DUE_ROW - 1,
    date_formulas: bool = True,
) -> bytes:
    """Tabla de 36 cuotas como las reales (sin caché: la guarda openpyxl)."""
    wb = openpyxl.Workbook()
    tasa = wb.active
    tasa.title = "TasaImplicita"
    tasa["A12"] = 1
    ws = wb.create_sheet(SHEET)
    wb.create_sheet("Hoja1")
    ws["B1"] = 245_629_566
    ws["B2"] = 36
    for col, label in _HEADERS.items():
        ws.cell(HEADER_ROW, col, label)
    for r in range(FIRST_ROW, LAST_ROW + 1):
        if r == FIRST_ROW:
            ws.cell(r, 1, "=+TasaImplicita!A12")
        else:
            ws.cell(r, 1, f'=IF(AND(A{r-1}<$B$2)*(A{r-1}<>0),1+A{r-1}," ")')
        ws.cell(r, 7, IBR_PAST if r < DUE_ROW else IBR_FUTURE)
        if not date_formulas:
            ws.cell(r, 12, "=DIA_PAGO")
            ws.cell(r, 13, "=MES_PAGO")
            ws.cell(r, 14, "=ANIO_PAGO")
        elif r == FIRST_ROW:
            ws.cell(r, 12, 28)
            ws.cell(r, 13, 2)
            ws.cell(r, 14, 2025)
        else:
            if r == FIRST_ROW + 1:
                ws.cell(r, 12, second_day)
            else:
                ws.cell(r, 12, f'=IF(A{r}<>" ",+L{r-1}," ")')
            ws.cell(r, 13, f'=IF(AND(A{r}<>" ")*(M{r-1}=12),1,1+M{r-1})')
            ws.cell(r, 14, f'=IF(AND(A{r}<>" ")*(M{r}=1),1+N{r-1},N{r-1})')
        if r <= paid_through_row:
            ws.cell(r, 17, datetime.combine(_cuota_date(r - FIRST_ROW, second_day=second_day), datetime.min.time()))
            ws.cell(r, 18, 2_000_000)
            ws.cell(r, 19, 6_000_000)
            ws.cell(r, 20, f"=+T{r-1}-S{r}" if r > FIRST_ROW else f"=+B1-S{r}")
            ws.cell(r, 22, f"=+R{r}+S{r}+U{r}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _accounting_text(fecha: date = FECHA) -> str:
    return f"""
    Comprobante 77 Fecha {fecha.strftime('%d/%m/%Y')}
    {ACCOUNT_VALOR_PAGADO_CLIENTE} 8,000,000.00
    544113410519 6,500,000.00
    544113430501 0.00
    544141502030 1,500,000.00
    """


def _files(table: bytes, *, fecha_limite: date = FECHA, manifest_extra: dict | None = None) -> dict[str, bytes]:
    output = {
        "id_pago": ID_PAGO,
        "cliente": "EQUINORTE",
        "credito": CREDITO,
        "asiento_pdf_path": ASIENTO,
        "extracto_pdf_path": "clientes/EQUINORTE/extracto.pdf",
        "output_relative_path": "OUT/consolidado.pdf",
    }
    output.update(manifest_extra or {})
    manifest = {
        "report_date_iso": FECHA.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [output],
        "skipped": [],
    }
    return {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{FECHA.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": _hist_bytes(
            ID_PAGO, CREDITO, TABLA, fecha_limite, fecha_banco=FECHA, monto_banco=8_000_000.0
        ),
        TABLA: table,
        ASIENTO: _asiento_pdf_placeholder(),
    }


@pytest.fixture
def accounting_text(monkeypatch):
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )


def _dry_run(files: dict[str, bytes]) -> dict:
    return asyncio.run(
        run_amortization_fill_dry_run(
            MockGraphDryRun(files),
            report_date_iso=FECHA.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )


def _apply(files: dict[str, bytes]) -> tuple[dict, MockGraphApply]:
    g = MockGraphApply(files)
    out = asyncio.run(
        run_amortization_fill_apply(
            g,
            report_date_iso=FECHA.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    return out, g


def _brief(item: dict) -> dict:
    keys = (
        "application_status",
        "error_code",
        "user_message",
        "due_date_row",
        "application_row",
        "warnings",
        "internal_notes",
    )
    return {k: item.get(k) for k in keys}


def _sheet(raw: bytes, *, data_only: bool):
    return openpyxl.load_workbook(io.BytesIO(raw), data_only=data_only)[SHEET]


# --- Evaluador de fórmulas -------------------------------------------------------


def _formula_sheets():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Main"
    ws["A1"] = 30
    ws["B1"] = "=A1"
    ws["C1"] = "=+A1+1"
    ws["D1"] = "=DATE(2026,9,31)"
    ws["E1"] = "=EDATE(DATE(2026,1,31),1)"
    ws["F1"] = "=SUM(A1:C1)"
    ws["G1"] = "=+Otra!A1*2"
    ws["H1"] = "='Hoja con espacio'!$B$2"
    ws["I1"] = "=PLAZO_TOTAL"
    ws["J1"] = '=IF(A1>10,"x"," ")'
    ws["K1"] = "=K1"
    otra = wb.create_sheet("Otra")
    otra["A1"] = 21
    esp = wb.create_sheet("Hoja con espacio")
    esp["B2"] = 7
    return wb


def test_evaluador_resuelve_formulas_simples_y_otra_hoja():
    wb = _formula_sheets()
    evaluators: dict[str, SimpleFormulaEvaluator] = {}

    def resolver(title: str):
        if title not in wb.sheetnames:
            return None
        if title not in evaluators:
            evaluators[title] = SimpleFormulaEvaluator(wb[title], sheet_resolver=resolver)
        return evaluators[title]

    ev = resolver("Main")
    assert ev.value(1, 2) == 30
    assert ev.value(1, 3) == 31
    # Serial de Excel (celda sin formato de fecha); misma semántica que Excel.
    assert from_excel(ev.value(1, 4)).date() == date(2026, 10, 1)
    assert from_excel(ev.value(1, 5)).date() == date(2026, 2, 28)
    assert ev.value(1, 6) == 91
    assert ev.value(1, 7) == 42
    assert ev.value(1, 8) == 7
    assert ev.value(1, 10) == "x"


@pytest.mark.parametrize("col", [9, 11])
def test_evaluador_no_inventa_valores(col):
    """Nombres definidos y ciclos no se evalúan: el llamador conserva su comportamiento."""
    ev = SimpleFormulaEvaluator(_formula_sheets()["Main"])
    with pytest.raises(FormulaNotEvaluable):
        ev.value(1, col)


def test_evaluador_sin_resolver_no_lee_otras_hojas():
    ev = SimpleFormulaEvaluator(_formula_sheets()["Main"])
    with pytest.raises(FormulaNotEvaluable):
        ev.value(1, 7)


# --- Fechas del cronograma sin caché (ítems 3 y 4) ---------------------------------


@pytest.mark.parametrize(
    ("second_day", "expected"),
    [(30, date(2026, 9, 30)), (31, date(2026, 9, 30))],
    ids=["dia_30_como_151", "dia_31_como_148"],
)
def test_fecha_de_fila_sin_cache_evalua_formulas(second_day, expected):
    raw = _table_bytes(second_day=second_day)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb[SHEET]
    headers = detect_headers(ws, header_row=HEADER_ROW)
    assert ws.cell(DUE_ROW, 12).value is None  # sin caché
    assert _row_due_date(ws, DUE_ROW, headers) is None
    enable_formula_fallback(wb, raw)
    assert _row_due_date(ws, DUE_ROW, headers) == expected
    assert _row_due_date(ws, FIRST_ROW, headers) == date(2025, 2, 28)
    # Febrero 2026 con día 30/31 → último día del mes.
    assert _row_due_date(ws, FIRST_ROW + 12, headers) == date(2026, 2, 28)


def test_cronograma_sin_ninguna_fecha_legible_no_inventa_fechas():
    raw = _table_bytes(second_day=31, date_formulas=False)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb[SHEET]
    enable_formula_fallback(wb, raw)
    headers = detect_headers(ws, header_row=HEADER_ROW)
    assert _row_due_date(ws, DUE_ROW, headers) is None
    schedule, _meta = _build_monthly_schedule(ws, headers, HEADER_ROW)
    assert DUE_ROW not in (schedule or {})


def test_cronograma_proyecta_dia_31_en_el_ultimo_dia_de_meses_cortos():
    """Solo las dos primeras cuotas tienen fecha legible (día 28 y 31): se proyecta el día 31."""
    raw = _table_bytes(second_day=31)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb[SHEET]
    headers = detect_headers(ws, header_row=HEADER_ROW)
    schedule, meta = _build_monthly_schedule(ws, headers, HEADER_ROW)
    assert meta.get("monthly_schedule_day") == 31
    assert schedule[DUE_ROW] == date(2026, 9, 30)
    assert schedule[DUE_ROW + 1] == date(2026, 10, 31)


@pytest.mark.parametrize("second_day", [30, 31], ids=["como_151", "como_148"])
def test_dry_run_tabla_sin_cache_encuentra_30_09_2026(second_day, accounting_text):
    out = _dry_run(_files(_table_bytes(second_day=second_day)))
    item = out["items"][0]
    assert item["application_status"] == "WOULD_APPLY", _brief(item)
    assert item["due_date_row"] == DUE_ROW
    assert item["application_row"] == DUE_ROW
    assert not item.get("warnings")


# --- Ya aplicado antes de buscar la fecha límite (ítem 1) ---------------------------


def _applied_table() -> bytes:
    """Tabla tal como la deja Apply (openpyxl, sin caché, con _AUTOMATION_LOG)."""
    out, g = _apply(_files(_table_bytes()))
    assert out["summary"]["applied"] == 1, out
    return g.uploaded[TABLA]


def test_credito_aplicado_guardado_sin_cache_sale_ya_aplicado(accounting_text):
    applied = _applied_table()
    assert _sheet(applied, data_only=True).cell(DUE_ROW, 12).value is None
    item = _dry_run(_files(applied))["items"][0]
    assert item["application_status"] == "ALREADY_APPLIED", _brief(item)
    assert not item.get("error_code")
    assert item["application_row"] == DUE_ROW


def test_credito_aplicado_sale_ya_aplicado_aunque_no_se_pueda_leer_la_fecha(accounting_text):
    """La identidad en el log decide antes que la búsqueda de la fila de vencimiento."""
    applied = _applied_table()
    wb = openpyxl.load_workbook(io.BytesIO(applied))
    ws = wb[SHEET]
    for r in range(FIRST_ROW, LAST_ROW + 1):
        ws.cell(r, 12, "=DIA_PAGO")
        ws.cell(r, 13, "=MES_PAGO")
        ws.cell(r, 14, "=ANIO_PAGO")
    buf = io.BytesIO()
    wb.save(buf)
    item = _dry_run(_files(buf.getvalue()))["items"][0]
    assert item["application_status"] == "ALREADY_APPLIED", _brief(item)
    assert item.get("error_code") != "DUE_DATE_ROW_NOT_FOUND"


def test_reintento_de_apply_no_duplica_la_linea(accounting_text):
    applied = _applied_table()
    out, g = _apply(_files(applied))
    assert out["summary"].get("applied", 0) == 0
    ws = _sheet(g.files[TABLA], data_only=False)
    assert ws.cell(DUE_ROW + 1, 17).value is None
    log = openpyxl.load_workbook(io.BytesIO(g.files[TABLA]))[AUTOMATION_LOG_SHEET]
    assert log.max_row == 2  # una sola fila de evento


# --- Fila de la fecha límite no encontrada (ítem 5) ---------------------------------


def test_sin_fila_de_fecha_limite_con_pagos_previos_no_bloquea(accounting_text):
    """El extracto trae una fecha que no está en el cronograma: se escribe tras el último pago."""
    item = _dry_run(_files(_table_bytes(), fecha_limite=date(2026, 9, 15)))["items"][0]
    assert item["application_status"] == "WOULD_APPLY", _brief(item)
    assert item["application_row"] == DUE_ROW
    assert item.get("due_date_row") is None
    assert not item.get("warnings")
    assert item.get("internal_notes")


def test_sin_fila_de_fecha_limite_sin_pagos_previos_sigue_bloqueando(accounting_text):
    raw = _table_bytes(paid_through_row=FIRST_ROW - 1)
    item = _dry_run(_files(raw, fecha_limite=date(2026, 9, 15)))["items"][0]
    assert item["application_status"] == "ERROR"
    assert item["error_code"] == "DUE_DATE_ROW_NOT_FOUND"


# --- IBR nunca se escribe (ítem 6) --------------------------------------------------


@pytest.mark.parametrize("actualiza_ibr", [True, False, None])
def test_apply_no_modifica_la_columna_ibr(actualiza_ibr, accounting_text, monkeypatch):
    # Aunque una política futura pidiera IBR, Apply no debe tocar la columna.
    monkeypatch.setattr(
        "app.application.services.review_schema.resolve_actualiza_ibr",
        lambda *_a, **_k: bool(actualiza_ibr),
    )
    raw = _table_bytes()
    before = _sheet(raw, data_only=False)
    ibr_before = [before.cell(r, 7).value for r in range(1, LAST_ROW + 1)]
    out, g = _apply(_files(raw, manifest_extra={"actualiza_ibr": actualiza_ibr}))
    assert out["summary"]["applied"] == 1, out
    after = _sheet(g.uploaded[TABLA], data_only=False)
    assert [after.cell(r, 7).value for r in range(1, LAST_ROW + 1)] == ibr_before
    assert after.cell(DUE_ROW, 7).value == IBR_FUTURE
    assert after.cell(DUE_ROW, 17).value is not None  # sí escribió el pago
    assert not g.put_calls or all("IBR_DIARIO" not in ep for ep, _ in g.put_calls)
