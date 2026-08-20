"""Notify Fase 3: filas unificadas Aplicacion_Pagos v3."""

import asyncio
from datetime import date
from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest
from openpyxl import Workbook, load_workbook

from app.application.services.historical_application_rows import (
    group_abono_rows_for_email,
    read_validated_abono_rows,
)
from app.application.services.review_schema import (
    AplicacionPagosCols,
    REVIEW_SCHEMA_VERSION,
    ReviewSheets,
    TipoAplicacionConfirmado,
    ValidarPago,
)
from app.application.use_cases.send_validar_extractos_notification import (
    _abono_table_from_groups,
    _build_html,
    send_validar_extractos_notification_email,
)
from tests.test_finalize_validation import make_distrib_row


def _bank_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Fecha", "Concepto", "Crédito", "Monto"])
    ws.append([date(2026, 6, 3), "pago", "258", "1000"])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _historico_v3_bytes() -> bytes:
    r, _ = make_distrib_row(
        id_pago="P1",
        validar_pago=ValidarPago.SI,
        ruta_pdf_internal="clientes/CLI/extractos/e1.pdf",
        tipo_aplicacion=TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL,
    )
    a1, _ = make_distrib_row(
        id_pago="AB123",
        cliente="EQUINORTE",
        credito="258",
        monto_banco=2_000_000,
        fecha_banco=date(2026, 6, 3),
        validar_pago=ValidarPago.SI,
        tipo_aplicacion=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
        abono_capital=1_000_000,
        valor_int=0,
        mora_a_aplicar=0,
    )
    a2, _ = make_distrib_row(
        id_pago="AB123",
        cliente="EQUINORTE",
        credito="265",
        monto_banco=2_000_000,
        fecha_banco=date(2026, 6, 3),
        validar_pago=ValidarPago.SI,
        tipo_aplicacion=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
        abono_capital=1_000_000,
        valor_int=0,
        mora_a_aplicar=0,
    )
    wb = Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(
        list(AplicacionPagosCols.HEADERS)
        + ["_ruta_extracto", "_ruta_unidad_credito", "_ruta_asientos_contables"]
    )
    ws.append(list(r) + ["clientes/CLI/ASIENTOS"])
    for row, cred in ((a1, "258"), (a2, "265")):
        ws.append(list(row) + [f"clientes/EQUINORTE/CREDITO# {cred}/ASIENTOS CONTABLES"])
    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append(["Campo", "Valor"])
    ws_meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_abono_table_single_row_per_id_pago():
    wb = load_workbook(BytesIO(_historico_v3_bytes()), data_only=True)
    groups = group_abono_rows_for_email(read_validated_abono_rows(wb))
    headers, rows = _abono_table_from_groups(groups)
    assert len(rows) == 1
    assert rows[0][0] == "AB123"
    assert "258" in rows[0][-1] and "265" in rows[0][-1]
    assert rows[0][3] == "2.000.000"


def test_build_html_omits_abono_section():
    html = _build_html(
        intro="<p>Hola</p>",
        bank_headers=["Fecha"],
        bank_rows=[["03/06/2026"]],
        abono_headers=["ID Pago", "Créditos seleccionados"],
        abono_rows=[["AB123", "258, 265"]],
    )
    assert "Abonos (sin extracto)" not in html
    assert "movimientos reportados con créditos seleccionados" not in html
    assert "AB123" not in html
    assert "background-color:#1F4E79" in html
    assert "Calibri,Segoe UI,Arial,sans-serif" in html
    assert "Reporte de pagos (Banco Bogotá)" not in html
    assert "03/06/2026" in html


class _NotifyGraph:
    def __init__(self, hist_bytes: bytes, extract_pdf: bytes = b"%PDF"):
        self.hist_bytes = hist_bytes
        self.extract_pdf = extract_pdf
        self.sent: list[dict] = []

    async def get(self, endpoint, params=None):
        if endpoint == "/sites":
            return {"value": [{"id": "s1"}]}
        if endpoint == "/sites/s1/drives":
            return {"value": [{"id": "d1", "name": "DRIVE"}]}
        return {"value": []}

    async def get_bytes(self, endpoint, params=None):
        ep = str(endpoint)
        if "HIST/" in ep or "cartera_validada" in ep:
            return self.hist_bytes
        if "extractos" in ep or "e1.pdf" in ep:
            return self.extract_pdf
        if ".pdf" in ep.lower():
            return b"%PDF-1.4 mock"
        return _bank_xlsx()

    async def post_json(self, endpoint, body):
        self.sent.append(body)
        return {}, 202

    async def put_bytes(self, endpoint, content, content_type=None):
        return {}


@pytest.fixture(autouse=True)
def _notify_env(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "SITIO")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "DRIVE")
    monkeypatch.setenv("GRAPH_SHAREPOINT_FILE_PATH", "banco.xlsx")
    monkeypatch.setenv("GRAPH_BANK_PAYMENTS_FILE_PATH", "banco.xlsx")
    monkeypatch.setenv("GRAPH_VALIDAR_NOTIFY_ATTACH_PDFS", "false")
    monkeypatch.setenv("GRAPH_VALIDAR_NOTIFY_EXPORT_EMAIL_PDF", "false")


def test_notify_reads_unified_aplicacion_pagos(monkeypatch):
    g = _NotifyGraph(_historico_v3_bytes())

    async def _fake_resolve(*_a, **_k):
        return {"site_id": "s1", "drive_id": "d1", "path_encoded": "x", "file_path": "banco.xlsx"}

    async def _fake_correos(*_a, **_k):
        return "sender@hbi.test", ["to@hbi.test"]

    async def _fake_snap(*_a, **_k):
        from types import SimpleNamespace

        return SimpleNamespace(
            process_key="payment-validation|banco_bogota|2026-06-03",
            estado_proceso="PENDIENTE_NOTIFICAR",
            email_pdf_path="",
            notify_idempotency_key="",
            is_active=True,
            historical_file_path="HIST/cartera_validada.xlsx",
            control_file_path="CTL/x.xlsx",
            bank_code="banco_bogota",
            bank_name="Banco Bogotá",
        )

    monkeypatch.setattr(
        "app.application.use_cases.send_validar_extractos_notification.resolve_sharepoint_from_env",
        _fake_resolve,
    )
    monkeypatch.setattr(
        "app.application.use_cases.send_validar_extractos_notification._load_sender_and_recipients_from_correos_xlsx",
        _fake_correos,
    )
    with patch(
        "app.application.use_cases.payment_validation_process_control.read_process_control_snapshot",
        new_callable=AsyncMock,
        side_effect=_fake_snap,
    ), patch(
        "app.application.use_cases.payment_validation_process_control.update_process_control_row2",
        new_callable=AsyncMock,
        return_value=True,
    ):
        result = asyncio.run(
            send_validar_extractos_notification_email(
                g,
                bank_code="banco_bogota",
                historical_file_path="HIST/cartera_validada.xlsx",
            )
        )
    assert result is not None
    assert g.sent, "debe enviar correo"
