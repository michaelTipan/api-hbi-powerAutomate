"""Regresión: dry-run/adopt no bloquea créditos ya aplicados debajo del vencimiento."""

from __future__ import annotations

from datetime import date

import openpyxl

from app.application.services.accounting_pdf_parser import (
    ACCOUNT_CAPITAL,
    ACCOUNT_VALOR_PAGADO_CLIENTE,
)
from app.application.services.amortization_workbook import (
    ADOPTADO_EXISTENTE,
    APLICADO,
    REVISION_MANUAL,
    PaymentApplicationWriteOptions,
    compare_existing_application,
    detect_amortization_sheet,
    detect_headers,
    find_application_row_detailed,
    find_bottom_most_occupied_payment_row,
    write_payment_application,
)
from app.application.ui.amortization_operational_issues import (
    attach_operational_issues_to_amortization_result,
    build_operational_issues_from_amortization_result,
)
from tests.test_amortization_workbook import _sample_event


def test_adopt_matching_payment_written_below_due_despite_incomplete_orphan_on_due():
    """Caso NACHONAL/93: Apply escribió debajo del vencimiento; due quedó huérfana."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NACHONAL"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
            "Saldo a capital",
        ]
    )
    # Filas 2-16: cronograma previo sin aplicación.
    for r in range(2, 17):
        ws.append([r, 8, 2026, None, None, None, None, None, None])
    # Fila 17 = vencimiento de la cuota: huérfano incompleto (solo un monto suelto).
    ws.append([7, 9, 2026, None, None, 1.0, None, None, None])
    # Filas 18-30: cronograma vacío de aplicación.
    for r in range(18, 31):
        ws.append([r - 10, 9, 2026, None, None, None, None, None, None])
    # Fila 31: hueco para escribir (con cronograma).
    ws.append([10, 10, 2026, None, None, None, None, None, None])

    headers = detect_headers(ws, header_row=1)
    ev = _sample_event(
        intereses=4288284.0,
        capital=7299796.0,
        mora=0.0,
        valor_pagado_cliente=11588080.0,
    )
    opts = PaymentApplicationWriteOptions(
        payment_date=date(2026, 9, 7),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    write_payment_application(
        ws, 31, headers, ev, write_options=opts, header_row=1
    )
    assert compare_existing_application(
        ws,
        31,
        headers,
        ev,
        payment_date=date(2026, 9, 7),
        detected_codes=opts.detected_codes,
    ) == ADOPTADO_EXISTENTE

    # Sin el fix: start cae en due=17 (huérfano) → REVISION_MANUAL y no ve la 31.
    result = find_application_row_detailed(
        ws,
        headers,
        ev,
        due_date_row=17,
        header_row=1,
        payment_date=date(2026, 9, 7),
        detected_codes=opts.detected_codes,
    )
    assert result.row == 31
    assert result.compare_status == ADOPTADO_EXISTENTE
    assert result.compare_status != REVISION_MANUAL


def test_revision_manual_issue_points_to_table_not_asiento_pdf():
    result = {
        "outcome": "requires_correction",
        "can_apply": False,
        "items": [
            {
                "id_pago": "P93",
                "credito": "93",
                "cliente": "NACHONAL",
                "application_status": "REVISION_MANUAL",
                "error_code": None,
                "tabla_amortizacion_path": (
                    "clientes/N/NACHONAL/tabla de amortizacion NACHONAL.xlsx"
                ),
                "asiento_pdf_path": "clientes/N/NACHONAL/ASIENTOS/93.pdf",
            }
        ],
    }
    enriched = attach_operational_issues_to_amortization_result(result)
    issues = enriched.get("operational_issues") or []
    assert len(issues) == 1
    issue = issues[0]
    assert "Tabla de amortización" in issue["title"]
    assert "93" in issue["title"]
    assert "93.pdf" not in (issue.get("user_message") or "")
    refs = {str(issue.get("technical_reference") or "").upper()}
    assert "PREFLIGHT_REVISION_MANUAL" in refs or "PREFLIGHT_REVISION_MANUAL" in {
        str(i.get("technical_reference") or "").upper() for i in issues
    }
    links = issue.get("links") or []
    assert any(l.get("rel") == "amortization_table" for l in links)
    assert not any(l.get("rel") == "asientos" for l in links)


def test_already_applied_counts_as_ready_in_summarize():
    from app.application.use_cases.amortization_fill_dry_run import _summarize

    items = [
        {"application_status": "ALREADY_APPLIED"},
        {"application_status": "WOULD_APPLY"},
    ]
    summary = _summarize(items)
    assert summary["would_adopt_existing"] == 1
    assert summary["would_apply"] == 1
    assert summary["revision_manual"] == 0
    assert summary["errors"] == 0


def test_false_automation_log_on_empty_row_surfaces_growth_blocked():
    """Log ADOPTADO en fila vacía + basura intermedia → APPLICATION_GROWTH_BLOCKED."""
    from app.application.services.amortization_workbook import (
        APPLICATION_GROWTH_BLOCKED,
        AUTOMATION_LOG_SHEET,
        append_automation_log,
        ensure_automation_log,
        find_application_row_detailed,
        is_payment_application_empty,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
        ]
    )
    ws.append([10, 8, 2026, date(2026, 8, 10), 10.0, 20.0, None, 30.0])  # 2 válido
    ws.append([None, None, None, date(2026, 6, 1), date(2026, 7, 1), None, None, None])  # 3 basura
    ws.append([10, 9, 2026, None, None, None, None, None])  # 4 vacía
    ws.append([10, 10, 2026, None, None, None, None, None])  # 5 vacía (log mentía fila 5)

    headers = detect_headers(ws, header_row=1)
    ev = _sample_event(
        intereses=100.0,
        capital=200.0,
        mora=0.0,
        valor_pagado_cliente=300.0,
    )
    assert is_payment_application_empty(ws, 5, headers) is True

    ensure_automation_log(wb)
    append_automation_log(
        wb,
        {
            "timestamp": "2026-09-10T13:15:00-05:00",
            "id_pago": "P1",
            "cliente": "BENAIA",
            "credito": "295",
            "fila": 5,
            "application_row": 5,
            "accion": ADOPTADO_EXISTENTE,
            "estado": ADOPTADO_EXISTENTE,
            "idempotency_key": "P1|295|path||",
            "asiento_pdf_path": "path",
        },
    )
    assert AUTOMATION_LOG_SHEET in wb.sheetnames

    # Motor de fila: con basura en 3 no adopta ni escribe en 5.
    result = find_application_row_detailed(
        ws,
        headers,
        ev,
        due_date_row=4,
        header_row=1,
        payment_date=date(2026, 9, 7),
        detected_codes=frozenset({ACCOUNT_VALOR_PAGADO_CLIENTE, ACCOUNT_CAPITAL}),
    )
    assert result.growth_blocked is True
    assert result.growth_block_code == APPLICATION_GROWTH_BLOCKED
    assert result.growth_block_row == 3
    assert result.compare_status != ADOPTADO_EXISTENTE
