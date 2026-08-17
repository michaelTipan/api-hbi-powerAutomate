import asyncio
import httpx
import io
import os
from datetime import date
from pathlib import Path
from urllib.parse import unquote

import openpyxl
import pytest

from app.application.services.review_schema import (
    AplicacionPagosCols,
    AsientosPendientesCols,
    InternalPathCols,
    REVIEW_SCHEMA_VERSION,
    ReviewSheets,
    SUPPORT_NOT_APPLICABLE,
    TipoAplicacion,
    TipoAplicacionConfirmado,
    ValidarPago)
from app.application.use_cases.payment_validation_finalize import (
    SECRETARY_FIRST_DATA_ROW,
    SECRETARY_HEADER_ROW,
    SECRETARY_HEADERS,
    SECRETARY_INSTRUCTION,
    SECRETARY_SHEET,
    SECRETARY_TITLE,
    _parse_extract_date_from_filename,
    _secretary_link_visible_text,
    finalize_payment_validation)
from app.application.config.payment_validation_settings import (
    BANK_CODE_BANCOLOMBIA,
    BANK_CODE_BOGOTA,
    resolve_bank_control_file_path)
from app.application.use_cases.setup_merge_control_workbook import (
    _build_process_control_workbook_bytes)

def _norm_fill_rgb(cell) -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    rgb = fill.fgColor.rgb
    if rgb is None:
        return None
    s = str(rgb).upper()
    return s[-6:] if len(s) > 6 else s

def _cell_has_hyperlink(cell) -> bool:
    hl = getattr(cell, "hyperlink", None)
    if hl is None:
        return False
    return bool(getattr(hl, "target", None))

# Formato real HBI de ruta interna al drive (adjuntos por correo)
HBI_CLIENTS_ROOT = (
    "INFORMACION CREDITOS-CLIENTES/01 COMWARE AUTOMATIZACION - INFORMACION CREDITOS CLIENTES"
)
HBI_EXPECTED_RUTA = (
    "INFORMACION CREDITOS-CLIENTES/01 COMWARE AUTOMATIZACION - INFORMACION CREDITOS CLIENTES/"
    "EQUINORTE/CREDITO # 264/Extracto 2026-01-01 CREDITO # 264.pdf"
)
HBI_EXTRACT_PDF = "Extracto 2026-01-01 CREDITO # 264.pdf"

class MockGraphClient:
    def __init__(self):
        self.children: list = []
        self.downloaded_files: dict[str, bytes] = {}
        self.uploaded_files: dict[str, bytes] = {}
        self.fail_upload_path: str | None = None
        self.requested_file_paths: list[str] = []
        self.requested_endpoints: list[str] = []
        self.put_calls: list[tuple[str, str]] = []
        self.children_overrides: dict[str, list] | None = None
        self.children_get_paths: list[str] = []
        self.drive_item_responses: dict[str, dict] = {}
        # Si se define, get_bytes lanza HTTP 404 cuando esta subcadena aparece en la ruta (simula Graph).
        self.get_bytes_404_substring: str | None = None
        self.dynamic_folder_children: dict[str, list[dict]] = {}

    async def post_json(self, endpoint: str, body: dict):
        if "children" not in endpoint or "/root:/" not in endpoint:
            return {}, 201
        raw = endpoint.split("/root:/", 1)[1].rsplit(":/children", 1)[0]
        path_decoded = unquote(raw)
        name = str(body.get("name", "")).strip()
        item = {
            "name": name,
            "folder": {},
            "webUrl": f"https://mock.invalid/asientos/{path_decoded}/{name}",
        }
        self.dynamic_folder_children.setdefault(path_decoded, []).append(item)
        return item, 201

    async def get(self, endpoint, params=None):
        ep_base = endpoint.split("?", 1)[0]
        if ep_base in self.drive_item_responses:
            return self.drive_item_responses[ep_base]
        if "children" in endpoint and "/root:/" in endpoint:
            raw = endpoint.split("/root:/", 1)[1].rsplit(":/children", 1)[0]
            path_decoded = unquote(raw)
            self.children_get_paths.append(path_decoded)
            if self.children_overrides and path_decoded in self.children_overrides:
                return {"value": self.children_overrides[path_decoded]}
            builtin = {
                "clientes/CLI": [
                    {"name": "CRED", "folder": {}},
                    {"name": "CRED_A", "folder": {}},
                    {"name": "CRED_B", "folder": {}},
                ],
                "clientes/CLI/CRED": [
                    {"name": "ASIENTOS CONTABLES", "folder": {}, "webUrl": "https://mock.invalid/asientos/CRED"},
                ],
                "clientes/CLI/CRED_A": [],
                "clientes/CLI/CRED_B": [],
            }
            extra = self.dynamic_folder_children.get(path_decoded, [])
            if path_decoded in builtin:
                return {"value": list(builtin[path_decoded]) + extra}
            if not path_decoded.startswith("clientes/"):
                return {"value": self.children}
            return {"value": extra}
        if "/drives" in endpoint:
            return {"value": [{"id": "dummy_drive", "name": "DRIVE"}]}
        if "/sites" in endpoint:
            return {"value": [{"id": "dummy_site"}]}
        return {}

    async def get_bytes(self, endpoint, params=None):
        self.requested_endpoints.append(endpoint)
        file_path = endpoint.split("/root:/", 1)[1].rsplit(":/content", 1)[0]
        file_path = unquote(file_path)
        self.requested_file_paths.append(file_path)
        if self.get_bytes_404_substring and self.get_bytes_404_substring in file_path:
            req = httpx.Request("GET", "https://graph.microsoft.com/v1.0/mock")
            resp = httpx.Response(404, request=req)
            raise httpx.HTTPStatusError("404", request=req, response=resp)
        if file_path in self.downloaded_files:
            return self.downloaded_files[file_path]
        if file_path == resolve_bank_control_file_path(BANK_CODE_BOGOTA):
            return _build_process_control_workbook_bytes("banco_bogota", "Banco de Bogotá")
        if file_path == resolve_bank_control_file_path(BANK_CODE_BANCOLOMBIA):
            return _build_process_control_workbook_bytes("banco_bancolombia", "Bancolombia")
        return b""

    async def put_bytes(self, endpoint, content, content_type):
        file_path = endpoint.split("/root:/", 1)[1].rsplit(":/content", 1)[0]
        file_path = unquote(file_path)
        if self.fail_upload_path and self.fail_upload_path in file_path:
            raise Exception("Mock Upload Error")
        self.put_calls.append((endpoint, content_type))
        self.uploaded_files[file_path] = content
        return {"id": "new_file_id", "webUrl": f"https://mock.invalid/web?path={file_path}"}

    async def delete(self, endpoint: str) -> None:
        return None

def set_env_vars():
    os.environ["GRAPH_BANK_PAYMENTS_FILE_PATH"] = "banco.xlsx"
    os.environ["GRAPH_CLIENTS_BASE_PATH"] = "clientes"
    os.environ["GRAPH_SHAREPOINT_SITE_SEARCH"] = "SITIO"
    os.environ["GRAPH_SHAREPOINT_DRIVE_NAME"] = "DRIVE"
    os.environ["GRAPH_PAYMENT_VALIDATION_REVIEW_PATH"] = "revision"
    os.environ["GRAPH_VALIDATION_FILE_PREFIX"] = "val"
    os.environ["GRAPH_PAYMENT_VALIDATION_HISTORY_PATH"] = "history"

def make_distrib_row(
    id_pago="ID1",
    cliente="CLI",
    credito="CRED",
    fecha_banco=None,
    fecha_limite=None,
    valor_int=70,
    mora_a_aplicar=10,
    abono_capital=20,
    otros=0,
    abono_k=None,
    mora=None,
    monto_banco=100,
    estado=None,
    validar_pago=ValidarPago.SI,
    observacion="",
    extract_route="clientes/CLI/CRED/extractos/e1.pdf",
    link_tabla_display="Ver tabla",
    tabla_hyperlink_target: str | None = None,
    ruta_pdf_internal: str = "",
    ruta_unidad_credito: str = "",
    tipo_aplicacion=TipoAplicacionConfirmado.PAGO_OBLIGACION_ACTUAL):
    """Fila Aplicacion_Pagos v3. Params legacy (estado/otros) se ignoran."""
    _ = (estado, otros)
    fecha_banco = fecha_banco or date(2026, 5, 10)
    fecha_limite = fecha_limite or date(2026, 5, 15)
    if abono_k is not None:
        abono_capital = abono_k
    if mora is not None:
        mora_a_aplicar = mora

    def _n(x):
        if isinstance(x, (int, float)):
            return float(x)
        return 0.0

    a = _n(valor_int)
    v = _n(mora_a_aplicar)
    k = _n(abono_capital)
    total = a + v + k
    vals = {
        AplicacionPagosCols.ID_PAGO: id_pago,
        AplicacionPagosCols.CLIENTE: cliente,
        AplicacionPagosCols.CREDITO: credito,
        AplicacionPagosCols.MONTO_BANCO: monto_banco,
        AplicacionPagosCols.FECHA_BANCO: fecha_banco,
        AplicacionPagosCols.FECHA_LIMITE: fecha_limite,
        AplicacionPagosCols.DIAS_RESPECTO_VENCIMIENTO: 0,
        AplicacionPagosCols.VALOR_OBLIGACION_ACTUAL: 100,
        AplicacionPagosCols.SALDO_VENCIDO: "",
        AplicacionPagosCols.VALIDAR_PAGO: validar_pago,
        AplicacionPagosCols.TIPO_APLICACION: tipo_aplicacion if validar_pago == ValidarPago.SI else "",
        AplicacionPagosCols.LINK_EXTRACTO: extract_route,
        AplicacionPagosCols.LINK_TABLA: link_tabla_display,
        AplicacionPagosCols.LINK_CARPETA_CREDITO: "",
        AplicacionPagosCols.OBSERVACION: observacion,
    }
    row = [vals.get(c, "") for c in AplicacionPagosCols.HEADERS]
    row.append(ruta_pdf_internal)
    row.append(ruta_unidad_credito or f"clientes/{cliente}/{credito}")
    return row, tabla_hyperlink_target

def make_abono_row(
    id_pago="AB1",
    cliente="CLI",
    credito="CRED",
    monto_banco=100000,
    fecha_banco=None,
    validar_abono=ValidarPago.SI,
    observacion="",
    link_tabla="https://mock.invalid/tabla",
    link_carpeta="https://mock.invalid/carpeta",
    ruta_unidad_credito="clientes/CLI/CRED",
    ruta_tabla_amortizacion="clientes/CLI/CRED/tabla.xlsx",
    credito_normalizado=None,
    tabla_hyperlink_target: str | None = None):
    """Compat: abonos viven en Aplicacion_Pagos v3 (sin hoja Distribucion_Abonos)."""
    _ = (link_carpeta, ruta_tabla_amortizacion, credito_normalizado)
    return make_distrib_row(
        id_pago=id_pago,
        cliente=cliente,
        credito=credito,
        monto_banco=monto_banco,
        fecha_banco=fecha_banco,
        validar_pago=validar_abono,
        observacion=observacion,
        link_tabla_display=link_tabla,
        ruta_unidad_credito=ruta_unidad_credito,
        tipo_aplicacion=TipoAplicacionConfirmado.ABONO_A_CAPITAL,
        tabla_hyperlink_target=tabla_hyperlink_target,
        valor_int=0,
        mora_a_aplicar=0,
        abono_capital=monto_banco)

def create_review_workbook(
    procesar="SI",
    estado="EN_REVISION",
    casos_data=None,
    distrib_specs=None,
    abono_specs=None,
    include_control=True,
    include_distrib=True,
    errores_rows=None,
    include_errores=False):
    """Workbook de revisión (Aplicacion_Pagos + _Meta + Errores)."""
    _ = (procesar, estado, casos_data, abono_specs, include_control)
    wb = openpyxl.Workbook()

    if include_distrib:
        ws_dist = wb.active
        ws_dist.title = ReviewSheets.APLICACION_PAGOS
        ws_dist.append(list(AplicacionPagosCols.HEADERS) + [InternalPathCols.RUTA_EXTRACTO, InternalPathCols.RUTA_UNIDAD_CREDITO])
        specs = distrib_specs or []
        if not specs:
            r, _hl = make_distrib_row()
            specs = [(r, None)]
        for spec in specs:
            tabla_hl = None
            if isinstance(spec, tuple) and len(spec) == 2:
                row_vals, tabla_hl = spec
            else:
                row_vals = spec
            row_vals = list(row_vals)
            if len(row_vals) < len(AplicacionPagosCols.HEADERS):
                row_vals = row_vals + [""] * (len(AplicacionPagosCols.HEADERS) - len(row_vals))
            r = ws_dist.max_row + 1
            for c, v in enumerate(row_vals, start=1):
                ws_dist.cell(r, c, v)
            if tabla_hl:
                lt_col = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_TABLA) + 1
                ctab = ws_dist.cell(r, lt_col)
                ctab.hyperlink = tabla_hl
                if not ctab.value:
                    ctab.value = "Tabla"
    else:
        ws = wb.active
        ws.title = "OtroDist"

    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append(["Campo", "Valor"])
    ws_meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    ws_meta.sheet_state = "hidden"

    if include_errores or errores_rows:
        from app.application.services.review_schema import ErroresCols

        ws_err = wb.create_sheet(ReviewSheets.ERRORES)
        ws_err.append(list(ErroresCols.HEADERS))
        for row in errores_rows or []:
            padded = list(row) + [""] * (len(ErroresCols.HEADERS) - len(row))
            ws_err.append(padded[: len(ErroresCols.HEADERS)])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def _append_distrib_ruta_column(wb_bytes: bytes, ruta_by_row: dict[int, str]) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb[ReviewSheets.APLICACION_PAGOS]
    col = ws.max_column + 1
    ws.cell(1, col, InternalPathCols.RUTA_EXTRACTO)
    for r, val in ruta_by_row.items():
        ws.cell(r, col, val)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def _run(coro):
    return asyncio.run(coro)





def test_finalize_blocks_when_errores_sheet_has_open_cases():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        error_row = [
            "622c5c5c-3f87-4750-af64-8bbcd0a3262b",
            "EQUINORTE",
            "CREDITO # 265",
            "Extracto",
            "No se pudo leer la fecha límite de pago del extracto del crédito.",
            "Revise el PDF del extracto.",
            "SI, si persiste",
            "",
            "",
            "fecha_limite_extracto_not_readable",
        ]
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, None)],
            errores_rows=[error_row])
        with pytest.raises(ValueError, match=r"review_has_open_errors\|1"):
            await finalize_payment_validation(client, "val_latest.xlsx")

    _run(run_test())

def test_finalize_allows_empty_errores_sheet():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(mora=0, valor_int=100, abono_k=0)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r, None)],
            include_errores=True)
        res = await finalize_payment_validation(
            client, "val_latest.xlsx", process_date=date(2026, 5, 10), bank_code="banco_bogota"
        )
        assert res["status"] == "success"

    _run(run_test())



def test_finalize_validar_accepts_zero_intereses_mora():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(mora=0, valor_int=100, abono_k=0)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        res = await finalize_payment_validation(
            client, "val_latest.xlsx", process_date=date(2026, 5, 10), bank_code="banco_bogota"
        )
        assert res["status"] == "success"

    _run(run_test())







def test_finalize_ignores_manual_distribution_and_uses_bank_amount():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(valor_int=150, abono_k=0, mora=0)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["status"] == "success"

    _run(run_test())

def test_finalize_si_does_not_require_positive_manual_total():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(valor_int=0, abono_k=0, mora=0)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["status"] == "success"

    _run(run_test())





def test_finalize_multi_credito_same_pago_cuadra():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r1, h1 = make_distrib_row(
            id_pago="ID1",
            credito="CRED_A",
            valor_int=60,
            abono_k=0,
            mora=0,
            extract_route="clientes/CLI/CRED_A/e.pdf")
        r2, h2 = make_distrib_row(
            id_pago="ID1",
            credito="CRED_B",
            valor_int=40,
            abono_k=0,
            mora=0,
            extract_route="clientes/CLI/CRED_B/e.pdf")
        r1 = list(r1)
        r2 = list(r2)
        monto_col = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.MONTO_BANCO)
        r1[monto_col] = 100
        r2[monto_col] = None
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            casos_data=[["ID1", date(2026, 5, 10), "CLI", "concepto", 100, ""]],
            distrib_specs=[(r1, h1), (r2, h2)])
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["status"] == "success"
        assert res["validated_rows"] == 2

    _run(run_test())

def test_finalize_respects_validation_file_path_when_provided():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        exact_path = "revision/subcarpeta/val_manual.xlsx"
        r, _ = make_distrib_row()
        client.downloaded_files[exact_path] = create_review_workbook(distrib_specs=[(r, None)])
        res = await finalize_payment_validation(
            client,
            validation_file_path=exact_path,
            process_date=date(2026, 5, 10))
        assert res["status"] == "success"
        assert res["validation_file_path"] == exact_path
        hist_keys = [k for k in client.uploaded_files if "cartera_validada_banco_bogota_2026-05-10_" in k]
        assert hist_keys, f"expected unique hist name, got {list(client.uploaded_files)}"
        assert hist_keys[0].startswith("history/2026/05/2026-05-10/cartera_validada_banco_bogota_2026-05-10_")
        # id corto (8) en nombre; UUID completo solo en Control/revisión.
        assert hist_keys[0].endswith(".xlsx")
        assert len(hist_keys[0].rsplit("_", 1)[-1].removesuffix(".xlsx")) == 8

    _run(run_test())


def test_finalize_does_not_update_amortization_tables():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, th = make_distrib_row()
        tab_url = "https://x.sharepoint.com/u?u=https://host/root:/clientes/CLI/CRED/tabla.xlsx:/"
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, tab_url)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert not any("tabla" in k and k.endswith(".xlsx") for k in client.uploaded_files)

    _run(run_test())

def test_finalize_does_not_clean_banco_bogota():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["bank_cleaned"] is False
        assert "banco.xlsx" not in client.uploaded_files

    _run(run_test())

def test_finalize_valid_excel_no_longer_fails_on_amortization_row_errors():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["status"] == "success"

    _run(run_test())

def test_finalize_historical_workbook_has_exact_ruta_column():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert headers.count(InternalPathCols.RUTA_EXTRACTO) == 1

    _run(run_test())

def test_finalize_historical_workbook_populates_ruta_for_validar_rows():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        path = "clientes/CLI/CRED/docs/ex.pdf"
        r, _ = make_distrib_row(extract_route=path)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        cmap = {str(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
        assert ws.cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value == path

    _run(run_test())

def test_finalize_missing_ruta_for_validar_blocks():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(extract_route="")
        wb_b = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(wb_b))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        le = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, le, "")
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        with pytest.raises(ValueError, match="missing_extract_route"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())

def test_finalize_creates_secretary_support_workbook():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        assert SECRETARY_SHEET in wb.sheetnames

    _run(run_test())


def test_secretary_workbook_includes_extract_link():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, tab = make_distrib_row()
        ext_url = "https://host/path/root:/clientes/CLI/CRED/x.pdf:/"
        wb0 = create_review_workbook(distrib_specs=[(r, tab)])
        wb = openpyxl.load_workbook(io.BytesIO(wb0))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        cell = ws.cell(2, c)
        cell.value = "Abrir"
        cell.hyperlink = ext_url
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wbs = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        wss = wbs[SECRETARY_SHEET]
        col = SECRETARY_HEADERS.index("Link extracto") + 1
        assert _cell_has_hyperlink(wss.cell(SECRETARY_FIRST_DATA_ROW, col))

    _run(run_test())

def test_secretary_workbook_includes_amortization_table_link():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        tab_url = "https://z/root:/clientes/CLI/CRED/tabla.xlsx:/"
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, tab_url)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wbs = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        col = SECRETARY_HEADERS.index("Link tabla amortización") + 1
        assert _cell_has_hyperlink(wbs[SECRETARY_SHEET].cell(SECRETARY_FIRST_DATA_ROW, col))

    _run(run_test())

def test_secretary_workbook_includes_asientos_folder_link_when_http_url():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        client.children_overrides = {
            "clientes/CLI": [{"name": "CRED", "folder": {}}],
            "clientes/CLI/CRED": [
                {
                    "name": "ASIENTOS CONTABLES",
                    "folder": {},
                    "webUrl": "https://sharepoint/asientos/CRED",
                }
            ],
        }
        r, tab = make_distrib_row()
        tab_url = "https://z/root:/clientes/CLI/CRED/tabla.xlsx:/"
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, tab_url)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wbs = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        wss = wbs[SECRETARY_SHEET]
        col = SECRETARY_HEADERS.index("Link carpeta asientos contables") + 1
        c = wss.cell(SECRETARY_FIRST_DATA_ROW, col)
        assert _cell_has_hyperlink(c)
        assert "CRED" in str(c.value)

    _run(run_test())

def test_secretary_workbook_asientos_link_after_finalize_provisions_folder():
    """Finalize crea ASIENTOS CONTABLES CRED {n} + EXTRACTOS y deja enlace en soporte secretaría."""
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        client.children_overrides = {
            "clientes/CLI": [{"name": "CRED", "folder": {}}],
            "clientes/CLI/CRED": [],
        }
        r, tab = make_distrib_row()
        tab_url = "https://z/root:/clientes/CLI/CRED/tabla.xlsx:/"
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, tab_url)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert "clientes/CLI/CRED" in client.dynamic_folder_children
        created_names = {
            str(it.get("name") or "")
            for it in client.dynamic_folder_children.get("clientes/CLI/CRED", [])
        }
        assert any(n.startswith("ASIENTOS CONTABLES CRED") for n in created_names)
        assert "EXTRACTOS" in created_names
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wbs = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        wss = wbs[SECRETARY_SHEET]
        col = SECRETARY_HEADERS.index("Link carpeta asientos contables") + 1
        c = wss.cell(SECRETARY_FIRST_DATA_ROW, col)
        assert c.value
        assert _cell_has_hyperlink(c)
        assert "CRED" in str(c.value)

    _run(run_test())

def test_finalize_writes_ruta_asientos_contables_on_historical_distrib():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(mora=0, valor_int=100, abono_k=0)
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        hdr = 1
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == AplicacionPagosCols.ID_PAGO:
                hdr = r
                break
        cmap = {
            str(ws.cell(hdr, c).value or "").strip(): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(hdr, c).value
        }
        assert InternalPathCols.RUTA_ASIENTOS_CONTABLES in cmap
        dr = hdr + 1
        rv = str(ws.cell(dr, cmap[InternalPathCols.RUTA_ASIENTOS_CONTABLES]).value or "")
        assert "ASIENTOS CONTABLES CRED" in rv
        assert rv.startswith("clientes/CLI/CRED/")

    _run(run_test())

def test_finalize_fails_when_ruta_unidad_credito_missing():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(mora=0, valor_int=100, abono_k=0, ruta_unidad_credito="")
        wb_bytes = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        col_uc = None
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(1, c).value or "").strip() == InternalPathCols.RUTA_UNIDAD_CREDITO:
                col_uc = c
                break
        assert col_uc is not None
        ws.cell(2, col_uc, "")
        # también vaciar el fallback por si make_distrib_row rellenó default
        out = io.BytesIO()
        wb.save(out)
        client.downloaded_files["revision/val_latest.xlsx"] = out.getvalue()
        with pytest.raises(ValueError, match="missing_ruta_unidad_credito"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())

def test_secretary_workbook_has_professional_title_and_instruction():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        assert ws.cell(1, 1).value == SECRETARY_TITLE
        assert ws.cell(2, 1).value == SECRETARY_INSTRUCTION
        assert ws.cell(2, 1).font.size == 12

    _run(run_test())

def test_secretary_workbook_headers_start_on_row_3():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        got = [ws.cell(SECRETARY_HEADER_ROW, c).value for c in range(1, len(SECRETARY_HEADERS) + 1)]
        assert got == SECRETARY_HEADERS

    _run(run_test())

def test_secretary_workbook_data_starts_on_row_4():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(id_pago="PID-ROW4")
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        assert ws.cell(SECRETARY_FIRST_DATA_ROW, SECRETARY_HEADERS.index("ID Pago") + 1).value == "PID-ROW4"

    _run(run_test())

def test_secretary_workbook_has_freeze_panes_after_credito():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        assert ws.freeze_panes == f"E{SECRETARY_FIRST_DATA_ROW}"

    _run(run_test())

def test_secretary_workbook_has_no_autofilter():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        assert ws.auto_filter.ref is None

    _run(run_test())


def test_secretary_workbook_formats_date_columns():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        r = SECRETARY_FIRST_DATA_ROW
        for name in ("Fecha banco", "Fecha límite"):
            c = SECRETARY_HEADERS.index(name) + 1
            assert ws.cell(r, c).number_format == "yyyy-mm-dd"

    _run(run_test())

def test_secretary_workbook_preserves_clickable_links():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        tab_url = "https://z/root:/clientes/CLI/CRED/tabla.xlsx:/"
        ext_url = "https://host/path/root:/clientes/CLI/CRED/x.pdf:/"
        wb0 = create_review_workbook(distrib_specs=[(r, tab_url)])
        wb = openpyxl.load_workbook(io.BytesIO(wb0))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        cell = ws.cell(2, c)
        cell.value = "Abrir"
        cell.hyperlink = ext_url
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wbs = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        wss = wbs[SECRETARY_SHEET]
        r = SECRETARY_FIRST_DATA_ROW
        assert _cell_has_hyperlink(wss.cell(r, SECRETARY_HEADERS.index("Link extracto") + 1))
        assert _cell_has_hyperlink(wss.cell(r, SECRETARY_HEADERS.index("Link tabla amortización") + 1))
        carpeta = wss.cell(r, SECRETARY_HEADERS.index("Link carpeta asientos contables") + 1)
        assert _cell_has_hyperlink(carpeta)

    _run(run_test())

def test_secretary_workbook_uses_descriptive_link_text_with_credito():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(credito="264")
        tab_url = "https://z/root:/clientes/CLI/264/tabla.xlsx:/"
        ext_url = "https://host/path/root:/clientes/CLI/264/x.pdf:/"
        wb0 = create_review_workbook(distrib_specs=[(r, tab_url)])
        wb = openpyxl.load_workbook(io.BytesIO(wb0))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        cell = ws.cell(2, c)
        cell.value = "Abrir"
        cell.hyperlink = ext_url
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wss = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))[SECRETARY_SHEET]
        r = SECRETARY_FIRST_DATA_ROW
        assert wss.cell(r, SECRETARY_HEADERS.index("Link extracto") + 1).value == "Ver extracto crédito 264"
        assert wss.cell(r, SECRETARY_HEADERS.index("Link tabla amortización") + 1).value == "Ver tabla crédito 264"
        hl = wss.cell(r, SECRETARY_HEADERS.index("Link extracto") + 1).hyperlink
        assert hl is not None and str(hl.target).startswith("https://")
        for col_name in ("Link extracto", "Link tabla amortización"):
            cell = wss.cell(r, SECRETARY_HEADERS.index(col_name) + 1)
            rgb = str(getattr(cell.font.color, "rgb", "") or "").upper()
            assert rgb.endswith("0563C1")
            assert cell.font.underline == "single"
            fill_rgb = str(getattr(cell.fill.fgColor, "rgb", "") or "").upper()
            assert fill_rgb.endswith("E8F4FC")

    _run(run_test())

def test_secretary_link_visible_text_non_numeric_credito():
    assert _secretary_link_visible_text("extracto", "ACIMOR", "CLI") == "Ver extracto ACIMOR"

def test_secretary_workbook_header_navy_style():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        hcell = ws.cell(SECRETARY_HEADER_ROW, 1)
        assert _norm_fill_rgb(hcell) == "002060"
        assert hcell.font.bold is True
        font_rgb = getattr(hcell.font.color, "rgb", None) if hcell.font.color else None
        assert font_rgb in ("FFFFFF", "00FFFFFF", None)

    _run(run_test())

def test_secretary_workbook_client_block_borders():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r1, _ = make_distrib_row(id_pago="A1", cliente="CLI-A", credito="100")
        r2, _ = make_distrib_row(id_pago="A2", cliente="CLI-A", credito="101")
        r3, _ = make_distrib_row(id_pago="B1", cliente="CLI-B", credito="200")
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(
            distrib_specs=[(r1, None), (r2, None), (r3, None)]
        )
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        r_first = SECRETARY_FIRST_DATA_ROW
        r_second = r_first + 1
        r_third = r_first + 2
        assert ws.cell(r_first, 1).border.top.style == "thin"
        assert ws.cell(r_second, 1).border.top.style == "thin"
        assert ws.cell(r_third, 1).border.top.style == "medium"
        assert ws.cell(r_second, 1).border.bottom.style == "medium"
        assert ws.cell(r_third, 1).border.bottom.style == "medium"

    _run(run_test())

def test_secretary_workbook_keeps_expected_columns():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        headers = [ws.cell(SECRETARY_HEADER_ROW, c).value for c in range(1, len(SECRETARY_HEADERS) + 1)]
        assert headers == SECRETARY_HEADERS

    _run(run_test())

def test_secretary_workbook_total_row_with_sum():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        sec_key = next(k for k in client.uploaded_files if "soporte_asientos_contables_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[sec_key]))
        ws = wb[SECRETARY_SHEET]
        tcol = SECRETARY_HEADERS.index("Total validado") + 1
        trow = SECRETARY_FIRST_DATA_ROW + 1
        assert ws.cell(trow, 1).value == "Total validado general"
        f = ws.cell(trow, tcol).value
        assert isinstance(f, str) and f.startswith("=SUM(")

    _run(run_test())

def test_finalize_result_includes_secretary_file_url():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["secretary_file_url"] and "soporte_asientos_contables" in res["secretary_file_url"]

    _run(run_test())

def test_finalize_result_flags_amortization_and_bank_false():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["amortization_updated"] is False
        assert res["bank_cleaned"] is False

    _run(run_test())


def test_finalize_falls_back_to_latest_prefixed_review_file_when_none_provided():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.children = [
            {"name": "~$val_temp.xlsx", "lastModifiedDateTime": "2026-05-10T09:00:00Z"},
            {"name": "otro_archivo.xlsx", "lastModifiedDateTime": "2026-05-10T09:30:00Z"},
            {"name": "val_old.xlsx", "lastModifiedDateTime": "2026-05-09T09:00:00Z"},
            {"name": "val_new.xlsx", "lastModifiedDateTime": "2026-05-10T10:00:00Z", "file": {}},
        ]
        client.downloaded_files["revision/val_new.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        # Phase 2: Finalize sin body resuelve primero por control por banco.
        ctrl = _build_process_control_workbook_bytes("banco_bogota", "Banco de Bogotá")
        wb = openpyxl.load_workbook(io.BytesIO(ctrl), data_only=False)
        try:
            ws = wb["Procesos"]
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            col = {h: i + 1 for i, h in enumerate(headers)}
            ws.cell(2, col["EstadoProceso"], value="REVISION_CREADA")
            ws.cell(2, col["IsActive"], value=True)
            ws.cell(2, col["ValidationFilePath"], value="revision/val_new.xlsx")
            ws.cell(2, col["ProcessKey"], value="payment-validation|banco_bogota|2026-05-10")
            buf = io.BytesIO()
            wb.save(buf)
            client.downloaded_files[resolve_bank_control_file_path(BANK_CODE_BOGOTA)] = buf.getvalue()
        finally:
            wb.close()
        res = await finalize_payment_validation(client, process_date=date(2026, 5, 10))
        assert res["validation_file"] == "val_new.xlsx"

    _run(run_test())

def test_finalize_uses_existing_ruta_when_present():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(extract_route="")
        base = create_review_workbook(distrib_specs=[(r, None)])
        client.downloaded_files["revision/val_latest.xlsx"] = _append_distrib_ruta_column(
            base, {2: "INFORMACION/precargada/extracto.pdf"}
        )
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        wb = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        cmap = {str(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
        assert ws.cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value == "INFORMACION/precargada/extracto.pdf"

    _run(run_test())

def test_finalize_extracts_ruta_from_hyperlink_root_path():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(extract_route="")
        raw = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, c, "Ver extracto")
        ws.cell(2, c).hyperlink = (
            "https://tenant.sharepoint.com/sites/x/_layouts/15/embed.aspx?"
            "u=https%3A%2F%2Fhost%2Froot%3A%2FINFORMACION%20CREDITOS%2FCLI%2FCRED%2Fex%2Epdf%3A%2F"
        )
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        cmap = {
            str(w[ReviewSheets.APLICACION_PAGOS].cell(1, c).value): c
            for c in range(1, w[ReviewSheets.APLICACION_PAGOS].max_column + 1)
            if w[ReviewSheets.APLICACION_PAGOS].cell(1, c).value
        }
        rv = w[ReviewSheets.APLICACION_PAGOS].cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value
        assert rv.endswith("ex.pdf")
        assert "INFORMACION CREDITOS" in str(rv)

    _run(run_test())

def test_finalize_extracts_filename_from_sharepoint_doc_url():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        client.children_overrides = {
            "clientes/CLI": [{"name": "CRED", "folder": {}}],
            "clientes/CLI/CRED": [
                {"name": "ASIENTOS CONTABLES", "folder": {}},
                {"name": "Extracto Feb Exacto 1.pdf", "file": {}},
            ],
        }
        r, _ = make_distrib_row(extract_route="")
        raw = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, c, "Ver extracto")
        ws.cell(2, c).hyperlink = (
            "https://tenant.sharepoint.com/sites/acme/_layouts/15/Doc.aspx"
            "?sourcedoc=%7B111%7D&file=Extracto%20Feb%20Exacto%201.pdf"
        )
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        cmap = {
            str(w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value): col
            for col in range(1, w[ReviewSheets.APLICACION_PAGOS].max_column + 1)
            if w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value
        }
        rv = w[ReviewSheets.APLICACION_PAGOS].cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value
        assert rv.endswith("Extracto Feb Exacto 1.pdf")
        assert rv.startswith("clientes/CLI/CRED/")

    _run(run_test())

def test_finalize_resolves_ruta_from_graph_drive_item_url():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        client.drive_item_responses["/drives/dummy_drive/items/ITEM99"] = {
            "name": "desde_graph.pdf",
            "parentReference": {"path": "/drives/dummy_drive/root:/clientes/CLI/CRED"},
        }
        r, _ = make_distrib_row(extract_route="")
        raw = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, c, "Ver extracto")
        ws.cell(2, c).hyperlink = "https://graph.microsoft.com/v1.0/drives/dummy_drive/items/ITEM99"
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        cmap = {
            str(w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value): col
            for col in range(1, w[ReviewSheets.APLICACION_PAGOS].max_column + 1)
            if w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value
        }
        rv = w[ReviewSheets.APLICACION_PAGOS].cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value
        assert rv == "clientes/CLI/CRED/desde_graph.pdf"

    _run(run_test())

def test_finalize_resolves_extract_route_by_exact_filename_in_credit_folder():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        client.children_overrides = {
            "clientes/CLI": [{"name": "CRED", "folder": {}}],
            "clientes/CLI/CRED": [
                {"name": "Otro.pdf", "file": {}},
                {"name": "Extracto Match Exacto.pdf", "file": {}},
            ],
        }
        r, _ = make_distrib_row(extract_route="")
        raw = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, c, "Ver extracto")
        ws.cell(2, c).hyperlink = "https://x/Doc.aspx?file=Extracto%20Match%20Exacto.pdf"
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        cmap = {
            str(w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value): col
            for col in range(1, w[ReviewSheets.APLICACION_PAGOS].max_column + 1)
            if w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value
        }
        rv = w[ReviewSheets.APLICACION_PAGOS].cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value
        assert str(rv).endswith("Extracto Match Exacto.pdf")

    _run(run_test())


def test_finalize_preserves_link_extracto_and_link_tabla():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        ext_u = "https://keep-extract/root:/clientes/CLI/CRED/x.pdf:/"
        tab_u = "https://keep-tabla/root:/clientes/CLI/CRED/tabla.xlsx:/"
        r, _ = make_distrib_row(extract_route="")
        raw = create_review_workbook(distrib_specs=[(r, tab_u)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        ce = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ct = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_TABLA) + 1
        ws.cell(2, ce, "Ver extracto")
        ws.cell(2, ce).hyperlink = ext_u
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws2 = w[ReviewSheets.APLICACION_PAGOS]
        assert str(ws2.cell(2, ce).hyperlink.target) == ext_u
        assert str(ws2.cell(2, ct).hyperlink.target) == tab_u

    _run(run_test())


def test_finalize_does_not_reactivate_amortization_or_bank_cleanup():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row()
        client.downloaded_files["revision/val_latest.xlsx"] = create_review_workbook(distrib_specs=[(r, None)])
        res = await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        assert res["amortization_updated"] is False
        assert res["bank_cleaned"] is False
        assert not any("tabla" in k and k.endswith(".xlsx") for k in client.uploaded_files)
        assert "banco.xlsx" not in client.uploaded_files

    _run(run_test())

def test_finalize_preserves_real_hbi_ruta_format_when_existing():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        r, _ = make_distrib_row(extract_route="")
        base = create_review_workbook(distrib_specs=[(r, None)])
        client.downloaded_files["revision/val_latest.xlsx"] = _append_distrib_ruta_column(base, {2: HBI_EXPECTED_RUTA})
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws = w[ReviewSheets.APLICACION_PAGOS]
        cmap = {str(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
        assert ws.cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value == HBI_EXPECTED_RUTA
        assert HBI_EXPECTED_RUTA.endswith(".pdf")
        assert "CREDITO #" in HBI_EXPECTED_RUTA
        assert "INFORMACION CREDITOS-CLIENTES/" in HBI_EXPECTED_RUTA

    _run(run_test())

def test_finalize_resolves_real_hbi_extract_path_from_credit_folder():
    async def run_test():
        set_env_vars()
        os.environ["GRAPH_CLIENTS_BASE_PATH"] = HBI_CLIENTS_ROOT
        client = MockGraphClient()
        client.children_overrides = {
            f"{HBI_CLIENTS_ROOT}/EQUINORTE": [{"name": "CREDITO # 264", "folder": {}}],
            f"{HBI_CLIENTS_ROOT}/EQUINORTE/CREDITO # 264": [
                {"name": HBI_EXTRACT_PDF, "file": {}},
            ],
        }
        r, _ = make_distrib_row(
            cliente="EQUINORTE",
            credito="264",
            extract_route="",
            fecha_banco=date(2026, 1, 1))
        raw = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, c, "Ver extracto")
        ws.cell(2, c).hyperlink = "https://tenant.sharepoint.com/sites/x/_layouts/15/Doc.aspx?id=onlyId"
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        ws2 = w[ReviewSheets.APLICACION_PAGOS]
        cmap = {str(ws2.cell(1, col).value): col for col in range(1, ws2.max_column + 1) if ws2.cell(1, col).value}
        rv = ws2.cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value
        assert rv == HBI_EXPECTED_RUTA
        assert "EQUINORTE/" in str(rv)
        assert str(rv).endswith(HBI_EXTRACT_PDF)

    _run(run_test())

def test_finalize_ruta_is_drive_relative_not_web_url():
    async def run_test():
        set_env_vars()
        os.environ["GRAPH_CLIENTS_BASE_PATH"] = HBI_CLIENTS_ROOT
        client = MockGraphClient()
        client.children_overrides = {
            f"{HBI_CLIENTS_ROOT}/EQUINORTE": [{"name": "CREDITO # 264", "folder": {}}],
            f"{HBI_CLIENTS_ROOT}/EQUINORTE/CREDITO # 264": [
                {"name": HBI_EXTRACT_PDF, "file": {}},
            ],
        }
        r, _ = make_distrib_row(
            cliente="EQUINORTE",
            credito="264",
            extract_route="",
            fecha_banco=date(2026, 1, 1))
        raw = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, c, "Ver extracto")
        ws.cell(2, c).hyperlink = "https://long.web.url/share?a=1"
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        cmap = {
            str(w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value): col
            for col in range(1, w[ReviewSheets.APLICACION_PAGOS].max_column + 1)
            if w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value
        }
        rv = str(w[ReviewSheets.APLICACION_PAGOS].cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value)
        assert not rv.lower().startswith("http")
        assert "://" not in rv
        assert rv.endswith(".pdf")

    _run(run_test())

def test_finalize_prefers_explicit_ruta_column_over_blank_link_hyperlink_resolve():
    """Con Ruta rellena, no hace falta hyperlink ni texto parseable en Link extracto."""

    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        client.children_overrides = {
            "clientes/RUTP": [{"name": "CREDX", "folder": {}}],
            "clientes/RUTP/CREDX": [{"name": "Extracto rutp.pdf", "file": {}}],
        }
        r, _ = make_distrib_row(
            cliente="RUTP",
            credito="CREDX",
            extract_route="",
            ruta_pdf_internal="clientes/RUTP/CREDX/from_ruta_col.pdf")
        raw = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, c, "solo texto amigable sin link")
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        cmap = {
            str(w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value): col
            for col in range(1, w[ReviewSheets.APLICACION_PAGOS].max_column + 1)
            if w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value
        }
        rv = w[ReviewSheets.APLICACION_PAGOS].cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value
        assert str(rv).replace("\\", "/") == "clientes/RUTP/CREDX/from_ruta_col.pdf"

    _run(run_test())

def test_finalize_fallback_extractos_folder_before_credit_root():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        client.children_overrides = {
            "clientes/XSUB": [{"name": "CR1", "folder": {}}],
            "clientes/XSUB/CR1": [
                {"name": "EXTRACTOS", "folder": {}},
                {"name": "solo_en_raiz_sin_extract_kw.pdf", "file": {}},
            ],
            "clientes/XSUB/CR1/EXTRACTOS": [
                {"name": "Extracto subcarpeta obligacion.pdf", "file": {}},
            ],
        }
        r, _ = make_distrib_row(
            cliente="XSUB",
            credito="CR1",
            extract_route="",
            ruta_pdf_internal="",
            link_tabla_display="tab")
        raw = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        c = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, c, "sin ruta tecnica previa")
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))
        hist_key = next(k for k in client.uploaded_files if "cartera_validada_" in k)
        w = openpyxl.load_workbook(io.BytesIO(client.uploaded_files[hist_key]))
        cmap = {
            str(w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value): col
            for col in range(1, w[ReviewSheets.APLICACION_PAGOS].max_column + 1)
            if w[ReviewSheets.APLICACION_PAGOS].cell(1, col).value
        }
        rv = str(w[ReviewSheets.APLICACION_PAGOS].cell(2, cmap[InternalPathCols.RUTA_EXTRACTO]).value)
        assert "/EXTRACTOS/" in rv.replace("\\", "/")
        assert rv.endswith(".pdf")

    _run(run_test())

def test_finalize_missing_route_when_pdf_names_lack_extract_keyword():
    async def run_test():
        set_env_vars()
        client = MockGraphClient()
        client.children_overrides = {
            "clientes/YBAD": [{"name": "CY", "folder": {}}],
            "clientes/YBAD/CY": [
                {"name": "FACTURA_recibo_nomina_kw.pdf", "file": {}},
            ],
        }
        r, _ = make_distrib_row(
            cliente="YBAD",
            credito="CY",
            extract_route="",
            ruta_pdf_internal="",
            link_tabla_display="tab")
        raw = create_review_workbook(distrib_specs=[(r, None)])
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb[ReviewSheets.APLICACION_PAGOS]
        ce = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
        ws.cell(2, ce, "x")
        buf = io.BytesIO()
        wb.save(buf)
        client.downloaded_files["revision/val_latest.xlsx"] = buf.getvalue()
        with pytest.raises(ValueError, match="missing_extract_route"):
            await finalize_payment_validation(client, "val_latest.xlsx", process_date=date(2026, 5, 10))

    _run(run_test())

def test_finalize_parse_extract_date_from_real_hbi_filename():
    assert _parse_extract_date_from_filename(HBI_EXTRACT_PDF) == date(2026, 1, 1)
    assert _parse_extract_date_from_filename("Extracto 2026-01-01 CREDITO # 264.pdf") == date(2026, 1, 1)
