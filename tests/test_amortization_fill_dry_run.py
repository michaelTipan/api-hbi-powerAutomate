"""Tests dry-run amortización (sin escritura en SharePoint)."""

import asyncio
import io
import json
from datetime import date, timedelta
from urllib.parse import unquote

import openpyxl
import pytest
from pypdf import PdfWriter

from app.application.services.accounting_pdf_parser import (
    ACCOUNT_CAPITAL,
    ACCOUNT_VALOR_PAGADO_CLIENTE,
)
from app.application.use_cases.amortization_fill_dry_run import (
    PDF_TEXT_NOT_EXTRACTABLE,
    TABLE_PATH_NOT_FOUND,
    run_amortization_fill_dry_run,
)


def _asiento_pdf_placeholder() -> bytes:
    """PDF mínimo; el texto del asiento se inyecta vía patch en tests."""
    return _empty_pdf_bytes()


def _empty_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _accounting_text() -> str:
    return f"""
    Comprobante 99 Fecha 22/05/2026
    {ACCOUNT_VALOR_PAGADO_CLIENTE} 50,000,000.00
    544113410519 49,118,143.00
    544113430501 0.00
    544141502030 881,857.00
    """


def _hist_bytes(
    id_pago: str,
    credito: str,
    tabla_path: str,
    fecha_limite: date,
    *,
    fecha_banco: date | None = None,
    tipo_aplicacion: str = "PAGO DE OBLIGACIÓN ACTUAL",
    monto_banco: float = 50_000_000.0,
) -> bytes:
    from app.application.services.review_schema import (
        REVIEW_SCHEMA_VERSION,
        AplicacionPagosCols,
        ReviewSheets,
        ValidarPago,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(
        list(AplicacionPagosCols.HEADERS)
        + ["_ruta_extracto", "_ruta_unidad_credito", "_ruta_tabla_amortizacion", "_ruta_asientos_contables"]
    )
    banco = fecha_banco if fecha_banco is not None else fecha_limite
    vals = {h: "" for h in AplicacionPagosCols.HEADERS}
    vals.update(
        {
            AplicacionPagosCols.ID_PAGO: id_pago,
            AplicacionPagosCols.CLIENTE: "EQUINORTE",
            AplicacionPagosCols.CREDITO: credito,
            AplicacionPagosCols.MONTO_BANCO: monto_banco,
            AplicacionPagosCols.FECHA_BANCO: banco,
            AplicacionPagosCols.FECHA_LIMITE: fecha_limite,
            AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.SI,
            AplicacionPagosCols.TIPO_APLICACION: tipo_aplicacion,
            AplicacionPagosCols.LINK_TABLA: tabla_path or "Ver tabla",
        }
    )
    row = [vals[h] for h in AplicacionPagosCols.HEADERS]
    row.extend(["x.pdf", f"clientes/EQUINORTE/{credito}", tabla_path or "", ""])
    ws.append(row)
    if tabla_path:
        lt = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_TABLA) + 1
        cell = ws.cell(2, lt)
        cell.hyperlink = f"https://sharepoint/root:/{tabla_path.replace('/', '%2F')}:"
    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append(["Campo", "Valor"])
    ws_meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _amort_table_bytes(fecha_limite: date) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Amort"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "IBR +i",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
            "Saldos Menores",
        ]
    )
    ws.append([fecha_limite.day, fecha_limite.month, fecha_limite.year, None, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _amort_table_with_existing_payment(fecha_limite: date) -> bytes:
    raw = _amort_table_bytes(fecha_limite)
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb.active
    ws.cell(2, 9, value=50_000_000.0)
    ws.cell(2, 7, value=49_118_143.0)
    ws.cell(2, 6, value=0.0)
    ws.cell(2, 8, value=881_857.0)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _amort_table_with_different_payment(fecha_limite: date) -> bytes:
    raw = _amort_table_bytes(fecha_limite)
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb.active
    ws.cell(2, 9, value=1.0)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ibr_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IBR"
    ws.append(["Inicio", "Fin", "Valor"])
    ws.append([date(2026, 5, 1), date(2026, 5, 31), "10.580%"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class MockGraphDryRun:
    def __init__(self, files: dict[str, bytes]) -> None:
        self.files = files
        self.put_calls: list[tuple[str, bytes]] = []
        self.children: dict[str, list[dict]] = {}

    def _key(self, endpoint: str) -> str | None:
        if ":/content" not in endpoint:
            return None
        path = endpoint.split("/root:/", 1)[1].rsplit(":/content", 1)[0]
        return unquote(path)

    async def get(self, endpoint: str, params=None):
        if "children" in endpoint and "/root:/" in endpoint:
            folder = endpoint.split("/root:/", 1)[1].rsplit(":/children", 1)[0]
            folder = unquote(folder)
            return {"value": self.children.get(folder, [])}
        return {"value": [{"id": "site"}, {"id": "drive", "name": "Doc"}]}

    async def get_bytes(self, endpoint: str, params=None):
        key = self._key(endpoint)
        if key and key in self.files:
            return self.files[key]
        raise FileNotFoundError(key)

    async def put_bytes(
        self,
        endpoint: str,
        content: bytes,
        content_type: str = "",
        if_match: str | None = None,
    ):
        self.put_calls.append((endpoint, content))
        return {"id": "new"}


@pytest.fixture(autouse=True)
def env_sharepoint(monkeypatch):
    monkeypatch.setenv("GRAPH_SHAREPOINT_SITE_SEARCH", "TEST")
    monkeypatch.setenv("GRAPH_SHAREPOINT_DRIVE_NAME", "")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_CONTROL_PATH", "CTL")
    monkeypatch.setenv("GRAPH_PAYMENT_VALIDATION_LOGS_PATH", "LOGS")
    monkeypatch.setenv("GRAPH_IBR_DIARIO_PATH", "CTL/IBR_DIARIO.xlsx")


def _base_files(
    *,
    hist: bytes,
    amort: bytes,
    asiento_pdf: bytes,
    ibr: bytes | None = None,
    fecha: date = date(2026, 5, 22),
) -> dict[str, bytes]:
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 258",
                "asiento_pdf_path": (
                    "clientes/EQUINORTE/CREDITO # 258/ASIENTOS CONTABLES CRED 258/asiento.pdf"
                ),
                "extracto_pdf_path": "clientes/EQUINORTE/extracto.pdf",
                "output_relative_path": "OUT/consolidado.pdf",
            }
        ],
        "skipped": [],
    }
    manifest_key = f"LOGS/merge_manifest_{fecha.isoformat()}.json"
    files = {
        "CTL/dummy.xlsx": b"x",
        manifest_key: json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort.xlsx": amort,
        "clientes/EQUINORTE/CREDITO # 258/ASIENTOS CONTABLES CRED 258/asiento.pdf": asiento_pdf,
    }
    if ibr is not None:
        files["CTL/IBR_DIARIO.xlsx"] = ibr
    return files


def test_dry_run_happy_path_would_apply_and_write_ibr(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphDryRun(
        _base_files(hist=hist, amort=_amort_table_bytes(fecha), asiento_pdf=_asiento_pdf_placeholder(), ibr=_ibr_bytes())
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )

    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["status"] == "ok"
    assert out["mode"] == "dry_run"
    assert g.put_calls == []
    item = out["items"][0]
    assert item["application_status"] == "WOULD_APPLY"
    assert item["target_row"] == 2
    assert item["ibr"]["status"] == "WOULD_WRITE_IBR"
    assert item["ibr"]["found"] is True
    assert item["payment_application"]["valor_pagado_cliente"] == 50_000_000.0


def test_dry_run_pending_ibr_without_ibr_file(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    files = _base_files(hist=hist, amort=_amort_table_bytes(fecha), asiento_pdf=_asiento_pdf_placeholder(), ibr=None)
    files.pop("CTL/IBR_DIARIO.xlsx", None)
    g = MockGraphDryRun(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )

    out = asyncio.run(run_amortization_fill_dry_run(g, report_date_iso="2026-05-22"))
    assert out["items"][0]["ibr"]["status"] == "PENDING_IBR"
    assert out["items"][0]["ibr"]["found"] is False


def test_dry_run_table_path_not_found(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "", fecha)
    g = MockGraphDryRun(
        _base_files(hist=hist, amort=_amort_table_bytes(fecha), asiento_pdf=_asiento_pdf_placeholder())
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(run_amortization_fill_dry_run(g, report_date_iso="2026-05-22"))
    assert out["items"][0]["application_status"] == "ERROR"
    assert out["items"][0]["error_code"] == TABLE_PATH_NOT_FOUND


def test_dry_run_pdf_not_extractable():
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphDryRun(
        _base_files(
            hist=hist,
            amort=_amort_table_bytes(fecha),
            asiento_pdf=_empty_pdf_bytes(),
            ibr=_ibr_bytes(),
        )
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["items"][0]["error_code"] == PDF_TEXT_NOT_EXTRACTABLE


def test_dry_run_would_adopt_existing(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphDryRun(
        _base_files(
            hist=hist,
            amort=_amort_table_with_existing_payment(fecha),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(run_amortization_fill_dry_run(g, report_date_iso="2026-05-22"))
    assert out["items"][0]["application_status"] == "WOULD_ADOPT_EXISTING"


def test_dry_run_revision_manual(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphDryRun(
        _base_files(
            hist=hist,
            amort=_amort_table_with_different_payment(fecha),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(run_amortization_fill_dry_run(g, report_date_iso="2026-05-22"))
    assert out["items"][0]["application_status"] == "REVISION_MANUAL"


def test_dry_run_never_uploads_files(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphDryRun(
        _base_files(hist=hist, amort=_amort_table_bytes(fecha), asiento_pdf=_asiento_pdf_placeholder())
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    asyncio.run(run_amortization_fill_dry_run(g, report_date_iso="2026-05-22"))
    assert g.put_calls == []


_AMORT_HEADER_ROW = [
    "dia",
    "mes",
    "año",
    "IBR +i",
    "Fecha pago",
    "Valor intereses",
    "Abono a K",
    "intereses mora",
    "Valor pagado cliente",
    "Saldos Menores",
]


def _amort_table_with_op_formulas(
    fecha_limite: date,
    due_row: int,
    *,
    formula_through_row: int = 7,
    max_row: int | None = None,
    sheet_title: str = "EQUINORTE",
) -> bytes:
    """Tabla con fórmulas O:P en filas de plantilla (p. ej. =+C{r}/30 y =+O{r}*8)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(_AMORT_HEADER_ROW)
    last_row = max_row if max_row is not None else max(due_row, formula_through_row)
    while (ws.max_row or 1) < last_row:
        ws.append([None] * len(_AMORT_HEADER_ROW))
    ws.cell(due_row, 1, fecha_limite.day)
    ws.cell(due_row, 2, fecha_limite.month)
    ws.cell(due_row, 3, fecha_limite.year)
    for r in range(4, formula_through_row + 1):
        ws.cell(r, 3, 30.0)
        ws.cell(r, 15, f"=+C{r}/30")
        mult = 8 if r < 7 else 10
        ws.cell(r, 16, f"=+O{r}*{mult}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _amort_table_date_at_row(fecha_limite: date, data_row: int, *, sheet_title: str = "EQUINORTE") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(_AMORT_HEADER_ROW)
    while (ws.max_row or 1) < data_row - 1:
        ws.append([None] * len(_AMORT_HEADER_ROW))
    ws.append(
        [
            fecha_limite.day,
            fecha_limite.month,
            fecha_limite.year,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _amort_table_two_dates_at_rows(
    fecha_limite: date, row_a: int, row_b: int, *, sheet_title: str = "EQUINORTE"
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(_AMORT_HEADER_ROW)
    for target_row in (row_a, row_b):
        while (ws.max_row or 1) < target_row - 1:
            ws.append([None] * len(_AMORT_HEADER_ROW))
        ws.append(
            [
                fecha_limite.day,
                fecha_limite.month,
                fecha_limite.year,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hist_bytes_multi_credit(
    rows: list[tuple],
    *,
    monto_banco: float = 100_000_000.0,
) -> bytes:
    from app.application.services.review_schema import (
        REVIEW_SCHEMA_VERSION,
        AplicacionPagosCols,
        ReviewSheets,
        ValidarPago,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(
        list(AplicacionPagosCols.HEADERS)
        + ["_ruta_extracto", "_ruta_unidad_credito", "_ruta_tabla_amortizacion", "_ruta_asientos_contables"]
    )
    for id_pago, cliente, credito, tabla_path, fecha_limite in rows:
        vals = {h: "" for h in AplicacionPagosCols.HEADERS}
        vals.update(
            {
                AplicacionPagosCols.ID_PAGO: id_pago,
                AplicacionPagosCols.CLIENTE: cliente,
                AplicacionPagosCols.CREDITO: credito,
                AplicacionPagosCols.MONTO_BANCO: monto_banco,
                AplicacionPagosCols.FECHA_BANCO: fecha_limite,
                AplicacionPagosCols.FECHA_LIMITE: fecha_limite,
                AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.SI,
                AplicacionPagosCols.TIPO_APLICACION: "PAGO DE OBLIGACIÓN ACTUAL",
                AplicacionPagosCols.LINK_TABLA: tabla_path,
            }
        )
        row = [vals[h] for h in AplicacionPagosCols.HEADERS]
        row.extend(["x.pdf", f"clientes/{cliente}/{credito}", tabla_path, ""])
        ws.append(row)
        lt = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_TABLA) + 1
        cell = ws.cell(ws.max_row, lt)
        cell.hyperlink = f"https://sharepoint/root:/{tabla_path.replace('/', '%2F')}:"
    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append(["Campo", "Valor"])
    ws_meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _amort_table_two_rows_same_date(fecha_limite: date) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Amort"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "IBR +i",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "intereses mora",
            "Valor pagado cliente",
            "Saldos Menores",
        ]
    )
    ws.append([fecha_limite.day, fecha_limite.month, fecha_limite.year, None, None, None, None, None, None, None])
    ws.append([fecha_limite.day, fecha_limite.month, fecha_limite.year, None, None, None, None, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _erp_asiento_text(numero: str, lineas: str) -> str:
    """Asiento como lo entrega pypdf: consecutivo en la 1ª línea y fecha invertida."""
    return f"""{numero}
    Año       Mes      Día
23 4 2026 Fecha  : SIN ENTIDAD Entidad : 0 Soporte :
{lineas}
"""


_TEXTO_PAGO_CUOTA_265 = _erp_asiento_text(
    "3494",
    """1,041,446.00 PAGO: No.Rad. 265 Linea 544 1 11100505
463,050.00 PAGO: No.Rad. 265 Linea 544 1 13410519
578,111.00 PAGO: No.Rad. 265 Linea 544 1 13430501
285.00 PAGO: No.Rad. 265 Linea 544 1 41502030""",
)

_TEXTO_AJUSTE_SALDOS_265 = _erp_asiento_text(
    "3495",
    """285.00 PAGO: No.Rad. 265 Linea 544 1 53159505
285.00 PAGO: No.Rad. 265 Linea 544 1 13410519""",
)


def test_dry_run_orders_pago_cuota_before_saldos_menores_adjustment(monkeypatch):
    """
    Caso EQUINORTE 265 (2026-04-23): el ajuste de saldos menores se llamaba
    ``...credito-265-evento-2.pdf`` y ordenaba antes que el pago de la cuota por
    orden alfabético. El abono a capital de la primera fila alimenta el capital base
    del período siguiente, así que el orden tiene efecto financiero.
    """
    fecha = date(2026, 5, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 265",
        "TABLAS/amort.xlsx",
        fecha,
        monto_banco=1_041_731.0,
    )
    base = "clientes/EQUINORTE/ASIENTOS CONTABLES CRED 265"
    asiento_ajuste = f"{base}/asiento_banco_bogota_credito-265-evento-2.pdf"
    asiento_cuota = f"{base}/asiento_banco_bogota_credito-265.pdf"
    bytes_ajuste = _asiento_pdf_placeholder() + b"%AJUSTE"
    bytes_cuota = _asiento_pdf_placeholder() + b"%CUOTA"

    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                # Orden alfabético, tal como lo entrega el merge.
                "asiento_pdf_paths": [asiento_ajuste, asiento_cuota],
                "extracto_pdf_path": "clientes/EQUINORTE/extracto.pdf",
                "output_relative_path": "OUT/consolidado.pdf",
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort.xlsx": _amort_table_two_rows_same_date(fecha),
        asiento_ajuste: bytes_ajuste,
        asiento_cuota: bytes_cuota,
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda b: _TEXTO_AJUSTE_SALDOS_265 if b.endswith(b"%AJUSTE") else _TEXTO_PAGO_CUOTA_265,
    )

    out = asyncio.run(
        run_amortization_fill_dry_run(
            MockGraphDryRun(files),
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )

    assert len(out["items"]) == 2
    primero, segundo = out["items"]
    assert primero["asiento_pdf_path"] == asiento_cuota
    assert primero["event_index"] == 1
    assert primero["application_row"] == 2
    assert primero["payment_application"]["intereses"] == 578111.0
    assert primero["payment_application"]["capital"] == 463050.0

    assert segundo["asiento_pdf_path"] == asiento_ajuste
    assert segundo["event_index"] == 2
    assert segundo["application_row"] == 3
    assert segundo["payment_application"]["saldos_menores"] == 285.0


def test_dry_run_flags_payment_date_differing_from_asiento(monkeypatch):
    """
    Fecha pago = Fecha banco. Si el asiento trae otra fecha se marca mismatch
    (auditoría), sin warning fail-closed.
    """
    fecha = date(2026, 5, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 265",
        "TABLAS/amort.xlsx",
        fecha,
        monto_banco=1_041_446.0,
    )
    asiento = "clientes/EQUINORTE/ASIENTOS/asiento_credito-265.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_paths": [asiento],
                "extracto_pdf_path": "clientes/EQUINORTE/extracto.pdf",
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort.xlsx": _amort_table_two_rows_same_date(fecha),
        asiento: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _TEXTO_PAGO_CUOTA_265,
    )

    out = asyncio.run(
        run_amortization_fill_dry_run(
            MockGraphDryRun(files),
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )

    item = out["items"][0]
    assert item["fecha_asiento"] == "2026-04-23"
    assert item["payment_date_iso"] == "2026-05-22"
    assert item["payment_date_matches_asiento"] is False
    assert item["warnings"] == []


def test_dry_run_errors_when_fecha_banco_missing(monkeypatch):
    """Sin Fecha banco no se planifica escritura (prohibido usar report_date)."""
    from app.application.use_cases.amortization_fill_dry_run import _load_historical_index

    limite = date(2026, 5, 22)
    report = date(2026, 7, 29)
    # Histórico sin columna Fecha banco / valor vacío → error.
    wb = openpyxl.Workbook()
    ws = wb.active
    from app.application.services.review_schema import (
        REVIEW_SCHEMA_VERSION,
        AplicacionPagosCols,
        ReviewSheets,
        ValidarPago,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    ws.append(
        list(AplicacionPagosCols.HEADERS)
        + ["_ruta_extracto", "_ruta_unidad_credito", "_ruta_tabla_amortizacion"]
    )
    vals = {h: "" for h in AplicacionPagosCols.HEADERS}
    vals.update(
        {
            AplicacionPagosCols.ID_PAGO: "7785e37e",
            AplicacionPagosCols.CLIENTE: "EQUINORTE",
            AplicacionPagosCols.CREDITO: "CREDITO # 265",
            AplicacionPagosCols.MONTO_BANCO: 1_041_446.0,
            AplicacionPagosCols.FECHA_LIMITE: limite,
            # Sin Fecha banco a propósito
            AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.SI,
            AplicacionPagosCols.TIPO_APLICACION: "PAGO DE OBLIGACIÓN ACTUAL",
            AplicacionPagosCols.LINK_TABLA: "TABLAS/amort.xlsx",
        }
    )
    row = [vals[h] for h in AplicacionPagosCols.HEADERS]
    row.extend(["x.pdf", "clientes/EQUINORTE/265", "TABLAS/amort.xlsx"])
    ws.append(row)
    cell = ws.cell(2, AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_TABLA) + 1)
    cell.hyperlink = "https://sharepoint/root:/TABLAS%2Famort.xlsx:"
    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append(["Campo", "Valor"])
    ws_meta.append(["ReviewSchemaVersion", REVIEW_SCHEMA_VERSION])
    buf = io.BytesIO()
    wb.save(buf)
    hist = buf.getvalue()
    assert _load_historical_index(hist)[("7785e37e", "265")].get("fecha_banco") is None

    asiento = "clientes/EQUINORTE/ASIENTOS/asiento_credito-265.pdf"
    manifest = {
        "report_date_iso": report.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_paths": [asiento],
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{report.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort.xlsx": _amort_table_two_rows_same_date(limite),
        asiento: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _TEXTO_PAGO_CUOTA_265,
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            MockGraphDryRun(files),
            report_date_iso=report.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    item = out["items"][0]
    assert item["error_code"] == "FECHA_BANCO_REQUIRED"
    assert "payment_date_iso" not in item or not item.get("payment_date_iso")


def test_parse_date_value_accepts_excel_serial():
    from openpyxl.utils.datetime import to_excel
    from app.application.use_cases.amortization_fill_dry_run import (
        _parse_date_value,
        _resolve_pago_payment_date_iso,
    )

    banco = date(2026, 4, 23)
    serial = float(to_excel(banco))
    assert _parse_date_value(serial) == banco
    assert (
        _resolve_pago_payment_date_iso(hist_row={"fecha_banco": serial})
        == "2026-04-23"
    )


def test_dry_run_payment_date_prefers_fecha_banco_over_report_date(monkeypatch):
    """Fecha pago debe ser Fecha banco del histórico, no el report_date del lote."""
    limite = date(2026, 5, 22)
    banco = date(2026, 4, 23)
    report = date(2026, 7, 29)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 265",
        "TABLAS/amort.xlsx",
        limite,
        fecha_banco=banco,
        monto_banco=1_041_446.0,
    )
    asiento = "clientes/EQUINORTE/ASIENTOS/asiento_credito-265.pdf"
    manifest = {
        "report_date_iso": report.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_paths": [asiento],
                "extracto_pdf_path": "clientes/EQUINORTE/extracto.pdf",
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{report.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort.xlsx": _amort_table_two_rows_same_date(limite),
        asiento: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _TEXTO_PAGO_CUOTA_265,
    )

    out = asyncio.run(
        run_amortization_fill_dry_run(
            MockGraphDryRun(files),
            report_date_iso=report.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )

    item = out["items"][0]
    assert item["payment_date_iso"] == banco.isoformat()
    assert item["fecha_asiento"] == "2026-04-23"
    assert item["payment_date_matches_asiento"] is True


def test_dry_run_two_asientos_produce_two_events(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 258",
        "TABLAS/amort.xlsx",
        fecha,
        monto_banco=60_000_000.0,
    )
    asiento_a = "clientes/EQUINORTE/ASIENTOS/Asiento cuota credito 258.pdf"
    asiento_b = "clientes/EQUINORTE/ASIENTOS/Asiento abono capital credito 258.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 258",
                "asiento_pdf_paths": [asiento_a, asiento_b],
                "asiento_pdf_path": asiento_a,
                "extracto_pdf_path": "clientes/EQUINORTE/extracto.pdf",
                "output_relative_path": "OUT/consolidado.pdf",
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort.xlsx": _amort_table_two_rows_same_date(fecha),
        asiento_a: _asiento_pdf_placeholder(),
        asiento_b: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)
    texts = iter([_accounting_text(), _accounting_text().replace("50,000,000", "10,000,000")])

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: next(texts),
    )

    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["summary"]["total_events"] == 2
    assert len(out["items"]) == 2
    assert out["items"][0]["event_index"] == 1
    assert out["items"][1]["event_index"] == 2
    assert out["items"][0]["asiento_pdf_path"] != out["items"][1]["asiento_pdf_path"]
    assert out["items"][0]["due_date_row"] == 2
    assert out["items"][1]["due_date_row"] == 2
    assert out["items"][0]["application_row"] == 2
    assert out["items"][1]["application_row"] == 3


def test_dry_run_single_pdf_with_two_asientos_produces_two_rows(monkeypatch):
    """Un solo PDF multi-asiento (caso MADERPOL) → dos filas de aplicación."""
    from app.application.services.accounting_pdf_parser import parse_accounting_text

    fecha = date(2026, 5, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 258",
        "TABLAS/amort.xlsx",
        fecha,
        monto_banco=60_000_000.0,
    )
    asiento_combo = "clientes/EQUINORTE/ASIENTOS/Asiento combo 258.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 258",
                "asiento_pdf_paths": [asiento_combo],
                "asiento_pdf_path": asiento_combo,
                "extracto_pdf_path": "clientes/EQUINORTE/extracto.pdf",
                "output_relative_path": "OUT/consolidado.pdf",
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort.xlsx": _amort_table_two_rows_same_date(fecha),
        asiento_combo: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)

    ret_text = """
4120
21 7 2026 Fecha : SIN ENTIDAD
574,411.00 Retenciones factura 1 13551503
574,411.00 PAGO: No.Rad. 258 Linea 544 1 13410519
"""
    pay_text = _accounting_text()

    def _fake_events(pdf_bytes, context):
        return [
            parse_accounting_text(pay_text, context),
            parse_accounting_text(ret_text, context),
        ]

    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run._events_from_asiento_pdf_bytes",
        _fake_events,
    )

    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["summary"]["total_events"] == 2
    assert len(out["items"]) == 2
    assert out["items"][0]["asiento_pdf_path"] == out["items"][1]["asiento_pdf_path"]
    assert out["items"][0]["application_row"] != out["items"][1]["application_row"]
    keys = {it["idempotency_key"] for it in out["items"]}
    assert len(keys) == 2
    # Con recaudo primero (orden de negocio), luego retenciones.
    first_pa = out["items"][0]["payment_application"]
    second_pa = out["items"][1]["payment_application"]
    assert float(first_pa.get("valor_pagado_cliente") or 0) > 0
    assert float(second_pa.get("retenciones") or 0) > 0


def test_dry_run_legacy_manifest_single_asiento_path(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphDryRun(
        _base_files(hist=hist, amort=_amort_table_bytes(fecha), asiento_pdf=_asiento_pdf_placeholder(), ibr=_ibr_bytes())
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(run_amortization_fill_dry_run(g, report_date_iso="2026-05-22"))
    assert len(out["items"]) == 1
    assert out["items"][0]["event_index"] == 1


def test_dry_run_credit_items_resolves_hist_per_individual_credit(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes_multi_credit(
        [
            ("8326b91b", "EQUINORTE", "CREDITO # 258", "TABLAS/amort_258.xlsx", fecha),
            ("8326b91b", "EQUINORTE", "CREDITO # 265", "TABLAS/amort_265.xlsx", fecha),
        ],
        monto_banco=150_000_000.0,
    )

    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "8326b91b",
                "cliente": "EQUINORTE",
                "credito": "258, 265",
                "credit_items": [
                    {
                        "credito": "258",
                        "asiento_pdf_paths": ["clientes/E/asiento_258.pdf"],
                        "extracto_pdf_paths": ["clientes/E/extracto_258.pdf"],
                        "extracto_pdf_path": "clientes/E/extracto_258.pdf",
                    },
                    {
                        "credito": "265",
                        "asiento_pdf_paths": [
                            "clientes/E/asiento_265_1.pdf",
                            "clientes/E/asiento_265_2.pdf",
                        ],
                        "extracto_pdf_paths": ["clientes/E/extracto_265.pdf"],
                        "extracto_pdf_path": "clientes/E/extracto_265.pdf",
                    },
                ],
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_258.xlsx": _amort_table_bytes(fecha),
        "TABLAS/amort_265.xlsx": _amort_table_two_rows_same_date(fecha),
        "clientes/E/asiento_258.pdf": _asiento_pdf_placeholder(),
        "clientes/E/asiento_265_1.pdf": _asiento_pdf_placeholder(),
        "clientes/E/asiento_265_2.pdf": _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["summary"]["total_events"] == 3
    creditos = [it["credito"] for it in out["items"]]
    assert creditos == ["258", "265", "265"]
    assert all(it["application_status"] != "ERROR" or it.get("error_code") != "TABLE_PATH_NOT_FOUND" for it in out["items"])
    assert out["items"][0]["tabla_amortizacion_path"] == "TABLAS/amort_258.xlsx"
    assert out["items"][1]["tabla_amortizacion_path"] == "TABLAS/amort_265.xlsx"
    assert out["items"][0]["due_date_row"] == 2
    assert out["items"][1]["due_date_row"] == 2
    assert out["items"][2]["due_date_row"] == 2
    assert out["items"][0]["application_row"] == 2
    assert out["items"][1]["application_row"] == 2
    assert out["items"][2]["application_row"] == 3
    assert out["items"][0]["application_status"] in (
        "WOULD_APPLY",
        "WOULD_ADOPT_EXISTING",
        "REVISION_MANUAL",
    )
    assert out["items"][1]["application_status"] in (
        "WOULD_APPLY",
        "WOULD_ADOPT_EXISTING",
        "REVISION_MANUAL",
    )


def test_dry_run_two_credits_different_tables_same_target_row(monkeypatch):
    fecha = date(2026, 4, 22)
    id_pago = "8326b91b"
    hist = _hist_bytes_multi_credit(
        [
            (id_pago, "EQUINORTE", "CREDITO # 258", "TABLAS/amort_258.xlsx", fecha),
            (id_pago, "EQUINORTE", "CREDITO # 265", "TABLAS/amort_265.xlsx", fecha),
        ]
    )
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": id_pago,
                "cliente": "EQUINORTE",
                "credito": "258, 265",
                "credit_items": [
                    {
                        "credito": "258",
                        "asiento_pdf_paths": ["clientes/E/asiento_258.pdf"],
                    },
                    {
                        "credito": "265",
                        "asiento_pdf_paths": ["clientes/E/asiento_265.pdf"],
                    },
                ],
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_258.xlsx": _amort_table_date_at_row(fecha, 8),
        "TABLAS/amort_265.xlsx": _amort_table_date_at_row(fecha, 8),
        "clientes/E/asiento_258.pdf": _asiento_pdf_placeholder(),
        "clientes/E/asiento_265.pdf": _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert len(out["items"]) == 2
    assert out["items"][0]["credito"] == "258"
    assert out["items"][1]["credito"] == "265"
    assert out["items"][0]["due_date_row"] == 8
    assert out["items"][1]["due_date_row"] == 8
    assert out["items"][0]["application_row"] == 8
    assert out["items"][1]["application_row"] == 8
    assert out["items"][0]["target_row"] == 8
    assert out["items"][1]["target_row"] == 8
    assert out["items"][0]["error_code"] is None
    assert out["items"][1]["error_code"] is None


def test_dry_run_same_table_second_asiento_excludes_used_row(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 265",
        "TABLAS/amort_265.xlsx",
        fecha,
        monto_banco=100_000_000.0,
    )
    asiento_a = "clientes/E/asiento_265_1.pdf"
    asiento_b = "clientes/E/asiento_265_2.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_paths": [asiento_a, asiento_b],
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_265.xlsx": _amort_table_two_dates_at_rows(fecha, 8, 9),
        asiento_a: _asiento_pdf_placeholder(),
        asiento_b: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["items"][0]["due_date_row"] == 8
    assert out["items"][1]["due_date_row"] == 8
    assert out["items"][0]["application_row"] == 8
    assert out["items"][1]["application_row"] == 9
    assert out["items"][0]["target_row"] == 8
    assert out["items"][1]["target_row"] == 9


def test_dry_run_row_already_assigned_warning_only_same_table(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 265",
        "TABLAS/amort_265.xlsx",
        fecha,
        monto_banco=100_000_000.0,
    )
    asiento_a = "clientes/E/a1.pdf"
    asiento_b = "clientes/E/a2.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_paths": [asiento_a, asiento_b],
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_265.xlsx": _amort_table_date_at_row(fecha, 8),
        asiento_a: _asiento_pdf_placeholder(),
        asiento_b: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["items"][0]["target_row"] == 8
    second = out["items"][1]
    assert second["error_code"] == "REQUIRES_APPLICATION_ROW"
    assert second["due_date_row"] == 8
    assert second["application_row"] is None
    assert any(
        "filas de aplicación ya reservadas" in w.lower() for w in second.get("warnings", [])
    )


def test_dry_run_multi_credit_real_row_8_scenario(monkeypatch):
    fecha = date(2026, 4, 22)
    id_pago = "8326b91b"
    hist = _hist_bytes_multi_credit(
        [
            (id_pago, "EQUINORTE", "CREDITO # 258", "TABLAS/amort_258.xlsx", fecha),
            (id_pago, "EQUINORTE", "CREDITO # 265", "TABLAS/amort_265.xlsx", fecha),
        ],
        monto_banco=150_000_000.0,
    )
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": id_pago,
                "cliente": "EQUINORTE",
                "credito": "258, 265",
                "credit_items": [
                    {
                        "credito": "258",
                        "asiento_pdf_paths": ["clientes/E/asiento_258.pdf"],
                    },
                    {
                        "credito": "265",
                        "asiento_pdf_paths": [
                            "clientes/E/asiento_265_1.pdf",
                            "clientes/E/asiento_265_2.pdf",
                        ],
                    },
                ],
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_258.xlsx": _amort_table_date_at_row(fecha, 8),
        "TABLAS/amort_265.xlsx": _amort_table_two_dates_at_rows(fecha, 8, 9),
        "clientes/E/asiento_258.pdf": _asiento_pdf_placeholder(),
        "clientes/E/asiento_265_1.pdf": _asiento_pdf_placeholder(),
        "clientes/E/asiento_265_2.pdf": _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert len(out["items"]) == 3
    assert out["items"][0]["credito"] == "258"
    assert out["items"][0]["due_date_row"] == 8
    assert out["items"][0]["application_row"] == 8
    assert out["items"][1]["credito"] == "265"
    assert out["items"][1]["due_date_row"] == 8
    assert out["items"][1]["application_row"] == 8
    assert out["items"][2]["credito"] == "265"
    if out["items"][1]["ibr"]["status"] == "WOULD_WRITE_IBR":
        assert out["items"][2]["ibr"]["status"] == "WOULD_SKIP_IBR_ALREADY_PLANNED"
    assert out["items"][2]["due_date_row"] == 8
    assert out["items"][2]["application_row"] == 9
    assert out["items"][2]["error_code"] is None


def _fill_application_cells(ws, row: int, *, amount: float = 1.0) -> None:
    ws.cell(row, 6, amount)
    ws.cell(row, 7, amount)
    ws.cell(row, 9, amount)


def _amort_table_displaced_application(
    fecha_limite: date,
    due_row: int,
    occupied_application_rows: list[int],
    *,
    free_application_row: int | None = None,
) -> bytes:
    """Cronograma en due_row; aplicación ocupada en filas sin fecha de corte."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EQUINORTE"
    ws.append(list(_AMORT_HEADER_ROW))
    max_row_needed = max([due_row, *occupied_application_rows, free_application_row or 0])
    while (ws.max_row or 1) < max_row_needed:
        ws.append([None] * len(_AMORT_HEADER_ROW))
    ws.cell(due_row, 1, fecha_limite.day)
    ws.cell(due_row, 2, fecha_limite.month)
    ws.cell(due_row, 3, fecha_limite.year)
    for r in occupied_application_rows:
        _fill_application_cells(ws, r)
    if free_application_row is not None:
        while (ws.max_row or 1) < free_application_row:
            ws.append([None] * len(_AMORT_HEADER_ROW))
        ws.cell(free_application_row, 5, "")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_dry_run_due_and_application_row_same_when_no_displacement(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphDryRun(
        _base_files(
            hist=hist,
            amort=_amort_table_date_at_row(fecha, 8),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    item = out["items"][0]
    assert item["due_date_row"] == 8
    assert item["ibr_row"] == 8
    assert item["application_row"] == 8
    assert item["target_row"] == 8


def test_dry_run_application_row_below_due_when_block_displaced(monkeypatch):
    fecha = date(2026, 4, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 265", "TABLAS/amort_265.xlsx", fecha)
    amort = _amort_table_displaced_application(
        fecha,
        due_row=8,
        occupied_application_rows=[8, 9],
        free_application_row=10,
    )
    asiento = "clientes/E/asiento_265.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "historico_excel_path": "HIST/cartera.xlsx",
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 265",
                "asiento_pdf_path": asiento,
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort_265.xlsx": amort,
        asiento: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso=fecha.isoformat(),
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    item = out["items"][0]
    assert item["due_date_row"] == 8
    assert item["ibr_row"] == 8
    assert item["application_row"] == 10
    assert item["target_row"] == 10


def test_dry_run_summary_counts_events_not_outputs(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 258",
        "TABLAS/amort.xlsx",
        fecha,
        monto_banco=100_000_000.0,
    )
    asiento_a = "clientes/EQUINORTE/a1.pdf"
    asiento_b = "clientes/EQUINORTE/a2.pdf"
    manifest = {
        "report_date_iso": fecha.isoformat(),
        "outputs": [
            {
                "id_pago": "7785e37e",
                "cliente": "EQUINORTE",
                "credito": "CREDITO # 258",
                "asiento_pdf_paths": [asiento_a, asiento_b],
                "asiento_pdf_path": asiento_a,
            }
        ],
    }
    files = {
        "CTL/dummy.xlsx": b"x",
        f"LOGS/merge_manifest_{fecha.isoformat()}.json": json.dumps(manifest).encode("utf-8"),
        "HIST/cartera.xlsx": hist,
        "TABLAS/amort.xlsx": _amort_table_two_rows_same_date(fecha),
        asiento_a: _asiento_pdf_placeholder(),
        asiento_b: _asiento_pdf_placeholder(),
        "CTL/IBR_DIARIO.xlsx": _ibr_bytes(),
    }
    g = MockGraphDryRun(files)
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["manifest_outputs_count"] == 1
    assert out["summary"]["total"] == 2


def _invalid_amort_table_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Total pagado", "Observación"])
    ws.append([100, "sin encabezados de amortización"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_dry_run_monthly_schedule_finds_target_row(monkeypatch):
    fecha = date(2026, 4, 22)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EQUINORTE"
    ws.append(
        [
            "dia",
            "mes",
            "año",
            "Fecha pago",
            "Valor intereses",
            "Abono a K",
            "Valor pagado cliente",
        ]
    )
    ws.append([22, 12, 2025, None, None, None, None])
    for r in range(3, 8):
        ws.cell(r, 1, f'=IF(A{r}<>" ",1,1)')
        ws.cell(r, 2, f'=IF(A{r}<>" ",1,1)')
        ws.cell(r, 3, f'=IF(A{r}<>" ",1,1)')
    buf = io.BytesIO()
    wb.save(buf)
    amort = buf.getvalue()

    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphDryRun(
        _base_files(
            hist=hist,
            amort=amort,
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-04-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["status"] == "ok"
    item = out["items"][0]
    assert item["application_status"] in ("WOULD_APPLY", "WOULD_ADOPT_EXISTING", "REVISION_MANUAL")
    assert item.get("error_code") is None
    assert item["due_date_row"] == 6
    assert item["ibr_row"] == 6
    assert item["application_row"] == 6
    assert item["target_row"] == 6
    assert item["fecha_limite_pago"] == "2026-04-22"


def test_dry_run_amortization_sheet_not_found_item_error_job_completes(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    g = MockGraphDryRun(
        _base_files(
            hist=hist,
            amort=_invalid_amort_table_bytes(),
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            fecha=fecha,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["status"] == "ok"
    assert out["mode"] == "dry_run"
    assert out["summary"]["errors"] == 1
    item = out["items"][0]
    assert item["application_status"] == "ERROR"
    assert item["error_code"] == "AMORTIZATION_SHEET_NOT_FOUND"
    assert any("encabezados" in w.lower() for w in item["warnings"])


# ---------------------------------------------------------------------------
# §36 / mandato §12: PAGO TOTAL — payoff real vs PAYOFF_NOT_ACHIEVED
# ---------------------------------------------------------------------------

_CANCELACION = "CANCELACIÓN / PAGO TOTAL"
_PAYOFF_CAPITAL = 10_000_000.0
_PAYOFF_TOTAL = 10_000_000.0


def _accounting_text_payoff(*, capital: float, total: float | None = None) -> str:
    paid = total if total is not None else capital
    return f"""
    Comprobante 99 Fecha 22/05/2026
    {ACCOUNT_VALOR_PAGADO_CLIENTE} {paid:,.2f}
    {ACCOUNT_CAPITAL} {capital:,.2f}
    544113430501 0.00
    544141502030 0.00
    """


def _amort_table_with_saldo_before(
    fecha_limite: date,
    *,
    saldo_before: float,
) -> bytes:
    """Tabla secretaria: fila previa con saldo + fila cuota = fecha_limite."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Amort"
    headers = [
        "dia",
        "mes",
        "año",
        "IBR +i",
        "Fecha pago",
        "Valor intereses",
        "Abono a K",
        "intereses mora",
        "Retenciones",
        "Valor pagado cliente",
        "Saldo a capital",
        "Saldos Menores",
    ]
    ws.append(headers)
    # Fila 2: periodo previo (saldo a capital previo al pago bajo prueba).
    prev = date(fecha_limite.year, fecha_limite.month, 1) - timedelta(days=1)
    ws.append(
        [
            prev.day,
            prev.month,
            prev.year,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            saldo_before,
            None,
        ]
    )
    # Fila 3: cuota actual (fecha límite del histórico).
    ws.append(
        [
            fecha_limite.day,
            fecha_limite.month,
            fecha_limite.year,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()



def _payoff_manifest_files(
    *,
    hist: bytes,
    amort: bytes,
    asiento_pdf: bytes,
    ibr: bytes | None = None,
    fecha: date = date(2026, 5, 22),
    tipo_aplicacion: str = _CANCELACION,
) -> dict[str, bytes]:
    files = _base_files(hist=hist, amort=amort, asiento_pdf=asiento_pdf, ibr=ibr, fecha=fecha)
    manifest_key = f"LOGS/merge_manifest_{fecha.isoformat()}.json"
    manifest = json.loads(files[manifest_key].decode("utf-8"))
    for out in manifest.get("outputs") or []:
        out["tipo_aplicacion"] = tipo_aplicacion
        out["tipo_aplicacion_original"] = tipo_aplicacion
    files[manifest_key] = json.dumps(manifest).encode("utf-8")
    return files


def test_dry_run_payoff_real_pass_when_saldo_matches_capital(monkeypatch):
    """§36 A: CANCELACIÓN / PAGO TOTAL con payoff real → dry-run PASS (sin PAYOFF_NOT_ACHIEVED)."""
    from app.application.use_cases.amortization_fill_dry_run import PAYOFF_NOT_ACHIEVED

    fecha = date(2026, 5, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 258",
        "TABLAS/amort.xlsx",
        fecha,
        tipo_aplicacion=_CANCELACION,
        monto_banco=_PAYOFF_TOTAL,
    )
    amort = _amort_table_with_saldo_before(fecha, saldo_before=_PAYOFF_CAPITAL)
    g = MockGraphDryRun(
        _payoff_manifest_files(
            hist=hist,
            amort=amort,
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            tipo_aplicacion=_CANCELACION,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text_payoff(capital=_PAYOFF_CAPITAL, total=_PAYOFF_TOTAL),
    )

    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["status"] == "ok"
    assert out["mode"] == "dry_run"
    assert g.put_calls == []
    item = out["items"][0]
    assert item["error_code"] is None
    assert item.get("error_code") != PAYOFF_NOT_ACHIEVED
    assert item["application_status"] == "WOULD_APPLY"
    assert item["payoff_expected"] is True
    assert out["can_apply"] is True


def test_dry_run_payoff_not_achieved_is_advisory_and_can_apply(monkeypatch):
    """PAGO TOTAL con saldo restante → aviso PAYOFF_NOT_ACHIEVED, Apply permitido."""
    from app.application.use_cases.amortization_fill_apply import _writable_planned_items
    from app.application.use_cases.amortization_fill_dry_run import PAYOFF_NOT_ACHIEVED
    from app.application.ui.amortization_operational_issues import (
        build_operational_issues_from_amortization_result,
    )

    fecha = date(2026, 5, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 258",
        "TABLAS/amort.xlsx",
        fecha,
        tipo_aplicacion=_CANCELACION,
        monto_banco=_PAYOFF_TOTAL,
    )
    # Saldo mucho mayor que el capital del asiento (obligación exacta ≠ payoff).
    amort = _amort_table_with_saldo_before(fecha, saldo_before=100_000_000.0)
    g = MockGraphDryRun(
        _payoff_manifest_files(
            hist=hist,
            amort=amort,
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            tipo_aplicacion=_CANCELACION,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text_payoff(capital=_PAYOFF_CAPITAL, total=_PAYOFF_TOTAL),
    )

    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["status"] == "ok"
    assert out["mode"] == "dry_run"
    assert g.put_calls == []
    item = out["items"][0]
    assert item["application_status"] == "WOULD_APPLY"
    assert item.get("error_code") in (None, "")
    assert item.get("advisory_code") == PAYOFF_NOT_ACHIEVED
    assert item["payoff_expected"] is True
    assert out["can_apply"] is True
    assert _writable_planned_items(out)
    assert out.get("operational_issues")
    assert any(
        i.get("technical_reference") == PAYOFF_NOT_ACHIEVED
        for i in (out.get("operational_issues") or [])
    )

    issues = build_operational_issues_from_amortization_result(out)
    assert any(i.get("technical_reference") == PAYOFF_NOT_ACHIEVED for i in issues)
    msg = " ".join(str(i.get("user_message") or "") for i in issues).lower()
    assert "cancel" in msg or "pago total" in msg


def test_dry_run_parse_failed_masks_payoff_until_asiento_legible(monkeypatch):
    """E15 gap: asiento ilegible -> PDF_TEXT_NOT_EXTRACTABLE (no PAYOFF_NOT_ACHIEVED)."""
    from app.application.services.accounting_pdf_parser import PdfTextNotExtractableError
    from app.application.use_cases.amortization_fill_dry_run import PAYOFF_NOT_ACHIEVED

    fecha = date(2026, 5, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 258",
        "TABLAS/amort.xlsx",
        fecha,
        tipo_aplicacion=_CANCELACION,
        monto_banco=_PAYOFF_TOTAL,
    )
    amort = _amort_table_with_saldo_before(fecha, saldo_before=100_000_000.0)
    g = MockGraphDryRun(
        _payoff_manifest_files(
            hist=hist,
            amort=amort,
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            tipo_aplicacion=_CANCELACION,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: (_ for _ in ()).throw(PdfTextNotExtractableError("scan")),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["can_apply"] is False
    codes = {it.get("error_code") for it in out["items"]}
    assert PAYOFF_NOT_ACHIEVED not in codes
    assert PDF_TEXT_NOT_EXTRACTABLE in codes


def test_dry_run_payoff_evaluated_when_asiento_parseable(monkeypatch):
    """Con asiento legible, dry-run evalua payoff (orden: parse -> payoff)."""
    from app.application.use_cases.amortization_fill_dry_run import PAYOFF_NOT_ACHIEVED

    fecha = date(2026, 5, 22)
    hist = _hist_bytes(
        "7785e37e",
        "CREDITO # 258",
        "TABLAS/amort.xlsx",
        fecha,
        tipo_aplicacion=_CANCELACION,
        monto_banco=_PAYOFF_TOTAL,
    )
    amort = _amort_table_with_saldo_before(fecha, saldo_before=100_000_000.0)
    g = MockGraphDryRun(
        _payoff_manifest_files(
            hist=hist,
            amort=amort,
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
            tipo_aplicacion=_CANCELACION,
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text_payoff(capital=_PAYOFF_CAPITAL, total=_PAYOFF_TOTAL),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["items"][0].get("advisory_code") == PAYOFF_NOT_ACHIEVED
    assert out["can_apply"] is True
    refs = [i.get("technical_reference") for i in (out.get("operational_issues") or [])]
    assert PAYOFF_NOT_ACHIEVED in refs


def test_dry_run_writes_ibr_when_cell_is_placeholder_zero(monkeypatch):
    """Plantilla con IBR+i=0 no debe contar como tasa ya presente (regresión ui-develop)."""
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    amort = _amort_table_bytes(fecha)
    wb = openpyxl.load_workbook(io.BytesIO(amort))
    ws = wb.active
    ws.cell(2, 4, value=0)  # IBR +i placeholder
    buf = io.BytesIO()
    wb.save(buf)
    amort = buf.getvalue()
    g = MockGraphDryRun(
        _base_files(
            hist=hist,
            amort=amort,
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["items"][0]["ibr"]["status"] == "WOULD_WRITE_IBR"


def test_dry_run_skips_ibr_when_real_rate_already_present(monkeypatch):
    fecha = date(2026, 5, 22)
    hist = _hist_bytes("7785e37e", "CREDITO # 258", "TABLAS/amort.xlsx", fecha)
    amort = _amort_table_bytes(fecha)
    wb = openpyxl.load_workbook(io.BytesIO(amort))
    ws = wb.active
    ws.cell(2, 4, value=0.1058)
    buf = io.BytesIO()
    wb.save(buf)
    amort = buf.getvalue()
    g = MockGraphDryRun(
        _base_files(
            hist=hist,
            amort=amort,
            asiento_pdf=_asiento_pdf_placeholder(),
            ibr=_ibr_bytes(),
        )
    )
    monkeypatch.setattr(
        "app.application.use_cases.amortization_fill_dry_run.extract_text_from_pdf",
        lambda _b: _accounting_text(),
    )
    out = asyncio.run(
        run_amortization_fill_dry_run(
            g,
            report_date_iso="2026-05-22",
            historical_file_path="HIST/cartera.xlsx",
        )
    )
    assert out["items"][0]["ibr"]["status"] == "WOULD_SKIP_IBR_ALREADY_PRESENT"
