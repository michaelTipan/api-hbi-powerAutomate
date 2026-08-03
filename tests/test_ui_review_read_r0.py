"""R0: parseo tipado del Excel de revisión (sin Graph)."""
from __future__ import annotations

from openpyxl import Workbook

from app.application.services.review_schema import (
    DistribucionAbonosCols,
    DistribucionCols,
    ErroresCols,
    ReviewSheets,
)
from app.application.ui.review_read import (
    make_row_key,
    parse_review_workbook,
)


def _build_minimal_review_wb() -> Workbook:
    wb = Workbook()
    # Control
    ws_c = wb.active
    ws_c.title = ReviewSheets.CONTROL
    ws_c.append(["Campo", "Valor"])
    ws_c.append(["ReviewSchemaVersion", 2])
    ws_c.append(["Procesar", "NO"])
    ws_c.append(["Estado", "EN_REVISION"])

    # Banner + headers + 1 pago
    ws_p = wb.create_sheet(ReviewSheets.DISTRIBUCION_PAGOS)
    ws_p.append(["banner"])
    ws_p.append(["banner2"])
    headers = [
        DistribucionCols.ID_PAGO,
        DistribucionCols.CLIENTE,
        DistribucionCols.CREDITO,
        DistribucionCols.MONTO_BANCO,
        DistribucionCols.APLICAR_A_EXTRACTO,
        DistribucionCols.MORA_A_APLICAR,
        DistribucionCols.ABONO_A_CAPITAL,
        DistribucionCols.OTROS_VALORES,
        DistribucionCols.TOTAL_APLICADO,
        DistribucionCols.SALDO_POR_ASIGNAR,
        DistribucionCols.ESTADO_PAGO,
        DistribucionCols.VALIDAR_PAGO,
        DistribucionCols.OBSERVACION,
        DistribucionCols.LINK_EXTRACTO,
        DistribucionCols.LINK_TABLA,
        DistribucionCols.LINK_CARPETA_CREDITO,
    ]
    ws_p.append(headers)
    ws_p.append(
        [
            "pago-1",
            "Cliente A",
            "265",
            1000.5,
            800.0,
            0.0,
            200.5,
            0.0,
            1000.5,
            0.0,
            "NORMAL",
            "SI",
            "",
            "Extracto",
            "Tabla",
            "Carpeta",
        ]
    )
    # Hyperlinks
    ws_p.cell(4, 14).hyperlink = "https://example.com/extracto.pdf"
    ws_p.cell(4, 15).hyperlink = "https://example.com/tabla.xlsx"
    ws_p.cell(4, 16).hyperlink = "https://example.com/carpeta"

    ws_a = wb.create_sheet(ReviewSheets.DISTRIBUCION_ABONOS)
    ws_a.append(["b1"])
    ws_a.append(["b2"])
    ws_a.append(
        [
            DistribucionAbonosCols.ID_PAGO,
            DistribucionAbonosCols.CLIENTE,
            DistribucionAbonosCols.CREDITO,
            DistribucionAbonosCols.MONTO_BANCO,
            DistribucionAbonosCols.VALIDAR_ABONO,
            DistribucionAbonosCols.LINK_CARPETA_CREDITO,
        ]
    )
    ws_a.append(["pago-2", "Cliente B", "320", 50.0, "NO", "Carpeta B"])
    ws_a.cell(4, 6).hyperlink = "https://example.com/folder-b"

    ws_e = wb.create_sheet(ReviewSheets.ERRORES)
    ws_e.append(["b"])
    ws_e.append(["b2"])
    ws_e.append(
        [
            ErroresCols.ID_PAGO,
            ErroresCols.CLIENTE,
            ErroresCols.CREDITO,
            ErroresCols.TIPO_CASO,
            ErroresCols.DESCRIPCION,
            ErroresCols.QUE_DEBE_HACER,
            ErroresCols.REQUIERE_SOPORTE,
            ErroresCols.LINK_EXTRACTO,
            ErroresCols.LINK_CARPETA_CREDITO,
            ErroresCols.CODIGO_TECNICO,
        ]
    )
    ws_e.append(
        [
            "pago-err",
            "Cliente E",
            "999",
            "Falta extracto",
            "No hay PDF",
            "Suba el extracto",
            "SI",
            "Extracto",
            "Carpeta",
            "EXTRACTO_MISSING",
        ]
    )
    ws_e.cell(4, 8).hyperlink = "https://example.com/err-extract"
    ws_e.cell(4, 9).hyperlink = "https://example.com/err-folder"
    return wb


def test_make_row_key_stable() -> None:
    assert make_row_key("Distribucion_Pagos", "p1", "265") == "Distribucion_Pagos|p1|265"


def test_parse_review_workbook_pagos_abonos_errors() -> None:
    wb = _build_minimal_review_wb()
    pagos, abonos, errors, schema = parse_review_workbook(wb)
    assert schema == 2
    assert len(pagos) == 1
    p = pagos[0]
    assert p.row_key == "Distribucion_Pagos|pago-1|265"
    assert p.monto_banco == 1000.5
    assert p.validar_pago == "SI"
    rels = {l.rel for l in p.links}
    assert rels == {"extract", "folder", "tabla"}
    assert any(l.web_url and "extracto" in l.web_url for l in p.links)

    assert len(abonos) == 1
    assert abonos[0].row_key == "Distribucion_Abonos|pago-2|320"
    assert abonos[0].validar_abono == "NO"
    assert any(l.rel == "folder" for l in abonos[0].links)

    assert len(errors) == 1
    e = errors[0]
    assert e.requires_regeneration is True
    assert e.credito == "999"
    assert {l.rel for l in e.links} == {"extract", "folder"}
