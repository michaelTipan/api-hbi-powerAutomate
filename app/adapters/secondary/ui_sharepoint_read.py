"""Adaptador SharePoint read-only para UI (solo GET / get_bytes vía UiGraphHttpReadPort)."""
from __future__ import annotations

import io
import json
import logging
import os
from typing import Any

from openpyxl import load_workbook

from app.application.config.payment_validation_settings import (
    DEFAULT_FOLLOWUP_ADELANTADOS,
    resolve_bank_control_file_path,
    resolve_followup_workbook_path,
    validate_bank_code,
)
from app.application.sharepoint_resolution import (
    encode_graph_drive_path,
    resolve_sharepoint_path,
    sharepoint_open_in_browser_url,
)
from app.application.ui.download_limits import (
    UiDownloadTooLargeError,
    assert_download_size_allowed,
    ui_max_download_bytes,
)
from app.application.ui.environment import resolve_active_environment
from app.application.ui.graph_endpoint import assert_relative_graph_endpoint
from app.application.ui.path_guard import (
    UiAllowedRoots,
    UiPathEscapeError,
    assert_path_allowed,
    collect_allowed_roots_from_env,
)
from app.application.ui.manifest_outputs import parse_manifest_output_pdfs
from app.application.ui.ports import (
    UiControlReadResult,
    UiDriveItemMeta,
    UiFileContent,
    UiGraphHttpReadPort,
    UiManifestSummary,
)
from app.application.ui.read_cache import UiReadCache, build_cache_key
from app.application.use_cases.payment_validation_process_control import (
    parse_process_control_row2,
)
from app.application.use_cases.setup_merge_control_workbook import (
    PROCESS_CONTROL_COLUMNS,
    SHEET_NAME,
)
from app.domain.exceptions import GraphConfigError

logger = logging.getLogger(__name__)


def _site_search() -> str:
    return (os.getenv("GRAPH_SHAREPOINT_SITE_SEARCH") or "").strip()


def _drive_name() -> str:
    return (os.getenv("GRAPH_SHAREPOINT_DRIVE_NAME") or "Documentos").strip() or "Documentos"


def _content_endpoint(site_id: str, drive_id: str, relative_path: str) -> str:
    enc = encode_graph_drive_path(relative_path)
    return assert_relative_graph_endpoint(
        f"/sites/{site_id}/drives/{drive_id}/root:/{enc}:/content"
    )


def _item_endpoint(site_id: str, drive_id: str, relative_path: str) -> str:
    enc = encode_graph_drive_path(relative_path)
    return assert_relative_graph_endpoint(
        f"/sites/{site_id}/drives/{drive_id}/root:/{enc}"
    )


def _cell_map_from_control(raw: bytes) -> dict[str, Any]:
    wb = load_workbook(filename=io.BytesIO(raw), data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            return {}
        ws = wb[SHEET_NAME]
        header_map: dict[str, int] = {}
        for c in range(1, (ws.max_column or 0) + 1):
            h = str(ws.cell(row=1, column=c).value or "").strip()
            if h and h not in header_map:
                header_map[h] = c
        out: dict[str, Any] = {}
        for name in PROCESS_CONTROL_COLUMNS:
            col = header_map.get(name)
            if col is None:
                continue
            out[name] = ws.cell(row=2, column=col).value
        return out
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()


def _nz(value: Any) -> str | None:
    text = str(value or "").strip()
    return text or None


class UiSharePointReadAdapter:
    """Implementación Graph del puerto UI. No expone métodos mutantes."""

    def __init__(
        self,
        http: UiGraphHttpReadPort,
        *,
        cache: UiReadCache | None = None,
        allowed_roots: UiAllowedRoots | None = None,
        site_search: str | None = None,
        drive_name: str | None = None,
        max_download_bytes: int | None = None,
    ) -> None:
        self._http = http
        self._cache = cache or UiReadCache(ttl_seconds=30.0)
        self._roots = allowed_roots
        self._site_search = site_search
        self._drive_name = drive_name
        self._max_download_bytes = (
            max_download_bytes if max_download_bytes is not None else ui_max_download_bytes()
        )
        # Guardrail estático: el adaptador no debe declarar mutaciones.
        for banned in (
            "upload",
            "put_bytes",
            "create_folder",
            "delete",
            "move",
            "rename",
            "copy",
            "patch",
            "send_mail",
            "replace_content",
        ):
            if hasattr(self, banned):
                raise RuntimeError(f"UiSharePointReadAdapter no puede exponer {banned}")

    def _environment(self) -> str:
        return resolve_active_environment().environment

    def _roots_bundle(self) -> UiAllowedRoots:
        return self._roots or collect_allowed_roots_from_env()

    def _guard(self, path: str) -> str:
        return assert_path_allowed(path, roots=self._roots_bundle())

    def _prepare_cache(self) -> str:
        env = self._environment()
        self._cache.bind_environment(env)
        return env

    async def _graph_get(self, endpoint: str, params: dict[str, Any] | None = None) -> dict[str, Any]:
        safe = assert_relative_graph_endpoint(endpoint)
        return await self._http.get(safe, params=params)

    async def _graph_get_bytes(self, endpoint: str, params: dict[str, Any] | None = None) -> bytes:
        safe = assert_relative_graph_endpoint(endpoint)
        return await self._http.get_bytes(safe, params=params)

    async def _resolve(self, relative_path: str) -> dict[str, str]:
        safe = self._guard(relative_path)
        return await resolve_sharepoint_path(
            self._http,  # type: ignore[arg-type]
            self._site_search if self._site_search is not None else _site_search(),
            self._drive_name if self._drive_name is not None else _drive_name(),
            safe,
        )

    def _key(
        self,
        *,
        environment: str,
        site_id: str,
        drive_id: str,
        relative_path: str,
        kind: str,
    ) -> str:
        return build_cache_key(
            environment=environment,
            site_id=site_id,
            drive_id=drive_id,
            relative_path=relative_path,
            kind=kind,
        )

    async def get_item_meta(self, relative_path: str) -> UiDriveItemMeta:
        safe = self._guard(relative_path)
        env = self._prepare_cache()
        try:
            info = await self._resolve(safe)
        except Exception as exc:
            logger.info("ui_read: resolve miss path=%s err=%s", safe, type(exc).__name__)
            return UiDriveItemMeta(path=safe, exists=False, name=safe.rsplit("/", 1)[-1])

        site_id = info["site_id"]
        drive_id = info["drive_id"]
        cache_key = self._key(
            environment=env,
            site_id=site_id,
            drive_id=drive_id,
            relative_path=safe,
            kind="meta",
        )
        cached = self._cache.get(cache_key)
        if isinstance(cached, UiDriveItemMeta):
            return cached
        try:
            endpoint = _item_endpoint(site_id, drive_id, safe)
            data = await self._graph_get(
                endpoint,
                params={"$select": "id,name,webUrl,eTag,cTag,size,lastModifiedDateTime"},
            )
        except Exception as exc:
            logger.info("ui_read: meta miss path=%s err=%s", safe, type(exc).__name__)
            return UiDriveItemMeta(path=safe, exists=False, name=safe.rsplit("/", 1)[-1])

        meta = UiDriveItemMeta(
            path=safe,
            name=_nz(data.get("name")),
            web_url=sharepoint_open_in_browser_url(_nz(data.get("webUrl"))) or _nz(data.get("webUrl")),
            etag=_nz(data.get("eTag")),
            ctag=_nz(data.get("cTag")),
            size=int(data["size"]) if isinstance(data.get("size"), int) else None,
            last_modified=_nz(data.get("lastModifiedDateTime")),
            exists=True,
        )
        self._cache.put(cache_key, meta, etag=meta.etag, environment=env)
        return meta

    async def get_web_url(self, relative_path: str) -> str | None:
        meta = await self.get_item_meta(relative_path)
        return meta.web_url if meta.exists else None

    async def download_bytes(self, relative_path: str) -> UiFileContent:
        safe = self._guard(relative_path)
        env = self._prepare_cache()
        meta = await self.get_item_meta(safe)
        assert_download_size_allowed(
            safe,
            meta.size,
            limit=self._max_download_bytes,
        )
        info = await self._resolve(safe)
        site_id = info["site_id"]
        drive_id = info["drive_id"]
        cache_key = self._key(
            environment=env,
            site_id=site_id,
            drive_id=drive_id,
            relative_path=safe,
            kind="bytes",
        )
        cached = self._cache.get(cache_key, etag=meta.etag)
        if isinstance(cached, UiFileContent):
            return cached
        content = await self._graph_get_bytes(
            _content_endpoint(site_id, drive_id, safe)
        )
        assert_download_size_allowed(
            safe,
            len(content),
            limit=self._max_download_bytes,
        )
        result = UiFileContent(
            path=safe,
            content=content,
            etag=meta.etag,
            content_type=None,
        )
        self._cache.put(cache_key, result, etag=meta.etag, environment=env)
        return result

    async def read_process_control(self, bank_code: str) -> UiControlReadResult:
        validate_bank_code(bank_code)
        path = resolve_bank_control_file_path(bank_code)
        safe = self._guard(path)
        file_content = await self.download_bytes(safe)
        snap = parse_process_control_row2(file_content.content, control_file_path=safe)
        cells = _cell_map_from_control(file_content.content)
        meta = await self.get_item_meta(safe)
        return UiControlReadResult(
            snapshot=snap,
            meta=meta,
            generate_job_id=_nz(cells.get("GenerateJobId")),
            finalize_job_id=_nz(cells.get("FinalizeJobId")),
            notify_job_id=_nz(cells.get("NotifyJobId")),
            merge_job_id=_nz(cells.get("MergeJobId")),
            apply_job_id=_nz(cells.get("ApplyJobId")),
            dry_run_job_id=_nz(cells.get("DryRunJobId")),
        )

    async def read_merge_manifest_summary(self, relative_path: str) -> UiManifestSummary:
        safe = self._guard(relative_path)
        env = self._prepare_cache()
        meta = await self.get_item_meta(safe)
        if not meta.exists:
            return UiManifestSummary(path=safe, exists=False, etag=None)
        info = await self._resolve(safe)
        cache_key = self._key(
            environment=env,
            site_id=info["site_id"],
            drive_id=info["drive_id"],
            relative_path=safe,
            kind="manifest",
        )
        cached = self._cache.get(cache_key, etag=meta.etag)
        if isinstance(cached, UiManifestSummary):
            return cached
        raw = await self.download_bytes(safe)
        try:
            data = json.loads(raw.content.decode("utf-8"))
        except Exception:
            summary = UiManifestSummary(path=safe, exists=True, etag=meta.etag, status="UNPARSEABLE")
            self._cache.put(cache_key, summary, etag=meta.etag, environment=env)
            return summary
        if not isinstance(data, dict):
            summary = UiManifestSummary(path=safe, exists=True, etag=meta.etag, status="INVALID")
            self._cache.put(cache_key, summary, etag=meta.etag, environment=env)
            return summary
        incomplete = data.get("incomplete_groups") or data.get("incomplete") or []
        outputs = data.get("outputs") or data.get("complete_groups") or []
        status = _nz(data.get("manifest_status") or data.get("status"))
        eligible = data.get("eligible_for_dry_run")
        output_pdfs = parse_manifest_output_pdfs(outputs)
        primary_output = output_pdfs[0].path if output_pdfs else None
        summary = UiManifestSummary(
            path=safe,
            exists=True,
            etag=meta.etag,
            status=status,
            incomplete_group_count=len(incomplete) if isinstance(incomplete, list) else 0,
            complete_group_count=len(outputs) if isinstance(outputs, list) else 0,
            eligible_for_dry_run=bool(eligible) if isinstance(eligible, bool) else None,
            primary_output_path=primary_output,
            output_pdfs=output_pdfs,
            raw_keys=tuple(sorted(str(k) for k in data.keys())),
        )
        self._cache.put(cache_key, summary, etag=meta.etag, environment=env)
        return summary

    async def adelantados_meta(self) -> UiDriveItemMeta | None:
        try:
            path = resolve_followup_workbook_path(DEFAULT_FOLLOWUP_ADELANTADOS)
        except Exception:
            return None
        try:
            return await self.get_item_meta(path)
        except (UiPathEscapeError, GraphConfigError, UiDownloadTooLargeError):
            return None
