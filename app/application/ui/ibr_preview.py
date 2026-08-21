"""Vista previa de IBR_DIARIO.xlsx para la fase de amortización (solo lectura).

La amortización toma la tasa por fecha de vencimiento (corte) de cada cuota.
Un mismo lote puede tener varios cortes; el modal muestra una tasa por fecha.
La decisión de actualizar IBR usa la misma regla que Apply (resolve_actualiza_ibr).
"""

from __future__ import annotations

import io
import logging
import os
import unicodedata
from datetime import date, datetime
from typing import Any

import openpyxl

from app.application.config.payment_validation_settings import (
    BANK_CODE_BANCOLOMBIA,
    BANK_CODE_BOGOTA,
    resolve_ibr_workbook_path,
)
from app.application.services.colombia_time import format_operator_date, format_operator_datetime
from app.application.services.ibr_workbook import find_ibr_for_date, normalize_ibr_value
from app.application.services.review_schema import (
    resolve_actualiza_ibr,
    resolve_policy_from_tipo_confirmado,
)
from app.application.sharepoint_resolution import (
    require_operations_site_config,
    resolve_sharepoint_path,
)
from app.application.use_cases.setup_ibr_workbook import IBR_COLUMNS, IBR_SHEET_NAME
from app.application.use_cases.setup_merge_control_workbook import (
    process_date_from_process_key,
)
from app.application.use_cases.validate_payment_report import (
    _graph_download_by_path,
    _graph_get_item_metadata_by_path,
)
from app.domain.ports.graph import GraphApiPort

logger = logging.getLogger(__name__)

_AUTOSAVE_HINT = (
    "Si acaba de editar IBR_DIARIO.xlsx en Excel Online, espere unos segundos, "
    "guarde y pulse Actualizar lectura antes de procesar la amortización."
)
_MAX_RANGES = 8


def _accent_fold_upper(text: str) -> str:
    raw = unicodedata.normalize("NFC", str(text or ""))
    nfkd = unicodedata.normalize("NFD", raw)
    return "".join(c for c in nfkd if unicodedata.category(c) != "Mn").strip().upper()


def _parse_excel_date_value(raw: Any) -> date | None:
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def format_operator_rate_pct(rate: float) -> str:
    """Porcentaje con coma decimal, sin ceros de más: «10,58 %»."""
    text = f"{rate * 100:.5f}".rstrip("0").rstrip(".")
    return f"{text.replace('.', ',')} %"


def _list_ibr_ranges(workbook_bytes: bytes, *, limit: int = _MAX_RANGES) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes), data_only=True)
    try:
        if IBR_SHEET_NAME not in wb.sheetnames:
            return []
        ws = wb[IBR_SHEET_NAME]
        col_inicio = col_fin = col_valor = None
        for c in range(1, (ws.max_column or 1) + 1):
            h = _accent_fold_upper(str(ws.cell(1, c).value or ""))
            if h == _accent_fold_upper(IBR_COLUMNS[0]):
                col_inicio = c
            elif h == _accent_fold_upper(IBR_COLUMNS[1]):
                col_fin = c
            elif h == _accent_fold_upper(IBR_COLUMNS[2]):
                col_valor = c
        if not col_inicio or not col_fin or not col_valor:
            return []
        rows: list[dict[str, Any]] = []
        for r in range(2, (ws.max_row or 1) + 1):
            inicio = _parse_excel_date_value(ws.cell(r, col_inicio).value)
            fin = _parse_excel_date_value(ws.cell(r, col_fin).value)
            valor = normalize_ibr_value(ws.cell(r, col_valor).value)
            if inicio is None or fin is None or valor is None:
                continue
            rows.append(
                {
                    "inicio": inicio.isoformat(),
                    "fin": fin.isoformat(),
                    "valor": valor,
                    "valor_pct": round(valor * 100, 5),
                }
            )
        return rows[-limit:]
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()


def resolve_ibr_rates_for_dates(
    workbook_bytes: bytes, dates: list[date]
) -> list[dict[str, Any]]:
    """Una entrada por fecha de corte (todas con updates_ibr=True), orden cronológico."""
    entries = [{"cut": d, "updates_ibr": True, "skip_reason": None} for d in dates]
    return resolve_ibr_rate_rows(workbook_bytes, entries)


def resolve_ibr_rate_rows(
    workbook_bytes: bytes, entries: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """
    Una fila por fecha de corte.
    Si el mismo corte aparece con y sin actualización, gana la que sí actualiza.
    """
    by_cut: dict[date, dict[str, Any]] = {}
    for raw in entries:
        cut = raw.get("cut")
        if not isinstance(cut, date):
            continue
        updates = bool(raw.get("updates_ibr"))
        prev = by_cut.get(cut)
        if prev is None or (updates and not prev.get("updates_ibr")):
            by_cut[cut] = {
                "cut": cut,
                "updates_ibr": updates,
                "skip_reason": None if updates else (raw.get("skip_reason") or "policy"),
            }
    out: list[dict[str, Any]] = []
    for cut in sorted(by_cut.keys()):
        meta = by_cut[cut]
        if not meta["updates_ibr"]:
            reason = str(meta.get("skip_reason") or "policy")
            status = (
                "skipped_adelantado" if reason == "adelantado" else "skipped_policy"
            )
            out.append(
                {
                    "date": cut.isoformat(),
                    "date_label": format_operator_date(cut),
                    "rate": None,
                    "rate_pct": None,
                    "rate_label": None,
                    "status": status,
                    "updates_ibr": False,
                    "skip_reason": reason,
                }
            )
            continue
        rate = find_ibr_for_date(workbook_bytes, cut)
        out.append(
            {
                "date": cut.isoformat(),
                "date_label": format_operator_date(cut),
                "rate": rate,
                "rate_pct": round(rate * 100, 5) if rate is not None else None,
                "rate_label": format_operator_rate_pct(rate) if rate is not None else None,
                "status": "found" if rate is not None else "missing",
                "updates_ibr": True,
                "skip_reason": None,
            }
        )
    return out


def build_ibr_operator_message(rates: list[dict[str, Any]]) -> str:
    """Texto del modal: tasas por corte, sin jerga de ProcessKey ni rangos."""
    if not rates:
        return "No hay fechas de corte para resolver la tasa IBR de este lote."
    needed = [row for row in rates if row.get("updates_ibr", True)]
    skipped = [row for row in rates if not row.get("updates_ibr", True)]
    found = [
        row
        for row in needed
        if row.get("status") == "found" and row.get("rate") is not None
    ]
    missing = [row for row in needed if row.get("status") == "missing"]
    if not needed and skipped:
        return (
            "En este lote no se actualizará IBR (pago adelantado o tipo que no lo requiere)."
        )
    if len(needed) == 1 and not skipped:
        row = needed[0]
        if row.get("status") == "found" and row.get("rate") is not None:
            return (
                f"Tasa IBR para el corte del {row['date_label']}: {row['rate_label']}."
            )
        return (
            f"No hay tasa IBR para el corte del {row['date_label']}. "
            "Actualice IBR_DIARIO.xlsx y vuelva a leer."
        )
    lines: list[str] = []
    for row in rates:
        if row.get("updates_ibr", True):
            lines.append(f"{row['date_label']}: {row['rate_label'] or 'sin tasa'}")
        elif row.get("status") == "skipped_adelantado":
            lines.append(f"{row['date_label']}: no se actualiza (pago antes del corte)")
        else:
            lines.append(f"{row['date_label']}: no se actualiza")
    joined = "; ".join(lines)
    if found and not missing:
        prefix = (
            "Este lote tiene cuotas con distintos cortes. Cada uno usa su tasa IBR: "
            if len(needed) > 1
            else "Tasas IBR del lote: "
        )
        return f"{prefix}{joined}."
    if found or skipped:
        return (
            "Tasas IBR del lote: "
            f"{joined}. Complete las fechas sin tasa en IBR_DIARIO.xlsx antes de amortizar."
            if missing
            else f"Tasas IBR del lote: {joined}."
        )
    return (
        "No hay tasa IBR para los cortes de este lote. "
        f"Fechas: {joined}. Actualice IBR_DIARIO.xlsx y vuelva a leer."
    )


def _bank_code_from_process_key(process_key: str) -> str | None:
    parts = [p.strip() for p in str(process_key or "").split("|") if str(p).strip()]
    for part in parts:
        if part in (BANK_CODE_BOGOTA, BANK_CODE_BANCOLOMBIA):
            return part
    return None


async def _collect_ibr_cut_entries(
    graph: GraphApiPort,
    *,
    process_key: str,
    site_id: str,
    drive_id: str,
    process_date: date | None,
) -> tuple[list[dict[str, Any]], str]:
    """
    Cortes del histórico con la misma decisión IBR que Apply.
    Fallback: fecha del lote (updates_ibr=True) si no hay histórico usable.
    """
    from app.application.services.historical_application_rows import (
        read_validated_application_rows,
    )
    from app.application.use_cases.payment_validation_process_control import (
        read_process_control_snapshot,
    )

    bank = _bank_code_from_process_key(process_key)
    if not bank:
        if process_date:
            return (
                [{"cut": process_date, "updates_ibr": True, "skip_reason": None}],
                "process_key",
            )
        return [], "process_key"
    try:
        snap = await read_process_control_snapshot(
            graph, site_id, drive_id, bank_code=bank
        )
        hist = (snap.historical_file_path or "").strip().strip("/")
        if not hist:
            if process_date:
                return (
                    [{"cut": process_date, "updates_ibr": True, "skip_reason": None}],
                    "process_key",
                )
            return [], "process_key"
        data = await _graph_download_by_path(graph, site_id, drive_id, hist)
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        try:
            rows = read_validated_application_rows(wb)
        finally:
            closer = getattr(wb, "close", None)
            if callable(closer):
                closer()
        entries: list[dict[str, Any]] = []
        for row in rows:
            fecha_lim = row.get("fecha_limite")
            fecha_banco = row.get("fecha_banco")
            cut = fecha_lim if isinstance(fecha_lim, date) else None
            if cut is None and isinstance(fecha_banco, date):
                cut = fecha_banco
            if cut is None:
                continue
            tipo = row.get("tipo_aplicacion_original") or row.get("tipo_aplicacion")
            try:
                policy = resolve_policy_from_tipo_confirmado(tipo)
            except ValueError:
                continue
            pay = fecha_banco if isinstance(fecha_banco, date) else None
            lim = fecha_lim if isinstance(fecha_lim, date) else None
            updates = resolve_actualiza_ibr(
                policy, payment_date=pay, fecha_limite=lim
            )
            skip_reason = None
            if not updates:
                if (
                    pay is not None
                    and lim is not None
                    and pay < lim
                ):
                    skip_reason = "adelantado"
                else:
                    skip_reason = "policy"
            entries.append(
                {
                    "cut": cut,
                    "updates_ibr": updates,
                    "skip_reason": skip_reason,
                }
            )
        if entries:
            return entries, "historical"
    except Exception:
        logger.info(
            "ibr_preview: no se pudieron leer cortes del histórico process_key=%s",
            process_key,
            exc_info=True,
        )
    if process_date:
        return (
            [{"cut": process_date, "updates_ibr": True, "skip_reason": None}],
            "process_key",
        )
    return [], "process_key"


async def load_ibr_preview(graph: GraphApiPort, *, process_key: str) -> dict[str, Any]:
    """Lee IBR_DIARIO y resuelve la tasa por cada fecha de corte del lote."""
    require_operations_site_config()
    key = (process_key or "").strip()
    process_date = process_date_from_process_key(key)
    rel = resolve_ibr_workbook_path().strip().strip("/")
    site_search = (os.getenv("GRAPH_SHAREPOINT_SITE_SEARCH") or "").strip()
    drive_name = (os.getenv("GRAPH_SHAREPOINT_DRIVE_NAME") or "").strip()
    info = await resolve_sharepoint_path(graph, site_search, drive_name, rel)
    site_id = str(info["site_id"])
    drive_id = str(info["drive_id"])
    meta = await _graph_get_item_metadata_by_path(graph, site_id, drive_id, rel)
    last_modified = str(meta.get("lastModifiedDateTime") or "").strip() or None
    data = await _graph_download_by_path(graph, site_id, drive_id, rel)
    ranges = _list_ibr_ranges(data)
    cut_entries, dates_source = await _collect_ibr_cut_entries(
        graph,
        process_key=key,
        site_id=site_id,
        drive_id=drive_id,
        process_date=process_date,
    )
    rates = resolve_ibr_rate_rows(data, cut_entries)
    needed = [row for row in rates if row.get("updates_ibr", True)]
    found = [row for row in needed if row.get("status") == "found"]
    missing = [row for row in needed if row.get("status") == "missing"]
    if not rates:
        rate_status = "no_process_date"
    elif not needed:
        rate_status = "not_required"
    elif found and not missing:
        rate_status = "found"
    elif found:
        rate_status = "partial"
    else:
        rate_status = "missing_for_date"
    primary = next(
        (
            row
            for row in needed
            if row.get("date")
            == (process_date.isoformat() if process_date else None)
        ),
        needed[0] if needed else (rates[0] if rates else None),
    )
    rate = primary.get("rate") if primary else None
    warnings: list[str] = [_AUTOSAVE_HINT]
    if rate_status == "missing_for_date":
        warnings.append(
            "No hay tasa IBR para las fechas de corte de este lote. "
            "Actualice IBR_DIARIO.xlsx antes de amortizar."
        )
    elif rate_status == "partial":
        warnings.append(
            "Falta la tasa IBR en al menos un corte de este lote. "
            "Complete IBR_DIARIO.xlsx antes de amortizar."
        )
    if dates_source == "historical" and len(needed) > 1:
        warnings.append(
            "Las tasas corresponden a la fecha de vencimiento de cada cuota, "
            "no a la fecha del movimiento en el banco."
        )
    user_message = build_ibr_operator_message(rates)
    return {
        "ok": True,
        "source_path": rel,
        "process_key": key,
        "process_date": process_date.isoformat() if process_date else None,
        "rate": rate,
        "rate_pct": round(rate * 100, 5) if rate is not None else None,
        "rate_status": rate_status,
        "dates_source": dates_source,
        "rates": rates,
        "ranges": ranges,
        "file_last_modified": format_operator_datetime(last_modified),
        "warnings": warnings,
        "user_message": user_message,
    }
