"""Lectura tipada del Excel de revisión para la UI (R0 solo lectura).

No escribe SharePoint. Identidad de fila: sheet + id_pago + credito.
"""
from __future__ import annotations

import io
import logging
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from app.application.services.review_schema import (
    DistribucionAbonosCols,
    DistribucionCols,
    ReviewSheets,
    apply_legacy_estado_migration,
    detect_distrib_schema_version_from_headers,
    find_distribucion_pagos_sheet,
    normalize_distrib_row_keys,
    read_control_review_schema_version,
)
from app.application.ui.ports import UiSharePointReadPort
from app.application.ui.review_errores_read import (
    parse_review_errores_workbook,
)
from app.application.ui.schemas import (
    UiLink,
    UiReviewAbonoRow,
    UiReviewErrorItem,
    UiReviewPagoRow,
    UiReviewResponse,
)
from app.application.use_cases.payment_validation_finalize import (
    _find_table_header_row,
)

logger = logging.getLogger(__name__)

# Headers Excel (referencia); el contrato API usa los nombres snake_case abajo.
EDITABLE_PAGO_FIELDS = [
    DistribucionCols.APLICAR_A_EXTRACTO,
    DistribucionCols.MORA_A_APLICAR,
    DistribucionCols.ABONO_A_CAPITAL,
    DistribucionCols.OTROS_VALORES,
    DistribucionCols.ESTADO_PAGO,
    DistribucionCols.VALIDAR_PAGO,
    DistribucionCols.OBSERVACION,
]

EDITABLE_ABONO_FIELDS = [DistribucionAbonosCols.VALIDAR_ABONO]

# Nombres de campo del contrato PATCH / GET editable_fields
EDITABLE_PAGO_API_FIELDS = [
    "aplicar_a_extracto",
    "mora_a_aplicar",
    "abono_a_capital",
    "otros_valores",
    "estado_pago",
    "validar_pago",
    "observacion",
]

EDITABLE_ABONO_API_FIELDS = ["validar_abono"]


class ReviewFileMissingError(LookupError):
    """No hay ValidationFilePath o el archivo no existe en SharePoint."""


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _hyperlink_target(cell: Any) -> str | None:
    try:
        hl = getattr(cell, "hyperlink", None)
        if hl is None:
            return None
        target = getattr(hl, "target", None) or getattr(hl, "ref", None)
        text = str(target or "").strip()
        return text or None
    except Exception:
        return None


def _to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, Decimal):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        return None


def _to_int(value: Any) -> int | None:
    n = _to_float(value)
    if n is None:
        return None
    try:
        return int(n)
    except (TypeError, ValueError):
        return None


def make_row_key(sheet: str, id_pago: str, credito: str, *, excel_row: int | None = None) -> str:
    base = f"{sheet}|{(id_pago or '').strip()}|{(credito or '').strip()}"
    if not (id_pago or "").strip() and excel_row is not None:
        return f"{base}|r{excel_row}"
    return base


def _link(rel: str, label: str, web_url: str | None) -> UiLink | None:
    url = (web_url or "").strip()
    lab = (label or "").strip() or rel
    if not url and not lab:
        return None
    return UiLink(
        rel=rel,
        label=lab if lab != rel else (
            "Abrir extracto" if rel == "extract"
            else "Abrir carpeta" if rel == "folder"
            else "Abrir tabla" if rel == "tabla"
            else lab
        ),
        path=None,
        web_url=url or None,
        open_mode="sharepoint",
    )


def _row_links_from_cells(
    by_header: dict[str, Any],
    *,
    extract_header: str | None,
    folder_header: str | None,
    tabla_header: str | None,
) -> list[UiLink]:
    links: list[UiLink] = []
    specs = (
        ("extract", extract_header),
        ("folder", folder_header),
        ("tabla", tabla_header),
    )
    for rel, header in specs:
        if not header:
            continue
        cell = by_header.get(header)
        if cell is None:
            continue
        label = _cell_text(getattr(cell, "value", None))
        url = _hyperlink_target(cell)
        item = _link(rel, label, url)
        if item and (item.web_url or item.label):
            # Sin URL: aún mostrar label si hay texto (operador ve que existe)
            if item.web_url:
                links.append(item)
            elif label and label.upper() not in {"", "N/A", "NO APLICA", "-"}:
                links.append(item)
    return links


def parse_distribucion_pagos(workbook: Any) -> tuple[list[UiReviewPagoRow], int | None]:
    try:
        ws = find_distribucion_pagos_sheet(workbook)
    except ValueError:
        return [], None
    try:
        header_row = _find_table_header_row(ws, DistribucionCols.ID_PAGO)
    except ValueError:
        return [], None

    headers: list[str] = []
    schema_version: int | None = None
    out: list[UiReviewPagoRow] = []
    for r_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
        if r_idx < header_row:
            continue
        if r_idx == header_row:
            headers = [_cell_text(c.value) for c in row]
            schema_version = detect_distrib_schema_version_from_headers(headers)
            continue
        values = [c.value for c in row]
        if not any(v is not None and str(v).strip() for v in values):
            continue
        raw = dict(zip(headers, values))
        rd = normalize_distrib_row_keys(raw, schema_version=schema_version)
        apply_legacy_estado_migration(rd)
        id_pago = _cell_text(rd.get(DistribucionCols.ID_PAGO))
        credito = _cell_text(rd.get(DistribucionCols.CREDITO))
        if not id_pago and not credito:
            continue
        by_header = {h: c for h, c in zip(headers, row)}
        links = _row_links_from_cells(
            by_header,
            extract_header=DistribucionCols.LINK_EXTRACTO,
            folder_header=DistribucionCols.LINK_CARPETA_CREDITO,
            tabla_header=DistribucionCols.LINK_TABLA,
        )
        out.append(
            UiReviewPagoRow(
                row_key=make_row_key(
                    ReviewSheets.DISTRIBUCION_PAGOS, id_pago, credito, excel_row=r_idx
                ),
                excel_row=r_idx,
                id_pago=id_pago,
                cliente=_cell_text(rd.get(DistribucionCols.CLIENTE)),
                credito=credito,
                monto_banco=_to_float(rd.get(DistribucionCols.MONTO_BANCO)),
                fecha_banco=_cell_text(rd.get(DistribucionCols.FECHA_BANCO)) or None,
                fecha_limite=_cell_text(rd.get(DistribucionCols.FECHA_LIMITE)) or None,
                dias_mora=_to_int(rd.get(DistribucionCols.DIAS_MORA)),
                valor_extracto=_to_float(rd.get(DistribucionCols.VALOR_EXTRACTO)),
                aplicar_a_extracto=_to_float(rd.get(DistribucionCols.APLICAR_A_EXTRACTO)),
                mora_a_aplicar=_to_float(rd.get(DistribucionCols.MORA_A_APLICAR)),
                abono_a_capital=_to_float(rd.get(DistribucionCols.ABONO_A_CAPITAL)),
                otros_valores=_to_float(rd.get(DistribucionCols.OTROS_VALORES)),
                total_aplicado=_to_float(rd.get(DistribucionCols.TOTAL_APLICADO)),
                saldo_por_asignar=_to_float(rd.get(DistribucionCols.SALDO_POR_ASIGNAR)),
                estado_pago=_cell_text(rd.get(DistribucionCols.ESTADO_PAGO)) or None,
                validar_pago=_cell_text(rd.get(DistribucionCols.VALIDAR_PAGO)) or None,
                observacion=_cell_text(rd.get(DistribucionCols.OBSERVACION)) or None,
                tipo_aplicacion_original=_cell_text(
                    rd.get(DistribucionCols.TIPO_APLICACION_ORIGINAL)
                )
                or None,
                editable_fields=list(EDITABLE_PAGO_API_FIELDS),
                links=links,
            )
        )
    return out, schema_version


def parse_distribucion_abonos(workbook: Any) -> list[UiReviewAbonoRow]:
    if ReviewSheets.DISTRIBUCION_ABONOS not in getattr(workbook, "sheetnames", []):
        return []
    ws = workbook[ReviewSheets.DISTRIBUCION_ABONOS]
    try:
        header_row = _find_table_header_row(ws, DistribucionAbonosCols.ID_PAGO)
    except ValueError:
        return []

    headers: list[str] = []
    out: list[UiReviewAbonoRow] = []
    for r_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
        if r_idx < header_row:
            continue
        if r_idx == header_row:
            headers = [_cell_text(c.value) for c in row]
            continue
        values = [c.value for c in row]
        if not any(v is not None and str(v).strip() for v in values):
            continue
        rd = dict(zip(headers, values))
        id_pago = _cell_text(rd.get(DistribucionAbonosCols.ID_PAGO))
        credito = _cell_text(rd.get(DistribucionAbonosCols.CREDITO))
        if not id_pago and not credito:
            continue
        by_header = {h: c for h, c in zip(headers, row)}
        links = _row_links_from_cells(
            by_header,
            extract_header=DistribucionAbonosCols.LINK_EXTRACTO,
            folder_header=DistribucionAbonosCols.LINK_CARPETA_CREDITO,
            tabla_header=DistribucionAbonosCols.LINK_TABLA,
        )
        out.append(
            UiReviewAbonoRow(
                row_key=make_row_key(
                    ReviewSheets.DISTRIBUCION_ABONOS, id_pago, credito, excel_row=r_idx
                ),
                excel_row=r_idx,
                id_pago=id_pago,
                cliente=_cell_text(rd.get(DistribucionAbonosCols.CLIENTE)),
                credito=credito,
                monto_banco=_to_float(rd.get(DistribucionAbonosCols.MONTO_BANCO)),
                fecha_banco=_cell_text(rd.get(DistribucionAbonosCols.FECHA_BANCO)) or None,
                validar_abono=_cell_text(rd.get(DistribucionAbonosCols.VALIDAR_ABONO))
                or None,
                observacion=_cell_text(rd.get(DistribucionAbonosCols.OBSERVACION)) or None,
                origen_credito=_cell_text(rd.get(DistribucionAbonosCols.ORIGEN_CREDITO))
                or None,
                tipo_aplicacion_original=_cell_text(
                    rd.get(DistribucionAbonosCols.TIPO_APLICACION_ORIGINAL)
                )
                or None,
                editable_fields=list(EDITABLE_ABONO_API_FIELDS),
                links=links,
            )
        )
    return out


def parse_review_errors_as_ui(workbook: Any) -> list[UiReviewErrorItem]:
    items: list[UiReviewErrorItem] = []
    for err in parse_review_errores_workbook(workbook):
        links: list[UiLink] = []
        if err.extract_url:
            links.append(
                UiLink(
                    rel="extract",
                    label=err.extract_label or "Abrir extracto",
                    web_url=err.extract_url,
                )
            )
        if err.folder_url:
            links.append(
                UiLink(
                    rel="folder",
                    label=err.folder_label or "Abrir carpeta",
                    web_url=err.folder_url,
                )
            )
        items.append(
            UiReviewErrorItem(
                row_key=make_row_key(
                    ReviewSheets.ERRORES,
                    err.id_pago,
                    err.credito,
                    excel_row=err.row_number,
                ),
                excel_row=err.row_number,
                id_pago=err.id_pago,
                cliente=err.cliente,
                credito=err.credito,
                tipo_caso=err.tipo_caso,
                descripcion=err.descripcion,
                que_debe_hacer=err.que_debe_hacer,
                codigo_tecnico=err.codigo_tecnico or None,
                requires_regeneration=True,
                links=links,
            )
        )
    return items


def parse_review_workbook(workbook: Any) -> tuple[
    list[UiReviewPagoRow],
    list[UiReviewAbonoRow],
    list[UiReviewErrorItem],
    int | None,
]:
    schema_from_control: int | None = None
    if ReviewSheets.CONTROL in getattr(workbook, "sheetnames", []):
        schema_from_control = read_control_review_schema_version(
            workbook[ReviewSheets.CONTROL]
        )
    pagos, schema_from_headers = parse_distribucion_pagos(workbook)
    abonos = parse_distribucion_abonos(workbook)
    errors = parse_review_errors_as_ui(workbook)
    schema = schema_from_control if schema_from_control is not None else schema_from_headers
    return pagos, abonos, errors, schema


async def load_ui_review_for_process(
    reader: UiSharePointReadPort,
    *,
    process_key: str,
    banks: tuple[str, ...],
) -> UiReviewResponse:
    """Resuelve Control por process_key y parsea el Excel de revisión."""
    key = (process_key or "").strip()
    matched_bank: str | None = None
    snap = None
    for bank in banks:
        try:
            control = await reader.read_process_control(bank)
        except Exception:
            logger.info("review_read: control falló bank=%s", bank, exc_info=True)
            continue
        candidate = control.snapshot
        if (candidate.process_key or "").strip() == key:
            matched_bank = bank
            snap = candidate
            break
    if snap is None or matched_bank is None:
        raise KeyError(key)

    path = (snap.validation_file_path or "").strip()
    if not path:
        raise ReviewFileMissingError("validation_file_path_empty")

    meta = await reader.get_item_meta(path)
    if meta is not None and meta.exists is False:
        raise ReviewFileMissingError("validation_file_missing")

    file_content = await reader.download_bytes(path)
    etag = (
        (file_content.etag if file_content else None)
        or (meta.etag if meta else None)
    )
    review_web = meta.web_url if meta else None
    if not review_web:
        try:
            review_web = await reader.get_web_url(path)
        except Exception:
            review_web = None

    try:
        wb = load_workbook(io.BytesIO(file_content.content), data_only=False)
    except Exception as exc:
        logger.info("review_read: workbook ilegible path=%s", path, exc_info=True)
        raise ValueError("review_workbook_unreadable") from exc

    pagos, abonos, errors, schema = parse_review_workbook(wb)
    requires_regen = len(errors) > 0
    read_only = True
    try:
        from app.application.ui.feature_flags import get_ui_feature_flags

        read_only = not get_ui_feature_flags().review_edit_allowed
    except Exception:
        read_only = True
    return UiReviewResponse(
        process_key=key,
        bank_code=matched_bank or (snap.bank_code or ""),
        validation_file_path=path,
        review_excel=UiLink(
            rel="review_excel",
            label="Abrir archivo de revisión",
            path=path,
            web_url=review_web,
        )
        if path
        else None,
        etag=etag,
        schema_version=schema,
        pagos=pagos,
        abonos=abonos,
        errors=errors,
        requires_regeneration=requires_regen,
        read_only=read_only,
        summary={
            "pagos": len(pagos),
            "abonos": len(abonos),
            "errors": len(errors),
        },
    )
