"""Preflight dry-run de reglas Finalize sobre el Excel de revisión (R1).

No modifica archivos SharePoint ni Control. Reutiliza colectores de Finalize.
"""
from __future__ import annotations

import io
import logging
from typing import Any

from openpyxl import load_workbook

from app.application.job_status_enrichment import finalize_message_for_code
from app.application.services.review_schema import (
    DistribucionCols,
    ReviewSheets,
    REVIEW_SCHEMA_VERSION,
    apply_legacy_estado_migration,
    detect_distrib_schema_version_from_headers,
    find_distribucion_pagos_sheet,
    normalize_distrib_row_keys,
    read_control_review_schema_version,
)
from app.application.ui.ports import UiSharePointReadPort
from app.application.ui.review_read import ReviewFileMissingError
from app.application.ui.schemas import (
    UiReviewPreflightIssue,
    UiReviewPreflightResponse,
)
from app.application.use_cases.payment_validation_finalize import (
    _collect_abono_issues,
    _collect_distribucion_pago_issues,
    _count_open_errores_rows,
    _find_table_header_row,
    _read_abono_distributions,
    _resolve_distrib_policy,
    _validate_no_duplicate_distrib_monto_banco,
)

logger = logging.getLogger(__name__)


def _issue_from_raw(raw: dict[str, Any]) -> UiReviewPreflightIssue:
    code = str(raw.get("error_code") or "unknown").strip() or "unknown"
    user_msg, _next = finalize_message_for_code(code)
    excel_row = raw.get("excel_row")
    row_int: int | None
    try:
        row_int = int(excel_row) if excel_row is not None else None
    except (TypeError, ValueError):
        row_int = None
    value = raw.get("value_found")
    return UiReviewPreflightIssue(
        error_code=code,
        sheet=str(raw.get("sheet") or "").strip() or None,
        excel_row=row_int,
        id_pago=str(raw.get("id_pago") or "").strip() or None,
        credito=str(raw.get("credito") or "").strip() or None,
        field=str(raw.get("field") or "").strip() or None,
        value_found=None if value is None else str(value),
        user_message=user_msg,
    )


def collect_preflight_issues_from_workbook(workbook: Any) -> list[dict[str, Any]]:
    """Misma lógica de prevalidación que Finalize, sin mutar ni exigir Procesar=SI."""
    issues: list[dict[str, Any]] = []

    open_errores = _count_open_errores_rows(workbook)
    if open_errores > 0:
        issues.append(
            {
                "error_code": "review_has_open_errors",
                "sheet": ReviewSheets.ERRORES,
                "value_found": open_errores,
            }
        )

    try:
        ws_dist = find_distribucion_pagos_sheet(workbook)
    except ValueError:
        issues.append(
            {
                "error_code": "review_unreadable",
                "sheet": ReviewSheets.DISTRIBUCION_PAGOS,
            }
        )
        return issues

    ws_ctrl = (
        workbook[ReviewSheets.CONTROL]
        if ReviewSheets.CONTROL in workbook.sheetnames
        else None
    )

    distributions: list[dict[str, Any]] = []
    try:
        dist_header_row = _find_table_header_row(ws_dist, DistribucionCols.ID_PAGO)
    except ValueError:
        issues.append(
            {
                "error_code": "review_unreadable",
                "sheet": ReviewSheets.DISTRIBUCION_PAGOS,
            }
        )
        return issues

    headers: list[str] = []
    for r_idx, row in enumerate(ws_dist.iter_rows(values_only=True), start=1):
        if r_idx < dist_header_row:
            continue
        if r_idx == dist_header_row:
            headers = [str(v).strip() if v else "" for v in row]
            distrib_schema_version = detect_distrib_schema_version_from_headers(headers)
            ctrl_schema_version = (
                read_control_review_schema_version(ws_ctrl) if ws_ctrl is not None else None
            )
            if distrib_schema_version < REVIEW_SCHEMA_VERSION or (
                ctrl_schema_version is not None
                and ctrl_schema_version < REVIEW_SCHEMA_VERSION
            ):
                issues.append(
                    {
                        "error_code": "review_schema_version_1_requires_regenerate",
                        "sheet": ReviewSheets.DISTRIBUCION_PAGOS,
                        "value_found": distrib_schema_version,
                    }
                )
            continue
        if not any(row):
            continue
        row_dict = normalize_distrib_row_keys(
            dict(zip(headers, row)),
            schema_version=detect_distrib_schema_version_from_headers(headers),
        )
        apply_legacy_estado_migration(row_dict)
        row_dict["_excel_row"] = r_idx
        row_dict["_policy"] = _resolve_distrib_policy(row_dict)
        distributions.append(row_dict)

    try:
        _validate_no_duplicate_distrib_monto_banco(distributions)
    except ValueError as exc:
        issues.append(
            {
                "error_code": str(exc).split("|", 1)[0] or "duplicate_bank_amount_in_payment_group",
                "sheet": ReviewSheets.DISTRIBUCION_PAGOS,
            }
        )

    issues.extend(_collect_distribucion_pago_issues(distributions))

    abono_rows_all: list[dict[str, Any]] = []
    if ReviewSheets.DISTRIBUCION_ABONOS in workbook.sheetnames:
        try:
            _header, abono_rows_all = _read_abono_distributions(
                workbook[ReviewSheets.DISTRIBUCION_ABONOS]
            )
        except ValueError as exc:
            code = str(exc).split("|", 1)[0] or "invalid_validar_abono"
            issues.append(
                {
                    "error_code": code,
                    "sheet": ReviewSheets.DISTRIBUCION_ABONOS,
                }
            )
            abono_rows_all = []

    issues.extend(_collect_abono_issues(abono_rows_all))
    return issues


async def run_ui_review_preflight(
    reader: UiSharePointReadPort,
    *,
    process_key: str,
    banks: tuple[str, ...],
) -> UiReviewPreflightResponse:
    key = (process_key or "").strip()
    matched_bank: str | None = None
    snap = None
    for bank in banks:
        try:
            control = await reader.read_process_control(bank)
        except Exception:
            continue
        if (control.snapshot.process_key or "").strip() == key:
            matched_bank = bank
            snap = control.snapshot
            break
    if snap is None or matched_bank is None:
        raise KeyError(key)

    path = (snap.validation_file_path or "").strip()
    if not path:
        raise ReviewFileMissingError("validation_file_path_empty")

    meta = await reader.get_item_meta(path)
    if meta is not None and meta.exists is False:
        raise ReviewFileMissingError("validation_file_missing")

    file_content = await reader.download_bytes(path)
    etag = (file_content.etag if file_content else None) or (
        meta.etag if meta else None
    )

    try:
        wb = load_workbook(io.BytesIO(file_content.content), data_only=False)
    except Exception as exc:
        raise ValueError("review_workbook_unreadable") from exc

    raw_issues = collect_preflight_issues_from_workbook(wb)
    mapped = [_issue_from_raw(i) for i in raw_issues]
    requires_regen = any(
        i.error_code
        in {
            "review_has_open_errors",
            "review_schema_version_1_requires_regenerate",
        }
        for i in mapped
    )
    return UiReviewPreflightResponse(
        process_key=key,
        etag=etag,
        ok=len(mapped) == 0,
        issue_count=len(mapped),
        issues=mapped,
        requires_regeneration=requires_regen,
    )
