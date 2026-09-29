"""Readiness de Merge (solo lectura): histórico + listado de ASIENTOS + lectura de asientos.

No adquiere mutex, no escribe control, no une PDFs ni sube manifiestos. Además de
comprobar que cada crédito tenga su PDF, lee los asientos (mismo parser de fase 4)
y arma la verificación por crédito (``merge_verification``) respetando las decisiones
del registro del proceso (exclusiones, asignaciones, tipo corregido, no usados).
"""
from __future__ import annotations

import asyncio
import hashlib
import logging
import os
import time
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from typing import Any, Protocol

from openpyxl import load_workbook

from app.application.job_manager import get_job_manager
from app.application.services.accounting_pdf_parser import parse_accounting_pdf_events
from app.application.services.amortization_decisions import (
    DecisionsRegistryUnavailableError,
    load_registry_for_process,
    unused_asiento_paths,
)
from app.application.services.asiento_format_gate import (
    FORMAT_GATE_CODES,
    classify_asiento_pdf_bytes,
)
from app.application.services.colombia_time import now_colombia_iso
from app.application.services.lot_conciliation import cents
from app.application.services.merge_verification import (
    AsientoRead,
    apply_tipo_corrections,
    evaluate_verification,
    event_facts,
    excluded_id_pagos,
    table_fact_requests,
)
from app.application.services.historical_application_rows import (
    group_rows_by_id_pago,
    read_validated_abono_rows,
    read_validated_payment_rows,
)
from app.application.services.merge_group_validation import (
    credit_hint_from_pdf_filename,
    validate_merge_group_completeness,
)
from app.application.services.review_schema import ExtractRole
from app.application.ui.merge_capabilities import control_indicates_already_merged
from app.application.ui.schemas import UiMergeReadiness
from app.application.use_cases.merge_composite_validado_pdfs import (
    _classify_asiento_pdf_names,
    _pdf_names_in_children,
    _ruta_asientos_from_cell,
    normalize_sharepoint_path,
)
from app.application.use_cases.payment_validation_process_control import (
    ProcessControlSnapshot,
)
from app.application.use_cases.send_validar_extractos_notification import (
    _collect_pdf_paths_from_ruta_cell,
    _list_drive_folder_children,
)
from app.application.use_cases.validate_payment_report import _graph_download_by_path
from app.application.sharepoint_resolution import resolve_sharepoint_from_env

logger = logging.getLogger(__name__)

_MSG_ALREADY = "Los asientos contables de este proceso ya fueron consolidados."
_MSG_INCOMPLETE = "Aún faltan documentos contables para uno o más créditos."
_MSG_INCOMPLETE_MISMATCH = (
    "Hay PDF en ASIENTOS cuyo nombre no coincide con el crédito."
)
_MSG_INCOMPLETE_FORMAT = (
    "Hay PDF de asiento ilegibles o que no tienen el formato del ERP."
)
_MSG_INCOMPLETE_MIXED = (
    "Faltan asientos, hay nombres que no coinciden o el formato del PDF no es el esperado."
)
_MSG_UNKNOWN = (
    "No se pudo verificar si los asientos contables están completos. "
    "Actualice e intente nuevamente."
)
_MSG_READY = "Los asientos contables están listos para consolidar."
_NEXT_LOAD = (
    "Cargue o corrija los archivos pendientes y verifique los asientos contables."
)
_NEXT_RETRY = "Actualice el detalle del proceso e intente nuevamente."
_NEXT_MERGE = "Puede generar el PDF consolidado desde la UI."
_NEXT_DONE = "Consulte los PDFs consolidados y continúe con amortización cuando corresponda."
_MSG_REGISTRY_UNAVAILABLE = (
    "No se pudo leer el registro de decisiones del proceso. "
    "Actualice e intente nuevamente."
)
_MSG_ALL_EXCLUDED = (
    "Todos los créditos están excluidos. Deshaga una exclusión para consolidar "
    "o cancele el proceso."
)
_NEXT_REVIEW = "Revise los créditos que requieren atención."

_BYTES_CACHE_MAX = 256
_EVENTS_CACHE_MAX = 512
_FACTS_TTL_SECONDS = 600.0
_bytes_cache: OrderedDict[tuple[str, str], bytes] = OrderedDict()
_events_cache: OrderedDict[str, list[dict[str, Any]] | None] = OrderedDict()
_facts_cache: dict[tuple[str, str], tuple[float, dict[str, Any] | None]] = {}


def clear_verification_caches() -> None:
    _bytes_cache.clear()
    _events_cache.clear()
    _facts_cache.clear()


def _lru_put(cache: OrderedDict, key: Any, value: Any, limit: int) -> None:
    cache[key] = value
    cache.move_to_end(key)
    while len(cache) > limit:
        cache.popitem(last=False)


def _incomplete_user_message(missing_items: list[dict[str, Any]]) -> str:
    """Mensaje corto según el tipo de faltante (detalle va en missing_items)."""
    codes = {str(item.get("error_code") or "").strip() for item in missing_items}
    has_mismatch = "asiento_contable_credit_mismatch" in codes
    has_format = bool(codes & FORMAT_GATE_CODES)
    has_absence = bool(
        codes
        & {
            "asiento_contable_not_found",
            "missing_ruta_asientos_contables",
            "document_missing",
            "extract_routes_missing",
            "asiento_download_failed",
            "ASIENTO_ASSIGNMENT_AMBIGUOUS",
            "ASIENTO_ASSIGNMENT_NO_MATCH",
            "ASIENTO_ASSIGNMENT_PARSE_FAILED",
            "ASIENTO_ASSIGNMENT_COMPLEXITY_LIMIT",
        }
    )
    kinds = sum((has_mismatch, has_format, has_absence))
    if kinds >= 2:
        return _MSG_INCOMPLETE_MIXED
    if has_mismatch:
        return _MSG_INCOMPLETE_MISMATCH
    if has_format:
        return _MSG_INCOMPLETE_FORMAT
    return _MSG_INCOMPLETE


class GraphLike(Protocol):
    async def get(self, *a: Any, **k: Any) -> Any: ...
    async def get_bytes(self, *a: Any, **k: Any) -> Any: ...


@dataclass
class MergeReadiness:
    status: str  # ready|incomplete|unknown|already_merged
    expected_groups: int = 0
    ready_groups: int = 0
    missing_groups: int = 0
    missing_items: list[dict[str, Any]] = field(default_factory=list)
    folder_links: list[dict[str, Any]] = field(default_factory=list)
    checked_at: str = ""
    user_message: str = ""
    next_action: str = ""
    verification: dict[str, Any] | None = None


def _now() -> str:
    return now_colombia_iso()


def _result(
    status: str,
    *,
    expected: int = 0,
    ready: int = 0,
    missing_items: list[dict[str, Any]] | None = None,
    folder_links: list[dict[str, Any]] | None = None,
    user_message: str,
    next_action: str,
    verification: dict[str, Any] | None = None,
) -> MergeReadiness:
    missing_groups = max(0, expected - ready)
    return MergeReadiness(
        status=status,
        expected_groups=expected,
        ready_groups=ready,
        missing_groups=missing_groups,
        missing_items=list(missing_items or []),
        folder_links=list(folder_links or []),
        checked_at=_now(),
        user_message=user_message,
        next_action=next_action,
        verification=verification,
    )


def readiness_to_ui(readiness: MergeReadiness) -> UiMergeReadiness:
    return UiMergeReadiness(
        status=readiness.status,  # type: ignore[arg-type]
        expected_groups=readiness.expected_groups,
        ready_groups=readiness.ready_groups,
        missing_groups=readiness.missing_groups,
        missing_items=list(readiness.missing_items),
        folder_links=list(readiness.folder_links),
        checked_at=readiness.checked_at or None,
        user_message=readiness.user_message,
        next_action=readiness.next_action,
        verification=readiness.verification,
    )


def collect_asientos_folder_links(
    all_groups: list[tuple[str, str, list[dict[str, Any]]]],
) -> list[dict[str, Any]]:
    """Una entrada por carpeta ASIENTOS requerida (dedupe por path normalizado).

    Alineado con lo que Finalize provisionó en histórico: cada crédito/ruta del
    lote PAGO+ABONO debe aparecer aunque Graph falle al listar hijos.
    Mismo crédito / misma ruta → un solo link.
    """
    folder_links: list[dict[str, Any]] = []
    seen: set[str] = set()
    for _id_pago, _tipo, rows in all_groups:
        for row in rows:
            credit_digits = str(row.get("credito_digits") or "").strip()
            asientos_dir = _ruta_asientos_from_cell(row.get("ruta_asientos_cell"))
            if not asientos_dir:
                continue
            key = normalize_sharepoint_path(asientos_dir)
            if key in seen:
                continue
            seen.add(key)
            folder_links.append(
                {
                    "rel": "asientos",
                    "path": asientos_dir,
                    "credito": credit_digits,
                    "label": "Carpeta ASIENTOS",
                }
            )
    return folder_links


def _observed_pdfs_from_children(children: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Metadata ligera de PDFs en el nivel de la carpeta (sin descargar contenido)."""
    out: list[dict[str, Any]] = []
    for child in children:
        if not isinstance(child, dict):
            continue
        # Graph marca archivos con clave ``file`` (puede ser {}).
        if "folder" in child or "file" not in child:
            continue
        name = str(child.get("name") or "").strip()
        if not name or not name.lower().endswith(".pdf") or name.startswith("~$"):
            continue
        size_raw = child.get("size")
        size = int(size_raw) if isinstance(size_raw, int) else None
        out.append(
            {
                "name": name,
                "size": size,
                "etag": str(child.get("eTag") or child.get("etag") or "").strip() or None,
                "last_modified": str(child.get("lastModifiedDateTime") or "").strip()
                or None,
            }
        )
    return out


async def _attach_observed_pdfs(
    graph: GraphLike,
    site_id: str,
    drive_id: str,
    folder_links: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Lista ASIENTOS por carpeta y adjunta observed_pdfs (informativo; no cambia status)."""
    enriched: list[dict[str, Any]] = []
    for fl in folder_links:
        row = dict(fl)
        path = str(row.get("path") or "").strip()
        if not path:
            enriched.append(row)
            continue
        try:
            children = await _list_drive_folder_children(graph, site_id, drive_id, path)
            row["observed_pdfs"] = _observed_pdfs_from_children(children)
            row["list_ok"] = True
        except Exception:
            logger.info(
                "merge_readiness: list asientos (already_merged) falló path=%s",
                path,
                exc_info=True,
            )
            row["observed_pdfs"] = []
            row["list_ok"] = False
        enriched.append(row)
    return enriched


async def _load_historico_groups(
    graph: GraphLike,
    historico: str,
) -> tuple[str, str, list[tuple[str, str, list[dict[str, Any]]]]] | None:
    """Descarga histórico y agrupa filas PAGO/ABONO. None si no hay grupos."""
    ctx = await resolve_sharepoint_from_env(graph)
    site_id = str(ctx["site_id"])
    drive_id = str(ctx["drive_id"])
    hist_bytes = await _graph_download_by_path(graph, site_id, drive_id, historico)
    estado = (
        os.getenv("GRAPH_VALIDAR_EXTRACTO_ESTADO_CONTAINS") or "VALIDAR"
    ).strip() or "VALIDAR"
    wb = load_workbook(filename=BytesIO(hist_bytes), data_only=True)
    try:
        payment_rows = read_validated_payment_rows(wb, legacy_estado_token=estado)
        abono_rows = read_validated_abono_rows(wb)
        payment_groups = group_rows_by_id_pago(payment_rows)
        abono_groups = group_rows_by_id_pago(abono_rows)
    finally:
        wb.close()

    all_groups: list[tuple[str, str, list[dict[str, Any]]]] = []
    for id_pago, rows in payment_groups.items():
        all_groups.append((str(id_pago), "PAGO", list(rows)))
    for id_pago, rows in abono_groups.items():
        all_groups.append((str(id_pago), "ABONO", list(rows)))
    if not all_groups:
        return None
    return site_id, drive_id, all_groups


def _parse_asiento_events(
    pdf_bytes: bytes, *, credit: str, path: str
) -> list[dict[str, Any]] | None:
    """Asientos del PDF con el parser de fase 4; None si no se pueden leer."""
    key = hashlib.sha1(pdf_bytes or b"").hexdigest()
    if key in _events_cache:
        _events_cache.move_to_end(key)
        return _events_cache[key]
    try:
        events = parse_accounting_pdf_events(
            pdf_bytes,
            {"id_pago": "", "cliente": "", "credito": credit, "asiento_pdf_path": path},
        )
        facts: list[dict[str, Any]] | None = [event_facts(ev) for ev in events] or None
    except Exception:
        facts = None
    _lru_put(_events_cache, key, facts, _EVENTS_CACHE_MAX)
    return facts


async def _download_asiento(
    graph: GraphLike, site_id: str, drive_id: str, path: str, etag: str
) -> bytes:
    key = (normalize_sharepoint_path(path), etag)
    if etag and key in _bytes_cache:
        _bytes_cache.move_to_end(key)
        return _bytes_cache[key]
    blob = await _graph_download_by_path(graph, site_id, drive_id, path)
    if etag:
        _lru_put(_bytes_cache, key, bytes(blob), _BYTES_CACHE_MAX)
    return blob


def _table_facts_from_bytes(blob: bytes, fecha_limite_iso: str) -> dict[str, Any] | None:
    from app.application.services.amortization_workbook import (
        _cell_float,
        detect_amortization_sheet,
        find_row_by_due_date_detailed,
        scheduled_cuota_amount,
    )

    wb = load_workbook(filename=BytesIO(blob), data_only=True)
    try:
        match = detect_amortization_sheet(wb)
        ws, headers, header_row = match.worksheet, match.headers, match.header_row
        found = find_row_by_due_date_detailed(
            ws, headers, date.fromisoformat(fecha_limite_iso), header_row=header_row
        )
        if found.row is None:
            return None
        cuota = scheduled_cuota_amount(ws, headers, found.row)
        saldo: float | None = None
        col = headers.get("saldo_a_capital")
        if col:
            for r in range(found.row - 1, header_row, -1):
                saldo = _cell_float(ws.cell(r, col).value)
                if saldo is not None:
                    break
        return {
            "cuota": cents(cuota) if cuota else None,
            "saldo_capital": cents(saldo) if saldo is not None else None,
        }
    finally:
        wb.close()


async def _load_table_facts(
    graph: GraphLike, site_id: str, drive_id: str, tabla: str, fecha_limite_iso: str
) -> dict[str, Any] | None:
    """Cuota y saldo de capital de la tabla (solo lectura; None si no se puede leer)."""
    key = (normalize_sharepoint_path(tabla), fecha_limite_iso)
    hit = _facts_cache.get(key)
    if hit is not None and time.monotonic() - hit[0] <= _FACTS_TTL_SECONDS:
        return hit[1]
    try:
        blob = await _graph_download_by_path(graph, site_id, drive_id, tabla)
        facts = await asyncio.to_thread(_table_facts_from_bytes, bytes(blob), fecha_limite_iso)
    except Exception:
        logger.info("merge_readiness: tabla no legible para tipo path=%s", tabla, exc_info=True)
        facts = None
    _facts_cache[key] = (time.monotonic(), facts)
    return facts


def _processed_names(children: list[dict[str, Any]]) -> dict[str, str]:
    out: dict[str, str] = {}
    for child in children:
        if not isinstance(child, dict) or "folder" in child:
            continue
        name = str(child.get("name") or "").strip()
        if name.lower().endswith(".pdf"):
            out[name.casefold()] = str(child.get("lastModifiedDateTime") or "")
    return out


def _child_info(children: list[dict[str, Any]], name: str) -> tuple[str, str]:
    want = name.casefold()
    for child in children:
        if isinstance(child, dict) and str(child.get("name") or "").strip().casefold() == want:
            return (
                str(child.get("eTag") or child.get("etag") or "").strip(),
                str(child.get("webUrl") or "").strip(),
            )
    return "", ""


def _verification_message(verification: dict[str, Any]) -> str:
    total = int(verification.get("credits_total") or 0) - int(
        verification.get("credits_excluded") or 0
    )
    ready = int(verification.get("credits_ready") or 0)
    return f"{ready} de {total} créditos listos para consolidar."


async def assess_merge_readiness(
    graph: GraphLike,
    snap: ProcessControlSnapshot,
    bank_code: str,
) -> MergeReadiness:
    """Evalúa si hay asientos contables suficientes para consolidar (solo lectura)."""
    pk = (snap.process_key or "").strip()
    already_merged = control_indicates_already_merged(snap) or (
        bool(pk) and get_job_manager().has_completed_merge(pk)
    )

    historico = (snap.historical_file_path or "").strip().strip("/")
    email_pdf = (snap.email_pdf_path or "").strip().strip("/")
    if not historico or not email_pdf:
        if already_merged:
            return _result(
                "already_merged",
                user_message=_MSG_ALREADY,
                next_action=_NEXT_DONE,
            )
        return _result(
            "incomplete",
            user_message=_MSG_INCOMPLETE,
            next_action=_NEXT_LOAD,
        )

    try:
        loaded = await _load_historico_groups(graph, historico)
        if loaded is None:
            if already_merged:
                return _result(
                    "already_merged",
                    user_message=_MSG_ALREADY,
                    next_action=_NEXT_DONE,
                )
            return _result(
                "incomplete",
                user_message=_MSG_INCOMPLETE,
                next_action=_NEXT_LOAD,
            )

        site_id, drive_id, all_groups = loaded
        # Carpetas primero: no depender del listado Graph ni de grupos incompletos.
        folder_links = collect_asientos_folder_links(all_groups)

        # already_merged: conservar status, pero devolver folder_links + listado
        # ligero (nombre/size/etag/fecha) para recovery UX. No cambia gates.
        if already_merged:
            folder_links = await _attach_observed_pdfs(
                graph, site_id, drive_id, folder_links
            )
            expected = len(all_groups)
            return _result(
                "already_merged",
                expected=expected,
                ready=expected,
                missing_items=[],
                folder_links=folder_links,
                user_message=_MSG_ALREADY,
                next_action=_NEXT_DONE,
            )

        registry: dict[str, Any] | None = None
        registry_path = (getattr(snap, "decisions_registry_path", "") or "").strip()
        if registry_path:
            try:
                registry = await load_registry_for_process(
                    graph,
                    site_id,
                    drive_id,
                    registry_path=registry_path,
                    process_key=pk,
                    use_cache=True,
                )
            except DecisionsRegistryUnavailableError:
                logger.warning(
                    "merge_readiness: registro de decisiones no legible path=%s", registry_path
                )
                return _result(
                    "unknown",
                    expected=len(all_groups),
                    folder_links=folder_links,
                    user_message=_MSG_REGISTRY_UNAVAILABLE,
                    next_action=_NEXT_RETRY,
                )
        all_groups = apply_tipo_corrections(all_groups, registry)
        excluded_pagos = excluded_id_pagos(all_groups, registry)
        unused_keys = unused_asiento_paths(registry)

        ready = 0
        missing_items: list[dict[str, Any]] = []
        list_failed = False
        children_cache: dict[str, list[dict[str, Any]]] = {}
        processed_cache: dict[str, dict[str, str]] = {}
        failed_dirs: set[str] = set()
        asientos_read: dict[str, AsientoRead] = {}

        for id_pago, tipo, rows in all_groups:
            if id_pago in excluded_pagos:
                continue
            credit_items: list[dict[str, Any]] = []
            pre_skips: list[str] = []
            for row in rows:
                credit_digits = str(row.get("credito_digits") or "").strip()
                asientos_dir = _ruta_asientos_from_cell(row.get("ruta_asientos_cell"))
                if not asientos_dir:
                    pre_skips.append(
                        f"missing_ruta_asientos_contables credito={credit_digits}"
                    )
                    continue
                dir_key = normalize_sharepoint_path(asientos_dir)
                children = children_cache.get(dir_key)
                if children is None and dir_key not in failed_dirs:
                    try:
                        children = await _list_drive_folder_children(
                            graph, site_id, drive_id, asientos_dir
                        )
                    except Exception:
                        logger.info(
                            "merge_readiness: list asientos falló path=%s",
                            asientos_dir,
                            exc_info=True,
                        )
                        failed_dirs.add(dir_key)
                    else:
                        children_cache[dir_key] = children
                        try:
                            processed_children = await _list_drive_folder_children(
                                graph, site_id, drive_id, f"{asientos_dir.rstrip('/')}/PROCESADOS"
                            )
                        except Exception:
                            processed_children = []
                        processed_cache[dir_key] = _processed_names(processed_children)
                if children is None:
                    list_failed = True
                    pre_skips.append(
                        f"asientos_list_failed credito={credit_digits}"
                    )
                    continue
                processed = processed_cache.get(dir_key, {})
                names = _pdf_names_in_children(children)
                valid_names, rejected = _classify_asiento_pdf_names(
                    names, credit_digits
                )
                candidates: list[str] = []
                for asiento_name in valid_names:
                    asiento_rel = f"{asientos_dir}/{asiento_name}".replace("//", "/")
                    key = normalize_sharepoint_path(asiento_rel)
                    etag, web_url = _child_info(children, asiento_name)
                    if asiento_name.casefold() in processed:
                        # Mismo nombre que uno ya aplicado: el job de merge no lo usa.
                        asientos_read.setdefault(
                            key,
                            AsientoRead(
                                credito=credit_digits,
                                path=asiento_rel,
                                name=asiento_name,
                                web_url=web_url,
                                processed=True,
                                processed_at=processed[asiento_name.casefold()],
                            ),
                        )
                        continue
                    if asiento_rel.strip().strip("/").casefold() in unused_keys:
                        asientos_read.setdefault(
                            key,
                            AsientoRead(
                                credito=credit_digits,
                                path=asiento_rel,
                                name=asiento_name,
                                web_url=web_url,
                            ),
                        )
                        continue
                    candidates.append(asiento_name)
                if not candidates:
                    # Misma distinción que el job de merge (_pick_single_asiento_pdf):
                    # PDF presente pero sin dígitos del crédito → mismatch; carpeta
                    # vacía o sin PDF usable → not_found.
                    if rejected:
                        first_rejected = rejected[0]
                        hint = credit_hint_from_pdf_filename(
                            first_rejected, credit_digits
                        )
                        # Tokens sin espacios (parser de skip lines).
                        safe_name = first_rejected.replace(" ", "_")
                        skip = (
                            f"asiento_contable_credit_mismatch credito={credit_digits}"
                            f" asiento_pdf_found={safe_name}"
                        )
                        if hint:
                            skip = f"{skip} found_credit={hint}"
                        pre_skips.append(skip)
                    else:
                        pre_skips.append(
                            f"asiento_contable_not_found credito={credit_digits}"
                        )
                    continue

                format_failed = False
                usable_names: list[str] = []
                for asiento_name in candidates:
                    asiento_rel = f"{asientos_dir}/{asiento_name}".replace("//", "/")
                    key = normalize_sharepoint_path(asiento_rel)
                    safe_name = asiento_name.replace(" ", "_")
                    known = asientos_read.get(key)
                    code: str | None
                    if known is not None:
                        code = known.format_error or None
                    else:
                        etag, web_url = _child_info(children, asiento_name)
                        try:
                            pdf_bytes = await _download_asiento(
                                graph, site_id, drive_id, asiento_rel, etag
                            )
                        except Exception:
                            logger.info(
                                "merge_readiness: download asiento falló path=%s",
                                asiento_rel,
                                exc_info=True,
                            )
                            code = "asiento_download_failed"
                            pdf_bytes = b""
                        else:
                            code = classify_asiento_pdf_bytes(
                                pdf_bytes, credit=credit_digits, path=asiento_rel
                            )
                        asientos_read[key] = AsientoRead(
                            credito=credit_digits,
                            path=asiento_rel,
                            name=asiento_name,
                            web_url=web_url,
                            format_error=code or "",
                            events=None
                            if code
                            else _parse_asiento_events(
                                pdf_bytes, credit=credit_digits, path=asiento_rel
                            ),
                        )
                    if code:
                        if not format_failed:
                            pre_skips.append(
                                f"{code} credito={credit_digits}"
                                f" asiento_pdf_found={safe_name}"
                            )
                        format_failed = True
                        continue
                    usable_names.append(asiento_name)
                if format_failed or not usable_names:
                    continue

                row_include_extract = bool(
                    row.get(
                        "include_extract_in_composite",
                        row.get("requiere_extracto"),
                    )
                )
                row_rol = str(row.get("rol_extracto") or ExtractRole.CIERRE_CUOTA).strip().upper()
                if (
                    tipo == "PAGO"
                    and row_include_extract
                    and row_rol
                    in (
                        ExtractRole.CIERRE_CUOTA,
                        ExtractRole.REFERENCIA_MORA,
                        ExtractRole.REFERENCIA_SALDO,
                    )
                ):
                    extract_paths: list[str] = []
                    for p in await _collect_pdf_paths_from_ruta_cell(
                        graph, site_id, drive_id, row.get("ruta_cell")
                    ):
                        extract_paths.append(p)
                    if not extract_paths:
                        pre_skips.append(
                            f"extract_routes_missing credito={credit_digits}"
                        )
                        continue

                credit_items.append(
                    {
                        "credito": credit_digits,
                        "asiento_pdf_paths": [
                            f"{asientos_dir}/{n}".replace("//", "/")
                            for n in usable_names
                        ],
                    }
                )

            validation = validate_merge_group_completeness(
                id_pago=id_pago,
                tipo_aplicacion=tipo,
                group_rows=rows,
                credit_items=credit_items,
                pre_skips=pre_skips,
            )
            if validation.is_complete:
                ready += 1
            else:
                for item in validation.missing_inputs:
                    missing_items.append(
                        {
                            "id_pago": id_pago,
                            "tipo_aplicacion": tipo,
                            **dict(item),
                        }
                    )

        evaluable_credits = {
            str(a.credito) for a in asientos_read.values() if a.events
        }
        table_facts: dict[str, dict[str, Any]] = {}
        for tok, (tabla, limite) in table_fact_requests(all_groups).items():
            if tok not in evaluable_credits:
                continue
            facts = await _load_table_facts(graph, site_id, drive_id, tabla, limite)
            if facts:
                table_facts[tok] = facts
        verification = evaluate_verification(
            all_groups,
            asientos=list(asientos_read.values()),
            missing_items=missing_items,
            registry=registry,
            table_facts=table_facts,
        )

        expected = sum(1 for g in all_groups if g[0] not in excluded_pagos)
        if list_failed:
            return _result(
                "unknown",
                expected=expected,
                ready=ready,
                missing_items=missing_items,
                folder_links=folder_links,
                user_message=_MSG_UNKNOWN,
                next_action=_NEXT_RETRY,
                verification=verification,
            )
        if verification["all_excluded"]:
            return _result(
                "incomplete",
                expected=expected,
                ready=ready,
                missing_items=missing_items,
                folder_links=folder_links,
                user_message=_MSG_ALL_EXCLUDED,
                next_action=_NEXT_REVIEW,
                verification=verification,
            )
        attention = int(verification.get("credits_attention") or 0) > 0
        if ready >= expected and expected > 0 and not missing_items and not attention:
            return _result(
                "ready",
                expected=expected,
                ready=ready,
                missing_items=[],
                folder_links=folder_links,
                user_message=_MSG_READY,
                next_action=_NEXT_MERGE,
                verification=verification,
            )
        return _result(
            "incomplete",
            expected=expected,
            ready=ready,
            missing_items=missing_items,
            folder_links=folder_links,
            user_message=(
                _incomplete_user_message(missing_items)
                if missing_items
                else _verification_message(verification)
            ),
            next_action=_NEXT_LOAD if missing_items else _NEXT_REVIEW,
            verification=verification,
        )
    except Exception:
        logger.info(
            "merge_readiness: fallo Graph/histórico bank=%s",
            bank_code,
            exc_info=True,
        )
        if already_merged:
            return _result(
                "already_merged",
                user_message=_MSG_ALREADY,
                next_action=_NEXT_DONE,
            )
        return _result(
            "unknown",
            user_message=_MSG_UNKNOWN,
            next_action=_NEXT_RETRY,
        )
