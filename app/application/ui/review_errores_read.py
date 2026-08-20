"""Lectura read-only de la hoja Errores del Excel de revisión (UI).

Expone filas abiertas con contexto operativo (crédito, archivos, carpeta)
para avisar al operador antes de Finalize.
"""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from typing import Any, Protocol

from app.application.services.review_schema import ErroresCols, ReviewSheets
from app.application.sharepoint_resolution import encode_graph_drive_path
from app.application.use_cases.payment_validation_finalize import (
    _find_table_header_row,
)

logger = logging.getLogger(__name__)


class _GraphBytes(Protocol):
    async def get_bytes(self, *a: Any, **k: Any) -> bytes: ...


@dataclass(frozen=True)
class ReviewErrorRow:
    """Fila abierta de la hoja Errores."""

    row_number: int
    id_pago: str
    cliente: str
    credito: str
    tipo_caso: str
    descripcion: str
    que_debe_hacer: str
    requiere_soporte: str
    codigo_tecnico: str
    extract_label: str
    folder_label: str
    extract_url: str | None
    folder_url: str | None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _hyperlink_target(cell: Any) -> str | None:
    try:
        hl = getattr(cell, "hyperlink", None)
        if hl is None:
            return None
        target = getattr(hl, "target", None) or getattr(hl, "ref", None)
        text = str(target or "").strip()
        return text or None
    except Exception:
        return None


def parse_review_errores_workbook(workbook: Any) -> list[ReviewErrorRow]:
    """Parsea filas abiertas de Errores (misma regla que Finalize)."""
    if ReviewSheets.ERRORES not in getattr(workbook, "sheetnames", []):
        return []
    ws = workbook[ReviewSheets.ERRORES]
    try:
        header_row = _find_table_header_row(ws, ErroresCols.ID_PAGO)
    except ValueError:
        return []

    headers: list[str] = []
    out: list[ReviewErrorRow] = []
    for r_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
        if r_idx < header_row:
            continue
        if r_idx == header_row:
            headers = [_cell_text(c.value) for c in row]
            continue
        values = [c.value for c in row]
        if not any(v is not None and str(v).strip() for v in values):
            continue
        rd = dict(zip(headers, values))
        id_pago = _cell_text(rd.get(ErroresCols.ID_PAGO))
        descripcion = _cell_text(rd.get(ErroresCols.DESCRIPCION))
        codigo = _cell_text(rd.get(ErroresCols.CODIGO_TECNICO))
        if not (id_pago or descripcion or codigo):
            continue

        by_header = {h: c for h, c in zip(headers, row)}
        ext_cell = by_header.get(ErroresCols.LINK_EXTRACTO)
        fold_cell = by_header.get(ErroresCols.LINK_CARPETA_CREDITO)
        extract_label = _cell_text(ext_cell.value) if ext_cell is not None else ""
        folder_label = _cell_text(fold_cell.value) if fold_cell is not None else ""
        extract_url = _hyperlink_target(ext_cell) if ext_cell is not None else None
        folder_url = _hyperlink_target(fold_cell) if fold_cell is not None else None
        if not extract_url and extract_label.lower().startswith("http"):
            extract_url = extract_label
        if not folder_url and folder_label.lower().startswith("http"):
            folder_url = folder_label

        out.append(
            ReviewErrorRow(
                row_number=r_idx,
                id_pago=id_pago,
                cliente=_cell_text(rd.get(ErroresCols.CLIENTE)),
                credito=_cell_text(rd.get(ErroresCols.CREDITO)),
                tipo_caso=_cell_text(rd.get(ErroresCols.TIPO_CASO)),
                descripcion=descripcion,
                que_debe_hacer=_cell_text(rd.get(ErroresCols.QUE_DEBE_HACER)),
                requiere_soporte=_cell_text(rd.get(ErroresCols.REQUIERE_SOPORTE)),
                codigo_tecnico=codigo,
                extract_label=extract_label,
                folder_label=folder_label,
                extract_url=extract_url,
                folder_url=folder_url,
            )
        )
    return out


async def load_review_errores_from_path(
    graph: _GraphBytes,
    *,
    site_id: str,
    drive_id: str,
    validation_file_path: str,
) -> list[ReviewErrorRow]:
    """Descarga el Excel de revisión y devuelve filas abiertas de Errores."""
    rel = (validation_file_path or "").strip().strip("/")
    if not rel or not site_id or not drive_id:
        return []
    endpoint = (
        f"/sites/{site_id}/drives/{drive_id}/root:"
        f"/{encode_graph_drive_path(rel)}:/content"
    )
    try:
        raw = await graph.get_bytes(endpoint)
    except Exception:
        logger.info(
            "review_errores: no se pudo descargar validation path=%s",
            rel,
            exc_info=True,
        )
        return []

    try:
        from openpyxl import load_workbook

        wb = load_workbook(filename=io.BytesIO(raw), data_only=False)
    except Exception:
        logger.info("review_errores: Excel ilegible path=%s", rel, exc_info=True)
        return []

    try:
        return parse_review_errores_workbook(wb)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def build_operational_issues_from_review_errores(
    rows: list[ReviewErrorRow],
    *,
    bank_input_link: Any | None = None,
    bank_folder_link: Any | None = None,
    clients_base_link: Any | None = None,
    file_name: str | None = None,
) -> list[Any]:
    """Convierte filas Errores → UiOperationalIssue (import diferido para evitar ciclos)."""
    from app.application.ui.review_error_links import (
        ReviewErrorLinkContext,
        build_review_error_issue_links,
    )
    from app.application.ui.schemas import (
        UiIssueLocation,
        UiIssueRetry,
        UiOperationalIssue,
    )

    link_ctx = ReviewErrorLinkContext(
        bank_input_link=bank_input_link,
        bank_folder_link=bank_folder_link,
        clients_base_link=clients_base_link,
    )

    issues: list[UiOperationalIssue] = []
    for idx, row in enumerate(rows):
        involved: list[str] = []
        if row.extract_label:
            involved.append(f"Extracto: {row.extract_label}")
        if row.folder_label:
            involved.append(f"Carpeta: {row.folder_label}")
        involved_txt = " · ".join(involved) if involved else None

        title = row.tipo_caso or "Caso en hoja Errores"
        if row.credito:
            from app.application.services.review_schema import normalize_credito_digits

            cred_disp = normalize_credito_digits(row.credito) or row.credito
            title = f"{title} · Crédito {cred_disp}"

        descripcion = row.descripcion
        que_hacer = row.que_debe_hacer
        if not descripcion or not que_hacer:
            from app.application.services.review_error_guide import guide_texts_for_ui

            tipo_g, descr_g, hacer_g = guide_texts_for_ui(
                row.codigo_tecnico,
                {
                    "cliente": row.cliente,
                    "credito": row.credito,
                    "code": row.codigo_tecnico,
                },
            )
            if not row.tipo_caso:
                title = tipo_g or title
                if row.credito:
                    title = f"{title} · Crédito {row.credito}"
            descripcion = descripcion or descr_g
            que_hacer = que_hacer or hacer_g

        message_parts = [p for p in (descripcion,) if p]
        user_message = " ".join(message_parts) if message_parts else (
            "Hay un caso pendiente en la hoja Errores del Excel de revisión."
        )
        if row.extract_label and row.extract_label not in user_message:
            user_message = f"{user_message} Archivo: «{row.extract_label}»."

        links = build_review_error_issue_links(row, context=link_ctx)

        issues.append(
            UiOperationalIssue(
                issue_id=f"review-errores-{idx}-{row.row_number}",
                stage="generate",
                category="correction_required",
                severity="business",
                recoverable=True,
                title=title,
                user_message=user_message,
                location=UiIssueLocation(
                    file_name=file_name,
                    sheet="Errores",
                    row=row.row_number,
                    credit=row.credito or None,
                    payment_id=row.id_pago or None,
                    client_name=row.cliente or None,
                ),
                value_found=row.codigo_tecnico or None,
                next_action=que_hacer
                or (
                    "Corrija o retire el archivo indicado en SharePoint y use "
                    "«Regenerar archivo de revisión»."
                ),
                retry=UiIssueRetry(
                    allowed=True,
                    action="regenerate",
                    label="Regenerar archivo de revisión",
                ),
                links=links,
                technical_reference=row.codigo_tecnico or None,
            )
        )
    return issues
