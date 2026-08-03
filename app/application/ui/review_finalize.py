"""Finalize atómico desde la UI de revisión (R2).

Cadena: ETag → patches opcionales → preflight → Procesar=SI → put Excel → enqueue.
Si preflight falla: no escribe ni encola. Si ETag falla: 409.
"""
from __future__ import annotations

import io
import logging
from typing import Any, Protocol

from fastapi import BackgroundTasks
from openpyxl import load_workbook

from app.application.job_manager import get_job_manager
from app.application.services.finalize_queue_service import (
    FinalizeQueueAccepted,
    FinalizeQueueBusyError,
    FinalizeQueueValidationError,
    get_finalize_queue_service,
)
from app.application.services.review_schema import ControlCols, ReviewSheets
from app.application.sharepoint_resolution import (
    encode_graph_drive_path,
    resolve_sharepoint_from_env,
)
from app.application.ui.finalize_resolve import (
    FinalizeProcessIdentityError,
    resolve_finalize_target_from_control,
)
from app.application.ui.ports import UiSharePointReadPort
from app.application.ui.review_preflight import (
    collect_preflight_issues_from_workbook,
)
from app.application.ui.review_read import ReviewFileMissingError
from app.application.ui.review_write import (
    ReviewEtagConflictError,
    ReviewFileLockedError,
    ReviewPatchValidationError,
    apply_patches_to_workbook,
    etags_match,
)
from app.application.ui.schemas import (
    UiReviewFinalizeAccepted,
    UiReviewFinalizeRequest,
    UiReviewPreflightIssue,
)
from app.application.job_status_enrichment import finalize_message_for_code

logger = logging.getLogger(__name__)


class ReviewPreflightBlockedError(Exception):
    """Preflight de negocio falló; no se escribe ni se encola Finalize."""

    def __init__(self, issues: list[UiReviewPreflightIssue]) -> None:
        self.issues = issues
        super().__init__("review_preflight_blocked")


class _GraphWrite(Protocol):
    async def put_bytes(self, *a: Any, **k: Any) -> Any: ...


def _issue_models(raw_issues: list[dict[str, Any]]) -> list[UiReviewPreflightIssue]:
    out: list[UiReviewPreflightIssue] = []
    for raw in raw_issues:
        code = str(raw.get("error_code") or "unknown").strip() or "unknown"
        user_msg, _next = finalize_message_for_code(code)
        excel_row = raw.get("excel_row")
        try:
            row_int = int(excel_row) if excel_row is not None else None
        except (TypeError, ValueError):
            row_int = None
        value = raw.get("value_found")
        out.append(
            UiReviewPreflightIssue(
                error_code=code,
                sheet=str(raw.get("sheet") or "").strip() or None,
                excel_row=row_int,
                id_pago=str(raw.get("id_pago") or "").strip() or None,
                credito=str(raw.get("credito") or "").strip() or None,
                field=str(raw.get("field") or "").strip() or None,
                value_found=None if value is None else str(value),
                user_message=user_msg,
            )
        )
    return out


def set_control_procesar_si(workbook: Any) -> bool:
    """Marca Procesar=SI y asegura Estado=EN_REVISION en hoja Control del review.

    Retorna True si encontró y actualizó la fila Procesar.
    """
    if ReviewSheets.CONTROL not in workbook.sheetnames:
        raise ReviewPatchValidationError(
            "missing_control_sheet",
            "El Excel de revisión no tiene hoja Control.",
        )
    ws = workbook[ReviewSheets.CONTROL]
    found_procesar = False
    for r in range(1, (ws.max_row or 1) + 1):
        label = str(ws.cell(r, 1).value or "").strip()
        if label == ControlCols.ROW_PROCESAR:
            ws.cell(r, 2).value = ControlCols.VAL_PROCESAR_SI
            found_procesar = True
        elif label == ControlCols.ROW_ESTADO:
            current = str(ws.cell(r, 2).value or "").strip().upper()
            if current != "EN_REVISION":
                ws.cell(r, 2).value = "EN_REVISION"
    if not found_procesar:
        raise ReviewPatchValidationError(
            "missing_procesar_row",
            "No se encontró la fila Procesar en la hoja Control.",
        )
    return True


async def finalize_ui_review_atomic(
    reader: UiSharePointReadPort,
    graph: _GraphWrite,
    *,
    process_key: str,
    banks: tuple[str, ...],
    body: UiReviewFinalizeRequest,
    if_match: str | None,
    background_tasks: BackgroundTasks,
    requested_by: str | None,
    ui_request_id: str | None,
) -> UiReviewFinalizeAccepted:
    if not (if_match or "").strip():
        raise ReviewPatchValidationError(
            "missing_if_match",
            "Debe enviar el header If-Match con el etag actual.",
        )

    if get_job_manager().is_generate_or_finalize_active():
        raise ReviewPatchValidationError(
            "mutation_in_progress",
            "Hay un Generate o Finalize en curso; espere a que termine.",
        )

    key = (process_key or "").strip()
    matched_bank: str | None = None
    snap = None
    for bank in banks:
        try:
            control = await reader.read_process_control(bank)
        except Exception:
            continue
        if (control.snapshot.process_key or "").strip() == key:
            matched_bank = bank
            snap = control.snapshot
            break
    if snap is None or matched_bank is None:
        raise KeyError(key)

    # Misma identidad que Finalize UI clásico (Control técnico)
    try:
        target = resolve_finalize_target_from_control(
            snap,
            bank_code=matched_bank,
            process_key=key,
        )
    except FinalizeProcessIdentityError:
        raise

    path = (target.validation_file_path or "").strip()
    if not path:
        raise ReviewFileMissingError("validation_file_path_empty")

    meta = await reader.get_item_meta(path)
    current_etag = meta.etag if meta else None
    if not etags_match(if_match, current_etag):
        raise ReviewEtagConflictError(current_etag)

    file_content = await reader.download_bytes(path)
    dl_etag = file_content.etag or current_etag
    if not etags_match(if_match, dl_etag):
        raise ReviewEtagConflictError(dl_etag)

    try:
        wb = load_workbook(io.BytesIO(file_content.content), data_only=False)
    except Exception as exc:
        raise ValueError("review_workbook_unreadable") from exc

    changes = list(body.changes or [])
    updated = apply_patches_to_workbook(wb, changes)

    raw_issues = collect_preflight_issues_from_workbook(wb)
    if raw_issues:
        raise ReviewPreflightBlockedError(_issue_models(raw_issues))

    set_control_procesar_si(wb)

    buf = io.BytesIO()
    wb.save(buf)
    payload = buf.getvalue()

    ctx = await resolve_sharepoint_from_env(graph)  # type: ignore[arg-type]
    site_id = str(ctx["site_id"])
    drive_id = str(ctx["drive_id"])
    enc = encode_graph_drive_path(path)
    endpoint = f"/sites/{site_id}/drives/{drive_id}/root:/{enc}:/content"
    # If-Match en el PUT: cierra la carrera entre el check local y SharePoint.
    try:
        put_resp = await graph.put_bytes(
            endpoint,
            payload,
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            if_match=if_match.strip(),
        )
    except Exception as exc:
        status = getattr(getattr(exc, "response", None), "status_code", None)
        msg = str(exc)
        if status == 423 or "423" in msg or "Locked" in msg or "bloqueado" in msg.lower():
            raise ReviewFileLockedError(msg[:400]) from exc
        raise
    new_etag: str | None = None
    if isinstance(put_resp, dict):
        new_etag = put_resp.get("eTag") or put_resp.get("etag")
    new_etag = new_etag or None

    svc = get_finalize_queue_service()
    try:
        accepted: FinalizeQueueAccepted = await svc.enqueue(
            graph=graph,
            background_tasks=background_tasks,
            bank_code=target.bank_code,
            validation_file_path=target.validation_file_path,
            process_key=target.process_key,
            trigger_source="web_ui",
            requested_by=requested_by,
            ui_request_id=ui_request_id,
        )
    except (FinalizeQueueBusyError, FinalizeQueueValidationError):
        raise

    return UiReviewFinalizeAccepted(
        accepted=True,
        action="finalize",
        bank_code=target.bank_code,
        process_key=target.process_key,
        job_id=accepted.job_id,
        status=accepted.status,
        poll_url=f"/api/ui/v1/jobs/{accepted.job_id}",
        etag=new_etag,
        updated_row_keys=updated,
    )
