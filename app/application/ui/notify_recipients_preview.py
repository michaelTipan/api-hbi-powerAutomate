"""Vista previa de EMISOR/RECEPTORES desde CORREOS.xlsx (solo lectura).

Misma fuente que Notify al enviar: no inventa overrides ni lee env de sandbox.
"""

from __future__ import annotations

import logging
import os
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.application.config.payment_validation_settings import resolve_correos_xlsx_path
from app.application.sharepoint_resolution import (
    require_operations_site_config,
    resolve_sharepoint_path,
)
from app.application.use_cases.send_validar_extractos_notification import (
    _collect_emails_from_cell,
    _dedupe_emails_preserve_order,
    _find_correos_header_row,
)
from app.application.use_cases.validate_payment_report import (
    _graph_download_by_path,
    _graph_get_item_metadata_by_path,
)
from app.domain.ports.graph import GraphApiPort

logger = logging.getLogger(__name__)

_AUTOSAVE_HINT = (
    "Si acabas de editar CORREOS.xlsx en Excel Online, espera unos segundos, "
    "guarda y pulsa Actualizar antes de enviar."
)


def _parse_correos_with_sheet(wb: Any) -> tuple[str, list[str], str]:
    """Devuelve (emisor, receptores, hoja) con la misma regla que el envío."""
    last_error: str | None = None
    for ws in wb.worksheets:
        found = _find_correos_header_row(ws)
        if not found:
            continue
        h_row, col_em, col_rec = found
        last = ws.max_row or h_row
        sender: str | None = None
        recipients: list[str] = []
        for r in range(h_row + 1, last + 1):
            if sender is None:
                em = _collect_emails_from_cell(ws.cell(row=r, column=col_em).value)
                if em:
                    sender = em[0]
            recipients.extend(_collect_emails_from_cell(ws.cell(row=r, column=col_rec).value))
        recipients = _dedupe_emails_preserve_order(recipients)
        if sender and recipients:
            return sender, recipients, str(ws.title)
        if sender and not recipients:
            last_error = f"Hoja {ws.title!r}: hay EMISOR pero RECEPTORES vacío."
        elif not sender and recipients:
            last_error = f"Hoja {ws.title!r}: hay RECEPTORES pero EMISOR vacío."
        else:
            last_error = f"Hoja {ws.title!r}: sin datos útiles."
    msg = (
        'No se encontró en CORREOS.xlsx una hoja con columnas "EMISOR" y "RECEPTORES" '
        "y al menos un remitente y un destinatario."
    )
    if last_error:
        msg += f" Detalle: {last_error}"
    raise ValueError(msg)


async def load_notify_recipients_preview(graph: GraphApiPort) -> dict[str, Any]:
    """Lee CORREOS.xlsx y metadatos; no envía correo ni adquiere locks."""
    require_operations_site_config()
    rel = resolve_correos_xlsx_path().strip().strip("/")
    site_search = (os.getenv("GRAPH_SHAREPOINT_SITE_SEARCH") or "").strip()
    drive_name = (os.getenv("GRAPH_SHAREPOINT_DRIVE_NAME") or "").strip()
    info = await resolve_sharepoint_path(graph, site_search, drive_name, rel)
    site_id = str(info["site_id"])
    drive_id = str(info["drive_id"])
    meta = await _graph_get_item_metadata_by_path(graph, site_id, drive_id, rel)
    last_modified = str(meta.get("lastModifiedDateTime") or "").strip() or None
    data = await _graph_download_by_path(graph, site_id, drive_id, rel)
    wb = load_workbook(filename=BytesIO(data), data_only=True)
    try:
        sender, recipients, sheet = _parse_correos_with_sheet(wb)
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()
    to_effective = [e for e in recipients if e.lower() != sender.lower()]
    warnings: list[str] = [_AUTOSAVE_HINT]
    if not to_effective:
        warnings.append(
            "Tras excluir el emisor de RECEPTORES no queda ningún destinatario válido."
        )
    return {
        "ok": True,
        "source_path": rel,
        "sheet": sheet,
        "emisor": sender,
        "receptores": to_effective,
        "receptores_raw_count": len(recipients),
        "file_last_modified": last_modified,
        "warnings": warnings,
        "user_message": (
            f"Se enviará desde {sender} a {', '.join(to_effective)}."
            if to_effective
            else "No hay destinatarios efectivos en CORREOS.xlsx."
        ),
    }
