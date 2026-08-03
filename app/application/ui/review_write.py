"""Escritura parcial del Excel de revisión (R1).

Solo columnas whitelist. Borrador incompleto permitido.
ETag obligatorio (If-Match). No marca Procesar=SI ni encola Finalize.
"""
from __future__ import annotations

import io
import logging
from decimal import Decimal, InvalidOperation
from typing import Any, Protocol

from openpyxl import load_workbook

from app.application.job_manager import get_job_manager
from app.application.services.review_schema import (
    DistribucionAbonosCols,
    DistribucionCols,
    ReviewSheets,
    find_distribucion_pagos_sheet,
)
from app.application.sharepoint_resolution import (
    encode_graph_drive_path,
    resolve_sharepoint_from_env,
)
from app.application.ui.ports import UiSharePointReadPort
from app.application.ui.review_read import (
    EDITABLE_ABONO_API_FIELDS,
    EDITABLE_PAGO_API_FIELDS,
    ReviewFileMissingError,
    load_ui_review_for_process,
    make_row_key,
)
from app.application.ui.schemas import (
    UiReviewPatchRequest,
    UiReviewPatchResponse,
    UiReviewRowPatch,
)
from app.application.use_cases.payment_validation_finalize import (
    _find_table_header_row,
)

logger = logging.getLogger(__name__)

# API field → Excel header (pagos)
_PAGO_FIELD_MAP: dict[str, str] = {
    "aplicar_a_extracto": DistribucionCols.APLICAR_A_EXTRACTO,
    "mora_a_aplicar": DistribucionCols.MORA_A_APLICAR,
    "abono_a_capital": DistribucionCols.ABONO_A_CAPITAL,
    "otros_valores": DistribucionCols.OTROS_VALORES,
    "estado_pago": DistribucionCols.ESTADO_PAGO,
    "validar_pago": DistribucionCols.VALIDAR_PAGO,
    "observacion": DistribucionCols.OBSERVACION,
}

_ABONO_FIELD_MAP: dict[str, str] = {
    "validar_abono": DistribucionAbonosCols.VALIDAR_ABONO,
}

_MONEY_FIELDS = {
    "aplicar_a_extracto",
    "mora_a_aplicar",
    "abono_a_capital",
    "otros_valores",
}

# Asegurar paridad con whitelist publicada en GET
assert set(_PAGO_FIELD_MAP) == set(EDITABLE_PAGO_API_FIELDS)
assert set(_ABONO_FIELD_MAP) == set(EDITABLE_ABONO_API_FIELDS)


class ReviewEtagConflictError(Exception):
    def __init__(self, current_etag: str | None) -> None:
        self.current_etag = current_etag
        super().__init__("review_etag_conflict")


class ReviewPatchValidationError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        super().__init__(message)


class _GraphWrite(Protocol):
    async def put_bytes(self, *a: Any, **k: Any) -> Any: ...


def normalize_etag(value: str | None) -> str:
    raw = (value or "").strip()
    if raw.upper().startswith("W/"):
        raw = raw[2:].strip()
    if len(raw) >= 2 and raw[0] == raw[-1] == '"':
        raw = raw[1:-1]
    return raw.strip()


def etags_match(expected: str | None, actual: str | None) -> bool:
    a = normalize_etag(expected)
    b = normalize_etag(actual)
    if not a or not b:
        return False
    return a == b


def parse_row_key(row_key: str) -> tuple[str, str, str]:
    parts = [p.strip() for p in (row_key or "").split("|")]
    if len(parts) < 3:
        raise ReviewPatchValidationError(
            "invalid_row_key",
            f"row_key inválido: {row_key!r}",
        )
    sheet, id_pago, credito = parts[0], parts[1], parts[2]
    if not sheet:
        raise ReviewPatchValidationError("invalid_row_key", "sheet vacío")
    return sheet, id_pago, credito


def _parse_money(value: Any) -> Decimal:
    if value is None or value == "":
        raise ReviewPatchValidationError("invalid_decimal", "Monto vacío")
    if isinstance(value, Decimal):
        return value.quantize(Decimal("0.01"))
    if isinstance(value, bool):
        raise ReviewPatchValidationError("invalid_decimal", "Monto inválido")
    if isinstance(value, (int, float)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = str(value).strip().replace(",", "")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError) as exc:
        raise ReviewPatchValidationError(
            "invalid_decimal", f"Monto inválido: {value!r}"
        ) from exc


def _excel_money(value: Decimal) -> float:
    return float(value.quantize(Decimal("0.01")))


def _column_map(ws: Any, header_row: int) -> dict[str, int]:
    m: dict[str, int] = {}
    for c in range(1, (ws.max_column or 1) + 1):
        h = str(ws.cell(header_row, c).value or "").strip()
        if h:
            m[h] = c
    return m


def _find_data_row(
    ws: Any,
    *,
    header_row: int,
    col_map: dict[str, int],
    id_pago: str,
    credito: str,
) -> int | None:
    id_col = col_map.get(DistribucionCols.ID_PAGO) or col_map.get(
        DistribucionAbonosCols.ID_PAGO
    )
    cred_col = col_map.get(DistribucionCols.CREDITO) or col_map.get(
        DistribucionAbonosCols.CREDITO
    )
    if not id_col or not cred_col:
        return None
    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        rid = str(ws.cell(r, id_col).value or "").strip()
        rcred = str(ws.cell(r, cred_col).value or "").strip()
        if rid == id_pago and rcred == credito:
            return r
    return None


def apply_patches_to_workbook(
    workbook: Any,
    changes: list[UiReviewRowPatch],
) -> list[str]:
    """Aplica cambios in-memory. Retorna row_keys actualizados."""
    if not changes:
        return []
    updated: list[str] = []

    try:
        ws_pagos = find_distribucion_pagos_sheet(workbook)
        pagos_header = _find_table_header_row(ws_pagos, DistribucionCols.ID_PAGO)
        pagos_map = _column_map(ws_pagos, pagos_header)
    except ValueError:
        ws_pagos = None
        pagos_header = 1
        pagos_map = {}

    ws_abonos = None
    abonos_header = 1
    abonos_map: dict[str, int] = {}
    if ReviewSheets.DISTRIBUCION_ABONOS in workbook.sheetnames:
        ws_abonos = workbook[ReviewSheets.DISTRIBUCION_ABONOS]
        try:
            abonos_header = _find_table_header_row(
                ws_abonos, DistribucionAbonosCols.ID_PAGO
            )
            abonos_map = _column_map(ws_abonos, abonos_header)
        except ValueError:
            ws_abonos = None

    for change in changes:
        sheet, id_pago, credito = parse_row_key(change.row_key)
        fields = change.fields or {}
        if not fields:
            continue

        is_abono = sheet == ReviewSheets.DISTRIBUCION_ABONOS
        if is_abono:
            if ws_abonos is None:
                raise ReviewPatchValidationError(
                    "row_not_found", f"Hoja abonos ausente para {change.row_key}"
                )
            field_map = _ABONO_FIELD_MAP
            ws = ws_abonos
            header_row = abonos_header
            col_map = abonos_map
        else:
            if ws_pagos is None:
                raise ReviewPatchValidationError(
                    "row_not_found", f"Hoja pagos ausente para {change.row_key}"
                )
            field_map = _PAGO_FIELD_MAP
            ws = ws_pagos
            header_row = pagos_header
            col_map = pagos_map

        for api_field in fields:
            if api_field not in field_map:
                raise ReviewPatchValidationError(
                    "field_not_editable",
                    f"Campo no editable: {api_field}",
                )

        row_idx = _find_data_row(
            ws,
            header_row=header_row,
            col_map=col_map,
            id_pago=id_pago,
            credito=credito,
        )
        if row_idx is None:
            raise ReviewPatchValidationError(
                "row_not_found",
                f"No se encontró la fila {change.row_key}",
            )

        for api_field, raw_val in fields.items():
            header = field_map[api_field]
            col = col_map.get(header)
            if not col:
                raise ReviewPatchValidationError(
                    "field_not_editable",
                    f"Columna ausente en Excel: {header}",
                )
            if api_field in _MONEY_FIELDS:
                if raw_val is None or raw_val == "":
                    ws.cell(row_idx, col).value = None
                else:
                    ws.cell(row_idx, col).value = _excel_money(_parse_money(raw_val))
            else:
                if raw_val is None:
                    ws.cell(row_idx, col).value = None
                else:
                    ws.cell(row_idx, col).value = str(raw_val).strip()

        canon_sheet = (
            ReviewSheets.DISTRIBUCION_ABONOS
            if is_abono
            else ReviewSheets.DISTRIBUCION_PAGOS
        )
        updated.append(make_row_key(canon_sheet, id_pago, credito))

    return updated


async def patch_ui_review(
    reader: UiSharePointReadPort,
    graph: _GraphWrite,
    *,
    process_key: str,
    banks: tuple[str, ...],
    body: UiReviewPatchRequest,
    if_match: str | None,
) -> UiReviewPatchResponse:
    if not (if_match or "").strip():
        raise ReviewPatchValidationError(
            "missing_if_match",
            "Debe enviar el header If-Match con el etag actual.",
        )

    if get_job_manager().is_generate_or_finalize_active():
        raise ReviewPatchValidationError(
            "mutation_in_progress",
            "Hay un Generate o Finalize en curso; espere a que termine.",
        )

    key = (process_key or "").strip()
    matched_bank: str | None = None
    snap = None
    for bank in banks:
        try:
            control = await reader.read_process_control(bank)
        except Exception:
            continue
        if (control.snapshot.process_key or "").strip() == key:
            matched_bank = bank
            snap = control.snapshot
            break
    if snap is None or matched_bank is None:
        raise KeyError(key)

    path = (snap.validation_file_path or "").strip()
    if not path:
        raise ReviewFileMissingError("validation_file_path_empty")

    meta = await reader.get_item_meta(path)
    current_etag = meta.etag if meta else None
    if not etags_match(if_match, current_etag):
        raise ReviewEtagConflictError(current_etag)

    file_content = await reader.download_bytes(path)
    dl_etag = file_content.etag or current_etag
    if not etags_match(if_match, dl_etag):
        raise ReviewEtagConflictError(dl_etag)

    try:
        wb = load_workbook(io.BytesIO(file_content.content), data_only=False)
    except Exception as exc:
        raise ValueError("review_workbook_unreadable") from exc

    updated = apply_patches_to_workbook(wb, list(body.changes or []))

    buf = io.BytesIO()
    wb.save(buf)
    payload = buf.getvalue()

    ctx = await resolve_sharepoint_from_env(graph)  # type: ignore[arg-type]
    site_id = str(ctx["site_id"])
    drive_id = str(ctx["drive_id"])
    enc = encode_graph_drive_path(path)
    endpoint = f"/sites/{site_id}/drives/{drive_id}/root:/{enc}:/content"
    # If-Match en el PUT: cierra la carrera entre el check local y SharePoint.
    await graph.put_bytes(
        endpoint,
        payload,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        if_match=if_match.strip(),
    )

    review = await load_ui_review_for_process(
        reader, process_key=key, banks=banks
    )
    return UiReviewPatchResponse(
        process_key=key,
        etag=review.etag,
        updated_row_keys=updated,
        review=review,
    )
