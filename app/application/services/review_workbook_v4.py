"""Builder del workbook de revisión schema v4 (sin distribución manual ni Aplicación sugerida)."""
from __future__ import annotations

import io
import re
from datetime import date
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.application.services.extract_snapshot_parser import (
    EVIDENCE_HEADERS_LABEL,
    EVIDENCE_META_FIELDS,
    EVIDENCE_ROW_PREFIX,
    serialize_evidence_meta_value,
)
from app.application.services.review_error_guide import error_record_to_sheet_row
from app.application.services.review_schema import (
    REVIEW_SCHEMA_VERSION,
    AplicacionPagosCols,
    ErroresCols,
    InternalPathCols,
    MetaCols,
    ReviewSheets,
    TipoAplicacionConfirmado,
    ValidarPago,
    dias_respecto_vencimiento,
)


_FILL_NAVY = PatternFill(fill_type="solid", fgColor="002060")
_HEADER_FILL = _FILL_NAVY
_HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_FONT_TITLE = Font(name="Calibri", bold=True, size=16, color="FFFFFF")
_FONT_INSTRUCTION = Font(name="Calibri", size=12, color="1A2F36")
_FONT_BODY = Font(name="Calibri", size=11)
_FONT_HLINK = Font(name="Calibri", color="0563C1", size=11, underline="single")
_FILL_INSTRUCTION = PatternFill(fill_type="solid", fgColor="E8F2FA")
_FILL_HLINK = PatternFill(fill_type="solid", fgColor="E8F4FC")
_FILL_EDITABLE = PatternFill(fill_type="solid", fgColor="FFF3CD")
_FILL_ZEBRA_A = PatternFill(fill_type="solid", fgColor="FCFCFD")
_FILL_ZEBRA_B = PatternFill(fill_type="solid", fgColor="F6F8FA")
_AMBIGUOUS_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
_FILL_DIAS_LATE = PatternFill(fill_type="solid", fgColor="F8D7DA")
_FILL_DIAS_ON_TIME = PatternFill(fill_type="solid", fgColor="D6EAF8")
_FILL_DIAS_EARLY = PatternFill(fill_type="solid", fgColor="D5F5E3")
_FONT_DIAS_LATE = Font(name="Calibri", size=11, color="9B1B30", bold=True)
_FONT_DIAS_ON_TIME = Font(name="Calibri", size=11, color="1A5276", bold=True)
_FONT_DIAS_EARLY = Font(name="Calibri", size=11, color="1E7A46", bold=True)
_THIN = Side(style="thin", color="C8C8C8")
_MEDIUM_CLIENT_EDGE = Side(style="medium", color="002060")
_BORDER_LIGHT = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_ALIGN_TITLE = Alignment(vertical="center", horizontal="center", wrap_text=True)
_ALIGN_WRAP = Alignment(vertical="center", horizontal="left", wrap_text=True)
_ALIGN_CENTER = Alignment(vertical="center", horizontal="center", wrap_text=True)
_ALIGN_VCENTER = Alignment(vertical="center", horizontal="left")
_LOCKED = Protection(locked=True)
_UNLOCKED = Protection(locked=False)
_FMT_MONEY = "#,##0.00"
_FMT_DATE = "yyyy-mm-dd"
_TAB_APLICACION = "00B050"
_TAB_ERRORES = "FF6969"

_TIPO_COL_WIDTH = max(28, max(len(opt) for opt in TipoAplicacionConfirmado.OPTIONS_ORDERED) + 2)

APLICACION_TITLE = "APLICACIÓN DE PAGOS"
APLICACION_HELP = (
    "Complete únicamente las celdas editables (fondo amarillo): Validar Pago y Tipo de aplicación. "
    "Marque SI solo en las filas que desea validar; el resto puede quedar en NO (o vacío). "
    "En las filas SI elija Tipo de aplicación. Observación es opcional. "
    "No ingrese montos: el banco y el asiento definen los valores. "
    "Al terminar, vuelva a la aplicación y pulse Finalizar."
)
ERRORES_TITLE = "REGISTRO DE ERRORES"
ERRORES_HELP = (
    "Revise estos casos manualmente. Use la descripción, la acción recomendada y los links "
    "para corregir documentos o carpetas antes de volver a generar."
)

REVIEW_HEADER_ROW = 3
REVIEW_FIRST_DATA_ROW = 4


def _http_url_only(raw: Any) -> str:
    s = str(raw or "").strip()
    return s if s.lower().startswith("http") else ""


def _link_label_suffix(credito: Any, cliente: Any) -> str:
    from app.application.services.review_schema import normalize_credito_digits

    cred = str(credito or "").strip()
    cli = str(cliente or "").strip()
    digits = normalize_credito_digits(cred) or re.sub(r"\D+", "", cred)
    if digits and (cred.upper().startswith("CREDITO") or re.fullmatch(r"\d+", cred)):
        return f"crédito {digits}"
    if cred:
        return cred
    return cli


def _format_distrib_link_visible_text(link_kind: str, credito: Any, cliente: Any) -> str:
    suffix = _link_label_suffix(credito, cliente)
    if link_kind == "extracto":
        return f"Ver extracto {suffix}".strip() if suffix else "Ver extracto"
    if link_kind == "tabla":
        return f"Ver tabla {suffix}".strip() if suffix else "Ver tabla"
    if link_kind == "carpeta":
        return f"Ver carpeta {suffix}".strip() if suffix else "Ver carpeta"
    return ""


def _format_errores_link_visible_text(link_kind: str, credito: Any, cliente: Any) -> str:
    return _format_distrib_link_visible_text(link_kind, credito, cliente)


def _style_cell_as_excel_hyperlink(cell: Any, target: str, display: str) -> None:
    cell.value = display
    cell.hyperlink = target
    cell.font = _FONT_HLINK
    cell.fill = _FILL_HLINK


def _distrib_row_border(*, client_top: bool = False, client_bottom: bool = False) -> Border:
    top = _MEDIUM_CLIENT_EDGE if client_top else _THIN
    bottom = _MEDIUM_CLIENT_EDGE if client_bottom else _THIN
    return Border(left=_THIN, right=_THIN, top=top, bottom=bottom)


def _sheet_client_block_edges(
    ws: Any,
    first_data_row: int,
    last_row: int,
    col_cliente: int,
    col_row_key: int,
) -> dict[int, tuple[bool, bool]]:
    indexed: list[tuple[int, str]] = []
    for r in range(first_data_row, last_row + 1):
        cliente = str(ws.cell(row=r, column=col_cliente).value or "").strip()
        row_key = ws.cell(row=r, column=col_row_key).value
        if not cliente and (row_key is None or str(row_key).strip() == ""):
            continue
        indexed.append((r, cliente))
    edges: dict[int, tuple[bool, bool]] = {}
    for i, (r, cliente) in enumerate(indexed):
        prev_cliente = indexed[i - 1][1] if i > 0 else None
        next_cliente = indexed[i + 1][1] if i + 1 < len(indexed) else None
        client_top = i > 0 and bool(cliente) and cliente != prev_cliente
        client_bottom = bool(cliente) and (next_cliente is None or cliente != next_cliente)
        edges[r] = (client_top, client_bottom)
    return edges


def _apply_banner(ws: Any, *, title: str, help_text: str, ncols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    t = ws.cell(1, 1, title)
    t.font = _FONT_TITLE
    t.fill = _FILL_NAVY
    t.alignment = _ALIGN_TITLE
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    s = ws.cell(2, 1, help_text)
    s.fill = _FILL_INSTRUCTION
    s.font = _FONT_INSTRUCTION
    s.alignment = _ALIGN_WRAP
    ws.row_dimensions[2].height = 48


def _style_header_row(ws: Any, row_idx: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row_idx, col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _ALIGN_CENTER
        cell.border = _BORDER_LIGHT
    ws.row_dimensions[row_idx].height = 28
    ws.auto_filter.ref = f"A{row_idx}:{get_column_letter(ncols)}{row_idx}"
    ws.freeze_panes = f"A{row_idx + 1}"
    ws.sheet_view.showGridLines = False


def _apply_aplicacion_hyperlinks(ws: Any, first_data: int, last_data: int) -> None:
    if last_data < first_data:
        return
    col_ext = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_EXTRACTO) + 1
    col_tab = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_TABLA) + 1
    col_fold = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.LINK_CARPETA_CREDITO) + 1
    col_credito = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.CREDITO) + 1
    col_cliente = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.CLIENTE) + 1
    mapping = (
        (col_ext, "extracto"),
        (col_tab, "tabla"),
        (col_fold, "carpeta"),
    )
    for r in range(first_data, last_data + 1):
        credito = ws.cell(r, col_credito).value
        cliente = ws.cell(r, col_cliente).value
        for col, kind in mapping:
            cell = ws.cell(r, col)
            url = _http_url_only(cell.value)
            if url:
                _style_cell_as_excel_hyperlink(
                    cell,
                    url,
                    _format_distrib_link_visible_text(kind, credito, cliente),
                )
            else:
                cell.value = ""
                cell.hyperlink = None


def _apply_errores_hyperlinks(
    ws: Any,
    first_data: int,
    error_records: list[dict[str, Any]],
) -> None:
    if not error_records:
        return
    col_ext = ErroresCols.HEADERS.index(ErroresCols.LINK_EXTRACTO) + 1
    col_fold = ErroresCols.HEADERS.index(ErroresCols.LINK_CARPETA_CREDITO) + 1
    for i, rec in enumerate(error_records):
        r = first_data + i
        ext_url = _http_url_only(rec.get("link_extracto_url") or rec.get("link_extracto"))
        fold_url = _http_url_only(
            rec.get("link_carpeta_credito_url") or rec.get("link_carpeta")
        )
        credito = rec.get("credito")
        cliente = rec.get("cliente")
        c_ext = ws.cell(r, col_ext)
        if ext_url:
            _style_cell_as_excel_hyperlink(
                c_ext,
                ext_url,
                _format_errores_link_visible_text("extracto", credito, cliente),
            )
        else:
            c_ext.value = ""
            c_ext.hyperlink = None
        c_fold = ws.cell(r, col_fold)
        if fold_url:
            _style_cell_as_excel_hyperlink(
                c_fold,
                fold_url,
                _format_errores_link_visible_text("carpeta", credito, cliente),
            )
        else:
            c_fold.value = ""
            c_fold.hyperlink = None


def _coerce_dias_int(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        try:
            return int(float(str(value).replace(",", ".")))
        except (TypeError, ValueError):
            return None


def _apply_aplicacion_body_style(
    ws: Any,
    *,
    first_data: int,
    last_data: int,
    ncols: int,
    aplicacion_rows: list[dict[str, Any]],
) -> None:
    money_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1
        for c in (
            AplicacionPagosCols.MONTO_BANCO,
            AplicacionPagosCols.VALOR_OBLIGACION_ACTUAL,
            AplicacionPagosCols.SALDO_VENCIDO,
        )
    }
    date_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1
        for c in (AplicacionPagosCols.FECHA_BANCO, AplicacionPagosCols.FECHA_LIMITE)
    }
    link_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1
        for c in (
            AplicacionPagosCols.LINK_EXTRACTO,
            AplicacionPagosCols.LINK_TABLA,
            AplicacionPagosCols.LINK_CARPETA_CREDITO,
        )
    }
    wrap_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1
        for c in (AplicacionPagosCols.OBSERVACION, AplicacionPagosCols.TIPO_APLICACION)
    }
    primary_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1 for c in AplicacionPagosCols.PRIMARY_EDITABLE
    }
    dias_col = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.DIAS_RESPECTO_VENCIMIENTO) + 1
    col_id = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.ID_PAGO) + 1
    col_cliente = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.CLIENTE) + 1
    block_edges = (
        _sheet_client_block_edges(ws, first_data, last_data, col_cliente, col_id)
        if last_data >= first_data
        else {}
    )
    seen: dict[str, int] = {}
    zebra_bit = 0
    for r in range(first_data, last_data + 1):
        pid = str(ws.cell(r, col_id).value or "").strip()
        if pid not in seen:
            seen[pid] = zebra_bit
            zebra_bit ^= 1
        zebra = _FILL_ZEBRA_A if seen[pid] == 0 else _FILL_ZEBRA_B
        row_off = r - first_data
        ambiguous = False
        if 0 <= row_off < len(aplicacion_rows):
            ambiguous = str(aplicacion_rows[row_off].get("_right_panel_role") or "") == "AMBIGUO"
        client_top, client_bottom = block_edges.get(r, (False, False))
        row_border = _distrib_row_border(client_top=client_top, client_bottom=client_bottom)
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.border = row_border
            cell.alignment = _ALIGN_WRAP if c in wrap_idx else _ALIGN_VCENTER
            if ambiguous:
                cell.fill = _AMBIGUOUS_FILL
            elif c in primary_idx:
                cell.fill = _FILL_EDITABLE
            elif c in link_idx and getattr(cell, "hyperlink", None) is not None:
                cell.fill = _FILL_HLINK
            else:
                cell.fill = zebra
            if c in money_idx:
                cell.number_format = _FMT_MONEY
            if c in date_idx:
                cell.number_format = _FMT_DATE
            if c in link_idx and getattr(cell, "hyperlink", None) is not None:
                cell.font = _FONT_HLINK
            elif c == dias_col:
                dias = _coerce_dias_int(cell.value)
                if dias is None:
                    cell.font = _FONT_BODY
                elif dias > 0:
                    cell.fill = _FILL_DIAS_LATE
                    cell.font = _FONT_DIAS_LATE
                    cell.alignment = _ALIGN_CENTER
                elif dias < 0:
                    cell.fill = _FILL_DIAS_EARLY
                    cell.font = _FONT_DIAS_EARLY
                    cell.alignment = _ALIGN_CENTER
                else:
                    cell.fill = _FILL_DIAS_ON_TIME
                    cell.font = _FONT_DIAS_ON_TIME
                    cell.alignment = _ALIGN_CENTER
            elif c not in link_idx:
                cell.font = _FONT_BODY
        ws.row_dimensions[r].height = 22

    widths = {
        AplicacionPagosCols.ID_PAGO: 38,
        AplicacionPagosCols.CLIENTE: 22,
        AplicacionPagosCols.CREDITO: 16,
        AplicacionPagosCols.MONTO_BANCO: 16,
        AplicacionPagosCols.FECHA_BANCO: 14,
        AplicacionPagosCols.FECHA_LIMITE: 14,
        AplicacionPagosCols.DIAS_RESPECTO_VENCIMIENTO: 14,
        AplicacionPagosCols.VALOR_OBLIGACION_ACTUAL: 18,
        AplicacionPagosCols.SALDO_VENCIDO: 16,
        AplicacionPagosCols.VALIDAR_PAGO: 16,
        AplicacionPagosCols.TIPO_APLICACION: _TIPO_COL_WIDTH,
        AplicacionPagosCols.LINK_EXTRACTO: 28,
        AplicacionPagosCols.LINK_TABLA: 28,
        AplicacionPagosCols.LINK_CARPETA_CREDITO: 32,
        AplicacionPagosCols.OBSERVACION: 40,
    }
    for idx, name in enumerate(AplicacionPagosCols.HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 14)
    freeze_col = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.CREDITO) + 2
    ws.freeze_panes = f"{get_column_letter(freeze_col)}{first_data}"
    if last_data >= first_data:
        ws.auto_filter.ref = (
            f"A{REVIEW_HEADER_ROW}:{get_column_letter(ncols)}{last_data}"
        )


def _apply_errores_body_style(ws: Any, *, first_data: int, last_data: int, ncols: int) -> None:
    wrap_names = {ErroresCols.DESCRIPCION, ErroresCols.QUE_DEBE_HACER}
    wrap_idx = {ErroresCols.HEADERS.index(n) + 1 for n in wrap_names}
    col_cliente = ErroresCols.HEADERS.index(ErroresCols.CLIENTE) + 1
    col_id = ErroresCols.HEADERS.index(ErroresCols.ID_PAGO) + 1
    block_edges = (
        _sheet_client_block_edges(ws, first_data, last_data, col_cliente, col_id)
        if last_data >= first_data
        else {}
    )
    for r in range(first_data, last_data + 1):
        zebra = _FILL_ZEBRA_A if (r - first_data) % 2 == 0 else _FILL_ZEBRA_B
        client_top, client_bottom = block_edges.get(r, (False, False))
        row_border = _distrib_row_border(client_top=client_top, client_bottom=client_bottom)
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.border = row_border
            if getattr(cell, "hyperlink", None) is not None:
                cell.font = _FONT_HLINK
                cell.fill = _FILL_HLINK
            else:
                cell.fill = zebra
                cell.font = _FONT_BODY
            cell.alignment = _ALIGN_WRAP if c in wrap_idx else _ALIGN_VCENTER
        ws.row_dimensions[r].height = 36
    widths = {
        ErroresCols.ID_PAGO: 38,
        ErroresCols.CLIENTE: 22,
        ErroresCols.CREDITO: 16,
        ErroresCols.TIPO_CASO: 18,
        ErroresCols.DESCRIPCION: 42,
        ErroresCols.QUE_DEBE_HACER: 42,
        ErroresCols.REQUIERE_SOPORTE: 16,
        ErroresCols.LINK_EXTRACTO: 32,
        ErroresCols.LINK_CARPETA_CREDITO: 36,
        ErroresCols.CODIGO_TECNICO: 28,
    }
    for idx, name in enumerate(ErroresCols.HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 14)
    if last_data >= first_data:
        ws.auto_filter.ref = (
            f"A{REVIEW_HEADER_ROW}:{get_column_letter(ncols)}{last_data}"
        )


def build_aplicacion_pagos_row(
    payment: dict[str, Any],
    candidate: dict[str, Any],
) -> dict[str, Any]:
    """Fila neutra v4: Validar Pago = NO; sin autoselección SI ni montos editables."""
    due = candidate.get("fecha_limite")
    fecha_banco = payment["fecha_banco"]
    dias = dias_respecto_vencimiento(fecha_banco, due)

    valor_oblig = candidate.get("valor_obligacion_actual")
    if valor_oblig is None:
        valor_oblig = candidate.get("valor_extracto")

    saldo_vis = candidate.get("saldo_vencido_visible")
    if "saldo_vencido_visible" not in candidate:
        saldo_vis = None

    row = {
        AplicacionPagosCols.ID_PAGO: payment["id_pago"],
        AplicacionPagosCols.CLIENTE: payment["cliente"],
        AplicacionPagosCols.CREDITO: candidate.get("credito"),
        AplicacionPagosCols.MONTO_BANCO: payment["monto_banco"],
        AplicacionPagosCols.FECHA_BANCO: fecha_banco.isoformat()
        if hasattr(fecha_banco, "isoformat")
        else fecha_banco,
        AplicacionPagosCols.FECHA_LIMITE: due.isoformat() if hasattr(due, "isoformat") else (due or ""),
        AplicacionPagosCols.DIAS_RESPECTO_VENCIMIENTO: dias if dias is not None else "",
        AplicacionPagosCols.VALOR_OBLIGACION_ACTUAL: valor_oblig if valor_oblig is not None else "",
        AplicacionPagosCols.SALDO_VENCIDO: saldo_vis if saldo_vis is not None else "",
        AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.NO,
        AplicacionPagosCols.TIPO_APLICACION: "",
        AplicacionPagosCols.LINK_EXTRACTO: candidate.get("link_extracto", ""),
        AplicacionPagosCols.LINK_TABLA: candidate.get("link_tabla", ""),
        AplicacionPagosCols.LINK_CARPETA_CREDITO: candidate.get("link_carpeta_credito", ""),
        AplicacionPagosCols.OBSERVACION: candidate.get("observacion_extra") or "",
    }
    row["_evidence"] = candidate.get("extract_evidence") or {}
    row["_right_panel_role"] = candidate.get("right_panel_role") or ""
    row["_parser_status"] = candidate.get("parser_status") or ""
    row["_ruta_extracto"] = candidate.get("ruta_extracto_pdf") or ""
    row["_ruta_unidad_credito"] = candidate.get("ruta_unidad_credito") or ""
    row["_ruta_tabla_amortizacion"] = candidate.get("ruta_tabla_amortizacion") or ""
    row["_credito_normalizado"] = candidate.get("credito_normalizado") or ""
    return row


def build_review_workbook_v4_bytes(
    *,
    process_id: str,
    process_date: date,
    bank_code: str,
    aplicacion_rows: list[dict[str, Any]],
    error_records: list[dict[str, Any]],
) -> bytes:
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    path_headers = [
        InternalPathCols.RUTA_EXTRACTO,
        InternalPathCols.RUTA_UNIDAD_CREDITO,
        InternalPathCols.RUTA_TABLA_AMORTIZACION,
        InternalPathCols.CREDITO_NORMALIZADO,
    ]
    visible_n = len(AplicacionPagosCols.HEADERS)
    _apply_banner(ws, title=APLICACION_TITLE, help_text=APLICACION_HELP, ncols=visible_n)
    ws.append(list(AplicacionPagosCols.HEADERS) + path_headers)
    _style_header_row(ws, REVIEW_HEADER_ROW, visible_n)

    seen_monto: set[str] = set()
    for row in aplicacion_rows:
        pid = str(row.get(AplicacionPagosCols.ID_PAGO) or "").strip()
        values = []
        for h in AplicacionPagosCols.HEADERS:
            val = row.get(h, "")
            if h == AplicacionPagosCols.MONTO_BANCO and pid:
                if pid in seen_monto:
                    val = None
                else:
                    seen_monto.add(pid)
            values.append(val)
        for ph in path_headers:
            values.append(row.get(ph, "") or "")
        ws.append(values)

    for offset in range(len(path_headers)):
        letter = get_column_letter(len(AplicacionPagosCols.HEADERS) + 1 + offset)
        ws.column_dimensions[letter].hidden = True

    first_data = REVIEW_FIRST_DATA_ROW
    last_data = first_data + len(aplicacion_rows) - 1 if aplicacion_rows else first_data - 1
    col_vp = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.VALIDAR_PAGO) + 1
    col_saldo_venc = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.SALDO_VENCIDO) + 1

    for r in range(first_data, last_data + 1):
        role = str(aplicacion_rows[r - first_data].get("_right_panel_role") or "")
        if role == "AMBIGUO":
            for c in range(1, len(AplicacionPagosCols.HEADERS) + 1):
                ws.cell(r, c).fill = _AMBIGUOUS_FILL
            cell_sv = ws.cell(r, col_saldo_venc)
            if cell_sv.comment is None:
                from openpyxl.comments import Comment

                cell_sv.comment = Comment(
                    "Panel derecho AMBIGUO: Saldo vencido vacío a propósito. "
                    "Revisar extracto; no se asume 0.",
                    "HBI",
                )

    ws_lists = wb.create_sheet(ReviewSheets.LISTAS)
    ws_lists.append(["ValidarPago"])
    for i, opt in enumerate(ValidarPago.OPTIONS_ORDERED, start=2):
        ws_lists.cell(i, 1, opt)
    ws_lists.cell(1, 2, "TipoAplicacion")
    for i, opt in enumerate(TipoAplicacionConfirmado.OPTIONS_ORDERED, start=2):
        ws_lists.cell(i, 2, opt)
    ws_lists.sheet_state = "hidden"

    dv_vp = DataValidation(
        type="list",
        formula1=f"={ReviewSheets.LISTAS}!$A$2:$A${1 + len(ValidarPago.OPTIONS_ORDERED)}",
        allow_blank=True,
    )
    ws.add_data_validation(dv_vp)
    dv_vp.add(
        f"{get_column_letter(col_vp)}{first_data}:{get_column_letter(col_vp)}{max(last_data, first_data + 50)}"
    )

    col_tipo = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.TIPO_APLICACION) + 1
    dv_tipo = DataValidation(
        type="list",
        formula1=f"={ReviewSheets.LISTAS}!$B$2:$B${1 + len(TipoAplicacionConfirmado.OPTIONS_ORDERED)}",
        allow_blank=True,
    )
    ws.add_data_validation(dv_tipo)
    dv_tipo.add(
        f"{get_column_letter(col_tipo)}{first_data}:{get_column_letter(col_tipo)}{max(last_data, first_data + 50)}"
    )

    editable_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1 for c in AplicacionPagosCols.SECRETARY_EDITABLE
    }
    for r in range(first_data, last_data + 1):
        for c in range(1, len(AplicacionPagosCols.HEADERS) + 1):
            cell = ws.cell(r, c)
            cell.protection = _UNLOCKED if c in editable_idx else _LOCKED
    ws.protection.sheet = True
    ws.sheet_properties.tabColor = Color(rgb=_TAB_APLICACION)
    _apply_aplicacion_hyperlinks(ws, first_data, last_data)
    _apply_aplicacion_body_style(
        ws,
        first_data=first_data,
        last_data=last_data,
        ncols=visible_n,
        aplicacion_rows=aplicacion_rows,
    )

    ws_err = wb.create_sheet(ReviewSheets.ERRORES)
    err_n = len(ErroresCols.HEADERS)
    _apply_banner(ws_err, title=ERRORES_TITLE, help_text=ERRORES_HELP, ncols=err_n)
    ws_err.append(list(ErroresCols.HEADERS))
    _style_header_row(ws_err, REVIEW_HEADER_ROW, err_n)
    for rec in error_records:
        ws_err.append(error_record_to_sheet_row(rec))
    if error_records:
        err_first = REVIEW_FIRST_DATA_ROW
        err_last = err_first + len(error_records) - 1
        _apply_errores_hyperlinks(ws_err, err_first, error_records)
        _apply_errores_body_style(
            ws_err, first_data=err_first, last_data=err_last, ncols=err_n
        )
        ws_err.sheet_properties.tabColor = Color(rgb=_TAB_ERRORES)
    else:
        ws_err.sheet_state = "hidden"

    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append([MetaCols.CAMPO, MetaCols.VALOR])
    ws_meta.append([MetaCols.ROW_REVIEW_SCHEMA_VERSION, REVIEW_SCHEMA_VERSION])
    ws_meta.append([MetaCols.ROW_PROCESS_ID, process_id])
    ws_meta.append([MetaCols.ROW_PROCESS_DATE, process_date.isoformat()])
    ws_meta.append([MetaCols.ROW_BANK_CODE, bank_code])
    ws_meta.append([EVIDENCE_HEADERS_LABEL, "|".join(EVIDENCE_META_FIELDS)])
    for idx, row in enumerate(aplicacion_rows, start=first_data):
        ev = row.get("_evidence") or {}
        ws_meta.append(
            [
                f"{EVIDENCE_ROW_PREFIX}{idx}",
                serialize_evidence_meta_value(
                    row_idx=idx,
                    evidence=ev,
                    right_panel_role=str(row.get("_right_panel_role") or ""),
                    parser_status=str(row.get("_parser_status") or ""),
                ),
            ]
        )
    ws_meta.sheet_state = "hidden"

    forbidden = {
        "Control",
        "Resumen",
        "Casos_Pago",
        "Distribucion_Pagos",
        "Distribucion_Abonos",
        "Distribucion",
    }
    for name in list(wb.sheetnames):
        if name in forbidden:
            del wb[name]

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
