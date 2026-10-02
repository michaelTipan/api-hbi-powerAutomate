"""
Evaluación acotada de fórmulas simples del cronograma (día / mes / año).

Apply guarda la tabla con openpyxl y Excel pierde los valores calculados de las
fórmulas hasta que alguien la vuelva a abrir. Para no depender de ese cache, la
lectura del cronograma puede evaluar fórmulas sencillas: referencias (``=L6``,
``=+L6``, ``=+TasaImplicita!A12`` a otra hoja del mismo libro), aritmética,
comparaciones, ``IF/SI``, ``AND/Y``, ``OR/O``, ``NOT/NO``, ``DATE/FECHA``,
``EDATE/FECHA.MES``, ``DAY/DIA``, ``MONTH/MES``, ``YEAR/AÑO`` y ``SUM/SUMA``.

Cualquier cosa fuera de ese subconjunto (otros libros, nombres definidos, funciones
desconocidas, ciclos, errores) levanta ``FormulaNotEvaluable`` y el llamador
conserva el comportamiento anterior: nunca se inventa un valor.
"""

from __future__ import annotations

import calendar
import re
from collections.abc import Callable
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

_EXCEL_EPOCH = datetime(1899, 12, 30)
_MAX_DEPTH = 12
_MAX_SUM_CELLS = 500


class FormulaNotEvaluable(Exception):
    """La fórmula no se puede evaluar de forma segura."""


_TOKEN_RE = re.compile(
    r"""
    (?P<ws>\s+)
  | (?P<string>"(?:[^"]|"")*")
  | (?P<sheetref>(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)!\$?[A-Za-z]{1,3}\$?\d+(?![A-Za-z0-9_(!:]))
  | (?P<range>\$?[A-Za-z]{1,3}\$?\d+\s*:\s*\$?[A-Za-z]{1,3}\$?\d+)
  | (?P<func>[A-Za-zÁÉÍÓÚÑáéíóúñ][A-Za-z0-9_.ÁÉÍÓÚÑáéíóúñ]*(?=\s*\())
  | (?P<ref>\$?[A-Za-z]{1,3}\$?\d+(?![A-Za-z0-9_(!]))
  | (?P<number>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?|\.\d+(?:[eE][+-]?\d+)?)
  | (?P<bool>(?i:TRUE|FALSE|VERDADERO|FALSO)(?![A-Za-z0-9_(]))
  | (?P<op><>|<=|>=|[-+*/^&=<>(),;%])
    """,
    re.VERBOSE,
)

_CELL_RE = re.compile(r"\$?([A-Za-z]{1,3})\$?(\d+)")

_FUNC_ALIASES = {
    "IF": "IF",
    "SI": "IF",
    "AND": "AND",
    "Y": "AND",
    "OR": "OR",
    "O": "OR",
    "NOT": "NOT",
    "NO": "NOT",
    "DATE": "DATE",
    "FECHA": "DATE",
    "EDATE": "EDATE",
    "FECHA.MES": "EDATE",
    "DAY": "DAY",
    "DIA": "DAY",
    "DÍA": "DAY",
    "MONTH": "MONTH",
    "MES": "MONTH",
    "YEAR": "YEAR",
    "AÑO": "YEAR",
    "ANO": "YEAR",
    "SUM": "SUM",
    "SUMA": "SUM",
}


def _tokenize(formula: str) -> list[tuple[str, str]]:
    text = formula.strip()
    if text.startswith("="):
        text = text[1:]
    tokens: list[tuple[str, str]] = []
    pos = 0
    while pos < len(text):
        m = _TOKEN_RE.match(text, pos)
        if not m:
            raise FormulaNotEvaluable(f"token no soportado en {formula!r}")
        kind = m.lastgroup or ""
        if kind != "ws":
            tokens.append((kind, m.group(kind)))
        pos = m.end()
    return tokens


def _to_serial(value: datetime | date) -> float:
    if isinstance(value, datetime):
        delta = value - _EXCEL_EPOCH
    else:
        delta = datetime(value.year, value.month, value.day) - _EXCEL_EPOCH
    return delta.days + delta.seconds / 86400.0


def _from_serial(serial: float) -> datetime:
    if serial < 1:
        raise FormulaNotEvaluable("serial de fecha fuera de rango")
    return _EXCEL_EPOCH + timedelta(days=float(serial))


def _num(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, (datetime, date)):
        return _to_serial(value)
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            raise FormulaNotEvaluable("texto vacío en operación numérica")
        try:
            return float(text.replace(",", "."))
        except ValueError as exc:
            raise FormulaNotEvaluable("texto no numérico en operación") from exc
    raise FormulaNotEvaluable(f"tipo no soportado: {type(value).__name__}")


def _truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, str):
        upper = value.strip().upper()
        if upper in ("TRUE", "VERDADERO"):
            return True
        if upper in ("FALSE", "FALSO"):
            return False
        raise FormulaNotEvaluable("texto usado como condición")
    return _num(value) != 0


def _compare(op: str, left: Any, right: Any) -> bool:
    def norm(v: Any) -> tuple[int, Any]:
        if v is None:
            return (-1, None)
        if isinstance(v, bool):
            return (2, v)
        if isinstance(v, (int, float)):
            return (0, float(v))
        if isinstance(v, (datetime, date)):
            return (0, _to_serial(v))
        return (1, str(v).casefold())

    lk, lv = norm(left)
    rk, rv = norm(right)
    # Celda vacía: vale "" frente a texto y 0 frente a número (semántica Excel).
    if lk == -1 and rk == -1:
        lk, lv, rk, rv = 0, 0.0, 0, 0.0
    elif lk == -1:
        lk, lv = (rk, "" if rk == 1 else (False if rk == 2 else 0.0))
    elif rk == -1:
        rk, rv = (lk, "" if lk == 1 else (False if lk == 2 else 0.0))
    if op == "=":
        return lk == rk and lv == rv
    if op == "<>":
        return not (lk == rk and lv == rv)
    if lk != rk:
        # Excel ordena número < texto < lógico.
        lv, rv = lk, rk
    if op == "<":
        return lv < rv
    if op == ">":
        return lv > rv
    if op == "<=":
        return lv <= rv
    if op == ">=":
        return lv >= rv
    raise FormulaNotEvaluable(f"operador desconocido {op}")


def _cell_coords(ref: str) -> tuple[int, int]:
    m = _CELL_RE.fullmatch(ref.strip())
    if not m:
        raise FormulaNotEvaluable(f"referencia inválida {ref!r}")
    return int(m.group(2)), column_index_from_string(m.group(1).upper())


class _Parser:
    def __init__(self, tokens: list[tuple[str, str]], evaluator: "SimpleFormulaEvaluator", depth: int):
        self.tokens = tokens
        self.pos = 0
        self.ev = evaluator
        self.depth = depth

    def peek(self) -> tuple[str, str] | None:
        return self.tokens[self.pos] if self.pos < len(self.tokens) else None

    def take(self) -> tuple[str, str]:
        tok = self.peek()
        if tok is None:
            raise FormulaNotEvaluable("fórmula incompleta")
        self.pos += 1
        return tok

    def expect_op(self, value: str) -> None:
        kind, text = self.take()
        if kind != "op" or text != value:
            raise FormulaNotEvaluable(f"se esperaba {value!r}")

    def is_op(self, *values: str) -> bool:
        tok = self.peek()
        return tok is not None and tok[0] == "op" and tok[1] in values

    def parse(self) -> Any:
        value = self.comparison()
        if self.peek() is not None:
            raise FormulaNotEvaluable("tokens sobrantes")
        return value

    def comparison(self) -> Any:
        left = self.concat()
        while self.is_op("=", "<>", "<", ">", "<=", ">="):
            op = self.take()[1]
            right = self.concat()
            left = _compare(op, left, right)
        return left

    def concat(self) -> Any:
        left = self.additive()
        while self.is_op("&"):
            self.take()
            right = self.additive()
            left = _text(left) + _text(right)
        return left

    def additive(self) -> Any:
        left = self.term()
        while self.is_op("+", "-"):
            op = self.take()[1]
            right = self.term()
            left = _num(left) + _num(right) if op == "+" else _num(left) - _num(right)
        return left

    def term(self) -> Any:
        left = self.power()
        while self.is_op("*", "/"):
            op = self.take()[1]
            right = self.power()
            if op == "*":
                left = _num(left) * _num(right)
            else:
                divisor = _num(right)
                if divisor == 0:
                    raise FormulaNotEvaluable("división por cero")
                left = _num(left) / divisor
        return left

    def power(self) -> Any:
        base = self.unary()
        if self.is_op("^"):
            self.take()
            exponent = self.unary()
            return _num(base) ** _num(exponent)
        return base

    def unary(self) -> Any:
        if self.is_op("+"):
            self.take()
            return self.unary()
        if self.is_op("-"):
            self.take()
            return -_num(self.unary())
        value = self.primary()
        if self.is_op("%"):
            self.take()
            return _num(value) / 100.0
        return value

    def primary(self) -> Any:
        kind, text = self.take()
        if kind == "number":
            return float(text)
        if kind == "string":
            return text[1:-1].replace('""', '"')
        if kind == "bool":
            return text.upper() in ("TRUE", "VERDADERO")
        if kind == "ref":
            row, col = _cell_coords(text)
            return self.ev._cell_value(row, col, self.depth + 1)
        if kind == "sheetref":
            sheet, ref = text.rsplit("!", 1)
            if sheet.startswith("'"):
                sheet = sheet[1:-1].replace("''", "'")
            other = self.ev._sheet_evaluator(sheet)
            row, col = _cell_coords(ref)
            return other._cell_value(row, col, self.depth + 1)
        if kind == "func":
            return self.function(text)
        if kind == "op" and text == "(":
            value = self.comparison()
            self.expect_op(")")
            return value
        raise FormulaNotEvaluable(f"token inesperado {text!r}")

    def args(self) -> list[Any]:
        """Argumentos sin evaluar (posiciones de tokens) para IF perezoso."""
        self.expect_op("(")
        spans: list[tuple[int, int]] = []
        if self.is_op(")"):
            self.take()
            return spans
        start = self.pos
        depth = 0
        while True:
            kind, text = self.take()
            if kind == "op" and text == "(":
                depth += 1
            elif kind == "op" and text == ")":
                if depth == 0:
                    spans.append((start, self.pos - 1))
                    return spans
                depth -= 1
            elif kind == "op" and text in (",", ";") and depth == 0:
                spans.append((start, self.pos - 1))
                start = self.pos

    def eval_span(self, span: tuple[int, int]) -> Any:
        start, end = span
        if start == end:
            return None
        sub = _Parser(self.tokens[start:end], self.ev, self.depth)
        return sub.parse()

    def range_values(self, span: tuple[int, int]) -> list[Any]:
        start, end = span
        if end - start == 1 and self.tokens[start][0] == "range":
            a, b = (p.strip() for p in self.tokens[start][1].split(":"))
            r1, c1 = _cell_coords(a)
            r2, c2 = _cell_coords(b)
            r_lo, r_hi = sorted((r1, r2))
            c_lo, c_hi = sorted((c1, c2))
            if (r_hi - r_lo + 1) * (c_hi - c_lo + 1) > _MAX_SUM_CELLS:
                raise FormulaNotEvaluable("rango demasiado grande")
            return [
                self.ev._cell_value(r, c, self.depth + 1)
                for r in range(r_lo, r_hi + 1)
                for c in range(c_lo, c_hi + 1)
            ]
        return [self.eval_span(span)]

    def function(self, name: str) -> Any:
        canonical = _FUNC_ALIASES.get(name.upper())
        if canonical is None:
            raise FormulaNotEvaluable(f"función no soportada {name}")
        spans = self.args()
        if canonical == "IF":
            if len(spans) not in (2, 3):
                raise FormulaNotEvaluable("IF con argumentos inválidos")
            if _truthy(self.eval_span(spans[0])):
                return self.eval_span(spans[1])
            return self.eval_span(spans[2]) if len(spans) == 3 else False
        if canonical in ("AND", "OR"):
            if not spans:
                raise FormulaNotEvaluable(f"{canonical} sin argumentos")
            values = [_truthy(self.eval_span(s)) for s in spans]
            return all(values) if canonical == "AND" else any(values)
        if canonical == "NOT":
            if len(spans) != 1:
                raise FormulaNotEvaluable("NOT con argumentos inválidos")
            return not _truthy(self.eval_span(spans[0]))
        if canonical == "SUM":
            total = 0.0
            for span in spans:
                for v in self.range_values(span):
                    if v is None or isinstance(v, str):
                        continue
                    total += _num(v)
            return total
        values = [self.eval_span(s) for s in spans]
        if canonical == "DATE":
            if len(values) != 3:
                raise FormulaNotEvaluable("DATE con argumentos inválidos")
            y, m, d = (int(_num(v)) for v in values)
            if y < 1900:
                y += 1900
            # Excel normaliza meses/días desbordados (mes 13 → enero siguiente).
            y += (m - 1) // 12
            m = (m - 1) % 12 + 1
            base = date(y, m, 1)
            return _to_serial(base) + (d - 1)
        if canonical == "EDATE":
            if len(values) != 2:
                raise FormulaNotEvaluable("EDATE con argumentos inválidos")
            start = _from_serial(_num(values[0])).date()
            months = int(_num(values[1]))
            total = start.month - 1 + months
            y = start.year + total // 12
            m = total % 12 + 1
            d = min(start.day, calendar.monthrange(y, m)[1])
            return _to_serial(date(y, m, d))
        if canonical in ("DAY", "MONTH", "YEAR"):
            if len(values) != 1:
                raise FormulaNotEvaluable(f"{canonical} con argumentos inválidos")
            v = values[0]
            d = v if isinstance(v, (datetime, date)) else _from_serial(_num(v))
            if canonical == "DAY":
                return float(d.day)
            if canonical == "MONTH":
                return float(d.month)
            return float(d.year)
        raise FormulaNotEvaluable(f"función no soportada {name}")


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.strip().startswith("=")


class SimpleFormulaEvaluator:
    """Evalúa celdas de una hoja con fórmulas (libro cargado con ``data_only=False``).

    ``ws_values`` (opcional) es la misma hoja con valores cacheados: si la celda
    referenciada tiene cache se usa ese valor.
    """

    def __init__(
        self,
        ws_formulas: Worksheet,
        ws_values: Worksheet | None = None,
        *,
        sheet_resolver: Callable[[str], "SimpleFormulaEvaluator | None"] | None = None,
    ):
        self.ws_formulas = ws_formulas
        self.ws_values = ws_values
        self._sheet_resolver = sheet_resolver
        self._memo: dict[tuple[int, int], Any] = {}
        self._failed: set[tuple[int, int]] = set()
        self._in_progress: set[tuple[int, int]] = set()
        self._warmed_upto: dict[int, int] = {}

    def value(self, row: int, col: int) -> Any:
        """Valor de la celda; levanta ``FormulaNotEvaluable`` si no es seguro."""
        self._warm_column(row, col)
        result = self._cell_value(row, col, 0)
        if isinstance(result, bool):
            raise FormulaNotEvaluable("resultado lógico, no es un número de fecha")
        if isinstance(result, float) and result.is_integer():
            return int(result)
        return result

    def _sheet_evaluator(self, title: str) -> "SimpleFormulaEvaluator":
        """Evaluador de otra hoja del mismo libro (solo si el llamador lo habilitó)."""
        if title == self.ws_formulas.title:
            return self
        other = self._sheet_resolver(title) if self._sheet_resolver else None
        if other is None:
            raise FormulaNotEvaluable(f"hoja no disponible {title!r}")
        return other

    def _cell_value(self, row: int, col: int, depth: int) -> Any:
        key = (row, col)
        if key in self._memo:
            return self._memo[key]
        if key in self._failed:
            raise FormulaNotEvaluable("celda no evaluable")
        if self.ws_values is not None:
            cached = self.ws_values.cell(row, col).value
            if cached is not None and not _is_formula(cached):
                return cached
        raw = self.ws_formulas.cell(row, col).value
        if not _is_formula(raw):
            return raw
        if depth > _MAX_DEPTH:
            # Cadena larga (=L6, =L7, ...): se resuelve de arriba hacia abajo.
            self._warm_column(row, col)
            if key in self._memo:
                return self._memo[key]
            raise FormulaNotEvaluable("profundidad máxima de referencias")
        if key in self._in_progress:
            raise FormulaNotEvaluable("referencia circular")
        self._in_progress.add(key)
        try:
            tokens = _tokenize(raw)
            result = _Parser(tokens, self, depth).parse()
            if isinstance(result, (int, float)) and not isinstance(result, bool):
                cell = self.ws_formulas.cell(row, col)
                if is_date_format(cell.number_format or "") and float(result) >= 1:
                    result = _from_serial(float(result))
        except FormulaNotEvaluable:
            self._failed.add(key)
            raise
        except (ValueError, OverflowError, ArithmeticError) as exc:
            self._failed.add(key)
            raise FormulaNotEvaluable(str(exc)) from exc
        finally:
            self._in_progress.discard(key)
        self._memo[key] = result
        return result

    def _warm_column(self, row: int, col: int) -> None:
        """Evalúa de arriba hacia abajo para que las cadenas =L6, =L7… no sean recursivas."""
        start = max(1, self._warmed_upto.get(col, 0) + 1, row - 5000)
        for r in range(start, row):
            if (r, col) in self._memo or (r, col) in self._failed:
                continue
            if (r, col) in self._in_progress:
                return
            if not _is_formula(self.ws_formulas.cell(r, col).value):
                continue
            try:
                self._cell_value(r, col, 0)
            except FormulaNotEvaluable:
                pass
        self._warmed_upto[col] = max(self._warmed_upto.get(col, 0), row - 1)
