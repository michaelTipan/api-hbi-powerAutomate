"""
Dry-run ABONO: preflight documental, cuadre financiero por ID Pago y planificación de fila.

Regla confirmada (Fase 5): ABONO no requiere extracto, fecha límite, cuota contractual ni IBR.
Solo añade una fila en Aplicación de Pagos con valores del asiento contable.

Campo canónico para cuadre: ``PaymentApplicationEvent.valor_pagado_cliente`` (mismo que PAGO).

Cardinalidad PDF → evento:
- ``parse_accounting_text`` produce un único ``PaymentApplicationEvent`` por PDF de asiento.
- ``valor_pagado_cliente`` es el total aplicable del comprobante (no se suman componentes).
- El cuadre del grupo suma ese valor **una vez por ruta normalizada de asiento**; rutas
  repetidas en el mismo grupo no incrementan el total (evita doble conteo si el manifest
  lista el mismo path más de una vez antes de la detección de duplicados).
"""

from __future__ import annotations

import io
import re
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

import openpyxl

from app.application.services.accounting_pdf_parser import (
    AccountingParseError,
    PdfTextNotExtractableError,
    PaymentApplicationEvent,
    extract_text_from_pdf,
    parse_accounting_text,
)
from app.application.services.review_schema import (
    ApplicationPolicy,
    ApplicationSubtype,
    DistribucionAbonosCols,
    ReviewSheets,
    TipoAplicacion,
    TipoAplicacionVisible,
    normalize_credito_digits,
    policy_observability_dict,
    policy_requires_reference_extract,
    resolve_application_policy,
    resolve_manifest_policy,
)
from app.application.services.accounting_pdf_processed_move import (
    list_processed_asiento_paths,
    resolve_asiento_pdf_bytes_with_procesados_fallback,
)
from app.application.services.amortization_apply_safety import compute_asiento_pdf_hash
from app.application.services.applied_abono_events import (
    APPLIED_ABONO_AMOUNT_UNRESOLVABLE,
    APPLIED_ABONO_EVENT_AMBIGUOUS,
    AppliedAbonoEventSnapshot,
    find_applied_abono_event,
    load_applied_abono_event_snapshots,
    resolve_applied_abono_amount,
)
from app.application.services.amortization_workbook import (
    ADOPTADO_EXISTENTE,
    APLICADO,
    REVISION_MANUAL,
    AmortizationSheetNotFoundError,
    build_amortization_idempotency_key,
    build_application_row_search_debug,
    detect_amortization_sheet,
    find_next_available_application_row,
    resolve_planned_application_row,
)
from app.application.use_cases.merge_composite_validado_pdfs import normalize_sharepoint_path

# Misma semántica que ``_AMOUNT_TOLERANCE`` en amortization_workbook (0.02).
ABONO_RECONCILIATION_TOLERANCE = Decimal("0.02")

ABONO_MANIFEST_INVALID = "ABONO_MANIFEST_INVALID"
ABONO_MONTO_BANCO_MISSING = "ABONO_MONTO_BANCO_MISSING"
ABONO_FECHA_BANCO_MISSING = "ABONO_FECHA_BANCO_MISSING"
ABONO_CREDIT_ITEMS_MISSING = "ABONO_CREDIT_ITEMS_MISSING"
ABONO_CREDITO_DUPLICADO = "ABONO_CREDITO_DUPLICADO"
ABONO_CREDITO_NO_DECLARADO = "ABONO_CREDITO_NO_DECLARADO"
ABONO_TABLA_AMORTIZACION_MISSING = "ABONO_TABLA_AMORTIZACION_MISSING"
ABONO_ASIENTO_FALTANTE = "ABONO_ASIENTO_FALTANTE"
ABONO_ASIENTO_DUPLICADO = "ABONO_ASIENTO_DUPLICADO"
ABONO_ASIENTO_TOTAL_NOT_FOUND = "ABONO_ASIENTO_TOTAL_NOT_FOUND"
ABONO_ASIENTOS_NO_CUADRAN = "ABONO_ASIENTOS_NO_CUADRAN"
ABONO_SCHEDULE_RULE_NOT_CONFIGURED = "ABONO_SCHEDULE_RULE_NOT_CONFIGURED"
SCHEDULE_NOT_REQUIRED = "NOT_REQUIRED"
ABONO_APPLICATION_ROW_STRATEGY = "NEXT_AVAILABLE_PAYMENT_ROW"
ABONO_APPLICATION_ROW_STRATEGY_EXISTING = "EXISTING_AUTOMATION_LOG"
ABONO_IBR_SKIPPED_REASON = "NOT_REQUIRED_FOR_ABONO"
APPLICATION_PAYMENT_SECTION_FULL = "APPLICATION_PAYMENT_SECTION_FULL"
ABONO_APPLICATION_STATUS_ALREADY_APPLIED = "ALREADY_APPLIED"
ABONO_MORA_REFERENCE_EXTRACT_MISSING = "ABONO_MORA_REFERENCE_EXTRACT_MISSING"
MORA_REFERENCE_TOLERANCE = ABONO_RECONCILIATION_TOLERANCE

_ABONO_STATUS_MAP = {
    APLICADO: "WOULD_APPLY",
    ADOPTADO_EXISTENTE: "WOULD_ADOPT_EXISTING",
}


@dataclass
class AbonoCreditItem:
    credito: str
    tipo_aplicacion: str
    ruta_tabla_amortizacion: str
    ruta_unidad_credito: str
    ruta_asientos_contables: str
    asiento_pdf_paths: list[str]
    extracto_pdf_paths: list[str]
    policy: ApplicationPolicy | None = None
    mora_reference_amount: Decimal | None = None


def compute_mora_reference_coverage(
    *,
    mora_reference_amount: Decimal | None,
    accounting_mora: Decimal | None,
    accounting_total: Decimal | None,
) -> dict[str, Any]:
    """Cobertura mora extracto vs asiento; MISMATCH es warning, no bloqueo."""
    ref = mora_reference_amount
    acct: Decimal | None = accounting_mora
    if acct is None and accounting_total is not None:
        acct = accounting_total
    if ref is None:
        return {
            "mora_reference_amount": None,
            "accounting_total": float(acct) if acct is not None else None,
            "mora_reference_difference": None,
            "mora_reference_coverage_status": "MISSING_REFERENCE",
        }
    if acct is None:
        return {
            "mora_reference_amount": float(ref),
            "accounting_total": None,
            "mora_reference_difference": None,
            "mora_reference_coverage_status": "MISSING_ACCOUNTING",
        }
    diff = acct - ref
    status = "MATCH" if abs(diff) <= MORA_REFERENCE_TOLERANCE else "MISMATCH"
    return {
        "mora_reference_amount": float(ref),
        "accounting_total": float(acct),
        "mora_reference_difference": float(diff),
        "mora_reference_coverage_status": status,
    }


@dataclass
class AbonoScheduleContextResult:
    status: str
    reference_date: date | None
    due_date_row: int | None
    application_row: int | None
    ibr_date: date | None
    error_code: str | None = None


@dataclass
class AbonoReconciliationResult:
    reconciliation_status: str
    monto_banco: Decimal | None
    total_asientos: Decimal
    diferencia: Decimal
    tolerancia: Decimal
    credit_amounts: dict[str, Decimal]
    accounting_pdf_amounts: dict[str, Decimal]
    blocking_errors: list[dict[str, Any]] = field(default_factory=list)
    already_applied_amount: Decimal = Decimal("0")
    pending_amount: Decimal = Decimal("0")
    total_reconciled_amount: Decimal = Decimal("0")
    already_applied_events_count: int = 0
    pending_events_count: int = 0


@dataclass
class AbonoEventReconcileState:
    credit_item: AbonoCreditItem
    asiento_path: str
    norm_path: str
    event_index: int
    status: str
    amount: Decimal | None = None
    applied_snapshot: AppliedAbonoEventSnapshot | None = None
    event: PaymentApplicationEvent | None = None
    error: dict[str, Any] | None = None
    resolved_pdf_source: str | None = None
    requires_pdf_download: bool = True
    idempotency_key: str = ""
    pdf_fingerprint: dict[str, Any] = field(default_factory=dict)


@dataclass
class AbonoDryRunGroup:
    bank_code: str
    process_key: str
    id_pago: str
    cliente: str
    monto_banco: Decimal | None
    fecha_banco: date | None
    creditos_seleccionados: list[str]
    credit_items: list[AbonoCreditItem]
    reconciliation_status: str = "PENDING"
    schedule_resolution_status: str = "PENDING"
    group_ready_for_apply: bool = False
    requires_business_rule: bool = False
    blocking_errors: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def infer_manifest_tipo_aplicacion(output: dict[str, Any]) -> str:
    """Tipo canónico (PAGO/ABONO) para enrutar dry-run."""
    return resolve_manifest_policy(output).tipo_aplicacion_canonica


def infer_manifest_requiere_extracto(output: dict[str, Any], tipo: str = "") -> bool:
    if "requiere_extracto" in output:
        return bool(output.get("requiere_extracto"))
    return resolve_manifest_policy(output).requiere_extracto


def resolve_abono_schedule_context(
    *,
    group: AbonoDryRunGroup,
    reconciliation: AbonoReconciliationResult,
) -> AbonoScheduleContextResult:
    """
    ABONO no requiere cuota contractual, fecha límite ni IBR.

    ``application_row`` se resuelve por tabla al planificar cada evento (dry-run/apply).
    """
    if reconciliation.reconciliation_status != "PASSED":
        return AbonoScheduleContextResult(
            status="SKIPPED",
            reference_date=None,
            due_date_row=None,
            application_row=None,
            ibr_date=None,
            error_code=None,
        )
    return AbonoScheduleContextResult(
        status=SCHEDULE_NOT_REQUIRED,
        reference_date=None,
        due_date_row=None,
        application_row=None,
        ibr_date=None,
        error_code="",
    )


def _parse_decimal_amount(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip()
    if not text:
        return None
    cleaned = text.replace(".", "").replace(",", ".") if re.search(r",\d{1,2}$", text) else text
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        try:
            return Decimal(re.sub(r"[^\d.\-]", "", text.replace(",", ".")))
        except InvalidOperation:
            return None


def _mora_reference_from_credit_source(ci: dict[str, Any]) -> Decimal | None:
    for key in ("mora_reference_amount", "valor_extracto", "otros_valores", "intereses_mora"):
        if key not in ci:
            continue
        parsed = _parse_decimal_amount(ci.get(key))
        if parsed is not None:
            return parsed
    return None


def _parse_banco_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    # Excel / SharePoint suelen guardar Fecha banco como serial numérico.
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            from openpyxl.utils.datetime import from_excel

            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted.date()
            if isinstance(converted, date):
                return converted
        except (ValueError, OverflowError, OSError):
            return None
        return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def _norm_credito_token(raw: str) -> str:
    digits = normalize_credito_digits(raw) or re.sub(r"\D", "", str(raw or ""))
    return digits or str(raw or "").strip()


def _blocking(
    code: str,
    message: str,
    *,
    id_pago: str = "",
    credito: str = "",
    creditos_seleccionados: list[str] | None = None,
    paths: list[str] | None = None,
    next_action: str = "",
) -> dict[str, Any]:
    return {
        "error_code": code,
        "message": message,
        "id_pago": id_pago,
        "credito": credito,
        "creditos_seleccionados": list(creditos_seleccionados or []),
        "paths": list(paths or []),
        "next_action": next_action,
    }


def _credit_items_from_output(output: dict[str, Any]) -> list[AbonoCreditItem]:
    items: list[AbonoCreditItem] = []
    raw_items = output.get("credit_items") or []
    if not isinstance(raw_items, list):
        return items
    for ci in raw_items:
        if not isinstance(ci, dict):
            continue
        paths = [
            str(p).strip().strip("/")
            for p in (ci.get("asiento_pdf_paths") or [])
            if str(p).strip()
        ]
        if not paths:
            single = str(ci.get("asiento_pdf_path") or "").strip().strip("/")
            if single:
                paths = [single]
        extractos = [
            str(p).strip().strip("/")
            for p in (ci.get("extracto_pdf_paths") or [])
            if str(p).strip()
        ]
        policy = resolve_manifest_policy(
            ci,
            default_canonical=TipoAplicacion.ABONO.value,
        )
        mora_ref = _mora_reference_from_credit_source(ci)
        items.append(
            AbonoCreditItem(
                credito=str(ci.get("credito") or "").strip(),
                tipo_aplicacion=policy.tipo_aplicacion_canonica,
                ruta_tabla_amortizacion=str(ci.get("ruta_tabla_amortizacion") or "").strip().strip("/"),
                ruta_unidad_credito=str(ci.get("ruta_unidad_credito") or "").strip().strip("/"),
                ruta_asientos_contables=str(ci.get("ruta_asientos_contables") or "").strip().strip("/"),
                asiento_pdf_paths=paths,
                extracto_pdf_paths=extractos,
                policy=policy,
                mora_reference_amount=mora_ref,
            )
        )
    return items


def build_abono_group_from_manifest_output(
    output: dict[str, Any],
    *,
    bank_code: str,
    process_key: str,
    abono_hist_index: dict[tuple[str, str], dict[str, Any]] | None = None,
) -> AbonoDryRunGroup:
    id_pago = str(output.get("id_pago") or "").strip()
    cliente = str(output.get("cliente") or "").strip()
    monto = _parse_decimal_amount(output.get("monto_banco"))
    fecha = _parse_banco_date(output.get("fecha_banco"))
    creditos_raw = output.get("creditos_seleccionados") or []
    creditos_sel = [
        _norm_credito_token(str(c))
        for c in creditos_raw
        if str(c).strip()
    ]
    if not creditos_sel:
        legacy = str(output.get("credito") or "").strip()
        if legacy:
            creditos_sel = [_norm_credito_token(p) for p in legacy.split(",") if p.strip()]

    credit_items = _credit_items_from_output(output)
    hist_index = abono_hist_index or {}

    enriched: list[AbonoCreditItem] = []
    for item in credit_items:
        tabla = item.ruta_tabla_amortizacion
        ruta_uc = item.ruta_unidad_credito
        cred_key = _norm_credito_token(item.credito)
        hist = hist_index.get((id_pago, cred_key)) or hist_index.get((id_pago, item.credito))
        if not tabla and hist:
            tabla = str(hist.get("tabla_amortizacion_path") or "").strip().strip("/")
        if not ruta_uc and hist:
            ruta_uc = str(hist.get("ruta_unidad_credito") or "").strip().strip("/")
        if fecha is None and hist:
            fecha = _parse_banco_date(hist.get("fecha_banco"))
        enriched.append(
            AbonoCreditItem(
                credito=item.credito or cred_key,
                tipo_aplicacion=item.tipo_aplicacion,
                ruta_tabla_amortizacion=tabla,
                ruta_unidad_credito=ruta_uc,
                ruta_asientos_contables=item.ruta_asientos_contables
                or str((hist or {}).get("ruta_asientos_contables") or "").strip().strip("/"),
                asiento_pdf_paths=list(item.asiento_pdf_paths),
                extracto_pdf_paths=list(item.extracto_pdf_paths),
                policy=item.policy,
                mora_reference_amount=item.mora_reference_amount,
            )
        )

    return AbonoDryRunGroup(
        bank_code=bank_code,
        process_key=process_key,
        id_pago=id_pago,
        cliente=cliente,
        monto_banco=monto,
        fecha_banco=fecha,
        creditos_seleccionados=creditos_sel,
        credit_items=enriched,
    )


def validate_abono_group_structure(group: AbonoDryRunGroup) -> list[dict[str, Any]]:
    errors: list[dict[str, Any]] = []
    id_pago = group.id_pago
    creditos_sel = list(group.creditos_seleccionados)

    if not id_pago:
        errors.append(
            _blocking(
                ABONO_MANIFEST_INVALID,
                "ID Pago vacío en manifest ABONO",
                id_pago=id_pago,
                creditos_seleccionados=creditos_sel,
            )
        )
    if group.monto_banco is None or group.monto_banco <= 0:
        errors.append(
            _blocking(
                ABONO_MONTO_BANCO_MISSING,
                "Monto bancario inválido o ausente",
                id_pago=id_pago,
                creditos_seleccionados=creditos_sel,
                next_action="Verifique el monto bancario en Distribucion_Abonos y vuelva a ejecutar la validación previa.",
            )
        )
    if group.fecha_banco is None:
        errors.append(
            _blocking(
                ABONO_FECHA_BANCO_MISSING,
                "Fecha bancaria inválida o ausente",
                id_pago=id_pago,
                creditos_seleccionados=creditos_sel,
            )
        )
    if not creditos_sel:
        errors.append(
            _blocking(
                ABONO_CREDIT_ITEMS_MISSING,
                "No hay créditos seleccionados en el grupo ABONO",
                id_pago=id_pago,
            )
        )
    if not group.credit_items:
        errors.append(
            _blocking(
                ABONO_CREDIT_ITEMS_MISSING,
                "credit_items vacío en manifest ABONO",
                id_pago=id_pago,
                creditos_seleccionados=creditos_sel,
            )
        )

    seen_credits: set[str] = set()
    declared: set[str] = set()
    for item in group.credit_items:
        cred = _norm_credito_token(item.credito)
        if not cred:
            errors.append(
                _blocking(
                    ABONO_MANIFEST_INVALID,
                    "credit item sin crédito",
                    id_pago=id_pago,
                    creditos_seleccionados=creditos_sel,
                )
            )
            continue
        if cred in seen_credits:
            errors.append(
                _blocking(
                    ABONO_CREDITO_DUPLICADO,
                    f"Crédito {cred} duplicado en credit_items",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                )
            )
        seen_credits.add(cred)
        declared.add(cred)

        if not item.ruta_tabla_amortizacion:
            errors.append(
                _blocking(
                    ABONO_TABLA_AMORTIZACION_MISSING,
                    f"Falta ruta de tabla de amortización para crédito {cred}",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                    next_action="Confirme la ruta de la tabla de amortización en el histórico o en SharePoint.",
                )
            )
        if not item.ruta_asientos_contables:
            errors.append(
                _blocking(
                    ABONO_ASIENTO_FALTANTE,
                    f"Falta RutaAsientosContables para crédito {cred}",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                )
            )
        if not item.asiento_pdf_paths:
            errors.append(
                _blocking(
                    ABONO_ASIENTO_FALTANTE,
                    f"Sin asiento_pdf_paths para crédito {cred}",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                    paths=[item.ruta_asientos_contables],
                    next_action="Cargue el PDF del asiento en la carpeta del crédito.",
                )
            )
        policy = item.policy or resolve_manifest_policy(
            {"tipo_aplicacion": item.tipo_aplicacion},
            default_canonical=TipoAplicacion.ABONO.value,
        )
        if policy_requires_reference_extract(policy) and not item.extracto_pdf_paths:
            errors.append(
                _blocking(
                    ABONO_MORA_REFERENCE_EXTRACT_MISSING,
                    f"ABONO MORA requiere extracto de referencia para crédito {cred}",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                    next_action="Confirme la ruta o enlace del extracto en el histórico o cargue el PDF en la carpeta del crédito.",
                )
            )

    for cred in creditos_sel:
        if cred not in declared:
            errors.append(
                _blocking(
                    ABONO_CREDITO_NO_DECLARADO,
                    f"Crédito seleccionado {cred} no tiene credit_item en manifest",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                )
            )
    for cred in declared:
        if cred not in set(creditos_sel):
            errors.append(
                _blocking(
                    ABONO_CREDITO_NO_DECLARADO,
                    f"credit_item {cred} no está en creditos_seleccionados",
                    id_pago=id_pago,
                    credito=cred,
                    creditos_seleccionados=creditos_sel,
                )
            )

    return errors


def detect_duplicate_asiento_paths(
    groups: list[AbonoDryRunGroup],
) -> dict[str, list[dict[str, str]]]:
    """
    Devuelve paths duplicados -> lista de {id_pago, credito}.
    Incluye duplicados intra-crédito (mismo path repetido en asiento_pdf_paths).
    """
    usage: dict[str, list[dict[str, str]]] = {}
    for group in groups:
        for item in group.credit_items:
            seen_in_credit: set[str] = set()
            for path in item.asiento_pdf_paths:
                norm = normalize_sharepoint_path(path)
                if not norm:
                    continue
                if norm in seen_in_credit:
                    usage.setdefault(norm, []).append(
                        {"id_pago": group.id_pago, "credito": item.credito, "kind": "intra_credit"}
                    )
                seen_in_credit.add(norm)
                usage.setdefault(norm, []).append(
                    {"id_pago": group.id_pago, "credito": item.credito, "kind": "cross_usage"}
                )
    dupes: dict[str, list[dict[str, str]]] = {}
    for path, refs in usage.items():
        if len(refs) > 1:
            dupes[path] = refs
    return dupes


def _canonical_amount_from_event(event: PaymentApplicationEvent) -> Decimal | None:
    """Campo canónico: valor_pagado_cliente (total aplicable del asiento)."""
    vp = event.valor_pagado_cliente
    if vp is None:
        return None
    try:
        amount = Decimal(str(vp))
    except InvalidOperation:
        return None
    if amount <= 0:
        return None
    return amount


async def _download_and_parse_asiento(
    download_fn,
    *,
    id_pago: str,
    cliente: str,
    credito: str,
    asiento_path: str,
) -> tuple[PaymentApplicationEvent | None, dict[str, Any] | None]:
    try:
        pdf_bytes = await download_fn(asiento_path)
        text = extract_text_from_pdf(pdf_bytes)
        event = parse_accounting_text(
            text,
            {
                "id_pago": id_pago,
                "cliente": cliente,
                "credito": credito,
                "asiento_pdf_path": asiento_path,
            },
        )
    except PdfTextNotExtractableError as exc:
        return None, _blocking(
            "PDF_TEXT_NOT_EXTRACTABLE",
            str(exc),
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        )
    except AccountingParseError as exc:
        return None, _blocking(
            getattr(exc, "error_code", None) or "ACCOUNTING_PARSE_FAILED",
            str(exc),
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        )
    except Exception as exc:
        return None, _blocking(
            "ASIENTO_DOWNLOAD_FAILED",
            str(exc)[:500],
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        )

    amount = _canonical_amount_from_event(event)
    if amount is None:
        return None, _blocking(
            ABONO_ASIENTO_TOTAL_NOT_FOUND,
            "No se pudo determinar valor_pagado_cliente del asiento",
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
            next_action="Revise que el PDF del asiento contenga el código 544111100505 con monto.",
        )
    return event, None


async def _download_abono_asiento_with_fallback(
    download_fn,
    *,
    asiento_path: str,
    id_pago: str,
    cliente: str,
    credito: str,
    bank_code: str,
    fecha_banco: date | None,
    event_index: int,
    use_event_suffix: bool = False,
    list_procesados_fn=None,
    extra_payment_dates: tuple[str, ...] | list[str] | None = None,
) -> tuple[PaymentApplicationEvent | None, dict[str, Any] | None, str | None, dict[str, Any]]:
    """
    Descarga asiento original; si falta, intenta ruta en PROCESADOS (solo parseo pendiente).

    Taxonomía alineada con PAGO / ``_download_and_parse_asiento``:
    faltante real → ``ABONO_ASIENTO_FALTANTE``; PDF dañado → parse codes.
    """
    from app.application.services.accounting_pdf_processed_move import (
        resolve_asiento_pdf_bytes_with_procesados_fallback,
    )

    fingerprint: dict[str, Any] = {}
    payment_iso = fecha_banco.isoformat() if fecha_banco else ""
    try:
        pdf_bytes, _resolved, source = await resolve_asiento_pdf_bytes_with_procesados_fallback(
            download_fn,
            asiento_path=asiento_path,
            payment_date_iso=payment_iso,
            bank_code=bank_code,
            credito=credito,
            id_pago=id_pago,
            event_index=event_index,
            use_event_suffix=use_event_suffix,
            extra_payment_dates=extra_payment_dates,
            list_procesados_fn=list_procesados_fn,
        )
    except FileNotFoundError:
        err = _blocking(
            ABONO_ASIENTO_FALTANTE,
            "No se encontró el PDF del asiento en ruta original ni en PROCESADOS",
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
            next_action="Verifique que el asiento exista o que el evento esté registrado en _AUTOMATION_LOG.",
        )
        return None, err, None, fingerprint
    except Exception as exc:
        err = _blocking(
            "ASIENTO_DOWNLOAD_FAILED",
            str(exc)[:500],
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        )
        return None, err, None, fingerprint

    fingerprint["asiento_pdf_hash"] = compute_asiento_pdf_hash(pdf_bytes)
    fingerprint["asiento_pdf_size"] = len(pdf_bytes)
    if source == "PROCESADOS":
        fingerprint["resolved_accounting_pdf_source"] = "PROCESADOS"

    try:
        text = extract_text_from_pdf(pdf_bytes)
        event = parse_accounting_text(
            text,
            {
                "id_pago": id_pago,
                "cliente": cliente,
                "credito": credito,
                "asiento_pdf_path": asiento_path,
            },
        )
        return event, None, source, fingerprint
    except PdfTextNotExtractableError as exc:
        return None, _blocking(
            "PDF_TEXT_NOT_EXTRACTABLE",
            str(exc),
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        ), source, fingerprint
    except AccountingParseError as exc:
        return None, _blocking(
            getattr(exc, "error_code", None) or "ACCOUNTING_PARSE_FAILED",
            str(exc),
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        ), source, fingerprint
    except Exception as exc:
        return None, _blocking(
            "ASIENTO_DOWNLOAD_FAILED",
            str(exc)[:500],
            id_pago=id_pago,
            credito=credito,
            paths=[asiento_path],
        ), source, fingerprint


async def _collect_abono_event_states(
    group: AbonoDryRunGroup,
    download_fn,
    table_download_fn: Callable[[str], Awaitable[bytes]],
    *,
    bank_code: str,
    duplicate_paths: dict[str, list[dict[str, str]]] | None = None,
    list_procesados_fn=None,
    extra_payment_dates: tuple[str, ...] | list[str] | None = None,
) -> list[AbonoEventReconcileState]:
    states: list[AbonoEventReconcileState] = []
    group_seen_pdf_paths: set[str] = set()
    event_counter = 0
    suffix_by_pago_cred: dict[tuple[str, str], int] = {}
    for ci in group.credit_items:
        for _ in ci.asiento_pdf_paths:
            cred_tok = _norm_credito_token(ci.credito)
            key = (group.id_pago, cred_tok or ci.credito)
            suffix_by_pago_cred[key] = suffix_by_pago_cred.get(key, 0) + 1
    use_suffix_flags = {k: v > 1 for k, v in suffix_by_pago_cred.items()}

    table_cache: dict[str, tuple[bytes, list[AppliedAbonoEventSnapshot], Any, dict[str, int], int]] = {}

    for ci in group.credit_items:
        cred = _norm_credito_token(ci.credito)
        tabla_path = (ci.ruta_tabla_amortizacion or "").strip().strip("/")
        for asiento_path in ci.asiento_pdf_paths:
            event_counter += 1
            norm = normalize_sharepoint_path(asiento_path)
            base_state = AbonoEventReconcileState(
                credit_item=ci,
                asiento_path=asiento_path,
                norm_path=norm or asiento_path,
                event_index=event_counter,
                status="PENDING",
            )
            if norm in group_seen_pdf_paths:
                base_state.status = "ERROR"
                base_state.error = _blocking(
                    ABONO_ASIENTO_DUPLICADO,
                    f"Asiento repetido en el grupo: {norm}",
                    id_pago=group.id_pago,
                    credito=cred,
                    creditos_seleccionados=group.creditos_seleccionados,
                    paths=[norm],
                )
                states.append(base_state)
                continue
            if duplicate_paths and norm in duplicate_paths:
                base_state.status = "ERROR"
                base_state.error = _blocking(
                    ABONO_ASIENTO_DUPLICADO,
                    f"Asiento reutilizado: {norm}",
                    id_pago=group.id_pago,
                    credito=cred,
                    creditos_seleccionados=group.creditos_seleccionados,
                    paths=[norm],
                    next_action="Use un PDF de asiento distinto por crédito y por ID Pago.",
                )
                states.append(base_state)
                continue
            group_seen_pdf_paths.add(norm)

            idem_key = build_amortization_idempotency_key(
                group.id_pago, cred or ci.credito, asiento_path, ""
            )
            base_state.idempotency_key = idem_key

            if not tabla_path:
                base_state.status = "ERROR"
                base_state.error = _blocking(
                    ABONO_TABLA_AMORTIZACION_MISSING,
                    "Falta ruta de tabla de amortización",
                    id_pago=group.id_pago,
                    credito=cred,
                )
                states.append(base_state)
                continue

            if tabla_path not in table_cache:
                try:
                    tabla_bytes = await table_download_fn(tabla_path)
                    wb = openpyxl.load_workbook(io.BytesIO(tabla_bytes), data_only=True)
                    try:
                        sheet_match = detect_amortization_sheet(
                            wb, tabla_amortizacion_path=tabla_path
                        )
                        snapshots = load_applied_abono_event_snapshots(
                            wb, id_pago_filter=group.id_pago
                        )
                        table_cache[tabla_path] = (
                            tabla_bytes,
                            snapshots,
                            sheet_match.worksheet,
                            sheet_match.headers,
                            sheet_match.header_row,
                        )
                    finally:
                        closer = getattr(wb, "close", None)
                        if callable(closer):
                            closer()
                except Exception as exc:
                    base_state.status = "ERROR"
                    base_state.error = _blocking(
                        "TABLE_DOWNLOAD_FAILED",
                        str(exc)[:500],
                        id_pago=group.id_pago,
                        credito=cred,
                        paths=[tabla_path],
                    )
                    states.append(base_state)
                    continue

            _, snapshots, ws, headers, _header_row = table_cache[tabla_path]
            match, amb_err = find_applied_abono_event(
                snapshots,
                id_pago=group.id_pago,
                credito=cred or ci.credito,
                asiento_pdf_path=asiento_path,
                idempotency_key=idem_key,
            )
            if amb_err:
                base_state.status = "ERROR"
                base_state.error = _blocking(
                    amb_err,
                    "Coincidencia ambigua en _AUTOMATION_LOG para el evento ABONO",
                    id_pago=group.id_pago,
                    credito=cred,
                    paths=[asiento_path],
                )
                states.append(base_state)
                continue

            if match is not None:
                amount, amount_err = resolve_applied_abono_amount(
                    match, ws=ws, headers=headers
                )
                if amount_err or amount is None:
                    base_state.status = "ERROR"
                    base_state.error = _blocking(
                        amount_err or APPLIED_ABONO_AMOUNT_UNRESOLVABLE,
                        "No se pudo recuperar el monto del evento ABONO ya aplicado",
                        id_pago=group.id_pago,
                        credito=cred,
                        paths=[asiento_path],
                        next_action="Complete ValorPagadoCliente en el log o verifique la fila de aplicación.",
                    )
                    states.append(base_state)
                    continue
                base_state.status = ABONO_APPLICATION_STATUS_ALREADY_APPLIED
                base_state.amount = amount
                base_state.applied_snapshot = match
                base_state.requires_pdf_download = False
                states.append(base_state)
                continue

            use_suffix = use_suffix_flags.get((group.id_pago, cred or ci.credito), False)
            event, err, pdf_source, fingerprint = await _download_abono_asiento_with_fallback(
                download_fn,
                asiento_path=asiento_path,
                id_pago=group.id_pago,
                cliente=group.cliente,
                credito=cred or ci.credito,
                bank_code=bank_code,
                fecha_banco=group.fecha_banco,
                event_index=event_counter,
                use_event_suffix=use_suffix,
                list_procesados_fn=list_procesados_fn,
                extra_payment_dates=extra_payment_dates,
            )
            base_state.pdf_fingerprint = fingerprint
            base_state.resolved_pdf_source = pdf_source
            if err or event is None:
                base_state.status = "ERROR"
                base_state.error = err or _blocking(
                    "ACCOUNTING_PARSE_FAILED",
                    "No se pudo parsear el asiento",
                    id_pago=group.id_pago,
                    credito=cred,
                    paths=[asiento_path],
                )
                states.append(base_state)
                continue
            amount = _canonical_amount_from_event(event)
            if amount is None:
                base_state.status = "ERROR"
                base_state.error = _blocking(
                    ABONO_ASIENTO_TOTAL_NOT_FOUND,
                    "No se pudo determinar valor_pagado_cliente del asiento",
                    id_pago=group.id_pago,
                    credito=cred,
                    paths=[asiento_path],
                )
                states.append(base_state)
                continue
            base_state.status = "PENDING"
            base_state.event = event
            base_state.amount = amount
            base_state.idempotency_key = build_amortization_idempotency_key(
                group.id_pago,
                cred or ci.credito,
                asiento_path,
                event.comprobante or "",
                pdf_hash=str(fingerprint.get("asiento_pdf_hash") or ""),
            )
            states.append(base_state)

    return states


def reconcile_abono_event_states(
    group: AbonoDryRunGroup,
    states: list[AbonoEventReconcileState],
) -> AbonoReconciliationResult:
    """Cuadre híbrido: montos ya aplicados + pendientes vs monto bancario."""
    blocking: list[dict[str, Any]] = []
    credit_amounts: dict[str, Decimal] = {}
    pdf_amounts: dict[str, Decimal] = {}
    already_total = Decimal("0")
    pending_total = Decimal("0")
    already_count = 0
    pending_count = 0

    for st in states:
        if st.error:
            blocking.append(st.error)
            continue
        if st.amount is None:
            continue
        cred = _norm_credito_token(st.credit_item.credito)
        if st.status == ABONO_APPLICATION_STATUS_ALREADY_APPLIED:
            already_total += st.amount
            already_count += 1
            pdf_amounts[st.norm_path] = st.amount
            if cred:
                credit_amounts[cred] = credit_amounts.get(cred, Decimal("0")) + st.amount
        elif st.status == "PENDING" and st.event is not None:
            pending_total += st.amount
            pending_count += 1
            pdf_amounts[st.norm_path] = st.amount
            if cred:
                credit_amounts[cred] = credit_amounts.get(cred, Decimal("0")) + st.amount

    total = already_total + pending_total
    monto_banco = group.monto_banco
    if monto_banco is None:
        return AbonoReconciliationResult(
            reconciliation_status="FAILED",
            monto_banco=None,
            total_asientos=total,
            diferencia=Decimal("0"),
            tolerancia=ABONO_RECONCILIATION_TOLERANCE,
            credit_amounts=credit_amounts,
            accounting_pdf_amounts=pdf_amounts,
            blocking_errors=blocking,
            already_applied_amount=already_total,
            pending_amount=pending_total,
            total_reconciled_amount=total,
            already_applied_events_count=already_count,
            pending_events_count=pending_count,
        )

    if blocking:
        return AbonoReconciliationResult(
            reconciliation_status="FAILED",
            monto_banco=monto_banco,
            total_asientos=total,
            diferencia=total - monto_banco,
            tolerancia=ABONO_RECONCILIATION_TOLERANCE,
            credit_amounts=credit_amounts,
            accounting_pdf_amounts=pdf_amounts,
            blocking_errors=blocking,
            already_applied_amount=already_total,
            pending_amount=pending_total,
            total_reconciled_amount=total,
            already_applied_events_count=already_count,
            pending_events_count=pending_count,
        )

    diferencia = total - monto_banco
    diff_abs = abs(diferencia)
    status = "PASSED" if diff_abs <= ABONO_RECONCILIATION_TOLERANCE else "FAILED"
    if status == "FAILED":
        blocking.append(
            _blocking(
                ABONO_ASIENTOS_NO_CUADRAN,
                "La suma híbrida (aplicados + pendientes) no cuadra con el monto bancario",
                id_pago=group.id_pago,
                creditos_seleccionados=group.creditos_seleccionados,
                next_action="Revise los montos de los asientos y el monto bancario en Distribucion_Abonos.",
            )
        )

    return AbonoReconciliationResult(
        reconciliation_status=status,
        monto_banco=monto_banco,
        total_asientos=total,
        diferencia=diferencia,
        tolerancia=ABONO_RECONCILIATION_TOLERANCE,
        credit_amounts=credit_amounts,
        accounting_pdf_amounts=pdf_amounts,
        blocking_errors=blocking,
        already_applied_amount=already_total,
        pending_amount=pending_total,
        total_reconciled_amount=total,
        already_applied_events_count=already_count,
        pending_events_count=pending_count,
    )


async def reconcile_abono_group(
    group: AbonoDryRunGroup,
    download_fn,
    *,
    duplicate_paths: dict[str, list[dict[str, str]]] | None = None,
) -> AbonoReconciliationResult:
    blocking: list[dict[str, Any]] = []
    credit_amounts: dict[str, Decimal] = {}
    pdf_amounts: dict[str, Decimal] = {}
    total = Decimal("0")
    group_seen_pdf_paths: set[str] = set()

    for item in group.credit_items:
        cred = _norm_credito_token(item.credito)
        credit_total = Decimal("0")
        for asiento_path in item.asiento_pdf_paths:
            norm = normalize_sharepoint_path(asiento_path)
            if norm in group_seen_pdf_paths:
                blocking.append(
                    _blocking(
                        ABONO_ASIENTO_DUPLICADO,
                        f"Asiento repetido en el grupo: {norm}",
                        id_pago=group.id_pago,
                        credito=cred,
                        creditos_seleccionados=group.creditos_seleccionados,
                        paths=[norm],
                    )
                )
                continue
            if duplicate_paths and norm in duplicate_paths:
                blocking.append(
                    _blocking(
                        ABONO_ASIENTO_DUPLICADO,
                        f"Asiento reutilizado: {norm}",
                        id_pago=group.id_pago,
                        credito=cred,
                        creditos_seleccionados=group.creditos_seleccionados,
                        paths=[norm],
                        next_action="Use un PDF de asiento distinto por crédito y por ID Pago.",
                    )
                )
                continue
            event, err = await _download_and_parse_asiento(
                download_fn,
                id_pago=group.id_pago,
                cliente=group.cliente,
                credito=cred,
                asiento_path=asiento_path,
            )
            if err:
                blocking.append(err)
                continue
            assert event is not None
            amount = _canonical_amount_from_event(event)
            assert amount is not None
            group_seen_pdf_paths.add(norm)
            credit_total += amount
            pdf_amounts[norm or asiento_path] = amount
            total += amount
        if cred:
            credit_amounts[cred] = credit_total

    monto_banco = group.monto_banco
    if monto_banco is None:
        return AbonoReconciliationResult(
            reconciliation_status="FAILED",
            monto_banco=None,
            total_asientos=total,
            diferencia=Decimal("0"),
            tolerancia=ABONO_RECONCILIATION_TOLERANCE,
            credit_amounts=credit_amounts,
            accounting_pdf_amounts=pdf_amounts,
            blocking_errors=blocking,
        )

    if blocking:
        return AbonoReconciliationResult(
            reconciliation_status="FAILED",
            monto_banco=monto_banco,
            total_asientos=total,
            diferencia=total - monto_banco,
            tolerancia=ABONO_RECONCILIATION_TOLERANCE,
            credit_amounts=credit_amounts,
            accounting_pdf_amounts=pdf_amounts,
            blocking_errors=blocking,
        )

    diferencia = total - monto_banco
    diff_abs = abs(diferencia)
    status = "PASSED" if diff_abs <= ABONO_RECONCILIATION_TOLERANCE else "FAILED"
    if status == "FAILED":
        blocking.append(
            _blocking(
                ABONO_ASIENTOS_NO_CUADRAN,
                "La suma de asientos no cuadra con el monto bancario",
                id_pago=group.id_pago,
                creditos_seleccionados=group.creditos_seleccionados,
                next_action="Revise los montos de los asientos y el monto bancario en Distribucion_Abonos.",
            )
        )

    return AbonoReconciliationResult(
        reconciliation_status=status,
        monto_banco=monto_banco,
        total_asientos=total,
        diferencia=diferencia,
        tolerancia=ABONO_RECONCILIATION_TOLERANCE,
        credit_amounts=credit_amounts,
        accounting_pdf_amounts=pdf_amounts,
        blocking_errors=blocking,
    )


def _payment_application_dict(event: PaymentApplicationEvent) -> dict[str, float]:
    return {
        "valor_pagado_cliente": event.valor_pagado_cliente,
        "capital": event.capital,
        "intereses": event.intereses,
        "mora": event.mora,
        "retenciones": event.retenciones,
        "saldos_menores": event.saldos_menores,
    }


def _abono_payment_date_fields(
    group: AbonoDryRunGroup,
    event: PaymentApplicationEvent,
) -> dict[str, Any]:
    """
    Fecha pago en ABONO: únicamente Fecha banco (Excel BANCO_* / histórico).

    La fecha del asiento queda solo como auditoría (`fecha_asiento` /
    `payment_date_matches_asiento`); nunca se escribe en «Fecha pago».
    """
    out: dict[str, Any] = {}
    if event.fecha_asiento is not None:
        out["fecha_asiento"] = event.fecha_asiento.isoformat()
    if group.fecha_banco is None:
        out["payment_date_source"] = "none"
        return out
    iso = group.fecha_banco.isoformat()
    out["payment_date_iso"] = iso
    out["payment_date_source"] = "fecha_banco"
    if event.fecha_asiento is not None:
        out["payment_date_matches_asiento"] = iso == event.fecha_asiento.isoformat()
    return out


def _abono_ibr_block(policy: ApplicationPolicy | None = None) -> dict[str, Any]:
    if policy is None or not policy.actualiza_ibr:
        return {
            "required_date": None,
            "found": False,
            "value": None,
            "status": "NOT_REQUIRED",
        }
    return {
        "required_date": None,
        "found": False,
        "value": None,
        "status": "PENDING_IBR",
    }


def _abono_observability_base(
    *,
    group: AbonoDryRunGroup,
    schedule: AbonoScheduleContextResult,
    policy: ApplicationPolicy | None = None,
) -> dict[str, Any]:
    resolved = policy or resolve_manifest_policy(
        {"tipo_aplicacion": TipoAplicacion.ABONO.value},
        default_canonical=TipoAplicacion.ABONO.value,
    )
    base = {
        **policy_observability_dict(resolved),
        "ibr_skipped_reason": (
            ABONO_IBR_SKIPPED_REASON if not resolved.actualiza_ibr else None
        ),
        "schedule_resolution_status": schedule.status,
        "schedule_context": schedule.status,
        "application_row_strategy": ABONO_APPLICATION_ROW_STRATEGY,
        "group_reconciliation_status": group.reconciliation_status,
        "group_ready_for_apply": group.group_ready_for_apply,
        "requires_business_rule": group.requires_business_rule,
        "due_date_row": None,
        "fecha_limite_pago": None,
        "ibr_row": None,
        "ibr": _abono_ibr_block(resolved),
    }
    return base


def _mora_coverage_for_event(
    credit_item: AbonoCreditItem,
    event: PaymentApplicationEvent | None,
) -> dict[str, Any]:
    policy = credit_item.policy
    if policy is None or policy.subtipo_aplicacion != ApplicationSubtype.MORA:
        return {}
    accounting_mora: Decimal | None = None
    accounting_total: Decimal | None = None
    if event is not None:
        if event.mora is not None:
            accounting_mora = _parse_decimal_amount(event.mora)
        accounting_total = _canonical_amount_from_event(event)
    return compute_mora_reference_coverage(
        mora_reference_amount=credit_item.mora_reference_amount,
        accounting_mora=accounting_mora,
        accounting_total=accounting_total,
    )


def _build_already_applied_abono_item(
    group: AbonoDryRunGroup,
    *,
    state: AbonoEventReconcileState,
    schedule: AbonoScheduleContextResult,
) -> dict[str, Any]:
    cred = _norm_credito_token(state.credit_item.credito)
    snap = state.applied_snapshot
    app_row = snap.application_row if snap else None
    amount = float(state.amount) if state.amount is not None else None
    extracto_path = (
        (state.credit_item.extracto_pdf_paths or [None])[0]
        if state.credit_item.extracto_pdf_paths
        else None
    )
    policy = state.credit_item.policy
    return {
        "id_pago": group.id_pago,
        "cliente": group.cliente,
        "credito": cred or state.credit_item.credito,
        "asiento_pdf_path": state.asiento_path,
        "extracto_pdf_path": extracto_path,
        "event_index": state.event_index,
        "comprobante": None,
        "tabla_amortizacion_path": (state.credit_item.ruta_tabla_amortizacion or "").strip("/") or None,
        "detected_codes": [],
        "parser_mode": None,
        "payment_application": {
            "valor_pagado_cliente": amount,
            "capital": None,
            "intereses": None,
            "mora": None,
            "retenciones": None,
            "saldos_menores": None,
        },
        "warnings": [],
        **_abono_observability_base(group=group, schedule=schedule, policy=policy),
        "application_row_strategy": ABONO_APPLICATION_ROW_STRATEGY_EXISTING,
        "requires_pdf_download": False,
        "requires_new_application_row": False,
        "idempotency_key": state.idempotency_key,
        "application_row": app_row,
        "target_row": app_row,
        "application_status": ABONO_APPLICATION_STATUS_ALREADY_APPLIED,
        "error_code": None,
        "already_applied_from_log": True,
        "accounting_pdf_move_status": "already_processed",
    }


async def _plan_abono_asiento_item(
    group: AbonoDryRunGroup,
    *,
    credit_item: AbonoCreditItem,
    asiento_path: str,
    event_index: int,
    event: PaymentApplicationEvent,
    pdf_fingerprint: dict[str, Any],
    schedule: AbonoScheduleContextResult,
    table_download_fn: Callable[[str], Awaitable[bytes]],
    used_application_rows_by_table: dict[str, set[int]],
) -> dict[str, Any]:
    cred = _norm_credito_token(credit_item.credito)
    tabla_path = (credit_item.ruta_tabla_amortizacion or "").strip().strip("/")
    policy = credit_item.policy
    extracto_path = (
        (credit_item.extracto_pdf_paths or [None])[0]
        if credit_item.extracto_pdf_paths
        else None
    )
    mora_coverage = _mora_coverage_for_event(credit_item, event)
    item_warnings = list(event.parse_warnings or [])
    if mora_coverage.get("mora_reference_coverage_status") == "MISMATCH":
        item_warnings.append(
            "mora_reference_mismatch:"
            f"ref={mora_coverage.get('mora_reference_amount')}"
            f",asiento={mora_coverage.get('accounting_total')}"
            f",diff={mora_coverage.get('mora_reference_difference')}"
        )
    base = {
        "id_pago": group.id_pago,
        "cliente": group.cliente,
        "credito": cred or credit_item.credito,
        "asiento_pdf_path": asiento_path,
        "extracto_pdf_path": extracto_path,
        "event_index": event_index,
        "comprobante": event.comprobante or None,
        "tabla_amortizacion_path": tabla_path or None,
        "detected_codes": list(event.detected_codes),
        "parser_mode": event.parser_mode or None,
        "payment_application": _payment_application_dict(event),
        "warnings": item_warnings,
        **_abono_observability_base(group=group, schedule=schedule, policy=policy),
        **mora_coverage,
        **_abono_payment_date_fields(group, event),
        **pdf_fingerprint,
        "idempotency_key": build_amortization_idempotency_key(
            group.id_pago, cred or credit_item.credito, asiento_path, event.comprobante
        ),
    }

    if not tabla_path:
        return {
            **base,
            "application_row": None,
            "target_row": None,
            "application_status": "ERROR",
            "error_code": ABONO_TABLA_AMORTIZACION_MISSING,
        }

    try:
        tabla_bytes = await table_download_fn(tabla_path)
    except Exception as exc:
        return {
            **base,
            "application_row": None,
            "target_row": None,
            "application_status": "ERROR",
            "error_code": "TABLE_DOWNLOAD_FAILED",
            "warnings": base.get("warnings", []) + [str(exc)[:500]],
        }

    wb = openpyxl.load_workbook(io.BytesIO(tabla_bytes), data_only=True)
    try:
        try:
            sheet_match = detect_amortization_sheet(wb, tabla_amortizacion_path=tabla_path)
        except AmortizationSheetNotFoundError as exc:
            return {
                **base,
                "application_row": None,
                "target_row": None,
                "application_status": "ERROR",
                "error_code": "AMORTIZATION_SHEET_NOT_FOUND",
                "warnings": base.get("warnings", [])
                + [str(exc), f"workbook_sheets={exc.workbook_sheets}"],
            }

        ws = sheet_match.worksheet
        headers = sheet_match.headers
        header_row = sheet_match.header_row
        base["sheet_name"] = ws.title

        tabla_key = normalize_sharepoint_path(tabla_path)
        reserved = used_application_rows_by_table.setdefault(tabla_key, set())

        payment_date: date | None = None
        raw_pd = str(base.get("payment_date_iso") or "").strip()
        if raw_pd:
            try:
                payment_date = date.fromisoformat(raw_pd)
            except ValueError:
                payment_date = None

        app_result = find_next_available_application_row(
            ws,
            headers,
            event,
            reserved_rows=reserved,
            header_row=header_row,
            payment_date=payment_date,
            detected_codes=frozenset(event.detected_codes),
            warnings=frozenset(event.parse_warnings or []),
        )
        application_row, compare_status = resolve_planned_application_row(
            app_result,
            allow_suggested_row=False,
        )

        if application_row is None:
            if app_result.requires_new_row:
                app_debug = build_application_row_search_debug(
                    ws,
                    headers,
                    event,
                    due_date_row=None,
                    header_row=header_row,
                    sheet_name=ws.title,
                    tabla_amortizacion_path=tabla_path,
                    exclude_rows=frozenset(reserved),
                    find_result=app_result,
                )
                occupied = sum(
                    1
                    for c in app_debug.get("candidates") or []
                    if c.get("has_application_data")
                )
                return {
                    **base,
                    "application_row": None,
                    "target_row": None,
                    "application_status": "ERROR",
                    "error_code": APPLICATION_PAYMENT_SECTION_FULL,
                    "user_message": (
                        "La sección de Aplicación de Pagos no tiene filas disponibles "
                        "para registrar el movimiento."
                    ),
                    "next_action": (
                        "Amplíe de forma controlada la sección en la plantilla de "
                        "amortización y vuelva a ejecutar la validación previa. Si el error continúa, contacte a soporte."
                    ),
                    "application_section_full": {
                        "tabla_amortizacion_path": tabla_path,
                        "sheet_name": ws.title,
                        "id_pago": group.id_pago,
                        "credito": cred or credit_item.credito,
                        "search_start_row": app_result.search_start_row,
                        "last_safe_row": ws.max_row,
                        "suggested_row": app_result.suggested_row,
                        "occupied_rows_count": occupied,
                    },
                }
            return {
                **base,
                "application_row": None,
                "target_row": None,
                "application_status": "ERROR",
                "error_code": "APPLICATION_ROW_NOT_FOUND",
            }

        if compare_status == REVISION_MANUAL:
            return {
                **base,
                "application_row": application_row,
                "target_row": application_row,
                "application_status": "REVISION_MANUAL",
                "error_code": None,
            }

        application_status = _ABONO_STATUS_MAP.get(compare_status or APLICADO, "ERROR")
        if application_status == "ERROR":
            return {
                **base,
                "application_row": application_row,
                "target_row": application_row,
                "application_status": "ERROR",
                "error_code": "APPLICATION_ROW_CONFLICT",
            }

        if compare_status == APLICADO:
            reserved.add(application_row)

        return {
            **base,
            "application_row": application_row,
            "target_row": application_row,
            "application_status": application_status,
            "error_code": None,
        }
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()


def abono_group_result_dict(
    group: AbonoDryRunGroup,
    reconciliation: AbonoReconciliationResult,
    schedule: AbonoScheduleContextResult,
) -> dict[str, Any]:
    return {
        "id_pago": group.id_pago,
        "bank_code": group.bank_code,
        "cliente": group.cliente,
        "monto_banco": float(group.monto_banco) if group.monto_banco is not None else None,
        "fecha_banco": group.fecha_banco.isoformat() if group.fecha_banco else None,
        "total_asientos": float(reconciliation.total_asientos),
        "diferencia": float(reconciliation.diferencia),
        "tolerancia": float(reconciliation.tolerancia),
        "creditos_seleccionados": list(group.creditos_seleccionados),
        "credit_amounts": {k: float(v) for k, v in reconciliation.credit_amounts.items()},
        "accounting_pdf_amounts": {k: float(v) for k, v in reconciliation.accounting_pdf_amounts.items()},
        "reconciliation_status": reconciliation.reconciliation_status,
        "schedule_resolution_status": schedule.status,
        "group_ready_for_apply": group.group_ready_for_apply,
        "requires_business_rule": group.requires_business_rule,
        "blocking_errors": list(group.blocking_errors),
        "warnings": list(group.warnings),
        "already_applied_amount": float(reconciliation.already_applied_amount),
        "pending_amount": float(reconciliation.pending_amount),
        "total_reconciled_amount": float(reconciliation.total_reconciled_amount),
        "already_applied_events_count": reconciliation.already_applied_events_count,
        "pending_events_count": reconciliation.pending_events_count,
    }


def load_abono_historical_index(hist_bytes: bytes) -> dict[tuple[str, str], dict[str, Any]]:
    """Índice (id_pago, crédito) desde hoja Distribucion_Abonos del histórico."""
    import io

    import openpyxl

    from app.application.services.historical_application_rows import (
        _find_abonos_header_row,
        _find_distribucion_abonos_sheet,
        _get_col_abono,
    )
    from app.application.use_cases.send_validar_extractos_notification import _excel_cell_display

    wb = openpyxl.load_workbook(io.BytesIO(hist_bytes), data_only=True)
    try:
        ws = _find_distribucion_abonos_sheet(wb)
        if ws is None:
            return {}
        h_row, header_map = _find_abonos_header_row(ws)
        col_id = _get_col_abono(header_map, DistribucionAbonosCols.ID_PAGO, "ID Pago")
        col_cred = _get_col_abono(header_map, DistribucionAbonosCols.CREDITO, "Crédito")
        col_cred_norm = _get_col_abono(
            header_map, DistribucionAbonosCols.CREDITO_NORMALIZADO, "CreditoNormalizado"
        )
        col_tabla = _get_col_abono(
            header_map, DistribucionAbonosCols.RUTA_TABLA_AMORTIZACION, "RutaTablaAmortizacion"
        )
        col_uc = _get_col_abono(
            header_map, DistribucionAbonosCols.RUTA_UNIDAD_CREDITO, "RutaUnidadCredito"
        )
        col_ra = _get_col_abono(
            header_map, DistribucionAbonosCols.RUTA_ASIENTOS_CONTABLES, "RutaAsientosContables"
        )
        if not col_id or not col_cred:
            return {}
        index: dict[tuple[str, str], dict[str, Any]] = {}
        for r in range(h_row + 1, (ws.max_row or h_row) + 1):
            id_p = _excel_cell_display(ws.cell(r, col_id).value).strip()
            cred_vis = _excel_cell_display(ws.cell(r, col_cred).value).strip()
            if not id_p:
                continue
            cred_norm = ""
            if col_cred_norm:
                cred_norm = _excel_cell_display(ws.cell(r, col_cred_norm).value).strip()
            if not cred_norm:
                cred_norm = _norm_credito_token(cred_vis)
            tabla = ""
            if col_tabla:
                tabla = _excel_cell_display(ws.cell(r, col_tabla).value).strip().strip("/")
            ruta_uc = ""
            if col_uc:
                ruta_uc = _excel_cell_display(ws.cell(r, col_uc).value).strip().strip("/")
            ruta_as = ""
            if col_ra:
                ruta_as = _excel_cell_display(ws.cell(r, col_ra).value).strip().strip("/")
            row_data = {
                "tabla_amortizacion_path": tabla,
                "ruta_unidad_credito": ruta_uc,
                "ruta_asientos_contables": ruta_as,
                "credito_normalizado": cred_norm,
            }
            index[(id_p, cred_norm or cred_vis)] = row_data
            if cred_vis and cred_vis != (cred_norm or cred_vis):
                index.setdefault((id_p, cred_vis), row_data)
        return index
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()


async def process_abono_manifest_outputs(
    outputs: list[dict[str, Any]],
    *,
    bank_code: str,
    process_key: str,
    download_fn,
    abono_hist_index: dict[tuple[str, str], dict[str, Any]] | None = None,
    table_download_fn: Callable[[str], Awaitable[bytes]] | None = None,
    asiento_metadata_fn: Callable[[str], Awaitable[dict[str, Any]]] | None = None,
    used_application_rows_by_table: dict[str, set[int]] | None = None,
    list_procesados_fn=None,
    extra_payment_dates: tuple[str, ...] | list[str] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    """
    Procesa outputs ABONO del manifest.

    Devuelve (items observabilidad, abono_group_results, contadores).
    """
    groups = [
        build_abono_group_from_manifest_output(
            out,
            bank_code=bank_code,
            process_key=process_key,
            abono_hist_index=abono_hist_index,
        )
        for out in outputs
        if isinstance(out, dict)
    ]

    counters = {
        "abono_groups_total": len(groups),
        "abono_groups_reconciled": 0,
        "abono_groups_not_reconciled": 0,
        "abono_groups_missing_accounting_pdf": 0,
        "abono_groups_schedule_rule_missing": 0,
        "abono_groups_ready": 0,
        "abono_credit_items_total": sum(len(g.credit_items) for g in groups),
    }
    reserved_rows = used_application_rows_by_table if used_application_rows_by_table is not None else {}

    duplicate_paths = detect_duplicate_asiento_paths(groups)
    all_items: list[dict[str, Any]] = []
    group_results: list[dict[str, Any]] = []

    for group in groups:
        struct_errors = validate_abono_group_structure(group)
        if struct_errors:
            group.blocking_errors.extend(struct_errors)
            if any(e.get("error_code") == ABONO_ASIENTO_FALTANTE for e in struct_errors):
                counters["abono_groups_missing_accounting_pdf"] += 1
            reconciliation = AbonoReconciliationResult(
                reconciliation_status="FAILED",
                monto_banco=group.monto_banco,
                total_asientos=Decimal("0"),
                diferencia=Decimal("0"),
                tolerancia=ABONO_RECONCILIATION_TOLERANCE,
                credit_amounts={},
                accounting_pdf_amounts={},
                blocking_errors=struct_errors,
            )
            schedule = AbonoScheduleContextResult(
                status="SKIPPED",
                reference_date=None,
                due_date_row=None,
                application_row=None,
                ibr_date=None,
            )
            group.reconciliation_status = "FAILED"
            group.schedule_resolution_status = schedule.status
            group.group_ready_for_apply = False
            counters["abono_groups_not_reconciled"] += 1
            group_results.append(abono_group_result_dict(group, reconciliation, schedule))
            all_items.extend(build_abono_observability_items(group, reconciliation, schedule, {}))
            continue

        event_states: list[AbonoEventReconcileState] = []
        if table_download_fn is not None:
            event_states = await _collect_abono_event_states(
                group,
                download_fn,
                table_download_fn,
                bank_code=bank_code,
                duplicate_paths=duplicate_paths,
                list_procesados_fn=list_procesados_fn,
                extra_payment_dates=extra_payment_dates,
            )
            reconciliation = reconcile_abono_event_states(group, event_states)
        else:
            reconciliation = await reconcile_abono_group(
                group,
                download_fn,
                duplicate_paths=duplicate_paths,
            )

        group.blocking_errors.extend(reconciliation.blocking_errors)
        group.reconciliation_status = reconciliation.reconciliation_status

        if reconciliation.reconciliation_status == "PASSED" and not reconciliation.blocking_errors:
            counters["abono_groups_reconciled"] += 1
        else:
            counters["abono_groups_not_reconciled"] += 1
            if any(
                e.get("error_code") in (ABONO_ASIENTO_FALTANTE, ABONO_ASIENTO_TOTAL_NOT_FOUND)
                for e in reconciliation.blocking_errors
            ):
                counters["abono_groups_missing_accounting_pdf"] += 1

        schedule = resolve_abono_schedule_context(group=group, reconciliation=reconciliation)
        group.schedule_resolution_status = schedule.status

        if reconciliation.reconciliation_status == "PASSED" and schedule.status == SCHEDULE_NOT_REQUIRED:
            group.requires_business_rule = False
            group.group_ready_for_apply = True
        else:
            group.group_ready_for_apply = False

        parsed_events: dict[str, PaymentApplicationEvent] = {}
        group_items: list[dict[str, Any]] = []

        if group.group_ready_for_apply and table_download_fn is not None and event_states:
            for st in event_states:
                if st.error:
                    cred_tok = _norm_credito_token(st.credit_item.credito)
                    group_items.append(
                        {
                            "id_pago": group.id_pago,
                            "cliente": group.cliente,
                            "credito": cred_tok or st.credit_item.credito,
                            "asiento_pdf_path": st.asiento_path,
                            "event_index": st.event_index,
                            "application_status": "ERROR",
                            "error_code": st.error.get("error_code"),
                            **_abono_observability_base(
                                group=group,
                                schedule=schedule,
                                policy=st.credit_item.policy,
                            ),
                        }
                    )
                    group.group_ready_for_apply = False
                    continue

                if st.status == ABONO_APPLICATION_STATUS_ALREADY_APPLIED:
                    group_items.append(
                        _build_already_applied_abono_item(
                            group, state=st, schedule=schedule
                        )
                    )
                    continue

                assert st.event is not None
                norm = st.norm_path
                parsed_events[norm] = st.event
                pdf_fingerprint = dict(st.pdf_fingerprint)
                if asiento_metadata_fn is not None and st.requires_pdf_download:
                    try:
                        meta = await asiento_metadata_fn(st.asiento_path)
                        pdf_fingerprint["asiento_pdf_etag"] = str(
                            meta.get("eTag") or meta.get("etag") or ""
                        )
                        pdf_fingerprint["asiento_pdf_last_modified"] = meta.get(
                            "lastModifiedDateTime"
                        )
                    except Exception:
                        pass
                if st.resolved_pdf_source:
                    pdf_fingerprint["resolved_accounting_pdf_source"] = st.resolved_pdf_source

                planned_item = await _plan_abono_asiento_item(
                    group,
                    credit_item=st.credit_item,
                    asiento_path=st.asiento_path,
                    event_index=st.event_index,
                    event=st.event,
                    pdf_fingerprint=pdf_fingerprint,
                    schedule=schedule,
                    table_download_fn=table_download_fn,
                    used_application_rows_by_table=reserved_rows,
                )
                if st.idempotency_key:
                    planned_item["idempotency_key"] = st.idempotency_key
                coverage_status = planned_item.get("mora_reference_coverage_status")
                if coverage_status == "MISMATCH":
                    group.warnings.append(
                        f"mora_reference_mismatch|id_pago={group.id_pago}|credito={st.credit_item.credito}"
                    )
                group_items.append(planned_item)
                if planned_item.get("application_status") not in (
                    "WOULD_APPLY",
                    "WOULD_ADOPT_EXISTING",
                    ABONO_APPLICATION_STATUS_ALREADY_APPLIED,
                ):
                    group.group_ready_for_apply = False
        elif not group.group_ready_for_apply:
            if reconciliation.reconciliation_status == "PASSED" and not event_states:
                for item in group.credit_items:
                    for asiento_path in item.asiento_pdf_paths:
                        norm = normalize_sharepoint_path(asiento_path)
                        event, _ = await _download_and_parse_asiento(
                            download_fn,
                            id_pago=group.id_pago,
                            cliente=group.cliente,
                            credito=item.credito,
                            asiento_path=asiento_path,
                        )
                        if event:
                            parsed_events[norm or asiento_path] = event
            group_items = build_abono_observability_items(
                group, reconciliation, schedule, parsed_events
            )
        else:
            group_items = build_abono_observability_items(
                group, reconciliation, schedule, parsed_events
            )

        if group.group_ready_for_apply:
            counters["abono_groups_ready"] += 1

        group_results.append(abono_group_result_dict(group, reconciliation, schedule))
        all_items.extend(group_items)

    return all_items, group_results, counters


def build_abono_observability_items(
    group: AbonoDryRunGroup,
    reconciliation: AbonoReconciliationResult,
    schedule: AbonoScheduleContextResult,
    parsed_events: dict[str, PaymentApplicationEvent],
) -> list[dict[str, Any]]:
    """Eventos de observabilidad para grupos bloqueados (sin planificación de tabla)."""
    items: list[dict[str, Any]] = []
    doc_blocking_errors = list(
        (group.blocking_errors or []) + (reconciliation.blocking_errors or [])
    )
    doc_blocked = reconciliation.reconciliation_status != "PASSED" or bool(doc_blocking_errors)

    for idx, ci in enumerate(group.credit_items, start=1):
        cred = _norm_credito_token(ci.credito)
        for asiento_path in ci.asiento_pdf_paths:
            norm = normalize_sharepoint_path(asiento_path)
            event = parsed_events.get(norm) or parsed_events.get(asiento_path)
            payment_app = _payment_application_dict(event) if event else None
            error_code = None
            application_status = "BLOCKED"
            if doc_blocked:
                application_status = "ERROR"
                error_code = (
                    doc_blocking_errors[0].get("error_code")
                    if doc_blocking_errors
                    else ABONO_MANIFEST_INVALID
                )

            extracto_path = (ci.extracto_pdf_paths or [None])[0] if ci.extracto_pdf_paths else None
            mora_coverage = _mora_coverage_for_event(ci, event)
            item: dict[str, Any] = {
                "id_pago": group.id_pago,
                "cliente": group.cliente,
                "credito": cred or ci.credito,
                "asiento_pdf_path": asiento_path,
                "extracto_pdf_path": extracto_path,
                "tabla_amortizacion_path": ci.ruta_tabla_amortizacion or None,
                "event_index": idx,
                "payment_application": payment_app,
                "application_row": None,
                "target_row": None,
                "application_status": application_status,
                "error_code": error_code,
                "warnings": list(group.warnings),
                **_abono_observability_base(group=group, schedule=schedule, policy=ci.policy),
                **mora_coverage,
            }
            if event:
                item.update(_abono_payment_date_fields(group, event))
            items.append(item)
    return items
