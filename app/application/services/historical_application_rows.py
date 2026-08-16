"""
Lectura compartida de filas validadas en el histórico (cartera_validada).

Schema v4: única hoja Aplicacion_Pagos. Notify y Merge usan estas funciones
para no divergir en reglas de inclusión.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from app.application.services.review_schema import (
    AplicacionPagosCols,
    ApplicationPolicy,
    ReviewSheets,
    TipoAplicacion,
    TipoAplicacionConfirmado,
    ValidarPago,
    find_aplicacion_pagos_sheet,
    normalize_credito_digits,
    normalize_validar_pago_value,
    resolve_policy_from_tipo_confirmado,
)


def _norm_key(s: str) -> str:
    return " ".join(str(s or "").strip().split())


def _accent_fold_upper(s: str) -> str:
    s = _norm_key(s)
    nfkd = unicodedata.normalize("NFD", s)
    return "".join(c for c in nfkd if unicodedata.category(c) != "Mn").upper()


def _excel_cell_display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _get_col(header_map: dict[str, int], *candidates: str) -> int | None:
    for c in candidates:
        key = _accent_fold_upper(c)
        if key in header_map:
            return header_map[key]
    return None


def _find_aplicacion_header_row(ws: Any) -> tuple[int, dict[str, int]]:
    marker = _accent_fold_upper(AplicacionPagosCols.ID_PAGO)
    marker_vp = _accent_fold_upper(AplicacionPagosCols.VALIDAR_PAGO)
    max_col = ws.max_column or 1
    max_scan = min(ws.max_row or 1, 50)
    for row_idx in range(1, max_scan + 1):
        raw_vals = [ws.cell(row=row_idx, column=c).value for c in range(1, max_col + 1)]
        cells = [_norm_key(str(v)) if v is not None and str(v).strip() else "" for v in raw_vals]
        if not any(cells):
            continue
        folded = [_accent_fold_upper(c) if c else "" for c in cells]
        if marker not in folded:
            continue
        if marker_vp not in folded and _accent_fold_upper("Validar pago") not in folded:
            continue
        header_map: dict[str, int] = {}
        for col, name in enumerate(cells, start=1):
            if name:
                header_map[_accent_fold_upper(name)] = col
        return row_idx, header_map
    raise ValueError("missing_aplicacion_pagos_headers")


def _coerce_historical_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            from openpyxl.utils.datetime import from_excel

            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted.date()
            if isinstance(converted, date):
                return converted
        except (ValueError, OverflowError, OSError):
            return None
        return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _coerce_historical_amount(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text.startswith("="):
        return None
    try:
        return float(text)
    except ValueError:
        try:
            return float(text.replace(".", "").replace(",", "."))
        except ValueError:
            return None


def _policy_fields_from_row(policy: ApplicationPolicy) -> dict[str, Any]:
    pd = policy.policy_dict()
    return {
        "tipo_aplicacion": policy.tipo_aplicacion_canonica,
        "tipo_aplicacion_original": pd["tipo_aplicacion_original"],
        "tipo_aplicacion_canonica": pd["tipo_aplicacion_canonica"],
        "subtipo_aplicacion": pd["subtipo_aplicacion"],
        "requiere_extracto": pd["requiere_extracto"],
        "rol_extracto": pd["rol_extracto"],
        "cierra_cuota": pd["cierra_cuota"],
        "actualiza_ibr": pd["actualiza_ibr"],
        "payoff_expected": pd["payoff_expected"],
        "include_extract_in_composite": pd["include_extract_in_composite"],
    }


def aplicacion_row_included_for_notify_merge(
    ws: Any,
    row: int,
    header_map: dict[str, int],
) -> bool:
    """Incluye filas con Validar Pago = SI y Tipo de aplicación confirmado."""
    col_vp = _get_col(
        header_map,
        AplicacionPagosCols.VALIDAR_PAGO,
        "Validar pago",
        "VALIDAR PAGO",
    )
    if col_vp is None:
        return False
    if normalize_validar_pago_value(ws.cell(row=row, column=col_vp).value) != ValidarPago.SI:
        return False

    col_tipo = _get_col(
        header_map,
        AplicacionPagosCols.TIPO_APLICACION,
        "Tipo de aplicacion",
        "Tipo aplicación",
    )
    if col_tipo is None:
        return False
    tipo_raw = ws.cell(row=row, column=col_tipo).value
    try:
        resolve_policy_from_tipo_confirmado(tipo_raw)
    except ValueError:
        return False
    return True


def read_validated_application_rows(wb: Any) -> list[dict[str, Any]]:
    """Filas SI de Aplicacion_Pagos con política documental v4."""
    ws = find_aplicacion_pagos_sheet(wb)
    h_row, header_map = _find_aplicacion_header_row(ws)

    col_ruta = _get_col(
        header_map,
        "_ruta_extracto",
        "Ruta",
        "RUTA",
        "Rutas",
        "RUTAS",
    )
    col_ruta_asientos = _get_col(
        header_map,
        "_ruta_asientos_contables",
        "RutaAsientosContables",
        "RUTA_ASIENTOS_CONTABLES",
        "Ruta asientos contables",
    )
    col_ruta_uc = _get_col(
        header_map,
        "_ruta_unidad_credito",
        "RutaUnidadCredito",
        "RUTA_UNIDAD_CREDITO",
    )
    col_ruta_tabla = _get_col(
        header_map,
        "_ruta_tabla_amortizacion",
        "RutaTablaAmortizacion",
        "RUTA_TABLA_AMORTIZACION",
    )
    col_cred_norm = _get_col(
        header_map,
        "_credito_normalizado",
        "CreditoNormalizado",
        "CREDITO_NORMALIZADO",
    )
    col_id = _get_col(header_map, AplicacionPagosCols.ID_PAGO, "ID pago", "ID_PAGO")
    col_cliente = _get_col(header_map, AplicacionPagosCols.CLIENTE, "CLIENTE")
    col_credito = _get_col(header_map, AplicacionPagosCols.CREDITO, "Credito", "CREDITO")
    col_monto = _get_col(header_map, AplicacionPagosCols.MONTO_BANCO, "Monto banco", "MONTO BANCO")
    col_fecha = _get_col(header_map, AplicacionPagosCols.FECHA_BANCO, "Fecha banco", "FECHA BANCO")
    col_limite = _get_col(header_map, AplicacionPagosCols.FECHA_LIMITE, "Fecha limite", "FECHA LIMITE")
    col_tipo = _get_col(header_map, AplicacionPagosCols.TIPO_APLICACION, "Tipo de aplicacion")
    col_obs = _get_col(header_map, AplicacionPagosCols.OBSERVACION, "Observacion")
    col_link_ext = _get_col(header_map, AplicacionPagosCols.LINK_EXTRACTO, "Link extracto")

    rows: list[dict[str, Any]] = []
    last = ws.max_row or h_row
    for r in range(h_row + 1, last + 1):
        if not aplicacion_row_included_for_notify_merge(ws, r, header_map):
            continue
        tipo_raw = ws.cell(row=r, column=col_tipo).value if col_tipo else None
        try:
            policy = resolve_policy_from_tipo_confirmado(tipo_raw)
        except ValueError:
            continue
        policy_fields = _policy_fields_from_row(policy)

        if policy.include_extract_in_composite and col_ruta is None and col_link_ext is None:
            # Histórico sin columna de ruta: fail-closed solo si la policy exige extracto en composite.
            raise ValueError("missing_aplicacion_pagos_route_column")

        id_raw = ws.cell(row=r, column=col_id).value if col_id else None
        id_str = _excel_cell_display(id_raw).strip() if col_id else ""
        if not id_str:
            id_str = f"__sin_id_fila_{r}__"
        cred_raw = ws.cell(row=r, column=col_credito).value if col_credito else None
        cred_digits = ""
        if col_cred_norm:
            cred_digits = _excel_cell_display(ws.cell(row=r, column=col_cred_norm).value).strip()
        if not cred_digits:
            cred_digits = normalize_credito_digits(cred_raw) if cred_raw is not None else ""
        if not cred_digits and cred_raw is not None:
            cred_digits = re.sub(r"\D", "", str(cred_raw).strip())

        cliente = ""
        if col_cliente:
            cliente = _excel_cell_display(ws.cell(row=r, column=col_cliente).value).strip()
        monto = _coerce_historical_amount(ws.cell(row=r, column=col_monto).value) if col_monto else None
        fecha = _coerce_historical_date(ws.cell(row=r, column=col_fecha).value) if col_fecha else None
        fecha_lim = (
            _coerce_historical_date(ws.cell(row=r, column=col_limite).value) if col_limite else None
        )
        ruta_cell = ws.cell(row=r, column=col_ruta).value if col_ruta else None
        # No inventar ruta desde Link extracto si la columna técnica existe (aunque vacía).
        if col_ruta is None and col_link_ext is not None:
            ruta_cell = ws.cell(row=r, column=col_link_ext).value

        rows.append(
            {
                **policy_fields,
                "excel_row": r,
                "id_pago": id_str,
                "cliente": cliente,
                "credito_raw": cred_raw,
                "credito_label": _excel_cell_display(cred_raw).strip() if cred_raw is not None else "",
                "credito_digits": cred_digits,
                "monto_banco": monto,
                "fecha_banco": fecha,
                "fecha_limite": fecha_lim,
                "observacion": (
                    _excel_cell_display(ws.cell(row=r, column=col_obs).value).strip()
                    if col_obs
                    else ""
                ),
                "ruta_cell": ruta_cell,
                "ruta_asientos_cell": (
                    ws.cell(row=r, column=col_ruta_asientos).value if col_ruta_asientos else None
                ),
                "ruta_unidad_credito": (
                    _excel_cell_display(ws.cell(row=r, column=col_ruta_uc).value).strip()
                    if col_ruta_uc
                    else ""
                ),
                "ruta_tabla_amortizacion": (
                    _excel_cell_display(ws.cell(row=r, column=col_ruta_tabla).value).strip()
                    if col_ruta_tabla
                    else ""
                ),
                "sheet": ReviewSheets.APLICACION_PAGOS,
            }
        )
    return rows


def read_validated_payment_rows(
    wb: Any,
    *,
    legacy_estado_token: str = "",
) -> list[dict[str, Any]]:
    """Compat: filas canónicas PAGO de Aplicacion_Pagos (ignora legacy_estado_token)."""
    _ = legacy_estado_token
    return [
        r
        for r in read_validated_application_rows(wb)
        if r.get("tipo_aplicacion") == TipoAplicacion.PAGO.value
    ]


def read_validated_abono_rows(wb: Any) -> list[dict[str, Any]]:
    """Compat: filas canónicas ABONO de Aplicacion_Pagos (sin hoja Distribucion_Abonos)."""
    return [
        r
        for r in read_validated_application_rows(wb)
        if r.get("tipo_aplicacion") == TipoAplicacion.ABONO.value
    ]


@dataclass(frozen=True)
class AbonoEmailGroup:
    id_pago: str
    cliente: str
    monto_banco: float | None
    fecha_banco: date | None
    creditos_seleccionados: tuple[str, ...]
    tipo_aplicacion_original: str = TipoAplicacionConfirmado.ABONO_A_CAPITAL


def group_abono_rows_for_email(abono_rows: list[dict[str, Any]]) -> list[AbonoEmailGroup]:
    by_id: dict[str, list[dict[str, Any]]] = {}
    for row in abono_rows:
        id_pago = str(row.get("id_pago") or "").strip()
        if not id_pago:
            continue
        by_id.setdefault(id_pago, []).append(row)

    groups: list[AbonoEmailGroup] = []
    for id_pago in sorted(by_id.keys()):
        members = by_id[id_pago]
        ref = members[0]
        cliente = str(ref.get("cliente") or "").strip()
        monto = ref.get("monto_banco")
        fecha = ref.get("fecha_banco")
        ref_tipo_original = str(
            ref.get("tipo_aplicacion_original") or TipoAplicacionConfirmado.ABONO_A_CAPITAL
        ).strip()
        for row in members[1:]:
            if str(row.get("cliente") or "").strip() != cliente:
                raise ValueError(f"abono_group_inconsistent|id_pago={id_pago}")
            if row.get("monto_banco") != monto:
                raise ValueError(f"abono_group_inconsistent|id_pago={id_pago}")
            if row.get("fecha_banco") != fecha:
                raise ValueError(f"abono_group_inconsistent|id_pago={id_pago}")
            if str(row.get("tipo_aplicacion_original") or "").strip() != ref_tipo_original:
                raise ValueError(f"abono_group_inconsistent|id_pago={id_pago}")
        creditos: list[str] = []
        seen: set[str] = set()
        for row in sorted(members, key=lambda x: str(x.get("credito_digits") or x.get("credito_label") or "")):
            c = str(row.get("credito_digits") or row.get("credito_label") or "").strip()
            if c and c not in seen:
                seen.add(c)
                creditos.append(c)
        groups.append(
            AbonoEmailGroup(
                id_pago=id_pago,
                cliente=cliente,
                monto_banco=monto if isinstance(monto, (int, float)) else None,
                fecha_banco=fecha if isinstance(fecha, date) else None,
                creditos_seleccionados=tuple(creditos),
                tipo_aplicacion_original=ref_tipo_original,
            )
        )
    return groups


def group_rows_by_id_pago(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    groups: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        id_pago = str(row.get("id_pago") or "").strip() or f"__sin_id_fila_{row.get('excel_row')}__"
        groups.setdefault(id_pago, []).append(row)
    return groups


def detect_application_type_group_conflict(
    payment_groups: dict[str, list[dict[str, Any]]],
    abono_groups: dict[str, list[dict[str, Any]]],
) -> None:
    """Legacy no-op check; v3 unifica por ID Pago (tipos distintos → APLICACION MULTIPLE)."""
    _ = (payment_groups, abono_groups)


# Aliases legacy usados por tests/callers aún no migrados.
abono_row_included_for_notify_merge = aplicacion_row_included_for_notify_merge
