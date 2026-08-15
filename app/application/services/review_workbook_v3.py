"""Builder del workbook de revisión schema v3 (columnas canónicas + presentación operador)."""
from __future__ import annotations

import io
from datetime import date
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.application.services.extract_snapshot_parser import (
    EVIDENCE_HEADERS_LABEL,
    EVIDENCE_META_FIELDS,
    EVIDENCE_ROW_PREFIX,
    serialize_evidence_meta_value,
)
from app.application.services.review_schema import (
    REVIEW_SCHEMA_VERSION,
    AplicacionPagosCols,
    AplicacionSugerida,
    ErroresCols,
    InternalPathCols,
    MetaCols,
    ReviewSheets,
    TipoAplicacionConfirmado,
    ValidarPago,
    compute_aplicacion_sugerida,
    dias_respecto_vencimiento,
)


# Mismo criterio visual que el soporte de asientos / control de merge.
_FILL_NAVY = PatternFill(fill_type="solid", fgColor="002060")
_HEADER_FILL = _FILL_NAVY
_HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_FONT_TITLE = Font(name="Calibri", bold=True, size=16, color="FFFFFF")
_FONT_INSTRUCTION = Font(name="Calibri", size=12, color="1A2F36")
_FONT_BODY = Font(name="Calibri", size=11)
_FONT_HLINK = Font(name="Calibri", color="0563C1", size=11, underline="single")
_FILL_INSTRUCTION = PatternFill(fill_type="solid", fgColor="E8F2FA")
_FILL_HLINK = PatternFill(fill_type="solid", fgColor="E8F4FC")
_FILL_EDITABLE = PatternFill(fill_type="solid", fgColor="E2F4E8")
_FILL_ZEBRA_A = PatternFill(fill_type="solid", fgColor="FCFCFD")
_FILL_ZEBRA_B = PatternFill(fill_type="solid", fgColor="F6F8FA")
_AMBIGUOUS_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
_THIN = Side(style="thin", color="C8C8C8")
_BORDER_LIGHT = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_ALIGN_TITLE = Alignment(vertical="center", horizontal="center", wrap_text=True)
_ALIGN_WRAP = Alignment(vertical="center", horizontal="left", wrap_text=True)
_ALIGN_CENTER = Alignment(vertical="center", horizontal="center", wrap_text=True)
_LOCKED = Protection(locked=True)
_UNLOCKED = Protection(locked=False)
_FMT_MONEY = "#,##0.00"
_FMT_DATE = "yyyy-mm-dd"
_TAB_APLICACION = "00B050"
_TAB_ERRORES = "FF6969"

APLICACION_TITLE = "APLICACIÓN DE PAGOS"
APLICACION_HELP = (
    "Complete únicamente las celdas editables (fondo verde): Validar Pago, montos a aplicar "
    "y Tipo de aplicación. Revise los links si necesita validar documentos. "
    "Al terminar, en Control cambie Procesar a SI."
)
ERRORES_TITLE = "REGISTRO DE ERRORES"
ERRORES_HELP = (
    "Revise estos casos manualmente. Use la descripción, la acción recomendada y los links "
    "para corregir documentos o carpetas antes de volver a generar."
)

# Filas 1–2: banner; fila 3: encabezados. Finalize localiza la fila de headers.
REVIEW_HEADER_ROW = 3
REVIEW_FIRST_DATA_ROW = 4


def _aplicacion_sugerida_excel_formula(
    *,
    row: int,
    col_vp: int,
    col_a: int,
    col_v: int,
    col_k: int,
    col_oblig: int,
) -> str:
    """
    Fórmula dinámica que replica compute_aplicacion_sugerida en Excel.
    Se recalcula al editar Validar Pago / A / V / K / valor obligación.
    """
    vp = f"{get_column_letter(col_vp)}{row}"
    a = f"{get_column_letter(col_a)}{row}"
    v = f"{get_column_letter(col_v)}{row}"
    k = f"{get_column_letter(col_k)}{row}"
    oblig = f"{get_column_letter(col_oblig)}{row}"
    # Nombres exactos del mandato (TipoAplicacionConfirmado / AplicacionSugerida).
    return (
        f'IF(OR({vp}="POR DEFINIR",{vp}=""),"POR DEFINIR",'
        f'IF({vp}="NO","NO APLICA",'
        f'IF(AND(N({a})=0,N({v})=0,N({k})=0),"POR DISTRIBUIR",'
        f'IF(AND(N({a})>0,N({v})=0,N({k})=0),'
        f'IF(AND({oblig}<>"",N({a})<N({oblig})),'
        f'"PAGO PARCIAL A OBLIGACIÓN ACTUAL","PAGO DE OBLIGACIÓN ACTUAL"),'
        f'IF(AND(N({a})=0,N({v})>0,N({k})=0),"APLICACIÓN A SALDO VENCIDO",'
        f'IF(AND(N({a})=0,N({v})=0,N({k})>0),"ABONO A CAPITAL",'
        f'IF(AND(N({a})>0,N({v})>0,N({k})=0),'
        f'"PAGO COMBINADO (SALDO VENCIDO + OBLIGACIÓN ACTUAL)",'
        f'IF(AND(N({a})>0,N({v})=0,N({k})>0),"PAGO Y ABONO A CAPITAL",'
        f'IF(AND(N({a})=0,N({v})>0,N({k})>0),'
        f'"APLICACIÓN A SALDO VENCIDO + ABONO A CAPITAL",'
        f'IF(AND(N({a})>0,N({v})>0,N({k})>0),"PAGO COMBINADO + ABONO A CAPITAL",'
        f'"POR DISTRIBUIR"))))))))))'
    )


def _errores_row(rec: dict[str, Any]) -> list[Any]:
    values = {
        ErroresCols.ID_PAGO: rec.get("id_pago"),
        ErroresCols.CLIENTE: rec.get("cliente"),
        ErroresCols.CREDITO: rec.get("credito"),
        ErroresCols.TIPO_CASO: rec.get("tipo_caso") or "",
        ErroresCols.DESCRIPCION: rec.get("descripcion") or rec.get("message") or "",
        ErroresCols.QUE_DEBE_HACER: rec.get("que_debe_hacer") or "",
        ErroresCols.REQUIERE_SOPORTE: rec.get("requiere_soporte") or "",
        ErroresCols.LINK_EXTRACTO: rec.get("link_extracto") or "",
        ErroresCols.LINK_CARPETA_CREDITO: rec.get("link_carpeta_credito") or "",
        ErroresCols.CODIGO_TECNICO: rec.get("codigo") or rec.get("code") or "",
    }
    return [values[c] for c in ErroresCols.HEADERS]


def _apply_banner(ws: Any, *, title: str, help_text: str, ncols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    t = ws.cell(1, 1, title)
    t.font = _FONT_TITLE
    t.fill = _FILL_NAVY
    t.alignment = _ALIGN_TITLE
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    s = ws.cell(2, 1, help_text)
    s.fill = _FILL_INSTRUCTION
    s.font = _FONT_INSTRUCTION
    s.alignment = _ALIGN_WRAP
    ws.row_dimensions[2].height = 48


def _style_header_row(ws: Any, row_idx: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row_idx, col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _ALIGN_CENTER
        cell.border = _BORDER_LIGHT
    ws.row_dimensions[row_idx].height = 28
    ws.auto_filter.ref = f"A{row_idx}:{get_column_letter(ncols)}{row_idx}"
    ws.freeze_panes = f"A{row_idx + 1}"
    ws.sheet_view.showGridLines = False


def _looks_like_url(value: Any) -> bool:
    text = str(value or "").strip()
    return text.lower().startswith("http://") or text.lower().startswith("https://")


def _apply_aplicacion_body_style(
    ws: Any,
    *,
    first_data: int,
    last_data: int,
    ncols: int,
    aplicacion_rows: list[dict[str, Any]],
) -> None:
    money_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1
        for c in (
            AplicacionPagosCols.MONTO_BANCO,
            AplicacionPagosCols.VALOR_OBLIGACION_ACTUAL,
            AplicacionPagosCols.SALDO_VENCIDO,
            AplicacionPagosCols.APLICAR_OBLIGACION_ACTUAL,
            AplicacionPagosCols.APLICAR_SALDO_VENCIDO,
            AplicacionPagosCols.ABONO_ADICIONAL_CAPITAL,
            AplicacionPagosCols.TOTAL_ASIGNADO,
            AplicacionPagosCols.SALDO_POR_ASIGNAR,
        )
    }
    date_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1
        for c in (AplicacionPagosCols.FECHA_BANCO, AplicacionPagosCols.FECHA_LIMITE)
    }
    link_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1
        for c in (
            AplicacionPagosCols.LINK_EXTRACTO,
            AplicacionPagosCols.LINK_TABLA,
            AplicacionPagosCols.LINK_CARPETA_CREDITO,
        )
    }
    wrap_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1
        for c in (AplicacionPagosCols.OBSERVACION, AplicacionPagosCols.APLICACION_SUGERIDA)
    }
    editable_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1 for c in AplicacionPagosCols.SECRETARY_EDITABLE
    }
    col_id = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.ID_PAGO) + 1
    seen: dict[str, int] = {}
    zebra_bit = 0
    for r in range(first_data, last_data + 1):
        pid = str(ws.cell(r, col_id).value or "").strip()
        if pid not in seen:
            seen[pid] = zebra_bit
            zebra_bit ^= 1
        zebra = _FILL_ZEBRA_A if seen[pid] == 0 else _FILL_ZEBRA_B
        row_off = r - first_data
        ambiguous = False
        if 0 <= row_off < len(aplicacion_rows):
            ambiguous = str(aplicacion_rows[row_off].get("_right_panel_role") or "") == "AMBIGUO"
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.border = _BORDER_LIGHT
            cell.alignment = _ALIGN_WRAP if c in wrap_idx else Alignment(vertical="center")
            if ambiguous:
                cell.fill = _AMBIGUOUS_FILL
            elif c in editable_idx:
                cell.fill = _FILL_EDITABLE
            else:
                cell.fill = zebra
            if c in money_idx:
                cell.number_format = _FMT_MONEY
            if c in date_idx:
                cell.number_format = _FMT_DATE
            if c in link_idx and _looks_like_url(cell.value):
                cell.font = _FONT_HLINK
                if not ambiguous:
                    cell.fill = _FILL_HLINK
                cell.hyperlink = str(cell.value).strip()
            elif c not in link_idx:
                cell.font = _FONT_BODY
        ws.row_dimensions[r].height = 20

    widths = {
        AplicacionPagosCols.ID_PAGO: 14,
        AplicacionPagosCols.CLIENTE: 28,
        AplicacionPagosCols.CREDITO: 12,
        AplicacionPagosCols.MONTO_BANCO: 16,
        AplicacionPagosCols.FECHA_BANCO: 14,
        AplicacionPagosCols.FECHA_LIMITE: 14,
        AplicacionPagosCols.DIAS_RESPECTO_VENCIMIENTO: 14,
        AplicacionPagosCols.VALOR_OBLIGACION_ACTUAL: 18,
        AplicacionPagosCols.SALDO_VENCIDO: 16,
        AplicacionPagosCols.VALIDAR_PAGO: 16,
        AplicacionPagosCols.APLICAR_OBLIGACION_ACTUAL: 18,
        AplicacionPagosCols.APLICAR_SALDO_VENCIDO: 18,
        AplicacionPagosCols.ABONO_ADICIONAL_CAPITAL: 18,
        AplicacionPagosCols.TOTAL_ASIGNADO: 18,
        AplicacionPagosCols.SALDO_POR_ASIGNAR: 16,
        AplicacionPagosCols.APLICACION_SUGERIDA: 36,
        AplicacionPagosCols.TIPO_APLICACION: 28,
        AplicacionPagosCols.LINK_EXTRACTO: 22,
        AplicacionPagosCols.LINK_TABLA: 22,
        AplicacionPagosCols.LINK_CARPETA_CREDITO: 22,
        AplicacionPagosCols.OBSERVACION: 40,
    }
    for idx, name in enumerate(AplicacionPagosCols.HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 14)
    freeze_col = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.CREDITO) + 2
    ws.freeze_panes = f"{get_column_letter(freeze_col)}{first_data}"
    if last_data >= first_data:
        ws.auto_filter.ref = (
            f"A{REVIEW_HEADER_ROW}:{get_column_letter(ncols)}{last_data}"
        )


def _apply_errores_body_style(ws: Any, *, first_data: int, last_data: int, ncols: int) -> None:
    link_names = {ErroresCols.LINK_EXTRACTO, ErroresCols.LINK_CARPETA_CREDITO}
    wrap_names = {ErroresCols.DESCRIPCION, ErroresCols.QUE_DEBE_HACER}
    link_idx = {ErroresCols.HEADERS.index(n) + 1 for n in link_names}
    wrap_idx = {ErroresCols.HEADERS.index(n) + 1 for n in wrap_names}
    for r in range(first_data, last_data + 1):
        zebra = _FILL_ZEBRA_A if (r - first_data) % 2 == 0 else _FILL_ZEBRA_B
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.border = _BORDER_LIGHT
            cell.fill = zebra
            cell.alignment = _ALIGN_WRAP if c in wrap_idx else Alignment(vertical="center")
            if c in link_idx and _looks_like_url(cell.value):
                cell.font = _FONT_HLINK
                cell.fill = _FILL_HLINK
                cell.hyperlink = str(cell.value).strip()
            else:
                cell.font = _FONT_BODY
        ws.row_dimensions[r].height = 22
    widths = {
        ErroresCols.ID_PAGO: 14,
        ErroresCols.CLIENTE: 28,
        ErroresCols.CREDITO: 12,
        ErroresCols.TIPO_CASO: 18,
        ErroresCols.DESCRIPCION: 42,
        ErroresCols.QUE_DEBE_HACER: 42,
        ErroresCols.REQUIERE_SOPORTE: 16,
        ErroresCols.LINK_EXTRACTO: 22,
        ErroresCols.LINK_CARPETA_CREDITO: 22,
        ErroresCols.CODIGO_TECNICO: 28,
    }
    for idx, name in enumerate(ErroresCols.HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 14)
    if last_data >= first_data:
        ws.auto_filter.ref = (
            f"A{REVIEW_HEADER_ROW}:{get_column_letter(ncols)}{last_data}"
        )


def build_aplicacion_pagos_row(
    payment: dict[str, Any],
    candidate: dict[str, Any],
) -> dict[str, Any]:
    """Fila neutra: Validar Pago = POR DEFINIR; sin autoselección SI."""
    due = candidate.get("fecha_limite")
    fecha_banco = payment["fecha_banco"]
    dias = dias_respecto_vencimiento(fecha_banco, due)

    valor_oblig = candidate.get("valor_obligacion_actual")
    if valor_oblig is None:
        valor_oblig = candidate.get("valor_extracto")

    saldo_vis = candidate.get("saldo_vencido_visible")
    if "saldo_vencido_visible" not in candidate:
        # Fail-closed: sin rol explícito no inventar saldo.
        saldo_vis = None

    row = {
        AplicacionPagosCols.ID_PAGO: payment["id_pago"],
        AplicacionPagosCols.CLIENTE: payment["cliente"],
        AplicacionPagosCols.CREDITO: candidate.get("credito"),
        AplicacionPagosCols.MONTO_BANCO: payment["monto_banco"],
        AplicacionPagosCols.FECHA_BANCO: fecha_banco.isoformat()
        if hasattr(fecha_banco, "isoformat")
        else fecha_banco,
        AplicacionPagosCols.FECHA_LIMITE: due.isoformat() if hasattr(due, "isoformat") else (due or ""),
        AplicacionPagosCols.DIAS_RESPECTO_VENCIMIENTO: dias if dias is not None else "",
        AplicacionPagosCols.VALOR_OBLIGACION_ACTUAL: valor_oblig if valor_oblig is not None else "",
        AplicacionPagosCols.SALDO_VENCIDO: saldo_vis if saldo_vis is not None else "",
        AplicacionPagosCols.VALIDAR_PAGO: ValidarPago.POR_DEFINIR,
        AplicacionPagosCols.APLICAR_OBLIGACION_ACTUAL: "",
        AplicacionPagosCols.APLICAR_SALDO_VENCIDO: "",
        AplicacionPagosCols.ABONO_ADICIONAL_CAPITAL: "",
        AplicacionPagosCols.TOTAL_ASIGNADO: "",
        AplicacionPagosCols.SALDO_POR_ASIGNAR: "",
        AplicacionPagosCols.APLICACION_SUGERIDA: AplicacionSugerida.POR_DEFINIR,
        AplicacionPagosCols.TIPO_APLICACION: "",
        AplicacionPagosCols.LINK_EXTRACTO: candidate.get("link_extracto", ""),
        AplicacionPagosCols.LINK_TABLA: candidate.get("link_tabla", ""),
        AplicacionPagosCols.LINK_CARPETA_CREDITO: candidate.get("link_carpeta_credito", ""),
        AplicacionPagosCols.OBSERVACION: candidate.get("observacion_extra") or "",
    }
    # Evidencia congelada (vive en _Meta / modelos; no columnas visibles).
    row["_evidence"] = candidate.get("extract_evidence") or {}
    row["_right_panel_role"] = candidate.get("right_panel_role") or ""
    row["_parser_status"] = candidate.get("parser_status") or ""
    row["_ruta_extracto"] = candidate.get("ruta_extracto_pdf") or ""
    row["_ruta_unidad_credito"] = candidate.get("ruta_unidad_credito") or ""
    row["_ruta_tabla_amortizacion"] = candidate.get("ruta_tabla_amortizacion") or ""
    row["_credito_normalizado"] = candidate.get("credito_normalizado") or ""
    return row


def build_review_workbook_v3_bytes(
    *,
    process_id: str,
    process_date: date,
    bank_code: str,
    aplicacion_rows: list[dict[str, Any]],
    error_records: list[dict[str, Any]],
) -> bytes:
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = ReviewSheets.APLICACION_PAGOS
    path_headers = [
        InternalPathCols.RUTA_EXTRACTO,
        InternalPathCols.RUTA_UNIDAD_CREDITO,
        InternalPathCols.RUTA_TABLA_AMORTIZACION,
        InternalPathCols.CREDITO_NORMALIZADO,
    ]
    visible_n = len(AplicacionPagosCols.HEADERS)
    _apply_banner(ws, title=APLICACION_TITLE, help_text=APLICACION_HELP, ncols=visible_n)
    ws.append(list(AplicacionPagosCols.HEADERS) + path_headers)
    _style_header_row(ws, REVIEW_HEADER_ROW, visible_n)

    # Agrupar por ID Pago para Saldo por asignar solo en primera fila SI/grupo.
    seen_monto: set[str] = set()
    for row in aplicacion_rows:
        pid = str(row.get(AplicacionPagosCols.ID_PAGO) or "").strip()
        values = []
        for h in AplicacionPagosCols.HEADERS:
            val = row.get(h, "")
            if h == AplicacionPagosCols.MONTO_BANCO and pid:
                if pid in seen_monto:
                    val = None
                else:
                    seen_monto.add(pid)
            if h == AplicacionPagosCols.APLICACION_SUGERIDA:
                val = compute_aplicacion_sugerida(
                    validar_pago=row.get(AplicacionPagosCols.VALIDAR_PAGO),
                    aplicar_obligacion=row.get(AplicacionPagosCols.APLICAR_OBLIGACION_ACTUAL),
                    aplicar_saldo_vencido=row.get(AplicacionPagosCols.APLICAR_SALDO_VENCIDO),
                    abono_capital=row.get(AplicacionPagosCols.ABONO_ADICIONAL_CAPITAL),
                    valor_obligacion_actual=row.get(AplicacionPagosCols.VALOR_OBLIGACION_ACTUAL),
                )
            values.append(val)
        for ph in path_headers:
            values.append(row.get(ph, "") or "")
        ws.append(values)

    # Ocultar columnas técnicas de ruta (tras las 21 visibles).
    for offset in range(len(path_headers)):
        letter = get_column_letter(len(AplicacionPagosCols.HEADERS) + 1 + offset)
        ws.column_dimensions[letter].hidden = True

    first_data = REVIEW_FIRST_DATA_ROW
    last_data = first_data + len(aplicacion_rows) - 1 if aplicacion_rows else first_data - 1
    # Fórmulas Total asignado / Saldo por asignar / Aplicación sugerida (dinámicas).
    col_total = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.TOTAL_ASIGNADO) + 1
    col_a = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.APLICAR_OBLIGACION_ACTUAL) + 1
    col_v = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.APLICAR_SALDO_VENCIDO) + 1
    col_k = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.ABONO_ADICIONAL_CAPITAL) + 1
    col_saldo = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.SALDO_POR_ASIGNAR) + 1
    col_monto = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.MONTO_BANCO) + 1
    col_id = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.ID_PAGO) + 1
    col_vp = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.VALIDAR_PAGO) + 1
    col_sug = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.APLICACION_SUGERIDA) + 1
    col_oblig = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.VALOR_OBLIGACION_ACTUAL) + 1
    col_saldo_venc = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.SALDO_VENCIDO) + 1

    formula_last = max(last_data, first_data) if aplicacion_rows else first_data
    for r in range(first_data, last_data + 1):
        ws.cell(r, col_total).value = (
            f"={get_column_letter(col_a)}{r}+{get_column_letter(col_v)}{r}+{get_column_letter(col_k)}{r}"
        )
        # Saldo por asignar: Monto banco − SUMIFS(Total, ID, id, ValidarPago, "SI").
        id_cell = f"{get_column_letter(col_id)}{r}"
        id_range = f"{get_column_letter(col_id)}${first_data}:{get_column_letter(col_id)}${formula_last}"
        vp_range = f"{get_column_letter(col_vp)}${first_data}:{get_column_letter(col_vp)}${formula_last}"
        tot_range = (
            f"{get_column_letter(col_total)}${first_data}:{get_column_letter(col_total)}${formula_last}"
        )
        ws.cell(r, col_saldo).value = (
            f"=IF({get_column_letter(col_monto)}{r}=\"\",\"\","
            f"{get_column_letter(col_monto)}{r}-SUMIFS({tot_range},{id_range},{id_cell},{vp_range},\"SI\"))"
        )
        ws.cell(r, col_sug).value = (
            "="
            + _aplicacion_sugerida_excel_formula(
                row=r,
                col_vp=col_vp,
                col_a=col_a,
                col_v=col_v,
                col_k=col_k,
                col_oblig=col_oblig,
            )
        )
        # Marca visual AMBIGUO: no asume saldo vencido = 0.
        role = str(aplicacion_rows[r - first_data].get("_right_panel_role") or "")
        if role == "AMBIGUO":
            for c in range(1, len(AplicacionPagosCols.HEADERS) + 1):
                ws.cell(r, c).fill = _AMBIGUOUS_FILL
            cell_sv = ws.cell(r, col_saldo_venc)
            if cell_sv.comment is None:
                from openpyxl.comments import Comment

                cell_sv.comment = Comment(
                    "Panel derecho AMBIGUO: Saldo vencido vacío a propósito. "
                    "Revisar extracto; no se asume 0.",
                    "HBI",
                )

    # Listas + dropdowns
    ws_lists = wb.create_sheet(ReviewSheets.LISTAS)
    ws_lists.append(["ValidarPago"])
    for i, opt in enumerate(ValidarPago.OPTIONS_ORDERED, start=2):
        ws_lists.cell(i, 1, opt)
    ws_lists.cell(1, 2, "TipoAplicacion")
    for i, opt in enumerate(TipoAplicacionConfirmado.OPTIONS_ORDERED, start=2):
        ws_lists.cell(i, 2, opt)
    ws_lists.sheet_state = "hidden"

    dv_vp = DataValidation(
        type="list",
        formula1=f"={ReviewSheets.LISTAS}!$A$2:$A${1 + len(ValidarPago.OPTIONS_ORDERED)}",
        allow_blank=False,
    )
    ws.add_data_validation(dv_vp)
    dv_vp.add(
        f"{get_column_letter(col_vp)}{first_data}:{get_column_letter(col_vp)}{max(last_data, first_data + 50)}"
    )

    col_tipo = AplicacionPagosCols.HEADERS.index(AplicacionPagosCols.TIPO_APLICACION) + 1
    dv_tipo = DataValidation(
        type="list",
        formula1=f"={ReviewSheets.LISTAS}!$B$2:$B${1 + len(TipoAplicacionConfirmado.OPTIONS_ORDERED)}",
        allow_blank=True,
    )
    ws.add_data_validation(dv_tipo)
    dv_tipo.add(
        f"{get_column_letter(col_tipo)}{first_data}:{get_column_letter(col_tipo)}{max(last_data, first_data + 50)}"
    )

    # Protección: desbloquear editables
    editable_idx = {
        AplicacionPagosCols.HEADERS.index(c) + 1 for c in AplicacionPagosCols.SECRETARY_EDITABLE
    }
    for r in range(first_data, last_data + 1):
        for c in range(1, len(AplicacionPagosCols.HEADERS) + 1):
            cell = ws.cell(r, c)
            cell.protection = _UNLOCKED if c in editable_idx else _LOCKED
    ws.protection.sheet = True
    ws.sheet_properties.tabColor = Color(rgb=_TAB_APLICACION)
    _apply_aplicacion_body_style(
        ws,
        first_data=first_data,
        last_data=last_data,
        ncols=visible_n,
        aplicacion_rows=aplicacion_rows,
    )

    # Errores
    ws_err = wb.create_sheet(ReviewSheets.ERRORES)
    err_n = len(ErroresCols.HEADERS)
    _apply_banner(ws_err, title=ERRORES_TITLE, help_text=ERRORES_HELP, ncols=err_n)
    ws_err.append(list(ErroresCols.HEADERS))
    _style_header_row(ws_err, REVIEW_HEADER_ROW, err_n)
    for rec in error_records:
        ws_err.append(_errores_row(rec))
    if error_records:
        err_first = REVIEW_FIRST_DATA_ROW
        err_last = err_first + len(error_records) - 1
        _apply_errores_body_style(
            ws_err, first_data=err_first, last_data=err_last, ncols=err_n
        )
        ws_err.sheet_properties.tabColor = Color(rgb=_TAB_ERRORES)
    else:
        ws_err.sheet_state = "hidden"

    # Meta + evidencia congelada
    ws_meta = wb.create_sheet(ReviewSheets.META)
    ws_meta.append([MetaCols.CAMPO, MetaCols.VALOR])
    ws_meta.append([MetaCols.ROW_REVIEW_SCHEMA_VERSION, REVIEW_SCHEMA_VERSION])
    ws_meta.append([MetaCols.ROW_PROCESS_ID, process_id])
    ws_meta.append([MetaCols.ROW_PROCESS_DATE, process_date.isoformat()])
    ws_meta.append([MetaCols.ROW_BANK_CODE, bank_code])
    ws_meta.append([EVIDENCE_HEADERS_LABEL, "|".join(EVIDENCE_META_FIELDS)])
    for idx, row in enumerate(aplicacion_rows, start=first_data):
        ev = row.get("_evidence") or {}
        ws_meta.append(
            [
                f"{EVIDENCE_ROW_PREFIX}{idx}",
                serialize_evidence_meta_value(
                    row_idx=idx,
                    evidence=ev,
                    right_panel_role=str(row.get("_right_panel_role") or ""),
                    parser_status=str(row.get("_parser_status") or ""),
                ),
            ]
        )
    ws_meta.sheet_state = "hidden"

    # Sin hojas legacy
    forbidden = {
        "Control",
        "Resumen",
        "Casos_Pago",
        "Distribucion_Pagos",
        "Distribucion_Abonos",
        "Distribucion",
    }
    for name in list(wb.sheetnames):
        if name in forbidden:
            del wb[name]

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
