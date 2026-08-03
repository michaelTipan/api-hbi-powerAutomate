"""
Validaciones de solo lectura sobre el Excel de revisión antes de Finalize.

Reutiliza reglas de negocio ya aplicadas en finalize/generate sin escribir en SharePoint.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl

from app.application.job_status_enrichment import resolve_business_error_messages
from app.application.services.excel_write_guard import ExcelWriteGuardError
from app.application.services.review_schema import (
    CasosPagoCols,
    ControlCols,
    DistribucionCols,
    ErroresCols,
    ReviewSheets,
    find_distribucion_pagos_sheet,
    is_validar_pago_si,
)


def _find_table_header_row(ws: Any, first_header_value: str) -> int:
    marker = str(first_header_value).strip()
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and str(row[0]).strip() == marker:
            return r_idx
    raise ValueError("missing_sheet_headers")


def count_open_errores_rows(wb: Any) -> int:
    """Cuenta filas abiertas en hoja Errores (0 si no existe)."""
    if ReviewSheets.ERRORES not in getattr(wb, "sheetnames", []):
        return 0
    ws = wb[ReviewSheets.ERRORES]
    try:
        header_row = _find_table_header_row(ws, ErroresCols.ID_PAGO)
    except ValueError:
        return 0

    headers: list[str] = []
    open_count = 0
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r_idx < header_row:
            continue
        if r_idx == header_row:
            headers = [str(v).strip() if v is not None else "" for v in row]
            continue
        if not any(v is not None and str(v).strip() for v in row):
            continue
        rd = dict(zip(headers, row))
        id_pago = str(rd.get(ErroresCols.ID_PAGO) or "").strip()
        descripcion = str(rd.get(ErroresCols.DESCRIPCION) or "").strip()
        codigo = str(rd.get(ErroresCols.CODIGO_TECNICO) or "").strip()
        if id_pago or descripcion or codigo:
            open_count += 1
    return open_count


def _issue_from_code(
    job_type: str,
    code: str,
    *,
    stage: str = "review",
    severity: str = "business",
    **hints: Any,
) -> dict[str, Any]:
    user_message, next_action = resolve_business_error_messages(
        job_type,
        code,
        full_message=code,
    )
    row: dict[str, Any] = {
        "stage": stage,
        "severity": severity,
        "error_code": code,
        "user_message": user_message,
        "next_action": next_action,
    }
    for key in ("excel_row", "field", "payment_id", "client_name", "credit", "open_count"):
        if hints.get(key) is not None:
            row[key] = hints[key]
    return row


def collect_review_workbook_preflight_issues(rev_bytes: bytes) -> list[dict[str, Any]]:
    """
    Barrera previa a encolar Finalize: PROCESAR, Errores, control y distribución básica.
    """
    issues: list[dict[str, Any]] = []
    wb = openpyxl.load_workbook(io.BytesIO(rev_bytes), data_only=True)

    if ReviewSheets.CONTROL not in wb.sheetnames:
        issues.append(_issue_from_code("finalize", "missing_control_sheet"))
        return issues

    ws_ctrl = wb[ReviewSheets.CONTROL]
    procesar: str | None = None
    estado_ctrl: str | None = None
    for row in ws_ctrl.iter_rows(values_only=True):
        if row and row[0] == ControlCols.ROW_PROCESAR:
            procesar = str(row[1]).strip().upper() if row[1] else ""
        if row and row[0] == ControlCols.ROW_ESTADO:
            estado_ctrl = str(row[1]).strip().upper() if row[1] else ""

    if procesar != ControlCols.VAL_PROCESAR_SI:
        issues.append(_issue_from_code("finalize", "process_not_approved"))

    if not estado_ctrl:
        issues.append(_issue_from_code("finalize", "missing_control_state"))
    elif estado_ctrl != "EN_REVISION":
        issues.append(_issue_from_code("finalize", "invalid_control_state"))

    open_errores = count_open_errores_rows(wb)
    if open_errores > 0:
        user_message, next_action = resolve_business_error_messages(
            "finalize",
            "review_has_open_errors",
            full_message=f"review_has_open_errors|{open_errores}",
        )
        issues.append(
            {
                "stage": "review",
                "severity": "business",
                "error_code": "review_has_open_errors",
                "user_message": user_message,
                "next_action": next_action,
                "open_count": open_errores,
            }
        )

    try:
        ws_dist = find_distribucion_pagos_sheet(wb)
    except ValueError:
        issues.append(_issue_from_code("finalize", "missing_distribucion_sheet"))
        return issues

    try:
        header_row = _find_table_header_row(ws_dist, DistribucionCols.ID_PAGO)
    except ValueError:
        issues.append(_issue_from_code("finalize", "missing_sheet_headers"))
        return issues

    headers: list[str] = []
    distrib_rows: list[dict[str, Any]] = []
    for r_idx, row in enumerate(ws_dist.iter_rows(values_only=True), start=1):
        if r_idx < header_row:
            continue
        if r_idx == header_row:
            headers = [str(v).strip() if v else "" for v in row]
            continue
        if not any(row):
            continue
        row_dict = dict(zip(headers, row))
        row_dict["_excel_row"] = r_idx
        distrib_rows.append(row_dict)

    editable_cols = (
        DistribucionCols.ESTADO_PAGO,
        DistribucionCols.APLICAR_A_EXTRACTO,
        DistribucionCols.MORA_A_APLICAR,
        DistribucionCols.ABONO_A_CAPITAL,
        DistribucionCols.OTROS_VALORES,
    )
    for dist in distrib_rows:
        if not is_validar_pago_si(dist):
            continue
        row_no = dist.get("_excel_row", "?")
        estado = str(dist.get(DistribucionCols.ESTADO_PAGO) or "").strip()
        if not estado:
            issues.append(
                _issue_from_code(
                    "finalize",
                    "empty_estado_pago",
                    excel_row=row_no,
                )
            )
            continue
        for col in editable_cols:
            if col == DistribucionCols.ESTADO_PAGO:
                continue
            val = dist.get(col)
            if val is None or str(val).strip() == "":
                code = {
                    DistribucionCols.APLICAR_A_EXTRACTO: "missing_valor_intereses",
                    DistribucionCols.MORA_A_APLICAR: "missing_mora_a_aplicar",
                    DistribucionCols.ABONO_A_CAPITAL: "missing_abono_capital",
                    DistribucionCols.OTROS_VALORES: "missing_otros_valores",
                }.get(col, "missing_valor_intereses")
                issues.append(
                    _issue_from_code(
                        "finalize",
                        code,
                        excel_row=row_no,
                        field=col,
                    )
                )

    # Validación compartida con excel_write_guard (casos de pago, rutas, abonos).
    from app.application.services.excel_write_guard import (
        validate_review_workbook_before_finalize_writes,
    )

    monto_casos: dict[str, float] = {}
    if ReviewSheets.CASOS_PAGO in wb.sheetnames:
        ws_casos = wb[ReviewSheets.CASOS_PAGO]
        try:
            casos_header = _find_table_header_row(ws_casos, CasosPagoCols.ID_PAGO)
        except ValueError:
            casos_header = None
        if casos_header:
            c_headers: list[str] = []
            for r_idx, row in enumerate(ws_casos.iter_rows(values_only=True), start=1):
                if r_idx < casos_header:
                    continue
                if r_idx == casos_header:
                    c_headers = [str(v).strip() if v else "" for v in row]
                    continue
                if not any(row):
                    continue
                rd = dict(zip(c_headers, row))
                id_pago = str(rd.get(CasosPagoCols.ID_PAGO) or "").strip()
                if id_pago:
                    try:
                        monto_casos[id_pago] = float(
                            rd.get(CasosPagoCols.MONTO_BANCO) or 0
                        )
                    except (TypeError, ValueError):
                        monto_casos[id_pago] = 0.0

    abono_rows: list[dict[str, Any]] = []
    if ReviewSheets.DISTRIBUCION_ABONOS in wb.sheetnames:
        ws_ab = wb[ReviewSheets.DISTRIBUCION_ABONOS]
        try:
            ab_header = _find_table_header_row(ws_ab, "ID Pago")
        except ValueError:
            ab_header = None
        if ab_header:
            ab_headers: list[str] = []
            for r_idx, row in enumerate(ws_ab.iter_rows(values_only=True), start=1):
                if r_idx < ab_header:
                    continue
                if r_idx == ab_header:
                    ab_headers = [str(v).strip() if v else "" for v in row]
                    continue
                if not any(row):
                    continue
                abono_rows.append(dict(zip(ab_headers, row)))

    try:
        validate_review_workbook_before_finalize_writes(
            distributions=distrib_rows,
            abono_rows=abono_rows,
            monto_casos=monto_casos,
        )
    except ExcelWriteGuardError as exc:
        details = dict(exc.details)
        issues.append(
            _issue_from_code(
                "finalize",
                exc.error_code,
                excel_row=details.get("excel_row"),
                field=details.get("field"),
                payment_id=details.get("id_pago"),
            )
        )

    return issues
