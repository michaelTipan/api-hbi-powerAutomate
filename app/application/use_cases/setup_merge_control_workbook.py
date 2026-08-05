"""
Setup idempotente de Excel de control del proceso de validación de pagos (oficial por banco).

Crea o repara (sin sobrescribir datos de fila 2) los controles oficiales:

* ``control_proceso_validacion_pagos_banco_bogota.xlsx``
* ``control_proceso_validacion_pagos_banco_bancolombia.xlsx``

Hoja ``Procesos``, tabla ``tblControlProcesosPagos``.
Solo crea/repara los controles oficiales por banco en la carpeta de control.
"""

from __future__ import annotations

import io
import logging
import os
from dataclasses import dataclass
from datetime import date
from typing import Any

import httpx
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.application.config.payment_validation_settings import (
    BANK_CODE_BANCOLOMBIA,
    BANK_CODE_BOGOTA,
    PaymentValidationFolderName,
    list_payment_banks,
    resolve_payment_validation_folder,
)
from app.application.sharepoint_resolution import (
    encode_graph_drive_path,
    require_operations_site_config,
    resolve_sharepoint_path,
)
from app.domain.ports.graph import GraphApiPort

logger = logging.getLogger(__name__)

MERGE_CONTROL_FOLDER_RELATIVE_PATH = resolve_payment_validation_folder(
    PaymentValidationFolderName.CONTROL
)

SHEET_NAME = "Procesos"

# Tabla oficial en controles por banco.
PROCESS_CONTROL_TABLE_DISPLAY_NAME = "tblControlProcesosPagos"

# Alias de tabla en workbooks por banco creados antes de la migración de nombre.
_PRE_MIGRATION_TABLE_DISPLAY_NAME = "tblControlMergePDFs"

MERGE_CONTROL_COLUMNS: tuple[str, ...] = (
    "Title",
    "EstadoProceso",
    "IsActive",
    "HistoricalFilePath",
    "EmailPdfPath",
    "MergeOutputCount",
    "MergeSkippedCount",
    "LastErrorUserMessage",
    "LastErrorNextAction",
    "CreatedAtProceso",
    "LastUpdatedAtProceso",
    "MergeManifestPath",
)

MERGE_CONTROL_AMORTIZATION_COLUMN = "MergeManifestPath"

PROCESS_CONTROL_EXTENSION_COLUMNS: tuple[str, ...] = (
    "ProcessKey",
    "ProcessDate",
    "BankCode",
    "BankName",
    "ProcessId",
    "ValidationFilePath",
    "SecretaryFilePath",
    "GenerateIdempotencyKey",
    "FinalizeIdempotencyKey",
    "NotifyIdempotencyKey",
    "MergeIdempotencyKey",
    "ApplyIdempotencyKey",
    "LastCompletedStep",
    "LastStepStatus",
    "LastStepErrorCode",
    "GenerateJobId",
    "FinalizeJobId",
    "NotifyJobId",
    "MergeJobId",
    "ApplyJobId",
    "ExecutionId",
    "ExecutionLogPath",
    "LastAmortizationAttemptJson",
)

PROCESS_CONTROL_COLUMNS: tuple[str, ...] = MERGE_CONTROL_COLUMNS + PROCESS_CONTROL_EXTENSION_COLUMNS

# Columnas que deben existir en posición fija 1..N para Notify / inicio de Merge (workbooks antiguos).
MERGE_CONTROL_PREFIX_COLUMNS: tuple[str, ...] = tuple(
    c for c in MERGE_CONTROL_COLUMNS if c != MERGE_CONTROL_AMORTIZATION_COLUMN
)

@dataclass(frozen=True)
class ProcessControlBankDefinition:
    bank_code: str
    bank_name: str
    control_file_path: str


PROCESS_CONTROL_BANKS: tuple[ProcessControlBankDefinition, ...] = tuple(
    ProcessControlBankDefinition(
        bank_code=cfg.bank_code,
        bank_name=cfg.bank_name,
        control_file_path=cfg.control_file_path,
    )
    for cfg in list_payment_banks()
)

PROCESS_CONTROL_BANK_FILE_BOGOTA = next(
    b.control_file_path for b in PROCESS_CONTROL_BANKS if b.bank_code == BANK_CODE_BOGOTA
)
PROCESS_CONTROL_BANK_FILE_BANCOLOMBIA = next(
    b.control_file_path for b in PROCESS_CONTROL_BANKS if b.bank_code == BANK_CODE_BANCOLOMBIA
)

SECURITY_WARNING = (
    "El archivo fue creado con protección de hoja tipo Generate para evitar edición manual, "
    "pero la restricción de apertura debe manejarse con permisos de SharePoint si se requiere."
)

SETUP_PHASE_NOTE = (
    "Este endpoint prepara y repara la estructura de los controles oficiales por banco. "
    "El flujo productivo Generate/Finalize/Notify/Merge/dry-run/apply ya usa estos controles "
    "para encadenamiento, trazabilidad e idempotencia."
)

_FILL_HEADER = PatternFill(fill_type="solid", fgColor="002060")
_FONT_HEADER = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_FONT_BODY = Font(name="Calibri", size=11)
_BORDER = Border(
    left=Side(style="thin", color="C8C8C8"),
    right=Side(style="thin", color="C8C8C8"),
    top=Side(style="thin", color="C8C8C8"),
    bottom=Side(style="thin", color="C8C8C8"),
)
_ALIGN_HEADER = Alignment(vertical="center", horizontal="center", wrap_text=True)
_ALIGN_BODY = Alignment(vertical="center", horizontal="left", wrap_text=True)
_PROT_LOCKED = Protection(locked=True)

_TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)


class MergeControlSetupError(Exception):
    """Error de negocio / Graph al preparar el workbook de control."""

    def __init__(
        self,
        *,
        http_status: int,
        user_message: str,
        next_action: str,
        technical_message: str,
    ) -> None:
        self.http_status = http_status
        self.user_message = user_message
        self.next_action = next_action
        self.technical_message = technical_message
        super().__init__(technical_message)


def build_payment_validation_process_key(
    bank_code: str,
    process_date: str,
    process_id: str | None = None,
) -> str:
    """
    Clave de proceso para idempotencia.

    - Legado (un lote/día): ``payment-validation|banco_bogota|2026-06-01``
    - Lote único (recomendado): ``payment-validation|banco_bogota|2026-06-01|{uuid}``

    El ``process_id`` (UUID del proceso) permite varios lotes el mismo día sin
    sobrescribir histórico ni confundir la idempotencia de Apply.
    """
    bc = (bank_code or "").strip()
    pd = (process_date or "").strip()
    if not bc or not pd:
        raise ValueError("bank_code_and_process_date_required")
    base = f"payment-validation|{bc}|{pd}"
    pid = (process_id or "").strip()
    if pid:
        return f"{base}|{pid}"
    return base


def process_date_from_process_key(process_key: str) -> date | None:
    """
    Extrae YYYY-MM-DD del ProcessKey.

    Acepta legado ``…|YYYY-MM-DD`` y lote ``…|YYYY-MM-DD|{uuid}``.
    """
    parts = [p.strip() for p in str(process_key or "").split("|") if str(p).strip()]
    for part in parts:
        try:
            return date.fromisoformat(part[:10])
        except ValueError:
            continue
    return None


def process_id_from_process_key(process_key: str) -> str:
    """Extrae el UUID de lote si el ProcessKey tiene 4 segmentos."""
    parts = [p.strip() for p in str(process_key or "").split("|") if str(p).strip()]
    if len(parts) >= 4:
        return parts[-1]
    return ""


def build_process_artifact_filename(
    *,
    kind: str,
    bank_code: str,
    process_date: str,
    process_id: str,
    use_short_id: bool = False,
) -> str:
    """
    Nombre de artefacto SharePoint (histórico, soporte, revisión).

    - Revisión: UUID completo (``use_short_id=False``, default).
    - Histórico/soporte: id de 8 (``use_short_id=True``) bajo carpetas fechadas.
    """
    from app.application.services.dated_artifact_layout import short_process_id

    kind_s = (kind or "").strip().strip("_")
    bc = (bank_code or "").strip()
    pd = (process_date or "").strip()
    pid = (process_id or "").strip()
    if not kind_s or not bc or not pd or not pid:
        raise ValueError("artifact_filename_requires_kind_bank_date_process_id")
    id_part = short_process_id(pid) if use_short_id else pid
    return f"{kind_s}_{bc}_{pd}_{id_part}.xlsx"


def _content_endpoint(site_id: str, drive_id: str, file_path: str) -> str:
    return f"/sites/{site_id}/drives/{drive_id}/root:/{encode_graph_drive_path(file_path)}:/content"


def _item_endpoint(site_id: str, drive_id: str, file_path: str) -> str:
    return f"/sites/{site_id}/drives/{drive_id}/root:/{encode_graph_drive_path(file_path)}:"


async def _list_folder_children(
    graph: GraphApiPort,
    site_id: str,
    drive_id: str,
    parent_rel: str,
) -> list[dict[str, Any]]:
    if parent_rel.strip():
        enc = encode_graph_drive_path(parent_rel.strip().strip("/"))
        endpoint = f"/sites/{site_id}/drives/{drive_id}/root:/{enc}:/children"
    else:
        endpoint = f"/sites/{site_id}/drives/{drive_id}/root/children"
    resp = await graph.get(endpoint)
    return list(resp.get("value") or [])


async def _create_folder_child(
    graph: GraphApiPort,
    site_id: str,
    drive_id: str,
    parent_rel: str,
    folder_name: str,
) -> None:
    if parent_rel.strip():
        enc = encode_graph_drive_path(parent_rel.strip().strip("/"))
        endpoint = f"/sites/{site_id}/drives/{drive_id}/root:/{enc}:/children"
    else:
        endpoint = f"/sites/{site_id}/drives/{drive_id}/root/children"
    body: dict[str, Any] = {
        "name": folder_name,
        "folder": {},
        "@microsoft.graph.conflictBehavior": "fail",
    }
    try:
        await graph.post_json(endpoint, body)
    except httpx.HTTPStatusError as exc:
        if exc.response is not None and exc.response.status_code == 409:
            logger.info("process-control-setup: carpeta %r ya existía (409)", folder_name)
            return
        raise


async def ensure_merge_control_folder_path(
    graph: GraphApiPort,
    site_id: str,
    drive_id: str,
    folder_rel: str,
) -> None:
    parts = [p for p in folder_rel.strip().split("/") if p.strip()]
    acc: list[str] = []
    for part in parts:
        parent = "/".join(acc) if acc else ""
        kids = await _list_folder_children(graph, site_id, drive_id, parent)
        folder_names = {
            str(it.get("name", ""))
            for it in kids
            if "folder" in it and str(it.get("name", "")).strip()
        }
        if part not in folder_names:
            try:
                await _create_folder_child(graph, site_id, drive_id, parent, part)
            except httpx.HTTPStatusError as exc:
                code = exc.response.status_code if exc.response else 0
                body_txt = (exc.response.text if exc.response else "") or ""
                raise MergeControlSetupError(
                    http_status=502,
                    user_message="No se pudo crear la carpeta de control en SharePoint.",
                    next_action=(
                        "Cree la carpeta manualmente o verifique permisos de escritura en la ruta de validación."
                    ),
                    technical_message=f"Graph HTTP {code}: {body_txt[:2000]}",
                ) from exc
        acc.append(part)


async def _file_exists(graph: GraphApiPort, site_id: str, drive_id: str, path: str) -> bool:
    try:
        await graph.get_bytes(_content_endpoint(site_id, drive_id, path))
        return True
    except httpx.HTTPStatusError as exc:
        if exc.response is not None and exc.response.status_code == 404:
            return False
        raise


async def _drive_item_web_url(graph: GraphApiPort, site_id: str, drive_id: str, path: str) -> str | None:
    try:
        meta = await graph.get(_item_endpoint(site_id, drive_id, path))
        wu = meta.get("webUrl")
        return str(wu).strip() if isinstance(wu, str) and wu.strip() else None
    except httpx.HTTPStatusError as exc:
        if exc.response is not None and exc.response.status_code == 404:
            return None
        raise


def _column_widths(ncols: int) -> dict[int, float]:
    base = {
        1: 28.0,
        2: 22.0,
        3: 12.0,
        4: 62.0,
        5: 62.0,
        6: 18.0,
        7: 20.0,
        8: 48.0,
        9: 48.0,
        10: 22.0,
        11: 22.0,
        12: 72.0,
    }
    for c in range(13, ncols + 1):
        base[c] = 24.0
    return base


def merge_control_header_column_map(ws: Any) -> dict[str, int]:
    col_map: dict[str, int] = {}
    for c in range(1, (ws.max_column or 0) + 1):
        name = str(ws.cell(row=1, column=c).value or "").strip()
        if name and name not in col_map:
            col_map[name] = c
    return col_map


def merge_control_prefix_headers_match(ws: Any) -> bool:
    """Valida encabezados 1..11 (plantilla histórica); MergeManifestPath puede faltar."""
    for i, exp in enumerate(MERGE_CONTROL_PREFIX_COLUMNS, start=1):
        if str(ws.cell(row=1, column=i).value or "").strip() != exp:
            return False
    return True


def _resize_control_table_and_filter(
    ws: Any,
    *,
    columns: tuple[str, ...],
    table_names: tuple[str, ...],
) -> None:
    ncols = len(columns)
    for col_idx, width in _column_widths(ncols).items():
        if col_idx <= ncols:
            ws.column_dimensions[get_column_letter(col_idx)].width = width
    ref = f"A1:{get_column_letter(ncols)}2"
    ws.auto_filter.ref = ref
    for tname in table_names:
        if tname in ws.tables:
            ws.tables[tname].ref = ref


def _resize_merge_control_table_and_filter(ws: Any) -> None:
    _resize_control_table_and_filter(
        ws, columns=MERGE_CONTROL_COLUMNS, table_names=(_PRE_MIGRATION_TABLE_DISPLAY_NAME,)
    )


def _resize_process_control_table_and_filter(ws: Any) -> None:
    names: list[str] = []
    if PROCESS_CONTROL_TABLE_DISPLAY_NAME in ws.tables:
        names.append(PROCESS_CONTROL_TABLE_DISPLAY_NAME)
    if _PRE_MIGRATION_TABLE_DISPLAY_NAME in ws.tables:
        names.append(_PRE_MIGRATION_TABLE_DISPLAY_NAME)
    _resize_control_table_and_filter(
        ws,
        columns=PROCESS_CONTROL_COLUMNS,
        table_names=tuple(names) if names else (PROCESS_CONTROL_TABLE_DISPLAY_NAME,),
    )


def ensure_merge_control_worksheet_columns(ws: Any) -> bool:
    """
    Añade columnas del contrato legacy que falten (p. ej. MergeManifestPath).
    Devuelve True si modificó la hoja.
    """
    col_map = merge_control_header_column_map(ws)
    modified = False
    next_col = max(col_map.values(), default=len(MERGE_CONTROL_PREFIX_COLUMNS)) + 1

    for name in MERGE_CONTROL_COLUMNS:
        if name in col_map:
            continue
        canonical = MERGE_CONTROL_COLUMNS.index(name) + 1
        if canonical <= len(MERGE_CONTROL_COLUMNS) and str(
            ws.cell(row=1, column=canonical).value or ""
        ).strip() == "":
            col = canonical
        else:
            col = next_col
            next_col += 1
        hcell = ws.cell(row=1, column=col, value=name)
        hcell.font = _FONT_HEADER
        hcell.fill = _FILL_HEADER
        hcell.alignment = _ALIGN_HEADER
        hcell.border = _BORDER
        dcell = ws.cell(row=2, column=col, value="")
        dcell.font = _FONT_BODY
        dcell.alignment = _ALIGN_BODY
        dcell.border = _BORDER
        col_map[name] = col
        modified = True

    if modified:
        _resize_merge_control_table_and_filter(ws)
    return modified


def ensure_process_control_worksheet_columns(ws: Any) -> tuple[bool, list[str]]:
    """
    Añade columnas legacy y de proceso que falten al final (sin duplicar).
    Devuelve (modified, lista de nombres de columnas agregadas).
    """
    col_map = merge_control_header_column_map(ws)
    added: list[str] = []
    next_col = max(col_map.values(), default=0) + 1

    for name in PROCESS_CONTROL_COLUMNS:
        if name in col_map:
            continue
        col = next_col
        next_col += 1
        hcell = ws.cell(row=1, column=col, value=name)
        hcell.font = _FONT_HEADER
        hcell.fill = _FILL_HEADER
        hcell.alignment = _ALIGN_HEADER
        hcell.border = _BORDER
        default_val = _default_process_control_cell_value(name)
        dcell = ws.cell(row=2, column=col, value=default_val)
        dcell.font = _FONT_BODY
        dcell.alignment = _ALIGN_BODY
        dcell.border = _BORDER
        col_map[name] = col
        added.append(name)

    modified = bool(added)
    if modified:
        _resize_process_control_table_and_filter(ws)
    return modified, added


def _default_process_control_cell_value(column_name: str) -> Any:
    if column_name in ("MergeOutputCount", "MergeSkippedCount"):
        return 0
    if column_name == "EstadoProceso":
        return "VACIO"
    if column_name == "IsActive":
        return "false"
    return ""


def _initial_process_control_row_values(bank_code: str, bank_name: str) -> dict[str, Any]:
    values: dict[str, Any] = {name: _default_process_control_cell_value(name) for name in PROCESS_CONTROL_COLUMNS}
    values["BankCode"] = bank_code
    values["BankName"] = bank_name
    return values


def apply_process_control_worksheet_protection(ws: Any) -> None:
    """Bloquea filas 1-2 del contrato completo de proceso."""
    ncols = len(PROCESS_CONTROL_COLUMNS)
    for r in (1, 2):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).protection = _PROT_LOCKED
    ws.protection.sheet = True


def _apply_sheet_cosmetics(ws: Any) -> None:
    try:
        ws.sheet_view.showGridLines = False
    except Exception:
        pass
    try:
        ws.sheet_properties.tabColor = Color(rgb="002060")
    except Exception:
        pass


def _add_table(ws: Any, *, display_name: str, ncols: int) -> None:
    ref = f"A1:{get_column_letter(ncols)}2"
    tab = Table(displayName=display_name, ref=ref)
    tab.tableStyleInfo = _TABLE_STYLE
    ws.add_table(tab)


def _build_workbook_with_base_columns_only() -> bytes:
    """Workbook de prueba con solo columnas base (sin extensiones de proceso)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    _apply_sheet_cosmetics(ws)

    ncols = len(MERGE_CONTROL_COLUMNS)
    for col_idx, name in enumerate(MERGE_CONTROL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = _FILL_HEADER
        cell.font = _FONT_HEADER
        cell.alignment = _ALIGN_HEADER
        cell.border = _BORDER

    row2_values: list[Any] = [
        "",
        "VACIO",
        "false",
        "",
        "",
        0,
        0,
        "",
        "",
        "",
        "",
    ]
    for col_idx, val in enumerate(row2_values, start=1):
        cell = ws.cell(row=2, column=col_idx, value=val)
        cell.font = _FONT_BODY
        cell.alignment = _ALIGN_BODY
        cell.border = _BORDER

    for col_idx, width in _column_widths(ncols).items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}2"
    ws.freeze_panes = "A2"
    _add_table(ws, display_name=_PRE_MIGRATION_TABLE_DISPLAY_NAME, ncols=ncols)
    apply_process_control_worksheet_protection(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_process_control_workbook_bytes(bank_code: str, bank_name: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    _apply_sheet_cosmetics(ws)

    row_values = _initial_process_control_row_values(bank_code, bank_name)
    ncols = len(PROCESS_CONTROL_COLUMNS)
    for col_idx, name in enumerate(PROCESS_CONTROL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = _FILL_HEADER
        cell.font = _FONT_HEADER
        cell.alignment = _ALIGN_HEADER
        cell.border = _BORDER
        dcell = ws.cell(row=2, column=col_idx, value=row_values[name])
        dcell.font = _FONT_BODY
        dcell.alignment = _ALIGN_BODY
        dcell.border = _BORDER

    for col_idx, width in _column_widths(ncols).items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}2"
    ws.freeze_panes = "A2"
    _add_table(ws, display_name=PROCESS_CONTROL_TABLE_DISPLAY_NAME, ncols=ncols)
    apply_process_control_worksheet_protection(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _process_control_header_row_matches(ws: Any) -> bool:
    if not merge_control_prefix_headers_match(ws):
        return False
    col_map = merge_control_header_column_map(ws)
    return all(name in col_map for name in PROCESS_CONTROL_COLUMNS)


def _process_control_table_ok(ws: Any) -> bool:
    return (
        PROCESS_CONTROL_TABLE_DISPLAY_NAME in ws.tables
        or _PRE_MIGRATION_TABLE_DISPLAY_NAME in ws.tables
    )


def _validate_and_repair_process_control_workbook(
    data: bytes,
) -> tuple[list[str], bool, bool, bytes | None, list[str]]:
    """
    Devuelve (warnings, structure_ok, excel_protection_applied, new_bytes_if_modified, columns_added).
    """
    warnings: list[str] = []
    columns_added: list[str] = []
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    modified = False
    try:
        if SHEET_NAME not in wb.sheetnames:
            warnings.append("needs_manual_review: falta la hoja Procesos")
            return warnings, False, False, None, columns_added

        ws = wb[SHEET_NAME]

        if not merge_control_prefix_headers_match(ws):
            warnings.append("needs_manual_review: encabezados legacy no coinciden con el contrato")
            return warnings, False, ws.protection.sheet is True, None, columns_added

        cols_modified, added = ensure_process_control_worksheet_columns(ws)
        if cols_modified:
            columns_added.extend(added)
            warnings.append("repaired: se añadieron columnas faltantes del contrato de proceso")
            modified = True

        if ws.max_row < 2:
            warnings.append("needs_manual_review: falta la fila 2 de control")
            return warnings, False, ws.protection.sheet is True, None, columns_added

        if not _process_control_table_ok(ws):
            ref = f"A1:{get_column_letter(len(PROCESS_CONTROL_COLUMNS))}2"
            tab = Table(displayName=PROCESS_CONTROL_TABLE_DISPLAY_NAME, ref=ref)
            tab.tableStyleInfo = _TABLE_STYLE
            ws.add_table(tab)
            warnings.append("repaired: se añadió la tabla tblControlProcesosPagos faltante")
            modified = True
        elif (
            PROCESS_CONTROL_TABLE_DISPLAY_NAME not in ws.tables
            and _PRE_MIGRATION_TABLE_DISPLAY_NAME in ws.tables
        ):
            _resize_process_control_table_and_filter(ws)
            modified = True

        if not ws.protection.sheet:
            warnings.append("repaired: se activó protección de hoja")
            modified = True

        need_relock = False
        for r in (1, 2):
            for c in range(1, len(PROCESS_CONTROL_COLUMNS) + 1):
                cell = ws.cell(row=r, column=c)
                if getattr(cell, "protection", None) is None or cell.protection.locked is not True:
                    need_relock = True
                    break
            if need_relock:
                break
        if need_relock:
            modified = True

        if modified:
            apply_process_control_worksheet_protection(ws)

        struct_ok = (
            _process_control_header_row_matches(ws)
            and ws.max_row >= 2
            and _process_control_table_ok(ws)
        )
        protection_on = ws.protection.sheet is True

        new_bytes: bytes | None = None
        if modified:
            buf = io.BytesIO()
            wb.save(buf)
            new_bytes = buf.getvalue()

        return warnings, struct_ok, protection_on, new_bytes, columns_added
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()


async def _upload_workbook(
    graph: GraphApiPort,
    site_id: str,
    drive_id: str,
    file_path: str,
    payload: bytes,
) -> str | None:
    resp = await graph.put_bytes(
        _content_endpoint(site_id, drive_id, file_path),
        payload,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    if isinstance(resp, dict):
        wu = resp.get("webUrl")
        return str(wu).strip() if isinstance(wu, str) and wu.strip() else None
    return None


def _graph_setup_error_from_http(exc: httpx.HTTPStatusError, *, creating: bool) -> MergeControlSetupError:
    code = exc.response.status_code if exc.response else 0
    body_txt = (exc.response.text if exc.response else "") or ""
    if code == 403:
        return MergeControlSetupError(
            http_status=403,
            user_message="No se pudo crear o actualizar el archivo de control en SharePoint.",
            next_action=(
                "Verifique que la aplicación tenga permisos suficientes para crear archivos en la carpeta "
                "de validación."
            ),
            technical_message=f"Graph HTTP {code}: {body_txt[:2000]}",
        )
    action = "crear" if creating else "actualizar"
    return MergeControlSetupError(
        http_status=502,
        user_message=f"No se pudo {action} el archivo de control en SharePoint.",
        next_action="Verifique espacio, bloqueos (423) o permisos en la biblioteca.",
        technical_message=f"Graph HTTP {code}: {body_txt[:2000]}",
    )


async def _setup_process_control_bank_workbook(
    graph: GraphApiPort,
    site_id: str,
    drive_id: str,
    bank: ProcessControlBankDefinition,
) -> dict[str, Any]:
    file_path = bank.control_file_path
    warnings: list[str] = []
    columns_added: list[str] = []
    repaired = False
    created = False
    file_url: str | None = None
    excel_protection = False

    try:
        exists = await _file_exists(graph, site_id, drive_id, file_path)
    except httpx.HTTPStatusError as exc:
        raise _graph_setup_error_from_http(exc, creating=True) from exc

    if exists:
        raw = await graph.get_bytes(_content_endpoint(site_id, drive_id, file_path))
        w, struct_ok, excel_protection, new_bytes, added = _validate_and_repair_process_control_workbook(
            raw
        )
        warnings.extend(w)
        columns_added.extend(added)
        if new_bytes is not None:
            repaired = True
            try:
                file_url = await _upload_workbook(graph, site_id, drive_id, file_path, new_bytes)
            except httpx.HTTPStatusError as exc:
                code = exc.response.status_code if exc.response else 0
                body_txt = (exc.response.text if exc.response else "") or ""
                warnings.append(
                    f"needs_manual_review: no se pudo guardar reparación (HTTP {code}): {body_txt[:500]}"
                )
        if file_url is None:
            file_url = await _drive_item_web_url(graph, site_id, drive_id, file_path)
        status = "success" if struct_ok else "needs_manual_review"
        if not struct_ok:
            warnings.append("needs_manual_review: estructura del archivo no válida o incompleta")
        return {
            "bank_code": bank.bank_code,
            "bank_name": bank.bank_name,
            "control_file_path": file_path,
            "created": created,
            "repaired": repaired,
            "columns_added": columns_added,
            "status": status,
            "file_url": file_url,
            "sheet_name": SHEET_NAME,
            "table_name": PROCESS_CONTROL_TABLE_DISPLAY_NAME,
            "previous_table_display_name": _PRE_MIGRATION_TABLE_DISPLAY_NAME,
            "warnings": warnings,
            "excel_protection_applied": excel_protection,
        }

    payload = _build_process_control_workbook_bytes(bank.bank_code, bank.bank_name)
    try:
        file_url = await _upload_workbook(graph, site_id, drive_id, file_path, payload)
    except httpx.HTTPStatusError as exc:
        raise _graph_setup_error_from_http(exc, creating=True) from exc

    if not file_url:
        file_url = await _drive_item_web_url(graph, site_id, drive_id, file_path)

    created = True
    excel_protection = True
    return {
        "bank_code": bank.bank_code,
        "bank_name": bank.bank_name,
        "control_file_path": file_path,
        "created": created,
        "repaired": repaired,
        "columns_added": columns_added,
        "status": "success",
        "file_url": file_url,
        "sheet_name": SHEET_NAME,
        "table_name": PROCESS_CONTROL_TABLE_DISPLAY_NAME,
        "previous_table_display_name": _PRE_MIGRATION_TABLE_DISPLAY_NAME,
        "warnings": warnings,
        "excel_protection_applied": excel_protection,
    }


async def setup_merge_control_workbook(graph: GraphApiPort) -> dict[str, Any]:
    site_search = os.getenv("GRAPH_SHAREPOINT_SITE_SEARCH", "").strip()
    drive_name = os.getenv("GRAPH_SHAREPOINT_DRIVE_NAME", "").strip()
    require_operations_site_config()

    base = await resolve_sharepoint_path(graph, site_search, drive_name, MERGE_CONTROL_FOLDER_RELATIVE_PATH)
    site_id = base["site_id"]
    drive_id = base["drive_id"]

    try:
        await ensure_merge_control_folder_path(graph, site_id, drive_id, MERGE_CONTROL_FOLDER_RELATIVE_PATH)
    except MergeControlSetupError:
        raise
    except httpx.HTTPStatusError as exc:
        code = exc.response.status_code if exc.response else 0
        body_txt = (exc.response.text if exc.response else "") or ""
        if code == 403:
            raise MergeControlSetupError(
                http_status=403,
                user_message="No se pudo preparar la carpeta de control en SharePoint.",
                next_action=(
                    "Verifique que la aplicación tenga permisos suficientes para crear archivos en la carpeta "
                    "de validación."
                ),
                technical_message=f"Graph HTTP {code}: {body_txt[:2000]}",
            ) from exc
        raise MergeControlSetupError(
            http_status=502,
            user_message="No se pudo preparar la carpeta de control en SharePoint.",
            next_action="Revise permisos y el detalle técnico devuelto por Graph.",
            technical_message=f"Graph HTTP {code}: {body_txt[:2000]}",
        ) from exc

    banks: list[dict[str, Any]] = []
    for bank_def in PROCESS_CONTROL_BANKS:
        banks.append(await _setup_process_control_bank_workbook(graph, site_id, drive_id, bank_def))

    overall_status = "success"
    if any(b.get("status") != "success" for b in banks):
        overall_status = "needs_manual_review"

    return {
        "status": overall_status,
        "banks": banks,
        "official_control_file_names": [
            "control_proceso_validacion_pagos_banco_bogota.xlsx",
            "control_proceso_validacion_pagos_banco_bancolombia.xlsx",
        ],
        "process_control_columns": list(PROCESS_CONTROL_COLUMNS),
        "base_columns": list(MERGE_CONTROL_COLUMNS),
        "row_mode": "single_active_row",
        "sharepoint_permission_applied": False,
        "security_warning": SECURITY_WARNING,
        "phase_note": SETUP_PHASE_NOTE,
    }
