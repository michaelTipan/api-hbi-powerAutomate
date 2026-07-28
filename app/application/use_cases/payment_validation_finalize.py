import json
import logging
import os
import io
import re
import unicodedata
from calendar import monthrange
from urllib.parse import unquote
from datetime import date, datetime
from typing import Any

import httpx
import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain.ports.graph import GraphApiPort
from app.application.sharepoint_resolution import encode_graph_drive_path, resolve_sharepoint_path
from app.application.services.payment_followup_finalize import register_payment_followups_after_finalize
from app.application.services.review_schema import (
    DISTRIBUCION_ABONOS_TECHNICAL_HIDDEN_COLUMNS,
    DISTRIBUCION_TECHNICAL_HIDDEN_COLUMNS,
    ApplicationSubtype,
    AsientosPendientesCols,
    DistribucionAbonosCols,
    ReviewSheets,
    ControlCols,
    CasosPagoCols,
    DistribucionCols,
    EstadoPago,
    SUPPORT_NOT_APPLICABLE,
    TipoAplicacion,
    TipoAplicacionVisible,
    ValidarAbono,
    apply_legacy_estado_migration,
    apply_policy_to_abonos_row,
    apply_policy_to_pagos_row,
    find_distribucion_pagos_sheet,
    is_validar_abono_si,
    is_validar_pago_si,
    REVIEW_SCHEMA_VERSION,
    detect_distrib_schema_version_from_headers,
    normalize_distrib_row_keys,
    policy_from_row,
    read_control_review_schema_version,
    policy_requires_closing_extract,
    policy_requires_reference_extract,
    require_validar_abono_value,
    resolve_application_policy,
)

logger = logging.getLogger(__name__)


def _include_in_validation_outputs(dist: dict[str, Any]) -> bool:
    """Filas que generan ruta extracto, filas en soporte secretaría y conteo «validadas»."""
    ep = str(dist.get(DistribucionCols.ESTADO_PAGO, "")).strip().upper()
    return ep in EstadoPago.SECRETARY_AND_RUTA and is_validar_pago_si(dist)

SECRETARY_SHEET = "Asientos_Pendientes"
ASIENTOS_FOLDER_LABEL = "ASIENTOS CONTABLES"
ASIENTOS_FOLDER_PER_CREDIT_PREFIX = "ASIENTOS CONTABLES CRED"
PENDIENTE_CREAR_ASIENTOS = "PENDIENTE_CREAR"
OBS_NO_ASIENTOS = "No se encontró carpeta ASIENTOS CONTABLES."

SECRETARY_HEADERS = list(AsientosPendientesCols.HEADERS)

SECRETARY_TITLE = "Soporte de asientos contables"
SECRETARY_INSTRUCTION = (
    "Use esta hoja para cargar o revisar los soportes de asientos contables por crédito. "
    "Abra el link de la carpeta correspondiente, cargue el PDF del asiento contable y luego "
    "continúe con la consolidación cuando todos los soportes estén completos."
)
SECRETARY_HEADER_ROW = 3
SECRETARY_FIRST_DATA_ROW = 4

_SEC_THIN = Side(style="thin", color="C8C8C8")
_SEC_BORDER_LIGHT = Border(left=_SEC_THIN, right=_SEC_THIN, top=_SEC_THIN, bottom=_SEC_THIN)
_SEC_FILL_NAVY = PatternFill(fill_type="solid", fgColor="002060")
_SEC_FILL_HEADER = _SEC_FILL_NAVY
_SEC_FONT_TITLE = Font(name="Calibri", bold=True, size=16, color="FFFFFF")
_SEC_FILL_INSTRUCTION = PatternFill(fill_type="solid", fgColor="E8F2FA")
_SEC_FONT_INSTRUCTION = Font(name="Calibri", size=12, color="1A2F36")
_SEC_MEDIUM_CLIENT_EDGE = Side(style="medium", color="002060")
_SEC_FONT_HEADER = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_SEC_FONT_BODY = Font(name="Calibri", size=11)
_SEC_FONT_HLINK = Font(name="Calibri", color="0563C1", size=11, underline="single")
_SEC_FONT_TOTAL_LABEL = Font(name="Calibri", bold=True, size=11)
_SEC_ALIGN_CENTER_WRAP = Alignment(vertical="center", horizontal="center", wrap_text=True)
_SEC_ALIGN_VCENTER = Alignment(vertical="center", horizontal="left")
_SEC_ALIGN_WRAP = Alignment(vertical="center", horizontal="left", wrap_text=True)
_SEC_FMT_MONEY = "#,##0.00"
_SEC_FMT_DATE = "yyyy-mm-dd"
_SEC_TAB_COLOR = "FF2E75B6"


def _secretary_coerce_date_value(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.date()
    return v


def _apply_secretary_banner(ws: Any, ncols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    t = ws.cell(1, 1, SECRETARY_TITLE)
    t.font = _SEC_FONT_TITLE
    t.fill = _SEC_FILL_NAVY
    t.alignment = _SEC_ALIGN_CENTER_WRAP
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    s = ws.cell(2, 1, SECRETARY_INSTRUCTION)
    s.fill = _SEC_FILL_INSTRUCTION
    s.font = _SEC_FONT_INSTRUCTION
    s.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
    ws.row_dimensions[2].height = 56


def _apply_secretary_header_style(ws: Any, row_idx: int, headers: list[str]) -> None:
    for col_idx, name in enumerate(headers, start=1):
        cell = ws.cell(row_idx, col_idx, value=name)
        cell.fill = _SEC_FILL_HEADER
        cell.font = _SEC_FONT_HEADER
        cell.alignment = _SEC_ALIGN_CENTER_WRAP
        cell.border = _SEC_BORDER_LIGHT
    ws.row_dimensions[row_idx].height = 22


def _visible_len_secretary(value: Any) -> int:
    if value is None:
        return 0
    s = str(value)
    if s.startswith("="):
        return min(42, len(s))
    return len(s)


def _auto_fit_secretary_columns(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for col in range(min_col, max_col + 1):
        best = 9.0
        for row in range(min_row, max_row + 1):
            cell = ws.cell(row=row, column=col)
            est = _visible_len_secretary(cell.value) * 1.12 + 2.5
            if est > best:
                best = est
        best = min(56.0, max(9.0, best))
        letter = get_column_letter(col)
        cur = ws.column_dimensions[letter].width
        if cur is None or cur < best:
            ws.column_dimensions[letter].width = best


def _apply_secretary_body_style(
    ws: Any,
    first_data_row: int,
    last_data_row: int,
    ncols: int,
    hmap: dict[str, int],
) -> None:
    if last_data_row < first_data_row:
        return
    wrap_cols = {
        hmap[name]
        for name in (
            AsientosPendientesCols.CLIENTE,
            AsientosPendientesCols.OBSERVACION,
            AsientosPendientesCols.LINK_CARPETA_ASIENTOS,
            AsientosPendientesCols.LINK_EXTRACTO,
            AsientosPendientesCols.LINK_TABLA,
        )
        if name in hmap
    }
    for r in range(first_data_row, last_data_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _SEC_BORDER_LIGHT
            cell.font = _SEC_FONT_BODY
            cell.alignment = _SEC_ALIGN_WRAP if c in wrap_cols else _SEC_ALIGN_VCENTER


def _apply_secretary_number_formats(ws: Any, first_data_row: int, last_data_row: int, hmap: dict[str, int]) -> None:
    if last_data_row < first_data_row:
        return
    money_cols = {hmap[AsientosPendientesCols.TOTAL_VALIDADO], hmap[AsientosPendientesCols.MONTO_BANCO]}
    date_cols = {hmap[AsientosPendientesCols.FECHA_BANCO], hmap[AsientosPendientesCols.FECHA_LIMITE]}
    for r in range(first_data_row, last_data_row + 1):
        for c in money_cols:
            cell = ws.cell(row=r, column=c)
            if cell.value not in (None, "", SUPPORT_NOT_APPLICABLE) and not (
                isinstance(cell.value, str) and str(cell.value).startswith("=")
            ):
                cell.number_format = _SEC_FMT_MONEY
        for c in date_cols:
            cell = ws.cell(row=r, column=c)
            v = cell.value
            if v not in (None, "", SUPPORT_NOT_APPLICABLE):
                coerced = _secretary_coerce_date_value(v)
                if coerced is not v:
                    cell.value = coerced
                if not (isinstance(cell.value, str) and str(cell.value).startswith("=")):
                    cell.number_format = _SEC_FMT_DATE


def _secretary_http_url(value: Any) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    if s.lower().startswith("http://") or s.lower().startswith("https://"):
        return s
    return None


def _secretary_link_label_suffix(credito: Any, cliente: Any) -> str:
    cred = str(credito or "").strip()
    cli = str(cliente or "").strip()
    if cred and re.fullmatch(r"\d+", cred):
        return f"crédito {cred}"
    if cred:
        return cred
    return cli


def _secretary_link_visible_text(link_kind: str, credito: Any, cliente: Any) -> str:
    suffix = _secretary_link_label_suffix(credito, cliente)
    if link_kind == "extracto":
        return f"Ver extracto {suffix}".strip() if suffix else "Ver extracto"
    if link_kind == "tabla":
        return f"Ver tabla {suffix}".strip() if suffix else "Ver tabla"
    if link_kind == "carpeta_asientos":
        return f"Ver carpeta asientos {suffix}".strip() if suffix else "Ver carpeta asientos"
    return ""


def _hyperlink_target_from_cell(cell: Any) -> str | None:
    hl = getattr(cell, "hyperlink", None)
    if hl is None:
        return None
    tgt = getattr(hl, "target", None) or getattr(hl, "ref", None)
    return str(tgt).strip() if tgt else None


def _set_secretary_link_cell(
    dst: Any,
    src: Any,
    link_kind: str,
    credito: Any,
    cliente: Any,
) -> None:
    tgt = _secretary_http_url(_hyperlink_target_from_cell(src))
    if tgt:
        dst.hyperlink = tgt
        dst.value = _secretary_link_visible_text(link_kind, credito, cliente)
        dst.font = _SEC_FONT_HLINK
    else:
        dst.value = None


def _set_secretary_url_link_cell(
    dst: Any,
    url: Any,
    link_kind: str,
    credito: Any,
    cliente: Any,
) -> None:
    tgt = _secretary_http_url(url)
    if tgt:
        dst.hyperlink = tgt
        dst.value = _secretary_link_visible_text(link_kind, credito, cliente)
        dst.font = _SEC_FONT_HLINK
    else:
        dst.value = None


def _secretary_row_border(*, client_top: bool = False, client_bottom: bool = False) -> Border:
    top = _SEC_MEDIUM_CLIENT_EDGE if client_top else _SEC_THIN
    bottom = _SEC_MEDIUM_CLIENT_EDGE if client_bottom else _SEC_THIN
    return Border(left=_SEC_THIN, right=_SEC_THIN, top=top, bottom=bottom)


def _apply_secretary_client_borders(
    ws: Any,
    first_data_row: int,
    last_data_row: int,
    ncols: int,
    col_cliente: int,
) -> None:
    if last_data_row < first_data_row:
        return
    indexed: list[tuple[int, str]] = []
    for r in range(first_data_row, last_data_row + 1):
        cliente = str(ws.cell(row=r, column=col_cliente).value or "").strip()
        if not cliente:
            continue
        indexed.append((r, cliente))
    if not indexed:
        return
    for i, (r, cliente) in enumerate(indexed):
        prev_cliente = indexed[i - 1][1] if i > 0 else None
        next_cliente = indexed[i + 1][1] if i + 1 < len(indexed) else None
        client_top = i > 0 and bool(cliente) and cliente != prev_cliente
        client_bottom = bool(cliente) and (i == len(indexed) - 1 or cliente != next_cliente)
        border = _secretary_row_border(client_top=client_top, client_bottom=client_bottom)
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).border = border


def _secretary_freeze_panes_cell(hmap: dict[str, int]) -> str:
    col_credito = hmap[AsientosPendientesCols.CREDITO]
    return f"{get_column_letter(col_credito)}{SECRETARY_FIRST_DATA_ROW}"


def _apply_secretary_total_row(
    ws: Any, hmap: dict[str, int], first_data_row: int, last_data_row: int
) -> int | None:
    if last_data_row < first_data_row:
        return None
    trow = last_data_row + 1
    tc = hmap[AsientosPendientesCols.TOTAL_VALIDADO]
    lett = get_column_letter(tc)
    ws.cell(trow, 1, "Total validado general")
    ws.cell(trow, tc, f"=SUM({lett}{first_data_row}:{lett}{last_data_row})")
    lbl = ws.cell(trow, 1)
    lbl.font = _SEC_FONT_TOTAL_LABEL
    lbl.border = _SEC_BORDER_LIGHT
    sum_cell = ws.cell(trow, tc)
    sum_cell.font = _SEC_FONT_TOTAL_LABEL
    sum_cell.number_format = _SEC_FMT_MONEY
    sum_cell.border = _SEC_BORDER_LIGHT
    top = Side(style="medium", color="888888")
    for c in range(1, len(SECRETARY_HEADERS) + 1):
        cell = ws.cell(trow, c)
        cell.border = Border(
            left=_SEC_THIN,
            right=_SEC_THIN,
            top=top,
            bottom=_SEC_THIN,
        )
    return trow


def _find_table_header_row(ws: Any, first_header_value: str) -> int:
    marker = str(first_header_value).strip()
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and str(row[0]).strip() == marker:
            return r_idx
    raise ValueError("missing_sheet_headers")


def _normalize_process_date(process_date: date | str | None) -> date:
    if process_date is None:
        return datetime.now().date()
    if isinstance(process_date, date):
        return process_date
    return date.fromisoformat(str(process_date))


def _resolve_validation_selection(children_items: list[dict[str, Any]], prefix: str) -> str:
    valid_files = []
    for item in children_items:
        name = str(item.get("name", "")).strip()
        if not name or name.startswith("~$") or not name.lower().endswith(".xlsx"):
            continue
        if prefix and not name.startswith(prefix):
            continue
        valid_files.append(item)

    if not valid_files:
        raise ValueError("no_validation_file_found")

    def sort_key(item: dict[str, Any]) -> tuple[str, str]:
        return (str(item.get("lastModifiedDateTime", "")), str(item.get("name", "")))

    valid_files.sort(key=sort_key, reverse=True)
    return str(valid_files[0]["name"])


def _build_content_endpoint(site_id: str, drive_id: str, file_path: str) -> str:
    return f"/sites/{site_id}/drives/{drive_id}/root:/{encode_graph_drive_path(file_path)}:/content"


def _accounting_cell_filled(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return True
    return str(value).strip() != ""


def _validate_no_duplicate_distrib_monto_banco(distributions: list[dict[str, Any]]) -> None:
    """Monto banco en Distribución debe aparecer como máximo una vez por ID Pago (fila líder)."""
    filled_rows_by_id: dict[str, int] = {}
    for dist in distributions:
        id_pago = str(dist.get(DistribucionCols.ID_PAGO) or "").strip()
        if not id_pago:
            continue
        raw = dist.get(DistribucionCols.MONTO_BANCO)
        if not _accounting_cell_filled(raw):
            continue
        filled_rows_by_id[id_pago] = filled_rows_by_id.get(id_pago, 0) + 1
        if filled_rows_by_id[id_pago] > 1:
            raise ValueError("duplicate_bank_amount_in_payment_group")


def _normalize_for_credit_match(text: str) -> str:
    if text is None:
        return ""
    value = unicodedata.normalize("NFD", str(text).strip().lower())
    value = "".join(ch for ch in value if unicodedata.category(ch) != "Mn")
    value = re.sub(r"\s+", " ", value)
    return value


def _extract_credit_id_for_match(credit_value: str) -> str | None:
    s = str(credit_value).strip()
    if not s:
        return None
    norm = _normalize_for_credit_match(s)
    m = re.search(r"(?:credito|obligacion)\s*#?\s*(\d+)", norm)
    if m:
        return m.group(1)
    if re.fullmatch(r"\d+", norm.strip()):
        return norm.strip()
    matches = list(re.finditer(r"(?<!\d)(\d+)(?!\d)", norm))
    if matches:
        return matches[-1].group(1)
    return None


async def _resolve_credit_folder_for_outputs(
    client: GraphApiPort,
    site_id: str,
    drive_id: str,
    client_folder_path: str,
    credit_value: str,
) -> str:
    """Resuelve el nombre real de carpeta de crédito (solo lectura / rutas de salida)."""
    encoded = encode_graph_drive_path(client_folder_path)
    endpoint = f"/sites/{site_id}/drives/{drive_id}/root:/{encoded}:/children"
    try:
        resp = await client.get(endpoint)
    except Exception:
        raise ValueError("credit_folder_not_found") from None

    folders = [
        it
        for it in resp.get("value", [])
        if isinstance(it, dict) and "folder" in it and str(it.get("name", "")).strip()
    ]
    names = [str(it["name"]).strip() for it in folders]
    if not names:
        raise ValueError("credit_folder_not_found")

    norm_credit = _normalize_for_credit_match(credit_value)
    exact = [n for n in names if _normalize_for_credit_match(n) == norm_credit]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        raise ValueError("credit_folder_ambiguous")

    credit_id = _extract_credit_id_for_match(credit_value)
    if not credit_id:
        raise ValueError("credit_folder_not_found")

    token_re = re.compile(rf"(?<!\d){re.escape(credit_id)}(?!\d)")
    token_matches = [n for n in names if token_re.search(_normalize_for_credit_match(n))]
    if len(token_matches) == 0:
        raise ValueError("credit_folder_not_found")
    if len(token_matches) > 1:
        raise ValueError("credit_folder_ambiguous")
    return token_matches[0]


def _dist_column_map(ws: Any, header_row: int) -> dict[str, int]:
    m: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(header_row, c).value or "").strip()
        if h:
            m[h] = c
    return m


def _extract_internal_path_from_root_url(url: str) -> str | None:
    decoded = unquote(str(url))
    if "root:/" not in decoded:
        return None
    inner = decoded.split("root:/", 1)[1]
    inner = inner.split(":/", 1)[0]
    path = unquote(inner).replace("\\", "/").strip().strip("/")
    return path or None


def _looks_like_pdf_internal_path(value: str) -> bool:
    s = str(value).strip()
    if not s or s.lower().startswith("http"):
        return False
    if "/" not in s:
        return False
    return s.lower().endswith(".pdf")


def _internal_pdf_path_from_ruta_column(cell_ruta: Any | None, dist: dict[str, Any]) -> str | None:
    raw: Any = None
    if cell_ruta is not None:
        v = cell_ruta.value
        if v is not None and str(v).strip():
            raw = v
    if raw is None and dist.get(DistribucionCols.RUTA) not in (None, ""):
        raw = dist.get(DistribucionCols.RUTA)
    if raw is None:
        return None
    norm = str(raw).strip().replace("\\", "/")
    if _looks_like_pdf_internal_path(norm):
        return norm
    return None


def _extract_pdf_filename_from_sharepoint_url(url: str) -> str | None:
    m = re.search(r"(?i)[?&#]file=([^&]+)", str(url))
    if not m:
        return None
    f = unquote(m.group(1).replace("+", " ")).strip().strip("/")
    if not f.lower().endswith(".pdf"):
        return None
    return f.rsplit("/", 1)[-1]


def _extract_graph_item_reference_from_url(url: str) -> tuple[str, str] | None:
    m = re.search(r"/drives/([^/]+)/items/([^/?#]+)", str(url))
    if not m:
        return None
    return unquote(m.group(1)), unquote(m.group(2))


def _parent_reference_path_to_drive_relative(path: str) -> str | None:
    if not path:
        return None
    raw = unquote(str(path)).replace("\\", "/")
    if "root:/" not in raw:
        return None
    inner = raw.split("root:/", 1)[1]
    inner = inner.split(":/", 1)[0]
    return inner.strip("/") or None


def _drive_item_to_internal_path(item: dict[str, Any]) -> str | None:
    name = str(item.get("name", "") or "").strip()
    if not name.lower().endswith(".pdf"):
        return None
    pr = item.get("parentReference") or {}
    parent = _parent_reference_path_to_drive_relative(str(pr.get("path", "") or ""))
    if not parent:
        return None
    return f"{parent}/{name}".replace("//", "/")


def _hyperlink_target(cell: Any) -> str | None:
    hl = getattr(cell, "hyperlink", None)
    if hl is None:
        return None
    t = getattr(hl, "target", None) or getattr(hl, "ref", None)
    return str(t).strip() if t else None


def _resolve_extract_route_without_network(cell: Any) -> str | None:
    v = cell.value
    if v is not None:
        s = str(v).strip()
        if _looks_like_pdf_internal_path(s):
            return s.replace("\\", "/")
    target = _hyperlink_target(cell)
    if not target:
        return None
    gp = _extract_internal_path_from_root_url(target)
    if gp and gp.lower().endswith(".pdf") and "/" in gp:
        return gp
    if not target.lower().startswith("http") and _looks_like_pdf_internal_path(target):
        return target.replace("\\", "/")
    return None


def _parse_extract_date_from_filename(filename: str) -> date | None:
    lower = filename.strip().lower()
    m = re.search(r"(\d{4}-\d{2}-\d{2})\.pdf\s*$", lower)
    if m:
        try:
            return datetime.strptime(m.group(1), "%Y-%m-%d").date()
        except ValueError:
            return None
    m = re.search(r"(\d{4}-\d{2})\.pdf\s*$", lower)
    if m:
        try:
            y, mo = map(int, m.group(1).split("-"))
            last = monthrange(y, mo)[1]
            return date(y, mo, last)
        except ValueError:
            return None
    found: list[date] = []
    for dm in re.finditer(r"\b(\d{4}-\d{2}-\d{2})\b", lower):
        try:
            found.append(datetime.strptime(dm.group(1), "%Y-%m-%d").date())
        except ValueError:
            continue
    if found:
        return found[-1]
    return None


def _coerce_fecha_banco_to_date(v: Any) -> date | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    text = str(v).strip()
    if not text:
        return None
    # Generate escribe ISO en Distribucion_Abonos; Excel puede dejarlo como texto.
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def _pick_extract_pdf_name(
    pdf_items: list[dict[str, Any]],
    extract_keyword: str,
    fecha_banco: date | None,
    exact_basename: str | None,
) -> str | None:
    pdfs = [it for it in pdf_items if str(it.get("name", "")).lower().endswith(".pdf")]
    if not pdfs:
        return None
    if exact_basename:
        eb = exact_basename.strip().lower()
        for it in pdfs:
            if str(it.get("name", "")).strip().lower() == eb:
                return str(it["name"])
    kw = extract_keyword.lower()
    pool = [it for it in pdfs if kw in str(it.get("name", "")).lower()]
    if not pool:
        return None
    if len(pool) == 1:
        return str(pool[0]["name"])
    if fecha_banco:
        scored: list[tuple[int, str]] = []
        for it in pool:
            name = str(it.get("name", ""))
            d = _parse_extract_date_from_filename(name)
            if d is None:
                scored.append((10**9, name))
            else:
                scored.append((abs((d - fecha_banco).days), name))
        scored.sort(key=lambda x: (x[0], x[1]))
        return scored[0][1]
    return str(sorted(pool, key=lambda x: str(x.get("name", "")))[0]["name"])


async def _resolve_extract_route_from_graph_item(
    client: GraphApiPort, item_drive_id: str, item_id: str
) -> str | None:
    try:
        item = await client.get(
            f"/drives/{item_drive_id}/items/{item_id}",
            params={"$select": "name,parentReference"},
        )
        if isinstance(item, dict):
            p = _drive_item_to_internal_path(item)
            return p
    except Exception:
        return None
    return None


async def _resolve_extract_route_by_credit_folder(
    client: GraphApiPort,
    site_id: str,
    drive_id: str,
    clients_path: str,
    cliente: str,
    credito: str,
    fecha_banco_raw: Any,
    extract_keyword: str,
    filename_hint: str | None,
) -> str | None:
    """
    Busca PDF de extracto; primero dentro de carpeta EXTRACTOS (si existe); si no, en raíz del crédito.
    Solo considera PDFs cuyo nombre contiene extract_keyword (vía _pick_extract_pdf_name).
    """
    if not cliente or not credito:
        return None

    credit_full: str
    try:
        real_credit = await _resolve_credit_folder_for_outputs(
            client, site_id, drive_id, f"{clients_path}/{cliente}", credito
        )
        credit_full = f"{clients_path}/{cliente}/{real_credit}"
    except Exception:
        return None

    fecha_d = _coerce_fecha_banco_to_date(fecha_banco_raw)
    hint_base = filename_hint.rsplit("/", 1)[-1] if filename_hint else None

    async def fetch_children(rel_path_full: str) -> list[dict[str, Any]]:
        enc_inner = encode_graph_drive_path(rel_path_full)
        try:
            resp = await client.get(
                f"/sites/{site_id}/drives/{drive_id}/root:/{enc_inner}:/children"
            )
            return list(resp.get("value") or [])
        except Exception:
            return []

    async def pdf_route_under(folder_full_path: str) -> str | None:
        inner = await fetch_children(folder_full_path)
        chosen = _pick_extract_pdf_name(inner, extract_keyword, fecha_d, hint_base)
        if not chosen:
            return None
        return f"{folder_full_path}/{chosen}".replace("//", "/")

    root_items = await fetch_children(credit_full)
    extractos_label = _normalize_folder_label("EXTRACTOS")
    extractos_name: str | None = None
    for it in root_items:
        if "folder" not in it:
            continue
        nm = str(it.get("name", "")).strip()
        if nm and _normalize_folder_label(nm) == extractos_label:
            extractos_name = nm
            break

    if extractos_name:
        inner_path = f"{credit_full}/{extractos_name}".replace("//", "/")
        routed = await pdf_route_under(inner_path)
        if routed:
            return routed

    return await pdf_route_under(credit_full.replace("//", "/"))


async def _resolve_extract_route_for_validar_row(
    client: GraphApiPort,
    drive_id: str,
    site_id: str,
    clients_path: str,
    extract_keyword: str,
    dist: dict[str, Any],
    cell_ext: Any,
) -> str | None:
    local = _resolve_extract_route_without_network(cell_ext)
    if local:
        return local
    target = _hyperlink_target(cell_ext)
    filename_hint: str | None = None
    if target:
        ref = _extract_graph_item_reference_from_url(target)
        if ref:
            id_drive, iid = ref
            gp = await _resolve_extract_route_from_graph_item(client, id_drive, iid)
            if gp:
                return gp
        filename_hint = _extract_pdf_filename_from_sharepoint_url(target)
    return await _resolve_extract_route_by_credit_folder(
        client,
        site_id,
        drive_id,
        clients_path,
        str(dist.get(DistribucionCols.CLIENTE, "")).strip(),
        str(dist.get(DistribucionCols.CREDITO, "")).strip(),
        dist.get(DistribucionCols.FECHA_BANCO),
        extract_keyword,
        filename_hint,
    )


def _dirname_internal_path(path: str) -> str | None:
    path = path.strip().replace("\\", "/")
    if "/" not in path:
        return None
    return path.rsplit("/", 1)[0]


def _credit_parent_dir_from_tabla_cell(cell: Any) -> str | None:
    hl = getattr(cell, "hyperlink", None)
    if hl is None:
        return None
    target = getattr(hl, "target", None)
    if not target:
        return None
    t = str(target).strip()
    p = _extract_internal_path_from_root_url(t) or (t if "/" in t and not t.lower().startswith("http") else None)
    if not p:
        return None
    return _dirname_internal_path(p)


def _normalize_folder_label(name: str) -> str:
    v = unicodedata.normalize("NFD", name.strip().lower())
    v = "".join(ch for ch in v if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", v)


def _asientos_folder_name_for_credit(credito: str, ruta_unidad_credito: str) -> str:
    """Nombre bajo la unidad de crédito: ASIENTOS CONTABLES CRED {n}."""
    credit_id = _extract_credit_id_for_match(str(credito or "").strip())
    if not credit_id:
        seg = ruta_unidad_credito.replace("\\", "/").strip("/").split("/")[-1]
        credit_id = _extract_credit_id_for_match(seg)
    if not credit_id:
        label = str(credito or "").strip()
        if not label:
            raise ValueError("credit_number_not_resolved")
        credit_id = label
    return f"{ASIENTOS_FOLDER_PER_CREDIT_PREFIX} {credit_id}"


def _folder_item_by_exact_name(children: list[dict[str, Any]], folder_name: str) -> dict[str, Any] | None:
    target = folder_name.casefold()
    for it in children:
        if "folder" not in it:
            continue
        nm = str(it.get("name", "")).strip()
        if nm.casefold() == target:
            return it
    return None


async def _list_folder_children_finalize(
    client: GraphApiPort,
    site_id: str,
    drive_id: str,
    parent_rel: str,
) -> list[dict[str, Any]]:
    enc = encode_graph_drive_path(parent_rel.strip().strip("/"))
    resp = await client.get(f"/sites/{site_id}/drives/{drive_id}/root:/{enc}:/children")
    return list(resp.get("value") or [])


async def _ensure_asientos_folder_under_credit_unit(
    client: GraphApiPort,
    site_id: str,
    drive_id: str,
    ruta_unidad_credito: str,
    folder_name: str,
) -> tuple[str, str | None]:
    """
    Crea la carpeta de asientos bajo la unidad de crédito si no existe.
    Retorna (ruta relativa al drive, webUrl opcional).
    """
    parent = ruta_unidad_credito.strip().replace("\\", "/").strip("/")
    if not parent:
        raise ValueError("missing_ruta_unidad_credito")
    rel_path = f"{parent}/{folder_name}".replace("//", "/")
    children = await _list_folder_children_finalize(client, site_id, drive_id, parent)
    existing = _folder_item_by_exact_name(children, folder_name)
    if existing is not None:
        web = existing.get("webUrl")
        return rel_path, str(web) if web else None

    enc = encode_graph_drive_path(parent)
    endpoint = f"/sites/{site_id}/drives/{drive_id}/root:/{enc}:/children"
    body: dict[str, Any] = {
        "name": folder_name,
        "folder": {},
        "@microsoft.graph.conflictBehavior": "fail",
    }
    try:
        resp, code = await client.post_json(endpoint, body)
    except httpx.HTTPStatusError as exc:
        if exc.response is not None and exc.response.status_code == 409:
            children = await _list_folder_children_finalize(client, site_id, drive_id, parent)
            existing = _folder_item_by_exact_name(children, folder_name)
            if existing is not None:
                web = existing.get("webUrl")
                return rel_path, str(web) if web else None
        raise
    else:
        if code not in (200, 201):
            raise RuntimeError(f"unexpected_graph_status_{code}")
        web = resp.get("webUrl") if isinstance(resp, dict) else None
        return rel_path, str(web) if web else None


async def _provision_asientos_folders_for_distribution_rows(
    client: GraphApiPort,
    site_id: str,
    drive_id: str,
    ws_dist: Any,
    dist_header_row: int,
    distributions: list[dict[str, Any]],
) -> dict[int, tuple[str, str]]:
    """Crea carpetas ASIENTOS CONTABLES CRED {n} y rellena RutaAsientosContables por fila."""
    colmap = _dist_column_map(ws_dist, dist_header_row)
    col_ruta_uc = colmap.get(DistribucionCols.RUTA_UNIDAD_CREDITO)
    asientos_by_row: dict[int, tuple[str, str]] = {}

    for dist in distributions:
        if not _include_in_validation_outputs(dist):
            continue
        r = int(dist["_excel_row"])
        ruta_uc = str(dist.get(DistribucionCols.RUTA_UNIDAD_CREDITO) or "").strip().replace("\\", "/")
        if not ruta_uc and col_ruta_uc:
            cell_val = ws_dist.cell(r, col_ruta_uc).value
            ruta_uc = str(cell_val or "").strip().replace("\\", "/")
        if not ruta_uc:
            logger.error(
                "finalize missing_ruta_unidad_credito: fila=%s id_pago=%r cliente=%r credito=%r",
                r,
                dist.get(DistribucionCols.ID_PAGO),
                dist.get(DistribucionCols.CLIENTE),
                dist.get(DistribucionCols.CREDITO),
            )
            raise ValueError("missing_ruta_unidad_credito")

        credito = str(dist.get(DistribucionCols.CREDITO, "")).strip()
        try:
            folder_name = _asientos_folder_name_for_credit(credito, ruta_uc)
        except ValueError:
            logger.error(
                "finalize credit_number_not_resolved: fila=%s id_pago=%r credito=%r ruta_unidad=%r",
                r,
                dist.get(DistribucionCols.ID_PAGO),
                credito,
                ruta_uc,
            )
            raise ValueError("credit_number_not_resolved") from None

        try:
            rel_path, web_url = await _ensure_asientos_folder_under_credit_unit(
                client, site_id, drive_id, ruta_uc, folder_name
            )
        except Exception as exc:
            logger.error(
                "finalize asientos_folder_create_failed: fila=%s id_pago=%r cliente=%r credito=%r "
                "ruta_unidad=%r folder_name=%r detail=%s",
                r,
                dist.get(DistribucionCols.ID_PAGO),
                dist.get(DistribucionCols.CLIENTE),
                credito,
                ruta_uc,
                folder_name,
                exc,
                exc_info=True,
            )
            raise ValueError("asientos_folder_create_failed") from exc

        dist[DistribucionCols.RUTA_ASIENTOS_CONTABLES] = rel_path
        link_txt = web_url if web_url else rel_path
        asientos_by_row[r] = (link_txt, "")

    return asientos_by_row


def _configure_hist_distrib_technical_path_columns(
    ws_distribution: Any,
    dist_header_row: int,
) -> None:
    colmap = _dist_column_map(ws_distribution, dist_header_row)
    for col_name in (
        *DISTRIBUCION_TECHNICAL_HIDDEN_COLUMNS,
        DistribucionCols.RUTA_ASIENTOS_CONTABLES,
    ):
        cidx = colmap.get(col_name)
        if not cidx:
            continue
        letter = get_column_letter(cidx)
        wd = ws_distribution.column_dimensions[letter]
        wd.hidden = True
        wd.width = min(float(wd.width or 9.0), 12.0)


async def _apply_ruta_asientos_column_on_hist_sheet(
    ws_dist: Any,
    dist_header_row: int,
    distributions: list[dict[str, Any]],
) -> None:
    colmap = _dist_column_map(ws_dist, dist_header_row)
    col_asientos = colmap.get(DistribucionCols.RUTA_ASIENTOS_CONTABLES)
    if col_asientos is None:
        col_asientos = ws_dist.max_column + 1
        ws_dist.cell(dist_header_row, col_asientos, DistribucionCols.RUTA_ASIENTOS_CONTABLES)
        colmap[DistribucionCols.RUTA_ASIENTOS_CONTABLES] = col_asientos

    for dist in distributions:
        if not _include_in_validation_outputs(dist):
            continue
        r = int(dist["_excel_row"])
        rel = str(dist.get(DistribucionCols.RUTA_ASIENTOS_CONTABLES) or "").strip()
        if not rel:
            raise ValueError("missing_ruta_asientos_contables")
        ws_dist.cell(r, col_asientos, rel)


def _resolve_distrib_policy(dist: dict[str, Any]) -> Any:
    try:
        return policy_from_row(dist)
    except ValueError:
        return resolve_application_policy(TipoAplicacionVisible.PAGO, from_bank=False)


def _resolve_abono_policy(abono: dict[str, Any]) -> Any:
    try:
        return policy_from_row(abono)
    except ValueError:
        return resolve_application_policy(TipoAplicacionVisible.ABONO_LEGACY, from_bank=False)


def _write_policy_technical_columns(
    ws: Any,
    header_row: int,
    row_idx: int,
    row_dict: dict[str, Any],
    *,
    is_abono: bool,
    policy: Any,
) -> None:
    apply_fn = apply_policy_to_abonos_row if is_abono else apply_policy_to_pagos_row
    apply_fn(row_dict, policy)
    colmap = _dist_column_map(ws, header_row)
    hidden_cols = (
        DISTRIBUCION_ABONOS_TECHNICAL_HIDDEN_COLUMNS
        if is_abono
        else DISTRIBUCION_TECHNICAL_HIDDEN_COLUMNS
    )
    for col_name in hidden_cols:
        if col_name not in row_dict:
            continue
        cidx = colmap.get(col_name)
        if cidx is None:
            cidx = ws.max_column + 1
            ws.cell(header_row, cidx, col_name)
            colmap[col_name] = cidx
        ws.cell(row_idx, cidx, row_dict.get(col_name))


async def _apply_ruta_column_on_hist_sheet(
    client: GraphApiPort,
    site_id: str,
    drive_id: str,
    clients_path: str,
    extract_keyword: str,
    ws_dist: Any,
    dist_header_row: int,
    distributions: list[dict[str, Any]],
) -> None:
    colmap = _dist_column_map(ws_dist, dist_header_row)
    col_ruta = colmap.get(DistribucionCols.RUTA)
    if col_ruta is None:
        col_ruta = ws_dist.max_column + 1
        ws_dist.cell(dist_header_row, col_ruta, DistribucionCols.RUTA)
    col_link_ext = colmap.get(DistribucionCols.LINK_EXTRACTO)
    if col_link_ext is None:
        logger.error(
            "finalize missing_extract_route: columna %r no encontrada (fila encabezados=%s)",
            DistribucionCols.LINK_EXTRACTO,
            dist_header_row,
        )
        raise ValueError("missing_extract_route")

    for dist in distributions:
        r = int(dist["_excel_row"])
        if not _include_in_validation_outputs(dist):
            continue
        policy = _resolve_distrib_policy(dist)
        if not (policy_requires_closing_extract(policy) or policy_requires_reference_extract(policy)):
            continue
        cell_ruta = ws_dist.cell(r, col_ruta)
        cell_ext = ws_dist.cell(r, col_link_ext)

        resolved: str | None = _internal_pdf_path_from_ruta_column(cell_ruta, dist)
        if not resolved:
            resolved = await _resolve_extract_route_for_validar_row(
                client,
                drive_id,
                site_id,
                clients_path,
                extract_keyword,
                dist,
                cell_ext,
            )

        if not resolved:
            logger.error(
                "finalize missing_extract_route: fila_excel=%s id_pago=%r cliente=%r credito=%r "
                "estado_linea=%r link_val=%r link_hyperlink=%r ruta_celda=%r dist_ruta=%r "
                "(se intentó fallback con subcarpeta EXTRACTOS antes que raíz de crédito)",
                r,
                dist.get(DistribucionCols.ID_PAGO),
                dist.get(DistribucionCols.CLIENTE),
                dist.get(DistribucionCols.CREDITO),
                dist.get(DistribucionCols.ESTADO_PAGO),
                cell_ext.value,
                _hyperlink_target(cell_ext),
                cell_ruta.value,
                dist.get(DistribucionCols.RUTA),
            )
            raise ValueError("missing_extract_route")
        cell_ruta.value = resolved


async def _apply_abono_reference_extract_routes_on_hist_sheet(
    client: GraphApiPort,
    site_id: str,
    drive_id: str,
    clients_path: str,
    extract_keyword: str,
    ws_abono: Any,
    abono_header_row: int,
    abono_rows: list[dict[str, Any]],
) -> None:
    colmap = _abono_column_map(ws_abono, abono_header_row)
    col_ruta = colmap.get(DistribucionAbonosCols.RUTA_EXTRACTO)
    if col_ruta is None:
        col_ruta = ws_abono.max_column + 1
        ws_abono.cell(abono_header_row, col_ruta, DistribucionAbonosCols.RUTA_EXTRACTO)
    col_link_ext = colmap.get(DistribucionAbonosCols.LINK_EXTRACTO)

    for abono in abono_rows:
        if not _include_abono_in_validation_outputs(abono):
            continue
        policy = _resolve_abono_policy(abono)
        if not policy_requires_reference_extract(policy):
            continue
        r = int(abono["_excel_row"])
        cell_ruta = ws_abono.cell(r, col_ruta)
        cell_ext = ws_abono.cell(r, col_link_ext) if col_link_ext else None

        resolved: str | None = _internal_pdf_path_from_ruta_column(cell_ruta, abono)
        if not resolved and cell_ext is not None:
            resolved = await _resolve_extract_route_for_validar_row(
                client,
                drive_id,
                site_id,
                clients_path,
                extract_keyword,
                abono,
                cell_ext,
            )
        if not resolved:
            _raise_finalize_detail(
                "missing_reference_extract_route",
                excel_row=r,
                id_pago=abono.get(DistribucionAbonosCols.ID_PAGO),
                credito=abono.get(DistribucionAbonosCols.CREDITO),
            )
        cell_ruta.value = resolved


def _raise_finalize_detail(code: str, **details: Any) -> None:
    payload = json.dumps(details, ensure_ascii=False, default=str)
    raise ValueError(f"{code}|{payload}")


def _abono_column_map(ws: Any, header_row: int) -> dict[str, int]:
    return _dist_column_map(ws, header_row)


def _include_abono_in_validation_outputs(abono: dict[str, Any]) -> bool:
    return is_validar_abono_si(abono)


def _read_abono_distributions(ws_abono: Any) -> tuple[int, list[dict[str, Any]]]:
    if ws_abono.max_row < 1:
        return 1, []
    try:
        header_row = _find_table_header_row(ws_abono, DistribucionAbonosCols.ID_PAGO)
    except ValueError:
        return 1, []
    headers: list[str] = []
    rows: list[dict[str, Any]] = []
    for r_idx, row in enumerate(ws_abono.iter_rows(values_only=True), start=1):
        if r_idx < header_row:
            continue
        if r_idx == header_row:
            headers = [str(v).strip() if v else "" for v in row]
            continue
        if not any(row):
            continue
        row_dict = dict(zip(headers, row))
        row_dict[DistribucionAbonosCols.VALIDAR_ABONO] = require_validar_abono_value(
            row_dict.get(DistribucionAbonosCols.VALIDAR_ABONO)
        )
        row_dict["_excel_row"] = r_idx
        rows.append(row_dict)
    return header_row, rows


def _coerce_abono_bank_amount(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text.startswith("="):
        return None
    try:
        return float(text)
    except ValueError:
        try:
            return float(text.replace(".", "").replace(",", "."))
        except ValueError:
            return None


def _coerce_abono_bank_date(value: Any) -> date | None:
    return _coerce_fecha_banco_to_date(value)


def _validate_selected_abono_row(abono: dict[str, Any]) -> None:
    excel_row = int(abono["_excel_row"])
    id_pago = str(abono.get(DistribucionAbonosCols.ID_PAGO) or "").strip()
    if not id_pago:
        _raise_finalize_detail("abono_group_inconsistent", excel_row=excel_row, field="ID Pago")
    for field in (
        DistribucionAbonosCols.CLIENTE,
        DistribucionAbonosCols.CREDITO,
    ):
        if not str(abono.get(field) or "").strip():
            _raise_finalize_detail(
                "abono_group_inconsistent",
                excel_row=excel_row,
                id_pago=id_pago,
                field=field,
            )
    if _coerce_abono_bank_amount(abono.get(DistribucionAbonosCols.MONTO_BANCO)) is None:
        _raise_finalize_detail("abono_missing_bank_amount", excel_row=excel_row, id_pago=id_pago)
    if _coerce_abono_bank_date(abono.get(DistribucionAbonosCols.FECHA_BANCO)) is None:
        _raise_finalize_detail("abono_missing_bank_date", excel_row=excel_row, id_pago=id_pago)
    if not str(abono.get(DistribucionAbonosCols.RUTA_UNIDAD_CREDITO) or "").strip():
        _raise_finalize_detail(
            "abono_credit_without_unit_path",
            excel_row=excel_row,
            id_pago=id_pago,
            credito=abono.get(DistribucionAbonosCols.CREDITO),
        )
    if not str(abono.get(DistribucionAbonosCols.RUTA_TABLA_AMORTIZACION) or "").strip():
        _raise_finalize_detail(
            "abono_credit_without_amortization_path",
            excel_row=excel_row,
            id_pago=id_pago,
            credito=abono.get(DistribucionAbonosCols.CREDITO),
        )
    if not str(abono.get(DistribucionAbonosCols.CREDITO_NORMALIZADO) or "").strip():
        _raise_finalize_detail(
            "abono_group_inconsistent",
            excel_row=excel_row,
            id_pago=id_pago,
            field=DistribucionAbonosCols.CREDITO_NORMALIZADO,
        )
    policy = _resolve_abono_policy(abono)
    if policy.canonical_enum != TipoAplicacion.ABONO:
        _raise_finalize_detail(
            "abono_invalid_application_type",
            excel_row=excel_row,
            id_pago=id_pago,
            value_found=abono.get(DistribucionAbonosCols.TIPO_APLICACION),
        )
    if policy_requires_reference_extract(policy):
        if _coerce_abono_bank_date(abono.get(DistribucionAbonosCols.FECHA_LIMITE)) is None:
            _raise_finalize_detail(
                "abono_mora_missing_reference_date",
                excel_row=excel_row,
                id_pago=id_pago,
                credito=abono.get(DistribucionAbonosCols.CREDITO),
            )
        has_extract = bool(
            str(abono.get(DistribucionAbonosCols.LINK_EXTRACTO) or "").strip()
            or str(abono.get(DistribucionAbonosCols.RUTA_EXTRACTO) or "").strip()
        )
        if not has_extract:
            _raise_finalize_detail(
                "abono_mora_missing_reference_extract",
                excel_row=excel_row,
                id_pago=id_pago,
                credito=abono.get(DistribucionAbonosCols.CREDITO),
            )
    abono["_policy"] = policy


def _validate_abono_groups(abono_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not abono_rows:
        return []
    groups: dict[str, list[dict[str, Any]]] = {}
    for row in abono_rows:
        id_pago = str(row.get(DistribucionAbonosCols.ID_PAGO) or "").strip()
        if not id_pago:
            continue
        groups.setdefault(id_pago, []).append(row)

    selected_rows: list[dict[str, Any]] = []
    for id_pago, members in groups.items():
        selected = [r for r in members if _include_abono_in_validation_outputs(r)]
        if not selected:
            _raise_finalize_detail("abono_without_selected_credit", id_pago=id_pago)
        seen_credits: dict[str, list[int]] = {}
        ref_cliente = str(selected[0].get(DistribucionAbonosCols.CLIENTE) or "").strip()
        ref_monto = _coerce_abono_bank_amount(selected[0].get(DistribucionAbonosCols.MONTO_BANCO))
        ref_fecha = _coerce_abono_bank_date(selected[0].get(DistribucionAbonosCols.FECHA_BANCO))
        ref_policy = _resolve_abono_policy(selected[0])
        ref_tipo_original = ref_policy.tipo_aplicacion_original
        for row in selected:
            _validate_selected_abono_row(row)
            cred_norm = str(row.get(DistribucionAbonosCols.CREDITO_NORMALIZADO) or "").strip()
            excel_row = int(row["_excel_row"])
            seen_credits.setdefault(cred_norm, []).append(excel_row)
            cliente = str(row.get(DistribucionAbonosCols.CLIENTE) or "").strip()
            monto = _coerce_abono_bank_amount(row.get(DistribucionAbonosCols.MONTO_BANCO))
            fecha = _coerce_abono_bank_date(row.get(DistribucionAbonosCols.FECHA_BANCO))
            row_policy = _resolve_abono_policy(row)
            if (
                cliente != ref_cliente
                or monto != ref_monto
                or fecha != ref_fecha
                or row_policy.tipo_aplicacion_original != ref_tipo_original
            ):
                _raise_finalize_detail(
                    "abono_group_inconsistent",
                    id_pago=id_pago,
                    excel_row=excel_row,
                )
        for cred_norm, excel_rows in seen_credits.items():
            if len(excel_rows) > 1:
                _raise_finalize_detail(
                    "abono_duplicate_selected_credit",
                    id_pago=id_pago,
                    credito=cred_norm,
                    excel_rows=excel_rows,
                )
        selected_rows.extend(selected)
    return selected_rows


async def _provision_asientos_folders_for_abono_rows(
    client: GraphApiPort,
    site_id: str,
    drive_id: str,
    ws_abono: Any,
    abono_header_row: int,
    abono_rows: list[dict[str, Any]],
) -> dict[int, tuple[str, str]]:
    colmap = _abono_column_map(ws_abono, abono_header_row)
    col_ruta_uc = colmap.get(DistribucionAbonosCols.RUTA_UNIDAD_CREDITO)
    asientos_by_row: dict[int, tuple[str, str]] = {}

    for abono in abono_rows:
        if not _include_abono_in_validation_outputs(abono):
            continue
        r = int(abono["_excel_row"])
        ruta_uc = str(abono.get(DistribucionAbonosCols.RUTA_UNIDAD_CREDITO) or "").strip().replace("\\", "/")
        if not ruta_uc and col_ruta_uc:
            cell_val = ws_abono.cell(r, col_ruta_uc).value
            ruta_uc = str(cell_val or "").strip().replace("\\", "/")
        if not ruta_uc:
            _raise_finalize_detail(
                "abono_credit_without_unit_path",
                excel_row=r,
                id_pago=abono.get(DistribucionAbonosCols.ID_PAGO),
            )
        credito = str(abono.get(DistribucionAbonosCols.CREDITO, "")).strip()
        try:
            folder_name = _asientos_folder_name_for_credit(credito, ruta_uc)
        except ValueError:
            _raise_finalize_detail(
                "abono_group_inconsistent",
                excel_row=r,
                id_pago=abono.get(DistribucionAbonosCols.ID_PAGO),
                field=DistribucionAbonosCols.CREDITO,
            )
        try:
            rel_path, web_url = await _ensure_asientos_folder_under_credit_unit(
                client, site_id, drive_id, ruta_uc, folder_name
            )
        except Exception as exc:
            logger.error(
                "finalize abono asientos_folder_create_failed: fila=%s id_pago=%r credito=%r",
                r,
                abono.get(DistribucionAbonosCols.ID_PAGO),
                credito,
                exc_info=True,
            )
            raise ValueError("asientos_folder_create_failed") from exc
        abono[DistribucionAbonosCols.RUTA_ASIENTOS_CONTABLES] = rel_path
        link_txt = web_url if web_url else rel_path
        asientos_by_row[r] = (link_txt, "")
    return asientos_by_row


def _configure_hist_abono_technical_columns(ws_abono: Any, abono_header_row: int) -> None:
    colmap = _abono_column_map(ws_abono, abono_header_row)
    for col_name in DISTRIBUCION_ABONOS_TECHNICAL_HIDDEN_COLUMNS:
        cidx = colmap.get(col_name)
        if not cidx:
            continue
        letter = get_column_letter(cidx)
        wd = ws_abono.column_dimensions[letter]
        wd.hidden = True
        wd.width = min(float(wd.width or 9.0), 12.0)


async def _apply_ruta_asientos_column_on_hist_abono_sheet(
    ws_abono: Any,
    abono_header_row: int,
    abono_rows: list[dict[str, Any]],
) -> None:
    colmap = _abono_column_map(ws_abono, abono_header_row)
    col_asientos = colmap.get(DistribucionAbonosCols.RUTA_ASIENTOS_CONTABLES)
    if col_asientos is None:
        col_asientos = ws_abono.max_column + 1
        ws_abono.cell(abono_header_row, col_asientos, DistribucionAbonosCols.RUTA_ASIENTOS_CONTABLES)
    for abono in abono_rows:
        if not _include_abono_in_validation_outputs(abono):
            continue
        r = int(abono["_excel_row"])
        rel = str(abono.get(DistribucionAbonosCols.RUTA_ASIENTOS_CONTABLES) or "").strip()
        if not rel:
            raise ValueError("missing_ruta_asientos_contables")
        ws_abono.cell(r, col_asientos, rel)


def _resolve_payment_monto_banco(
    dist: dict[str, Any],
    monto_casos: dict[str, float],
) -> float | None:
    id_pago = str(dist.get(DistribucionCols.ID_PAGO) or "").strip()
    raw = dist.get(DistribucionCols.MONTO_BANCO)
    monto = _coerce_abono_bank_amount(raw)
    if monto is None and id_pago in monto_casos:
        monto = float(monto_casos[id_pago])
    return monto


def _build_secretary_workbook(
    ws_src_dist: Any,
    dist_header_row: int,
    distributions: list[dict[str, Any]],
    asientos_by_row: dict[int, tuple[str, str]],
    *,
    bank_name: str,
    monto_casos: dict[str, float],
    ws_src_abono: Any | None = None,
    abono_header_row: int = 1,
    abono_rows: list[dict[str, Any]] | None = None,
    abono_asientos_by_row: dict[int, tuple[str, str]] | None = None,
) -> bytes:
    colmap = _dist_column_map(ws_src_dist, dist_header_row)
    abono_colmap = (
        _abono_column_map(ws_src_abono, abono_header_row) if ws_src_abono is not None else {}
    )
    abono_asientos_by_row = abono_asientos_by_row or {}
    abono_rows = abono_rows or []
    ncols = len(SECRETARY_HEADERS)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SECRETARY_SHEET
    _apply_secretary_banner(ws, ncols)
    _apply_secretary_header_style(ws, SECRETARY_HEADER_ROW, SECRETARY_HEADERS)
    hmap = {h: i + 1 for i, h in enumerate(SECRETARY_HEADERS)}

    row_idx = SECRETARY_FIRST_DATA_ROW
    monto_shown_for_id: set[str] = set()
    for dist in distributions:
        if not _include_in_validation_outputs(dist):
            continue
        r = int(dist["_excel_row"])
        total_v = float(dist.get("total_f", 0))
        cliente = dist.get(DistribucionCols.CLIENTE)
        credito = dist.get(DistribucionCols.CREDITO)
        obs_as, _obs_note = asientos_by_row.get(r, (PENDIENTE_CREAR_ASIENTOS, OBS_NO_ASIENTOS))
        id_pago = str(dist.get(DistribucionCols.ID_PAGO) or "").strip()
        raw_monto = dist.get(DistribucionCols.MONTO_BANCO)
        if _accounting_cell_filled(raw_monto):
            monto_banco = _coerce_abono_bank_amount(raw_monto)
        elif id_pago and id_pago in monto_casos and id_pago not in monto_shown_for_id:
            monto_banco = float(monto_casos[id_pago])
        else:
            monto_banco = None
        if monto_banco is not None and id_pago:
            monto_shown_for_id.add(id_pago)

        pay_policy = _resolve_distrib_policy(dist)
        ws.cell(row_idx, hmap[AsientosPendientesCols.TIPO_APLICACION], pay_policy.tipo_aplicacion_original)
        ws.cell(row_idx, hmap[AsientosPendientesCols.ID_PAGO], dist.get(DistribucionCols.ID_PAGO))
        ws.cell(row_idx, hmap[AsientosPendientesCols.BANCO], bank_name)
        ws.cell(row_idx, hmap[AsientosPendientesCols.CLIENTE], cliente)
        ws.cell(row_idx, hmap[AsientosPendientesCols.CREDITO], credito)
        ws.cell(row_idx, hmap[AsientosPendientesCols.MONTO_BANCO], monto_banco)
        ws.cell(row_idx, hmap[AsientosPendientesCols.FECHA_BANCO], dist.get(DistribucionCols.FECHA_BANCO))
        ws.cell(row_idx, hmap[AsientosPendientesCols.FECHA_LIMITE], dist.get(DistribucionCols.FECHA_LIMITE))
        ws.cell(row_idx, hmap[AsientosPendientesCols.TOTAL_VALIDADO], total_v)
        ws.cell(
            row_idx,
            hmap[AsientosPendientesCols.OBSERVACION],
            dist.get(DistribucionCols.OBSERVACION) or "",
        )

        _set_secretary_url_link_cell(
            ws.cell(row_idx, hmap[AsientosPendientesCols.LINK_CARPETA_ASIENTOS]),
            obs_as,
            "carpeta_asientos",
            credito,
            cliente,
        )

        c_le = colmap.get(DistribucionCols.LINK_EXTRACTO)
        c_lt = colmap.get(DistribucionCols.LINK_TABLA)
        if c_le:
            _set_secretary_link_cell(
                ws.cell(row_idx, hmap[AsientosPendientesCols.LINK_EXTRACTO]),
                ws_src_dist.cell(r, c_le),
                "extracto",
                credito,
                cliente,
            )
        if c_lt:
            _set_secretary_link_cell(
                ws.cell(row_idx, hmap[AsientosPendientesCols.LINK_TABLA]),
                ws_src_dist.cell(r, c_lt),
                "tabla",
                credito,
                cliente,
            )

        row_idx += 1

    for abono in abono_rows:
        if not _include_abono_in_validation_outputs(abono):
            continue
        r = int(abono["_excel_row"])
        cliente = abono.get(DistribucionAbonosCols.CLIENTE)
        credito = abono.get(DistribucionAbonosCols.CREDITO)
        obs_as, _ = abono_asientos_by_row.get(r, (PENDIENTE_CREAR_ASIENTOS, OBS_NO_ASIENTOS))
        monto_banco = _coerce_abono_bank_amount(abono.get(DistribucionAbonosCols.MONTO_BANCO))

        abono_policy = abono.get("_policy") or _resolve_abono_policy(abono)
        ws.cell(row_idx, hmap[AsientosPendientesCols.TIPO_APLICACION], abono_policy.tipo_aplicacion_original)
        ws.cell(row_idx, hmap[AsientosPendientesCols.ID_PAGO], abono.get(DistribucionAbonosCols.ID_PAGO))
        ws.cell(row_idx, hmap[AsientosPendientesCols.BANCO], bank_name)
        ws.cell(row_idx, hmap[AsientosPendientesCols.CLIENTE], cliente)
        ws.cell(row_idx, hmap[AsientosPendientesCols.CREDITO], credito)
        ws.cell(row_idx, hmap[AsientosPendientesCols.MONTO_BANCO], monto_banco)
        ws.cell(row_idx, hmap[AsientosPendientesCols.FECHA_BANCO], abono.get(DistribucionAbonosCols.FECHA_BANCO))
        if policy_requires_reference_extract(abono_policy):
            ws.cell(row_idx, hmap[AsientosPendientesCols.FECHA_LIMITE], abono.get(DistribucionAbonosCols.FECHA_LIMITE))
            ws.cell(row_idx, hmap[AsientosPendientesCols.TOTAL_VALIDADO], SUPPORT_NOT_APPLICABLE)
            obs_text = str(abono.get(DistribucionAbonosCols.OBSERVACION) or "").strip()
            if not obs_text:
                obs_text = "Abono a mora con extracto de referencia."
            ws.cell(row_idx, hmap[AsientosPendientesCols.OBSERVACION], obs_text)
        else:
            ws.cell(row_idx, hmap[AsientosPendientesCols.FECHA_LIMITE], SUPPORT_NOT_APPLICABLE)
            ws.cell(row_idx, hmap[AsientosPendientesCols.TOTAL_VALIDADO], SUPPORT_NOT_APPLICABLE)
            ws.cell(
                row_idx,
                hmap[AsientosPendientesCols.OBSERVACION],
                abono.get(DistribucionAbonosCols.OBSERVACION) or "",
            )

        _set_secretary_url_link_cell(
            ws.cell(row_idx, hmap[AsientosPendientesCols.LINK_CARPETA_ASIENTOS]),
            obs_as,
            "carpeta_asientos",
            credito,
            cliente,
        )
        c_le = abono_colmap.get(DistribucionAbonosCols.LINK_EXTRACTO)
        if policy_requires_reference_extract(abono_policy) and c_le and ws_src_abono is not None:
            _set_secretary_link_cell(
                ws.cell(row_idx, hmap[AsientosPendientesCols.LINK_EXTRACTO]),
                ws_src_abono.cell(r, c_le),
                "extracto",
                credito,
                cliente,
            )
        else:
            ws.cell(row_idx, hmap[AsientosPendientesCols.LINK_EXTRACTO], SUPPORT_NOT_APPLICABLE)

        c_lt = abono_colmap.get(DistribucionAbonosCols.LINK_TABLA)
        if c_lt and ws_src_abono is not None:
            _set_secretary_link_cell(
                ws.cell(row_idx, hmap[AsientosPendientesCols.LINK_TABLA]),
                ws_src_abono.cell(r, c_lt),
                "tabla",
                credito,
                cliente,
            )

        row_idx += 1

    last_data_row = row_idx - 1
    first_data_row = SECRETARY_FIRST_DATA_ROW

    try:
        ws.sheet_view.showGridLines = False
    except Exception:
        pass
    try:
        ws.sheet_properties.tabColor = Color(rgb=_SEC_TAB_COLOR)
    except Exception:
        pass

    if last_data_row >= first_data_row:
        _apply_secretary_body_style(ws, first_data_row, last_data_row, ncols, hmap)
        _apply_secretary_number_formats(ws, first_data_row, last_data_row, hmap)
        _apply_secretary_client_borders(
            ws, first_data_row, last_data_row, ncols, hmap["Cliente"]
        )
        trow = _apply_secretary_total_row(ws, hmap, first_data_row, last_data_row)
        max_r = trow if trow else last_data_row
        _auto_fit_secretary_columns(ws, 1, max_r, 1, ncols)
    else:
        _auto_fit_secretary_columns(ws, 1, SECRETARY_HEADER_ROW, 1, ncols)

    ws.freeze_panes = _secretary_freeze_panes_cell(hmap)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def finalize_payment_validation(
    client: GraphApiPort,
    validation_file: str = None,
    validation_file_path: str | None = None,
    process_date: date | str | None = None,
    bank_code: str | None = None,
    job_id: str | None = None,
) -> dict[str, Any]:
    site_search = os.getenv("GRAPH_SHAREPOINT_SITE_SEARCH", "").strip()
    drive_name = os.getenv("GRAPH_SHAREPOINT_DRIVE_NAME", "").strip()
    from app.application.config.payment_validation_settings import (
        BANK_CODE_BANCOLOMBIA,
        BANK_CODE_BOGOTA,
        get_payment_validation_paths,
        normalize_bank_code,
        resolve_bank_display_name,
    )

    paths = get_payment_validation_paths()
    review_path = os.getenv("GRAPH_PAYMENT_VALIDATION_REVIEW_PATH", "").strip() or paths.review
    history_path = os.getenv("GRAPH_PAYMENT_VALIDATION_HISTORY_PATH", "").strip() or paths.historical
    validation_prefix = os.getenv("GRAPH_VALIDATION_FILE_PREFIX", "").strip()
    clients_path = os.getenv("GRAPH_CLIENTS_BASE_PATH", "").strip()
    extract_keyword = os.getenv("GRAPH_EXTRACT_KEYWORD", "Extracto").strip() or "Extracto"
    effective_process_date = _normalize_process_date(process_date)

    from app.application.use_cases.payment_validation_process_control import (
        read_process_control_snapshot,
        resolve_process_control_path_for_bank,
        update_process_control_row2,
        utc_now_iso,
        validate_bank_code,
    )
    from app.application.use_cases.setup_merge_control_workbook import (
        build_payment_validation_process_key,
    )

    def _infer_bank_code_from_path(p: str | None) -> str | None:
        low = (p or "").lower()
        if "banco_bancolombia" in low:
            return BANK_CODE_BANCOLOMBIA
        if "banco_bogota" in low:
            return BANK_CODE_BOGOTA
        return None

    bank_code_source = "body" if (bank_code or "").strip() else "auto_detected"
    bank_code = (bank_code or "").strip() or None
    if bank_code:
        validate_bank_code(bank_code)

    rev_info = await resolve_sharepoint_path(client, site_search, drive_name, review_path)

    site_id = rev_info["site_id"]
    drive_id = rev_info["drive_id"]

    # Si hay override manual (validation_file / validation_file_path) y no hay bank_code,
    # intentamos inferirlo por el nombre; si no se puede, caemos en bogotá por compatibilidad.
    manual_override = bool((validation_file_path or "").strip() or (validation_file or "").strip())
    if manual_override and not bank_code:
        inferred = _infer_bank_code_from_path(validation_file_path) or _infer_bank_code_from_path(validation_file)
        bank_code = inferred or BANK_CODE_BOGOTA
        bank_code_source = "body"

    # Auto-detección si no hay bank_code y no hay override manual.
    ready_banks_detected: list[str] = []
    validation_file_source = "body" if (validation_file_path or "").strip() or (validation_file or "").strip() else "control"
    if not bank_code and not manual_override:
        candidates: list[str] = []
        for bc in (BANK_CODE_BOGOTA, BANK_CODE_BANCOLOMBIA):
            snap = await read_process_control_snapshot(client, site_id, drive_id, bank_code=bc)
            if (
                (snap.estado_proceso or "").strip() == "REVISION_CREADA"
                and snap.is_active
                and (snap.validation_file_path or "").strip()
            ):
                candidates.append(bc)
        if not candidates:
            raise ValueError("NO_READY_PROCESS")
        if len(candidates) > 1:
            raise ValueError("MULTIPLE_READY_PROCESSES|" + ",".join(candidates))
        bank_code = candidates[0]
        bank_code_source = "auto_detected"
        ready_banks_detected = list(candidates)

    if bank_code:
        validate_bank_code(bank_code)
    bank_name = resolve_bank_display_name(bank_code)
    process_control_file_path = resolve_process_control_path_for_bank(bank_code).strip().strip("/")

    snap = await read_process_control_snapshot(client, site_id, drive_id, bank_code=bank_code)
    estado_control = (snap.estado_proceso or "").strip()

    # Resolver ProcessKey desde control (o construir si falta)
    process_key = (snap.process_key or "").strip()
    if not process_key:
        process_key = build_payment_validation_process_key(bank_code, effective_process_date.isoformat())

    # Idempotencia: si ya finalizado y paths existen, reusar.
    if (
        (snap.process_key or "").strip() == process_key
        and estado_control == "FINALIZADO"
        and (snap.historical_file_path or "").strip()
        and (snap.secretary_file_path or "").strip()
    ):
        return {
            "status": "success",
            "historical_file_path": snap.historical_file_path,
            "historical_file_url": None,
            "secretary_file_path": snap.secretary_file_path,
            "secretary_file_url": None,
            "validated_rows": 0,
            "validated_payment_rows": 0,
            "validated_abono_groups": 0,
            "validated_abono_credit_rows": 0,
            "abonos_without_selection": 0,
            "support_payment_rows": 0,
            "support_abono_rows": 0,
            "payments_and_abonos_supported": True,
            "amortization_updated": False,
            "bank_cleaned": False,
            "history_file": snap.historical_file_path.rsplit("/", 1)[-1],
            "validation_file": validation_file or "",
            "validation_file_path": validation_file_path or snap.validation_file_path,
            "process_date": effective_process_date.isoformat(),
            "payment_followup_warnings": [],
            "bank_code": bank_code,
            "bank_name": bank_name,
            "bank_code_source": bank_code_source,
            "ready_banks_detected": ready_banks_detected,
            "process_key": process_key,
            "process_control_file_path": process_control_file_path,
            "process_control_updated": False,
            "process_control_estado": "FINALIZADO",
            "validation_file_source": validation_file_source,
            "already_finalized": True,
            "file_action": "reused",
            "finalize_idempotency_key": process_key,
        }

    terminal = {
        "VACIO",
        "CONSOLIDADO",
        "MERGE_PARCIAL",
        "AMORTIZACION_APLICADA",
        "ERROR_GENERATE",
        "ERROR_FINALIZE",
        "ERROR_NOTIFY",
        "ERROR_MERGE",
        "ERROR_APPLY",
    }
    if snap.is_active and estado_control and (estado_control not in terminal) and (snap.process_key or "").strip() != process_key:
        raise ValueError(f"active_process_exists|{snap.process_key}|{estado_control}")

    # Resolver archivo de revisión: body override (validation_file_path o validation_file) > control.
    if not (validation_file_path or "").strip() and not (validation_file or "").strip():
        if estado_control != "REVISION_CREADA" or not snap.is_active:
            raise ValueError("control_not_ready_for_finalize")
        if not (snap.validation_file_path or "").strip():
            raise ValueError("missing_validation_file_path")
        validation_file_path = snap.validation_file_path
        validation_file_source = "control"

    if validation_file_path:
        file_path = validation_file_path
        validation_file = validation_file or file_path.rsplit("/", 1)[-1]
    else:
        if not validation_file:
            children = await client.get(
                f"/sites/{rev_info['site_id']}/drives/{rev_info['drive_id']}/root:/{rev_info['path_encoded']}:/children"
            )
            children_items = children.get("value", [])
            validation_file = _resolve_validation_selection(children_items, validation_prefix)

        file_path = f"{review_path}/{validation_file}"

    rev_bytes = await client.get_bytes(
        _build_content_endpoint(rev_info["site_id"], rev_info["drive_id"], file_path)
    )

    wb_rev = openpyxl.load_workbook(io.BytesIO(rev_bytes), data_only=False)

    if ReviewSheets.CONTROL not in wb_rev.sheetnames:
        raise ValueError("missing_control_sheet")
    try:
        ws_dist = find_distribucion_pagos_sheet(wb_rev)
    except ValueError:
        raise ValueError("missing_distribucion_sheet") from None

    ws_ctrl = wb_rev[ReviewSheets.CONTROL]
    procesar = None
    estado_ctrl = None
    for row in ws_ctrl.iter_rows(values_only=True):
        if row and row[0] == ControlCols.ROW_PROCESAR:
            procesar = str(row[1]).strip().upper() if row[1] else ""
        if row and row[0] == ControlCols.ROW_ESTADO:
            estado_ctrl = str(row[1]).strip().upper() if row[1] else ""

    if procesar != ControlCols.VAL_PROCESAR_SI:
        raise ValueError("process_not_approved")

    if not estado_ctrl:
        raise ValueError("missing_control_state")
    if estado_ctrl != "EN_REVISION":
        raise ValueError("invalid_control_state")

    monto_casos: dict[str, float] = {}
    if ReviewSheets.CASOS_PAGO in wb_rev.sheetnames:
        ws_casos = wb_rev[ReviewSheets.CASOS_PAGO]
        casos_header_row = _find_table_header_row(ws_casos, CasosPagoCols.ID_PAGO)
        c_headers: list[str] = []
        for r_idx, row in enumerate(ws_casos.iter_rows(values_only=True), start=1):
            if r_idx < casos_header_row:
                continue
            if r_idx == casos_header_row:
                c_headers = [str(v).strip() if v else "" for v in row]
                continue
            if not any(row):
                continue
            rd = dict(zip(c_headers, row))
            monto_casos[str(rd.get(CasosPagoCols.ID_PAGO))] = float(rd.get(CasosPagoCols.MONTO_BANCO, 0) or 0)

    distributions: list[dict[str, Any]] = []

    dist_header_row = _find_table_header_row(ws_dist, DistribucionCols.ID_PAGO)
    headers: list[str] = []
    distrib_schema_version = REVIEW_SCHEMA_VERSION
    for r_idx, row in enumerate(ws_dist.iter_rows(values_only=True), start=1):
        if r_idx < dist_header_row:
            continue
        if r_idx == dist_header_row:
            headers = [str(v).strip() if v else "" for v in row]
            distrib_schema_version = detect_distrib_schema_version_from_headers(headers)
            ctrl_schema_version = read_control_review_schema_version(ws_ctrl)
            if distrib_schema_version < REVIEW_SCHEMA_VERSION or (
                ctrl_schema_version is not None and ctrl_schema_version < REVIEW_SCHEMA_VERSION
            ):
                raise ValueError("review_schema_version_1_requires_regenerate")
            continue

        if not any(row):
            continue

        row_dict = normalize_distrib_row_keys(
            dict(zip(headers, row)),
            schema_version=distrib_schema_version,
        )
        apply_legacy_estado_migration(row_dict)
        row_dict["_excel_row"] = r_idx
        row_dict["_policy"] = _resolve_distrib_policy(row_dict)
        distributions.append(row_dict)

    _validate_no_duplicate_distrib_monto_banco(distributions)

    sum_aplicado: dict[str, float] = {}

    def safe_float(v: Any) -> float:
        if v is None:
            return 0.0
        if isinstance(v, (int, float)):
            return float(v)
        text = str(v).strip()
        if not text or text.startswith("="):
            return 0.0
        try:
            return float(text)
        except ValueError:
            try:
                return float(text.replace(".", "").replace(",", "."))
            except ValueError:
                return 0.0

    for dist in distributions:
        estado = str(dist.get(DistribucionCols.ESTADO_PAGO, "")).strip().upper()
        if not estado or estado == "NONE":
            raise ValueError("empty_estado_pago")
        if estado == "INCOMPLETO":
            raise ValueError("INCOMPLETO_NOT_SUPPORTED")
        if estado not in EstadoPago.ALLOWED:
            raise ValueError("invalid_estado_pago")
        if estado in EstadoPago.FINALIZE_FORBIDDEN:
            raise ValueError("estado_pago_no_finalizable")

        id_pago = str(dist.get(DistribucionCols.ID_PAGO))
        raw_mora_aplicar = dist.get(DistribucionCols.MORA_A_APLICAR)
        raw_capital = dist.get(DistribucionCols.ABONO_A_CAPITAL)
        raw_otros = dist.get(DistribucionCols.OTROS_VALORES)
        raw_vi = dist.get(DistribucionCols.APLICAR_A_EXTRACTO)
        obs = dist.get(DistribucionCols.OBSERVACION)
        vp_si = is_validar_pago_si(dist)

        policy = dist.get("_policy") or _resolve_distrib_policy(dist)
        dist["_policy"] = policy

        if estado in EstadoPago.COUNTERS_POSITIVE_TOTAL and vp_si:
            if not _accounting_cell_filled(raw_vi):
                raise ValueError("missing_valor_intereses")
            if not _accounting_cell_filled(raw_mora_aplicar):
                raise ValueError("missing_mora_a_aplicar")
            if not _accounting_cell_filled(raw_capital):
                raise ValueError("missing_abono_capital")
            if not _accounting_cell_filled(raw_otros):
                raise ValueError("missing_otros_valores")

        mora_aplicar_f = safe_float(raw_mora_aplicar)
        capital_f = safe_float(raw_capital)
        otros_f = safe_float(raw_otros)
        int_f = safe_float(raw_vi)

        if (
            policy.subtipo_aplicacion == ApplicationSubtype.CUOTA_MAS_CAPITAL
            and vp_si
            and estado in EstadoPago.COUNTERS_POSITIVE_TOTAL
        ):
            if int_f <= 0:
                raise ValueError("pago_y_abono_capital_missing_parte_cuota")
            if capital_f <= 0:
                raise ValueError("pago_y_abono_capital_missing_capital")
            saldo_f = safe_float(dist.get(DistribucionCols.SALDO_POR_ASIGNAR))
            if abs(saldo_f) > 0.01:
                raise ValueError("pago_y_abono_capital_saldo_must_be_zero")

        total_f = int_f + mora_aplicar_f + capital_f + otros_f
        dist[DistribucionCols.TOTAL_APLICADO] = total_f

        if id_pago not in monto_casos:
            monto_casos[id_pago] = safe_float(dist.get(DistribucionCols.MONTO_BANCO))

        if estado == EstadoPago.NORMAL and not vp_si:
            if not obs or str(obs).strip() == "":
                raise ValueError("no_validar_requires_observation")
        elif estado in EstadoPago.COUNTERS_POSITIVE_TOTAL and vp_si:
            if total_f <= 0:
                raise ValueError("validar_requires_positive_total")
            sum_aplicado[id_pago] = sum_aplicado.get(id_pago, 0) + total_f

        dist["mora_aplicar_f"] = mora_aplicar_f
        dist["mora_f"] = mora_aplicar_f
        dist["int_f"] = int_f
        dist["abono_f"] = capital_f
        dist["capital_f"] = capital_f
        dist["otros_f"] = otros_f
        dist["total_f"] = total_f

    for idp, sum_ap in sum_aplicado.items():
        if idp in monto_casos:
            if abs(sum_ap - monto_casos[idp]) > 0.01:
                raise ValueError("amount_mismatch")

    validated_payment_rows = sum(1 for d in distributions if _include_in_validation_outputs(d))

    abono_rows_all: list[dict[str, Any]] = []
    abono_header_row = 1
    ws_abono_rev: Any | None = None
    if ReviewSheets.DISTRIBUCION_ABONOS in wb_rev.sheetnames:
        ws_abono_rev = wb_rev[ReviewSheets.DISTRIBUCION_ABONOS]
        abono_header_row, abono_rows_all = _read_abono_distributions(ws_abono_rev)

    selected_abono_rows = _validate_abono_groups(abono_rows_all)
    validated_abono_credit_rows = len(selected_abono_rows)
    validated_abono_groups = len(
        {str(r.get(DistribucionAbonosCols.ID_PAGO) or "").strip() for r in selected_abono_rows}
    )
    validated_rows = validated_payment_rows + validated_abono_credit_rows

    clients_info = await resolve_sharepoint_path(client, site_search, drive_name, clients_path)
    clients_drive_id = clients_info["drive_id"]
    clients_site_id = clients_info["site_id"]
    asientos_by_row = await _provision_asientos_folders_for_distribution_rows(
        client,
        clients_site_id,
        clients_drive_id,
        ws_dist,
        dist_header_row,
        distributions,
    )

    abono_asientos_by_row: dict[int, tuple[str, str]] = {}
    if selected_abono_rows and ws_abono_rev is not None:
        abono_asientos_by_row = await _provision_asientos_folders_for_abono_rows(
            client,
            clients_site_id,
            clients_drive_id,
            ws_abono_rev,
            abono_header_row,
            selected_abono_rows,
        )

    wb_hist = openpyxl.load_workbook(io.BytesIO(rev_bytes), data_only=False)
    ws_hist_dist = find_distribucion_pagos_sheet(wb_hist)
    hist_dist_header = _find_table_header_row(ws_hist_dist, DistribucionCols.ID_PAGO)
    for dist in distributions:
        if not _include_in_validation_outputs(dist):
            continue
        policy = dist.get("_policy") or _resolve_distrib_policy(dist)
        _write_policy_technical_columns(
            ws_hist_dist,
            hist_dist_header,
            int(dist["_excel_row"]),
            dist,
            is_abono=False,
            policy=policy,
        )
    await _apply_ruta_column_on_hist_sheet(
        client,
        clients_site_id,
        clients_drive_id,
        clients_path,
        extract_keyword,
        ws_hist_dist,
        hist_dist_header,
        distributions,
    )
    await _apply_ruta_asientos_column_on_hist_sheet(
        ws_hist_dist,
        hist_dist_header,
        distributions,
    )
    _configure_hist_distrib_technical_path_columns(ws_hist_dist, hist_dist_header)

    ws_hist_abono: Any | None = None
    hist_abono_header = abono_header_row
    if ReviewSheets.DISTRIBUCION_ABONOS in wb_hist.sheetnames:
        ws_hist_abono = wb_hist[ReviewSheets.DISTRIBUCION_ABONOS]
        hist_abono_header = _find_table_header_row(ws_hist_abono, DistribucionAbonosCols.ID_PAGO)
        if selected_abono_rows:
            for abono in selected_abono_rows:
                if not _include_abono_in_validation_outputs(abono):
                    continue
                policy = abono.get("_policy") or _resolve_abono_policy(abono)
                _write_policy_technical_columns(
                    ws_hist_abono,
                    hist_abono_header,
                    int(abono["_excel_row"]),
                    abono,
                    is_abono=True,
                    policy=policy,
                )
            await _apply_abono_reference_extract_routes_on_hist_sheet(
                client,
                clients_site_id,
                clients_drive_id,
                clients_path,
                extract_keyword,
                ws_hist_abono,
                hist_abono_header,
                selected_abono_rows,
            )
            await _apply_ruta_asientos_column_on_hist_abono_sheet(
                ws_hist_abono,
                hist_abono_header,
                selected_abono_rows,
            )
        _configure_hist_abono_technical_columns(ws_hist_abono, hist_abono_header)

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws_hist_ctrl = wb_hist[ReviewSheets.CONTROL]
    ws_hist_ctrl.append([ControlCols.ROW_ESTADO_PROCESO, ControlCols.VAL_PROCESADO])
    ws_hist_ctrl.append([ControlCols.ROW_FECHA_PROCESAMIENTO, now_str])
    ws_hist_ctrl.append([ControlCols.ROW_RESULTADO, ControlCols.VAL_FINALIZADO])

    out_hist = io.BytesIO()
    wb_hist.save(out_hist)
    hist_bytes = out_hist.getvalue()

    sec_bytes = _build_secretary_workbook(
        ws_hist_dist,
        hist_dist_header,
        distributions,
        asientos_by_row,
        bank_name=bank_name,
        monto_casos=monto_casos,
        ws_src_abono=ws_hist_abono,
        abono_header_row=hist_abono_header,
        abono_rows=selected_abono_rows,
        abono_asientos_by_row=abono_asientos_by_row,
    )

    hist_name = f"cartera_validada_{bank_code}_{effective_process_date.isoformat()}.xlsx"
    sec_name = f"soporte_asientos_contables_{bank_code}_{effective_process_date.isoformat()}.xlsx"

    hist_info = await resolve_sharepoint_path(client, site_search, drive_name, history_path)
    hist_full_path = f"{history_path}/{hist_name}"
    sec_full_path = f"{history_path}/{sec_name}"

    payment_followup_warnings = await register_payment_followups_after_finalize(
        client,
        hist_info["site_id"],
        hist_info["drive_id"],
        process_date=effective_process_date,
        distributions=distributions,
        historical_relative_path=hist_full_path,
    )

    historical_file_url: str | None = None
    secretary_file_url: str | None = None
    try:
        hist_resp = await client.put_bytes(
            _build_content_endpoint(hist_info["site_id"], hist_info["drive_id"], hist_full_path),
            hist_bytes,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        historical_file_url = hist_resp.get("webUrl") if isinstance(hist_resp, dict) else None
        sec_resp = await client.put_bytes(
            _build_content_endpoint(hist_info["site_id"], hist_info["drive_id"], sec_full_path),
            sec_bytes,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        secretary_file_url = sec_resp.get("webUrl") if isinstance(sec_resp, dict) else None
    except Exception as e:
        raise Exception(f"upload_failed|{hist_full_path}|{sec_full_path}|{str(e)}") from e

    # Escribir control por banco al éxito.
    now_iso = utc_now_iso()
    updates: dict[str, Any] = {
        "ProcessKey": process_key,
        "ProcessDate": effective_process_date.isoformat(),
        "BankCode": bank_code,
        "BankName": bank_name,
        "HistoricalFilePath": hist_full_path.strip().strip("/"),
        "SecretaryFilePath": sec_full_path.strip().strip("/"),
        "EstadoProceso": "FINALIZADO",
        "IsActive": True,
        "FinalizeIdempotencyKey": process_key,
        "FinalizeJobId": job_id or "",
        "LastCompletedStep": "FINALIZE",
        "LastStepStatus": "COMPLETED",
        "LastStepErrorCode": "",
        "LastUpdatedAtProceso": now_iso,
    }
    await update_process_control_row2(client, site_id, drive_id, bank_code=bank_code, updates=updates)

    support_payment_rows = validated_payment_rows
    support_abono_rows = validated_abono_credit_rows
    out: dict[str, Any] = {
        "status": "success",
        "historical_file_path": hist_full_path,
        "historical_file_url": historical_file_url,
        "secretary_file_path": sec_full_path,
        "secretary_file_url": secretary_file_url,
        "validated_rows": validated_rows,
        "validated_payment_rows": validated_payment_rows,
        "validated_abono_groups": validated_abono_groups,
        "validated_abono_credit_rows": validated_abono_credit_rows,
        "abonos_without_selection": 0,
        "support_payment_rows": support_payment_rows,
        "support_abono_rows": support_abono_rows,
        "payments_and_abonos_supported": True,
        "amortization_updated": False,
        "bank_cleaned": False,
        "history_file": hist_name,
        "validation_file": validation_file,
        "validation_file_path": file_path,
        "process_date": effective_process_date.isoformat(),
        "payment_followup_warnings": payment_followup_warnings,
        "bank_code": bank_code,
        "bank_name": bank_name,
        "bank_code_source": bank_code_source,
        "ready_banks_detected": ready_banks_detected,
        "process_key": process_key,
        "process_control_file_path": process_control_file_path,
        "process_control_updated": True,
        "process_control_estado": "FINALIZADO",
        "validation_file_source": validation_file_source,
        "already_finalized": False,
        "file_action": "created",
        "finalize_idempotency_key": process_key,
    }
    if validated_abono_groups > 0:
        out["user_message"] = (
            "La revisión de pagos y abonos quedó cerrada correctamente. "
            "Se generó el histórico y el soporte de asientos contables."
        )
        out["next_action"] = (
            "Abra el soporte y cargue un asiento contable en la carpeta de cada crédito incluido."
        )
    return out
