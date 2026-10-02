"""
Lectura/actualización del control oficial por banco (fila 2 / espejo en disco).

Única fuente operativa del flujo Generate → Apply.
En ``local`` (default): solo disco App Service; si falta el JSON se siembra una
plantilla canónica VACIO (compatible con el contrato Excel, sin Graph).
En ``dual``: lectura disco + escritura Excel y disco; bootstrap Excel→disco si falta.
En ``excel``: solo SharePoint (legacy).
"""

from __future__ import annotations

import asyncio
import io
import logging
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from app.application.config.payment_validation_settings import (
    BANK_CODE_BANCOLOMBIA,
    BANK_CODE_BOGOTA,
    get_payment_bank_config,
    normalize_bank_code,
    resolve_bank_control_file_path,
    validate_bank_code,
)
from app.application.process_control_store import (
    build_store_payload,
    load_control_payload,
    persist_control_payload,
    process_control_store_mode,
    resolve_active_environment_name,
)
from app.application.services.colombia_time import now_colombia_iso
from app.application.sharepoint_resolution import encode_graph_drive_path
from app.application.use_cases.setup_merge_control_workbook import (
    PROCESS_CONTROL_COLUMNS,
    SHEET_NAME,
    _initial_process_control_row_values,
)
from app.domain.ports.graph import GraphApiPort

logger = logging.getLogger(__name__)

__all__ = [
    "BANK_CODE_BANCOLOMBIA",
    "BANK_CODE_BOGOTA",
    "ProcessControlSnapshot",
    "ProcessControlStoreRecord",
    "normalize_bank_code",
    "validate_bank_code",
    "resolve_process_control_path_for_bank",
    "parse_process_control_row2",
    "extract_process_control_cells",
    "snapshot_from_control_cells",
    "download_process_control_bytes",
    "canonical_empty_control_cells",
    "seed_empty_control_on_disk",
    "read_process_control_record",
    "read_process_control_snapshot",
    "update_process_control_row2",
    "utc_now_iso",
]


def resolve_process_control_path_for_bank(bank_code: str) -> str:
    return resolve_bank_control_file_path(bank_code)


def utc_now_iso() -> str:
    """
    Marca de tiempo operativa (ISO con offset de Colombia).

    El nombre histórico ``utc_now_iso`` se conserva por compatibilidad de imports;
    el valor es siempre America/Bogota (``-05:00``), no UTC.
    """
    return now_colombia_iso()


def _content_endpoint(site_id: str, drive_id: str, rel_path: str) -> str:
    enc = encode_graph_drive_path(rel_path.strip().strip("/"))
    return f"/sites/{site_id}/drives/{drive_id}/root:/{enc}:/content"


def _col_index(name: str) -> int:
    return PROCESS_CONTROL_COLUMNS.index(name) + 1


def _is_active_cell(value: Any) -> bool:
    if value is True:
        return True
    if value in (False, None):
        return False
    s = str(value).strip().lower()
    # Excel en español muestra VERDADERO/FALSO; aceptar ambas convenciones.
    return s in ("true", "1", "yes", "si", "sí", "x", "verdadero")


def _timestamp_cell(value: Any) -> str:
    if value is None:
        return ""
    iso = getattr(value, "isoformat", None)
    if callable(iso):
        return str(iso())
    return str(value).strip()


@dataclass(frozen=True)
class ProcessControlSnapshot:
    control_file_path: str
    estado_proceso: str
    is_active: bool
    process_key: str
    process_id: str
    validation_file_path: str
    historical_file_path: str
    secretary_file_path: str
    email_pdf_path: str
    notify_idempotency_key: str
    merge_manifest_path: str
    merge_idempotency_key: str
    apply_idempotency_key: str
    bank_code: str
    bank_name: str
    execution_id: str = ""
    execution_log_path: str = ""
    last_amortization_attempt_json: str = ""
    last_completed_step: str = ""
    decisions_registry_path: str = ""
    last_updated_at_proceso: str = ""


@dataclass(frozen=True)
class ProcessControlStoreRecord:
    """Snapshot operativo + celdas de fila 2 (job ids y columnas extra)."""

    snapshot: ProcessControlSnapshot
    cells: dict[str, Any]
    source: str
    updated_at: str = ""


def _header_map_from_ws(ws: Any) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for c in range(1, (ws.max_column or 0) + 1):
        h = str(ws.cell(row=1, column=c).value or "").strip()
        if h and h not in header_map:
            header_map[h] = c
    return header_map


def _append_header_column(ws: Any, name: str) -> int:
    last = 0
    for c in range(1, (ws.max_column or 0) + 1):
        if str(ws.cell(row=1, column=c).value or "").strip():
            last = c
    col = last + 1
    ws.cell(row=1, column=col, value=name)
    return col


def extract_process_control_cells_from_ws(ws: Any) -> dict[str, Any]:
    header_map = _header_map_from_ws(ws)
    out: dict[str, Any] = {}
    names = list(PROCESS_CONTROL_COLUMNS)
    for name in header_map:
        if name not in names:
            names.append(name)
    for name in names:
        col = header_map.get(name)
        if col is None:
            continue
        if col > (ws.max_column or 0):
            continue
        out[name] = ws.cell(row=2, column=col).value
    return out


def extract_process_control_cells(raw: bytes) -> dict[str, Any]:
    wb = load_workbook(filename=io.BytesIO(raw), data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError("process_control_invalid_structure")
        ws = wb[SHEET_NAME]
        if ws.max_row < 2:
            raise ValueError("process_control_invalid_structure")
        return extract_process_control_cells_from_ws(ws)
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()


def snapshot_from_control_cells(
    cells: dict[str, Any], *, control_file_path: str
) -> ProcessControlSnapshot:
    def cell(name: str) -> Any:
        return cells.get(name)

    return ProcessControlSnapshot(
        control_file_path=control_file_path,
        estado_proceso=str(cell("EstadoProceso") or "").strip(),
        is_active=_is_active_cell(cell("IsActive")),
        process_key=str(cell("ProcessKey") or "").strip(),
        process_id=str(cell("ProcessId") or "").strip(),
        validation_file_path=str(cell("ValidationFilePath") or "").strip().strip("/"),
        historical_file_path=str(cell("HistoricalFilePath") or "").strip().strip("/"),
        secretary_file_path=str(cell("SecretaryFilePath") or "").strip().strip("/"),
        email_pdf_path=str(cell("EmailPdfPath") or "").strip().strip("/"),
        notify_idempotency_key=str(cell("NotifyIdempotencyKey") or "").strip(),
        merge_manifest_path=str(cell("MergeManifestPath") or "").strip().strip("/"),
        merge_idempotency_key=str(cell("MergeIdempotencyKey") or "").strip(),
        apply_idempotency_key=str(cell("ApplyIdempotencyKey") or "").strip(),
        bank_code=str(cell("BankCode") or "").strip(),
        bank_name=str(cell("BankName") or "").strip(),
        execution_id=str(cell("ExecutionId") or "").strip(),
        execution_log_path=str(cell("ExecutionLogPath") or "").strip().strip("/"),
        last_amortization_attempt_json=str(cell("LastAmortizationAttemptJson") or "").strip(),
        last_completed_step=str(cell("LastCompletedStep") or "").strip(),
        decisions_registry_path=str(cell("DecisionsRegistryPath") or "").strip().strip("/"),
        last_updated_at_proceso=_timestamp_cell(cell("LastUpdatedAtProceso")),
    )


def parse_process_control_row2(raw: bytes, *, control_file_path: str) -> ProcessControlSnapshot:
    cells = extract_process_control_cells(raw)
    return snapshot_from_control_cells(cells, control_file_path=control_file_path)


def _record_from_payload(payload: dict[str, Any], *, source: str) -> ProcessControlStoreRecord:
    cells = dict(payload.get("cells") or {})
    rel = str(payload.get("control_file_path") or "").strip()
    return ProcessControlStoreRecord(
        snapshot=snapshot_from_control_cells(cells, control_file_path=rel),
        cells=cells,
        source=source,
        updated_at=str(payload.get("updated_at") or ""),
    )


def _payload_from_cells(
    *,
    bank_code: str,
    control_file_path: str,
    cells: dict[str, Any],
) -> dict[str, Any]:
    env = resolve_active_environment_name()
    return build_store_payload(
        environment=env,
        bank_code=bank_code,
        control_file_path=control_file_path,
        cells=cells,
    )


async def _persist_store_safe(payload: dict[str, Any]) -> None:
    await asyncio.to_thread(persist_control_payload, payload)


def canonical_empty_control_cells(bank_code: str) -> dict[str, Any]:
    """Plantilla fila 2: EstadoProceso=VACIO, IsActive=false, columnas del contrato."""
    validate_bank_code(bank_code)
    cfg = get_payment_bank_config(bank_code)
    return _initial_process_control_row_values(cfg.bank_code, cfg.bank_name)


def seed_empty_control_on_disk(*, bank_code: str) -> dict[str, Any]:
    """Si no hay JSON, crea plantilla canónica en disco y la devuelve.

    No llama Graph. Idempotente: si ya existe payload, lo reutiliza.
    """
    validate_bank_code(bank_code)
    env = resolve_active_environment_name()
    existing = load_control_payload(environment=env, bank_code=bank_code)
    if existing is not None:
        return existing
    rel = resolve_process_control_path_for_bank(bank_code).strip().strip("/")
    cells = canonical_empty_control_cells(bank_code)
    payload = build_store_payload(
        environment=env,
        bank_code=bank_code,
        control_file_path=rel,
        cells=cells,
    )
    persist_control_payload(payload)
    return payload


async def download_process_control_bytes(
    graph: GraphApiPort, site_id: str, drive_id: str, *, bank_code: str
) -> tuple[str, bytes]:
    rel = resolve_process_control_path_for_bank(bank_code).strip().strip("/")
    raw = await graph.get_bytes(_content_endpoint(site_id, drive_id, rel))
    return rel, raw


async def _load_record_from_graph(
    graph: GraphApiPort, site_id: str, drive_id: str, *, bank_code: str
) -> ProcessControlStoreRecord:
    rel, raw = await download_process_control_bytes(graph, site_id, drive_id, bank_code=bank_code)
    cells = extract_process_control_cells(raw)
    snap = snapshot_from_control_cells(cells, control_file_path=rel)
    return ProcessControlStoreRecord(snapshot=snap, cells=cells, source="excel")


async def read_process_control_record(
    graph: GraphApiPort, site_id: str, drive_id: str, *, bank_code: str
) -> ProcessControlStoreRecord:
    validate_bank_code(bank_code)
    mode = process_control_store_mode()
    env = resolve_active_environment_name()
    if mode == "local":
        payload = await asyncio.to_thread(
            load_control_payload, environment=env, bank_code=bank_code
        )
        if payload is not None:
            return _record_from_payload(payload, source="disk")
        seeded = await asyncio.to_thread(seed_empty_control_on_disk, bank_code=bank_code)
        return _record_from_payload(seeded, source="seed")
    if mode == "dual":
        payload = await asyncio.to_thread(
            load_control_payload, environment=env, bank_code=bank_code
        )
        if payload is not None:
            return _record_from_payload(payload, source="disk")
        record = await _load_record_from_graph(
            graph, site_id, drive_id, bank_code=bank_code
        )
        await _persist_store_safe(
            _payload_from_cells(
                bank_code=bank_code,
                control_file_path=record.snapshot.control_file_path,
                cells=record.cells,
            )
        )
        return record
    return await _load_record_from_graph(graph, site_id, drive_id, bank_code=bank_code)


async def read_process_control_snapshot(
    graph: GraphApiPort, site_id: str, drive_id: str, *, bank_code: str
) -> ProcessControlSnapshot:
    record = await read_process_control_record(
        graph, site_id, drive_id, bank_code=bank_code
    )
    return record.snapshot


def _apply_updates_to_cells(cells: dict[str, Any], updates: dict[str, Any]) -> dict[str, Any]:
    out = dict(cells)
    for key, val in updates.items():
        if key not in PROCESS_CONTROL_COLUMNS:
            continue
        out[key] = val
    return out


async def update_process_control_row2(
    graph: GraphApiPort,
    site_id: str,
    drive_id: str,
    *,
    bank_code: str,
    updates: dict[str, Any],
) -> None:
    validate_bank_code(bank_code)
    mode = process_control_store_mode()
    rel = resolve_process_control_path_for_bank(bank_code).strip().strip("/")
    env = resolve_active_environment_name()

    if mode == "local":
        payload = await asyncio.to_thread(
            load_control_payload, environment=env, bank_code=bank_code
        )
        if payload is None:
            payload = await asyncio.to_thread(seed_empty_control_on_disk, bank_code=bank_code)
        cells = dict(payload.get("cells") or {})
        rel = str(payload.get("control_file_path") or rel)
        cells = _apply_updates_to_cells(cells, updates)
        await _persist_store_safe(
            _payload_from_cells(bank_code=bank_code, control_file_path=rel, cells=cells)
        )
        return

    raw = await graph.get_bytes(_content_endpoint(site_id, drive_id, rel))
    wb = load_workbook(filename=io.BytesIO(raw), data_only=False)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError("process_control_invalid_structure")
        ws = wb[SHEET_NAME]
        if ws.max_row < 2:
            raise ValueError("process_control_invalid_structure")

        header_map = _header_map_from_ws(ws)

        for key, val in updates.items():
            if key not in PROCESS_CONTROL_COLUMNS:
                continue
            col = header_map.get(key)
            if col is None:
                # Workbook antiguo sin la columna: se agrega al final, nunca se
                # pisa la posición canónica (puede estar ocupada por otra columna).
                col = _append_header_column(ws, key)
                header_map[key] = col
            ws.cell(row=2, column=col, value=val)

        cells = extract_process_control_cells_from_ws(ws)
        buf = io.BytesIO()
        wb.save(buf)
        out = buf.getvalue()
    finally:
        closer = getattr(wb, "close", None)
        if callable(closer):
            closer()

    await graph.put_bytes(
        _content_endpoint(site_id, drive_id, rel),
        out,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    if mode == "dual":
        await _persist_store_safe(
            _payload_from_cells(bank_code=bank_code, control_file_path=rel, cells=cells)
        )
